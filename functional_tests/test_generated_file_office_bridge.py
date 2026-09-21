# test_generated_file_office_bridge.py
"""
Functional tests for the generated-file facade's explicit Office source bridge.
Version: 0.261.126
Implemented in: 0.261.126
Refs: microsoft/simplechat#1509; renderer preparation, not runtime activation.

Reopen real binaries, validate the public prepared-deck JSON schema, and exercise
source/cancellation fences and stream ownership without any cloud or provider.
"""

import copy
import hashlib
import importlib
import io
import json
import os
import subprocess
import sys
from datetime import date
from pathlib import Path
from unittest.mock import patch

import fitz
import pytest
from docx import Document
from jsonschema import Draft202012Validator
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation

from test_generated_file_structured_renderers import CleanupFailingIterator, RecordSource, TextSource, ValueSource


ROOT = Path(__file__).resolve().parents[1]
APP_ROOT = ROOT / 'application' / 'single_app'
with patch.object(sys, 'path', [str(APP_ROOT), *sys.path]):
    exports = importlib.import_module('functions_generated_file_exports')
    adapters = importlib.import_module('functions_generated_office_adapters')
    office = importlib.import_module('functions_office_file_renderers')


FORMATS = ('xlsx', 'docx', 'pdf', 'pptx')
PROFILES = {
    'xlsx': 'tabular_workbook_v1', 'docx': 'prepared_report_v1',
    'pdf': 'prepared_report_v1', 'pptx': 'prepared_slide_deck_v1',
}


def deck_value():
    return {
        'schema_version': 'prepared_slide_deck_v1',
        'slide_count': 2,
        'title': 'Prepared results',
        'size': 'wide',
        'slides': [
            {
                'layout': 'title_and_content',
                'title': 'First prepared slide',
                'notes': 'Prepared notes.\nSecond line.',
                'shapes': [
                    {
                        'type': 'text_box', 'box': {'left': 0.5, 'top': 1.5, 'width': 5, 'height': 3},
                        'paragraphs': [
                            {'text': 'Literal <&> text', 'bold': True},
                            {'text': 'First finding', 'list_kind': 'bullet'},
                            {'text': 'Nested detail', 'list_kind': 'bullet', 'level': 1},
                            {'text': 'Citation', 'italic': True, 'url': 'https://example.test/evidence'},
                        ],
                    },
                    {
                        'type': 'table', 'box': {'left': 6, 'top': 1.5, 'width': 6, 'height': 3},
                        'rows': [['Name', 'Count'], ['First', '001']], 'header': True, 'font_size': 12,
                    },
                ],
            },
            {
                'layout': 'blank', 'title': '', 'notes': 'Last notes',
                'shapes': [{
                    'type': 'text_box', 'box': {'left': 1, 'top': 1, 'width': 8, 'height': 3},
                    'paragraphs': [{'text': 'Last prepared slide', 'list_kind': 'number'}],
                }],
            },
        ],
    }


def source_for(output_format, *, text_kind='markdown'):
    if output_format == 'xlsx':
        return RecordSource([{'id': '001', 'value': 2}], 1)
    if output_format == 'pptx':
        return ValueSource(deck_value())
    text = '# Complete\n\nPrepared **content** with λ.'
    return TextSource([text[:10], text[10:]], len(text), text_kind)


def request_for(output_format, **changes):
    options = {'columns': ('id', 'value'), 'sheet_name': 'Records'} if output_format == 'xlsx' else {}
    options.update(changes)
    return exports.GeneratedFileExportRequest(output_format, PROFILES[output_format], **options)


def render(output_format, source=None, request=None, **options):
    source = source_for(output_format) if source is None else source
    request = request_for(output_format) if request is None else request
    return exports.build_generated_file_export(
        source=source, export_request=request, max_output_bytes=options.pop('max_output_bytes', 32 * 1024 * 1024),
        **options,
    )


@pytest.fixture
def office_buffers(monkeypatch):
    original = office._RenderContext.buffer
    buffers = []

    def tracked(context):
        result = original(context)
        buffers.append(result)
        return result

    monkeypatch.setattr(office._RenderContext, 'buffer', tracked)
    return buffers


def png_bytes():
    with Image.new('RGB', (24, 16), (20, 120, 80)) as image, io.BytesIO() as stream:
        image.save(stream, format='PNG')
        return stream.getvalue()


@pytest.mark.parametrize('output_format', FORMATS)
def test_every_advertised_office_format_reopens_with_correct_stream_metadata(output_format):
    source = source_for(output_format)
    with render(output_format, source) as result:
        payload = result.file_content.read()
        result.file_content.seek(0)
        if output_format == 'xlsx':
            book = load_workbook(result.file_content, read_only=True)
            try:
                values = list(book['Records'].values)
            finally:
                book.close()
            assert values == [('id', 'value'), ('001', 2)]
            assert result.record_count == 1
        elif output_format == 'docx':
            document = Document(result.file_content)
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            assert paragraphs == ['Complete', 'Prepared content with λ.']
            assert result.character_count == source.character_count
        elif output_format == 'pdf':
            with fitz.open(stream=payload, filetype='pdf') as document:
                text = ''.join(page.get_text() for page in document)
            assert 'Complete' in text and 'Prepared content with λ.' in text
            assert result.metadata['text_coverage_verified'] is True
        else:
            deck = Presentation(result.file_content)
            last_text = deck.slides[-1].shapes[0].text
            first_notes = deck.slides[0].notes_slide.notes_text_frame.text
            assert len(deck.slides) == 2 and last_text == 'Last prepared slide'
            assert first_notes == 'Prepared notes.\nSecond line.'
            assert result.metadata['slide_count'] == 2 and result.metadata['table_count'] == 1
        assert result.output_format == output_format and result.file_extension == output_format
        assert result.profile == PROFILES[output_format]
        assert result.metadata['renderer_profile'] == 'prepared_office_v1'
        assert result.metadata['complete'] is True
        assert result.size_bytes == len(payload) and result.content_sha256 == hashlib.sha256(payload).hexdigest()
    assert result.file_content.closed
    assert source.checks >= 2


def test_catalog_and_preparation_use_the_same_strict_json_schema():
    catalog = exports.get_generated_file_export_catalog()
    catalog_by_format = {item['format_id']: item for item in catalog}
    schema = exports.get_prepared_slide_deck_schema()
    Draft202012Validator.check_schema(schema)
    restored = json.loads(json.dumps(catalog))
    assert len(restored) == 10
    assert catalog_by_format['pptx']['profiles'][0]['input_schema'] == schema
    assert catalog_by_format['xlsx']['profiles'][0]['required_options'] == ['columns', 'sheet_name']
    assert catalog_by_format['docx']['profiles'][0]['supported_options'] == ['title']
    assert catalog_by_format['pdf']['retryable_failure_codes'] == ['source_unavailable', 'render_io']
    value = deck_value()
    snapshot = copy.deepcopy(value)
    prepared = adapters.prepare_generated_slide_deck(value)
    assert value == snapshot
    assert type(prepared) is office.PreparedDeck and prepared.expected_slide_count == 2
    assert type(prepared.slides[0].shapes[0]) is office.SlideTextBox
    assert type(prepared.slides[0].shapes[1]) is office.SlideTable
    schema['properties']['schema_version']['const'] = 'invalid'
    unchanged = exports.get_prepared_slide_deck_schema()
    assert unchanged['properties']['schema_version']['const'] == 'prepared_slide_deck_v1'


def test_standard_deck_size_is_an_explicit_supported_mapping():
    value = deck_value()
    value['size'] = 'standard'
    value['slides'] = [value['slides'][1]]
    value['slide_count'] = 1
    with render('pptx', ValueSource(value)) as result:
        deck = Presentation(result.file_content)
        width, height = deck.slide_width.inches, deck.slide_height.inches
    assert width == 10 and height == 7.5 and len(deck.slides) == 1


def test_large_xlsx_keeps_complete_ordered_single_pass_records_and_scalar_types():
    count = 30_017
    source = RecordSource(
        ({
            'id': index, 'text': f'{index:06d}', 'enabled': index % 2 == 0,
            'value': index + 0.5, 'date_text': '2026-09-21',
        } for index in range(count)),
        count,
    )
    columns = ('text', 'id', 'value', 'enabled', 'date_text')
    request = request_for('xlsx', columns=columns, sheet_name='Ordered data')
    observations = []
    with render('xlsx', source, request, check=lambda: observations.append(source.reads)) as result:
        digest = hashlib.sha256()
        size = 0
        for chunk in iter(lambda: result.file_content.read(65536), b''):
            size += len(chunk)
            digest.update(chunk)
        result.file_content.seek(0)
        book = load_workbook(result.file_content, read_only=True)
        try:
            rows = book['Ordered data'].iter_rows(values_only=True)
            header = next(rows)
            first = next(rows)
            actual_count = 1
            last = first
            for last in rows:
                actual_count += 1
        finally:
            book.close()
        assert header == columns and actual_count == count
        assert first == ('000000', 0, 0.5, True, '2026-09-21')
        assert last == (f'{count - 1:06d}', count - 1, count - 0.5, True, '2026-09-21')
        assert result.record_count == count and result.metadata['record_count'] == count
        assert result.metadata['cell_count'] == (count + 1) * len(columns)
        assert size == result.size_bytes and digest.hexdigest() == result.content_sha256
    assert source.reads == count and source.iterations == 1 and source.closed
    assert observations[0] == 0 and observations[-1] == count
    assert any(0 < reads < count for reads in observations)


def test_xlsx_formula_strings_and_headers_are_literal_not_prefixed_or_executed():
    values = ['=SUM(A1:A2)', '+1+1', '-1+1', '@cmd', '#DIV/0!', '001', '', 'λ\n"<&>"']
    source = RecordSource(({'=heading': value} for value in values), len(values))
    request = request_for('xlsx', columns=('=heading',))
    with render('xlsx', source, request) as result:
        book = load_workbook(result.file_content, read_only=True, data_only=False)
        try:
            rows = list(book.active.iter_rows())
            actual = [row[0].value for row in rows]
            cell_types = [row[0].data_type for row in rows]
        finally:
            book.close()
    assert actual == ['=heading', *values]
    assert cell_types == ['s'] * (len(values) + 1)


@pytest.mark.parametrize('bad,error_type,code', [
    ({'id': [1], 'value': 2}, office.OfficeRenderError, 'unsupported_content'),
    ({'id': {'nested': True}, 'value': 2}, office.OfficeRenderError, 'unsupported_content'),
    ({'id': date(2026, 9, 21), 'value': 2}, exports.GeneratedFileExportError, 'invalid_data'),
    ({'id': float('nan'), 'value': 2}, exports.GeneratedFileExportError, 'invalid_data'),
])
def test_xlsx_does_not_flatten_or_guess_date_conversions(bad, error_type, code, office_buffers):
    source = RecordSource([bad], 1)
    with pytest.raises(error_type) as failure:
        render('xlsx', source)
    assert failure.value.code == code and failure.value.retryable is False
    assert source.closed and all(buffer.closed for buffer in office_buffers)


@pytest.mark.parametrize('value', [9007199254740993, 1.2345678901234567, -0.0])
def test_xlsx_rejects_numbers_that_excel_cannot_represent_exactly(value, office_buffers):
    source = RecordSource([{'id': value, 'value': 2}], 1)
    with pytest.raises(office.OfficeRenderError) as failure:
        render('xlsx', source)
    assert failure.value.code == 'unsupported_content' and failure.value.retryable is False
    assert source.closed and all(buffer.closed for buffer in office_buffers)


@pytest.mark.parametrize('record', [{'id': 1}, {'id': 1, 'value': 2, 'extra': 3}])
def test_workbook_rows_cannot_drop_or_invent_declared_fields(record, office_buffers):
    source = RecordSource([record], 1)
    with pytest.raises(exports.GeneratedFileExportError) as failure:
        render('xlsx', source)
    assert failure.value.code == 'invalid_data' and source.closed
    assert all(buffer.closed for buffer in office_buffers)


@pytest.mark.parametrize('output_format', ['docx', 'pdf'])
@pytest.mark.parametrize('kind', ['text', 'markdown'])
def test_report_kind_controls_literal_text_versus_markdown(output_format, kind):
    content = '# Heading\n\n**Bold** and a literal λ.'
    source = TextSource([content], len(content), kind)
    request = request_for(output_format, title='Declared report title')
    with render(output_format, source, request) as result:
        if output_format == 'docx':
            document = Document(result.file_content)
            text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
            title = document.core_properties.title
        else:
            payload = result.file_content.read()
            with fitz.open(stream=payload, filetype='pdf') as document:
                text = ''.join(page.get_text() for page in document)
                title = document.metadata['title']
    assert title == 'Declared report title'
    assert ('# Heading' in text and '**Bold**' in text) if kind == 'text' else (
        '# Heading' not in text and '**Bold**' not in text and 'Bold' in text
    )
    assert source.reads == 1 and source.closed


@pytest.mark.parametrize('output_format', ['xlsx', 'docx', 'pdf'])
def test_empty_complete_workbook_or_report_is_valid(output_format):
    source = RecordSource([], 0) if output_format == 'xlsx' else TextSource([], 0, 'text')
    with render(output_format, source) as result:
        assert result.size_bytes > 0
        assert result.record_count == 0
        assert result.metadata['complete'] is True


@pytest.mark.parametrize('change', [
    {'unexpected': True},
    {'schema_version': 'outline-v1'},
    {'slide_count': True},
    {'slide_count': 3},
    {'slide_count': 0, 'slides': []},
    {'size': 'portrait'},
    {'images': {'asset:figure': 'filesystem-path'}},
])
def test_deck_root_schema_and_declared_count_are_strict(change, office_buffers):
    value = deck_value()
    value.update(change)
    with pytest.raises(exports.GeneratedFileExportError):
        render('pptx', ValueSource(value))
    assert not office_buffers


@pytest.mark.parametrize('change', [
    {'layout': 'comparison'},
    {'layout': 'blank'},
    {'title': ''},
    {'unknown': 'not ignored'},
    {'shapes': [{'type': 'chart', 'data': [1, 2]}]},
])
def test_deck_layouts_and_unknown_shapes_are_not_inferred(change, office_buffers):
    value = deck_value()
    value['slides'][0].update(change)
    with pytest.raises(exports.GeneratedFileExportError):
        render('pptx', ValueSource(value))
    assert not office_buffers


@pytest.mark.parametrize('bad_shape', [
    {'type': 'text_box', 'box': {'left': 1, 'top': 2, 'width': 2, 'height': 2}, 'paragraphs': ['prose']},
    {'type': 'text_box', 'box': {'left': 1, 'top': 2, 'width': 2, 'height': 2},
     'paragraphs': [{'text': 'value', 'level': 1}]},
    {'type': 'table', 'box': {'left': 1, 'top': 2, 'width': 2, 'height': 2}, 'rows': [['a', 'b'], ['c']]},
    {'type': 'table', 'box': {'left': 1, 'top': 2, 'width': 2, 'height': 2}, 'rows': [[1]]},
    {'type': 'image', 'box': {'left': 1, 'top': 2, 'width': 2, 'height': 2}, 'source': 'https://example.test/x.png'},
    {'type': 'image', 'box': {'left': 1, 'top': 2, 'width': 2, 'height': 2}, 'source': r'C:\image.png'},
    {'type': 'text_box', 'box': {'left': True, 'top': 2, 'width': 2, 'height': 2},
     'paragraphs': [{'text': 'value'}]},
])
def test_deck_shapes_have_exact_types_and_no_arbitrary_handles(bad_shape):
    value = deck_value()
    value['slides'][0]['shapes'] = [bad_shape]
    with pytest.raises(exports.GeneratedFileExportError):
        adapters.prepare_generated_slide_deck(value)


@pytest.mark.parametrize('case', ['outside', 'overlap', 'title_overlap'])
def test_deck_geometry_is_validated_before_persistence(case):
    value = deck_value()
    shape = value['slides'][0]['shapes'][0]
    if case == 'outside':
        shape['box']['left'] = 20
    elif case == 'overlap':
        value['slides'][0]['shapes'].append(copy.deepcopy(shape))
    else:
        shape['box']['top'] = 0.5
    with pytest.raises(office.OfficeRenderError) as failure:
        adapters.prepare_generated_slide_deck(value)
    assert failure.value.code == 'layout_overflow' and failure.value.retryable is False


@pytest.mark.parametrize('value', ['Write a beautiful presentation', {'slides': ['First', 'Second']}])
def test_arbitrary_prose_and_legacy_outlines_are_not_composed(value):
    with pytest.raises(exports.GeneratedFileExportError):
        render('pptx', ValueSource(value))


@pytest.mark.parametrize('output_format', ['xlsx', 'docx', 'pdf', 'pptx', 'json', 'csv', 'md', 'txt', 'yaml', 'xml'])
def test_profiles_reject_options_they_do_not_own(output_format):
    if output_format in FORMATS:
        request = request_for(
            output_format, **({'title': 'not a workbook option'} if output_format == 'xlsx' else {'sheet_name': 'Data'}),
        )
        source = source_for(output_format)
    else:
        if output_format in ('md', 'txt'):
            source = TextSource([], 0, 'markdown' if output_format == 'md' else 'text')
            profile = 'prepared_text_v1'
        else:
            source = RecordSource([], 0)
            profile = {
                'json': 'exact_records_v1', 'yaml': 'structured_records_v1',
                'csv': 'tabular_records_v1', 'xml': 'typed_xml_v1',
            }[output_format]
        request = exports.GeneratedFileExportRequest(
            output_format, profile, columns=('value',) if output_format == 'csv' else None,
            title='not a structured option',
        )
    with pytest.raises(exports.GeneratedFileExportError) as failure:
        exports.build_generated_file_export(source=source, export_request=request, max_output_bytes=1024)
    assert failure.value.code == 'invalid_options'
    assert source.reads == source.checks == 0


@pytest.mark.parametrize('name', ['XLSX', '.docx', 'word', 'powerpoint', 'excel', 'PPTX'])
def test_office_format_aliases_are_not_guessed(name):
    source = source_for('xlsx')
    request = exports.GeneratedFileExportRequest(name, 'tabular_workbook_v1', ('id', 'value'), sheet_name='Data')
    with pytest.raises(exports.GeneratedFileExportError) as failure:
        render('xlsx', source, request)
    assert failure.value.code == 'unsupported_format' and source.reads == source.checks == 0


@pytest.mark.parametrize('export_request', [
    exports.GeneratedFileExportRequest('xlsx', 'tabular_workbook_v1', ('id', 'value')),
    exports.GeneratedFileExportRequest('xlsx', 'tabular_workbook_v1', sheet_name='Data'),
    exports.GeneratedFileExportRequest('xlsx', 'tabular_workbook_v1', ('id', 'value'), sheet_name='bad/name'),
    exports.GeneratedFileExportRequest('docx', 'prepared_report_v1', title=123),
])
def test_required_options_and_option_types_fail_before_source_read(export_request):
    source = source_for(export_request.output_format)
    with pytest.raises(exports.GeneratedFileExportError) as failure:
        render(export_request.output_format, source, export_request)
    assert failure.value.code == 'invalid_options' and source.reads == source.checks == 0


def test_server_office_arguments_are_rejected_by_other_paths():
    source = RecordSource([], 0)
    with pytest.raises(exports.GeneratedFileExportError) as explicit_failure:
        exports.build_generated_file_export(
            source=source, export_request=exports.GeneratedFileExportRequest('json'),
            max_output_bytes=1024, office_limits=office.OfficeRenderLimits(),
        )
    with pytest.raises(exports.GeneratedFileExportError) as legacy_failure:
        exports.build_generated_file_export('Create a PDF', 'Prepared content', image_resolver=lambda ref: b'')
    with pytest.raises(exports.GeneratedFileExportError) as workbook_failure:
        render('xlsx', image_resolver=lambda ref: b'')
    with pytest.raises(exports.GeneratedFileExportError) as resolver_failure:
        render('docx', image_resolver='not a server callback')
    assert all(failure.value.code == 'invalid_options' for failure in (
        explicit_failure, legacy_failure, workbook_failure, resolver_failure,
    ))
    assert source.reads == source.checks == 0


@pytest.mark.parametrize('output_format', FORMATS)
@pytest.mark.parametrize('readiness', [
    exports.GeneratedFileExportReadiness(state='pending'),
    exports.GeneratedFileExportReadiness(is_complete=False),
    exports.GeneratedFileExportReadiness(is_preview=True),
])
def test_office_does_not_promote_incomplete_sources(output_format, readiness, office_buffers):
    source = source_for(output_format)
    source.readiness = readiness
    with pytest.raises(exports.GeneratedFileExportError) as failure:
        render(output_format, source)
    assert failure.value.code == 'incomplete_source' and source.reads == 0
    assert not office_buffers


@pytest.mark.parametrize('output_format', ['xlsx', 'docx', 'pdf'])
@pytest.mark.parametrize('delta', [-1, 1])
def test_count_mismatches_do_not_return_partial_binaries(output_format, delta, office_buffers):
    source = source_for(output_format)
    if output_format == 'xlsx':
        source.record_count += delta
    else:
        source.character_count += delta
    with pytest.raises(exports.GeneratedFileExportError) as failure:
        render(output_format, source)
    assert failure.value.code == 'count_mismatch' and source.closed
    assert all(buffer.closed for buffer in office_buffers)


@pytest.mark.parametrize('output_format', FORMATS)
def test_output_quota_and_office_error_classification_are_preserved(output_format, office_buffers):
    with pytest.raises(office.OfficeRenderError) as failure:
        render(output_format, max_output_bytes=100)
    assert failure.value.code == 'limit_exceeded' and failure.value.retryable is False
    assert office_buffers and all(buffer.closed for buffer in office_buffers)


@pytest.mark.parametrize('output_format', ['docx', 'pdf'])
def test_report_accumulation_enforces_exact_utf8_input_budget(output_format, monkeypatch):
    source = TextSource(['λ' * 100], 100)
    buffers = []
    original = adapters._OutputWriter

    class TrackedWriter(original):
        def __init__(self, stream, *args, **kwargs):
            buffers.append(stream)
            super().__init__(stream, *args, **kwargs)

    monkeypatch.setattr(adapters, '_OutputWriter', TrackedWriter)
    with pytest.raises(exports.GeneratedFileExportError) as failure:
        render(output_format, source, office_limits=office.OfficeRenderLimits(max_input_bytes=150))
    assert failure.value.code == 'value_limit'
    assert source.closed and buffers and all(buffer.closed for buffer in buffers)


@pytest.mark.parametrize('output_format', FORMATS)
@pytest.mark.parametrize('exception_type', [PermissionError, InterruptedError, RuntimeError])
def test_initial_rechecks_preserve_caller_exception_identity(output_format, exception_type, office_buffers):
    source = source_for(output_format)
    reason = exception_type('caller-owned failure')

    def fail():
        raise reason

    source.on_recheck = fail
    with pytest.raises(exception_type) as failure:
        render(output_format, source)
    assert failure.value is reason and source.reads == 0 and not office_buffers


@pytest.mark.parametrize('exception_type', [PermissionError, InterruptedError, RuntimeError, TimeoutError])
def test_mid_read_xlsx_failures_are_not_reclassified_as_retryable(exception_type, office_buffers):
    source = RecordSource(({'id': str(index), 'value': index} for index in range(500)), 500)
    reason = exception_type('caller-owned source/cancellation failure')

    def fail():
        if source.reads >= 100:
            raise reason

    source.on_recheck = fail
    with pytest.raises(exception_type) as failure:
        render('xlsx', source)
    assert failure.value is reason and source.reads == 100 and source.closed
    assert office_buffers and all(buffer.closed for buffer in office_buffers)


@pytest.mark.parametrize('exception_type', [PermissionError, InterruptedError, TimeoutError])
def test_reader_iterator_failures_preserve_identity_and_original_cause(exception_type, office_buffers):
    reason = exception_type('source iterator failed')
    cause = ValueError('original caller cause')

    def records():
        yield {'id': '001', 'value': 1}
        raise reason from cause

    source = RecordSource(records(), 2)
    with pytest.raises(exception_type) as failure:
        render('xlsx', source)
    assert failure.value is reason and failure.value.__cause__ is cause
    assert source.closed and all(buffer.closed for buffer in office_buffers)


@pytest.mark.parametrize('output_format', ['xlsx', 'docx', 'pdf'])
@pytest.mark.parametrize('exception_type', [PermissionError, InterruptedError])
def test_office_source_errors_survive_an_additional_iterator_cleanup_failure(
    output_format, exception_type, office_buffers,
):
    primary = exception_type('primary access or cancellation failure')
    cause = ValueError('original cause')
    items = [{'id': '001', 'value': 1}] if output_format == 'xlsx' else ['x']
    iterator = CleanupFailingIterator(items, primary=primary, cause=cause)
    if output_format == 'xlsx':
        source = RecordSource([], 2)
        source.iter_records = lambda: iterator
    else:
        source = TextSource([], 2)
        source.iter_text = lambda: iterator
    with pytest.raises(exception_type) as failure:
        render(output_format, source)
    assert failure.value is primary and failure.value.__cause__ is cause
    assert iterator.close_calls == 1
    assert all(buffer.closed for buffer in office_buffers)
    assert any('cleanup' in note for note in getattr(primary, '__notes__', []))


@pytest.mark.parametrize('output_format', ['docx', 'pdf'])
def test_mid_read_text_cancellation_closes_accumulator(output_format, monkeypatch):
    source = TextSource(('a' for _ in range(500)), 500)
    reason = InterruptedError('cancelled during text read')
    buffers = []
    original = adapters._OutputWriter

    class TrackedWriter(original):
        def __init__(self, stream, *args, **kwargs):
            buffers.append(stream)
            super().__init__(stream, *args, **kwargs)

    def cancel():
        if source.reads >= 100:
            raise reason

    monkeypatch.setattr(adapters, '_OutputWriter', TrackedWriter)
    with pytest.raises(InterruptedError) as failure:
        render(output_format, source, check=cancel)
    assert failure.value is reason and source.closed
    assert buffers and all(buffer.closed for buffer in buffers)


@pytest.mark.parametrize('output_format', FORMATS)
@pytest.mark.parametrize('exception_type', [PermissionError, InterruptedError])
def test_mid_render_rechecks_preserve_identity_and_helper_cleanup(output_format, exception_type, monkeypatch):
    source = source_for(output_format)
    reason = exception_type('revoked while the binary renderer owns buffers')
    buffers = []
    original = office._RenderContext.buffer

    def tracked(context):
        stream = original(context)
        buffers.append(stream)
        return stream

    def fail():
        if buffers:
            raise reason

    monkeypatch.setattr(office._RenderContext, 'buffer', tracked)
    with pytest.raises(exception_type) as failure:
        render(output_format, source, check=fail)
    assert failure.value is reason
    assert buffers and all(buffer.closed for buffer in buffers)


def test_readiness_changes_during_binary_rendering_keep_the_shared_failure_code(office_buffers):
    source = source_for('docx')

    def become_partial():
        if office_buffers:
            source.readiness = exports.GeneratedFileExportReadiness(is_complete=False)

    source.on_recheck = become_partial
    with pytest.raises(exports.GeneratedFileExportError) as failure:
        render('docx', source)
    assert failure.value.code == 'incomplete_source'
    assert office_buffers and all(buffer.closed for buffer in office_buffers)


@pytest.mark.parametrize('output_format', FORMATS)
@pytest.mark.parametrize('exception_type', [PermissionError, InterruptedError])
def test_final_adapter_recheck_closes_the_transferred_candidate(output_format, exception_type, monkeypatch):
    original = getattr(office, f'render_prepared_{output_format}')
    candidates = []
    reason = exception_type('revoked after the Office helper returned')

    def tracked(*args, **kwargs):
        result = original(*args, **kwargs)
        candidates.append(result)
        return result

    def fail():
        if candidates:
            raise reason

    monkeypatch.setattr(office, f'render_prepared_{output_format}', tracked)
    with pytest.raises(exception_type) as failure:
        render(output_format, check=fail)
    assert failure.value is reason
    assert candidates and all(item.file_content.closed for item in candidates)


def test_success_transfers_one_stream_and_wrapper_failure_closes_it(monkeypatch):
    original = office.render_prepared_docx
    candidates = []

    def tracked(*args, **kwargs):
        output = original(*args, **kwargs)
        candidates.append(output)
        return output

    monkeypatch.setattr(office, 'render_prepared_docx', tracked)
    with render('docx') as result:
        assert result.file_content is candidates[0].file_content
        assert not result.file_content.closed
    assert candidates[0].file_content.closed

    def fail_wrapper(**kwargs):
        raise RuntimeError('adapter metadata construction failed')

    monkeypatch.setattr(adapters, 'GeneratedFileExportStream', fail_wrapper)
    with pytest.raises(RuntimeError, match='metadata construction'):
        render('docx')
    assert len(candidates) == 2 and candidates[1].file_content.closed


@pytest.mark.parametrize('code', ['source_unavailable', 'render_io', 'layout_overflow'])
def test_native_office_errors_keep_their_original_retry_policy(code, monkeypatch):
    reason = office.OfficeRenderError(code)

    def fail(*args, **kwargs):
        raise reason

    monkeypatch.setattr(office, 'render_prepared_docx', fail)
    with pytest.raises(office.OfficeRenderError) as failure:
        render('docx')
    assert failure.value is reason
    assert failure.value.retryable is (code in ('source_unavailable', 'render_io'))


@pytest.mark.parametrize('output_format', ['docx', 'pdf', 'pptx'])
def test_images_use_only_explicit_server_resolution(output_format):
    source = source_for(output_format)
    if output_format == 'pptx':
        source.value['slides'][0]['shapes'].append({
            'type': 'image', 'box': {'left': 0.5, 'top': 5, 'width': 2, 'height': 1.5},
            'source': 'asset:figure', 'alt': 'Prepared visual',
        })
    else:
        content = '![Prepared visual](asset:figure)'
        source = TextSource([content], len(content), 'markdown')
    calls = []
    content = png_bytes()

    def resolve(reference):
        calls.append(reference)
        return content

    with render(output_format, source, image_resolver=resolve) as result:
        if output_format == 'docx':
            document = Document(result.file_content)
            count = len(document.inline_shapes)
        elif output_format == 'pdf':
            payload = result.file_content.read()
            with fitz.open(stream=payload, filetype='pdf') as document:
                count = sum(len(page.get_images()) for page in document)
        else:
            deck = Presentation(result.file_content)
            count = sum(1 for slide in deck.slides for shape in slide.shapes if shape.shape_type == 13)
        assert result.metadata['image_count'] == count == 1
    assert calls == ['asset:figure']


@pytest.mark.parametrize('reference', ['https://example.test/picture.png', r'C:\private.png', 'file:///private.png'])
def test_report_image_handles_never_reach_the_server_resolver(reference):
    content = f'![image]({reference})'
    source = TextSource([content], len(content), 'markdown')
    calls = []

    def resolve(value):
        calls.append(value)
        return png_bytes()

    with pytest.raises(office.OfficeRenderError):
        render('docx', source, image_resolver=resolve)
    assert not calls


def test_an_unresolved_image_fails_instead_of_fetching_or_becoming_a_placeholder():
    content = '![image](asset:missing)'
    with pytest.raises(office.OfficeRenderError) as failure:
        render('pdf', TextSource([content], len(content), 'markdown'))
    assert failure.value.code == 'unsupported_content'


@pytest.mark.parametrize('exception_type', [PermissionError, InterruptedError])
def test_revocation_after_image_resolution_cannot_return_an_artifact(exception_type, office_buffers):
    content = '![image](asset:figure)'
    source = TextSource([content], len(content), 'markdown')
    reason = exception_type('image source revoked')
    resolved = []
    image = png_bytes()

    def resolve(reference):
        resolved.append(reference)
        return image

    def revoke():
        if resolved:
            raise reason

    source.on_recheck = revoke
    with pytest.raises(exception_type) as failure:
        render('docx', source, image_resolver=resolve)
    assert failure.value is reason and resolved == ['asset:figure']
    assert all(buffer.closed for buffer in office_buffers)


@pytest.mark.parametrize('optimized', [False, True])
def test_cold_office_dispatch_has_no_network_bootstrap_model_or_publication_path(optimized):
    script = f'''
import importlib.abc
import sys
sys.path.insert(0, {str(APP_ROOT)!r})
forbidden = {{
    "config", "flask", "azure", "openai", "semantic_kernel", "requests",
    "functions_settings", "functions_simplechat_operations", "functions_artifact_publication",
    "route_backend_conversation_export", "functions_orchestration_results",
}}
class DenyOwners(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname.split(".")[0] in forbidden or fullname.startswith("route_"):
            raise RuntimeError("Forbidden dependency: " + fullname)
sys.meta_path.insert(0, DenyOwners())
def audit(event, args):
    if event in {{"socket.connect", "socket.getaddrinfo", "urllib.Request"}}:
        raise RuntimeError("Network I/O is forbidden")
sys.addaudithook(audit)
import functions_generated_file_exports as e
class Source:
    readiness = e.GeneratedFileExportReadiness()
    record_count = 1
    character_count = 8
    def __init__(self, kind):
        self.kind = kind
    def recheck(self):
        return None
    def iter_records(self):
        yield {{"id": "001"}}
    def iter_text(self):
        yield "Complete"
    def read_value(self):
        return {{
            "schema_version": "prepared_slide_deck_v1", "slide_count": 1,
            "slides": [{{"layout": "title_and_content", "title": "Complete", "shapes": [], "notes": "Complete notes"}}],
        }}
    def preview(self):
        raise RuntimeError("A preview is never an export source")
for kind, request in (
    ("records", e.GeneratedFileExportRequest("xlsx", "tabular_workbook_v1", ("id",), sheet_name="Data")),
    ("text", e.GeneratedFileExportRequest("docx", "prepared_report_v1")),
    ("markdown", e.GeneratedFileExportRequest("pdf", "prepared_report_v1")),
    ("structured_value", e.GeneratedFileExportRequest("pptx", "prepared_slide_deck_v1")),
):
    with e.build_generated_file_export(source=Source(kind), export_request=request, max_output_bytes=1048576) as output:
        if output.size_bytes <= 0 or output.metadata["complete"] is not True:
            raise RuntimeError("No complete binary output")
        payload = output.file_content.read()
        if len(payload) != output.size_bytes:
            raise RuntimeError("Wrong byte count")
    if not output.file_content.closed:
        raise RuntimeError("Leaked output stream")
if forbidden.intersection(sys.modules):
    raise RuntimeError("Imported a forbidden owner")
'''
    command = [sys.executable] + (['-O'] if optimized else []) + ['-c', script]
    outcome = subprocess.run(
        command, cwd=ROOT, capture_output=True, text=True, encoding='utf-8', timeout=60,
        env={**os.environ, 'PYTHONUTF8': '1', 'PYTHONDONTWRITEBYTECODE': '1'}, check=False,
    )
    assert outcome.returncode == 0, outcome.stdout + outcome.stderr


if __name__ == '__main__':
    raise SystemExit(pytest.main([__file__, '-q']))
