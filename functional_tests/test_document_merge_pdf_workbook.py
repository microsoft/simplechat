#!/usr/bin/env python3
# test_document_merge_pdf_workbook.py
"""
Functional test for PDF and workbook assembly in V2 file merge.
Version: 0.261.221
Implemented in: 0.261.221

This test ensures that functions_document_merge assembles PDFs in order with one
bookmark per source (keeping each source's own bookmarks beneath it), honors page
ranges, refuses encrypted or damaged PDFs, removes scripts and actions that could run
code, open programs or submit forms while keeping links, and builds workbooks with one
sheet per source that keep Excel cell types and number formats, copy CSV values as text,
write formula-like text literally, remove control characters Excel can't store, and name
sheets uniquely within Excel's rules. The same files always give the same bytes, library
failures never quote file content, and every limit and cancellation fails closed.
"""

import io
import os
import sys
import time
import zipfile
from datetime import datetime

import pytest

sys.path.append(os.path.dirname(os.path.abspath(__file__)))
sys.path.append(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "application", "single_app"))

from test_support.versioning import assert_app_version_at_least

from functions_document_merge import (
    DocumentMergeCancelled,
    DocumentMergeError,
    DocumentMergeLimits,
    DocumentMergeOptions,
    DocumentMergePart,
    merge_documents,
)


def pdf_bytes(label, pages, *, toc=True, encrypt=False):
    import fitz

    document = fitz.open()
    for number in range(pages):
        document.new_page().insert_text((72, 72), f"{label} page {number + 1}")
    if toc:
        document.set_toc([[1, f"{label} intro", 1]])
    options = {}
    if encrypt:
        options = {"encryption": fitz.PDF_ENCRYPT_AES_256, "owner_pw": "owner", "user_pw": ""}
    data = document.tobytes(**options)
    document.close()
    return data


def pdf_with_script(label):
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(pdf_bytes(label, 1))))
    writer.add_js("app.alert('hello');")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def pdf_with_active_content():
    """A one-page PDF with scripts on its page, links, a form field and an XFA form."""
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import ArrayObject, DictionaryObject, NameObject, NumberObject, TextStringObject

    def name(value):
        return NameObject(value)

    def script(code):
        return DictionaryObject({name("/S"): name("/JavaScript"), name("/JS"): TextStringObject(code)})

    def web(target, **extra):
        return DictionaryObject({name("/S"): name("/URI"), name("/URI"): TextStringObject(target), **extra})

    def box():
        return ArrayObject([NumberObject(0), NumberObject(0), NumberObject(50), NumberObject(50)])

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(pdf_bytes("S", 1))))
    page = writer.pages[0]
    page[name("/AA")] = DictionaryObject({name("/O"): script("app.alert('page')")})

    def link(action):
        return writer._add_object(DictionaryObject({
            name("/Type"): name("/Annot"), name("/Subtype"): name("/Link"), name("/Rect"): box(), name("/A"): action,
        }))

    field = DictionaryObject({
        name("/FT"): name("/Tx"), name("/T"): TextStringObject("total"),
        name("/AA"): DictionaryObject({name("/C"): script("calculate()")}),
    })
    field_reference = writer._add_object(field)
    widget_reference = writer._add_object(DictionaryObject({
        name("/Type"): name("/Annot"), name("/Subtype"): name("/Widget"), name("/Rect"): box(),
        name("/Parent"): field_reference, name("/AA"): DictionaryObject({name("/K"): script("keystroke()")}),
    }))
    field[name("/Kids")] = ArrayObject([widget_reference])
    page[name("/Annots")] = ArrayObject([
        link(script("app.alert('link')")),
        link(web("https://example.com/kept")),
        link(web("https://example.com/chained", **{name("/Next"): script("next()")})),
        link(web(" Java Script:alert(1)")),
        link(web("\x01javascript:alert(1)")),
        link(web("file:///C:/Windows/System32/calc.exe")),
        link(web("ms-msdt:/id PCWDiagnostic")),
        link(DictionaryObject({name("/S"): name("/GoToR"), name("/F"): TextStringObject("\\\\server\\share\\x.pdf")})),
        link(DictionaryObject({name("/S"): name("/Launch"), name("/F"): TextStringObject("calc.exe")})),
        link(DictionaryObject({name("/S"): name("/Named"), name("/N"): name("/Print")})),
        link(DictionaryObject({name("/S"): name("/Named"), name("/N"): name("/NextPage")})),
        link(web("MAILTO:team@example.com")),
        widget_reference,
    ])
    writer.root_object[name("/AcroForm")] = DictionaryObject({
        name("/Fields"): ArrayObject([field_reference]), name("/XFA"): TextStringObject("<xdp:xdp/>"),
    })
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def part(source_id, file_name, data, **kwargs):
    return DocumentMergePart(source_id, file_name, lambda data=data: data, **kwargs)


def read_pdf(result):
    from pypdf import PdfReader

    return PdfReader(io.BytesIO(result.read_bytes()))


def outline_titles(reader):
    titles = []
    for item in reader.outline:
        titles.append([entry.title for entry in item] if isinstance(item, list) else item.title)
    return titles


def xlsx_bytes(sheets):
    from openpyxl import Workbook

    workbook = Workbook()
    first = True
    for title, rows, state, formats in sheets:
        worksheet = workbook.active if first else workbook.create_sheet()
        first = False
        worksheet.title = title
        worksheet.sheet_state = state
        for row in rows:
            worksheet.append(row)
        for cell, number_format in (formats or {}).items():
            worksheet[cell].number_format = number_format
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_version_includes_pdf_and_workbook_merges():
    assert_app_version_at_least("0.261.221")


def test_pdfs_merge_in_order_with_a_bookmark_per_file():
    parts = [part("a", "Alpha.pdf", pdf_bytes("A", 2)), part("b", "Beta.pdf", pdf_bytes("B", 3))]
    with merge_documents("pdf", parts) as result:
        reader = read_pdf(result)
        assert len(reader.pages) == 5
        assert "A page 1" in reader.pages[0].extract_text()
        assert "B page 3" in reader.pages[4].extract_text()
        assert outline_titles(reader) == ["Alpha.pdf", ["A intro"], "Beta.pdf", ["B intro"]]
        assert result.media_type == "application/pdf" and result.file_extension == ".pdf"
        report = result.report
        assert report["status"] == "merged"
        assert report["totals"]["pages"] == 5
        assert [entry["pages"] for entry in report["parts"]] == [2, 3]


def test_page_ranges_select_and_order_pages_and_bookmarks_can_be_turned_off():
    parts = [
        part("a", "Alpha.pdf", pdf_bytes("A", 3), pages=((3, 3), (1, 1))),
        part("b", "Beta.pdf", pdf_bytes("B", 2)),
    ]
    with merge_documents("pdf", parts, options=DocumentMergeOptions(bookmarks=False)) as result:
        reader = read_pdf(result)
        assert [page.extract_text().strip() for page in reader.pages[:2]] == ["A page 3", "A page 1"]
        assert reader.outline == []
        assert result.report["parts"][0]["source_pages"] == 3
    for ranges in (((0, 1),), ((2, 9),), ((2, 1),), ((1, 2), (2, 3))):
        with pytest.raises(DocumentMergeError) as caught:
            merge_documents("pdf", [part("a", "a.pdf", pdf_bytes("A", 3), pages=ranges), parts[1]])
        assert caught.value.code == "range_invalid"


def test_encrypted_damaged_and_mislabeled_pdfs_are_refused():
    good = part("b", "b.pdf", pdf_bytes("B", 1))
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("pdf", [part("a", "locked.pdf", pdf_bytes("A", 1, encrypt=True)), good])
    assert caught.value.code == "encrypted_document"
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("pdf", [part("a", "broken.pdf", b"not a pdf"), good])
    assert caught.value.code == "unreadable_document"
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("pdf", [part("a", "notes.docx", b"PK"), good])
    assert caught.value.code == "unsupported_format"


def test_document_level_scripts_are_not_carried_over():
    parts = [part("a", "a.pdf", pdf_with_script("A")), part("b", "b.pdf", pdf_bytes("B", 1))]
    with merge_documents("pdf", parts) as result:
        root = read_pdf(result).trailer["/Root"]
        names = root.get("/Names")
        assert names is None or "/JavaScript" not in names.get_object()
        assert "/OpenAction" not in root


def test_actions_that_could_run_code_are_removed_and_links_are_kept():
    source = pdf_with_active_content()
    assert b"/JavaScript" in source and b"/XFA" in source
    parts = [part("s", "scripted.pdf", source), part("b", "b.pdf", pdf_bytes("B", 1))]
    with merge_documents("pdf", parts) as result:
        merged = result.read_bytes()
        warnings = result.report["warnings"]
    for token in (b"/JavaScript", b"/JS", b"/AA", b"/XFA", b"/Launch", b"/GoToR", b"/Print", b"/OpenAction"):
        assert token not in merged, token
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(merged))
    assert "S page 1" in reader.pages[0].extract_text()
    annotations = [annotation.get_object() for annotation in reader.pages[0]["/Annots"]]
    actions = [annotation["/A"] for annotation in annotations if "/A" in annotation]
    # pypdf escapes these strings in the file, so the kept links are checked parsed.
    assert [(action["/S"], action.get("/URI", action.get("/N"))) for action in actions] == [
        ("/URI", "https://example.com/kept"), ("/Named", "/NextPage"), ("/URI", "MAILTO:team@example.com"),
    ]
    [widget] = [annotation for annotation in annotations if annotation["/Subtype"] == "/Widget"]
    assert "/AA" not in widget and "/AA" not in widget["/Parent"]
    form = reader.trailer["/Root"]["/AcroForm"]
    assert "/XFA" not in form and len(form["/Fields"]) == 1
    assert [warning["code"] for warning in warnings] == ["active_content_removed"]
    assert "links to pages and web or email addresses were kept" in warnings[0]["message"]


def test_the_same_files_always_give_the_same_bytes():
    workbook_parts = [
        part("x", "east.xlsx", xlsx_bytes([("Data", [["Code"], ["007"]], "visible", None)])),
        part("y", "west.csv", b"Code\n008\n"),
    ]
    pdf_parts = [part("a", "a.pdf", pdf_bytes("A", 2)), part("b", "b.pdf", pdf_bytes("B", 1))]
    with merge_documents("workbook", workbook_parts) as first_workbook, merge_documents("pdf", pdf_parts) as first_pdf:
        workbook_bytes, pdf_content = first_workbook.read_bytes(), first_pdf.read_bytes()
    # Two seconds is the resolution of a ZIP entry's date.
    time.sleep(2.1)
    with merge_documents("workbook", workbook_parts) as second_workbook, merge_documents("pdf", pdf_parts) as second_pdf:
        assert second_workbook.read_bytes() == workbook_bytes
        assert second_pdf.read_bytes() == pdf_content
    with zipfile.ZipFile(io.BytesIO(workbook_bytes)) as package:
        assert {entry.date_time for entry in package.infolist()} == {(2000, 1, 1, 0, 0, 0)}
        properties = package.read("docProps/core.xml").decode("utf-8")
    assert properties.count("2000-01-01T00:00:00Z") == 2


def test_control_characters_excel_cant_store_are_removed_and_counted():
    from openpyxl import load_workbook

    parts = [
        part("x", "legacy\x07export.csv", b"Code,Name\n001,Al\x0bice\x1a\n002,Bo\n"),
        part("y", "west.csv", b"Code\n008\n"),
    ]
    with merge_documents("workbook", parts) as result:
        workbook = load_workbook(io.BytesIO(result.read_bytes()))
        warnings = result.report["parts"][0]["warnings"]
    assert workbook.sheetnames == ["legacy export", "west"]
    assert workbook["legacy export"]["B2"].value == "Alice"
    assert warnings == [{
        "code": "control_characters_removed", "message": "2 control character(s) that Excel can't store were removed.",
    }]


def test_library_failures_name_only_the_file(monkeypatch):
    import functions_document_merge_workbook
    import pypdf

    def leak(*_args, **_kwargs):
        raise RuntimeError("PRIVATE cell text")

    workbook_parts = [part("x", "east.csv", b"A\n1\n"), part("y", "west.csv", b"A\n2\n")]
    monkeypatch.setattr(functions_document_merge_workbook, "_write_rows", leak)
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("workbook", workbook_parts)
    assert (caught.value.code, caught.value.message) == ("unreadable_document", "east.csv couldn't be copied into the workbook.")
    monkeypatch.undo()

    pdf_parts = [part("a", "a.pdf", pdf_bytes("A", 1)), part("b", "b.pdf", pdf_bytes("B", 1))]
    monkeypatch.setattr(pypdf.PdfWriter, "append", leak)
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("pdf", pdf_parts)
    assert (caught.value.code, caught.value.message) == ("unreadable_document", "a.pdf couldn't be read as a PDF.")
    monkeypatch.undo()

    monkeypatch.setattr(pypdf.PdfWriter, "write", leak)
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("pdf", pdf_parts)
    assert (caught.value.code, caught.value.message) == ("merge_failed", "The merged PDF couldn't be written.")
    monkeypatch.undo()

    def disk_full(*_args, **_kwargs):
        raise OSError("No space left on device")

    # A full temporary disk is the host's problem, not the file's, so it passes through to be retried.
    monkeypatch.setattr(functions_document_merge_workbook, "_write_rows", disk_full)
    with pytest.raises(OSError):
        merge_documents("workbook", workbook_parts)
    monkeypatch.setattr(pypdf.PdfWriter, "write", disk_full)
    with pytest.raises(OSError):
        merge_documents("pdf", pdf_parts)


@pytest.mark.parametrize("failing_call", [2, 5])
def test_a_failing_cancellation_check_is_not_blamed_on_a_file(failing_call):
    class RunStoreUnavailable(RuntimeError):
        pass

    calls = []

    def cancel_requested():
        calls.append(len(calls) + 1)
        if len(calls) >= failing_call:
            raise RunStoreUnavailable("The run store is unavailable.")
        return False

    # Call 2 is the first sheet of one.csv; call 5 is while the package is written.
    parts = [part("x", "one.csv", b"A\n1\n"), part("y", "two.csv", b"A\n2\n")]
    with pytest.raises(RunStoreUnavailable):
        merge_documents("workbook", parts, cancel_requested=cancel_requested)
    assert len(calls) == failing_call


def test_pdf_limits_and_cancellation_fail_closed():
    parts = [part("a", "a.pdf", pdf_bytes("A", 3)), part("b", "b.pdf", pdf_bytes("B", 3))]
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("pdf", parts, limits=DocumentMergeLimits(max_pages=5))
    assert caught.value.code == "page_limit_exceeded"
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("pdf", parts, limits=DocumentMergeLimits(max_output_bytes=10))
    assert caught.value.code == "output_too_large"
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("pdf", parts, limits=DocumentMergeLimits(max_source_bytes=10))
    assert caught.value.code == "source_too_large"
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("pdf", parts + [part("c", "c.pdf", pdf_bytes("C", 1))], limits=DocumentMergeLimits(max_parts=2))
    assert caught.value.code == "too_many_sources"
    with pytest.raises(DocumentMergeCancelled):
        merge_documents("pdf", parts, cancel_requested=lambda: True)
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("pdf", [parts[0]])
    assert caught.value.code == "too_few_sources"
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("pdf", [parts[0], DocumentMergePart("a", "again.pdf", parts[1].load_bytes)])
    assert caught.value.code == "duplicate_source"


def test_progress_reports_each_file():
    events = []
    parts = [part("a", "a.pdf", pdf_bytes("A", 1)), part("b", "b.pdf", pdf_bytes("B", 2))]
    with merge_documents("pdf", parts, on_progress=events.append):
        pass
    assert [(event["index"], event["pages"]) for event in events] == [(1, 1), (2, 2)]


def test_workbook_gets_one_typed_sheet_per_source():
    from openpyxl import load_workbook

    excel = xlsx_bytes([(
        "Data",
        [["Code", "When", "Rate", "Note"], ["007", datetime(2024, 1, 2), 0.125, "=SUM(A1:A2)"]],
        "visible", {"C2": "0.00%"},
    )])
    # openpyxl stores "=SUM(...)" as a formula; a literal formula-like string comes from CSV.
    csv_data = b"Code,Note\n007,=HYPERLINK(\"http://example\")\n"
    parts = [part("x", "Sales East.xlsx", excel), part("y", "west:data?.csv", csv_data)]
    with merge_documents("workbook", parts) as result:
        workbook = load_workbook(io.BytesIO(result.read_bytes()))
        assert workbook.sheetnames == ["Sales East", "west data"]
        east = workbook["Sales East"]
        assert east["A2"].value == "007" and east["B2"].value == datetime(2024, 1, 2)
        assert east["C2"].value == 0.125 and east["C2"].number_format == "0.00%"
        assert east["D2"].value is None, "A formula without a cached value has no calculated value to copy."
        west = workbook["west data"]
        assert west["A2"].value == "007"
        assert west["B2"].value == '=HYPERLINK("http://example")' and west["B2"].data_type == "s"
        assert result.media_type.endswith("spreadsheetml.sheet")
        assert result.report["totals"]["sheets"] == 2
        assert any("CSV files are copied as text" in item for item in result.report["limitations"])


def test_workbook_sheet_modes_names_and_hidden_sheets():
    from openpyxl import load_workbook

    first = xlsx_bytes([
        ("Data", [["A"], [1]], "visible", None),
        ("Hidden", [["H"], [2]], "hidden", None),
        ("Notes", [["N"], [3]], "visible", None),
    ])
    second = xlsx_bytes([("Data", [["A"], [4]], "visible", None)])
    parts = [part("x", "Report.xlsx", first), part("y", "Report.xlsx", second)]
    with merge_documents("workbook", parts, options=DocumentMergeOptions(sheets="all")) as result:
        names = load_workbook(io.BytesIO(result.read_bytes())).sheetnames
        assert names == ["Report - Data", "Report - Notes", "Report - Data (2)"]
        assert result.report["warnings"][0]["code"] == "hidden_sheets_skipped"
    with merge_documents("workbook", parts, options=DocumentMergeOptions(sheet="data")) as result:
        names = load_workbook(io.BytesIO(result.read_bytes())).sheetnames
        assert names == ["Report", "Report (2)"]
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("workbook", parts, options=DocumentMergeOptions(sheet="Totals"))
    assert caught.value.code == "sheet_not_found"
    long_name = "A" * 40 + ".csv"
    with merge_documents("workbook", [part("l", long_name, b"A\n1\n"), part("m", long_name, b"A\n2\n")]) as result:
        names = load_workbook(io.BytesIO(result.read_bytes())).sheetnames
        assert names == ["A" * 31, "A" * 27 + " (2)"]


def test_workbook_limits_and_damaged_inputs():
    csv_part = part("y", "y.csv", b"A\n1\n2\n3\n")
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("workbook", [csv_part, part("z", "z.csv", b"A\n1\n")], limits=DocumentMergeLimits(max_total_cells=3))
    assert caught.value.code == "cell_limit_exceeded"
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("workbook", [csv_part, part("z", "z.csv", b"A\n1\n")], limits=DocumentMergeLimits(max_sheets=1))
    assert caught.value.code == "sheet_limit_exceeded"
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("workbook", [csv_part, part("z", "locked.xlsx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 32)])
    assert caught.value.code == "encrypted_document"
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("workbook", [csv_part, part("z", "broken.xls", b"<html></html>")])
    assert caught.value.code == "unreadable_document"
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("workbook", [csv_part, part("z", "binary.csv", b"\x00\x01\x02")])
    assert caught.value.code == "unreadable_document"
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("workbook", [csv_part, part("z", "z.pdf", b"%PDF-")])
    assert caught.value.code == "unsupported_format"


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
