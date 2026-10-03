#!/usr/bin/env python3
# test_workflow_merge_task.py
"""
Functional test for workflow Merge tasks.
Version: 0.261.224
Implemented in: 0.261.220
Scale and formula checks added in: 0.261.224
Refs: microsoft/simplechat#1619

This test ensures that a workflow task can merge many CSV and Excel files into one CSV or
Excel file with code. The task's merge contract is normalized and validated with the same
rules chat merges use, and only workflow callers accept it. Files are kept in the order the
owner chose, or found when the run starts (every matching file, recently added files, or the
files a File Sync run changed) and merged in file-name order. Every file is authorized again
and read through the access boundary; the merged file is rendered by the shared export
framework and attached to the run's conversation once; cancellation stops the merge before
anything is published; and failures name what the owner can fix. One hundred files of a
thousand rows merge in one run, and merged CSV and Excel files never carry a live formula.
A chat-proposed blueprint builds the same merge action. No model, Azure service or network
is used.
"""

import copy
import csv
import hashlib
import importlib
import io
import socket
import sys
from datetime import datetime, timezone
from pathlib import Path
from types import ModuleType, SimpleNamespace
from unittest.mock import patch

import pytest

sys.path.append(str(Path(__file__).resolve().parent))

from test_support.app_source import run_definitions  # noqa: E402
from test_support.app_stubs import stubbed_app_imports  # noqa: E402
from test_support.versioning import assert_app_version_at_least  # noqa: E402
from test_workflow_draft_service import ORIGIN, USER_INFO, DraftHarness  # noqa: E402
from test_workflow_draft_save_parity import OWNER_ID  # noqa: E402


APP_ROOT = Path(__file__).resolve().parents[1] / "application" / "single_app"
# A merge loads its spreadsheet engine lazily, by module name.
if str(APP_ROOT) not in sys.path:
    sys.path.insert(0, str(APP_ROOT))

OWNER = "owner-1"
CONVERSATION = "conversation-1"
RUN_ID = "run-1"
TASK_ID = "task-1"
SETTINGS = {"document_action_capabilities": {"merge": {"enabled": True}}}
ALL_TYPES = {"none", "search", "analyze", "comparison", "merge"}

NORTH = ("north.csv", b"Region,Amount,Rep\r\nNorth,10,Ana\r\nNorth,12,Bo\r\n")
SOUTH = ("South.csv", b"Amount,Region,Rep\n7,South,Cy\n")
WEST = ("west.csv", b"Region,Total\nWest,3\n")
NOTES = ("notes.docx", b"not a spreadsheet")


def xlsx_bytes(sheets):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets:
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


EAST = ("east.xlsx", xlsx_bytes([("Sales", [["Region", "Amount", "Rep"], ["East", 5, "Di"]])]))


def _seam_modules():
    """The two service-backed modules functions_document_actions imports, without Azure."""
    search_source = run_definitions(
        "functions_search.py", ("VALID_SEARCH_SCOPES", "normalize_search_scope", "normalize_search_id_list"), {},
    )
    search = ModuleType("functions_search")
    search.normalize_search_scope = search_source["normalize_search_scope"]
    search.normalize_search_id_list = search_source["normalize_search_id_list"]

    analysis = ModuleType("functions_document_analysis")
    analysis.CHAT_DOCUMENT_ANALYSIS_MAX_DOCUMENTS = 3
    analysis.WORKFLOW_DOCUMENT_ANALYSIS_MAX_DOCUMENTS = 10

    def refuse(**kwargs):
        raise AssertionError("A merge never normalizes analysis targets.")

    analysis.normalize_document_analysis_targets = refuse
    return {"functions_search": search, "functions_document_analysis": analysis}


@pytest.fixture(scope="module")
def app():
    """The real merge modules over stand-ins for the services they import; nothing leaves the process."""
    before = set(sys.modules)
    seams = _seam_modules()
    installed = [name for name in seams if name not in sys.modules]
    for name in installed:
        sys.modules[name] = seams[name]
    try:
        with stubbed_app_imports(), patch.object(
            socket.socket, "connect", side_effect=AssertionError("Unexpected external I/O"),
        ):
            yield SimpleNamespace(
                actions=importlib.import_module("functions_document_actions"),
                merge=importlib.import_module("functions_workflow_merge"),
                orchestration_merge=importlib.import_module("functions_orchestration_merge"),
                exports=importlib.import_module("functions_generated_file_exports"),
                contracts=importlib.import_module("functions_generated_export_contracts"),
                office=importlib.import_module("functions_office_file_renderers"),
                search=seams["functions_search"],
            )
    finally:
        # Only modules this fixture loaded from the application, and its own stand-ins, are removed.
        for name in set(sys.modules) - before:
            module = sys.modules[name]
            if name in installed or str(getattr(module, "__file__", "") or "").startswith(str(APP_ROOT)):
                sys.modules.pop(name, None)


def normalize(app, payload, **kwargs):
    kwargs.setdefault("allowed_action_types", ALL_TYPES)
    return app.actions.normalize_document_action_config(copy.deepcopy(payload), **kwargs)


def merge_action(**changes):
    action = {"type": "merge", "document_ids": ["north", "south"]}
    action.update(changes)
    return action


class MergeWorld:
    """The files one owner can read, and the services a workflow merge calls, recorded."""

    def __init__(self, files, *, collected=(), unavailable=()):
        self.files = dict(files)
        self.collected = list(collected)
        self.unavailable = set(unavailable)
        self.collect_calls = []
        self.resolve_calls = []
        self.reads = []
        self.published = []

    def collect(self, action, extensions, max_documents):
        self.collect_calls.append((copy.deepcopy(action), tuple(extensions), max_documents))
        return list(self.collected)

    def resolve(self, document_ids, **kwargs):
        self.resolve_calls.append((list(document_ids), dict(kwargs)))
        manifest = []
        for document_id in document_ids:
            file_name = self.files[document_id][0]
            tabular = file_name.lower().endswith((".csv", ".xlsx", ".xlsm", ".xls"))
            manifest.append({
                "document_id": document_id, "file_name": file_name, "scope": "personal", "scope_id": OWNER,
                "source_kind": "tabular" if tabular else "narrative",
                "authorization_status": "unresolved" if document_id in self.unavailable else "authorized",
            })
        return manifest

    def read(self, source, user_id, group_id=None, public_workspace_id=None, *, purpose):
        assert user_id == OWNER and purpose == "native"
        self.reads.append(source["document_id"])
        return {"id": source["document_id"]}, self.files[source["document_id"]][1]

    def publish(self, *, user_id, conversation_id, file_name, output_format, rendered, row_count, summary,
                idempotency_key):
        content = rendered.file_content.read()
        assert len(content) == rendered.size_bytes
        self.published.append({
            "user_id": user_id, "conversation_id": conversation_id, "file_name": file_name,
            "output_format": output_format, "content": content, "row_count": row_count, "summary": summary,
            "idempotency_key": idempotency_key,
        })
        return {
            "capability": "file_merge", "artifact_message_id": f"artifact-{len(self.published)}",
            "conversation_id": conversation_id, "file_name": file_name, "output_format": output_format,
            "row_count": row_count, "summary": summary,
        }

    def run(self, app, action, *, settings=None, cancel=None, progress=None, render=None):
        def real_render(source, request, max_output_bytes):
            return app.exports.build_generated_file_export(
                source=source, export_request=request, max_output_bytes=max_output_bytes,
            )

        return app.merge.execute_workflow_merge(
            copy.deepcopy(action), settings or SETTINGS, user_id=OWNER, conversation_id=CONVERSATION,
            run_id=RUN_ID, task_id=TASK_ID, collect_documents=self.collect, resolve_manifest=self.resolve,
            render_file=render or real_render, publish_file=self.publish, byte_reader=self.read,
            cancel_requested=cancel, report_progress=progress,
        )


def world(**kwargs):
    files = {"north": NORTH, "south": SOUTH, "west": WEST, "notes": NOTES, "east": EAST}
    return MergeWorld(files, **kwargs)


def csv_rows(content):
    return list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))


def test_version_includes_workflow_merge_tasks():
    assert_app_version_at_least("0.261.220")


# ---------------------------------------------------------------------------
# The task contract
# ---------------------------------------------------------------------------

def test_a_selected_merge_keeps_the_chosen_order_and_its_settings(app):
    action = normalize(app, merge_action(
        document_ids=["south", "north", "west", "north"], output_format="XLSX",
        output_file_name="  Q3   sales.xlsx ", merge_options={"dedupe": "exact_rows", "sort_by": None},
        doc_scope="Personal", active_group_ids=["g1", "g1"],
    ))
    assert action["type"] == "merge" and action["merge_kind"] == "tabular"
    assert action["target_mode"] == "selected"
    assert action["document_ids"] == ["south", "north", "west"]
    assert (action["doc_scope"], action["active_group_ids"]) == ("personal", ["g1"])
    assert (action["output_format"], action["output_file_name"]) == ("xlsx", "Q3 sales")
    # Unset options are dropped; the rest are kept exactly as chat merges read them.
    assert action["merge_options"] == {"dedupe": "exact_rows"}
    assert "merge_targets_resolved" not in action

    defaults = normalize(app, merge_action())
    assert (defaults["output_format"], defaults["merge_options"], defaults["doc_scope"]) == ("csv", {}, "all")
    assert "output_file_name" not in defaults


@pytest.mark.parametrize("target_mode", ["all", "recent", "changed"])
def test_files_found_at_run_time_are_never_stored_with_the_task(app, target_mode):
    action = normalize(app, merge_action(target_mode=target_mode, document_ids=["north", "south", "west"]))
    assert action["target_mode"] == target_mode and action["document_ids"] == []
    if target_mode == "recent":
        assert action["recent_window_minutes"] == 60
        assert normalize(app, merge_action(target_mode="recent", recent_window_minutes=15))["recent_window_minutes"] == 15
    # The runner pins the files it found for one run; only then are ids kept, with that mark.
    resolved = normalize(app, merge_action(target_mode=target_mode, document_ids=["west"], merge_targets_resolved=True))
    assert resolved["document_ids"] == ["west"] and resolved["merge_targets_resolved"] is True


@pytest.mark.parametrize(("changes", "message"), [
    ({"document_ids": ["north"]}, "Select at least two files to merge."),
    ({"merge_kind": "zip"}, "Merge kind must be one of: tabular, workbook, pdf, docx, pptx."),
    ({"target_mode": "current_item"}, "Merge files must be one of: selected, all, recent, changed."),
    ({"doc_scope": "tenant"}, "The merge workspace scope must be all, personal, group or public."),
    ({"output_format": "pdf"}, "A spreadsheet merge creates CSV or XLSX files."),
    ({"output_file_name": ".hidden"}, "must not start with a dot"),
    ({"output_file_name": "a/b"}, 'must not contain / \\ : * ? " < > |'),
    ({"output_file_name": "x" * 101}, "at most 100 characters"),
    ({"output_file_name": 5}, "The merged file name must be text."),
    ({"merge_options": "by_name"}, "Merge options must be an object."),
    ({"merge_options": {"bookmarks": True}}, "Some merge options don't apply to spreadsheet merges. Remove them."),
    ({"merge_options": {"schema_policy": "mapped"}}, "column"),
    ({"merge_options": {"dedupe": "key_columns"}}, "column"),
])
def test_a_merge_the_contract_cannot_run_is_refused_with_what_to_fix(app, changes, message):
    with pytest.raises(ValueError) as caught:
        normalize(app, merge_action(**changes))
    assert message in str(caught.value)


def test_the_merge_file_limit_applies_to_selected_files(app):
    with pytest.raises(ValueError, match="A merge supports up to 2 files at a time."):
        normalize(app, merge_action(document_ids=["a", "b", "c"]), max_documents_by_type={"merge": 2})
    assert normalize(app, merge_action(), max_documents_by_type={"merge": 2})["document_ids"] == ["north", "south"]
    # Files a run found are counted by the merge, which names the count; the runner normalizes
    # the action again before dispatch, so the contract must not refuse them first.
    found = normalize(
        app, merge_action(target_mode="changed", document_ids=["a", "b", "c"], merge_targets_resolved=True),
        max_documents_by_type={"merge": 2},
    )
    assert found["document_ids"] == ["a", "b", "c"]


def test_only_workflow_callers_accept_a_merge(app):
    actions = app.actions
    chat_types = actions.get_enabled_document_action_types(SETTINGS)
    assert "merge" not in chat_types
    with pytest.raises(ValueError, match="File merging is turned off in admin settings, or is not available here."):
        normalize(app, merge_action(), allowed_action_types=chat_types)

    workflow_types = actions.get_enabled_document_action_types(SETTINGS, include_merge=True)
    assert "merge" in workflow_types
    assert normalize(app, merge_action(), allowed_action_types=workflow_types)["type"] == "merge"

    disabled = {"document_action_capabilities": {"merge": {"enabled": False}}}
    assert "merge" not in actions.get_enabled_document_action_types(disabled, include_merge=True)


def test_document_merge_options_follow_their_own_rules(app):
    options = app.actions._normalize_merge_options
    assert options("docx", {"formatting": "use_first", "page_breaks": True, "source_headings": None}) == {
        "formatting": "use_first", "page_breaks": True,
    }
    assert options("pptx", None) == {}
    assert options("workbook", {"sheets": "all"}) == {"sheets": "all"}
    for kind, value, message in (
        ("pdf", {"bookmarks": "yes"}, "The bookmarks merge option must be true or false."),
        ("pptx", {"formatting": "mixed"}, "Formatting must be keep_source or use_first."),
        ("workbook", {"sheets": "some"}, "Sheets must be first or all."),
        ("workbook", {"sheet": "x" * 32}, "A sheet name must be 1 to 31 characters."),
        ("workbook", {"sheet": "Sales", "sheets": "all"}, "Name one sheet, or read all sheets, but not both."),
        ("pdf", {"formatting": "keep_source"}, "Some merge options don't apply to PDF merges. Remove them."),
    ):
        with pytest.raises(ValueError) as caught:
            options(kind, value)
        assert message in str(caught.value)


# ---------------------------------------------------------------------------
# Running a merge
# ---------------------------------------------------------------------------

def test_a_selected_merge_creates_one_csv_in_the_chosen_order(app):
    files = world()
    updates = []
    result = files.run(app, normalize(app, merge_action(document_ids=["south", "north"])), progress=updates.append)

    [published] = files.published
    assert published["file_name"] == "merged.csv" and published["output_format"] == "csv"
    # The key names the run, the task and the bytes: a replay of the same merge reuses its file.
    assert published["idempotency_key"] == (
        f"workflow-merge:{RUN_ID}:{TASK_ID}:{hashlib.sha256(published['content']).hexdigest()}"
    )
    assert (published["user_id"], published["conversation_id"]) == (OWNER, CONVERSATION)
    assert published["row_count"] == 3
    assert published["summary"] == "Merged 3 row(s) from 2 file(s)."
    assert csv_rows(published["content"]) == [
        ["Source File", "Amount", "Region", "Rep"],
        ["South.csv", "7", "South", "Cy"],
        ["north.csv", "10", "North", "Ana"],
        ["north.csv", "12", "North", "Bo"],
    ]
    assert files.reads == ["south", "north"]
    [(resolved_ids, resolve_options)] = files.resolve_calls
    assert resolved_ids == ["south", "north"]
    assert resolve_options["user_id"] == OWNER and resolve_options["request_correlation_id"] == RUN_ID
    assert not files.collect_calls
    assert [(update["index"], update["total"], update["file_name"]) for update in updates] == [
        (1, 2, "South.csv"), (2, 2, "north.csv"),
    ]

    assert result["generated_analysis_artifacts"] == result["generated_tabular_outputs"] == [{
        "capability": "file_merge", "artifact_message_id": "artifact-1", "conversation_id": CONVERSATION,
        "file_name": "merged.csv", "output_format": "csv", "row_count": 3, "summary": "Merged 3 row(s) from 2 file(s).",
    }]
    assert result["token_usage"] == {} and result["model_deployment_name"] is None and result["provider"] is None
    assert result["merge_summary"] == {
        "status": "merged", "kind": "tabular", "files": 2, "rows": 3, "columns": 4, "duplicates_removed": 0,
        "excluded": 0, "skipped": 0, "file_name": "merged.csv",
    }
    assert result["reply"].startswith("Merged **3 row(s)** from **2 file(s)** into **merged.csv** with 4 column(s).")
    assert result["reply"].endswith("Files merged, in order: South.csv, north.csv.")


def test_an_excel_merge_holds_the_same_rows_and_reads_excel_files(app):
    from openpyxl import load_workbook

    files = world()
    action = normalize(app, merge_action(
        document_ids=["north", "east"], output_format="xlsx", output_file_name="Regional sales",
        merge_options={"include_source_column": False},
    ))
    result = files.run(app, action)
    [published] = files.published
    assert published["file_name"] == "Regional sales.xlsx"
    rows = [list(row) for row in load_workbook(io.BytesIO(published["content"])).active.iter_rows(values_only=True)]
    assert [str(value) for value in rows[0]] == ["Region", "Amount", "Rep"]
    assert [[str(value) for value in row] for row in rows[1:]] == [
        ["North", "10", "Ana"], ["North", "12", "Bo"], ["East", "5", "Di"],
    ]
    # Only CSV results are offered to later tasks as tabular inputs.
    assert result["generated_tabular_outputs"] == []
    assert result["merge_summary"]["file_name"] == "Regional sales.xlsx"


def test_one_hundred_files_merge_in_one_workflow_run(app):
    import time

    files = {
        f"file-{index:03d}": (
            f"region-{index:03d}.csv",
            ("Region,Amount,Rep\r\n" + "".join(
                f"R{index},{row},Rep {row % 7}\r\n" for row in range(1000)
            )).encode("utf-8"),
        )
        for index in range(100)
    }
    merge_world = MergeWorld(files)
    started = time.monotonic()
    result = merge_world.run(app, normalize(app, merge_action(document_ids=list(files))))
    elapsed = time.monotonic() - started

    [published] = merge_world.published
    assert merge_world.reads == list(files)
    assert published["row_count"] == 100_000
    rows = csv_rows(published["content"])
    assert len(rows) == 100_001
    assert rows[0] == ["Source File", "Region", "Amount", "Rep"]
    assert rows[1] == ["region-000.csv", "R0", "0", "Rep 0"]
    assert rows[-1] == ["region-099.csv", "R99", "999", "Rep 5"]
    assert result["merge_summary"]["files"] == 100 and result["merge_summary"]["rows"] == 100_000
    # A generous bound that still catches work that grows faster than the rows do.
    assert elapsed < 120, elapsed


@pytest.mark.parametrize("output_format", ["csv", "xlsx"])
def test_merged_files_never_carry_live_formulas(app, output_format):
    from openpyxl import load_workbook

    files = MergeWorld({
        "first": ("first.csv", b"Name,Value\r\nA,=1+1\r\nB,+cmd|' /C calc'!A0\r\nC,-2\r\n"),
        "second": ("second.csv", b"Name,Value\r\nD,@SUM(A1:A2)\r\nE,\"\t=HYPERLINK(\"\"http://x\"\")\"\r\nF,-10.5\r\n"),
    })
    files.run(app, normalize(app, merge_action(
        document_ids=["first", "second"], output_format=output_format,
        merge_options={"include_source_column": False},
    )))
    [published] = files.published
    if output_format == "csv":
        values = [row[1] for row in csv_rows(published["content"])[1:]]
        # Formula-like text is prefixed so spreadsheet apps show it as text; signed numbers stay numbers.
        assert values == [
            "'=1+1", "'+cmd|' /C calc'!A0", "-2", "'@SUM(A1:A2)", "'\t=HYPERLINK(\"http://x\")", "-10.5",
        ]
    else:
        sheet = load_workbook(io.BytesIO(published["content"])).active
        cells = [row[1] for row in sheet.iter_rows(min_row=2)]
        assert [cell.value for cell in cells] == [
            "=1+1", "+cmd|' /C calc'!A0", "-2", "@SUM(A1:A2)", "\t=HYPERLINK(\"http://x\")", "-10.5",
        ]
        # Every value is stored as text; none is a formula Excel would calculate.
        assert {cell.data_type for cell in cells} == {"s"}


def test_files_found_at_run_time_skip_what_cannot_be_merged_and_say_so(app):
    files = world(collected=["north", "notes", "south", "west"], unavailable=["west"])
    action = normalize(app, merge_action(target_mode="all", doc_scope="personal"))
    result = files.run(app, action)

    [(collected_action, extensions, max_documents)] = files.collect_calls
    assert collected_action["target_mode"] == "all"
    assert extensions == (".csv", ".xlsx", ".xlsm", ".xls") and max_documents == 100
    assert files.reads == ["north", "south"]
    assert result["merge_summary"]["files"] == 2 and result["merge_summary"]["skipped"] == 2
    assert "- Skipped 1 file(s) that aren't the right type: notes.docx." in result["reply"]
    assert "- Skipped 1 file(s) that aren't available right now: west.csv." in result["reply"]


def test_one_changed_file_is_still_merged_into_a_file(app):
    files = world()
    action = normalize(app, merge_action(target_mode="changed", document_ids=["west"], merge_targets_resolved=True))
    result = files.run(app, action)
    assert csv_rows(files.published[0]["content"]) == [["Source File", "Region", "Total"], ["west.csv", "West", "3"]]
    assert result["merge_summary"]["files"] == 1


@pytest.mark.parametrize(("target_mode", "collected", "document_ids", "text"), [
    ("all", [], [], "No files in the workspace matched this merge"),
    ("recent", [], [], "No files were added or updated in the merge's time window"),
    ("changed", [], [], "No new or changed files arrived with this sync"),
    ("all", ["notes"], [], "No files in the workspace matched this merge"),
])
def test_nothing_to_merge_creates_no_file(app, target_mode, collected, document_ids, text):
    files = world(collected=collected)
    action = normalize(app, merge_action(
        target_mode=target_mode, document_ids=document_ids, merge_targets_resolved=target_mode == "changed",
    ))
    result = files.run(app, action)
    assert result["reply"] == f"{text}, so there was nothing to merge and no file was created."
    assert result["merge_summary"] == {"status": "nothing_to_merge", "files": 0, "rows": 0}
    assert result["generated_analysis_artifacts"] == [] and not files.published and not files.reads


def test_a_selected_merge_fails_closed(app):
    files = world(unavailable=["north"])
    with pytest.raises(app.merge.WorkflowMergeAccessError, match="no longer available"):
        files.run(app, normalize(app, merge_action(document_ids=["south", "north"])))

    files = world()
    with pytest.raises(app.merge.WorkflowMergeError) as wrong_type:
        files.run(app, normalize(app, merge_action(document_ids=["north", "notes"])))
    assert str(wrong_type.value) == (
        "notes.docx isn't a CSV or Excel file (.csv, .xlsx, .xlsm or .xls), so it can't be merged."
    )

    files = world()
    files.resolve = lambda document_ids, **kwargs: MergeWorld.resolve(files, list(reversed(document_ids)), **kwargs)
    with pytest.raises(app.merge.WorkflowMergeAccessError, match="could not be authorized"):
        files.run(app, normalize(app, merge_action()))
    # A refusal is still a PermissionError for any caller that handles those.
    assert issubclass(app.merge.WorkflowMergeAccessError, PermissionError)

    with pytest.raises(app.merge.WorkflowMergeError, match="Select at least two files to merge."):
        world().run(app, {"type": "merge", "merge_kind": "tabular", "target_mode": "selected", "document_ids": []})
    assert not files.published and not files.reads


def test_more_files_than_the_limit_fail_instead_of_merging_some(app):
    limited = {"document_action_capabilities": {"merge": {"enabled": True, "workflow_max_documents": 2}}}
    files = world()
    # As the runner does: the stored action is normalized again with the administrator's limit.
    action = normalize(
        app, merge_action(target_mode="changed", document_ids=["north", "south", "west"], merge_targets_resolved=True),
        max_documents_by_type=app.actions.get_document_action_max_documents_by_type("workflow", settings=limited),
    )
    with pytest.raises(app.merge.WorkflowMergeError) as caught:
        files.run(app, action, settings=limited)
    assert str(caught.value).startswith("3 files match this merge; at most 2 can be merged in one workflow run.")
    assert not files.resolve_calls and not files.published


def test_run_time_candidates_are_merged_in_file_name_order(app):
    candidates = app.merge.collect_merge_candidates
    documents = [
        {"id": "b", "file_name": "beta.CSV"}, {"id": "a", "file_name": "Alpha.xlsx"},
        {"id": "c", "file_name": "charlie.pdf"}, {"id": "b", "file_name": "beta.CSV"},
        {"document_id": "d", "file_name": "alpha.csv"}, {"id": "", "file_name": "nameless.csv"},
    ]
    assert candidates(documents, (".csv", ".xlsx"), 10) == ["d", "a", "b"]
    with pytest.raises(app.merge.WorkflowMergeError, match="More than 2 files match this merge"):
        candidates(documents, (".csv", ".xlsx"), 2)


def test_merge_failures_name_what_to_fix(app):
    files = world()
    with pytest.raises(app.merge.WorkflowMergeError) as engine:
        files.run(app, normalize(app, merge_action(
            document_ids=["north", "west"], merge_options={"schema_policy": "exact_order"},
        )))
    # The owner's own file name makes the failure actionable on the task.
    assert "west.csv" in str(engine.value)
    assert not files.published

    def failing(error):
        def render(source, request, max_output_bytes):
            raise error
        return render

    for error, message in (
        (app.contracts.GeneratedFileExportError("record_limit", "x"),
         "The merged table has more rows than one file can hold. Merge fewer files per run."),
        (app.contracts.GeneratedFileExportError("render_failed", "x"), "The merged file couldn't be created."),
        (app.office.OfficeRenderError("limit_exceeded"),
         "The merged table is too large for an Excel file (at most 1,048,575 rows, 5,000,000 cells and 32 MB). "
         "Choose CSV output instead."),
        (app.office.OfficeRenderError("render_failed"), "The merged Excel file couldn't be created."),
    ):
        with pytest.raises(app.merge.WorkflowMergeError) as caught:
            world().run(app, normalize(app, merge_action()), render=failing(error))
        assert str(caught.value) == message

    with pytest.raises(ValueError, match="unexpected"):
        world().run(app, normalize(app, merge_action()), render=failing(ValueError("unexpected")))


def test_a_withheld_kind_is_refused_by_the_contract_and_the_merge(app, monkeypatch):
    # Every kind is available from 0.261.223; a kind left out of the list must still be refused everywhere.
    withheld = tuple(kind for kind in app.actions.MERGE_KINDS if kind != "pptx")
    monkeypatch.setattr(app.actions, "MERGE_KINDS_AVAILABLE", withheld)
    monkeypatch.setattr(app.merge, "MERGE_KINDS_AVAILABLE", withheld)
    with pytest.raises(app.actions.MergeActionError, match="Merging PowerPoint files is not available yet."):
        normalize(app, merge_action(merge_kind="pptx"))
    with pytest.raises(app.merge.WorkflowMergeError, match="Merging PowerPoint files is not available yet."):
        world().run(app, {"type": "merge", "merge_kind": "pptx", "document_ids": ["a", "b"]})


def test_cancellation_stops_the_merge_before_anything_is_published(app):
    files = world()
    with pytest.raises(app.merge.WorkflowMergeCancelled):
        files.run(app, normalize(app, merge_action()), cancel=lambda: True)
    assert not files.resolve_calls and not files.reads

    files = world()
    state = {"cancelled": False}

    def progress(update):
        state["cancelled"] = True

    with pytest.raises(app.merge.WorkflowMergeCancelled):
        files.run(
            app, normalize(app, merge_action(document_ids=["north", "south", "west"])),
            cancel=lambda: state["cancelled"], progress=progress,
        )
    assert files.reads == ["north"] and not files.published


def test_the_reply_lists_what_was_left_out(app):
    report = {
        "totals": {"rows": 1500, "sources": 25, "columns": 14, "duplicates_removed": 1200},
        "sources": [
            {"file_name": f"file-{index:02}.csv", "status": "merged"} for index in range(22)
        ] + [
            {"file_name": "old.xlsx", "sheet": "Notes", "status": "excluded"},
            {"file_name": "empty.csv", "status": "skipped"},
        ],
        "nullable_columns": [f"Column {index}" for index in range(14)] + ["Sheet"],
        "sheet_column": "Sheet",
    }
    reply = app.merge.build_merge_reply(
        report, "all.csv", skipped_files=["a.pdf"], unavailable_files=["gone.csv"],
    )
    lines = reply.splitlines()
    assert lines[0] == "Merged **1,500 row(s)** from **25 file(s)** into **all.csv** with 14 column(s)."
    assert "- Removed 1,200 duplicate row(s)." in lines
    assert "- Left out 2 file(s) or sheet(s): old.xlsx (Notes), empty.csv." in lines
    assert "- Skipped 1 file(s) that aren't the right type: a.pdf." in lines
    assert "- Skipped 1 file(s) that aren't available right now: gone.csv." in lines
    blank = next(line for line in lines if line.startswith("- Columns some files don't have"))
    assert blank.endswith("Column 10, Column 11 and 2 more.") and "Sheet" not in blank
    assert lines[-1].startswith("Files merged, in order: file-00.csv,")
    assert lines[-1].endswith("file-19.csv and 2 more.")


# ---------------------------------------------------------------------------
# The runner's glue, run unchanged from its source
# ---------------------------------------------------------------------------

RUNNER = "functions_workflow_runner.py"
RUNNER_CONSTANTS = {
    "DOCUMENT_ACTION_TYPE_MERGE": "merge",
    "DOCUMENT_ACTION_TYPE_ANALYZE": "analyze",
    "DOCUMENT_ACTION_TYPE_NONE": "none",
    "DOCUMENT_ACTION_TARGET_MODE_RECENT": "recent",
    "DEFAULT_RECENT_DOCUMENT_WINDOW_MINUTES": 10,
    "MERGE_TARGET_MODE_CHANGED": "changed",
}


def test_file_sync_changes_point_only_changed_file_merges_at_the_changed_files():
    namespace = run_definitions(RUNNER, ("_apply_file_sync_changed_documents_to_action",), dict(RUNNER_CONSTANTS))
    apply = namespace["_apply_file_sync_changed_documents_to_action"]
    changed = {"type": "merge", "target_mode": "changed", "document_ids": [], "doc_scope": "personal"}
    assert apply(changed, ["d2", "d1"], ["g1"], ["p1"]) == {
        "type": "merge", "target_mode": "changed", "document_ids": ["d2", "d1"], "doc_scope": "all",
        "active_group_ids": ["g1"], "active_public_workspace_id": ["p1"], "merge_targets_resolved": True,
    }
    # With no changed files the merge still runs, and reports there was nothing to merge.
    assert apply(changed, [], [], [])["document_ids"] == []
    assert changed["document_ids"] == []
    assert apply({"type": "merge", "target_mode": "all"}, ["d1"], [], []) is None
    assert apply({"type": "analyze"}, [], [], []) == {"type": "none"}
    assert apply({"type": "analyze"}, ["d1"], [], [])["document_ids"] == ["d1"]


def test_cancellation_is_read_at_most_every_two_seconds_and_sticks():
    clock = {"now": 100.0}
    namespace = run_definitions(
        RUNNER, ("_throttled_cancellation_check",), {"time": SimpleNamespace(monotonic=lambda: clock["now"])},
    )
    answers = iter([False, True])
    reads = []

    def check():
        reads.append(clock["now"])
        return next(answers)

    cancelled = namespace["_throttled_cancellation_check"](check)
    assert cancelled() is False
    clock["now"] = 101.9
    assert cancelled() is False
    clock["now"] = 102.0
    assert cancelled() is True
    clock["now"] = 500.0
    assert cancelled() is True
    assert reads == [100.0, 102.0]


class QueryContainer:
    def __init__(self, documents):
        self.documents = documents
        self.queries = []

    def query_items(self, query, parameters, enable_cross_partition_query):
        assert enable_cross_partition_query is True
        self.queries.append((query, {item["name"]: item["value"] for item in parameters}))
        return iter(copy.deepcopy(self.documents))


def collector(app, *, roles=None, groups=("g1",), workspaces=("p1",)):
    containers = {
        "user": QueryContainer([{"id": "u1", "file_name": "a.csv"}]),
        "group": QueryContainer([{"id": "g1-doc", "file_name": "b.xlsx"}]),
        "public": QueryContainer([{"id": "p1-doc", "file_name": "c.csv"}]),
    }
    calls = []

    class FixedClock(datetime):
        @classmethod
        def now(cls, tz=None):
            return datetime(2026, 9, 28, 12, 0, tzinfo=timezone.utc)

    def assert_group_role(user_id, group_id, allowed_roles):
        calls.append(("assert_group_role", user_id, group_id, tuple(allowed_roles)))
        if (roles or {}).get(group_id) is None:
            raise PermissionError("Not a member of this group.")

    namespace = run_definitions(RUNNER, (
        "_get_workflow_group_id", "_coerce_workflow_recent_window_minutes", "_query_merge_candidate_documents",
        "_collect_merge_workflow_documents",
    ), {
        **RUNNER_CONSTANTS,
        "normalize_search_scope": app.search.normalize_search_scope,
        "normalize_search_id_list": app.search.normalize_search_id_list,
        "assert_group_role": assert_group_role,
        "_resolve_recent_authorized_group_ids": lambda user_id, ids: [value for value in ids if value in groups],
        "_resolve_recent_authorized_public_workspace_ids": (
            lambda user_id, ids: [value for value in ids if value in workspaces]
        ),
        "cosmos_user_documents_container": containers["user"],
        "cosmos_group_documents_container": containers["group"],
        "cosmos_public_documents_container": containers["public"],
        "select_current_documents": lambda documents: calls.append(("select", [d["id"] for d in documents])) or documents,
        "datetime": FixedClock,
        "timezone": timezone,
    })
    return namespace["_collect_merge_workflow_documents"], containers, calls


def test_run_time_discovery_reads_only_authorized_scopes_for_the_merge_file_types(app):
    collect, containers, calls = collector(app)
    workflow = {"user_id": OWNER}
    action = {"target_mode": "all", "doc_scope": "all", "active_group_ids": ["g1", "g9"],
              "active_public_workspace_id": ["p1", "p9"]}
    documents = collect(workflow, action, SETTINGS, (".csv", ".xlsx"))
    assert [document["id"] for document in documents] == ["u1", "g1-doc", "p1-doc"]
    assert ("select", ["u1", "g1-doc", "p1-doc"]) in calls
    [(query, parameters)] = containers["user"].queries
    assert query == (
        "SELECT * FROM c WHERE c.user_id = @user_id AND "
        "(ENDSWITH(LOWER(c.file_name), @merge_extension_0) OR ENDSWITH(LOWER(c.file_name), @merge_extension_1))"
    )
    assert parameters == {"@user_id": OWNER, "@merge_extension_0": ".csv", "@merge_extension_1": ".xlsx"}
    # Unauthorized group and public workspace ids are dropped before any query.
    assert [parameters["@group_id"] for _, parameters in containers["group"].queries] == ["g1"]
    assert [parameters["@workspace_id"] for _, parameters in containers["public"].queries] == ["p1"]

    collect, containers, _calls = collector(app)
    collect(workflow, {"target_mode": "recent", "doc_scope": "personal", "recent_window_minutes": 30}, SETTINGS, (".csv",))
    [(query, parameters)] = containers["user"].queries
    assert query.endswith(" AND c._ts >= @cutoff_ts")
    assert parameters["@cutoff_ts"] == int(datetime(2026, 9, 28, 11, 30, tzinfo=timezone.utc).timestamp())
    assert not containers["group"].queries and not containers["public"].queries


def test_a_group_workflow_merges_only_its_own_group_files(app):
    collect, containers, calls = collector(app, roles={"team": "User"})
    workflow = {"user_id": OWNER, "group_id": "team"}
    action = {"target_mode": "all", "doc_scope": "all", "active_group_ids": ["g1"], "active_public_workspace_id": ["p1"]}
    collect(workflow, action, SETTINGS, (".csv",))
    assert calls[0] == ("assert_group_role", OWNER, "team", ("Owner", "Admin", "DocumentManager", "User"))
    assert [parameters["@group_id"] for _, parameters in containers["group"].queries] == ["team"]
    assert not containers["user"].queries and not containers["public"].queries

    collect, containers, _calls = collector(app, roles={})
    with pytest.raises(PermissionError):
        collect(workflow, action, SETTINGS, (".csv",))
    assert not any(container.queries for container in containers.values())


class WorkflowRunCancelledError(BaseException):
    pass


class WorkflowInputError(ValueError):
    def __init__(self, public_message):
        self.public_message = public_message
        super().__init__(public_message)


def runner_merge(app, monkeypatch, files, *, cancelled=False):
    uploads = []
    thoughts = []

    def upload(user_id, conversation_id, file_name, stream, size, *, capability, output_format, summary,
               artifact_idempotency_key):
        uploads.append({
            "user_id": user_id, "conversation_id": conversation_id, "file_name": file_name,
            "content": stream.read(), "size": size, "capability": capability, "output_format": output_format,
            "summary": summary, "key": artifact_idempotency_key,
        })
        return {"message": {"id": "artifact-message-1", "file_name": file_name}}

    operations = ModuleType("functions_simplechat_operations")
    operations.upload_generated_file_artifact_stream_for_user = upload
    monkeypatch.setitem(sys.modules, "functions_simplechat_operations", operations)
    # The runner reads originals through the access boundary; here that boundary reads the test files.
    monkeypatch.setattr(app.orchestration_merge, "read_available_document_bytes", files.read)

    def thought(tracker, workflow, run_id, **fields):
        thoughts.append(fields)

    namespace = run_definitions(RUNNER, ("_throttled_cancellation_check", "_execute_document_merge_workflow"), {
        "time": SimpleNamespace(monotonic=lambda: 0.0),
        "_collect_merge_workflow_documents": lambda *args: (_ for _ in ()).throw(AssertionError("not collected")),
        "build_generated_file_export": app.exports.build_generated_file_export,
        "build_generated_file_artifact_metadata": app.exports.build_generated_file_artifact_metadata,
        "_add_workflow_activity_thought": thought,
        "_is_workflow_run_cancellation_requested": lambda workflow, run_id: cancelled,
        "resolve_analysis_source_manifest": files.resolve,
        "WorkflowRunCancelledError": WorkflowRunCancelledError,
        "WorkflowInputError": WorkflowInputError,
        "WORKFLOW_RUN_CANCELLED_MESSAGE": "Workflow cancellation was requested.",
    })
    workflow = {"id": "wf-1", "user_id": OWNER, "active_task": {"id": TASK_ID, "name": "Merge sales"}}
    run = namespace["_execute_document_merge_workflow"]
    return lambda action: run(workflow, action, SETTINGS, conversation_id=CONVERSATION, run_id=RUN_ID,
                              thought_tracker=object()), uploads, thoughts


def test_the_runner_attaches_the_merged_file_and_reports_each_file(app, monkeypatch):
    files = world()
    run, uploads, thoughts = runner_merge(app, monkeypatch, files)
    result = run(normalize(app, merge_action(output_file_name="Sales")))

    [upload] = uploads
    assert upload["file_name"] == "Sales.csv" and upload["size"] == len(upload["content"])
    assert (upload["capability"], upload["output_format"]) == ("file_merge", "csv")
    assert upload["key"] == f"workflow-merge:{RUN_ID}:{TASK_ID}:{hashlib.sha256(upload['content']).hexdigest()}"
    [artifact] = result["generated_analysis_artifacts"]
    assert artifact["artifact_message_id"] == "artifact-message-1"
    assert (artifact["capability"], artifact["file_name"], artifact["conversation_id"]) == (
        "file_merge", "Sales.csv", CONVERSATION,
    )
    assert artifact["row_count"] == 3
    assert result["generated_tabular_outputs"] == [artifact]
    assert [(entry["status"], entry["content"]) for entry in thoughts] == [
        ("running", "Merged file 1 of 2"), ("running", "Merged file 2 of 2"), ("completed", "Merged files"),
    ]
    assert {entry["activity_key"] for entry in thoughts} == {f"file-merge:{RUN_ID}:{TASK_ID}"}
    assert {(entry["kind"], entry["title"]) for entry in thoughts} == {("file_merge", "Merge files")}


def test_a_cancelled_run_raises_the_runner_cancellation(app, monkeypatch):
    files = world()
    run, uploads, thoughts = runner_merge(app, monkeypatch, files, cancelled=True)
    with pytest.raises(WorkflowRunCancelledError):
        run(normalize(app, merge_action()))
    assert not uploads and not thoughts and not files.reads


@pytest.mark.parametrize(("files", "action", "message"), [
    (world(unavailable=["north"]), merge_action(document_ids=["south", "north"]),
     "A file selected for this merge is no longer available to the workflow's owner."),
    (world(), merge_action(document_ids=["north", "notes"]),
     "notes.docx isn't a CSV or Excel file (.csv, .xlsx, .xlsm or .xls), so it can't be merged."),
    (world(), merge_action(document_ids=["north", "west"], merge_options={"schema_policy": "exact_order"}), "west.csv"),
])
def test_a_merge_failure_is_shown_on_the_task_and_not_retried(app, monkeypatch, files, action, message):
    # The task loop shows a WorkflowInputError's public message and stops retrying; any other
    # error becomes a generic message about models and is retried.
    run, uploads, thoughts = runner_merge(app, monkeypatch, files)
    with pytest.raises(WorkflowInputError) as caught:
        run(normalize(app, action))
    assert message in caught.value.public_message
    assert not uploads and not [entry for entry in thoughts if entry["status"] == "completed"]


# ---------------------------------------------------------------------------
# Saving a workflow
# ---------------------------------------------------------------------------

def test_merging_changed_files_needs_a_file_sync_trigger_that_passes_them_on():
    namespace = run_definitions("functions_personal_workflows.py", ("_require_changed_file_merge_trigger",), {
        "DOCUMENT_ACTION_TYPE_MERGE": "merge", "MERGE_TARGET_MODE_CHANGED": "changed",
        "WorkflowPublicValidationError": WorkflowInputError,
    })
    require = namespace["_require_changed_file_merge_trigger"]
    with pytest.raises(WorkflowInputError, match="needs a File Sync trigger that uses changed documents"):
        require({"type": "Merge", "target_mode": "CHANGED"}, False)
    require({"type": "merge", "target_mode": "changed"}, True)
    require({"type": "merge", "target_mode": "all"}, False)
    require({"type": "analyze", "target_mode": "changed"}, False)
    require(None, False)


def test_a_refused_merge_save_says_why_without_repeating_the_task_name(app):
    # Plain ValueErrors from a save become "Invalid workflow settings"; a merge's reviewed reasons
    # become public validation errors that name the task only by its position.
    namespace = run_definitions(
        "functions_personal_workflows.py",
        ("_normalize_text", "normalize_workflow_max_tasks", "_normalize_workflow_tasks", "_normalize_merge_errors_for_save"),
        {
            "uuid": __import__("uuid"), "WORKFLOW_TASK_LIMIT_DEFAULT": 50, "WORKFLOW_TASK_LIMIT_MIN": 1,
            "WORKFLOW_TASK_LIMIT_MAX": 100, "WORKFLOW_MAX_TASKS": 50, "WORKFLOW_TASK_INSTRUCTIONS_MAX_LENGTH": 12000,
            "WORKFLOW_TASK_NAME_MAX_LENGTH": 120, "WORKFLOW_TASK_RUNNER_TYPES": {"inherit", "agent", "model"},
            "MergeActionError": app.actions.MergeActionError, "WorkflowPublicValidationError": WorkflowInputError,
        },
    )
    save_errors = namespace["_normalize_merge_errors_for_save"]

    def normalizer(action):
        return save_errors(lambda: normalize(app, action))

    tasks = [{"id": "t1", "name": "Merge <b>secret</b> sales", "instructions": "Merge.",
              "document_action": merge_action(document_ids=["north"])}]
    with pytest.raises(WorkflowInputError) as refused:
        namespace["_normalize_workflow_tasks"]({"tasks": tasks}, task_document_action_normalizer=normalizer)
    assert refused.value.public_message == "Workflow task 1: Select at least two files to merge."

    # Option errors from the merge engine can quote column names, so they stay private.
    tasks[0]["document_action"] = merge_action(merge_options={"dedupe": "key_columns"})
    with pytest.raises(ValueError) as private:
        namespace["_normalize_workflow_tasks"]({"tasks": tasks}, task_document_action_normalizer=normalizer)
    assert not isinstance(private.value, WorkflowInputError)
    assert str(private.value).startswith("Workflow task 1 (Merge <b>secret</b> sales): ")

    for changes in ({"merge_kind": "zip"}, {"output_file_name": "a/b"}, {"merge_options": {"bookmarks": True}}):
        with pytest.raises(app.actions.MergeActionError):
            normalize(app, merge_action(**changes))
    with pytest.raises(app.actions.MergeActionError, match="File merging is turned off"):
        normalize(app, merge_action(), allowed_action_types={"none", "analyze"})


# ---------------------------------------------------------------------------
# Workflows proposed from chat
# ---------------------------------------------------------------------------

@pytest.fixture
def drafts():
    harness = DraftHarness()
    harness.documents.update({
        ("personal", OWNER_ID, "doc-north"): {"id": "doc-north", "user_id": OWNER_ID},
        ("personal", OWNER_ID, "doc-south"): {"id": "doc-south", "user_id": OWNER_ID},
    })
    return harness


MERGE_HANDLES = {"documents": {
    "north": {"document_id": "doc-north", "scope_type": "personal"},
    "south": {"document_id": "doc-south", "scope_type": "personal"},
}}


def merge_blueprint(merge, *, inputs=(), trigger=None, **task_changes):
    task = {"title": "Merge regional sales", "instructions": "Merge the regional sales files.", "merge": merge}
    if inputs:
        task["inputs"] = list(inputs)
    task.update(task_changes)
    return {
        "name": "Regional sales merge",
        "trigger": trigger or {"type": "calendar", "frequency": "weekly", "days_of_week": ["monday"],
                               "time_of_day": "08:00", "timezone": "America/New_York"},
        "tasks": [task],
        "alerts": {"mode": "failures_only"},
    }


def dry_run(harness, blueprint, handles=None):
    return harness.call(
        "dry_run_workflow_blueprint", OWNER_ID, copy.deepcopy(blueprint), copy.deepcopy(handles or {}),
        origin=ORIGIN, settings=harness.settings, user_info=USER_INFO,
    )


def test_a_proposed_merge_task_builds_a_workflow_merge_action(drafts):
    blueprint = merge_blueprint({
        "files": "inputs", "output_format": "xlsx", "file_name": "Regional sales",
        "options": {"schema_policy": "union", "column_aliases": [{"column": "Amount", "aliases": ["Total"]}]},
    }, inputs=["south", "north"])
    result = dry_run(drafts, blueprint, MERGE_HANDLES)
    assert result["ok"] is True, result["errors"]
    [task] = result["workflow"]["tasks"]
    action = task["document_action"]
    assert (action["type"], action["merge_kind"], action["target_mode"]) == ("merge", "tabular", "selected")
    assert action["document_ids"] == ["doc-south", "doc-north"]
    assert (action["doc_scope"], action["output_format"], action["output_file_name"]) == (
        "personal", "xlsx", "Regional sales",
    )
    # The stored options use the editor's column-to-names map.
    assert action["merge_options"] == {"schema_policy": "union", "column_aliases": {"Amount": ["Total"]}}
    # A merge's files are what it merges, never reference documents read by a model.
    assert task["reference_ids"] == [] and result["workflow"]["reference_inputs"] == []
    assert drafts.writes() == {}

    every = dry_run(drafts, merge_blueprint({"files": "recent", "recent_window_minutes": 120}))
    assert every["ok"] is True, every["errors"]
    action = every["workflow"]["tasks"][0]["document_action"]
    assert (action["target_mode"], action["doc_scope"], action["recent_window_minutes"]) == ("recent", "personal", 120)
    assert action["document_ids"] == []


def test_a_proposed_merge_of_changed_files_runs_on_each_sync(drafts):
    trigger = {
        "type": "file_sync", "source_ids": ["contracts"],
        "schedule": {"kind": "calendar", "frequency": "weekdays", "time_of_day": "07:00", "timezone": "Europe/London"},
    }
    handles = {"sources": {"contracts": {"scope_type": "personal", "source_id": "my-share"}}}
    result = dry_run(drafts, merge_blueprint({"files": "changed"}, trigger=trigger), handles)
    assert result["ok"] is True, result["errors"]
    workflow = result["workflow"]
    assert workflow["trigger_type"] == "file_sync" and workflow["file_sync"]["use_changed_documents"] is True
    action = workflow["tasks"][0]["document_action"]
    assert (action["type"], action["target_mode"], action["document_ids"]) == ("merge", "changed", [])


def _codes(result):
    return [(error["code"], error["path"]) for error in result["errors"]]


def test_proposed_merge_rules_are_repairable_and_never_echo_input(drafts):
    hostile = "<img src=x onerror=alert(1)>"
    one_input = merge_blueprint({"files": "inputs"}, inputs=["north"])
    assert _codes(dry_run(drafts, one_input, MERGE_HANDLES)) == [("merge_inputs_required", "/tasks/0/inputs")]

    changed = merge_blueprint({"files": "changed"})
    assert _codes(dry_run(drafts, changed)) == [("merge_trigger_required", "/tasks/0/merge/files")]

    bad_options = merge_blueprint({"files": "all", "options": {"schema_policy": "mapped", "dedupe_columns": [hostile]}})
    result = dry_run(drafts, bad_options)
    assert _codes(result) == [("merge_options_invalid", "/tasks/0/merge/options")]
    assert hostile not in str(result) and "onerror" not in str(result)

    repeated = merge_blueprint({"files": "all", "options": {"column_aliases": [
        {"column": "Amount", "aliases": ["Total"]}, {"column": "amount", "aliases": ["Sum"]},
    ]}})
    assert _codes(dry_run(drafts, repeated)) == [("merge_options_invalid", "/tasks/0/merge/options")]

    agent = merge_blueprint({"files": "all"}, runner={"type": "agent", "agent_ref": "mail_agent"})
    handles = {"agents": {"mail_agent": {"id": "agent-researcher", "name": "researcher", "is_global": False}}}
    assert ("merge_runner_invalid", "/tasks/0/runner") in _codes(dry_run(drafts, agent, handles))

    drafts.settings["document_action_capabilities"] = {"merge": {"enabled": False}}
    assert _codes(dry_run(drafts, merge_blueprint({"files": "all"}))) == [("merge_unavailable", "/tasks/0/merge")]


def test_the_blueprint_merge_schema_is_closed(drafts):
    schema = drafts.drafts.workflow_blueprint_schema()
    merge = schema["$defs"]["merge"]
    assert merge["additionalProperties"] is False and merge["required"] == ["files"]
    assert merge["properties"]["files"] == {"enum": ["inputs", "changed", "all", "recent"]}
    aliases = schema["$defs"]["merge_options"]["properties"]["column_aliases"]
    assert aliases["type"] == "array" and aliases["items"]["additionalProperties"] is False
    unknown = merge_blueprint({"files": "all", "destination": "https://example.com"})
    assert _codes(dry_run(drafts, unknown)) == [("unsupported_field", "/tasks/0/merge/destination")]


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
