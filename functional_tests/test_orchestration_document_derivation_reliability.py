# test_orchestration_document_derivation_reliability.py
"""Source-grounded schema correction and producer-aware document generation recovery.

Version: 0.261.316
Implemented in: 0.261.309

Literal source annotations reach Word publication without correction as of 0.261.314.
Markdown-escaped source punctuation reaches publication as of 0.261.316.

Real collectors, producers, checkpoints, orchestration and renderers run offline.
Malformed metadata must not erase uncertainty or rewrite accepted source findings.
"""

import importlib
import io
import json
from concurrent.futures import ThreadPoolExecutor

import pytest
from docx import Document

from test_document_analysis_final_results import run_analysis
from test_document_analysis_work_recovery import make_checkpoints, run_saved
from test_orchestration_deliverables import _claim_retry, _retry_request
from test_orchestration_harness_execution import harness, initialized_application  # noqa: F401
from test_orchestration_reason_render_pipeline import download_outputs, narrative_io  # noqa: F401
from test_support.document_analysis import (
    FixtureAnalysisClient,
    document_analysis_runtime,
    extract_fixture_findings,
    original_document,
)
from test_support.orchestration_harness_execution import compose_step, decoded_frames, input_binding, render_step
from test_support.orchestration_research import _definitions


QUALIFICATION = 'The period is not stated.'


def correction_reply(prompt):
    if '<ResponseShapeCorrection>' in prompt:
        raw = prompt.split('Previous response:\n', 1)[1].rsplit('\n</ResponseShapeCorrection>', 1)[0]
        payload = json.loads(raw)
        for container in [payload, *payload['findings']]:
            for field in ('issues', 'caveats', 'notes'):
                if isinstance(container.get(field), str):
                    container[field] = [container[field]]
        return payload
    return json.loads(extract_fixture_findings(prompt))


@pytest.mark.parametrize('field', ['caveats', 'issues', 'notes'])
def test_malformed_text_metadata_is_corrected_without_erasing_uncertainty(field, monkeypatch):
    documents = {'one': original_document('one', ['The service owner is Mira.'])}

    def reply(prompt):
        payload = correction_reply(prompt)
        if '<ResponseShapeCorrection>' not in prompt:
            target = payload if field == 'notes' else payload['findings'][0]
            target[field] = QUALIFICATION
        return json.dumps(payload)

    with document_analysis_runtime(documents) as runtime:
        monkeypatch.setattr(runtime.producer, '_wait_analysis_retry', lambda *args: None)
        result, client = run_analysis(
            runtime, documents, FixtureAnalysisClient(reply), max_retries_per_window=1,
        )
    assert len(client.calls) == 2
    assert all(call['metadata']['json_output'] for call in client.calls)
    assert client.calls[1]['metadata']['analysis_response_correction']
    assert result['analysis_validation']['coverage']['completed_work_units'] == 1
    assert result['analysis_validation']['status'] == ('partial' if field == 'issues' else 'valid')
    if field == 'caveats':
        assert result['authoritative_result']['value'][0]['caveats'] == [QUALIFICATION]
    elif field == 'issues':
        assert result['authoritative_result']['value'] == []
        assert QUALIFICATION in json.dumps(result['analysis_diagnostics'])
    else:
        assert QUALIFICATION in result['analysis_reply']


@pytest.mark.parametrize('change', ['values', 'evidence', 'status', 'drop_finding', 'drop_caveat', 'extra_call'])
def test_unusable_correction_retains_original_partial_result_and_is_bounded(change, monkeypatch):
    documents = {'one': original_document('one', ['The service owner is Mira.'])}

    def reply(prompt):
        payload = correction_reply(prompt)
        if '<ResponseShapeCorrection>' not in prompt or change == 'extra_call':
            payload['findings'][0]['caveats'] = QUALIFICATION
        elif change == 'drop_finding':
            payload['findings'] = []
        elif change == 'drop_caveat':
            payload['findings'][0]['caveats'] = []
        elif change == 'values':
            payload['findings'][0]['values']['owner'] = 'Invented owner'
        elif change == 'evidence':
            payload['findings'][0]['evidence'][0]['quote'] = 'Invented quote'
        else:
            payload['findings'][0]['status'] = 'unresolved'
        return json.dumps(payload)

    with document_analysis_runtime(documents) as runtime:
        monkeypatch.setattr(runtime.producer, '_wait_analysis_retry', lambda *args: None)
        result, client = run_analysis(
            runtime, documents, FixtureAnalysisClient(reply), max_retries_per_window=1,
        )
    assert len(client.calls) == 2
    assert result['coverage']['retries'] == 1
    assert result['analysis_validation']['status'] == 'partial'
    assert result['analysis_validation']['coverage']['completed_work_units'] == 1
    assert result['analysis_validation']['issues'][0]['code'] == 'invalid_caveats'
    assert result['analysis_diagnostics']['candidates'][0]['values'] == {'owner': 'Mira'}


@pytest.mark.parametrize('kind', ['chat', 'workflow', 'orchestration'])
def test_retry_repairs_only_malformed_saved_window_without_reextracting_good_windows(kind, monkeypatch):
    documents = {'one': original_document('one', [
        'The service owner is Mira.', 'Reviews take place every quarter.',
    ])}

    def first_reply(prompt):
        payload = json.loads(extract_fixture_findings(prompt))
        if '[Page 1, Chunk 1]' in prompt:
            payload['findings'][0]['caveats'] = QUALIFICATION
        return json.dumps(payload)

    with document_analysis_runtime(documents) as runtime:
        monkeypatch.setattr(runtime.producer, '_wait_analysis_retry', lambda *args: None)
        first = make_checkpoints(documents, kind=kind)
        initial, original_client = run_saved(
            runtime, documents, first, client=FixtureAnalysisClient(first_reply),
        )
        assert initial['analysis_validation']['status'] == 'partial'
        second = make_checkpoints(
            documents, kind=kind, store=first.store, attempt='assistant-2', previous='assistant-1',
        )
        result, repair_client = run_saved(
            runtime, documents, second,
            client=FixtureAnalysisClient(lambda prompt: json.dumps(correction_reply(prompt))),
            max_retries_per_window=1,
        )
    assert len(original_client.calls) == 2
    assert len(repair_client.calls) == 1
    assert '<ResponseShapeCorrection>' in repair_client.calls[0]['prompt']
    assert result['analysis_validation']['status'] == 'valid'
    assert result['authoritative_result']['value'][0]['values'] == {
        'owner': 'Mira', 'review_frequency': 'quarterly',
    }
    assert result['authoritative_result']['value'][0]['caveats'] == [QUALIFICATION]
    assert result['analysis_metrics']['recovery']['reused_windows'] == 1


@pytest.mark.parametrize('supported', [False, True])
def test_unknown_support_status_cannot_be_promoted_by_schema_correction(supported, monkeypatch):
    documents = {'one': original_document('one', ['The service owner is Mira.'])}

    def reply(prompt):
        payload = correction_reply(prompt)
        payload['findings'][0]['status'] = (
            'supported' if supported else 'unresolved'
        ) if '<ResponseShapeCorrection>' in prompt else 'unknown'
        return json.dumps(payload)

    with document_analysis_runtime(documents) as runtime:
        monkeypatch.setattr(runtime.producer, '_wait_analysis_retry', lambda *args: None)
        result, client = run_analysis(
            runtime, documents, FixtureAnalysisClient(reply), max_retries_per_window=1,
        )
    assert len(client.calls) == 2
    assert result['analysis_validation']['status'] == 'partial'
    assert result['authoritative_result']['value'] == []
    assert result['analysis_validation']['issues'][0]['code'] == (
        'invalid_status' if supported else 'unresolved_finding'
    )


@pytest.mark.parametrize('malformed', ['non_object', 'values'])
def test_malformed_findings_preserve_other_accepted_candidates_without_inventing_data(malformed, monkeypatch):
    documents = {'one': original_document('one', ['The service owner is Mira.'])}

    def reply(prompt):
        payload = correction_reply(prompt)
        if '<ResponseShapeCorrection>' not in prompt:
            payload['findings'].append(
                'Uninterpreted source content' if malformed == 'non_object' else {
                    **payload['findings'][0], 'finding_key': 'other', 'values': 'Unknown',
                }
            )
        else:
            payload['findings'][1] = {
                **payload['findings'][0], 'finding_key': 'other', 'values': {'owner': 'Invented'},
            }
        return json.dumps(payload)

    with document_analysis_runtime(documents) as runtime:
        monkeypatch.setattr(runtime.producer, '_wait_analysis_retry', lambda *args: None)
        result, client = run_analysis(
            runtime, documents, FixtureAnalysisClient(reply), max_retries_per_window=1,
        )
    assert len(client.calls) == 2
    assert result['analysis_validation']['status'] == 'partial'
    assert result['analysis_validation']['coverage']['completed_work_units'] == 1
    assert result['authoritative_result']['value'][0]['values'] == {'owner': 'Mira'}
    assert 'Invented' not in json.dumps(result['authoritative_result'])


def test_concurrent_windows_use_bounded_correction_and_safe_correlated_events(monkeypatch):
    documents = {'one': original_document('one', [
        'The service owner is Mira.', 'Reviews take place every quarter.',
    ])}
    events = []
    calls = []

    def factory(metadata):
        def invoke(prompt, **kwargs):
            calls.append(kwargs['metadata'])
            payload = correction_reply(prompt)
            if '<ResponseShapeCorrection>' not in prompt:
                payload['findings'][0]['caveats'] = QUALIFICATION
            return json.dumps(payload)
        return invoke

    with document_analysis_runtime(documents) as runtime, ThreadPoolExecutor(max_workers=2) as executor:
        monkeypatch.setattr(runtime.producer, '_wait_analysis_retry', lambda *args: None)
        monkeypatch.setattr(runtime.producer, 'log_event', lambda message, **kwargs: events.append((message, kwargs)))
        monkeypatch.setattr(
            runtime.producer, 'workflow_log_context',
            _definitions('functions_appinsights.py', names={'workflow_log_context'})['workflow_log_context'],
        )
        result, _ = run_analysis(
            runtime, documents, max_retries_per_window=1, max_window_concurrency=2,
            invoke_prompt_factory=factory, executor=executor, conversation_id='private-conversation',
        )
    assert result['analysis_validation']['status'] == 'valid'
    assert len(calls) == 4
    assert all(call['json_output'] for call in calls)
    assert sum(bool(call.get('analysis_response_correction')) for call in calls) == 2
    correction_events = [kwargs['extra'] for message, kwargs in events if 'Response correction finished' in message]
    assert len(correction_events) == 2
    assert all(event['status'] == 'corrected' and event['conversation_id_hash'] for event in correction_events)
    assert QUALIFICATION not in json.dumps(correction_events)
    assert 'private-conversation' not in json.dumps(correction_events)


@pytest.mark.parametrize('interruption', ['cancel', 'access'])
def test_correction_does_not_publish_after_cancellation_or_access_loss(interruption, monkeypatch):
    documents = {'one': original_document('one', ['The service owner is Mira.'])}
    cancelled = [False]

    def reply(prompt):
        payload = correction_reply(prompt)
        if '<ResponseShapeCorrection>' not in prompt:
            payload['findings'][0]['caveats'] = QUALIFICATION
        elif interruption == 'access':
            raise PermissionError('Current source access is unavailable.')
        else:
            cancelled[0] = True
        return json.dumps(payload)

    with document_analysis_runtime(documents) as runtime:
        monkeypatch.setattr(runtime.producer, '_wait_analysis_retry', lambda *args: None)
        error = PermissionError if interruption == 'access' else runtime.cancellation.MixedSourceCancellationError
        with pytest.raises(error):
            run_analysis(
                runtime, documents, FixtureAnalysisClient(reply), max_retries_per_window=1,
                cancel_requested=lambda: cancelled[0],
            )


def analysis_step():
    return {
        'step_id': 'analyze', 'capability_id': 'document_analyze',
        'arguments': {
            'document_ids': ['document-1'], 'doc_scope': 'personal',
            'analysis_prompt': 'Read every source page and retain its findings.',
        },
        'outputs': [{'name': 'findings', 'kind': 'records-v1'}, {'name': 'coverage', 'kind': 'structured-v1'}],
    }


def document_steps(output_format):
    draft = compose_step(inputs={
        'findings': {'binding': input_binding('analyze', 'findings'), 'allow_partial': False},
    })
    profile = 'prepared_report_v1' if output_format == 'docx' else 'prepared_text_v1'
    return [analysis_step(), draft, render_step('report', output_format, profile=profile)]


@pytest.mark.parametrize('output_format', ['md', 'docx'])
@pytest.mark.parametrize('search_sources', [False, True])
def test_repaired_findings_reach_actual_composition_and_durable_document(
    harness, narrative_io, output_format, search_sources, monkeypatch,
):
    def reply():
        prompt = harness.model_calls[-1]['messages'][-1]['content']
        if '<DocumentSlice>' in prompt:
            payload = json.loads(narrative_io['reply']())
            if '<ResponseShapeCorrection>' in prompt:
                raw = prompt.split('Previous response:\n', 1)[1].rsplit('\n</ResponseShapeCorrection>', 1)[0]
                payload = json.loads(raw)
                payload['findings'][0]['caveats'] = [QUALIFICATION]
            elif '[Page 1, Chunk 1]' in prompt:
                payload['findings'][0]['caveats'] = QUALIFICATION
            return json.dumps(payload)
        return '# Derived document\n\nComplete finding 006. The period is not stated.'

    steps = document_steps(output_format)
    if search_sources:
        search = importlib.import_module('functions_search')
        monkeypatch.setattr(search, 'hybrid_search', lambda *args, **kwargs: [{
            'document_id': 'document-1', 'chunk_id': 'document-1-chunk-1',
            'chunk_sequence': 1, 'page_number': 1,
            'chunk_text': 'Item 000 contains original source detail.', 'file_name': 'document-1.txt',
        }])
        steps[0]['arguments'].pop('document_ids')
        steps[0]['inputs'] = {'sources': {'binding': input_binding('search', 'sources')}}
        steps.insert(0, {
            'step_id': 'search', 'capability_id': 'document_search',
            'arguments': {'query': 'Find the source report.', 'doc_scope': 'personal'},
        })
    harness.create(
        steps, replies=[reply] * 10,
        seeds={'document_ids': ['document-1'], 'doc_scope': 'personal'},
        original_seeds={'document_ids': ['document-1'], 'doc_scope': 'personal'},
    )
    execution = harness.prepare()
    frames = execution.execute()
    done = decoded_frames(frames)[-1]
    assert done['status'] == 'completed', done
    outputs, artifacts, payloads = download_outputs(harness)
    assert len(outputs) == len(artifacts) == 1
    content = payloads[f'report.{output_format}']
    if output_format == 'docx':
        text = '\n'.join(paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs)
    else:
        text = content.decode('utf-8')
    assert 'Complete finding 006.' in text and QUALIFICATION in text
    calls = harness.model_calls
    assert all(call['response_format'] == {'type': 'json_object'} for call in calls[:-1])
    assert len([call for call in calls if '<ResponseShapeCorrection>' in str(call['messages'])]) == 1


@pytest.mark.parametrize('malformed', [True, False])
def test_recovery_revisits_malformed_producer_but_does_not_offer_pointless_uncertainty_retry(
    harness, narrative_io, malformed,
):
    def reply():
        payload = json.loads(narrative_io['reply']())
        if malformed:
            payload['findings'][0]['caveats'] = QUALIFICATION
        else:
            payload['findings'][0]['issues'] = ['The source has contradictory requested values.']
        return json.dumps(payload)

    harness.create(
        document_steps('md'), replies=[reply] * 10,
        seeds={'document_ids': ['document-1'], 'doc_scope': 'personal'},
        original_seeds={'document_ids': ['document-1'], 'doc_scope': 'personal'},
    )
    execution = harness.prepare()
    frames = execution.execute()
    done = decoded_frames(frames)[-1]
    assert done['status'] != 'completed'
    assert harness.blobs.file_uploads == 0
    projection = harness.recovery.recovery_projection(harness.read())
    if malformed:
        assert projection['eligible'], projection
        assert 'analyze' in projection['retry_step_ids']
        assert 'analyze' not in projection['reused_step_ids']
        child = _retry_request(harness, execution, 'repair-malformed-producer')

        def repaired_reply():
            prompt = harness.model_calls[-1]['messages'][-1]['content']
            if '<ResponseShapeCorrection>' in prompt:
                return json.dumps(correction_reply(prompt))
            return '# Recovered document\n\nComplete finding 006.'

        previous_calls = len(harness.model_calls)
        harness.replies = [repaired_reply] * 10
        resumed = _claim_retry(harness, child)
        resumed_frames = resumed.execute()
        resumed_done = decoded_frames(resumed_frames)[-1]
        assert resumed_done['status'] == 'completed', resumed_done
        new_calls = harness.model_calls[previous_calls:]
        assert len(new_calls) == 3
        assert all('<ResponseShapeCorrection>' in str(call['messages']) for call in new_calls[:-1])
        assert harness.blobs.file_uploads == 1
    else:
        assert not projection['eligible'], projection
        assert projection['reason_code'] == 'input_partial_not_recoverable'


@pytest.mark.parametrize('annotation', [
    '<!-- PageFooter: Example report - Test Data Only -->',
    '<!-- This report contains test data. -->',
    '6\\) TELEPHONE NUMBER: ( 555 ) 123-4567',
    '\u2612\n2\\. Partnership',
])
def test_source_annotation_or_escape_reaches_composition_and_published_word_without_repairs(
    harness, narrative_io, monkeypatch, annotation,
):
    search = importlib.import_module('functions_search_service')
    original_read = search.get_ordered_document_chunks

    def read_chunks(*args, **kwargs):
        chunks = original_read(*args, **kwargs)
        chunks[0]['chunk_text'] += f'\n{annotation}'
        return chunks

    monkeypatch.setattr(search, 'get_ordered_document_chunks', read_chunks)

    def reply():
        prompt = harness.model_calls[-1]['messages'][-1]['content']
        if '<DocumentSlice>' in prompt:
            payload = json.loads(narrative_io['reply']())
            if '[Page 1, Chunk 1]' in prompt:
                quote = annotation.replace('\\)', ')').replace('\\.', '.')
                payload['findings'][0]['evidence'].append({'chunk_sequence': 1, 'quote': quote})
            return json.dumps(payload)
        report = '# Source-backed report\n\nComplete finding 006. This report contains test data.'
        if '\\' in annotation:
            rendered = annotation.replace('\\)', ')').replace('\\.', '.').replace('\n', ' ')
            report += f'\n\nEvidence: {rendered}'
        return report

    harness.create(
        document_steps('docx'), replies=[reply] * 3,
        seeds={'document_ids': ['document-1'], 'doc_scope': 'personal'},
        original_seeds={'document_ids': ['document-1'], 'doc_scope': 'personal'},
    )
    execution = harness.prepare()
    frames = execution.execute()
    done = decoded_frames(frames)[-1]
    assert done['status'] == 'completed', done
    analysis = execution.context.task_results['analyze']
    assert analysis.status == 'complete'
    assert analysis.output('findings').item_count == 7
    outputs, artifacts, payloads = download_outputs(harness)
    assert len(outputs) == len(artifacts) == harness.blobs.file_uploads == 1
    text = '\n'.join(paragraph.text for paragraph in Document(io.BytesIO(payloads['report.docx'])).paragraphs)
    assert 'Complete finding 006.' in text and 'This report contains test data.' in text
    if '\\' in annotation:
        assert annotation.replace('\\)', ')').replace('\\.', '.').replace('\n', ' ') in text
    calls = harness.model_calls
    assert len(calls) == 3
    assert sum('<DocumentSlice>' in str(call['messages']) for call in calls) == 2
    assert all('<ResponseShapeCorrection>' not in str(call['messages']) for call in calls)
