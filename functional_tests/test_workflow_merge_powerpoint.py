#!/usr/bin/env python3
# test_workflow_merge_powerpoint.py
"""
Functional test for PowerPoint Merge tasks in workflows.
Version: 0.261.223
Implemented in: 0.261.223
Refs: microsoft/simplechat#1619

This test ensures that a workflow Merge task can append the slides of PowerPoint decks, in
order, into one deck with a section per deck. The task contract accepts only PowerPoint
options and the PPTX format, files of the wrong type are refused when selected and skipped
when found at run time, a damaged deck fails the task with its name, a repeated merge
reuses its file, the runner attaches the deck and reports slides as it goes, and chat can
propose a PowerPoint merge. No model, Azure service or network is used.
"""

import io
import re
import sys
import zipfile
from pathlib import Path

import pytest

sys.path.append(str(Path(__file__).resolve().parent))

from test_support.versioning import assert_app_version_at_least  # noqa: E402
from test_workflow_merge_task import (  # noqa: E402,F401
    MERGE_HANDLES,
    NORTH,
    MergeWorld,
    WorkflowInputError,
    _codes,
    app,
    drafts,
    dry_run,
    merge_blueprint,
    normalize,
    runner_merge,
)


def deck_bytes(label, slides):
    from pptx import Presentation

    presentation = Presentation()
    for number in range(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1 if number else 0])
        slide.shapes.title.text = f"{label} slide {number + 1}"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


FILES = {
    "q1": ("Q1 review.pptx", deck_bytes("Q1", 2)),
    "q2": ("Q2 review.pptx", deck_bytes("Q2", 3)),
    "broken": ("broken.pptx", b"PK\x03\x04broken"),
    "north": NORTH,
    "notes": ("notes.docx", b"PK"),
}


def world(**kwargs):
    return MergeWorld(dict(FILES), **kwargs)


def deck_action(**changes):
    action = {"type": "merge", "merge_kind": "pptx", "document_ids": ["q1", "q2"]}
    action.update(changes)
    return action


def titles(content):
    from pptx import Presentation

    return [slide.shapes.title.text for slide in Presentation(io.BytesIO(content)).slides]


def section_names(content):
    with zipfile.ZipFile(io.BytesIO(content)) as package:
        return re.findall(r'section name="([^"]+)"', package.read("ppt/presentation.xml").decode())


def test_version_includes_powerpoint_merge_tasks():
    assert_app_version_at_least("0.261.223")


def test_a_powerpoint_task_takes_only_powerpoint_options_and_creates_pptx(app):
    action = normalize(app, deck_action(merge_options={"formatting": "use_first", "sections": False}))
    assert (action["merge_kind"], action["output_format"]) == ("pptx", "pptx")
    assert action["merge_options"] == {"formatting": "use_first", "sections": False}
    for payload, message in (
        (deck_action(output_format="pdf"), "A PowerPoint merge creates PPTX files."),
        (deck_action(merge_options={"bookmarks": True}), "Some merge options don't apply to PowerPoint merges. Remove them."),
        (deck_action(merge_options={"sections": "yes"}), "The sections merge option must be true or false."),
    ):
        with pytest.raises(app.actions.MergeActionError) as caught:
            normalize(app, payload)
        assert message in str(caught.value)


def test_a_powerpoint_merge_appends_slides_with_a_section_per_deck(app):
    files = world()
    updates = []
    result = files.run(app, normalize(app, deck_action(output_file_name="Reviews")), progress=updates.append)

    [published] = files.published
    assert (published["file_name"], published["output_format"]) == ("Reviews.pptx", "pptx")
    assert published["summary"] == "Merged 2 file(s) into one PowerPoint file."
    assert titles(published["content"]) == [
        "Q1 slide 1", "Q1 slide 2", "Q2 slide 1", "Q2 slide 2", "Q2 slide 3",
    ]
    assert section_names(published["content"]) == ["Q1 review", "Q2 review"]
    assert files.reads == ["q1", "q2"]
    assert [(update["index"], update["slides"]) for update in updates] == [(1, 2), (2, 3)]
    assert result["merge_summary"] == {
        "status": "merged", "kind": "pptx", "files": 2, "pages": 0, "slides": 5, "sheets": 0, "skipped": 0,
        "file_name": "Reviews.pptx",
    }
    assert result["reply"].startswith("Merged **2 file(s)** into **Reviews.pptx** with 5 slide(s).")


def test_files_found_at_run_time_skip_the_wrong_type(app):
    files = world(collected=["q1", "north", "notes", "q2"])
    result = files.run(app, normalize(app, deck_action(target_mode="all", doc_scope="personal")))
    [(_action, extensions, _limit)] = files.collect_calls
    assert extensions == (".pptx",)
    assert files.reads == ["q1", "q2"]
    assert result["merge_summary"]["files"] == 2 and result["merge_summary"]["skipped"] == 2
    assert "- Skipped 2 file(s) that aren't the right type: north.csv, notes.docx." in result["reply"]


def test_failures_name_the_deck_and_a_repeated_merge_reuses_its_file(app):
    with pytest.raises(app.merge.WorkflowMergeError) as wrong_type:
        world().run(app, normalize(app, deck_action(document_ids=["q1", "notes"])))
    assert str(wrong_type.value) == "notes.docx isn't a PowerPoint deck (.pptx), so it can't be merged."

    files = world()
    with pytest.raises(app.merge.WorkflowMergeError) as broken:
        files.run(app, normalize(app, deck_action(document_ids=["q1", "broken"])))
    assert str(broken.value) == "broken.pptx isn't a valid PowerPoint presentation."
    assert not files.published

    first, second = world(), world()
    first.run(app, normalize(app, deck_action()))
    second.run(app, normalize(app, deck_action()))
    assert first.published[0]["content"] == second.published[0]["content"]
    assert first.published[0]["idempotency_key"] == second.published[0]["idempotency_key"]


def test_the_runner_attaches_the_deck_and_reports_slides(app, monkeypatch):
    run, uploads, thoughts = runner_merge(app, monkeypatch, world())
    result = run(normalize(app, deck_action()))
    [upload] = uploads
    assert (upload["file_name"], upload["output_format"], upload["capability"]) == ("merged.pptx", "pptx", "file_merge")
    assert upload["content"].startswith(b"PK")
    [artifact] = result["generated_analysis_artifacts"]
    assert (artifact["file_name"], artifact["output_format"]) == ("merged.pptx", "pptx")
    assert [(entry["content"], entry["detail"]) for entry in thoughts] == [
        ("Merged file 1 of 2", "slides=2"), ("Merged file 2 of 2", "slides=3"), ("Merged files", "files=2"),
    ]

    run, uploads, _thoughts = runner_merge(app, monkeypatch, world())
    with pytest.raises(WorkflowInputError) as refused:
        run(normalize(app, deck_action(document_ids=["q1", "broken"])))
    assert refused.value.public_message == "broken.pptx isn't a valid PowerPoint presentation."
    assert not uploads


def test_chat_can_propose_a_powerpoint_merge_task(drafts):
    result = dry_run(drafts, merge_blueprint(
        {"kind": "pptx", "files": "inputs", "options": {"sections": False}}, inputs=["south", "north"],
    ), MERGE_HANDLES)
    assert result["ok"] is True, result["errors"]
    action = result["workflow"]["tasks"][0]["document_action"]
    assert (action["merge_kind"], action["output_format"], action["merge_options"]) == (
        "pptx", "pptx", {"sections": False},
    )

    wrong_option = dry_run(drafts, merge_blueprint({"kind": "pptx", "files": "all", "options": {"bookmarks": True}}))
    assert _codes(wrong_option) == [("merge_options_invalid", "/tasks/0/merge/options")]
    assert "pptx takes formatting, sections" in wrong_option["errors"][0]["message"]
    wrong_format = merge_blueprint({"kind": "pptx", "files": "all", "output_format": "docx"})
    assert _codes(dry_run(drafts, wrong_format)) == [("merge_format_invalid", "/tasks/0/merge/output_format")]


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
