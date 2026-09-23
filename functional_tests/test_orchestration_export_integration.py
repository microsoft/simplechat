# test_orchestration_export_integration.py
"""Full retained-result, private transport, and downloaded-byte export matrix.

Version: 0.261.127
Implemented in: 0.261.127

Exercises every shared format through the durable output service, and two full
30,000-record exports after restart. External storage is doubled. These tests
cover the result-to-file boundary, not planner or scheduler activation.
"""

import csv
import hashlib
import io
import json
from dataclasses import replace

import fitz
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from functions_orchestration_result_contracts import RecordColumn
from functions_orchestration_results import NamedOutput
from test_generated_file_office_bridge import deck_value
from test_generated_file_structured_renderers import readback
from test_orchestration_output_lifecycle import lifecycle, production_modules  # noqa: F401
from test_support.orchestration_results import COLUMNS, ROWS, complete


def retain_prepared(lifecycle, outputs):
    producer = replace(
        lifecycle.results.producer, step_id="prepared_content",
        capability_id="compose", contract_version="compose-v1",
    )
    lifecycle.results.add_producer(producer)
    task = lifecycle.results.save(producer=producer, grounded=False, outputs=outputs)
    run = lifecycle.runs.read_item("run-1", "conversation-1")
    run["plan"]["steps"].append({
        "step_id": producer.step_id, "capability_id": producer.capability_id, "enabled": True,
    })
    lifecycle.runs.upsert_item(run)
    return task


def test_all_ten_formats_reopen_after_private_commit_and_service_restart(lifecycle):
    markdown = (
        "# Prepared findings\n\nFirst retained record: **001**.\n\n"
        "| id | amount |\n| --- | --- |\n| 001 | 0 |\n| last | missing |\n\n"
        "- Complete caf\u00e9 finding.\n- Last retained record: **last**.\n"
    )
    text = "First retained record: 001.\nLast retained record: last.\ncaf\u00e9\n"
    task = retain_prepared(lifecycle, [
        NamedOutput("records", "records-v1", ROWS, complete(len(ROWS)), COLUMNS),
        NamedOutput("report", "markdown-v1", markdown, complete(1)),
        NamedOutput("text", "text-v1", text, complete(1)),
        NamedOutput("slides", "structured-v1", deck_value(), complete(1)),
    ])
    writes = lifecycle.results.container.sequence
    columns = ("id", "amount", "enabled")
    requests = [
        ("csv", "tabular_records_v1", "records", {"columns": columns}),
        ("xlsx", "tabular_workbook_v1", "records", {"columns": columns, "sheet_name": "Findings"}),
        ("json", "exact_records_v1", "records", {}),
        ("xml", "typed_xml_v1", "records", {}),
        ("yaml", "structured_records_v1", "records", {}),
        ("md", "prepared_text_v1", "report", {}),
        ("txt", "prepared_text_v1", "text", {}),
        ("docx", "prepared_report_v1", "report", {"title": "Approved findings"}),
        ("pdf", "prepared_report_v1", "report", {"title": "Approved findings"}),
        ("pptx", "prepared_slide_deck_v1", "slides", {}),
    ]
    payloads = {}
    for output_format, profile, name, options in requests:
        lifecycle.restart()
        requested = lifecycle.prepare(
            output_format, profile=profile, reference=task.output(name), **options,
        )
        completed = lifecycle.run(requested)
        assert completed["state"] == "completed", json.dumps(completed, sort_keys=True)
        lifecycle.restart()
        payload = lifecycle.download(completed)
        descriptor = lifecycle.raw(completed)["committed_intent"]
        repeated = lifecycle.run(completed)
        assert descriptor["size_bytes"] == completed["size_bytes"] == len(payload)
        assert descriptor["content_sha256"] == hashlib.sha256(payload).hexdigest()
        assert repeated == completed
        payloads[output_format] = payload

    for output_format in ("json", "xml", "yaml"):
        restored = readback(payloads[output_format], output_format, records=True)
        assert json.dumps(restored, sort_keys=True) == json.dumps(ROWS, sort_keys=True)
    csv_rows = list(csv.reader(io.StringIO(payloads["csv"].decode("utf-8"), newline="")))
    assert csv_rows == [
        list(columns), ["001", "0", "false"], ["'=SUM(A1:A2)", "1.25", "true"], ["last", "", "false"],
    ]
    with io.BytesIO(payloads["xlsx"]) as content:
        workbook = load_workbook(content, read_only=True, data_only=False)
        try:
            sheet_names = workbook.sheetnames
            workbook_rows = list(workbook["Findings"].values)
            formula_types = [cell.data_type for row in workbook["Findings"] for cell in row]
        finally:
            workbook.close()
    assert sheet_names == ["Findings"]
    assert workbook_rows == [columns, *[tuple(row[name] for name in columns) for row in ROWS]]
    assert "f" not in formula_types
    assert payloads["md"] == markdown.encode("utf-8")
    assert payloads["txt"] == text.encode("utf-8")

    document = Document(io.BytesIO(payloads["docx"]))
    paragraphs = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "First retained record: 001." in paragraphs and "Last retained record: last." in paragraphs
    assert document.core_properties.title == "Approved findings"
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows] == [
        ["id", "amount"], ["001", "0"], ["last", "missing"],
    ]
    with fitz.open(stream=payloads["pdf"], filetype="pdf") as report:
        pdf_text = "".join(page.get_text() for page in report)
        pdf_title = report.metadata["title"]
    assert "First retained record: 001." in pdf_text and "Last retained record: last." in pdf_text
    assert "caf\u00e9" in pdf_text and pdf_title == "Approved findings"

    deck = Presentation(io.BytesIO(payloads["pptx"]))
    assert len(deck.slides) == 2
    assert deck.slides[0].shapes.title.text == "First prepared slide"
    assert deck.slides[-1].shapes[0].text == "Last prepared slide"
    assert deck.slides[0].notes_slide.notes_text_frame.text == "Prepared notes.\nSecond line."
    tables = [shape.table for slide in deck.slides for shape in slide.shapes if shape.has_table]
    assert [[cell.text for cell in row.cells] for row in tables[0].rows] == [
        ["Name", "Count"], ["First", "001"],
    ]

    public = lifecycle.service.list_public_outputs("run-1")
    artifacts = lifecycle.service.committed_artifacts("run-1")
    assert len(public) == len(artifacts) == 10
    assert all(output["state"] == "completed" and output["available"] for output in public)
    assert lifecycle.blobs.uploads == len(lifecycle.render_calls) == 10
    assert lifecycle.results.container.sequence == writes
    encoded = json.dumps(public)
    assert all(private not in encoded for private in ("blob_path", "source_ref", "content_sha256", "lease"))


def test_large_results_export_twice_after_restart_without_rewriting_producer(lifecycle):
    count = 30_000
    detail = "\u03bb" * 160
    # Legacy exact JSON escapes Unicode, so its bytes exceed the retained UTF-8 size.
    output_limit = 64 * 1024 * 1024

    def rows():
        for index in range(count):
            yield {"id": f"{index:06d}", "detail": f"{detail}-{index}"}

    task = retain_prepared(lifecycle, [
        NamedOutput(
            "records", "records-v1", rows(), complete(count),
            (RecordColumn("id", "string"), RecordColumn("detail", "string")),
        ),
    ])
    reference = task.output("records")
    assert reference.size_bytes > 8 * 1024 * 1024
    writes = lifecycle.results.container.sequence
    for output_format, profile in (("json", "exact_records_v1"), ("csv", "tabular_records_v1")):
        lifecycle.restart(max_output_bytes=output_limit)
        requested = lifecycle.prepare(
            output_format, profile=profile, reference=reference,
            **({"columns": ("id", "detail")} if output_format == "csv" else {}),
        )
        completed = lifecycle.run(requested)
        assert completed["state"] == "completed", json.dumps(completed, sort_keys=True)
        lifecycle.restart(max_output_bytes=output_limit)
        downloaded = lifecycle.download(completed)
        descriptor = lifecycle.raw(completed)["committed_intent"]
        assert completed["row_count"] == count
        assert descriptor["content_sha256"] == hashlib.sha256(downloaded).hexdigest()
        assert descriptor["size_bytes"] == len(downloaded) > 8 * 1024 * 1024
        if output_format == "json":
            restored = json.loads(downloaded)
            assert len(restored) == count
            assert restored[0] == {"id": "000000", "detail": f"{detail}-0"}
            assert restored[-1] == {"id": "029999", "detail": f"{detail}-29999"}
        else:
            restored = list(csv.reader(io.StringIO(downloaded.decode("utf-8"), newline="")))
            assert len(restored) == count + 1
            assert restored[0] == ["id", "detail"]
            assert restored[1] == ["000000", f"{detail}-0"]
            assert restored[-1] == ["029999", f"{detail}-29999"]
    assert lifecycle.results.container.sequence == writes
    assert lifecycle.blobs.uploads == len(lifecycle.render_calls) == 2
