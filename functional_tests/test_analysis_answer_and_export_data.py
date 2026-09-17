# test_analysis_answer_and_export_data.py
"""
Functional tests for readable Analyze answers and authoritative export values.
Version: 0.261.113
Implemented in: 0.261.109

Display text is not the data handoff, and intermediate notes cannot replace
the final values used to build an export.
"""

import ast
import csv
import io
import json
import re
import sys
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

from docx import Document
from lxml import etree
from types import SimpleNamespace
import uuid

import pytest

from test_analyze_backend_saved_integration import final_analysis, load_functions, saved
from test_saved_analysis_service import saved_chat
from test_support.app_stubs import import_app_module
from test_workflow_result_contract import SerializedSections


exports = import_app_module("functions_generated_file_exports")
workflow_execution = import_app_module("functions_workflow_execution")
APP = Path(__file__).resolve().parents[1] / "application" / "single_app"


@pytest.fixture
def docx_markdown_renderer(monkeypatch):
    """Use the existing operations renderer without initializing its Azure clients."""
    path = APP / "functions_simplechat_operations.py"
    tree = ast.parse(path.read_text(encoding="utf-8"), filename=str(path))
    names = {"_clean_markdown_like_text", "_append_markdown_like_content_to_docx"}
    nodes = [node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name in names]
    assert {node.name for node in nodes} == names
    namespace = {"re": re}
    exec(compile(ast.Module(body=nodes, type_ignores=[]), str(path), "exec"), namespace)
    monkeypatch.setitem(sys.modules, "functions_simplechat_operations", SimpleNamespace(
        _append_markdown_like_content_to_docx=namespace["_append_markdown_like_content_to_docx"],
    ))


def analysis_result():
    return {
        "reply": "Download the attached analysis file.",
        "analysis_result": {
            "analysis_result_version": "analyze-final-v1",
            "analysis_reply": "## Findings\n\nThe reviewed controls need clearer ownership.",
            "authoritative_result": {
                "kind": "records",
                "value": [
                    {
                        "record_id": "finding-1",
                        "document_id": "doc-1",
                        "source": {"file_name": "Controls A", "scope_id": "INTERNAL_OWNER", "blob_path": "INTERNAL_STORAGE"},
                        "internal_lineage": {"candidate": "INTERNAL_CANDIDATE"},
                        "values": {"control": "Access review", "score": 81.25},
                        "evidence_refs": ["evidence-1"],
                    },
                    {
                        "record_id": "finding-2",
                        "document_id": "doc-2",
                        "source": {"file_name": "Controls B"},
                        "values": {"control": "Incident response", "score": 92.5},
                        "evidence_refs": ["evidence-2"],
                    },
                ],
            },
            "raw_analysis_items": [
                {"text": '{"control":"Access review","score":900}'},
            ],
            "analysis_validation": {
                "status": "valid",
                "coverage": {
                    "assigned_sources": 2, "completed_sources": 2,
                    "assigned_work_units": 2, "completed_work_units": 2,
                    "failed_work_units": 0, "pending_work_units": 0,
                },
                "limitations": ["No independent factual review was performed."],
            },
            "analysis_diagnostics": {"rejected": [{"score": 99998765, "private": "DIAGNOSTIC_ONLY"}]},
        },
    }


def test_artifact_announcement_cannot_replace_the_readable_answer():
    result = analysis_result()
    answer = exports.get_assistant_presentation_content(result)
    assert answer == result["analysis_result"]["analysis_reply"]
    assert "Download" not in answer
    assert "900" not in answer


def test_exports_use_only_final_public_values():
    result = analysis_result()
    assert exports.get_analysis_export_rows(result) == [
        {"control": "Access review", "score": 81.25},
        {"control": "Incident response", "score": 92.5},
    ]
    assert exports.get_analysis_export_rows(result["analysis_result"]) == [
        {"control": "Access review", "score": 81.25},
        {"control": "Incident response", "score": 92.5},
    ]


def test_zero_findings_are_not_replaced_with_diagnostic_rows():
    result = analysis_result()
    result["analysis_result"]["authoritative_result"]["value"] = []
    assert exports.get_analysis_export_rows(result) == []


def test_missing_final_records_do_not_fall_back_to_raw_notes():
    result = analysis_result()
    del result["analysis_result"]["authoritative_result"]
    with pytest.raises(ValueError, match="final analysis records"):
        exports.get_analysis_export_rows(result)


def test_missing_report_is_explicit():
    result = analysis_result()
    result["analysis_result"]["analysis_reply"] = ""
    with pytest.raises(ValueError, match="report is unavailable"):
        exports.get_assistant_presentation_content(result)


def test_legacy_data_and_presentation_accessors_remain_distinct():
    result = {
        "reply": "A legacy artifact is attached.",
        "analysis_result": {"analysis_reply": '{"legacy_field":1}'},
    }
    assert exports.get_assistant_presentation_content(result) == result["reply"]
    assert exports.get_generated_file_export_content(result) == '{"legacy_field":1}'
    assert exports.get_analysis_export_rows(result) is None


@pytest.mark.parametrize("output_format", ["pdf", "docx", "csv"])
def test_workflow_explanation_exports_keep_the_original_source_restrictions(saved_chat, output_format):
    results = import_app_module("functions_workflow_results")
    workflow = {
        "id": "workflow-1", "user_id": "owner",
        "task_prompt": f"Analyze the controls; export as {output_format.upper()}.",
    }
    store = SerializedSections()
    envelope = results.build_workflow_task_result(
        {"analysis_result": final_analysis(saved_chat), "reply": "Analysis of a restricted source."},
        workflow=workflow, run_id="run-1", task={"id": "analyze"},
    )
    manifest, reference = results.persist_workflow_task_result(
        envelope, workflow=workflow, run_id="run-1", task_id="analyze", save_result=store.save,
    )
    result = {
        "reply": "Saved controls lack owners.",
        "analysis_consumption": {"mode": "complete_input"},
        "analysis_origin_results": [results.workflow_result_summary(manifest, reference)],
    }
    uploads = []
    artifact_docs = []

    def upload(**kwargs):
        uploads.append(kwargs)
        artifact_docs.append({
            "id": "new-export", "conversation_id": kwargs["conversation_id"], "role": "file",
            "metadata": {
                "is_generated_chat_artifact": True,
                **saved.analysis_artifact_metadata(kwargs.get("analysis_producer")),
            },
        })
        return {"message": {"id": "new-export", "file_name": kwargs["file_name"]}}

    def bind_contexts(artifact, contexts, producer):
        artifact_docs[0]["metadata"]["analysis_result_contexts"] = contexts

    noop = lambda *args, **kwargs: None
    namespace = {
        "uuid": uuid,
        "assert_workflow_execution_owned": workflow_execution.assert_workflow_execution_owned,
        "workflow_saved_analysis_descriptor": saved.workflow_saved_analysis_descriptor,
        "saved_analysis_context": saved.saved_analysis_context,
        "get_analysis_export_rows": exports.get_analysis_export_rows,
        "get_generated_file_export_content": exports.get_generated_file_export_content,
        "get_requested_generated_file_format": exports.get_requested_generated_file_format,
        "has_generated_tabular_csv_output": lambda *args: False,
        "has_generated_file_output": exports.has_generated_file_output,
        "build_generated_file_export": lambda prompt, text, **kwargs: {
            "file_name": f"analysis.{output_format}", "output_format": output_format,
            "file_content": text.encode(), "row_count": 200000, "_structured_rows": [{"finding": "No owner"}],
        },
        "get_settings": lambda: {},
        "build_tabular_generated_output_row_batches": lambda rows, **kwargs: [rows],
        "should_queue_tabular_generated_output_background": lambda *args: True,
        "queue_tabular_generated_output_run": lambda **kwargs: pytest.fail("A source-bound export cannot become an unbound background copy."),
        "upload_generated_analysis_artifact_for_user": upload,
        "build_generated_file_artifact_metadata": exports.build_generated_file_artifact_metadata,
        "analysis_artifact_metadata": saved.analysis_artifact_metadata,
        "_bind_analysis_projection_contexts": bind_contexts,
        "_get_document_action_config": lambda value: {"type": "analyze"},
        "_get_workflow_scope": lambda value: "personal",
        "_get_workflow_group_id": lambda value: None,
        "_utc_now_iso": lambda: "2026-09-16T20:00:00Z",
        "build_cited_source_subsets": lambda *args, **kwargs: {
            "cited_hybrid_citations": [], "cited_web_search_citations": [],
        },
        "_persist_agent_citation_artifacts": lambda **kwargs: [],
        "cosmos_messages_container": SimpleNamespace(upsert_item=lambda item: item),
        "cosmos_conversations_container": SimpleNamespace(upsert_item=lambda item: item),
        **{name: noop for name in (
            "log_event", "apply_agent_document_citations", "initialize_conversation_used_document_tracking",
            "merge_cited_documents_into_conversation",
        )},
    }
    load_functions("functions_workflow_runner.py", {
        "_create_assistant_message", "_maybe_create_workflow_generated_file_output",
    }, namespace)
    assistant = namespace["_create_assistant_message"](
        {"id": "conversation-1"}, workflow, result, "manual", "run-1", {},
        assistant_message_id="final-assistant",
    )
    assert len(uploads) == 1
    assert assistant["metadata"]["saved_analyses"]
    reads = []

    def revoked(*args):
        reads.append(args)
        raise PermissionError("Source access revoked.")

    for for_publication in (False, True):
        with pytest.raises(PermissionError):
            saved.authorize_analysis_artifact(
                "owner", artifact_docs[0], result_reader=revoked,
                parents_loader=lambda *args: [assistant], for_publication=for_publication,
            )
    assert len(reads) == 2


@pytest.mark.parametrize("output_format", ["json", "csv", "docx"])
def test_generated_artifact_bytes_only_contain_accepted_public_values(output_format, docx_markdown_renderer):
    result = analysis_result()["analysis_result"]
    result["authoritative_result"]["value"][0]["values"]["control"] = "Access & `review`_[literal]*"
    public_rows = exports.get_analysis_export_rows(result)
    artifact = exports.build_saved_analysis_export(result, output_format)
    if output_format == "json":
        assert json.loads(artifact["file_content"]) == public_rows
        rendered = artifact["file_content"]
    elif output_format == "csv":
        reader = csv.DictReader(io.StringIO(artifact["file_content"]))
        assert reader.fieldnames == ["control", "score"]
        assert list(reader) == [
            {key: str(value) for key, value in row.items()} for row in public_rows
        ]
        rendered = artifact["file_content"]
    else:
        document = Document(io.BytesIO(artifact["file_content"]))
        assert len(document.tables) == 1
        table = document.tables[0]
        assert [cell.text for cell in table.rows[0].cells] == ["control", "score"]
        assert [[cell.text for cell in row.cells] for row in table.rows[1:]] == [
            [str(value) for value in row.values()] for row in public_rows
        ]
        headings = [paragraph.text for paragraph in document.paragraphs if paragraph.style.name.startswith("Heading")]
        assert {"Coverage", "Limitations", "Accepted findings"}.issubset(headings)
        assert all(not paragraph.text.startswith(("#", "- **")) for paragraph in document.paragraphs)
        rendered = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for row in table.rows for cell in row.cells]
        )
        with ZipFile(io.BytesIO(artifact["file_content"])) as package:
            assert "word/document.xml" in package.namelist()
            for name in package.namelist():
                if name.endswith(".xml"):
                    etree.fromstring(package.read(name))
    for internal in ("record_id", "document_id", "internal_lineage", "evidence_refs", "INTERNAL_OWNER", "INTERNAL_STORAGE", "INTERNAL_CANDIDATE", "DIAGNOSTIC_ONLY", "99998765", "900"):
        assert internal not in rendered


@pytest.mark.parametrize("output_format", ["csv", "docx"])
def test_generic_export_uses_explicit_final_analysis_without_raw_or_history_fallback(
    output_format, docx_markdown_renderer, monkeypatch,
):
    result = analysis_result()

    def no_raw_rows(*args, **kwargs):
        raise AssertionError("Accepted final Analyze exports must not read raw or legacy rows.")

    monkeypatch.setattr(exports, "extract_authorized_function_result_rows", no_raw_rows)
    artifact = exports.build_generated_file_export(
        f"Export this analysis as {output_format}.",
        json.dumps({"record_id": "RAW_ENVELOPE", "source": "RAW_SOURCE", "score": 900}),
        function_results=[{"function_result": {"rows": [{"score": 99998765}]}}],
        prior_function_results_loader=no_raw_rows,
        analysis_result=result,
    )
    assert artifact["_structured_rows"] == exports.get_analysis_export_rows(result)
    assert list(artifact["preview_rows"][0]) == ["control", "score"]
    assert artifact["row_count"] == 2
    assert artifact["passthrough_reason_code"] == "analysis_final_records"


@pytest.mark.parametrize("value", [{}, {"analysis_reply": "Legacy report."}, {"analysis_result_version": "analyze-final-v1"}])
def test_explicit_analysis_export_never_falls_back_when_final_data_is_missing(value):
    with pytest.raises(ValueError):
        exports.build_generated_file_export(
            "Export as CSV.", '{"legacy_column": 900}', analysis_result=value,
            function_results=[{"function_result": {"rows": [{"legacy_column": 900}]}}],
        )


def test_empty_final_csv_does_not_fall_back_to_available_legacy_rows():
    result = analysis_result()
    result["analysis_result"]["authoritative_result"]["value"] = []
    artifact = exports.build_generated_file_export(
        "Export as CSV.", '{"legacy_column": 900}', analysis_result=result,
        function_results=[{"function_result": {"rows": [{"legacy_column": 900}]}}],
    )
    assert artifact["row_count"] == 0
    assert artifact["_structured_rows"] == []
    assert "legacy_column" not in artifact["file_content"]


@pytest.mark.parametrize("output_format", ["json", "xml"])
def test_chat_structured_export_uses_final_values_not_envelope_or_raw_fallback(output_format):
    result = analysis_result()
    source = APP / "route_backend_chats.py"
    tree = ast.parse(source.read_text(encoding="utf-8"))
    node = next(node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name == "maybe_create_assistant_file_generated_output")
    uploads = []

    def forbidden(*args, **kwargs):
        raise AssertionError("A modern Analyze artifact must never use raw or legacy fallback.")

    def upload(**kwargs):
        uploads.append(kwargs)
        return {"message": {"id": "artifact", "file_name": kwargs["file_name"]}}

    namespace = {
        "get_tabular_generated_output_format": lambda prompt: output_format,
        "_has_generated_file_output": lambda *args: False,
        "_assistant_content_disclaims_complete_file": forbidden,
        "_build_structured_artifact_rows_payload": forbidden,
        "get_analysis_export_rows": exports.get_analysis_export_rows,
        "serialize_generated_json": exports.serialize_generated_json,
        "serialize_generated_xml": exports.serialize_generated_xml,
        "_build_assistant_file_preview_lines": lambda value: [],
        "_build_assistant_file_export_name": lambda value: f"analysis.{value}",
        "upload_generated_analysis_artifact_for_current_user": upload,
        "log_event": lambda *args, **kwargs: None,
    }
    exec(compile(ast.Module(body=[node], type_ignores=[]), str(source), "exec"), namespace)
    artifact = namespace["maybe_create_assistant_file_generated_output"](
        f"Export as {output_format}.",
        json.dumps(result["analysis_result"]["authoritative_result"]["value"]),
        "conversation-1", analysis_result=result,
    )
    assert artifact["artifact_message_id"] == "artifact"
    content = uploads[0]["file_content"]
    if output_format == "json":
        assert json.loads(content) == exports.get_analysis_export_rows(result)
        assert artifact["preview_columns"] == ["control", "score"]
    else:
        rows = etree.fromstring(content.encode("utf-8")).findall("Record")
        assert len(rows) == 2
        assert [child.tag for child in rows[0]] == ["control", "score"]
    assert "record_id" not in content and "source" not in content and "DIAGNOSTIC_ONLY" not in content


@pytest.mark.parametrize("output_format", ["csv", "docx"])
def test_workflow_file_finalizer_uploads_projected_values_and_retains_saved_lineage(
    output_format, docx_markdown_renderer,
):
    result = analysis_result()
    source = APP / "functions_workflow_runner.py"
    tree = ast.parse(source.read_text(encoding="utf-8"))
    node = next(
        node for node in tree.body
        if isinstance(node, ast.FunctionDef) and node.name == "_maybe_create_workflow_generated_file_output"
    )
    uploads = []
    bindings = []
    producer = {"kind": "workflow", "workflow_id": "workflow-1", "run_id": "run-1", "task_id": "analyze"}
    context = {"conversation_id": "conversation-1", "message_id": "assistant-1", "result_sha256": "a" * 64}

    def forbidden(*args, **kwargs):
        raise AssertionError("Already accepted analysis data must not be re-queued through native generation.")

    def upload(**kwargs):
        uploads.append(kwargs)
        return {"message": {"id": "artifact-1", "file_name": kwargs["file_name"]}}

    namespace = {
        "get_requested_generated_file_format": exports.get_requested_generated_file_format,
        "has_generated_tabular_csv_output": lambda *args: False,
        "has_generated_file_output": lambda *args: False,
        "build_generated_file_export": exports.build_generated_file_export,
        "get_settings": lambda: {},
        "build_tabular_generated_output_row_batches": forbidden,
        "should_queue_tabular_generated_output_background": forbidden,
        "upload_generated_analysis_artifact_for_user": upload,
        "build_generated_file_artifact_metadata": exports.build_generated_file_artifact_metadata,
        "analysis_artifact_metadata": lambda value: {"analysis_result_required": True, "analysis_producer": value},
        "_bind_analysis_projection_contexts": lambda *args: bindings.append(args),
        "log_event": lambda *args, **kwargs: None,
    }
    exec(compile(ast.Module(body=[node], type_ignores=[]), str(source), "exec"), namespace)
    artifact = namespace["_maybe_create_workflow_generated_file_output"](
        {"user_id": "owner"}, "conversation-1", f"Export as {output_format}.",
        '{"record_id":"RAW_ENVELOPE","score":900}', analysis_result=result,
        analysis_producer=producer, analysis_contexts=[context],
    )
    assert artifact["analysis_producer"] == producer
    assert uploads[0]["analysis_producer"] == producer
    assert bindings == [(artifact, [context], producer)]
    if output_format == "csv":
        rows = list(csv.DictReader(io.StringIO(uploads[0]["file_content"])))
        assert rows == [{"control": "Access review", "score": "81.25"}, {"control": "Incident response", "score": "92.5"}]
    else:
        document = Document(io.BytesIO(uploads[0]["file_content"]))
        assert [cell.text for cell in document.tables[0].rows[0].cells] == ["control", "score"]
        assert document.tables[0].rows[1].cells[1].text == "81.25"
