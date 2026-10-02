#!/usr/bin/env python3
# test_tabular_merge_engine.py
"""
Functional test for the deterministic tabular merge engine.
Version: 0.261.218
Implemented in: 0.261.218

This test ensures that functions_tabular_merge appends the rows of same-structure
CSV and Excel files exactly: values stay text (leading zeros, codes and Unicode
survive), columns match by name or exact order, every mismatching file is reported
before anything is merged, encodings and delimiters are detected, workbook sheets are
selected deterministically, and every limit, guard and cancellation fails closed.
"""

import codecs
import io
import os
import sys
from datetime import date, datetime

import pytest

sys.path.append(os.path.dirname(os.path.abspath(__file__)))
sys.path.append(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "application", "single_app"))

from test_support.versioning import assert_app_version_at_least

import functions_tabular_merge as merge
from functions_tabular_merge import (
    TabularMergeCancelled,
    TabularMergeError,
    TabularMergeLimits,
    TabularMergeOptions,
    TabularMergeSource,
    merge_tabular_sources,
)


def csv_source(name, text, encoding="utf-8", source_id=None):
    data = text.encode(encoding) if isinstance(text, str) else text
    return TabularMergeSource(source_id or name, name, lambda data=data: data)


def bytes_source(name, data, source_id=None):
    return TabularMergeSource(source_id or name, name, lambda data=data: data)


def xlsx_bytes(sheets):
    from openpyxl import Workbook

    workbook = Workbook()
    first = True
    for title, rows, state in sheets:
        worksheet = workbook.active if first else workbook.create_sheet()
        first = False
        worksheet.title = title
        worksheet.sheet_state = state
        for row in rows:
            worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def records(result):
    return list(result.iter_records())


def test_version_includes_the_merge_engine():
    assert_app_version_at_least("0.261.218")


def test_appends_rows_in_file_order_with_a_source_column():
    first = csv_source("east.csv", "Region,Code,Amount\r\nEast,007,10.50\r\nEast,008,\"1,200\"\r\n")
    second = csv_source("west.csv", "Region,Code,Amount\nWest,0099,5\n")
    with merge_tabular_sources([first, second]) as result:
        assert result.columns == ("Source File", "Region", "Code", "Amount")
        assert result.row_count == 3
        assert records(result) == [
            {"Source File": "east.csv", "Region": "East", "Code": "007", "Amount": "10.50"},
            {"Source File": "east.csv", "Region": "East", "Code": "008", "Amount": "1,200"},
            {"Source File": "west.csv", "Region": "West", "Code": "0099", "Amount": "5"},
        ]
        report = result.report
        assert report["version"] == "tabular-merge-report-v1"
        assert report["status"] == "merged"
        assert report["totals"] == {"sources": 2, "rows": 3, "columns": 4, "blank_rows_skipped": 0}
        assert [entry["rows"] for entry in report["sources"]] == [2, 1]
        assert report["limitations"] == []


def test_matches_columns_by_name_in_any_order_and_reports_reordering():
    first = csv_source("a.csv", "Region,Code\nNorth,1\n")
    second = csv_source("b.csv", "  code ;REGION\n2;South\n")
    with merge_tabular_sources([first, second], options=TabularMergeOptions(include_source_column=False)) as result:
        assert result.columns == ("Region", "Code")
        assert records(result) == [{"Region": "North", "Code": "1"}, {"Region": "South", "Code": "2"}]
        assert [entry["order_differs"] for entry in result.report["sources"]] == [False, True]
        assert result.report["sources"][1]["delimiter"] == ";"


def test_exact_order_rejects_reordered_columns():
    first = csv_source("a.csv", "Region,Code\nNorth,1\n")
    second = csv_source("b.csv", "Code,Region\n2,South\n")
    with pytest.raises(TabularMergeError) as caught:
        merge_tabular_sources([first, second], options=TabularMergeOptions(schema_policy="exact_order"))
    assert caught.value.code == "schema_mismatch"
    assert caught.value.report["sources"][1]["order_differs"] is True
    assert "different order" in caught.value.message


def test_schema_mismatch_reports_every_file_and_merges_nothing():
    reference = csv_source("ref.csv", "Region,Code,Amount\nEast,1,2\n")
    extra = csv_source("extra.csv", "Region,Code,Total,Notes\nX,1,2,3\n")
    matching = csv_source("ok.csv", "Amount,Code,Region\n3,4,West\n")
    missing = csv_source("missing.csv", "Region,Amount\nY,1\n")
    with pytest.raises(TabularMergeError) as caught:
        merge_tabular_sources([reference, extra, matching, missing])
    error = caught.value
    assert error.code == "schema_mismatch"
    statuses = [(entry["file_name"], entry["status"]) for entry in error.report["sources"]]
    assert statuses == [
        ("ref.csv", "compatible"), ("extra.csv", "schema_mismatch"),
        ("ok.csv", "compatible"), ("missing.csv", "schema_mismatch"),
    ]
    assert error.report["sources"][1]["missing_columns"] == ["Amount"]
    assert error.report["sources"][1]["extra_columns"] == ["Total", "Notes"]
    assert error.report["sources"][3]["missing_columns"] == ["Code"]
    assert all(entry["rows"] == 0 for entry in error.report["sources"])
    assert error.report["status"] == "failed"
    assert "Nothing was merged" in error.message


@pytest.mark.parametrize(
    ("payload", "encoding_name"),
    [
        ("Name,City\nZoë,Köln\n".encode("utf-8"), "utf-8"),
        (codecs.BOM_UTF8 + "Name,City\nZoë,Köln\n".encode("utf-8"), "utf-8-sig"),
        ("Name,City\nZoë,Köln\n".encode("utf-16"), "utf-16"),
        ("Name,City\nZoë,Köln\n".encode("cp1252"), "cp1252"),
    ],
)
def test_detects_text_encodings(payload, encoding_name):
    other = csv_source("other.csv", "Name,City\nAnn,Oslo\n")
    with merge_tabular_sources([bytes_source("enc.csv", payload), other]) as result:
        assert records(result)[0] == {"Source File": "enc.csv", "Name": "Zoë", "City": "Köln"}
        assert result.report["sources"][0]["encoding"] == encoding_name


@pytest.mark.parametrize("delimiter", [",", ";", "\t", "|"])
def test_detects_delimiters_and_keeps_quoted_delimiters(delimiter):
    text = f'Name{delimiter}Note\n"Smith{delimiter} Jr"{delimiter}"line one\nline two"\n'
    other = csv_source("b.csv", "Name,Note\nAnn,x\n")
    with merge_tabular_sources([csv_source("a.csv", text), other]) as result:
        first = records(result)[0]
        assert first["Name"] == f"Smith{delimiter} Jr"
        assert first["Note"] == "line one\nline two"
        assert result.report["sources"][0]["delimiter"] == delimiter


def test_rejects_binary_data_mislabeled_as_csv():
    with pytest.raises(TabularMergeError) as caught:
        merge_tabular_sources([bytes_source("bad.csv", b"PK\x03\x04\x00\x00junk"), csv_source("b.csv", "A\n1\n")])
    assert caught.value.code == "unreadable_csv"


def test_reads_xlsx_values_as_canonical_text():
    workbook = xlsx_bytes([(
        "Data",
        [
            ["Code", "Amount", "When", "Flag", "Ratio", "Formula"],
            ["007", 3.0, datetime(2024, 1, 31), True, 0.1, "=1+1"],
            ["008", 12, datetime(2024, 2, 1, 9, 30, 15), False, 1e-07, None],
            [None, None, None, None, None, None],
            ["009", -2.5, date(2023, 12, 25), None, 123456789012.0, None],
        ],
        "visible",
    )])
    other = csv_source("b.csv", "Code,Amount,When,Flag,Ratio,Formula\n010,1,2024-03-01,TRUE,1,\n")
    with merge_tabular_sources([bytes_source("book.xlsx", workbook), other]) as result:
        rows = records(result)
        assert rows[0] == {
            "Source File": "book.xlsx", "Code": "007", "Amount": "3", "When": "2024-01-31",
            "Flag": "TRUE", "Ratio": "0.1", "Formula": "",
        }
        assert rows[1]["When"] == "2024-02-01T09:30:15"
        assert rows[1]["Amount"] == "12"
        assert rows[1]["Flag"] == "FALSE"
        assert rows[1]["Ratio"] == "1e-07"
        assert rows[2]["When"] == "2023-12-25"
        assert rows[2]["Amount"] == "-2.5"
        assert rows[2]["Ratio"] == "123456789012"
        report = result.report
        assert report["sources"][0]["sheet"] == "Data"
        assert report["sources"][0]["blank_rows_skipped"] == 1
        assert report["limitations"], "Workbook merges must disclose how formulas and types were read."


def test_selects_the_first_visible_sheet_or_a_named_sheet():
    workbook = xlsx_bytes([
        ("Hidden", [["A"], ["hidden"]], "hidden"),
        ("Visible", [["A"], ["visible"]], "visible"),
        ("Second", [["A"], ["second"]], "visible"),
    ])
    other = csv_source("b.csv", "A\ncsv\n")
    with merge_tabular_sources([bytes_source("w.xlsx", workbook), other]) as result:
        assert records(result)[0]["A"] == "visible"
    options = TabularMergeOptions(sheet="second")
    with merge_tabular_sources([bytes_source("w.xlsx", workbook), other], options=options) as result:
        assert records(result)[0]["A"] == "second"
        assert result.report["sources"][0]["sheet"] == "Second"


def test_missing_named_sheet_lists_the_available_sheets():
    workbook = xlsx_bytes([("Data", [["A"], ["1"]], "visible")])
    with pytest.raises(TabularMergeError) as caught:
        merge_tabular_sources(
            [bytes_source("w.xlsx", workbook), csv_source("b.csv", "A\n2\n")],
            options=TabularMergeOptions(sheet="Totals"),
        )
    assert caught.value.code == "sheet_not_found"
    assert caught.value.details["available_sheets"] == ["Data"]


def test_xls_reader_converts_legacy_cell_types(monkeypatch):
    import xlrd

    class Cell:
        def __init__(self, ctype, value):
            self.ctype = ctype
            self.value = value

    rows = [
        [Cell(xlrd.XL_CELL_TEXT, "Code"), Cell(xlrd.XL_CELL_TEXT, "Amount"), Cell(xlrd.XL_CELL_TEXT, "When"),
         Cell(xlrd.XL_CELL_TEXT, "Flag"), Cell(xlrd.XL_CELL_TEXT, "Error")],
        [Cell(xlrd.XL_CELL_TEXT, "007"), Cell(xlrd.XL_CELL_NUMBER, 4.0), Cell(xlrd.XL_CELL_DATE, 45322.0),
         Cell(xlrd.XL_CELL_BOOLEAN, 1), Cell(xlrd.XL_CELL_ERROR, 0x2A)],
        [Cell(xlrd.XL_CELL_EMPTY, ""), Cell(xlrd.XL_CELL_BLANK, ""), Cell(xlrd.XL_CELL_EMPTY, ""),
         Cell(xlrd.XL_CELL_EMPTY, ""), Cell(xlrd.XL_CELL_EMPTY, "")],
    ]

    class Sheet:
        visibility = 0
        nrows = len(rows)

        def row(self, index):
            return rows[index]

    class Book:
        datemode = 0

        def sheet_names(self):
            return ["Legacy"]

        def sheet_by_index(self, index):
            return Sheet()

        def sheet_by_name(self, name):
            assert name == "Legacy"
            return Sheet()

        def release_resources(self):
            return None

    monkeypatch.setattr(xlrd, "open_workbook", lambda **kwargs: Book())
    legacy = bytes_source("old.xls", merge._OLE_MAGIC + b"\x00" * 64)
    other = csv_source("b.csv", "Code,Amount,When,Flag,Error\n1,2,3,4,5\n")
    with merge_tabular_sources([legacy, other]) as result:
        first = records(result)[0]
        assert first == {
            "Source File": "old.xls", "Code": "007", "Amount": "4", "When": "2024-01-31",
            "Flag": "TRUE", "Error": "#N/A",
        }
        assert result.report["sources"][0]["blank_rows_skipped"] == 1


def test_password_protected_and_damaged_workbooks_fail_clearly():
    other = csv_source("b.csv", "A\n1\n")
    with pytest.raises(TabularMergeError) as caught:
        merge_tabular_sources([bytes_source("locked.xlsx", merge._OLE_MAGIC + b"\x00" * 32), other])
    assert caught.value.code == "encrypted_workbook"
    with pytest.raises(TabularMergeError) as caught:
        merge_tabular_sources([bytes_source("broken.xlsx", b"not a workbook"), other])
    assert caught.value.code == "unreadable_workbook"
    with pytest.raises(TabularMergeError) as caught:
        merge_tabular_sources([bytes_source("broken.xls", b"<html>not xls</html>"), other])
    assert caught.value.code == "unreadable_workbook"


def test_workbook_archive_guard_rejects_oversized_expansion():
    workbook = xlsx_bytes([("Data", [["A"], ["1"]], "visible")])
    limits = TabularMergeLimits(max_workbook_uncompressed_bytes=64)
    with pytest.raises(TabularMergeError) as caught:
        merge_tabular_sources([bytes_source("w.xlsx", workbook), csv_source("b.csv", "A\n2\n")], limits=limits)
    assert caught.value.code == "workbook_too_large"


def test_blank_rows_are_skipped_short_rows_padded_and_extra_values_refused():
    first = csv_source("a.csv", "A,B,C,,\n1,2,3,,\n,,\n4\n")
    second = csv_source("b.csv", "A,B,C\n5,6,7\n")
    with merge_tabular_sources([first, second], options=TabularMergeOptions(include_source_column=False)) as result:
        assert records(result) == [
            {"A": "1", "B": "2", "C": "3"},
            {"A": "4", "B": "", "C": ""},
            {"A": "5", "B": "6", "C": "7"},
        ]
        source_entry = result.report["sources"][0]
        assert source_entry["blank_rows_skipped"] == 1
        assert source_entry["short_rows_padded"] == 1
    bad = csv_source("bad.csv", "A,B,C\n1,2,3,surprise\n")
    with pytest.raises(TabularMergeError) as caught:
        merge_tabular_sources([second, bad])
    assert caught.value.code == "row_has_extra_values"


def test_header_rules():
    with pytest.raises(TabularMergeError) as caught:
        merge_tabular_sources([csv_source("dup.csv", "A,a\n1,2\n"), csv_source("b.csv", "A\n1\n")])
    assert caught.value.code == "duplicate_columns"
    with merge_tabular_sources(
        [csv_source("blank.csv", "A,,C\n1,2,3\n"), csv_source("b.csv", "A,Column 2,C\n4,5,6\n")],
    ) as result:
        assert result.columns == ("Source File", "A", "Column 2", "C")
        assert result.report["sources"][0]["warnings"][0]["code"] == "blank_header"
    with pytest.raises(TabularMergeError) as caught:
        merge_tabular_sources([csv_source("empty.csv", "\n\n"), csv_source("b.csv", "A\n1\n")])
    assert caught.value.code == "empty_source"
    long_header = "é" * 129  # 258 UTF-8 bytes: record column names are capped at 256 bytes.
    with pytest.raises(TabularMergeError) as caught:
        merge_tabular_sources([csv_source("long.csv", f"{long_header}\n1\n"), csv_source("b.csv", "A\n1\n")])
    assert caught.value.code == "header_too_long"


def test_source_column_is_renamed_on_collision_or_omitted():
    first = csv_source("a.csv", "Source File,Value\nx,1\n")
    second = csv_source("b.csv", "Source File,Value\ny,2\n")
    with merge_tabular_sources([first, second]) as result:
        assert result.columns == ("Source File (2)", "Source File", "Value")
        assert result.report["warnings"][0]["code"] == "source_column_renamed"
        assert records(result)[0] == {"Source File (2)": "a.csv", "Source File": "x", "Value": "1"}
    options = TabularMergeOptions(include_source_column=False)
    with merge_tabular_sources([first, second], options=options) as result:
        assert result.columns == ("Source File", "Value")
        assert result.report["source_column"] is None


def test_header_only_files_merge_to_an_empty_complete_table():
    with merge_tabular_sources([csv_source("a.csv", "A,B\n"), csv_source("b.csv", "A,B\r\n")]) as result:
        assert result.row_count == 0
        assert records(result) == []
        assert result.report["status"] == "merged"


def test_source_validation():
    with pytest.raises(TabularMergeError) as caught:
        merge_tabular_sources([csv_source("a.csv", "A\n1\n")])
    assert caught.value.code == "too_few_sources"
    with pytest.raises(TabularMergeError) as caught:
        merge_tabular_sources([csv_source("a.csv", "A\n1\n", source_id="same"), csv_source("b.csv", "A\n2\n", source_id="same")])
    assert caught.value.code == "duplicate_source"
    with pytest.raises(TabularMergeError) as caught:
        merge_tabular_sources([csv_source("a.csv", "A\n1\n"), csv_source("notes.docx", "A\n2\n")])
    assert caught.value.code == "unsupported_format"
    with pytest.raises(TabularMergeError) as caught:
        TabularMergeOptions(schema_policy="union")
    assert caught.value.code == "invalid_options"


def test_limits_fail_closed():
    sources = [csv_source(f"{index}.csv", "A\n1\n") for index in range(3)]
    with pytest.raises(TabularMergeError) as caught:
        merge_tabular_sources(sources, limits=TabularMergeLimits(max_sources=2))
    assert caught.value.code == "too_many_sources"
    many_rows = csv_source("big.csv", "A\n" + "\n".join(str(value) for value in range(30)) + "\n")
    with pytest.raises(TabularMergeError) as caught:
        merge_tabular_sources([many_rows, sources[0]], limits=TabularMergeLimits(max_total_rows=10))
    assert caught.value.code == "row_limit_exceeded"
    with pytest.raises(TabularMergeError) as caught:
        merge_tabular_sources([many_rows, sources[0]], limits=TabularMergeLimits(max_source_bytes=8))
    assert caught.value.code == "source_too_large"
    wide = ",".join(f"C{index}" for index in range(5))
    with pytest.raises(TabularMergeError) as caught:
        merge_tabular_sources(
            [csv_source("w.csv", f"{wide}\n"), csv_source("x.csv", f"{wide}\n")],
            limits=TabularMergeLimits(max_columns=5),
        )
    assert caught.value.code == "too_many_columns"
    with pytest.raises(TabularMergeError) as caught:
        merge_tabular_sources(
            [csv_source("c.csv", "A\n" + "x" * 50 + "\n"), sources[0]],
            limits=TabularMergeLimits(max_cell_chars=10),
        )
    assert caught.value.code == "cell_too_large"


def test_cancellation_stops_between_sources():
    calls = {"count": 0}

    def cancel_after_first():
        calls["count"] += 1
        return calls["count"] > 2

    sources = [csv_source(f"{index}.csv", "A\n1\n") for index in range(3)]
    with pytest.raises(TabularMergeCancelled):
        merge_tabular_sources(sources, cancel_requested=cancel_after_first)


def test_rows_can_be_read_more_than_once_until_closed_and_progress_is_reported():
    events = []
    result = merge_tabular_sources(
        [csv_source("a.csv", "A\n1\n2\n"), csv_source("b.csv", "A\n3\n")], on_progress=events.append,
    )
    try:
        assert records(result) == records(result)
        assert [event["index"] for event in events] == [1, 2]
        assert events[-1] == {"index": 2, "total": 2, "file_name": "b.csv", "rows": 1}
    finally:
        result.close()
    with pytest.raises(TabularMergeError) as caught:
        list(result.iter_rows())
    assert caught.value.code == "result_closed"


def test_large_merges_spill_to_disk_and_keep_exact_counts():
    rows = "\n".join(f"{index:07d},{'x' * 40}" for index in range(120_000))
    first = csv_source("big1.csv", f"Code,Text\n{rows}\n")
    second = csv_source("big2.csv", "Code,Text\n0000001,y\n")
    with merge_tabular_sources([first, second]) as result:
        assert result.row_count == 120_001
        count = 0
        for record in result.iter_records():
            count += 1
        assert count == 120_001
        assert records(result)[0]["Code"] == "0000000"


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
