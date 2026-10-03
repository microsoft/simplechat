#!/usr/bin/env python3
# test_document_merge_assembly.py
"""
Functional test for the retained document merge description and its re-assembly.
Version: 0.261.224
Implemented in: 0.261.224

This test ensures that functions_document_merge_assembly accepts only the settings of
the chosen merge kind, describes a checked merge by its files, options, size and
SHA-256 without keeping its bytes, refuses any description it did not write, and
assembles the same files again into byte-identical output for every kind, both in this
process and in fresh processes with different hash seeds, because Render may run in
another worker. A changed file fails as source_changed, a smaller limit fails before
any file is read, caller and access failures pass through unchanged, and importing the
merge engine, the description module or the spreadsheet engine loads no Office or PDF
library, with or without optimization.
"""

import hashlib
import io
import json
import os
import subprocess
import sys
from pathlib import Path

import pytest

sys.path.append(os.path.dirname(os.path.abspath(__file__)))
APP = Path(__file__).resolve().parent.parent / "application" / "single_app"
sys.path.append(str(APP))

from test_support.versioning import assert_app_version_at_least

from functions_document_merge import DocumentMergeError, DocumentMergeOptions, DocumentMergePart, merge_documents
from functions_document_merge_assembly import (
    DOCUMENT_ASSEMBLY_PROFILE,
    DocumentAssemblyError,
    assembled_output_format,
    build_document_assembly,
    document_merge_options_from_arguments,
    reassemble_document,
    validate_document_assembly,
)


def pdf_bytes(pages):
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def docx_bytes(label, color):
    from docx import Document
    from docx.shared import RGBColor

    document = Document()
    document.styles["Heading 1"].font.color.rgb = RGBColor(*color)
    document.add_heading(f"{label} title", level=1)
    document.add_paragraph(f"{label} body")
    document.add_paragraph("first item", style="List Number")
    # Word numbers ordinary lists directly; a merge copies such a list with an ID of its own.
    listed = document.add_paragraph(f"{label} listed")
    numbering = listed._p.get_or_add_pPr().get_or_add_numPr()
    numbering.get_or_add_ilvl().val = 0
    numbering.get_or_add_numId().val = 1
    document.add_table(rows=1, cols=2).cell(0, 0).text = f"{label} cell"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pptx_bytes(label, slides):
    from pptx import Presentation

    presentation = Presentation()
    for number in range(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1 if number else 0])
        slide.shapes.title.text = f"{label} slide {number + 1}"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def xlsx_bytes(title):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(["Region", "Amount", "When"])
    sheet.append(["West", 5.25, "2024-01-31"])
    workbook.create_sheet("Notes").append(["kept"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


FIXTURES = {
    "pdf": [("a.pdf", pdf_bytes(2)), ("b.pdf", pdf_bytes(1))],
    "docx": [("a.docx", docx_bytes("A", (200, 0, 0))), ("b.docx", docx_bytes("B", (0, 0, 200)))],
    "pptx": [("a.pptx", pptx_bytes("A", 2)), ("b.pptx", pptx_bytes("B", 2))],
    "workbook": [("east.csv", b"Region,Amount\r\nEast,10\r\n"), ("west.xlsx", xlsx_bytes("West"))],
}
OPTIONS = {
    "pdf": {"bookmarks": True},
    "docx": {"formatting": "keep_source", "source_headings": True},
    "pptx": {"formatting": "keep_source", "sections": True},
    "workbook": {"sheets": "all"},
}


def parts_for(kind, files=None):
    files = files or {name: data for name, data in FIXTURES[kind]}
    return [
        DocumentMergePart(f"doc-{name}", name, lambda name=name: files[name])
        for name, _ in FIXTURES[kind]
    ]


def checked(kind, **arguments):
    """Merge once and describe the result, like a chat document_merge step does."""
    kind, options = document_merge_options_from_arguments({"kind": kind, **OPTIONS[kind], **arguments})
    parts = parts_for(kind)
    with merge_documents(kind, parts, options=options) as result:
        content = result.read_bytes()
        description = build_document_assembly(kind, options, parts, result, hashlib.sha256(content).hexdigest())
    return description, content


def reader_for(kind, overrides=None, reads=None):
    files = {f"doc-{name}": data for name, data in FIXTURES[kind]}
    files.update(overrides or {})

    def read(document_id):
        if reads is not None:
            reads.append(document_id)
        return files[document_id]

    return read


def test_version_includes_document_merge_assembly():
    assert_app_version_at_least("0.261.224")


def test_options_follow_the_kind_and_leave_defaults_to_the_engine():
    kind, options = document_merge_options_from_arguments({"kind": "pdf", "document_ids": ["a", "b"]})
    assert kind == "pdf" and options == DocumentMergeOptions()
    kind, options = document_merge_options_from_arguments({
        "kind": "docx", "formatting": "use_first", "page_breaks": False, "source_headings": True, "doc_scope": "all",
    })
    assert (options.formatting, options.page_breaks, options.source_headings) == ("use_first", False, True)
    kind, options = document_merge_options_from_arguments({"kind": "workbook", "sheet": "Data", "sheets": None})
    assert (options.sheet, options.sheets) == ("Data", "first")
    for arguments in (
        {}, {"kind": "zip"}, {"kind": "PDF"}, {"kind": "pdf", "formatting": "use_first"},
        {"kind": "pdf", "sections": True}, {"kind": "docx", "sheets": "all"}, {"kind": "pptx", "bookmarks": False},
        {"kind": "workbook", "page_breaks": True}, {"kind": "workbook", "sheet": "Data", "sheets": "all"},
        {"kind": "docx", "formatting": "mine"}, {"kind": "pptx", "sections": "yes"}, {"kind": "workbook", "sheet": ""},
    ):
        with pytest.raises(DocumentMergeError) as refused:
            document_merge_options_from_arguments(arguments)
        assert refused.value.code == "invalid_options", arguments


def test_a_description_names_files_options_size_and_digest_only():
    description, content = checked("pdf")
    assert description == {
        "profile": DOCUMENT_ASSEMBLY_PROFILE,
        "kind": "pdf",
        "options": {
            "formatting": "keep_source", "bookmarks": True, "sections": True, "page_breaks": True,
            "source_headings": False, "sheets": "first", "sheet": None,
        },
        "parts": [{"document_id": "doc-a.pdf", "file_name": "a.pdf"}, {"document_id": "doc-b.pdf", "file_name": "b.pdf"}],
        "output": {
            "format": "pdf", "media_type": "application/pdf", "size_bytes": len(content),
            "content_sha256": hashlib.sha256(content).hexdigest(),
        },
    }
    assert assembled_output_format(description) == "pdf"
    copy = validate_document_assembly(description)
    copy["parts"][0]["file_name"] = "changed.pdf"
    assert description["parts"][0]["file_name"] == "a.pdf"


@pytest.mark.parametrize("change", [
    lambda value: value.update(extra=True),
    lambda value: value.update(profile="document_assembly_v2"),
    lambda value: value.update(kind="zip"),
    lambda value: value["options"].pop("sheet"),
    lambda value: value["options"].update(formatting="mine"),
    lambda value: value["options"].update(bookmarks="yes"),
    lambda value: value.update(parts=value["parts"][:1]),
    lambda value: value["parts"].append(dict(value["parts"][0])),
    lambda value: value["parts"][0].update(document_id=" "),
    lambda value: value["parts"][0].update(document_id=7),
    lambda value: value["parts"][0].update(extra="x"),
    lambda value: value["parts"][0].update(file_name="x" * 2000),
    lambda value: value["output"].update(format="docx"),
    lambda value: value["output"].update(media_type="text/plain"),
    lambda value: value["output"].update(size_bytes=0),
    lambda value: value["output"].update(size_bytes=True),
    lambda value: value["output"].update(size_bytes=1.5),
    lambda value: value["output"].update(content_sha256="A" * 64),
    lambda value: value["output"].update(content_sha256="a" * 63),
    lambda value: value["output"].pop("media_type"),
])
def test_descriptions_the_service_did_not_write_are_refused(change):
    description, _content = checked("pdf")
    change(description)
    with pytest.raises(DocumentAssemblyError) as refused:
        validate_document_assembly(description)
    assert refused.value.code == "invalid_source"
    with pytest.raises(DocumentAssemblyError):
        reassemble_document(description, reader_for("pdf"), max_output_bytes=10 * 1024 * 1024)


@pytest.mark.parametrize("kind", ["pdf", "docx", "pptx", "workbook"])
def test_the_same_files_assemble_into_the_same_bytes(kind):
    description, content = checked(kind)
    reads = []
    result, digest = reassemble_document(
        description, reader_for(kind, reads=reads), max_output_bytes=description["output"]["size_bytes"],
    )
    with result:
        assert result.read_bytes() == content
    assert digest == description["output"]["content_sha256"]
    assert reads == [part["document_id"] for part in description["parts"]]


def test_other_processes_assemble_the_same_bytes(tmp_path):
    expected = {}
    spec = {}
    for kind, files in FIXTURES.items():
        description, _content = checked(kind)
        expected[kind] = description["output"]["content_sha256"]
        entries = []
        for name, data in files:
            path = tmp_path / name
            path.write_bytes(data)
            entries.append([name, str(path)])
        spec[kind] = {"files": entries, "options": OPTIONS[kind]}
    script = (
        "import hashlib, json, sys\n"
        "from pathlib import Path\n"
        "from functions_document_merge import DocumentMergeOptions, DocumentMergePart, merge_documents\n"
        "spec = json.loads(Path(sys.argv[1]).read_text())\n"
        "digests = {}\n"
        "for kind, item in spec.items():\n"
        "    parts = [DocumentMergePart('doc-' + name, name, lambda path=path: Path(path).read_bytes())\n"
        "             for name, path in item['files']]\n"
        "    with merge_documents(kind, parts, options=DocumentMergeOptions(**item['options'])) as merged:\n"
        "        digests[kind] = hashlib.sha256(merged.read_bytes()).hexdigest()\n"
        "print(json.dumps(digests))\n"
    )
    spec_path = tmp_path / "spec.json"
    spec_path.write_text(json.dumps(spec))
    for seed in ("0", "4242"):
        environment = {
            **os.environ, "PYTHONHASHSEED": seed,
            "PYTHONPATH": os.pathsep.join(filter(None, [str(APP), os.environ.get("PYTHONPATH")])),
        }
        completed = subprocess.run(
            [sys.executable, "-c", script, str(spec_path)], capture_output=True, text=True,
            env=environment, timeout=300, check=False,
        )
        assert completed.returncode == 0, completed.stderr[-2000:]
        assert json.loads(completed.stdout.strip().splitlines()[-1]) == expected, seed


def test_a_changed_file_is_refused_as_changed():
    description, _content = checked("pdf")
    with pytest.raises(DocumentAssemblyError) as changed:
        reassemble_document(
            description, reader_for("pdf", {"doc-b.pdf": pdf_bytes(3)}), max_output_bytes=10 * 1024 * 1024,
        )
    assert changed.value.code == "source_changed"
    with pytest.raises(DocumentAssemblyError) as damaged:
        reassemble_document(description, reader_for("pdf", {"doc-b.pdf": b"%PDF-1.7 damaged"}), max_output_bytes=10 * 1024 * 1024)
    assert damaged.value.code == "source_changed"
    # A file that grew past the limit the checked file fits is a changed file, not a size problem.
    size = description["output"]["size_bytes"]
    with pytest.raises(DocumentAssemblyError) as grown:
        reassemble_document(description, reader_for("pdf", {"doc-b.pdf": pdf_bytes(400)}), max_output_bytes=size)
    assert grown.value.code == "source_changed"


def test_limits_and_readers_are_checked_before_any_file_is_read():
    description, _content = checked("docx")
    reads = []
    with pytest.raises(DocumentAssemblyError) as small:
        reassemble_document(
            description, reader_for("docx", reads=reads), max_output_bytes=description["output"]["size_bytes"] - 1,
        )
    assert small.value.code == "size_limit" and reads == []
    for limit in (0, -1, True, 1.5, None):
        with pytest.raises(DocumentAssemblyError) as invalid:
            reassemble_document(description, reader_for("docx", reads=reads), max_output_bytes=limit)
        assert invalid.value.code == "invalid_limit"
    with pytest.raises(DocumentAssemblyError) as missing:
        reassemble_document(description, None, max_output_bytes=10 * 1024 * 1024)
    assert missing.value.code == "unsupported_source"
    assert reads == []


def test_caller_and_access_failures_pass_through_unchanged():
    description, _content = checked("pptx")

    class Stopped(Exception):
        pass

    class Revoked(PermissionError):
        pass

    calls = []

    def stop():
        calls.append(True)
        if len(calls) == 2:
            raise Stopped()

    with pytest.raises(Stopped):
        reassemble_document(description, reader_for("pptx"), max_output_bytes=10 * 1024 * 1024, check=stop)

    def revoked(document_id):
        raise Revoked("access removed")

    with pytest.raises(Revoked):
        reassemble_document(description, revoked, max_output_bytes=10 * 1024 * 1024)

    checks = []
    result, _digest = reassemble_document(
        description, reader_for("pptx"), max_output_bytes=10 * 1024 * 1024, check=lambda: checks.append(True),
    )
    result.close()
    assert len(checks) >= 2


@pytest.mark.parametrize("optimize", [False, True])
def test_merge_modules_import_without_office_or_pdf_libraries(optimize):
    script = (
        "import sys\n"
        "import functions_document_merge, functions_document_merge_assembly, functions_tabular_merge\n"
        "heavy = sorted(name for name in ('pypdf', 'openpyxl', 'docx', 'docxcompose', 'pptx', 'lxml', 'xlrd',"
        " 'pandas', 'fitz') if name in sys.modules)\n"
        "if heavy:\n"
        "    raise SystemExit('loaded: ' + ', '.join(heavy))\n"
        "print('cold')\n"
    )
    environment = {
        **os.environ, "PYTHONPATH": os.pathsep.join(filter(None, [str(APP), os.environ.get("PYTHONPATH")])),
    }
    command = [sys.executable] + (["-O"] if optimize else []) + ["-c", script]
    completed = subprocess.run(command, capture_output=True, text=True, env=environment, timeout=120, check=False)
    assert completed.returncode == 0, completed.stdout + completed.stderr
    assert completed.stdout.strip().splitlines()[-1] == "cold"


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
