# test_orchestration_deliverables.py
"""The deliverables contract and planned image generation in Gather / Reason / Render plans.

Version: 0.261.140
Implemented in: 0.261.138
Single orchestration contract updated in: 0.261.139
Planner kind-specific fields and absent answer bindings updated in: 0.261.140

Uses the initialized headless harness (real bootstrap, planner, schema, executor, result
store, chat image persistence, rendering service and Office renderers) with the planning
and answering models replaced by offline replies and the image service call replaced by
fixed PNG bytes. Covers:

- deliverables validation: coverage by a capable step, render formats that must match the
  file deliverable, unavailable reasons that must match the server, the implicit answer,
  unambiguous linking, derived visuals, and the Image control;
- the server truth the planner receives and the single repair call;
- generate_image gating, the per-plan budget, persistence as an image message tied to the
  orchestrated answer, and the retained image-asset-v1 result;
- DOCX, PDF and PPTX files embedding exactly the images their source consumed, as
  renditions the Office renderers accept (a 4.7 MB PNG and a WEBP image never fail a file);
- answers owning the images they show: a retry reuses existing images, generates the
  missing one, keeps the earlier answer's images, and is not offered when it could only
  resend a refused request; approving a planned image's card never pays twice;
- scenarios: a CSV of states and capitals, a Word report with an image of each of the first
  three presidents, the same report without a file, and a failing search, image and render.
"""

import base64
import hashlib
import importlib
import io
import json
from copy import deepcopy

import pytest

from test_orchestration_harness_execution import harness, initialized_application  # noqa: F401
from test_support.orchestration_harness_execution import compose_step, decoded_frames, input_binding, render_step
from test_support.versioning import assert_app_version_at_least


IMAGE_SETTINGS = {
    "enable_image_generation": True,
    "image_gen_model": {"selected": [{"deploymentName": "gpt-image-1", "modelName": "gpt-image-1"}]},
    "azure_openai_image_gen_endpoint": "https://offline.invalid",
}
AVAILABLE = ["compose", "render_file", "generate_image", "web_search", "document_search"]
PRESIDENTS = (("washington", "George Washington"), ("adams", "John Adams"), ("jefferson", "Thomas Jefferson"))
REPORT_TEXT = (
    "# The first three presidents\n\n"
    "## George Washington\n\n[[image:washington]]\n\nWashington served from 1789 to 1797.\n\n"
    "## John Adams\n\n[[image:adams]]\n\nAdams served from 1797 to 1801.\n\n"
    "## Thomas Jefferson\n\n[[image:jefferson]]\n\nJefferson served from 1801 to 1809."
)


def test_version_includes_the_deliverables_contract():
    assert_app_version_at_least("0.261.138")


# ------------------------------------------------------------------------------------------
# Fixtures and builders
# ------------------------------------------------------------------------------------------

def _png(color, size=(48, 32)):
    # The image library is an execution dependency of the renderers under test.
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", size, color).save(buffer, format="PNG")
    return buffer.getvalue()


def _noise_png(size=(1536, 1024)):
    """A PNG that does not compress, like a detailed 1536x1024 illustration over 4 MB."""
    import random

    from PIL import Image

    buffer = io.BytesIO()
    Image.frombytes("RGB", size, random.Random(7).randbytes(size[0] * size[1] * 3)).save(buffer, format="PNG")
    return buffer.getvalue()


def _webp(color=(20, 40, 160, 180), size=(640, 480)):
    """A semi-transparent WEBP, which no Office renderer accepts as is."""
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGBA", size, color).save(buffer, format="WEBP")
    return buffer.getvalue()


def _data_url(mime_type, data):
    return f"data:{mime_type};base64," + base64.b64encode(data).decode("ascii")


def _enable_images(harness, monkeypatch, failures=None, images=None):
    """Real chat image persistence; only the image service request is doubled."""
    harness.settings.update(deepcopy(IMAGE_SETTINGS))
    generation = importlib.import_module("functions_image_generation")
    calls = []
    colors = iter(["red", "green", "blue", "orange", "purple", "teal", "gray", "olive"])

    def source(settings, prompt, size="", quality="", background=""):
        calls.append({"prompt": prompt, "size": size, "quality": quality, "background": background})
        failure = (failures or {}).get(len(calls))
        if failure is not None:
            raise failure
        custom = (images or {}).get(len(calls))
        return custom if custom is not None else _data_url("image/png", _png(next(colors)))

    monkeypatch.setattr(generation, "request_generated_image_source", source)
    return calls


def _deliverable(identifier, kind, description, *, requested="explicit", status="planned", **extra):
    return {
        "id": identifier, "kind": kind, "requested": requested, "description": description,
        "status": status, **extra,
    }


def _image_step(step_id, title, *, delivers=("portraits",), **arguments):
    return {
        "step_id": step_id, "capability_id": "generate_image", "title": f"Illustrate {title}",
        "arguments": {
            "prompt": f"An illustrated portrait of {title}, painted in a late 18th-century style.",
            "title": title, **arguments,
        },
        **({"delivers": list(delivers)} if delivers else {}),
    }


def _image_inputs(step_ids):
    return {step_id: {"binding": input_binding(step_id, "image"), "allow_partial": False} for step_id in step_ids}


def _report_step(inputs=None, *, delivers=("report",), outputs=None, basis="general_knowledge"):
    step = compose_step("report", inputs=inputs or {}, outputs=outputs or [{"name": "report", "kind": "markdown-v1"}])
    step["arguments"] = {
        "instruction": "Write a short report on the first three U.S. presidents.", "knowledge_basis": basis,
    }
    if delivers:
        step["delivers"] = list(delivers)
    return step


def _word_step(output_format="docx", *, delivers=("word_file",)):
    step = render_step("word", output_format, source="report", output="report", profile="prepared_report_v1")
    if delivers:
        step["delivers"] = list(delivers)
    return step


def _president_plan(*, with_file=True, with_search=False, output_format="docx"):
    deliverables = [
        _deliverable("report", "answer", "A report on the first three presidents"),
        _deliverable("portraits", "image", "An image of each president", quantity=3),
    ]
    steps = [_image_step(step_id, name) for step_id, name in PRESIDENTS]
    inputs = _image_inputs([step_id for step_id, _ in PRESIDENTS])
    basis = "general_knowledge"
    if with_search:
        steps.insert(0, {"step_id": "search", "capability_id": "web_search", "arguments": {
            "query": "First three presidents of the United States and their terms of office",
        }})
        inputs["findings"] = {"binding": input_binding("search", "prepared"), "allow_partial": False, "optional": True}
        basis = "sources_and_general_knowledge"
    steps.append(_report_step(inputs, basis=basis))
    if with_file:
        deliverables.append(_deliverable("word_file", "file", "The report as a file", format=output_format))
        steps.append(_word_step(output_format))
    return {"deliverables": deliverables, "steps": steps, "final_response": input_binding("report", "report")}


def _store(harness, plan, *, replies=()):
    """Save a validated plan as run-1, exactly as the harness saves its own fixtures."""
    run_store = importlib.import_module("functions_orchestration_runs")
    context = importlib.import_module("functions_orchestration_context")
    memory = importlib.import_module("functions_orchestration_memory")
    plan = deepcopy(plan)
    plan.update(run_id="run-1", plan_id="plan-1", turn_id="turn-1")
    normalized = context.normalize_history_message(harness.turn)
    turn_context = {
        "turn_id": "turn-1", "user_message": harness.turn["content"],
        "user_message_id": harness.turn["id"], "user_message_fingerprint": normalized["fingerprint"],
        "resolved_message": harness.turn["content"], "seeds": {}, "original_seeds": {},
        "answered_questions": [], "planning_token_usage": {},
        "conversation_context": context.build_conversation_snapshot([], harness.settings),
        "memory_audience": memory.validate_memory_audience(harness.conversation, "owner"),
        "memory_scope": None,
    }
    harness.replies = list(replies)
    return run_store.create_orchestration_run(
        plan, "owner", "conversation-1", turn_index=1, turn_context=turn_context,
    )


def _normalize(harness, raw, *, availability=None, image_selected=False, available=AVAILABLE):
    return harness.schema.normalize_plan(
        {"run_id": "run-1", "plan_id": "plan-1", "turn_id": "turn-1", **deepcopy(raw)},
        "conversation-1", "owner", settings=harness.settings, contract_version=2,
        available_capability_ids=available, deliverable_availability=availability,
        image_selected=image_selected, composition_profiles=harness.service_bindings.composition_profiles(),
    )


def _create(harness, raw, *, replies=()):
    return _store(harness, _normalize(harness, raw), replies=replies)


def _truth(harness, *, allowed=None, export_catalog=None, bindings=True):
    """The server truth a real plan request would compute for this caller."""
    registry = importlib.import_module("functions_orchestration_registry")
    deliverables = importlib.import_module("functions_orchestration_deliverables")
    unavailable = {}
    capabilities = registry.resolve_available_capabilities(
        harness.settings, allowed_ids=allowed, unavailable=unavailable, contract_version=2,
        export_catalog=export_catalog,
        request_context=harness.services().capability_request_bindings() if bindings else None,
    )
    return deliverables.build_deliverable_availability(
        harness.settings, capabilities=capabilities, unavailable=unavailable, export_catalog=export_catalog,
    )


def _plan_request(harness, replies, *, seeds=None, message="Please help with this request."):
    harness.replies = [json.dumps(reply) if isinstance(reply, dict) else reply for reply in replies]
    return harness.planner.plan_request(
        message, {}, "conversation-1", "owner", settings=harness.settings, seeds=seeds or {},
        contract_version=2, request_context=harness.services().capability_request_bindings(),
    )


def _saved_steps(harness):
    return {
        record["step_id"]: record for record in harness.steps.items.values()
        if record.get("run_id") == "run-1" and record.get("step_id") and "status" in record
    }


def _compose_call(harness):
    calls = [call for call in harness.model_calls if call["messages"][0]["content"].startswith("Prepare only")]
    assert len(calls) == 1
    return calls[0]["messages"]


def _files(harness, extension):
    return [
        record["data"] for (_container, name), record in harness.blobs.records.items()
        if name.endswith(extension) and "/images/" not in name
    ]


def _file_bytes(harness, extension):
    found = _files(harness, extension)
    assert len(found) == 1, sorted(name for _container, name in harness.blobs.records)
    return found[0]


def _image_messages(harness):
    return sorted(
        (deepcopy(message) for message in harness.messages.items.values() if message.get("role") == "image"),
        key=lambda message: message["metadata"]["image_proposal"]["visualId"],
    )


def _embedded_images(output_format, data):
    if output_format == "docx":
        from docx import Document

        return len(Document(io.BytesIO(data)).inline_shapes)
    if output_format == "pdf":
        import fitz

        with fitz.open(stream=data, filetype="pdf") as document:
            return sum(len(page.get_images()) for page in document)
    from pptx import Presentation

    deck = Presentation(io.BytesIO(data))
    return sum(1 for slide in deck.slides for shape in slide.shapes if shape.shape_type == 13)


# ------------------------------------------------------------------------------------------
# Validation
# ------------------------------------------------------------------------------------------

def test_a_plan_without_deliverables_gets_the_implicit_answer(harness):
    plan = _normalize(harness, {"steps": [compose_step()], "final_response": input_binding("prepare")})
    assert plan["deliverables"] == [{
        "id": "answer", "kind": "answer", "requested": "explicit", "status": "planned",
        "description": "An answer to your request.", "implicit": True,
    }]
    assert "delivers" not in plan["steps"][0]
    # Revalidating a saved plan derives the same implicit deliverable again.
    again = harness.schema.validate_plan(plan, settings=harness.settings, available_capability_ids=AVAILABLE)
    assert again["deliverables"] == plan["deliverables"]


def test_a_file_needs_a_render_step_in_its_own_format(harness):
    schema = harness.schema
    rows = compose_step("rows", outputs=[{"name": "rows", "kind": "records-v1", "columns": [
        {"name": "State", "value_type": "string", "nullable": False},
    ]}])
    render = render_step("capitals", "csv", source="rows", output="rows")
    render["arguments"]["options"] = {"columns": ["State"]}
    wanted = [_deliverable("sheet", "file", "A spreadsheet", format="xlsx")]

    with pytest.raises(schema.PlanValidationError) as mismatch:
        _normalize(harness, {"deliverables": wanted, "steps": [rows, {**render, "delivers": ["sheet"]}]})
    assert mismatch.value.code == "deliverables_invalid"
    assert "output_format must match" in str(mismatch.value)

    with pytest.raises(schema.PlanValidationError) as uncovered:
        _normalize(harness, {"deliverables": wanted, "steps": [compose_step()]})
    assert "No step delivers the planned file deliverable" in str(uncovered.value)

    with pytest.raises(schema.PlanValidationError) as undeclared:
        _normalize(harness, {"deliverables": [_deliverable("answer", "answer", "An answer")], "steps": [
            {**compose_step(), "delivers": ["answer"]}, rows, render,
        ], "final_response": input_binding("prepare")})
    assert "no deliverable declares" in str(undeclared.value)

    with pytest.raises(schema.PlanValidationError) as unknown:
        _normalize(harness, {"deliverables": wanted, "steps": [rows, {**render, "delivers": ["missing"]}]})
    assert "not a declared deliverable" in str(unknown.value)


def test_unambiguous_steps_are_linked_to_their_deliverables(harness):
    rows = compose_step("rows", outputs=[{"name": "rows", "kind": "records-v1", "columns": [
        {"name": "State", "value_type": "string", "nullable": False},
    ]}])
    render = render_step("capitals", "csv", source="rows", output="rows")
    render["arguments"]["options"] = {"columns": ["State"]}
    answer = compose_step("prepare")
    plan = _normalize(harness, {
        "deliverables": [
            _deliverable("csv_file", "file", "The CSV", format="CSV"),
            _deliverable("reply", "answer", "A short reply"),
            _deliverable("pictures", "image", "Two images", quantity=2),
        ],
        "steps": [rows, render, answer, _image_step("one", "One", delivers=()), _image_step("two", "Two", delivers=())],
        "final_response": input_binding("prepare"),
    })
    steps = {step["step_id"]: step for step in plan["steps"]}
    assert plan["deliverables"][0]["format"] == "csv"
    assert steps["capitals"]["delivers"] == ["csv_file"]
    assert steps["prepare"]["delivers"] == ["reply"]
    assert steps["one"]["delivers"] == steps["two"]["delivers"] == ["pictures"]


def test_unavailable_reasons_must_match_the_server(harness):
    schema = harness.schema
    truth = _truth(harness)
    assert truth["file"]["status"] == "available"
    assert truth["image"]["explicit"] == {"status": "unavailable", "reason": "image_generation_disabled"}
    answer = [{**compose_step(), "delivers": ["answer"]}]
    raw = {"steps": answer, "final_response": input_binding("prepare")}

    def planned(*extra):
        return {**raw, "deliverables": [_deliverable("answer", "answer", "An answer"), *extra]}

    with pytest.raises(schema.PlanValidationError) as claimed:
        _normalize(harness, planned(_deliverable(
            "doc", "file", "A Word file", format="docx", status="unavailable",
            unavailable_reason="file_rendering_unavailable",
        )), availability=truth)
    assert "The server can produce" in str(claimed.value)

    with pytest.raises(schema.PlanValidationError) as wrong:
        _normalize(harness, planned(_deliverable(
            "art", "image", "A picture", status="unavailable", unavailable_reason="image_budget_exceeded",
        )), availability=truth)
    assert '"image_generation_disabled", not "image_budget_exceeded"' in str(wrong.value)

    accepted = _normalize(harness, planned(
        _deliverable("art", "image", "A picture", status="unavailable", unavailable_reason="image_generation_disabled"),
        _deliverable("clip", "file", "A video", format="mp4", status="unavailable", unavailable_reason="format_not_supported"),
    ), availability=truth)
    unavailable = {item["id"]: item for item in accepted["deliverables"] if item["status"] == "unavailable"}
    assert unavailable["art"]["unavailable_message"] == "Image generation is turned off for this deployment."
    assert unavailable["clip"]["unavailable_message"] == "SimpleChat cannot create files in this format."
    # The answer step is told about them, so it neither promises nor apologizes for them.
    brief = accepted["steps"][0]["deliverable_context"]
    assert {entry["id"] for entry in brief if entry["relation"] == "unavailable"} == {"art", "clip"}

    with pytest.raises(schema.PlanValidationError) as planned_unsupported:
        _normalize(harness, planned(_deliverable("clip", "file", "A video", format="mp4")), availability=truth)
    assert "format_not_supported" in str(planned_unsupported.value)


def test_allowlisted_and_unadmitted_formats_carry_their_own_reasons(harness):
    truth = _truth(harness, allowed=["compose"])
    assert truth["file"]["reason"] == "capability_not_enabled_for_orchestration"
    assert truth["file"]["formats"]["docx"]["reason"] == "capability_not_enabled_for_orchestration"
    assert _truth(harness, bindings=False)["file"]["reason"] == "file_rendering_unavailable"
    registry = importlib.import_module("functions_orchestration_registry")
    catalog = [entry for entry in registry.resolve_admitted_export_catalog() if entry["format_id"] == "csv"]
    narrowed = _truth(harness, export_catalog=catalog)
    assert narrowed["file"]["formats"]["csv"]["status"] == "available"
    assert narrowed["file"]["formats"]["pdf"]["status"] == "unavailable"
    assert narrowed["file"]["formats"]["pdf"]["reason"] == "format_not_admitted"


def test_the_image_control_requires_an_explicit_image_deliverable(harness, monkeypatch):
    _enable_images(harness, monkeypatch)
    truth = _truth(harness)
    raw = {
        "deliverables": [_deliverable("answer", "answer", "An answer")],
        "steps": [{**compose_step(), "delivers": ["answer"]}], "final_response": input_binding("prepare"),
    }
    with pytest.raises(harness.schema.PlanValidationError) as missing:
        _normalize(harness, raw, availability=truth, image_selected=True)
    assert "Image control" in str(missing.value)
    raw["deliverables"].append(_deliverable("art", "image", "A picture"))
    raw["steps"].insert(0, _image_step("art_image", "A lighthouse", delivers=("art",)))
    raw["steps"][1]["inputs"] = _image_inputs(["art_image"])
    assert _normalize(harness, raw, availability=truth, image_selected=True)["status"] == "awaiting_approval"


def test_image_inputs_are_optional_and_charts_become_structured_visuals(harness, monkeypatch):
    _enable_images(harness, monkeypatch)
    step = _report_step(_image_inputs(["washington"]), basis="sources", delivers=("report", "chart"))
    plan = _normalize(harness, {
        "deliverables": [
            _deliverable("report", "answer", "The report"),
            _deliverable("chart", "chart", "A chart of term lengths"),
            _deliverable("portraits", "image", "A portrait"),
        ],
        "steps": [_image_step("washington", "George Washington"), step],
        "final_response": input_binding("report", "report"),
    }, availability=_truth(harness))
    steps = {step["step_id"]: step for step in plan["steps"]}
    assert steps["report"]["inputs"]["washington"]["optional"] is True
    assert steps["report"]["arguments"]["visuals"] == ["chart"]
    assert steps["washington"]["optional"] is True
    relations = {(entry["relation"], entry["id"]) for entry in steps["report"]["deliverable_context"]}
    assert relations == {("delivers", "report"), ("delivers", "chart"), ("images", "portraits")}


def test_planned_images_are_bounded_per_plan(harness, monkeypatch):
    _enable_images(harness, monkeypatch)
    truth = _truth(harness)
    registry = importlib.import_module("functions_orchestration_registry")
    budget = registry.MAX_GENERATED_IMAGES_PER_PLAN
    many = [_image_step(f"image_{index}", f"City {index}", delivers=("cities",)) for index in range(budget + 1)]
    with pytest.raises(harness.schema.PlanValidationError) as exceeded:
        _normalize(harness, {"deliverables": [_deliverable("cities", "image", "City images")], "steps": many})
    assert exceeded.value.code == "result_step_limit"

    rest = _deliverable(
        "more_cities", "image", "The remaining city images", status="unavailable",
        unavailable_reason="image_budget_exceeded",
    )
    partial = {"deliverables": [_deliverable("cities", "image", "City images", quantity=budget), rest],
               "steps": many[:budget]}
    plan = _normalize(harness, partial, availability=truth)
    assert [item["status"] for item in plan["deliverables"]] == ["planned", "unavailable"]
    # The budget is a reason only once the plan already uses all of it.
    with pytest.raises(harness.schema.PlanValidationError):
        _normalize(harness, {**partial, "deliverables": [
            _deliverable("cities", "image", "City images", quantity=budget - 1), rest,
        ], "steps": many[:budget - 1]}, availability=truth)


def test_image_options_must_be_ones_the_configured_model_supports(harness, monkeypatch):
    _enable_images(harness, monkeypatch)
    truth = _truth(harness)
    assert "1024x1024" in truth["image"]["explicit"]["options"]["sizes"]
    raw = {"deliverables": [_deliverable("art", "image", "A picture")],
           "steps": [_image_step("art_image", "A lighthouse", delivers=("art",), size="999x999")]}
    with pytest.raises(harness.schema.PlanValidationError) as unsupported:
        _normalize(harness, raw, availability=truth)
    assert "does not support" in str(unsupported.value)
    raw["steps"][0]["arguments"]["size"] = "1024x1024"
    assert _normalize(harness, raw, availability=truth)["steps"][0]["arguments"]["size"] == "1024x1024"


# ------------------------------------------------------------------------------------------
# Planning: server truth and the single repair call
# ------------------------------------------------------------------------------------------

def _csv_plan(*, render_format="csv"):
    return {
        "kind": "plan", "intent": {"summary": "List the U.S. states and capitals as a CSV file.", "complexity": "simple"},
        "assumptions": [],
        "deliverables": [_deliverable("capitals_csv", "file", "A CSV of the states and their capitals", format="csv")],
        "steps": [
            {
                "step_id": "rows", "capability_id": "compose", "title": "List states and capitals",
                "arguments": {"instruction": "List every U.S. state with its capital.",
                              "knowledge_basis": "general_knowledge"},
                "inputs": {}, "outputs": [{"name": "rows", "kind": "records-v1", "columns": [
                    {"name": "State", "value_type": "string", "nullable": False},
                    {"name": "Capital", "value_type": "string", "nullable": False},
                ]}],
            },
            {
                "step_id": "capitals", "capability_id": "render_file", "title": "Save the CSV file",
                "arguments": {
                    "file_name": f"us_states_capitals.{render_format}", "output_format": render_format,
                    "profile": "tabular_records_v1" if render_format == "csv" else "tabular_workbook_v1",
                    "options": {"columns": ["State", "Capital"], **(
                        {"sheet_name": "Capitals"} if render_format == "xlsx" else {}
                    )},
                },
                "inputs": {"source": {"binding": input_binding("rows", "rows"), "allow_partial": False}},
                "outputs": [], "delivers": ["capitals_csv"],
            },
        ],
    }


def test_the_planner_receives_the_server_truth_and_the_deliverables_contract(harness):
    kind, plan = _plan_request(harness, [_csv_plan()], message="create a csv of states and capitals")
    assert kind == "plan"
    system = harness.model_calls[0]["messages"][0]["content"]
    assert "List \"deliverables\" before the steps" in system
    assert "never state a limitation only in \"assumptions\"" in system
    assert harness.model_calls[0]["max_tokens"] == harness.planner.PLANNER_MAX_TOKENS == 4000
    context = json.loads(harness.model_calls[0]["messages"][1]["content"])
    truth = context["capability_availability"]["deliverables"]
    assert truth["file"]["formats"]["docx"]["embeds_images"] is True
    assert truth["file"]["formats"]["csv"]["source_kinds"] == ["records-v1"]
    assert any("return text and links only" in fact for fact in truth["facts"])
    assert {recipe["for"] for recipe in truth["recipes"]} >= {"CSV or XLSX file", "DOCX or PDF document", "PPTX deck"}
    assert set(truth["unavailable_reasons"]) >= {"format_not_admitted", "image_generation_disabled"}
    assert plan["deliverables"][0]["id"] == "capitals_csv"
    assert next(step for step in plan["steps"] if step["step_id"] == "capitals")["delivers"] == ["capitals_csv"]


def test_an_invalid_deliverables_plan_gets_exactly_one_repair_call(harness):
    kind, plan = _plan_request(
        harness, [_csv_plan(render_format="xlsx"), _csv_plan()], message="create a csv of states and capitals",
    )
    assert kind == "plan" and len(harness.model_calls) == 2
    repair = harness.model_calls[1]["messages"]
    assert repair[-2]["role"] == "assistant"
    assert repair[-1]["role"] == "user" and repair[-1]["content"].startswith("The server rejected that plan:")
    assert "output_format must match" in repair[-1]["content"]
    assert plan["token_usage"] == {"prompt_tokens": 14, "completion_tokens": 6, "total_tokens": 20}

    with pytest.raises(harness.planner.PlannerError) as failure:
        _plan_request(harness, [_csv_plan(render_format="xlsx"), _csv_plan(render_format="xlsx")])
    assert failure.value.reason == "invalid_plan_or_missing_requirement"
    assert failure.value.message == harness.planner.DELIVERABLES_FAILURE_MESSAGE


@pytest.mark.parametrize("comparison", [False, True])
def test_answer_declaration_fields_match_the_prompt_and_single_repair(harness, comparison):
    raw = {
        "kind": "plan",
        "deliverables": [_deliverable("answer", "answer", "The requested answer")],
        "steps": [{**compose_step("draft"), "delivers": ["answer"]}],
        "final_response": input_binding("draft"),
    }
    if comparison:
        harness.settings["document_action_capabilities"] = {
            "comparison": {"enabled": True, "chat_max_documents": 2},
        }
        raw["steps"].insert(0, {
            "step_id": "compare", "capability_id": "document_compare",
            "arguments": {
                "comparison_prompt": "Compare the two selected documents.",
                "left_document_id": "left", "right_document_ids": ["right"],
            },
        })
        raw["steps"][-1]["inputs"] = {
            "comparison": {"binding": input_binding("compare", "comparison"), "allow_partial": False},
        }
        raw["steps"][-1]["arguments"]["knowledge_basis"] = "sources"
    rejected = deepcopy(raw)
    rejected["deliverables"][0].update(format="markdown-v1", quantity=1)
    format_only_repair = deepcopy(rejected)
    format_only_repair["deliverables"][0].pop("format")
    available = ["compose", "document_compare"]
    with pytest.raises(harness.schema.PlanValidationError) as first:
        _normalize(harness, rejected, available=available)
    with pytest.raises(harness.schema.PlanValidationError) as second:
        _normalize(harness, format_only_repair, available=available)
    assert first.value.rule == "non_file_format"
    assert second.value.rule == "invalid_quantity"

    kind, plan = _plan_request(harness, [rejected, raw])
    assert kind == "plan" and len(harness.model_calls) == 2
    assert plan["deliverables"] == raw["deliverables"]
    prompt = harness.model_calls[0]["messages"][0]["content"]
    repair = harness.model_calls[1]["messages"][-1]["content"]
    assert "answer, chart, and diagram: omit BOTH format and quantity" in prompt
    assert "quantity is optional and counts files, not records, rows, pages, or answers" in prompt
    assert "Recheck every deliverable's kind-specific fields" in repair
    assert "not just the first field reported above" in repair
    if comparison:
        selected = harness.schema.plan_document_ids(plan)
        assert set(selected) == {"left", "right"}
        assert plan["steps"][-1]["depends_on"] == ["compare"]


def test_null_final_response_is_canonical_absence_for_file_only_work(harness):
    raw = {**_csv_plan(), "final_response": None}
    original = deepcopy(raw)
    kind, plan = _plan_request(harness, [raw])
    assert kind == "plan" and len(harness.model_calls) == 1
    assert "final_response" not in plan
    assert plan["deliverables"][0]["kind"] == "file"
    assert [step["capability_id"] for step in plan["steps"]] == ["compose", "render_file"]
    revalidated = harness.schema.validate_plan(
        {**plan, "final_response": None}, settings=harness.settings,
        available_capability_ids=AVAILABLE,
    )
    assert "final_response" not in revalidated
    assert raw == original


@pytest.mark.parametrize("invalid", [False, 0, "", [], {}])
def test_only_null_is_absent_and_malformed_final_bindings_still_fail(harness, invalid):
    with pytest.raises(harness.schema.PlanValidationError):
        _normalize(harness, {**_csv_plan(), "final_response": invalid})


def test_null_final_response_cannot_drop_a_declared_answer(harness):
    raw = {
        "deliverables": [_deliverable("answer", "answer", "The required answer")],
        "steps": [{**compose_step(), "delivers": ["answer"]}],
        "final_response": None,
    }
    with pytest.raises(harness.schema.PlanValidationError) as failure:
        _normalize(harness, raw)
    assert failure.value.code == "deliverables_invalid"
    assert failure.value.rule == "answer_producer_mismatch"


def test_the_image_control_asks_the_planner_for_explicit_images(harness, monkeypatch):
    _enable_images(harness, monkeypatch)
    reply = {"kind": "plan", "intent": {"summary": "Illustrate a lighthouse."}, "assumptions": [], **{
        "deliverables": [_deliverable("art", "image", "A lighthouse illustration")],
        "steps": [_image_step("art_image", "A lighthouse", delivers=("art",))],
    }}
    kind, plan = _plan_request(harness, [reply], seeds={"image_generation": True})
    context = json.loads(harness.model_calls[0]["messages"][1]["content"])
    assert kind == "plan"
    assert context["user_selected"]["images"] is True
    assert "image_proposals" not in context["user_selected"]
    assert context["capability_availability"]["deliverables"]["image"]["explicit"]["status"] == "available"
    assert [step["capability_id"] for step in plan["steps"]] == ["generate_image"]


# ------------------------------------------------------------------------------------------
# generate_image
# ------------------------------------------------------------------------------------------

def test_image_generation_is_gated_by_the_image_settings(harness, monkeypatch):
    registry = importlib.import_module("functions_orchestration_registry")
    unavailable = {}
    registry.resolve_available_capabilities(harness.settings, unavailable=unavailable, contract_version=2)
    assert unavailable["generate_image"] == "feature_disabled"
    harness.settings["enable_image_generation"] = True
    unavailable = {}
    registry.resolve_available_capabilities(harness.settings, unavailable=unavailable, contract_version=2)
    assert unavailable["generate_image"] == "image_generation_unavailable"
    assert _truth(harness)["image"]["explicit"]["reason"] == "image_generation_unavailable"
    _enable_images(harness, monkeypatch)
    capability = next(
        item for item in registry.resolve_available_capabilities(harness.settings, contract_version=2)
        if item["id"] == "generate_image"
    )
    assert capability["inputs"]["properties"]["size"]["enum"] == ["1024x1024", "1536x1024", "1024x1536"]
    assert capability["max_per_plan"] == registry.MAX_GENERATED_IMAGES_PER_PLAN
    # Only Render and planned images may publish; every other Gather or Reason step stays unable to.
    publishers = {
        item["id"] for item in registry.capabilities_for_contract(2) if item.get("publishes_generated_images")
    }
    assert publishers == {"generate_image"}


def test_a_generated_image_is_saved_with_the_answer_that_shows_it(harness, monkeypatch):
    calls = _enable_images(harness, monkeypatch)
    checkpoints = importlib.import_module("functions_orchestration_checkpoints")
    _create(harness, {
        "deliverables": [
            _deliverable("reply", "answer", "A short caption"),
            _deliverable("art", "image", "A lighthouse illustration"),
        ],
        "steps": [
            _image_step("lighthouse", "A lighthouse at dusk", delivers=("art",), size="1024x1024"),
            {**_report_step(_image_inputs(["lighthouse"]), delivers=("reply",))},
        ],
        "final_response": input_binding("report", "report"),
    }, replies=["Here is the lighthouse.\n\n[[image:lighthouse]]"])

    frames = decoded_frames(harness.prepare().execute())
    saved = harness.read()
    answer = harness.assistant_messages()[0]
    [image] = _image_messages(harness)
    proposal = image["metadata"]["image_proposal"]
    stored = harness.blobs.records[("harness-chat", image["blob_path"])]["data"]
    contracts = importlib.import_module("functions_orchestration_result_contracts")
    task = contracts.TaskResult.from_dict(saved["task_results"]["lighthouse"])
    value = harness.services().results.open_result(task.output("image"), require_current_sources=True).read_value()

    assert saved["status"] == "completed", saved.get("failure")
    assert calls == [{"prompt": calls[0]["prompt"], "size": "1024x1024", "quality": "", "background": ""}]
    assert proposal["visualId"] == "lighthouse" and proposal["title"] == "A lighthouse at dusk"
    assert proposal["source_assistant_message_id"] == answer["id"] == checkpoints.orchestration_answer_message_id("run-1")
    assert proposal["orchestration"] == {"run_id": "run-1", "step_id": "lighthouse"}
    assert image["file_content_source"] == "blob" and harness.blobs.image_uploads == 1
    assert task.output("image").kind == "image-asset-v1"
    assert value["message_id"] == image["id"] and value["ai_generated"] is True
    assert value["content_sha256"] == hashlib.sha256(stored).hexdigest()
    assert '"visualId": "lighthouse"' in answer["content"] and "```simpleimage" in answer["content"]
    assert "*AI-generated illustration: A lighthouse at dusk*" in answer["content"]
    assert answer["metadata"]["orchestration"]["generated_images"] == [
        {"visual_id": "lighthouse", "message_id": image["id"]},
    ]
    # The browser loads the image messages from the terminal frame, not only from the thread.
    [done] = [frame for frame in frames if frame.get("done")]
    assert done["generated_images"] == answer["metadata"]["orchestration"]["generated_images"]
    assert done["metadata"]["orchestration"]["generated_images"] == done["generated_images"]
    payload = json.loads(_compose_call(harness)[-1]["content"])
    assert payload["images"] == [{"token": "[[image:lighthouse]]", "title": "A lighthouse at dusk"}]
    assert "lighthouse" not in payload["inputs"]


def test_image_prompts_pass_the_chat_output_check_before_generation(harness, monkeypatch):
    calls = _enable_images(harness, monkeypatch)
    checks = importlib.import_module("functions_chat_content_checks")
    original = checks.check_chat_content
    blocked = checks.ChatContentDecision("chat_output", "findings", "block", {}, "removed")

    def check(text, checkpoint, **kwargs):
        # Only the planned image prompt is refused; the published reply is checked as usual.
        return blocked if "illustrated portrait of A lighthouse" in text else original(text, checkpoint, **kwargs)

    monkeypatch.setattr(checks, "check_chat_content", check)
    _create(harness, {
        "deliverables": [_deliverable("art", "image", "A picture")],
        "steps": [_image_step("art_image", "A lighthouse", delivers=("art",))],
    })
    harness.prepare().execute()
    record = _saved_steps(harness)["art_image"]
    assert calls == [] and _image_messages(harness) == []
    assert record["status"] == "failed" and record["failure"]["code"] == "image_content_refused"
    assert harness.read()["status"] == "failed"


# ------------------------------------------------------------------------------------------
# Images in files
# ------------------------------------------------------------------------------------------

def _image_media(output_format, data):
    """Each embedded image part of a DOCX or PPTX as (extension, size in bytes)."""
    import zipfile

    folder = "word/media/" if output_format == "docx" else "ppt/media/"
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return sorted(
            (info.filename.rsplit(".", 1)[-1].lower(), info.file_size)
            for info in archive.infolist() if info.filename.startswith(folder)
        )


def test_documents_embed_renditions_the_office_renderers_accept():
    rendering = importlib.import_module("functions_orchestration_rendering")
    output_store = importlib.import_module("functions_orchestration_output_store")
    from PIL import Image

    def opened(data):
        with Image.open(io.BytesIO(data)) as image:
            return image.format, image.mode, image.size

    limit = 4 * 1024 * 1024
    small = _png("red")
    assert rendering.document_image_bytes(small, max_bytes=limit, max_pixels=10_000_000) is small
    large = _noise_png()
    assert len(large) > limit
    fitted = rendering.document_image_bytes(large, max_bytes=limit, max_pixels=10_000_000)
    assert len(fitted) <= limit and opened(fitted) == ("JPEG", "RGB", (1536, 1024))
    # WEBP is re-encoded losslessly when it fits, keeping its transparency.
    assert opened(rendering.document_image_bytes(_webp(), max_bytes=limit, max_pixels=10_000_000)) == (
        "PNG", "RGBA", (640, 480),
    )
    # Images share the renderer's pixel budget, so each one fits its share.
    shared = rendering.document_image_bytes(small, max_bytes=limit, max_pixels=1_000)
    assert opened(shared)[0] == "PNG" and opened(shared)[2][0] * opened(shared)[2][1] <= 1_000
    with pytest.raises(output_store.OutputError):
        rendering.document_image_bytes(b"not an image", max_bytes=limit, max_pixels=10_000_000)


def test_the_image_asset_contract_admits_only_what_generation_validates():
    contracts = importlib.import_module("functions_orchestration_result_contracts")
    image_edit = importlib.import_module("functions_image_edit")
    # Generation decodes each image within these bounds; files embed a rendition of it.
    assert contracts.MAX_IMAGE_ASSET_BYTES == image_edit.MAX_SOURCE_IMAGE_BYTES
    assert contracts.MAX_IMAGE_ASSET_PIXELS == image_edit.MAX_IMAGE_PIXELS
    assert contracts._IMAGE_ASSET_MIME_TYPES == {"image/png", "image/jpeg", "image/webp"}


def test_a_large_png_and_a_webp_never_fail_a_word_file(harness, monkeypatch):
    """A 4.7 MB PNG and a WEBP image are embedded as renditions; chat keeps the originals."""
    large, webp = _noise_png(), _webp()
    _enable_images(harness, monkeypatch, images={
        1: _data_url("image/png", large), 2: _data_url("image/webp", webp),
    })
    _create(harness, _president_plan(), replies=[REPORT_TEXT])

    harness.prepare().execute()
    saved = harness.read()
    data = _file_bytes(harness, ".docx")
    stored = {
        image["metadata"]["image_proposal"]["visualId"]: harness.blobs.records[("harness-chat", image["blob_path"])]["data"]
        for image in _image_messages(harness)
    }
    media = _image_media("docx", data)

    assert saved["status"] == "completed", saved.get("failure")
    assert _embedded_images("docx", data) == 3
    assert stored["washington"] == large and stored["adams"] == webp
    assert len(media) == 3 and all(
        extension in {"png", "jpeg", "jpg"} and size <= 4 * 1024 * 1024 for extension, size in media
    ), media


def test_an_image_stays_linked_to_the_answer_that_generated_it(harness, monkeypatch):
    """Answers own the list of images they show; publishing never moves an image elsewhere."""
    _enable_images(harness, monkeypatch)
    checkpoints = importlib.import_module("functions_orchestration_checkpoints")
    _create(harness, _president_plan(with_file=False), replies=[REPORT_TEXT])

    harness.prepare().execute()
    answer = harness.assistant_messages()[0]

    assert answer["id"] == checkpoints.orchestration_answer_message_id("run-1")
    assert {
        image["id"]: image["metadata"]["image_proposal"]["source_assistant_message_id"]
        for image in _image_messages(harness)
    } == {entry["message_id"]: answer["id"] for entry in answer["metadata"]["orchestration"]["generated_images"]}
    assert not hasattr(harness.execution.HarnessExecution, "_link_generated_images")


def test_approving_a_planned_image_card_again_returns_the_saved_image(harness):
    generation = importlib.import_module("functions_image_generation")
    answer = {
        "id": "assistant_orchestration_1", "conversation_id": "conversation-1", "role": "assistant",
        "metadata": {"orchestration": {"generated_images": [
            "not-an-entry", {"visual_id": "washington", "message_id": 7},
            {"visual_id": "washington", "message_id": "conversation-1_image_1"},
        ]}},
    }
    image = {
        "id": "conversation-1_image_1", "conversation_id": "conversation-1", "role": "image",
        "content": "/api/image/conversation-1_image_1", "model_deployment_name": "gpt-image-1",
        "metadata": {"image_proposal": {"visualId": "washington", "title": "George Washington"}},
    }
    stored = {"conversation-1_image_1": image}
    read = lambda message_id: deepcopy(stored.get(message_id))
    proposal = {"visualId": "washington", "prompt": "A portrait"}

    assert generation.find_planned_proposal_image(answer, proposal, read)["id"] == image["id"]
    assert generation.find_planned_proposal_image(answer, {"visualId": "adams", "prompt": "x"}, read) is None
    assert generation.find_planned_proposal_image({**answer, "metadata": {}}, proposal, read) is None
    assert generation.find_planned_proposal_image(None, proposal, read) is None
    for changed in (
        {"metadata": {"image_proposal": {"visualId": "washington"}, "is_deleted": True}},
        {"conversation_id": "conversation-2"},
        {"metadata": {"image_proposal": {"visualId": "jefferson"}}},
        {"role": "assistant"},
    ):
        stored["conversation-1_image_1"] = {**image, **changed}
        assert generation.find_planned_proposal_image(answer, proposal, read) is None
    stored.clear()
    assert generation.find_planned_proposal_image(answer, proposal, read) is None


def _load_app_functions(file_name, names, namespace):
    """Execute the real selected functions of an application module without its startup graph."""
    import ast
    from pathlib import Path

    generation = importlib.import_module("functions_image_generation")
    source = Path(generation.__file__).with_name(file_name).read_text(encoding="utf-8-sig")
    tree = ast.parse(source)
    nodes = []
    for name in names:
        node = next(item for item in ast.walk(tree) if isinstance(item, ast.FunctionDef) and item.name == name)
        node.decorator_list = []
        nodes.append(node)
    exec(compile(ast.Module(body=nodes, type_ignores=[]), file_name, "exec"), namespace)
    return namespace


def test_the_approval_route_never_pays_twice_for_a_planned_image():
    """The real handler returns the answer's saved image and never calls the image service."""
    from unittest.mock import Mock

    from flask import Flask, jsonify, request

    generation = importlib.import_module("functions_image_generation")
    answer = {
        "id": "assistant_orchestration_1", "conversation_id": "conversation-1", "role": "assistant",
        "metadata": {"orchestration": {"generated_images": [
            {"visual_id": "washington", "message_id": "conversation-1_image_1"},
        ]}},
    }
    image = {
        "id": "conversation-1_image_1", "conversation_id": "conversation-1", "role": "image",
        "content": "/api/image/conversation-1_image_1", "prompt": "A portrait",
        "model_deployment_name": "gpt-image-1",
        "metadata": {"image_proposal": {"visualId": "washington", "title": "George Washington"}},
    }
    missing = type("NotFoundForTest", (Exception,), {})

    class Messages:
        def read_item(self, item, partition_key):
            found = {answer["id"]: answer, image["id"]: image}.get(item)
            if found is None or partition_key != "conversation-1":
                raise missing()
            return deepcopy(found)

    generate = Mock(side_effect=AssertionError("The image service must not be called."))
    namespace = _load_app_functions("route_backend_chats.py", ("generate_image_from_proposal",), {
        "request": request, "jsonify": jsonify, "logging": importlib.import_module("logging"),
        "datetime": importlib.import_module("datetime").datetime,
        "get_settings": lambda: {"enable_image_generation": True},
        "image_generation_is_enabled": lambda settings: True,
        "get_current_user_id": lambda: "owner", "get_current_user_info": lambda: {},
        "_authorize_personal_conversation_access": Mock(return_value={"id": "conversation-1", "title": "Presidents"}),
        "normalize_image_proposal": generation.normalize_image_proposal,
        "find_planned_proposal_image": generation.find_planned_proposal_image,
        "generate_chat_image_message": generate, "cosmos_messages_container": Messages(),
        "cosmos_conversations_container": Mock(), "invalidate_conversation_cache_for_item": Mock(),
        "CosmosResourceNotFoundError": missing, "AIConnectionError": type("AIConnectionForTest", (Exception,), {}),
        "image_generation_error_response": generation.image_generation_error_response,
        "image_generation_error_log_context": generation.image_generation_error_log_context,
        "log_event": Mock(),
    })
    app = Flask(__name__)
    with app.test_request_context(json={
        "conversation_id": "conversation-1", "assistant_message_id": answer["id"],
        "proposal": {"visualId": "washington", "title": "George Washington", "prompt": "A portrait"},
    }):
        response, status = namespace["generate_image_from_proposal"]()
    payload = response.get_json()

    assert status == 200 and generate.call_count == 0
    assert payload["already_generated"] is True and payload["message_id"] == image["id"]
    assert payload["image_message"]["id"] == image["id"]
    assert payload["image_message"]["content"] == image["content"]
    namespace["cosmos_conversations_container"].replace_item.assert_not_called()


def test_a_conversation_export_includes_the_images_each_answer_lists():
    """A retry's answer exports the image it reused, and the earlier answer keeps it too."""
    from typing import Any, Dict, List, Optional

    image = {
        "id": "image_w", "conversation_id": "conversation-1", "role": "image",
        "metadata": {"image_proposal": {"visualId": "washington", "source_assistant_message_id": "assistant_1"}},
    }
    other = {
        "id": "image_x", "conversation_id": "conversation-1", "role": "image",
        "metadata": {"image_proposal": {"visualId": "other", "source_assistant_message_id": "assistant_9"}},
    }

    class Messages:
        def query_items(self, query, parameters, partition_key):
            return [deepcopy(image), deepcopy(other)]

    namespace = _load_app_functions("route_backend_conversation_export.py", (
        "_attach_generated_image_proposal_assets", "_orchestration_generated_image_message_ids",
        "_load_generated_image_proposal_assets",
    ), {
        "Any": Any, "Dict": Dict, "List": List, "Optional": Optional, "cosmos_messages_container": Messages(),
        "debug_print": lambda *args, **kwargs: None,
        "_build_export_image_asset_from_message": lambda conversation_id, image_message, proposal: {
            "message_id": image_message["id"], "visual_id": proposal["visualId"],
        },
    })
    attach = namespace["_attach_generated_image_proposal_assets"]
    earlier = {"id": "assistant_1", "role": "assistant", "content": "", "metadata": {}}
    retried = {"id": "assistant_2", "role": "assistant", "content": "", "metadata": {"orchestration": {
        "generated_images": [{"visual_id": "washington", "message_id": "image_w"}, {"message_id": 7}],
    }}}

    assert attach(earlier, "conversation-1")["_export_generated_image_assets"] == [
        {"message_id": "image_w", "visual_id": "washington"},
    ]
    assert attach(retried, "conversation-1")["_export_generated_image_assets"] == [
        {"message_id": "image_w", "visual_id": "washington"},
    ]
    assert "_export_generated_image_assets" not in attach({**retried, "role": "user"}, "conversation-1")


@pytest.mark.parametrize("output_format", ["docx", "pdf", "pptx"])
def test_files_embed_exactly_the_images_their_source_consumed(harness, monkeypatch, output_format):
    # A WEBP image, which no Office renderer accepts as is, is embedded as a PNG rendition.
    _enable_images(harness, monkeypatch, images={1: _data_url("image/webp", _webp())})
    deck = output_format == "pptx"
    report = _report_step(
        _image_inputs(["washington"]), delivers=() if deck else ("report",),
        outputs=[{"name": "report", "kind": "structured-v1", "profile": "prepared_slide_deck_v1"}] if deck else None,
    )
    render = render_step(
        "word", output_format, source="report", output="report",
        profile="prepared_slide_deck_v1" if deck else "prepared_report_v1",
    )
    reply = json.dumps({"report": {
        "schema_version": "prepared_slide_deck_v1", "slide_count": 1, "slides": [{
            "layout": "title_and_content", "title": "George Washington", "shapes": [{
                "type": "image", "box": {"left": 1, "top": 1.6, "width": 3, "height": 2},
                "source": "asset:washington", "alt": "Portrait",
            }],
        }],
    }}) if deck else "# George Washington\n\n[[image:washington]]\n\nThe first president."
    _create(harness, {
        "deliverables": [
            *([] if deck else [_deliverable("report", "answer", "A short report")]),
            _deliverable("portraits", "image", "A portrait"),
            _deliverable("word_file", "file", "The file", format=output_format),
        ],
        "steps": [_image_step("washington", "George Washington"), report, {**render, "delivers": ["word_file"]}],
        **({} if deck else {"final_response": input_binding("report", "report")}),
    }, replies=[reply])

    harness.prepare().execute()
    saved = harness.read()
    data = _file_bytes(harness, f".{output_format}")

    assert saved["status"] == "completed", saved.get("failure")
    assert _embedded_images(output_format, data) == 1
    if output_format != "pdf":
        assert [extension for extension, _size in _image_media(output_format, data)] == ["png"]


def test_a_deck_that_did_not_place_its_image_gets_a_slide_for_it(harness, monkeypatch):
    _enable_images(harness, monkeypatch)
    report = _report_step(
        _image_inputs(["washington"]), delivers=(),
        outputs=[{"name": "report", "kind": "structured-v1", "profile": "prepared_slide_deck_v1"}],
    )
    render = render_step("word", "pptx", source="report", output="report", profile="prepared_slide_deck_v1")
    _create(harness, {
        "deliverables": [
            _deliverable("portraits", "image", "A portrait"),
            _deliverable("word_file", "file", "The deck", format="pptx"),
        ],
        "steps": [_image_step("washington", "George Washington"), report, {**render, "delivers": ["word_file"]}],
    }, replies=[json.dumps({"report": {
        "schema_version": "prepared_slide_deck_v1", "slide_count": 1, "slides": [{
            "layout": "title_and_content", "title": "The first president", "shapes": [{
                "type": "text_box", "box": {"left": 0.5, "top": 1.6, "width": 6, "height": 1},
                "paragraphs": [{"text": "George Washington served from 1789 to 1797."}],
            }],
        }],
    }})])

    harness.prepare().execute()

    assert harness.read()["status"] == "completed"
    assert _embedded_images("pptx", _file_bytes(harness, ".pptx")) == 1


def test_a_file_cannot_embed_an_image_its_source_did_not_consume(harness, monkeypatch):
    _enable_images(harness, monkeypatch)
    _create(harness, {
        "deliverables": [
            _deliverable("report", "answer", "A short report"),
            _deliverable("portraits", "image", "A portrait"),
            _deliverable("word_file", "file", "The report as a Word file", format="docx"),
        ],
        # The image exists in this run, but the report was not prepared from it.
        "steps": [_image_step("washington", "George Washington"), _report_step(), _word_step()],
        "final_response": input_binding("report", "report"),
    }, replies=["# Report\n\n![Portrait](asset:washington)\n\nThe first president."])

    harness.prepare().execute()
    saved = harness.read()
    [output] = saved["outputs"]

    assert saved["status"] == "failed"
    assert output["state"] == "failed" and output["error_code"] == "output_source_unavailable"
    assert harness.blobs.file_uploads == 0


def test_the_image_reader_refuses_a_deleted_or_moved_image(harness, monkeypatch):
    bootstrap = importlib.import_module("functions_orchestration_bootstrap")
    output_store = importlib.import_module("functions_orchestration_output_store")
    data = _png("red")
    harness.messages.create_item({
        "id": "conversation-1_image_1", "conversation_id": "conversation-1", "role": "image",
        "file_content_source": "blob", "blob_container": "harness-chat",
        "blob_path": "owner/conversation-1/images/conversation-1_image_1/art.png",
        "metadata": {"image_proposal": {"visualId": "art"}},
    })
    harness.blobs.records[("harness-chat", "owner/conversation-1/images/conversation-1_image_1/art.png")] = {
        "data": data, "metadata": {}, "content_type": "image/png", "etag": "etag-image",
    }
    asset = {"message_id": "conversation-1_image_1", "asset_id": "art", "size_bytes": len(data)}
    reader = bootstrap.build_image_asset_reader("owner", "conversation-1")
    assert reader(asset) == data
    message = harness.messages.read_item("conversation-1_image_1", "conversation-1")
    message["metadata"]["is_deleted"] = True
    harness.messages.upsert_item(message)
    with pytest.raises(output_store.OutputUnavailableError):
        reader(asset)
    message["metadata"].pop("is_deleted")
    message["blob_path"] = "someone-else/conversation-1/images/conversation-1_image_1/art.png"
    harness.messages.upsert_item(message)
    with pytest.raises(output_store.OutputUnavailableError):
        reader(asset)


# ------------------------------------------------------------------------------------------
# Scenarios
# ------------------------------------------------------------------------------------------

def test_scenario_a_csv_of_states_and_capitals_is_a_real_file(harness):
    kind, plan = _plan_request(harness, [_csv_plan()], message="create a csv of states and capitals")
    assert kind == "plan"
    _store(harness, plan, replies=[json.dumps({"rows": [
        {"State": "Alabama", "Capital": "Montgomery"}, {"State": "Alaska", "Capital": "Juneau"},
    ]})])

    harness.prepare().execute()
    saved = harness.read()
    guidance = "\n".join(message["content"] for message in _compose_call(harness) if message["role"] == "system")
    csv_text = _file_bytes(harness, ".csv").decode("utf-8")

    assert saved["status"] == "completed", saved.get("failure")
    assert 'A later step saves output "rows" as us_states_capitals.csv, a CSV file.' in guidance
    assert "Never say files cannot be created." in guidance
    assert csv_text.splitlines() == ["State,Capital", "Alabama,Montgomery", "Alaska,Juneau"]
    assert r"us\_states\_capitals\.csv: ready" in harness.assistant_messages()[0]["content"]
    assert "Delivery notes" not in harness.assistant_messages()[0]["content"]


def _succeeding_search(harness, monkeypatch, notes):
    """A completed search whose findings are retained without external-source admission."""
    executor = importlib.import_module("functions_orchestration_executor")
    schema = importlib.import_module("functions_orchestration_schema")
    results = importlib.import_module("functions_orchestration_results")
    contracts = importlib.import_module("functions_orchestration_result_contracts")
    original = executor._dependency_adapter

    def search(step, context, **kwargs):
        return schema.build_step_result(status=schema.STEP_STATUS_COMPLETED, summary="Found sources.", notes=[notes])

    def retain(step, context, result, *, source_manifest):
        complete = contracts.Completeness(
            "complete", 1, 1, contracts.Coverage(1, 1, "work_units"), "valid", ("test_search",), (),
        )
        return context.result_service.persist_task_result(
            producer=context.result_producer(step), role="gather", status="complete",
            outputs=[results.NamedOutput("prepared", "structured-v1", {"notes": result["notes"]}, complete)],
            sources=[], origin="generated", guard_token=context.result_guard_token_for_step(step["step_id"]),
            input_fingerprint=context.result_input_fingerprint_for_step(step["step_id"]),
        )

    monkeypatch.setattr(
        executor, "_dependency_adapter",
        lambda capability_id: search if capability_id == "web_search" else original(capability_id),
    )
    monkeypatch.setattr(executor, "retain_gather_result", retain)
    harness.settings["enable_web_search"] = True


def _failing_search(harness, monkeypatch):
    executor = importlib.import_module("functions_orchestration_executor")
    schema = importlib.import_module("functions_orchestration_schema")
    original = executor._dependency_adapter

    def failing(step, context, **kwargs):
        failure = schema.build_failure("provider_not_configured")
        return schema.build_step_result(status=schema.STEP_STATUS_FAILED, failure=failure, summary=failure["message"])

    monkeypatch.setattr(
        executor, "_dependency_adapter",
        lambda capability_id: failing if capability_id == "web_search" else original(capability_id),
    )
    harness.settings["enable_web_search"] = True


def test_scenario_a_word_report_with_an_image_of_each_president(harness, monkeypatch):
    calls = _enable_images(harness, monkeypatch)
    _succeeding_search(harness, monkeypatch, "Washington 1789-1797; Adams 1797-1801; Jefferson 1801-1809.")
    _create(harness, _president_plan(with_search=True), replies=[REPORT_TEXT])

    harness.prepare().execute()
    saved = harness.read()
    answer = harness.assistant_messages()[0]["content"]
    images = _image_messages(harness)
    payload = json.loads(_compose_call(harness)[-1]["content"])

    assert saved["status"] == "completed", saved.get("failure")
    assert len(calls) == 3 and [image["metadata"]["image_proposal"]["visualId"] for image in images] == [
        "adams", "jefferson", "washington",
    ]
    assert payload["inputs"]["findings"]["value"]["notes"][0].startswith("Washington 1789")
    assert [image["token"] for image in payload["images"]] == [
        "[[image:washington]]", "[[image:adams]]", "[[image:jefferson]]",
    ]
    assert _embedded_images("docx", _file_bytes(harness, ".docx")) == 3
    assert answer.count("```simpleimage") == 3 and "[[image:" not in answer and "asset:" not in answer
    assert answer.startswith("# The first three presidents")
    assert r"word\.docx: ready" in answer and "Delivery notes" not in answer


def test_scenario_a_report_with_images_and_no_file(harness, monkeypatch):
    _enable_images(harness, monkeypatch)
    _create(harness, _president_plan(with_file=False), replies=[REPORT_TEXT])

    harness.prepare().execute()
    saved = harness.read()
    answer = harness.assistant_messages()[0]["content"]

    assert saved["status"] == "completed", saved.get("failure")
    assert harness.blobs.file_uploads == 0 and saved["outputs"] == []
    assert answer.count("```simpleimage") == 3 and "Files:" not in answer


def test_scenario_a_failed_search_still_writes_the_report_from_general_knowledge(harness, monkeypatch):
    _enable_images(harness, monkeypatch)
    _failing_search(harness, monkeypatch)
    _create(harness, _president_plan(with_search=True), replies=[REPORT_TEXT + "\n\nWeb sources were unavailable."])

    harness.prepare().execute()
    saved = harness.read()
    payload = json.loads(_compose_call(harness)[-1]["content"])

    assert saved["status"] == "completed", saved.get("failure")
    assert payload["unavailable_inputs"][0]["name"] == "findings"
    assert _saved_steps(harness)["search"]["status"] == "failed"
    assert _embedded_images("docx", _file_bytes(harness, ".docx")) == 3


def test_scenario_a_failed_image_is_reported_and_never_delivered(harness, monkeypatch):
    route = importlib.import_module("functions_image_api_route")
    refused = route.ImageGenerationError("refused", "image_content_refused", 400)
    _enable_images(harness, monkeypatch, failures={2: refused})
    _create(harness, _president_plan(), replies=[REPORT_TEXT])

    harness.prepare().execute()
    saved = harness.read()
    answer = harness.assistant_messages()[0]["content"]
    payload = json.loads(_compose_call(harness)[-1]["content"])
    steps = _saved_steps(harness)

    assert steps["adams"]["status"] == "failed" and steps["adams"]["failure"]["code"] == "image_content_refused"
    assert steps["report"]["status"] == "completed" and steps["word"]["status"] == "completed"
    assert payload["unavailable_images"][0]["name"] == "adams"
    assert saved["status"] == "failed" and saved["outcome"] == "partial"
    assert "- Not delivered: An image of each president. 2 of 3 images were generated." in answer
    assert answer.count("```simpleimage") == 2
    assert _embedded_images("docx", _file_bytes(harness, ".docx")) == 2


def _authorize(harness):
    return lambda: harness.bootstrap.read_owned_conversation("owner", "conversation-1")


def _retry_request(harness, execution, submission_id):
    """The checkpoint retry of run-1, exactly as the retry route publishes it."""
    from copy import copy

    parent = harness.read()
    probe = copy(execution.context)
    probe.result_service = harness.services().results
    return harness.recovery.prepare_retry(
        "run-1", "owner", {
            "conversation_id": "conversation-1", "submission_id": submission_id,
            "expected_version": parent["recovery_version"],
        },
        authorize=_authorize(harness), message_container=harness.messages,
        validate=lambda current: harness.recovery.validate_resume(
            current, probe, harness.settings, _authorize(harness), source_run_id=current["id"],
        ),
    )


def _claim_retry(harness, child):
    services = harness.services()
    claimed = harness.revisions.claim_plan_run(
        child["id"], "owner", "conversation-1", expected_version=child["edit_version"],
        settings=harness.settings,
        result_alias_resolver=lambda current: harness.service_bindings.admitted_result_aliases(
            current, services.results,
        ),
    )
    lease = harness.recovery.ExecutionLease(claimed, _authorize(harness), message_container=harness.messages)
    return harness.execution.prepare_harness_execution(claimed, settings=harness.settings, lease=lease)


def test_a_retry_generates_the_missing_image_and_delivers_all_of_them(harness, monkeypatch):
    """Version 0.261.138: a retry reuses the images that exist and rewrites the report and file."""
    route = importlib.import_module("functions_image_api_route")
    checkpoints = importlib.import_module("functions_orchestration_checkpoints")
    calls = _enable_images(harness, monkeypatch, failures={
        2: route.ImageGenerationError("busy", "image_rate_limited", 429),
    })
    _create(harness, _president_plan(), replies=[REPORT_TEXT])

    execution = harness.prepare()
    execution.execute()
    parent = harness.read()
    recovery = harness.recovery.recovery_projection(parent)
    earlier = harness.messages.read_item(checkpoints.orchestration_answer_message_id("run-1"), "conversation-1")

    assert parent["status"] == "failed" and parent["outcome"] == "partial"
    assert earlier["content"].count("```simpleimage") == 2
    assert recovery["eligible"], recovery
    assert recovery["reused_step_ids"] == ["washington", "jefferson"]
    assert recovery["retry_step_ids"] == ["adams", "report", "word"]

    child = _retry_request(harness, execution, "retry-missing-image")
    harness.replies.append(REPORT_TEXT)
    _claim_retry(harness, child).execute()
    finished = harness.runs.read_item(child["id"], "conversation-1")
    answer = harness.messages.read_item(checkpoints.orchestration_answer_message_id(child["id"]), "conversation-1")
    images = {image["metadata"]["image_proposal"]["visualId"]: image for image in _image_messages(harness)}
    listed = {
        entry["visual_id"]: entry["message_id"]
        for entry in answer["metadata"]["orchestration"]["generated_images"]
    }

    assert finished["status"] == "completed", finished.get("failure")
    assert len(calls) == 4 and len(images) == 3
    assert answer["content"].count("```simpleimage") == 3 and "Delivery notes" not in answer["content"]
    assert sorted(_embedded_images("docx", data) for data in _files(harness, ".docx")) == [2, 3]
    assert listed == {visual_id: image["id"] for visual_id, image in images.items()}
    # The earlier attempt's answer still shows its own images; nothing was moved away from it.
    assert {
        visual_id: image["metadata"]["image_proposal"]["source_assistant_message_id"]
        for visual_id, image in images.items()
    } == {"washington": earlier["id"], "jefferson": earlier["id"], "adams": answer["id"]}
    assert harness.messages.read_item(earlier["id"], "conversation-1")["metadata"]["orchestration"][
        "generated_images"
    ] == earlier["metadata"]["orchestration"]["generated_images"]


def test_a_retry_that_could_only_resend_a_refused_image_is_not_offered(harness, monkeypatch):
    route = importlib.import_module("functions_image_api_route")
    _enable_images(harness, monkeypatch, failures={
        2: route.ImageGenerationError("refused", "image_content_refused", 400),
    })
    _create(harness, _president_plan(), replies=[REPORT_TEXT])

    execution = harness.prepare()
    execution.execute()
    recovery = harness.recovery.recovery_projection(harness.read())

    assert recovery["eligible"] is False and recovery["reason_code"] == "retry_would_repeat"
    assert recovery["retry_step_ids"] == ["adams", "report", "word"]
    with pytest.raises(harness.recovery.RecoveryError) as refused:
        _retry_request(harness, execution, "retry-refused-image")
    assert refused.value.message == harness.schema.FAILURE_MESSAGES["retry_would_repeat"]


def test_a_retry_stays_available_while_other_work_could_succeed(harness, monkeypatch):
    route = importlib.import_module("functions_image_api_route")
    _enable_images(harness, monkeypatch, failures={
        2: route.ImageGenerationError("refused", "image_content_refused", 400),
        3: route.ImageGenerationError("busy", "image_rate_limited", 429),
    })
    _create(harness, _president_plan(), replies=[REPORT_TEXT])

    harness.prepare().execute()
    recovery = harness.recovery.recovery_projection(harness.read())

    assert recovery["eligible"], recovery
    assert recovery["retry_step_ids"] == ["adams", "jefferson", "report", "word"]


def test_the_refusal_rule_only_counts_work_that_depends_on_a_refusal():
    recovery = importlib.import_module("functions_orchestration_recovery")
    repeats = recovery._retry_repeats_refusal
    record = {"plan": {"planner_contract_version": 2, "steps": []}}
    refused = {"step_id": "adams", "status": "failed", "failure": {"code": "image_content_refused"}}
    stale = {"step_id": "report", "status": "completed"}
    blocked = {"step_id": "word", "status": "skipped", "failure": {"code": "dependency_unavailable"}}
    timed_out = {"step_id": "word", "status": "failed", "failure": {"code": "provider_timeout"}}
    retry = ["adams", "report", "word"]

    assert repeats(record, [refused, stale, blocked], retry) is True
    assert repeats(record, [refused, stale, timed_out], retry) is False
    assert repeats(record, [refused, stale, blocked], [*retry, "never_ran"]) is False
    assert repeats(record, [stale, blocked], ["report", "word"]) is False
    # A run from the removed earlier contract never reaches this rule: its recovery
    # projection is refused first with the legacy-plan reason.
    legacy = recovery.recovery_projection({"id": "legacy", "plan": {"planner_contract_version": 1, "steps": []}})
    assert legacy["eligible"] is False and legacy["reason_code"] == "legacy_plan"


def test_scenario_a_failed_render_is_never_reported_as_delivered(harness, monkeypatch):
    _enable_images(harness, monkeypatch)
    adapters = importlib.import_module("functions_generated_office_adapters")
    contracts = importlib.import_module("functions_generated_export_contracts")

    def failing(*args, **kwargs):
        raise contracts.GeneratedFileExportError("invalid_data", "Offline render failure.")

    monkeypatch.setattr(adapters, "_render_generated_office_source", failing)
    _create(harness, _president_plan(), replies=[REPORT_TEXT])

    harness.prepare().execute()
    saved = harness.read()
    answer = harness.assistant_messages()[0]["content"]

    assert saved["status"] == "failed" and saved["outcome"] == "partial"
    assert answer.count("```simpleimage") == 3
    assert r"word\.docx: could not be created" in answer
    assert harness.blobs.file_uploads == 0


def test_a_turned_off_file_step_is_named_in_the_delivery_notes(harness, monkeypatch):
    _enable_images(harness, monkeypatch)
    plan = _normalize(harness, _president_plan())
    for step in plan["steps"]:
        if step["step_id"] == "word":
            step["enabled"] = False
    _store(harness, plan, replies=[REPORT_TEXT])

    harness.prepare().execute()
    answer = harness.assistant_messages()[0]["content"]

    assert harness.read()["status"] == "completed"
    assert "- Not delivered: The report as a file. Its step was turned off." in answer

def test_turning_off_a_file_at_approval_keeps_the_run_retryable(harness, monkeypatch):
    """The saved plan's deliverable briefs match what runs, so checkpoints still match on retry."""
    from copy import copy

    route = importlib.import_module("functions_image_api_route")
    _enable_images(harness, monkeypatch, failures={2: route.ImageGenerationError("busy", "image_rate_limited", 429)})
    record = _create(harness, _president_plan(), replies=[REPORT_TEXT])
    services = harness.services()
    claimed = harness.revisions.claim_plan_run(
        record["id"], "owner", "conversation-1", expected_version=record.get("edit_version"),
        settings=harness.settings, edits={"disabled_step_ids": ["word"], "removed_document_ids": {}},
        result_alias_resolver=lambda current: harness.service_bindings.admitted_result_aliases(
            current, services.results,
        ),
    )
    brief = next(step for step in claimed["plan"]["steps"] if step["step_id"] == "report")["deliverable_context"]
    assert [entry["enabled"] for entry in brief if entry["relation"] == "rendered_as"] == [False]

    def authorize():
        return harness.bootstrap.read_owned_conversation("owner", "conversation-1")

    lease = harness.recovery.ExecutionLease(claimed, authorize, message_container=harness.messages)
    execution = harness.execution.prepare_harness_execution(claimed, settings=harness.settings, lease=lease)
    execution.execute()
    parent = harness.read()
    assert parent["status"] == "failed" and parent["outcome"] == "partial"
    probe = copy(execution.context)
    probe.result_service = harness.services().results

    child = harness.recovery.prepare_retry(
        "run-1", "owner", {
            "conversation_id": "conversation-1", "submission_id": "retry-after-turning-off-a-file",
            "expected_version": parent["recovery_version"],
        },
        authorize=authorize, message_container=harness.messages,
        validate=lambda current: harness.recovery.validate_resume(
            current, probe, harness.settings, authorize, source_run_id=current["id"],
        ),
    )
    assert child["retry_of_run_id"] == "run-1"


def test_when_every_image_fails_the_file_holds_no_tokens_or_broken_images(harness, monkeypatch):
    route = importlib.import_module("functions_image_api_route")
    refused = route.ImageGenerationError("refused", "image_content_refused", 400)
    _enable_images(harness, monkeypatch, failures={1: refused, 2: refused, 3: refused})
    _create(harness, _president_plan(), replies=[REPORT_TEXT + "\n\n![Portrait](asset:washington)"])

    harness.prepare().execute()
    saved = harness.read()
    data = _file_bytes(harness, ".docx")
    from docx import Document

    text = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(data)).paragraphs)
    assert [output["state"] for output in saved["outputs"]] == ["completed"]
    assert _embedded_images("docx", data) == 0
    assert "[[image:" not in text and "asset:" not in text and "Jefferson served" in text
    assert "0 of 3 images were generated" in harness.assistant_messages()[0]["content"]


def test_a_revision_may_drop_images_chosen_with_the_image_control(harness, monkeypatch):
    _enable_images(harness, monkeypatch)
    reply = {
        "kind": "plan", "intent": {"summary": "Write the report without images."}, "assumptions": [],
        "revised_request": "Write the report on the first three presidents, without images.",
        "deliverables": [_deliverable("report", "answer", "The report")],
        "steps": [_report_step()], "final_response": input_binding("report", "report"),
    }
    harness.replies = [json.dumps(reply)]
    kind, plan = harness.planner.plan_request(
        "Remove the images.", {}, "conversation-1", "owner", settings=harness.settings,
        seeds={"image_generation": True}, contract_version=2,
        request_context=harness.services().capability_request_bindings(),
        edit_context={"current_plan": {}, "current_request": "A report with images.",
                      "instruction": "Remove the images.", "chat": []},
    )
    assert kind == "plan" and len(harness.model_calls) == 1
    assert plan["validation"]["repairs"] == [
        "The plan no longer includes the images selected with the Image control. Review this change before running.",
    ]


def test_an_echoed_implicit_answer_is_kept_beside_real_deliverables(harness):
    rows = compose_step("rows", outputs=[{"name": "rows", "kind": "records-v1", "columns": [
        {"name": "State", "value_type": "string", "nullable": False},
    ]}])
    render = render_step("capitals", "csv", source="rows", output="rows")
    render["arguments"]["options"] = {"columns": ["State"]}
    implicit = _normalize(harness, {"steps": [compose_step()], "final_response": input_binding("prepare")})
    plan = _normalize(harness, {
        "deliverables": [*implicit["deliverables"], _deliverable("csv_file", "file", "The CSV", format="csv")],
        "steps": [{**compose_step(), "delivers": ["answer"]}, rows, {**render, "delivers": ["csv_file"]}],
        "final_response": input_binding("prepare"),
    })
    assert [deliverable["id"] for deliverable in plan["deliverables"]] == ["answer", "csv_file"]
    assert "implicit" not in plan["deliverables"][0]
