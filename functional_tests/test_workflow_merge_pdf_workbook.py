#!/usr/bin/env python3
# test_workflow_merge_pdf_workbook.py
"""
Functional test for PDF and workbook Merge tasks in workflows.
Version: 0.261.221
Implemented in: 0.261.221
Refs: microsoft/simplechat#1619

This test ensures that a workflow Merge task can assemble PDFs into one PDF, with a bookmark
per file, and CSV or Excel files into one workbook with a sheet per file. The task contract
accepts only each kind's own options and output format, files of the wrong type are refused
when selected and skipped when found at run time, one file found at run time is still
assembled, engine failures name the file to fix, cancellation stops the assembly before
anything is attached, and the runner reports pages or sheets as it goes. No model, Azure
service or network is used.
"""

import hashlib
import io
import sys
from pathlib import Path

import pytest

sys.path.append(str(Path(__file__).resolve().parent))

from test_support.versioning import assert_app_version_at_least  # noqa: E402
from test_workflow_merge_task import (  # noqa: E402,F401
    CONVERSATION,
    EAST,
    MERGE_HANDLES,
    NORTH,
    OWNER,
    RUN_ID,
    TASK_ID,
    MergeWorld,
    WorkflowInputError,
    _codes,
    app,
    drafts,
    dry_run,
    merge_blueprint,
    normalize,
    runner_merge,
)


def pdf_bytes(pages, *, size=200):
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=size, height=size)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


FILES = {
    "brief": ("Brief.pdf", pdf_bytes(2)),
    "appendix": ("appendix.pdf", pdf_bytes(1, size=300)),
    "broken": ("broken.pdf", b"%PDF-1.7\nthis is not a real PDF"),
    "north": NORTH,
    "east": EAST,
    "notes": ("notes.docx", b"not a pdf"),
}


def world(**kwargs):
    files = MergeWorld(dict(FILES), **kwargs)
    resolve = files.resolve

    def resolve_documents(document_ids, **options):
        manifest = resolve(document_ids, **options)
        for source in manifest:
            # A PDF is a narrative source, never a table.
            if source["file_name"].lower().endswith(".pdf"):
                source["source_kind"] = "narrative"
        return manifest

    files.resolve = resolve_documents
    return files


def pdf_action(**changes):
    action = {"type": "merge", "merge_kind": "pdf", "document_ids": ["brief", "appendix"]}
    action.update(changes)
    return action


def workbook_action(**changes):
    action = {"type": "merge", "merge_kind": "workbook", "document_ids": ["north", "east"]}
    action.update(changes)
    return action


def test_version_includes_pdf_and_workbook_merge_tasks():
    assert_app_version_at_least("0.261.221")


def test_pdf_and_workbook_tasks_take_only_their_own_options_and_output(app):
    action = normalize(app, pdf_action(merge_options={"bookmarks": False}))
    assert (action["merge_kind"], action["output_format"], action["merge_options"]) == ("pdf", "pdf", {"bookmarks": False})
    action = normalize(app, workbook_action(merge_options={"sheets": "all"}))
    assert (action["merge_kind"], action["output_format"], action["merge_options"]) == ("workbook", "xlsx", {"sheets": "all"})
    for payload, message in (
        (pdf_action(output_format="xlsx"), "A PDF merge creates PDF files."),
        (workbook_action(output_format="csv"), "A workbook merge creates XLSX files."),
        (pdf_action(merge_options={"sheets": "all"}), "Some merge options don't apply to PDF merges. Remove them."),
        (workbook_action(merge_options={"bookmarks": True}), "Some merge options don't apply to workbook merges."),
        (workbook_action(merge_options={"sheet": "Sales", "sheets": "all"}), "Name one sheet, or read all sheets"),
    ):
        with pytest.raises(app.actions.MergeActionError) as caught:
            normalize(app, payload)
        assert message in str(caught.value)


def test_a_pdf_merge_keeps_every_page_in_order_with_a_bookmark_per_file(app):
    from pypdf import PdfReader

    files = world()
    updates = []
    result = files.run(app, normalize(app, pdf_action(output_file_name="Board pack")), progress=updates.append)

    [published] = files.published
    assert (published["file_name"], published["output_format"]) == ("Board pack.pdf", "pdf")
    assert published["idempotency_key"] == (
        f"workflow-merge:{RUN_ID}:{TASK_ID}:{hashlib.sha256(published['content']).hexdigest()}"
    )
    assert published["summary"] == "Merged 2 file(s) into one PDF file."
    reader = PdfReader(io.BytesIO(published["content"]))
    assert [round(float(page.mediabox.width)) for page in reader.pages] == [200, 200, 300]
    assert [item.title for item in reader.outline] == ["Brief.pdf", "appendix.pdf"]
    assert files.reads == ["brief", "appendix"]
    assert [(update["index"], update["pages"]) for update in updates] == [(1, 2), (2, 1)]

    assert result["generated_tabular_outputs"] == []
    assert result["merge_summary"] == {
        "status": "merged", "kind": "pdf", "files": 2, "pages": 3, "slides": 0, "sheets": 0, "skipped": 0,
        "file_name": "Board pack.pdf",
    }
    assert result["reply"].startswith("Merged **2 file(s)** into **Board pack.pdf** with 3 page(s).")
    assert result["reply"].endswith("Files merged, in order: Brief.pdf, appendix.pdf.")


def test_a_workbook_merge_puts_each_file_on_its_own_sheet(app):
    from openpyxl import load_workbook

    files = world()
    result = files.run(app, normalize(app, workbook_action()))
    [published] = files.published
    assert (published["file_name"], published["output_format"]) == ("merged.xlsx", "xlsx")
    workbook = load_workbook(io.BytesIO(published["content"]))
    assert workbook.sheetnames == ["north", "east"]
    assert [[str(value) for value in row] for row in workbook["north"].iter_rows(values_only=True)] == [
        ["Region", "Amount", "Rep"], ["North", "10", "Ana"], ["North", "12", "Bo"],
    ]
    # Excel sources keep their cell types; CSV values stay text.
    assert list(workbook["east"].iter_rows(values_only=True))[1] == ("East", 5, "Di")
    assert result["merge_summary"]["sheets"] == 2
    assert result["reply"].startswith("Merged **2 file(s)** into **merged.xlsx** with 2 sheet(s).")


def test_files_found_at_run_time_skip_the_wrong_type_and_one_pdf_is_still_assembled(app):
    files = world(collected=["brief", "north", "notes", "appendix"], unavailable=["appendix"])
    result = files.run(app, normalize(app, pdf_action(target_mode="all", doc_scope="personal")))
    [(collected_action, extensions, _limit)] = files.collect_calls
    assert extensions == (".pdf",) and collected_action["merge_kind"] == "pdf"
    assert files.reads == ["brief"]
    assert result["merge_summary"]["files"] == 1 and result["merge_summary"]["skipped"] == 3
    assert "- Skipped 2 file(s) that aren't the right type: north.csv, notes.docx." in result["reply"]
    assert "- Skipped 1 file(s) that aren't available right now: appendix.pdf." in result["reply"]


def test_document_merge_failures_name_the_file_to_fix(app):
    with pytest.raises(app.merge.WorkflowMergeError) as wrong_type:
        world().run(app, normalize(app, pdf_action(document_ids=["brief", "notes"])))
    assert str(wrong_type.value) == "notes.docx isn't a PDF file (.pdf), so it can't be merged."

    files = world()
    with pytest.raises(app.merge.WorkflowMergeError) as broken:
        files.run(app, normalize(app, pdf_action(document_ids=["brief", "broken"])))
    assert str(broken.value) == "broken.pdf couldn't be read as a PDF."
    assert not files.published

    with pytest.raises(app.merge.WorkflowMergeError, match='east.xlsx has no sheet named "Totals"'):
        world().run(app, normalize(app, workbook_action(merge_options={"sheet": "Totals"})))


def test_cancelling_stops_the_assembly_before_anything_is_attached(app):
    files = world()
    state = {"cancelled": False}

    def progress(update):
        state["cancelled"] = True

    with pytest.raises(app.merge.WorkflowMergeCancelled):
        files.run(app, normalize(app, pdf_action()), cancel=lambda: state["cancelled"], progress=progress)
    assert files.reads == ["brief"] and not files.published


def test_a_repeated_merge_reuses_its_file_and_pdf_input_stays_bounded(app):
    for action in (pdf_action(), workbook_action()):
        first, second = world(), world()
        first.run(app, normalize(app, action))
        second.run(app, normalize(app, action))
        assert first.published[0]["content"] == second.published[0]["content"]
        assert first.published[0]["idempotency_key"] == second.published[0]["idempotency_key"]

    mebibyte = 1024 * 1024
    settings = {"max_generated_chat_artifact_size_mb": 2000}
    # pypdf keeps every PDF in memory until the file is written; a workbook is copied a file at a time.
    assert app.merge._document_limits(settings, 100, "pdf").max_total_input_bytes == 300 * mebibyte
    assert app.merge._document_limits(settings, 100, "workbook").max_total_input_bytes == 4000 * mebibyte
    assert app.merge._document_limits({}, 100, "workbook").max_total_input_bytes == 1000 * mebibyte


def test_the_runner_attaches_the_pdf_and_reports_pages(app, monkeypatch):
    files = world()
    run, uploads, thoughts = runner_merge(app, monkeypatch, files)
    result = run(normalize(app, pdf_action()))
    [upload] = uploads
    assert (upload["file_name"], upload["output_format"], upload["capability"]) == ("merged.pdf", "pdf", "file_merge")
    assert upload["content"].startswith(b"%PDF-")
    [artifact] = result["generated_analysis_artifacts"]
    assert (artifact["file_name"], artifact["output_format"]) == ("merged.pdf", "pdf")
    assert [(entry["content"], entry["detail"]) for entry in thoughts] == [
        ("Merged file 1 of 2", "pages=2"), ("Merged file 2 of 2", "pages=1"), ("Merged files", "files=2"),
    ]

    run, uploads, _thoughts = runner_merge(app, monkeypatch, world())
    with pytest.raises(WorkflowInputError) as refused:
        run(normalize(app, pdf_action(document_ids=["brief", "broken"])))
    assert refused.value.public_message == "broken.pdf couldn't be read as a PDF."
    assert not uploads


def test_chat_can_propose_pdf_and_workbook_merge_tasks(drafts):
    result = dry_run(drafts, merge_blueprint(
        {"kind": "pdf", "files": "inputs", "options": {"bookmarks": False}}, inputs=["south", "north"],
    ), MERGE_HANDLES)
    assert result["ok"] is True, result["errors"]
    action = result["workflow"]["tasks"][0]["document_action"]
    assert (action["merge_kind"], action["output_format"], action["merge_options"]) == ("pdf", "pdf", {"bookmarks": False})
    assert action["document_ids"] == ["doc-south", "doc-north"]

    result = dry_run(drafts, merge_blueprint({"kind": "workbook", "files": "all", "options": {"sheets": "all"}}))
    assert result["ok"] is True, result["errors"]
    action = result["workflow"]["tasks"][0]["document_action"]
    assert (action["merge_kind"], action["output_format"], action["target_mode"]) == ("workbook", "xlsx", "all")

    wrong_format = merge_blueprint({"kind": "pdf", "files": "all", "output_format": "csv"})
    assert _codes(dry_run(drafts, wrong_format)) == [("merge_format_invalid", "/tasks/0/merge/output_format")]
    wrong_option = merge_blueprint({"kind": "pdf", "files": "all", "options": {"sheets": "all"}})
    assert _codes(dry_run(drafts, wrong_option)) == [("merge_options_invalid", "/tasks/0/merge/options")]
    unknown_kind = merge_blueprint({"kind": "zip", "files": "all"})
    assert _codes(dry_run(drafts, unknown_kind))[0][0] == "blueprint_invalid"


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
