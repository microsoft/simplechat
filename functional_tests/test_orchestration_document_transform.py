# test_orchestration_document_transform.py
"""Functional tests for native document inspection and the two preparation tools.

Version: 0.261.321
Implemented in: 0.261.321

Exercises real retained results, compiler, executor and composition with network blocked.
"""

import hashlib
import importlib
import io
import json
from copy import deepcopy
from dataclasses import replace
from datetime import datetime, timedelta, timezone

import pytest
from docx import Document
from pptx import Presentation

from test_orchestration_dependency_runtime import binding, execute, runtime  # noqa: F401


AVAILABLE = ['document_inspect', 'document_create_from_reference', 'document_refresh', 'compose']
PROFILE = 'document_transform_v1'


def reference_bytes():
    document = Document()
    document.add_paragraph('Reference example')
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class InspectionAdapter:
    @staticmethod
    def inspect_package(package, *, limits, check):
        return {
            'targets': [{'id': 'body:1', 'type': 'paragraph', 'text': 'Reference example'}],
            'layouts': [], 'limitations': [],
        }

    @staticmethod
    def transform_package(package, prepared, *, limits, check, image_resolver):
        return {'applied': [], 'skipped': [], 'unchanged': [], 'warnings': []}


def prepared(content, mode='refresh'):
    return {
        'schema_version': PROFILE, 'output_format': 'docx', 'mode': mode,
        'reference': {
            'handle': {'kind': 'document', 'document_id': 'reference'},
            'content_sha256': hashlib.sha256(content).hexdigest(), 'extension': '.docx',
        },
        'operations': [{'op': 'replace_text', 'target': 'body:1', 'text': 'New information'}],
        'unresolved': [{'target': 'chart:1', 'reason': 'The latest figure was not provided.'}],
    }


def steps(capability='document_refresh'):
    return [
        {
            'step_id': 'inspect', 'capability_id': 'document_inspect',
            'arguments': {'reference': {'kind': 'document', 'document_id': 'reference'}},
        },
        {
            'step_id': 'update', 'capability_id': capability,
            'arguments': {'instruction': 'Update the document; disclose missing information.'},
            'inputs': {'inspection': {'binding': binding('inspect', 'inspection'), 'allow_partial': False}},
        },
        {
            'step_id': 'answer', 'capability_id': 'compose',
            'arguments': {'instruction': 'Describe the prepared work without claiming a file exists.'},
            'inputs': {'prepared': {'binding': binding('update', 'transformation'), 'allow_partial': False}},
            'outputs': [{'name': 'answer', 'kind': 'markdown-v1'}],
        },
    ]


def setup_case(runtime, monkeypatch, capability='document_refresh', replies=None):
    import functions_document_transform as transform

    monkeypatch.setitem(transform._ADAPTERS, 'docx', InspectionAdapter)
    content = reference_bytes()
    mode = 'refresh' if capability == 'document_refresh' else 'create_from_reference'
    value = prepared(content, mode)
    case = runtime.make(
        steps(capability), replies or [json.dumps({'transformation': value}), 'Prepared a draft with missing information.'],
        available=AVAILABLE, profiles={PROFILE: transform.get_document_transform_schema()},
        profile_validator=lambda profile, value: transform.validate_document_transform(value),
        final_response=binding('answer'),
    )
    case.fixture.sources['reference'] = {
        'document_id': 'reference', 'scope': 'personal', 'scope_id': 'owner',
        'source_version': 1, 'source_revision': 'reference-rev', 'file_name': 'reference.docx',
        'source_kind': 'narrative',
    }
    case.context.document_transform_source_resolver = lambda handle: {'file_name': 'reference.docx', 'source': None}
    case.reads = []

    def read(handle):
        case.reads.append(deepcopy(handle))
        return content

    case.context.document_transform_source_reader = read
    return case, value


@pytest.mark.parametrize('capability', ['document_create_from_reference', 'document_refresh'])
def test_distinct_tools_reuse_real_composition_and_bound_inspection(runtime, monkeypatch, capability):
    case, value = setup_case(runtime, monkeypatch, capability)
    result = execute(runtime, case)
    assert result['status'] == 'completed', result
    assert len(case.reads) == 1
    assert len(case.model.calls) == 2
    task = runtime.contracts.TaskResult.from_dict(result['task_results']['update'])
    stored = case.fixture.restart().open_result(task.output('transformation')).read_value()
    assert stored == value
    assert task.role == 'reason' and task.outputs[0].completeness.status == 'complete'
    policy = case.model.calls[0][0][0]['content']
    assert 'useful document draft' in policy and f'Required mode: {value["mode"]}' in policy
    assert result['artifacts'] == []
    assert task.outputs[0].producer.capability_id == capability
    document_ids = runtime.schema.plan_document_ids(case.plan)
    assert document_ids == ['reference']


@pytest.mark.parametrize('field', ['mode', 'reference'])
def test_model_cannot_switch_mode_or_reference_and_receives_one_correction(runtime, monkeypatch, field):
    case, value = setup_case(runtime, monkeypatch)
    bad = deepcopy(value)
    if field == 'mode':
        bad[field] = 'create_from_reference'
    else:
        bad[field]['handle']['document_id'] = 'another'
    case.model.replies = [
        json.dumps({'transformation': bad}), json.dumps({'transformation': value}), 'Prepared a draft.',
    ]
    result = execute(runtime, case)
    assert result['status'] == 'completed', result
    assert len(case.model.calls) == 3
    correction = case.model.calls[1][0][-1]['content']
    assert 'declared columns or schema' in correction
    assert len(case.reads) == 1


def test_unsupported_edit_and_unresolved_content_are_complete_prepared_draft(runtime, monkeypatch):
    case, value = setup_case(runtime, monkeypatch)
    value['operations'].append({'op': 'replace_text', 'target': 'unsupported:diagram', 'text': 'New diagram'})
    case.model.replies[0] = json.dumps({'transformation': value})
    result = execute(runtime, case)
    assert result['status'] == 'completed', result
    task = runtime.contracts.TaskResult.from_dict(result['task_results']['update'])
    stored = case.fixture.restart().open_result(task.output('transformation')).read_value()
    assert stored['unresolved'] == value['unresolved']
    assert task.outputs[0].completeness.status == 'complete'


def test_inspection_access_failure_prevents_model_and_native_reads(runtime, monkeypatch):
    case, _ = setup_case(runtime, monkeypatch)

    def denied(handle):
        raise PermissionError('Private server-only explanation')

    case.context.document_transform_source_resolver = denied
    result = execute(runtime, case)
    assert result['status'] == 'failed'
    assert not case.model.calls and not case.reads
    assert 'Private server-only explanation' not in json.dumps(result)


def test_fake_inspection_producer_is_rejected_by_compiler(runtime):
    import functions_document_transform as transform

    forged = steps()
    forged[0] = {
        'step_id': 'inspect', 'capability_id': 'compose', 'arguments': {'instruction': 'Invent an inspection'},
        'outputs': [{'name': 'inspection', 'kind': 'structured-v1'}],
    }
    with pytest.raises(runtime.schema.PlanValidationError, match='inspection'):
        runtime.make(
            forged, available=AVAILABLE, profiles={PROFILE: transform.get_document_transform_schema()},
            final_response=binding('answer'),
        )


def test_inspection_runtime_discovery_needs_owner_callbacks(runtime):
    settings = {'enable_chat_orchestration': True}
    disabled = runtime.registry.resolve_available_capability_ids(settings, request_context={}, include_runtime_bindings=True)
    assert 'document_inspect' not in disabled
    available = runtime.registry.resolve_available_capability_ids(
        settings, request_context={
            'document_transform_source_reader': lambda handle: b'',
            'document_transform_source_resolver': lambda handle: {},
        }, include_runtime_bindings=True,
    )
    assert 'document_inspect' in available


def test_preparation_image_sources_are_exact_authorized_assets(runtime):
    from functions_orchestration_document_transform import validate_transform_preparation

    value = prepared(reference_bytes())
    value['operations'] = [{'op': 'replace_image', 'target': 'image:1', 'source': 'asset:generated-1'}]
    context = {
        'mode': 'refresh', 'inspection': {'reference': value['reference'], 'output_format': 'docx'},
        'image_sources': {'asset:generated-1'},
    }
    validate_transform_preparation({'transformation': value}, context)
    value['operations'][0]['source'] = 'asset:not-acquired'
    with pytest.raises(runtime.contracts.ResultContractError):
        validate_transform_preparation({'transformation': value}, context)


@pytest.mark.parametrize('output_format,template', [
    ('docx', False), ('docx', True), ('pptx', False), ('pptx', True),
])
@pytest.mark.parametrize('mode', ['refresh', 'create_from_reference'])
def test_complete_native_pipeline_delivers_actual_edits_and_persistent_report(
    runtime, monkeypatch, output_format, template, mode,
):
    import functions_document_transform as transform
    from test_document_transform_rendering import RenderCase, package_bytes
    from test_support.orchestration_results import source
    from test_support.orchestration_revisions import AtomicMemoryContainer

    adapter = importlib.import_module(f'functions_document_transform_{output_format}')
    monkeypatch.setitem(transform._ADAPTERS, output_format, adapter)
    content = package_bytes(output_format, template=template)
    extension = ('.dotx' if output_format == 'docx' else '.potx') if template else f'.{output_format}'
    handle = {'kind': 'document', 'document_id': 'reference'}
    inspection = transform.inspect_document(content, f'reference{extension}', handle)
    target = next(item['id'] for item in inspection['targets'] if item.get('text') == 'Original')
    value = {
        'schema_version': PROFILE, 'output_format': output_format, 'mode': mode,
        'reference': inspection['reference'],
        'operations': [
            {'op': 'replace_text', 'target': target, 'text': 'Verified update'},
            {'op': 'replace_text', 'target': 'uneditable:chart', 'text': 'Unsupported chart update'},
        ],
        'unresolved': [{'target': 'missing:figure', 'reason': 'The latest figure was not supplied.'}],
    }
    capability = 'document_refresh' if mode == 'refresh' else 'document_create_from_reference'
    pipeline = steps(capability)
    pipeline.insert(2, {
        'step_id': 'file', 'capability_id': 'render_file',
        'arguments': {
            'output_format': output_format, 'profile': PROFILE, 'file_name': f'draft.{output_format}',
        },
        'inputs': {'source': {'binding': binding('update', 'transformation'), 'allow_partial': False}},
    })
    case = runtime.make(
        pipeline, [json.dumps({'transformation': value}), 'Prepared the requested document.'],
        available=[*AVAILABLE, 'render_file'],
        profiles={PROFILE: transform.get_document_transform_schema()},
        profile_validator=lambda profile, value: transform.validate_document_transform(value),
        final_response=binding('answer'),
    )
    files = object.__new__(RenderCase)
    case.model.model_metadata = replace(case.model.model_metadata, context_window=250000, input_limit=200000)
    files.content = content
    files.authorized_handle = handle
    files.now = datetime.now(timezone.utc)
    files.fixture = case.fixture
    files.fixture.sources['reference'] = source('reference')
    files.runs = AtomicMemoryContainer('conversation_id')
    run = deepcopy(case.fixture.runs['run-1'])
    run.update(
        record_type='run', approval={'state': 'approved'},
        execution_deadline_at=(files.now + timedelta(hours=1)).isoformat(),
    )
    files.runs.create_item(run)
    files.messages, files.bytes = {}, {}
    files.reads, files.failures, files.streams = [], [], []
    files.uploads = 0
    files.restart()
    case.context.result_service = files.service.results
    case.context.rendering_service = files.service
    case.context.approved_work_id = 'run-1'
    case.context.document_transform_source_reader = files.read_native
    case.context.document_transform_source_resolver = lambda selected: {'file_name': f'reference{extension}'}
    result = execute(runtime, case)
    assert result['status'] == 'completed', result
    assert len(case.model.calls) == 2 and files.uploads == 1
    assert len(files.reads) >= 2 and all(selected == handle for selected in files.reads)
    payload = next(iter(files.bytes.values()))
    document = Document(io.BytesIO(payload)) if output_format == 'docx' else Presentation(io.BytesIO(payload))
    text = document.paragraphs[0].text if output_format == 'docx' else document.slides[0].shapes.title.text
    assert text == 'Verified update'
    assert files.content == content and all(stream.closed for stream in files.streams)
    output = files.service.store.list_outputs('run-1')[0]
    assert output['state'] == 'completed'
    report = output['renderer_metadata']['transformation_report']
    assert report['applied'][0]['target'] == target
    assert report['skipped'][0]['target'] == 'uneditable:chart'
    assert report['unresolved'][0]['target'] == 'missing:figure'
    assert report['status'] == 'draft'
    assert output['committed_intent']['content_sha256'] == hashlib.sha256(payload).hexdigest()
    from functions_orchestration_deliverables import delivery_notes

    notes = delivery_notes(case.plan, {}, file_outputs=result['delivery_facts']['files'])
    assert 'skipped' in notes.lower() and 'unresolved' in notes.lower()
    saved_report = files.restart().store.list_outputs('run-1')[0]['renderer_metadata']['transformation_report']
    assert saved_report == report
