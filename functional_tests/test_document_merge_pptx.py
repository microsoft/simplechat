#!/usr/bin/env python3
# test_document_merge_pptx.py
"""
Functional test for PowerPoint assembly in V2 file merge.
Version: 0.261.223
Implemented in: 0.261.223

This test ensures that functions_document_merge appends slides in order at the package
level: keep_source carries each deck's layouts, masters and themes (reusing identical
ones), use_first places slides on the first deck's matching layouts, charts with their
embedded workbooks, pictures, tables and speaker notes survive, each deck becomes a
section, slide selection and order are honored for every deck, links to slides that
were not merged are removed, and the merged package keeps unique IDs, a content type
for every part and no orphaned slides. Damaged, encrypted, macro-enabled and oversized
inputs fail closed, the same decks always give the same bytes, library failures name
only the file, and host failures pass through unchanged.
"""

import io
import os
import re
import sys
import time
import zipfile

import pytest

sys.path.append(os.path.dirname(os.path.abspath(__file__)))
sys.path.append(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "application", "single_app"))

from test_support.versioning import assert_app_version_at_least

from functions_document_merge import (
    DocumentMergeError,
    DocumentMergeLimits,
    DocumentMergeOptions,
    DocumentMergePart,
    merge_documents,
)


THEME_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"


def png_bytes():
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (40, 30), (200, 30, 30)).save(buffer, "PNG")
    return buffer.getvalue()


def rewrite(data, editor):
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            target.writestr(info, editor(info.filename, source.read(info)))
    return output.getvalue()


def deck_bytes(label, *, slides=3, theme=None, extras=False, size=None, external_link=False):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    presentation = Presentation()
    if size:
        presentation.slide_width, presentation.slide_height = size
    made = []
    for number in range(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1 if number else 0])
        slide.shapes.title.text = f"{label} slide {number + 1}"
        made.append(slide)
    if extras:
        slide = made[1]
        slide.shapes.add_picture(io.BytesIO(png_bytes()), Inches(1), Inches(3))
        chart_data = CategoryChartData()
        chart_data.categories = ["Q1", "Q2"]
        chart_data.add_series("Sales", (1, 2))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4), Inches(3), Inches(3), Inches(2), chart_data)
        slide.shapes.add_table(2, 2, Inches(0.5), Inches(5), Inches(3), Inches(1)).table.cell(0, 0).text = "cell"
        slide.notes_slide.notes_text_frame.text = f"{label} notes"
        made[0].shapes.title.click_action.target_slide = made[2]
    if external_link:
        made[0].shapes.title.click_action.hyperlink.address = "https://example.com/report"
    buffer = io.BytesIO()
    presentation.save(buffer)
    data = buffer.getvalue()
    if theme:
        def recolor(name, raw):
            if name.startswith("ppt/theme/theme1"):
                raw = re.sub(rb'(<a:accent1>\s*<a:srgbClr val=")[0-9A-F]{6}', rb"\g<1>" + theme, raw)
                raw = raw.replace(b'name="Office Theme"', b'name="' + label.encode() + b' Theme"')
            return raw

        data = rewrite(data, recolor)
    return data


def part(source_id, file_name, data, pages=None):
    return DocumentMergePart(source_id, file_name, lambda data=data: data, pages=pages)


def open_merged(parts, **options):
    from pptx import Presentation

    with merge_documents("pptx", parts, options=DocumentMergeOptions(**options)) as result:
        data = result.read_bytes()
        return Presentation(io.BytesIO(data)), zipfile.ZipFile(io.BytesIO(data)), result.report


def theme_name(slide):
    theme = slide.slide_layout.slide_master.part.part_related_by(THEME_REL)
    return re.search(rb'name="([^"]+)"', theme.blob).group(1).decode()


def titles(presentation):
    return [slide.shapes.title.text for slide in presentation.slides]


def assert_package_is_consistent(archive):
    names = set(archive.namelist())
    presentation = archive.read("ppt/presentation.xml").decode()
    master_ids = [int(value) for value in re.findall(r'<p:sldMasterId id="(\d+)"', presentation)]
    layout_ids = []
    for name in names:
        if name.startswith("ppt/slideMasters/") and name.endswith(".xml"):
            layout_ids += [int(value) for value in re.findall(r'<p:sldLayoutId id="(\d+)"', archive.read(name).decode())]
    combined = master_ids + layout_ids
    assert len(combined) == len(set(combined)), "Master and layout IDs must be unique across the presentation."
    assert all(value >= 2147483648 for value in combined)
    slide_ids = [int(value) for value in re.findall(r'<p:sldId id="(\d+)"', presentation)]
    assert len(slide_ids) == len(set(slide_ids)) and all(256 <= value < 2147483648 for value in slide_ids)
    types = archive.read("[Content_Types].xml").decode()
    defaults = set(re.findall(r'Extension="([^"]+)"', types))
    overrides = set(re.findall(r'PartName="/([^"]+)"', types))
    for name in names:
        if name == "[Content_Types].xml" or name.endswith("/"):
            continue
        extension = name.rsplit(".", 1)[-1].lower()
        assert name in overrides or extension in defaults, f"{name} has no content type."
    for name in names:
        if not name.endswith(".rels"):
            continue
        directory = os.path.dirname(os.path.dirname(name))
        for target, mode in re.findall(r'Target="([^"]+)"(?:\s+TargetMode="([^"]+)")?', archive.read(name).decode()):
            if mode == "External":
                continue
            resolved = target.lstrip("/") if target.startswith("/") else os.path.normpath(
                os.path.join(directory, target),
            ).replace("\\", "/")
            assert resolved in names, f"{name} points at missing part {resolved}."


def test_version_includes_powerpoint_merges():
    assert_app_version_at_least("0.261.223")


def test_keep_source_keeps_each_decks_look_and_reuses_identical_masters():
    parts = [
        part("a", "Alpha.pptx", deck_bytes("Alpha", slides=2)),
        part("b", "Beta.pptx", deck_bytes("Beta", theme=b"00FF00", extras=True)),
        part("c", "Gamma.pptx", deck_bytes("Gamma", slides=1)),
    ]
    presentation, archive, report = open_merged(parts)
    assert titles(presentation) == [
        "Alpha slide 1", "Alpha slide 2", "Beta slide 1", "Beta slide 2", "Beta slide 3", "Gamma slide 1",
    ]
    assert [theme_name(slide) for slide in presentation.slides] == [
        "Office Theme", "Office Theme", "Beta Theme", "Beta Theme", "Beta Theme", "Office Theme",
    ]
    assert len(presentation.slide_masters) == 2, "Gamma shares Alpha's template, so its master is reused."
    beta = presentation.slides[3]
    kinds = {str(shape.shape_type).split(".")[-1].split(" ")[0] for shape in beta.shapes if shape.shape_type}
    assert {"CHART", "PICTURE", "TABLE"} <= kinds
    assert beta.notes_slide.notes_text_frame.text == "Beta notes"
    assert any(name.startswith("ppt/embeddings/") for name in archive.namelist())
    sections = re.findall(r'section name="([^"]+)"', archive.read("ppt/presentation.xml").decode())
    assert sections == ["Alpha", "Beta", "Gamma"]
    assert report["totals"]["slides"] == 6
    assert_package_is_consistent(archive)


def test_use_first_places_slides_on_the_first_decks_layouts():
    parts = [
        part("a", "Alpha.pptx", deck_bytes("Alpha", slides=2)),
        part("b", "Beta.pptx", deck_bytes("Beta", theme=b"00FF00", extras=True)),
    ]
    presentation, archive, report = open_merged(parts, formatting="use_first", sections=False)
    assert len(presentation.slide_masters) == 1
    assert {theme_name(slide) for slide in presentation.slides} == {"Office Theme"}
    assert [slide.slide_layout.name for slide in presentation.slides][2:4] == ["Title Slide", "Title and Content"]
    assert "sectionLst" not in archive.read("ppt/presentation.xml").decode()
    assert report["options"] == {"formatting": "use_first", "sections": False}
    assert_package_is_consistent(archive)


def test_slide_selection_order_links_and_unselected_parts():
    parts = [
        part("a", "Alpha.pptx", deck_bytes("Alpha", slides=3), pages=((3, 3), (1, 1))),
        part("b", "Beta.pptx", deck_bytes("Beta", extras=True), pages=((1, 2),)),
    ]
    presentation, archive, report = open_merged(parts)
    assert titles(presentation) == ["Alpha slide 3", "Alpha slide 1", "Beta slide 1", "Beta slide 2"]
    slides = [name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)]
    assert len(slides) == 4, "The first deck's unselected slide must not remain in the package."
    assert [warning["code"] for warning in report["parts"][1]["warnings"]] == ["links_removed"]
    assert_package_is_consistent(archive)
    with pytest.raises(DocumentMergeError) as caught:
        open_merged([part("a", "a.pptx", deck_bytes("A"), pages=((1, 2), (2, 3))), parts[1]])
    assert caught.value.code == "range_invalid"


def test_size_differences_and_external_links_are_reported():
    from pptx.util import Inches

    parts = [
        part("a", "Alpha.pptx", deck_bytes("Alpha", slides=1)),
        part("b", "Wide.pptx", deck_bytes("Wide", slides=1, size=(Inches(13.333), Inches(7.5)), external_link=True)),
    ]
    presentation, archive, report = open_merged(parts)
    codes = [warning["code"] for warning in report["parts"][1]["warnings"]]
    assert "slide_size_differs" in codes
    assert presentation.slides[1].shapes.title.click_action.hyperlink.address == "https://example.com/report"
    assert_package_is_consistent(archive)


def test_damaged_encrypted_macro_enabled_and_oversized_decks_fail_closed():
    good = part("a", "Alpha.pptx", deck_bytes("Alpha", slides=2))
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("pptx", [good, part("b", "locked.pptx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 32)])
    assert caught.value.code == "encrypted_document"
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("pptx", [good, part("b", "broken.pptx", b"PK\x03\x04broken")])
    assert caught.value.code == "unreadable_document"

    def as_macro_enabled(name, raw):
        if name == "[Content_Types].xml":
            return raw.replace(
                b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
                b"application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml",
            )
        return raw

    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("pptx", [good, part("b", "macro.pptx", rewrite(deck_bytes("M", slides=1), as_macro_enabled))])
    assert caught.value.code == "unreadable_document"
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("pptx", [good, part("b", "b.pptx", deck_bytes("B", slides=2))], limits=DocumentMergeLimits(max_slides=3))
    assert caught.value.code == "slide_limit_exceeded"
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents(
            "pptx", [good, part("b", "b.pptx", deck_bytes("B", slides=1))],
            limits=DocumentMergeLimits(max_package_uncompressed_bytes=1000),
        )
    assert caught.value.code == "source_too_large"


def test_the_same_decks_always_give_the_same_bytes():
    parts = [
        part("a", "Alpha.pptx", deck_bytes("Alpha", slides=2)),
        part("b", "Beta.pptx", deck_bytes("Beta", theme=b"00FF00", extras=True)),
    ]
    with merge_documents("pptx", parts) as first:
        content = first.read_bytes()
    # Two seconds is the resolution of a ZIP entry's date.
    time.sleep(2.1)
    with merge_documents("pptx", parts) as second:
        assert second.read_bytes() == content
    with zipfile.ZipFile(io.BytesIO(content)) as package:
        assert {entry.date_time for entry in package.infolist()} == {(2000, 1, 1, 0, 0, 0)}
        sections = re.findall(r'section name="[^"]+" id="(\{[0-9A-F-]+\})"', package.read("ppt/presentation.xml").decode())
    assert len(sections) == 2 and len(set(sections)) == 2


def test_library_failures_name_only_the_file_and_host_failures_pass_through(monkeypatch):
    import functions_document_merge_pptx

    parts = [
        part("a", "Alpha.pptx", deck_bytes("Alpha", slides=1)),
        part("b", "Beta.pptx", deck_bytes("Beta", slides=1)),
    ]

    def leak(*_args, **_kwargs):
        raise RuntimeError("PRIVATE slide text")

    monkeypatch.setattr(functions_document_merge_pptx._DeckCopier, "copy_slide", leak)
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("pptx", parts)
    assert (caught.value.code, caught.value.message) == (
        "unreadable_document", "Beta.pptx couldn't be combined with the other decks.",
    )
    monkeypatch.undo()

    monkeypatch.setattr(functions_document_merge_pptx, "_validate", leak)
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("pptx", parts)
    assert (caught.value.code, caught.value.message) == ("merge_failed", "The merged PowerPoint deck couldn't be written.")
    monkeypatch.undo()

    def disk_full(*_args, **_kwargs):
        raise OSError("No space left on device")

    # A full temporary disk, or a failing cancellation check, is the host's problem, not the file's.
    monkeypatch.setattr(functions_document_merge_pptx._DeckCopier, "copy_slide", disk_full)
    with pytest.raises(OSError):
        merge_documents("pptx", parts)
    monkeypatch.undo()

    class RunStoreUnavailable(RuntimeError):
        pass

    calls = []

    def cancel_requested():
        calls.append(1)
        if len(calls) > 4:
            raise RunStoreUnavailable("The run store is unavailable.")
        return False

    # Call 5 is the check before Beta's first slide is copied.
    with pytest.raises(RunStoreUnavailable):
        merge_documents("pptx", parts, cancel_requested=cancel_requested)


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
