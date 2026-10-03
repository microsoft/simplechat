#!/usr/bin/env python3
# test_document_merge_docx.py
"""
Functional test for Word assembly in V2 file merge.
Version: 0.261.224
Implemented in: 0.261.222
Derived list IDs added in: 0.261.224

This test ensures that functions_document_merge appends Word documents in order with
docxcompose: keep_source formatting keeps a differently styled document's look by
copying clashing styles under new names, use_first maps them to the first document's
styles, page breaks and file-name headings are optional, tables, numbered lists and
images survive, comment anchors that would be orphaned are removed, damaged, encrypted
or macro-enabled files are refused, the same documents always give the same bytes, even
when later documents bring lists of their own, library failures name only the file, and
host failures pass through unchanged.
"""

import io
import os
import sys
import time
import zipfile

import pytest

sys.path.append(os.path.dirname(os.path.abspath(__file__)))
sys.path.append(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "application", "single_app"))

from test_support.versioning import assert_app_version_at_least

from functions_document_merge import (
    DocumentMergeError,
    DocumentMergeOptions,
    DocumentMergePart,
    merge_documents,
)


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def png_bytes():
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (20, 10), (10, 120, 200)).save(buffer, "PNG")
    return buffer.getvalue()


def docx_bytes(label, color, size, *, image=False, comment=False):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from lxml import etree

    document = Document()
    heading = document.styles["Heading 1"]
    heading.font.color.rgb = RGBColor(*color)
    heading.font.size = Pt(size)
    document.add_heading(f"{label} title", level=1)
    paragraph = document.add_paragraph(f"{label} body")
    document.add_paragraph("first item", style="List Number")
    document.add_paragraph("second item", style="List Number")
    document.add_table(rows=1, cols=2).cell(0, 0).text = f"{label} cell"
    if image:
        document.add_picture(io.BytesIO(png_bytes()), width=Inches(1))
    if comment:
        start = etree.SubElement(paragraph._p, f"{W}commentRangeStart")
        start.set(f"{W}id", "0")
        end = etree.SubElement(paragraph._p, f"{W}commentRangeEnd")
        end.set(f"{W}id", "0")
        run = etree.SubElement(paragraph._p, f"{W}r")
        reference = etree.SubElement(run, f"{W}commentReference")
        reference.set(f"{W}id", "0")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def part(source_id, file_name, data):
    return DocumentMergePart(source_id, file_name, lambda data=data: data)


def merged(parts, **options):
    from docx import Document

    with merge_documents("docx", parts, options=DocumentMergeOptions(**options)) as result:
        return Document(io.BytesIO(result.read_bytes())), result.report, result


def test_version_includes_word_merges():
    assert_app_version_at_least("0.261.222")


def test_keep_source_formatting_copies_clashing_styles():
    parts = [
        part("a", "Alpha.docx", docx_bytes("Alpha", (255, 0, 0), 20)),
        part("b", "Beta.docx", docx_bytes("Beta", (0, 0, 255), 30, image=True)),
    ]
    document, report, result = merged(parts)
    texts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    assert texts[:2] == ["Alpha title", "Alpha body"]
    assert "Beta title" in texts and texts.index("Beta title") > texts.index("Alpha body")
    styles = {paragraph.text: paragraph.style for paragraph in document.paragraphs if paragraph.text}
    assert styles["Alpha title"].name == "Heading 1"
    beta_style = styles["Beta title"]
    assert beta_style.name != "Heading 1", "The differently styled heading must keep its own definition."
    assert beta_style.font.size.pt == 30 and str(beta_style.font.color.rgb) == "0000FF"
    assert len(document.tables) == 2
    assert len(document.inline_shapes) == 1
    assert [paragraph.style.name for paragraph in document.paragraphs if paragraph.text == "first item"] == ["List Number"] * 2
    assert report["status"] == "merged"
    assert result.media_type.endswith("wordprocessingml.document")
    assert any("keeps its look" in item for item in report["limitations"])


def test_use_first_formatting_maps_styles_to_the_first_document():
    parts = [
        part("a", "Alpha.docx", docx_bytes("Alpha", (255, 0, 0), 20)),
        part("b", "Beta.docx", docx_bytes("Beta", (0, 0, 255), 30)),
    ]
    document, report, _ = merged(parts, formatting="use_first")
    styles = {paragraph.text: paragraph.style.name for paragraph in document.paragraphs if paragraph.text}
    assert styles["Beta title"] == "Heading 1"
    assert report["options"]["formatting"] == "use_first"


def test_page_breaks_and_headings_are_optional():
    parts = [
        part("a", "Alpha.docx", docx_bytes("Alpha", (255, 0, 0), 20)),
        part("b", "Beta.docx", docx_bytes("Beta", (0, 0, 255), 30)),
    ]
    document, _, _ = merged(parts, source_headings=True)
    texts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    assert texts[0] == "Alpha.docx" and "Beta.docx" in texts
    breaks = document.element.body.findall(f".//{W}br[@{W}type='page']")
    assert len(breaks) == 1
    document, _, _ = merged(parts, page_breaks=False)
    assert document.element.body.findall(f".//{W}br[@{W}type='page']") == []


def test_orphaned_comment_anchors_are_removed():
    parts = [
        part("a", "Alpha.docx", docx_bytes("Alpha", (255, 0, 0), 20)),
        part("b", "Beta.docx", docx_bytes("Beta", (0, 0, 255), 30, comment=True)),
    ]
    document, report, _ = merged(parts)
    assert document.element.body.findall(f".//{W}commentReference") == []
    assert document.element.body.findall(f".//{W}commentRangeStart") == []
    assert report["parts"][1]["warnings"][0]["code"] == "comments_removed"


def test_damaged_encrypted_and_wrong_files_are_refused():
    good = part("a", "Alpha.docx", docx_bytes("Alpha", (255, 0, 0), 20))
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("docx", [good, part("b", "locked.docx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 32)])
    assert caught.value.code == "encrypted_document"
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("docx", [good, part("b", "broken.docx", b"PK\x03\x04garbage")])
    assert caught.value.code == "unreadable_document"
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("docx", [good, part("b", "old.doc", b"\xd0\xcf\x11\xe0")])
    assert caught.value.code == "unsupported_format"
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("docx", [good, part("b", "deck.pptx", b"PK")])
    assert caught.value.code == "unsupported_format"


def test_the_same_documents_always_give_the_same_bytes():
    parts = [
        part("a", "Alpha.docx", docx_bytes("Alpha", (255, 0, 0), 20)),
        part("b", "Beta.docx", docx_bytes("Beta", (0, 0, 255), 30, image=True)),
    ]
    with merge_documents("docx", parts) as first:
        content = first.read_bytes()
    # Two seconds is the resolution of a ZIP entry's date.
    time.sleep(2.1)
    with merge_documents("docx", parts) as second:
        assert second.read_bytes() == content
    with zipfile.ZipFile(io.BytesIO(content)) as package:
        assert {entry.date_time for entry in package.infolist()} == {(2000, 1, 1, 0, 0, 0)}
        assert "2000-01-01T00:00:00Z" in package.read("docProps/core.xml").decode("utf-8")


def listed_docx_bytes(label, items=2):
    """A document whose paragraphs are numbered directly, the way Word writes an ordinary list."""
    from docx import Document

    document = Document()
    document.add_paragraph(f"{label} intro")
    for number in range(items):
        paragraph = document.add_paragraph(f"{label} item {number + 1}")
        numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        numbering.get_or_add_ilvl().val = 0
        numbering.get_or_add_numId().val = 1
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def list_identities(content):
    """For each numbered paragraph, the ID of its list definition (w:nsid)."""
    from lxml import etree

    with zipfile.ZipFile(io.BytesIO(content)) as package:
        body = etree.fromstring(package.read("word/document.xml"))
        numbering = etree.fromstring(package.read("word/numbering.xml"))
    nsids = {
        element.get(f"{W}abstractNumId"): element.find(f"{W}nsid").get(f"{W}val")
        for element in numbering.iter(f"{W}abstractNum") if element.find(f"{W}nsid") is not None
    }
    definitions = {
        element.get(f"{W}numId"): element.find(f"{W}abstractNumId").get(f"{W}val")
        for element in numbering.iter(f"{W}num")
    }
    identities = {}
    for paragraph in body.iter(f"{W}p"):
        number = paragraph.find(f"{W}pPr/{W}numPr/{W}numId")
        if number is not None:
            text = "".join(paragraph.itertext())
            identities[text] = nsids[definitions[number.get(f"{W}val")]]
    return identities, list(nsids.values())


def test_lists_copied_from_later_documents_restart_and_give_the_same_bytes():
    parts = [
        part("a", "Alpha.docx", listed_docx_bytes("Alpha")),
        part("b", "Beta.docx", listed_docx_bytes("Beta")),
        part("c", "Gamma.docx", listed_docx_bytes("Gamma")),
    ]
    merged = []
    for _ in range(3):
        with merge_documents("docx", parts) as result:
            merged.append(result.read_bytes())
    # docxcompose picks a random ID for each list it copies; the merge derives it instead.
    assert merged[0] == merged[1] == merged[2]
    identities, every_id = list_identities(merged[0])
    assert len(every_id) == len(set(every_id))
    # Each document's list keeps one ID of its own, so its numbering restarts at 1.
    assert identities["Alpha item 1"] == identities["Alpha item 2"]
    assert identities["Beta item 1"] == identities["Beta item 2"]
    assert len({identities["Alpha item 1"], identities["Beta item 1"], identities["Gamma item 1"]}) == 3
    # The first document's own list keeps the ID it had.
    original, _ = list_identities(listed_docx_bytes("Alpha"))
    assert identities["Alpha item 1"] == original["Alpha item 1"]


def test_library_failures_name_only_the_file_and_host_failures_pass_through(monkeypatch):
    from docxcompose.composer import Composer

    parts = [
        part("a", "Alpha.docx", docx_bytes("Alpha", (255, 0, 0), 20)),
        part("b", "Beta.docx", docx_bytes("Beta", (0, 0, 255), 30)),
    ]

    def leak(*_args, **_kwargs):
        raise RuntimeError("PRIVATE document text")

    monkeypatch.setattr(Composer, "append", leak)
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("docx", parts)
    assert (caught.value.code, caught.value.message) == (
        "unreadable_document", "Beta.docx couldn't be combined with the other documents.",
    )
    monkeypatch.undo()

    monkeypatch.setattr(Composer, "save", leak)
    with pytest.raises(DocumentMergeError) as caught:
        merge_documents("docx", parts)
    assert (caught.value.code, caught.value.message) == ("merge_failed", "The merged Word document couldn't be written.")
    monkeypatch.undo()

    def disk_full(*_args, **_kwargs):
        raise OSError("No space left on device")

    # A full temporary disk, or a failing cancellation check, is the host's problem, not the file's.
    monkeypatch.setattr(Composer, "save", disk_full)
    with pytest.raises(OSError):
        merge_documents("docx", parts)
    monkeypatch.undo()

    class RunStoreUnavailable(RuntimeError):
        pass

    calls = []

    def cancel_requested():
        calls.append(1)
        if len(calls) > 4:
            raise RunStoreUnavailable("The run store is unavailable.")
        return False

    with pytest.raises(RunStoreUnavailable):
        merge_documents("docx", parts, cancel_requested=cancel_requested)


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
