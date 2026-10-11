# test_document_transform_pptx.py
"""
Functional tests for package-preserving PowerPoint reference transformations.
Version: 0.261.321
Implemented in: 0.261.321

Real python-pptx/Pillow-generated PPTX and POTX packages exercise both modes,
multiple masters/layouts, grouped shapes, repeated labels, mixed run formatting,
tables, pictures, notes, variable slide plans and unsupported preserved parts.
No Azure, browser, network, temporary-directory fixtures or desktop Office is used.
"""

from copy import deepcopy
import base64
import hashlib
import io
import json
import os
from pathlib import Path
import posixpath
import socket
import subprocess
import sys
import zipfile

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt
import pytest


APP_PATH = str(Path(__file__).resolve().parents[1] / "application" / "single_app")
if APP_PATH not in sys.path:
    sys.path.insert(0, APP_PATH)

from functions_document_transform import (
    DocumentTransformCancelled,
    DocumentTransformError,
    DocumentTransformLimits,
    GuardedPackage,
    inspect_document,
    transform_document,
)
from functions_document_transform_pptx import transform_package


NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
}
REL_PREFIX = NS["r"] + "/"
HANDLE = {"kind": "document", "document_id": "authorized-reference"}


def _q(prefix, local):
    return f"{{{NS[prefix]}}}{local}"


def _xml(data):
    return etree.fromstring(data, etree.XMLParser(resolve_entities=False, no_network=True))


def _serialize(root):
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _files(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _zip(files):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name in sorted(files):
            info = zipfile.ZipInfo(name, date_time=(2020, 1, 1, 0, 0, 0))
            archive.writestr(info, files[name])
    return output.getvalue()


def _png(color="red"):
    output = io.BytesIO()
    Image.new("RGB", (48, 32), color).save(output, "PNG")
    return output.getvalue()


def _relationships(*items):
    root = etree.Element(_q("rel", "Relationships"), nsmap={None: NS["rel"]})
    for rid, kind, target in items:
        etree.SubElement(root, _q("rel", "Relationship"), {
            "Id": rid, "Type": REL_PREFIX + kind, "Target": target,
        })
    return _serialize(root)


def _add_master(files):
    """Make a second real master/layout/theme graph without reserializing the deck."""
    master = _xml(files["ppt/slideMasters/slideMaster1.xml"])
    layouts = master.find(_q("p", "sldLayoutIdLst"))
    for child in list(layouts):
        layouts.remove(child)
    etree.SubElement(layouts, _q("p", "sldLayoutId"), {"id": "2147483700", _q("r", "id"): "rId1"})
    master.find(_q("p", "cSld")).set("name", "Second master")
    layout = _xml(files["ppt/slideLayouts/slideLayout1.xml"])
    layout.find(_q("p", "cSld")).set("name", "Alternate title")
    files["ppt/slideMasters/slideMaster2.xml"] = _serialize(master)
    files["ppt/slideMasters/_rels/slideMaster2.xml.rels"] = _relationships(
        ("rId1", "slideLayout", "../slideLayouts/slideLayout12.xml"),
        ("rId2", "theme", "../theme/alternate.xml"),
    )
    files["ppt/slideLayouts/slideLayout12.xml"] = _serialize(layout)
    files["ppt/slideLayouts/_rels/slideLayout12.xml.rels"] = _relationships(
        ("rId1", "slideMaster", "../slideMasters/slideMaster2.xml"),
    )
    theme = files["ppt/theme/theme1.xml"].replace(b'name="Office Theme"', b'name="Alternate Theme"')
    files["ppt/theme/alternate.xml"] = theme
    presentation = _xml(files["ppt/presentation.xml"])
    rels = _xml(files["ppt/_rels/presentation.xml.rels"])
    etree.SubElement(rels, _q("rel", "Relationship"), {
        "Id": "rIdSecondMaster", "Type": REL_PREFIX + "slideMaster", "Target": "slideMasters/slideMaster2.xml",
    })
    etree.SubElement(presentation.find(_q("p", "sldMasterIdLst")), _q("p", "sldMasterId"), {
        "id": "2147483701", _q("r", "id"): "rIdSecondMaster",
    })
    files["ppt/presentation.xml"] = _serialize(presentation)
    files["ppt/_rels/presentation.xml.rels"] = _serialize(rels)
    types = _xml(files["[Content_Types].xml"])
    for part, kind in (
        ("ppt/slideMasters/slideMaster2.xml", "presentationml.slideMaster"),
        ("ppt/slideLayouts/slideLayout12.xml", "presentationml.slideLayout"),
        ("ppt/theme/alternate.xml", "theme"),
    ):
        etree.SubElement(types, _q("ct", "Override"), {
            "PartName": f"/{part}", "ContentType": f"application/vnd.openxmlformats-officedocument.{kind}+xml",
        })
    files["[Content_Types].xml"] = _serialize(types)
    if "ppt/slides/_rels/slide3.xml.rels" in files:
        rels = _xml(files["ppt/slides/_rels/slide3.xml.rels"])
        for rel in rels:
            if rel.get("Type").endswith("/slideLayout"):
                rel.set("Target", "../slideLayouts/slideLayout12.xml")
        files["ppt/slides/_rels/slide3.xml.rels"] = _serialize(rels)


def _deck(*, blank=False, template=False, multi_master=True, notes=True, linked=False):
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(12), Inches(7)
    if not blank:
        slides = [presentation.slides.add_slide(presentation.slide_layouts[0]) for _ in range(3)]
        for slide in slides:
            slide.shapes.title.text = "Repeated label"
        slide = slides[0]
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1.2))
        box.name = "Mixed body"
        paragraph = box.text_frame.paragraphs[0]
        paragraph.level = 1
        first = paragraph.add_run()
        first.text, first.font.bold, first.font.size = "Bold", True, Pt(24)
        second = paragraph.add_run()
        second.text, second.font.italic, second.font.size = " normal", True, Pt(18)
        box.text_frame.add_paragraph().text = "Second paragraph"
        group = slide.shapes.add_group_shape()
        group.name = "Grouped content"
        grouped = group.shapes.add_textbox(Inches(7), Inches(1), Inches(3), Inches(1))
        grouped.name, grouped.text = "Grouped title", "Repeated label"
        nested = group.shapes.add_group_shape()
        nested.shapes.add_textbox(Inches(8), Inches(2), Inches(2), Inches(1)).text = "Nested label"
        picture = slide.shapes.add_picture(io.BytesIO(_png()), Inches(6), Inches(3), Inches(2), Inches(1))
        picture.name = "Photo"
        picture.crop_left = 0.1
        shared = slides[1].shapes.add_picture(io.BytesIO(_png()), Inches(6), Inches(3), Inches(2), Inches(1))
        shared.name = "Shared photo"
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
        for row in range(2):
            for col in range(2):
                table.cell(row, col).text = f"Cell {row + 1}:{col + 1}"
        data = CategoryChartData()
        data.categories = ["Q1", "Q2"]
        data.add_series("Sales", (5, 9))
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7), Inches(4), Inches(3), Inches(2), data,
        )
        chart.name = "Protected chart"
        if notes:
            slide.notes_slide.notes_text_frame.text = "Original notes\nNote two"
        if linked:
            slide.shapes.title.click_action.target_slide = slides[2]
    output = io.BytesIO()
    presentation.save(output)
    files = _files(output.getvalue())
    if multi_master:
        _add_master(files)
    if not blank:
        slide = _xml(files["ppt/slides/slide1.xml"])
        extension_list = etree.SubElement(slide, _q("p", "extLst"))
        extension = etree.SubElement(extension_list, _q("p", "ext"), {"uri": "urn:simplechat:test:preserved"})
        etree.SubElement(extension, "{urn:simplechat:test}unknown").text = "preserve this extension"
        files["ppt/slides/slide1.xml"] = _serialize(slide)
        files["customXml/preserved.xml"] = b'<data xmlns="urn:simplechat:test">untouched custom XML</data>'
        rels = _xml(files["ppt/slides/_rels/slide1.xml.rels"])
        etree.SubElement(rels, _q("rel", "Relationship"), {
            "Id": "rIdCustomXml", "Type": REL_PREFIX + "customXml", "Target": "../../customXml/preserved.xml",
        })
        files["ppt/slides/_rels/slide1.xml.rels"] = _serialize(rels)
    if template:
        types = _xml(files["[Content_Types].xml"])
        for entry in types:
            if entry.get("PartName") == "/ppt/presentation.xml":
                entry.set("ContentType", "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml")
        files["[Content_Types].xml"] = _serialize(types)
    return _zip(files)


def _inspection(data, extension=".pptx", **kwargs):
    return inspect_document(data, "reference" + extension, HANDLE, **kwargs)


def _prepared(inspection, operations, mode="refresh", unresolved=None):
    return {
        "schema_version": "document_transform_v1", "output_format": "pptx",
        "mode": mode, "reference": inspection["reference"], "operations": operations,
        "unresolved": unresolved or [],
    }


def _op(op, target, **kwargs):
    return {"op": op, "target": target, **kwargs}


def _named(inspection, name):
    return next(item for item in inspection["targets"] if item.get("name") == name)


def _slide_targets(inspection):
    return [item["id"] for item in inspection["targets"] if item["type"] == "slide"]


def _transform(data, inspection, operations, mode="refresh", **kwargs):
    prepared = _prepared(inspection, operations, mode, kwargs.pop("unresolved", None))
    with transform_document(data, prepared, **kwargs) as result:
        return result.file_content.getvalue(), deepcopy(result.metadata)


def _consistent(data):
    package = GuardedPackage(data, ".pptx")
    names = set(package.files)
    for name in names:
        if name.endswith(".rels"):
            owner = "" if name == "_rels/.rels" else posixpath.join(
                posixpath.dirname(posixpath.dirname(name)), posixpath.basename(name)[:-5],
            )
            for rel in package.relationships(owner):
                if not rel["external"]:
                    assert rel["target"] in names
    presentation = Presentation(io.BytesIO(data))
    root = package.parse_xml(package.main_part)
    ids = [entry.get("id") for entry in root.findall(f"{_q('p', 'sldIdLst')}/{_q('p', 'sldId')}")]
    assert len(ids) == len(set(ids)) == len(presentation.slides)
    for slide in presentation.slides:
        identifiers = [node.get("id") for node in slide.element.iter(_q("p", "cNvPr"))]
        assert len(identifiers) == len(set(identifiers))
    return package, presentation


def test_inspects_actual_shapes_groups_notes_and_all_master_layouts():
    data = _deck()
    inspection = _inspection(data)
    repeated = [item for item in inspection["targets"] if item.get("text") == "Repeated label"]
    assert len(repeated) >= 8
    assert len({item["id"] for item in repeated}) == len(repeated)
    assert _slide_targets(inspection) == ["slide:256", "slide:257", "slide:258"]
    assert len(inspection["layouts"]) == 12
    assert len({item["master_part"] for item in inspection["layouts"]}) == 2
    group = _named(inspection, "Grouped content")
    child = _named(inspection, "Grouped title")
    assert child["parent"] == group["id"]
    assert child["id"].startswith("slide:256:shape:")
    assert not child["id"].startswith(group["id"] + ":")
    assert any(item["id"] == "slide:256:notes" and item["text"] == "Original notes\nNote two" for item in inspection["targets"])
    assert any(item["id"] == "slide:257:notes" and item["text"] == "" for item in inspection["targets"])
    assert inspection["slide_size"] == {"cx": str(Inches(12)), "cy": str(Inches(7)), "type": "screen4x3"}
    assert any(item["type"] == "unsupported" and item.get("name") == "Protected chart" for item in inspection["targets"])


@pytest.mark.parametrize("mode", ["refresh", "create_from_reference"])
@pytest.mark.parametrize("template", [False, True])
def test_selected_text_table_group_picture_and_notes_edits_preserve_other_parts(mode, template):
    source = _deck(template=template)
    snapshot = bytes(source)
    inspection = _inspection(source, ".potx" if template else ".pptx")
    mixed = _named(inspection, "Mixed body")["id"]
    grouped = _named(inspection, "Grouped title")["id"]
    photo = _named(inspection, "Photo")["id"]
    cell = next(item["id"] for item in inspection["targets"] if item.get("text") == "Cell 1:2" and item["type"] == "table_cell")
    resolver_calls = []

    def resolve(token):
        resolver_calls.append(token)
        return {"content": _png("blue")}

    operations = [
        _op("replace_text", mixed + ":p:1", text="New mixed words"),
        _op("append_text", mixed, text="New paragraph"),
        _op("replace_text", grouped, text="Changed group"),
        _op("replace_text", cell, text="Updated cell\nSecond line"),
        _op("replace_image", photo, source="asset:photo"),
        _op("replace_text", "slide:256:notes", text="Refreshed notes\nDetails"),
        _op("replace_text", "slide:257:notes", text="Created notes"),
    ]
    output, metadata = _transform(source, inspection, operations, mode, image_resolver=resolve)
    package, presentation = _consistent(output)
    assert source == snapshot
    assert resolver_calls == ["asset:photo"]
    assert len(metadata["transformation_report"]["applied"]) == len(operations)
    assert metadata["transformation_report"]["skipped"] == []
    original = _files(source)
    assert package.files["ppt/slides/slide2.xml"] == original["ppt/slides/slide2.xml"]
    assert package.files["ppt/slides/slide3.xml"] == original["ppt/slides/slide3.xml"]
    for name in original:
        if name.startswith(("ppt/slideMasters/", "ppt/slideLayouts/", "ppt/theme/", "ppt/charts/", "ppt/embeddings/", "customXml/")):
            assert package.files[name] == original[name], name
    slide = presentation.slides[0]
    shape = next(shape for shape in slide.shapes if shape.name == "Mixed body")
    assert shape.text == "New mixed words\nSecond paragraph\nNew paragraph"
    first = shape.text_frame.paragraphs[0]
    assert first.level == 1 and first.runs[0].font.bold and first.runs[0].font.size == Pt(24)
    assert first.runs[1].font.italic and first.runs[1].font.size == Pt(18)
    table = next(shape.table for shape in slide.shapes if shape.has_table)
    assert table.cell(0, 1).text == "Updated cell\nSecond line"
    picture = next(shape for shape in slide.shapes if shape.name == "Photo")
    assert picture.image.blob == _png("blue") and picture.crop_left == 0.1
    assert picture.width == Inches(2) and picture.height == Inches(1)
    shared = next(shape for shape in presentation.slides[1].shapes if shape.name == "Shared photo")
    assert shared.image.blob == _png()
    assert slide.notes_slide.notes_text_frame.text == "Refreshed notes\nDetails"
    assert presentation.slides[1].notes_slide.notes_text_frame.text == "Created notes"
    assert any("Mixed run styles" in message for message in metadata["transformation_report"]["warnings"])
    assert any("overflow" in message for message in metadata["transformation_report"]["warnings"])
    assert b"preserve this extension" in package.files["ppt/slides/slide1.xml"]
    assert package.main_content_type.endswith("presentation.main+xml")


@pytest.mark.parametrize("mode", ["refresh", "create_from_reference"])
def test_slide_clone_removal_and_edits_have_predictable_target_trees(mode):
    source = _deck(linked=True)
    inspection = _inspection(source)
    body = _named(inspection, "Mixed body")["id"]
    title = next(item["id"] for item in inspection["targets"] if item.get("text") == "Repeated label" and item["type"] == "shape" and item["parent"] == "slide:256")
    operations = [
        _op("clone", "slide:256", clone_id="agenda"),
        _op("replace_text", "agenda" + title[len("slide:256"):], text="New agenda"),
        _op("replace_text", "agenda" + body[len("slide:256"):], text="New body"),
        _op("replace_text", "agenda:notes", text="Independent cloned notes"),
        _op("clone", "agenda", clone_id="closing"),
        _op("replace_text", "closing" + title[len("slide:256"):], text="Closing"),
        _op("replace_text", "closing:notes", text="Independent closing notes"),
        _op("remove", "slide:257"),
        _op("remove", "slide:258"),
    ]
    output, metadata = _transform(source, inspection, operations, mode)
    package, presentation = _consistent(output)
    assert [slide.shapes.title.text for slide in presentation.slides] == ["Repeated label", "New agenda", "Closing"]
    assert [slide.notes_slide.notes_text_frame.text for slide in presentation.slides] == [
        "Original notes\nNote two", "Independent cloned notes", "Independent closing notes",
    ]
    assert len(metadata["cloned_targets"]) == 2
    assert metadata["cloned_targets"][0]["target"] == "agenda"
    assert metadata["cloned_targets"][0]["output_id"] == "slide:259"
    assert "agenda:notes" in metadata["cloned_targets"][0]["descendants"]
    assert "slide:257" in metadata["removed_targets"]
    originals = _files(source)
    for name in originals:
        if name.startswith(("ppt/slides/", "ppt/charts/", "ppt/embeddings/", "ppt/notesSlides/")):
            assert package.files[name] == originals[name]
    notes_parts = []
    root = package.parse_xml(package.main_part)
    by_id = {rel["id"]: rel["target"] for rel in package.relationships(package.main_part)}
    parts = [
        by_id[entry.get(_q("r", "id"))]
        for entry in root.findall(f"{_q('p', 'sldIdLst')}/{_q('p', 'sldId')}")
    ]
    for part in parts:
        notes = next(rel for rel in package.relationships(part) if rel["type"].endswith("/notesSlide"))
        notes_parts.append(notes["target"])
        back = next(rel for rel in package.relationships(notes["target"]) if rel["type"].endswith("/slide"))
        assert back["target"] == part
    assert len(set(notes_parts)) == 3


def test_blank_potx_layout_plan_creates_variable_length_deck_and_notes_master():
    source = _deck(blank=True, template=True)
    inspection = _inspection(source, ".potx")
    assert _slide_targets(inspection) == [] and len(inspection["layouts"]) == 12
    title_layout = next(item for item in inspection["layouts"] if item["name"] == "Title Slide")
    alternate = next(item for item in inspection["layouts"] if item["name"] == "Alternate title")
    title_shape = next(
        item for item in title_layout["clone_targets"]
        if item["type"] == "shape" and "replace_text" in item["operations"]
    )
    source_target = next(item for item in inspection["targets"] if item["id"] == title_shape["source_target"])
    assert source_target["operations"] == []
    assert source_target["clone_operations"] == ["replace_text", "append_text", "remove"]
    assert source_target["clone_target_suffix"] == title_shape["target_suffix"]
    assert title_layout["clone_root_is_clone_id"]
    assert any(item["target_suffix"] == ":notes" for item in title_layout["clone_targets"])
    operations = [
        _op("clone", title_layout["id"], clone_id="opening"),
        _op("replace_text", "opening" + title_shape["target_suffix"], text="Opening"),
        _op("replace_text", "opening:notes", text="Brand new notes"),
        _op("clone", alternate["id"], clone_id="discard"),
        _op("clone", alternate["id"], clone_id="end"),
        _op("replace_text", "end" + title_shape["target_suffix"], text="Ending"),
        _op("remove", "discard"),
    ]
    output, metadata = _transform(source, inspection, operations, "create_from_reference")
    package, presentation = _consistent(output)
    assert [slide.shapes.title.text for slide in presentation.slides] == ["Opening", "Ending"]
    assert [slide.slide_layout.name for slide in presentation.slides] == ["Title Slide", "Alternate title"]
    assert presentation.slides[0].notes_slide.notes_text_frame.text == "Brand new notes"
    assert len(presentation.slide_masters) == 2
    assert metadata["slide_count"] == 2
    assert len(metadata["transformation_report"]["applied"]) == len(operations)
    source_parts = _files(source)
    for name in source_parts:
        if name.startswith(("ppt/slideMasters/", "ppt/slideLayouts/", "ppt/theme/")):
            assert package.files[name] == source_parts[name]
    assert package.main_content_type.endswith("presentation.main+xml")
    assert len([name for name in package.files if name.startswith("ppt/notesMasters/") and name.endswith(".xml")]) == 1
    note_master = next(
        package.parse_xml(name) for name in package.files
        if name.startswith("ppt/notesMasters/") and name.endswith(".xml")
    )
    assert len(list(note_master.iter(_q("p", "ph")))) == 2
    assert all(shape.find(f"{_q('p', 'spPr')}/{_q('a', 'xfrm')}") is not None for shape in note_master.iter(_q("p", "sp")))


def test_blank_layout_picture_placeholder_accepts_authorized_image_and_clears_examples():
    source = _deck(blank=True, template=True)
    inspection = _inspection(source, ".potx")
    layout = next(item for item in inspection["layouts"] if item["name"] == "Picture with Caption")
    picture = next(item for item in layout["clone_targets"] if item["type"] == "picture")
    assert picture["operations"] == ["replace_image", "remove"]
    output, metadata = _transform(source, inspection, [
        _op("clone", layout["id"], clone_id="photo-slide"),
        _op("replace_image", "photo-slide" + picture["target_suffix"], source="asset:photo"),
    ], "create_from_reference", image_resolver=lambda token: _png("purple"))
    package, presentation = _consistent(output)
    slide = presentation.slides[0]
    photos = [shape for shape in slide.shapes if shape.element.tag == _q("p", "pic")]
    assert len(photos) == 1 and photos[0].image.blob == _png("purple")
    assert photos[0].width > 0 and photos[0].height > 0
    assert all(not shape.text for shape in slide.shapes if shape.has_text_frame)
    assert len(metadata["transformation_report"]["applied"]) == 2
    for name, original in _files(source).items():
        if name.startswith("ppt/slideLayouts/"):
            assert package.files[name] == original


def test_unsupported_missing_and_unresolved_data_still_delivers_valid_draft():
    source = _deck()
    inspection = _inspection(source)
    chart = _named(inspection, "Protected chart")["id"]
    body = _named(inspection, "Mixed body")["id"]
    layout = inspection["layouts"][0]["id"]
    operations = [
        _op("replace_text", chart, text="Unsupported chart refresh"),
        _op("replace_text", "missing-target", text="Missing target"),
        _op("replace_text", layout, text="Do not rewrite masters"),
        _op("clone", body, clone_id="unsupported-shape-clone"),
        _op("replace_image", _named(inspection, "Photo")["id"], source="asset:missing"),
        _op("replace_text", body, text="Useful supported update"),
    ]
    output, metadata = _transform(
        source, inspection, operations, "create_from_reference",
        unresolved=[{"target": chart, "reason": "No new revenue data supplied."}],
        image_resolver=lambda token: None,
    )
    package, _presentation = _consistent(output)
    report = metadata["transformation_report"]
    assert len(report["skipped"]) == 5 and len(report["applied"]) == 1
    assert report["status"] == "draft"
    assert report["unresolved"] == [{"target": chart, "reason": "No new revenue data supplied."}]
    assert all(item["reason"] for item in report["skipped"])
    assert package.files["ppt/charts/chart1.xml"] == _files(source)["ppt/charts/chart1.xml"]


def test_unchanged_and_skipped_operations_are_honest_and_do_not_reserialize_parts():
    source = _deck()
    inspection = _inspection(source)
    body = _named(inspection, "Mixed body")
    operations = [
        _op("replace_text", body["id"], text=body["text"]),
        _op("append_text", body["id"], text=""),
        _op("replace_image", _named(inspection, "Photo")["id"], source="asset:same"),
        _op("replace_text", "not-there", text=""),
    ]
    output, metadata = _transform(source, inspection, operations, image_resolver=lambda token: _png())
    _consistent(output)
    assert _files(output) == _files(source)
    report = metadata["transformation_report"]
    assert report["applied"] == [] and len([item for item in report["unchanged"] if "op_index" in item]) == 3
    assert any("No changes applied" in warning for warning in report["warnings"])


def test_paragraph_replacement_append_and_remove_keep_other_paragraph_targets_stable():
    source = _deck()
    inspection = _inspection(source)
    body = _named(inspection, "Mixed body")["id"]
    operations = [
        _op("remove", body + ":p:1"),
        _op("append_text", body + ":p:2", text=" appended"),
        _op("replace_text", body + ":p:2", text="Preserved slot\nLine break"),
    ]
    output, metadata = _transform(source, inspection, operations)
    _, presentation = _consistent(output)
    shape = next(shape for shape in presentation.slides[0].shapes if shape.name == "Mixed body")
    assert shape.text == "Preserved slot\vLine break"
    assert len(shape.text_frame.paragraphs) == 1
    assert len(metadata["transformation_report"]["applied"]) == 3


def test_appending_after_removal_and_cloning_preserve_live_paragraph_slot_ids():
    source = _deck()
    inspection = _inspection(source)
    body = _named(inspection, "Mixed body")["id"]
    cloned_body = "slots" + body[len("slide:256"):]
    operations = [
        _op("remove", body + ":p:1"),
        _op("append_text", body, text="Third slot"),
        _op("replace_text", body + ":p:2", text="Original second slot"),
        _op("clone", "slide:256", clone_id="slots"),
        _op("replace_text", cloned_body + ":p:2", text="Cloned second slot"),
        _op("replace_text", cloned_body + ":p:3", text="Cloned third slot"),
        _op("replace_text", body + ":p:1", text="Must remain unavailable"),
    ]
    output, metadata = _transform(source, inspection, operations)
    _, presentation = _consistent(output)
    first = next(shape for shape in presentation.slides[0].shapes if shape.name == "Mixed body")
    cloned = next(shape for shape in presentation.slides[-1].shapes if shape.name == "Mixed body")
    assert first.text == "Original second slot\nThird slot"
    assert cloned.text == "Cloned second slot\nCloned third slot"
    assert len(metadata["transformation_report"]["applied"]) == 6
    assert len(metadata["transformation_report"]["skipped"]) == 1
    descendants = metadata["cloned_targets"][0]["descendants"]
    assert cloned_body + ":p:2" in descendants and cloned_body + ":p:3" in descendants
    assert cloned_body + ":p:1" not in descendants


def test_clone_collision_missing_child_and_removed_group_are_skipped_without_wrong_match():
    source = _deck()
    inspection = _inspection(source)
    group = _named(inspection, "Grouped content")["id"]
    child = _named(inspection, "Grouped title")["id"]
    operations = [
        _op("clone", "slide:256", clone_id="slide:257"),
        _op("remove", group),
        _op("replace_text", child, text="Must not find a repeated label elsewhere"),
        _op("replace_text", "slide:257:shape:2", text="Only second slide"),
    ]
    output, metadata = _transform(source, inspection, operations)
    _, presentation = _consistent(output)
    assert [slide.shapes.title.text for slide in presentation.slides] == ["Repeated label", "Only second slide", "Repeated label"]
    assert len(metadata["transformation_report"]["skipped"]) == 2
    assert not any(shape.name == "Grouped content" for shape in presentation.slides[0].shapes)
    assert child in metadata["removed_targets"]


def test_remove_updates_custom_shows_sections_and_never_reuses_slide_ids():
    source = _deck()
    files = _files(source)
    main = _xml(files["ppt/presentation.xml"])
    entries = main.findall(f"{_q('p', 'sldIdLst')}/{_q('p', 'sldId')}")
    shows = etree.SubElement(main, _q("p", "custShowLst"))
    for index, selected in enumerate(((1,), (1, 2)), 1):
        show = etree.SubElement(shows, _q("p", "custShow"), {"name": f"Show {index}", "id": str(index)})
        slide_list = etree.SubElement(show, _q("p", "sldLst"))
        for number in selected:
            etree.SubElement(slide_list, _q("p", "sld"), {_q("r", "id"): entries[number].get(_q("r", "id"))})
    extension_list = etree.SubElement(main, _q("p", "extLst"))
    extension = etree.SubElement(extension_list, _q("p", "ext"), {"uri": "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"})
    section_ns = "{http://schemas.microsoft.com/office/powerpoint/2010/main}"
    sections = etree.SubElement(extension, section_ns + "sectionLst")
    for index, selected in enumerate(((257,), (256, 258)), 1):
        section = etree.SubElement(sections, section_ns + "section", {
            "name": f"Section {index}", "id": f"{{00000000-0000-0000-0000-00000000000{index}}}",
        })
        slide_list = etree.SubElement(section, section_ns + "sldIdLst")
        for number in selected:
            etree.SubElement(slide_list, section_ns + "sldId", {"id": str(number)})
    files["ppt/presentation.xml"] = _serialize(main)
    source = _zip(files)
    inspection = _inspection(source)
    output, metadata = _transform(source, inspection, [
        _op("remove", "slide:257"), _op("remove", "slide:258"),
        _op("clone", "slide:256", clone_id="after-removal"),
    ], "create_from_reference")
    package, presentation = _consistent(output)
    main = package.parse_xml(package.main_part)
    assert len(presentation.slides) == 2
    assert metadata["cloned_targets"][0]["output_id"] == "slide:259"
    assert main.find(_q("p", "custShowLst")) is None
    assert [node.get("id") for node in main.iter(section_ns + "sldId")] == ["256"]
    assert [node.get("name") for node in main.iter(section_ns + "section")] == ["Section 2"]


def test_unknown_presentation_reference_blocks_only_that_slide_removal():
    source = _deck()
    files = _files(source)
    main = _xml(files["ppt/presentation.xml"])
    rid = main.findall(f"{_q('p', 'sldIdLst')}/{_q('p', 'sldId')}")[1].get(_q("r", "id"))
    extensions = etree.SubElement(main, _q("p", "extLst"))
    extension = etree.SubElement(extensions, _q("p", "ext"), {"uri": "urn:simplechat:unknown"})
    etree.SubElement(extension, "{urn:simplechat:test}reference", {_q("r", "id"): rid})
    files["ppt/presentation.xml"] = _serialize(main)
    source = _zip(files)
    inspection = _inspection(source)
    output, metadata = _transform(source, inspection, [
        _op("remove", "slide:257"), _op("replace_text", "slide:256:shape:2", text="Other work"),
    ])
    package, presentation = _consistent(output)
    assert len(presentation.slides) == 3
    assert package.files["ppt/presentation.xml"] == files["ppt/presentation.xml"]
    assert len(metadata["transformation_report"]["skipped"]) == 1


def test_layout_clone_preserves_hyperlink_relationships_on_instantiated_placeholders():
    source = _deck(blank=True, template=True, multi_master=False)
    files = _files(source)
    layout_part = "ppt/slideLayouts/slideLayout1.xml"
    layout = _xml(files[layout_part])
    props = next(layout.iter(_q("p", "cNvPr")))
    hyperlink = etree.SubElement(props, _q("a", "hlinkClick"), {_q("r", "id"): "rIdHref"})
    hyperlink.set("tooltip", "A harmless external link")
    # The tree's group properties are not instantiated; use a real placeholder.
    props.remove(hyperlink)
    shape_props = layout.find(f"{_q('p', 'cSld')}/{_q('p', 'spTree')}/{_q('p', 'sp')}/{_q('p', 'nvSpPr')}/{_q('p', 'cNvPr')}")
    shape_props.append(hyperlink)
    files[layout_part] = _serialize(layout)
    rel_part = "ppt/slideLayouts/_rels/slideLayout1.xml.rels"
    rels = _xml(files[rel_part])
    etree.SubElement(rels, _q("rel", "Relationship"), {
        "Id": "rIdHref", "Type": REL_PREFIX + "hyperlink", "Target": "https://example.com/read", "TargetMode": "External",
    })
    files[rel_part] = _serialize(rels)
    source = _zip(files)
    inspection = _inspection(source, ".potx")
    layout_id = next(item["id"] for item in inspection["layouts"] if item["name"] == "Title Slide")
    output, _metadata = _transform(source, inspection, [_op("clone", layout_id, clone_id="linked")], "create_from_reference")
    package, presentation = _consistent(output)
    slide_part = presentation.slides[0].part.partname[1:]
    assert any(rel["external"] and rel["target"] == "https://example.com/read" for rel in package.relationships(slide_part))
    assert presentation.slides[0].shapes.title.click_action.hyperlink.address == "https://example.com/read"


def test_required_shape_animation_and_dynamic_text_are_preserved():
    source = _deck()
    inspection = _inspection(source)
    body = _named(inspection, "Mixed body")["id"]
    files = _files(source)
    slide = _xml(files["ppt/slides/slide1.xml"])
    extension = slide.find(f"{_q('p', 'extLst')}/{_q('p', 'ext')}")
    etree.SubElement(extension, _q("p", "spTgt"), {"spid": body.rsplit(":", 1)[1]})
    text_body = next(
        shape.find(_q("p", "txBody")) for shape in slide.iter(_q("p", "sp"))
        if shape.find(f"{_q('p', 'nvSpPr')}/{_q('p', 'cNvPr')}").get("name") == "Mixed body"
    )
    paragraph = text_body.find(_q("a", "p"))
    field = etree.SubElement(paragraph, _q("a", "fld"), {"id": "{00000000-0000-0000-0000-000000000001}", "type": "datetime1"})
    etree.SubElement(field, _q("a", "t")).text = "Automatic date"
    files["ppt/slides/slide1.xml"] = _serialize(slide)
    source = _zip(files)
    inspection = _inspection(source)
    output, metadata = _transform(source, inspection, [
        _op("remove", body), _op("replace_text", body, text="Do not break fields"),
        _op("replace_text", "slide:257:shape:2", text="Other work proceeds"),
    ])
    package, _presentation = _consistent(output)
    assert package.files["ppt/slides/slide1.xml"] == files["ppt/slides/slide1.xml"]
    assert len(metadata["transformation_report"]["skipped"]) == 2


def test_unknown_adapter_operation_is_skipped_and_other_work_proceeds():
    source = _deck()
    inspection = _inspection(source)
    package = GuardedPackage(source, ".pptx")
    prepared = _prepared(inspection, [
        _op("future_operation", "slide:256"),
        _op("replace_text", "slide:257:shape:2", text="Still works"),
    ])
    report = transform_package(
        package, prepared, limits=DocumentTransformLimits(), check=None, image_resolver=None,
    )
    packed = package.pack()
    try:
        _, presentation = _consistent(packed.getvalue())
    finally:
        packed.close()
    assert len(report["skipped"]) == 1 and len(report["applied"]) == 1
    assert presentation.slides[1].shapes.title.text == "Still works"


@pytest.mark.parametrize("failure", ["permission", "cancel", "invalid_image", "oversized_image"])
def test_image_authorization_corruption_and_resource_failures_are_hard_errors(failure):
    source = _deck()
    inspection = _inspection(source)
    photo = _named(inspection, "Photo")["id"]
    host_error = PermissionError("authorization revoked")

    def resolve(token):
        if failure == "permission":
            raise host_error
        if failure == "cancel":
            raise DocumentTransformCancelled()
        if failure == "invalid_image":
            return b"not an image"
        return b"x" * (32 * 1024 * 1024 + 1)

    expected = PermissionError if failure == "permission" else DocumentTransformError
    with pytest.raises(expected) as caught:
        _transform(source, inspection, [_op("replace_image", photo, source="asset:safe")], image_resolver=resolve)
    if failure == "permission":
        assert caught.value is host_error
    else:
        assert caught.value.code == {"cancel": "cancelled", "invalid_image": "invalid_image", "oversized_image": "resource_limit"}[failure]


def test_missing_resolver_and_non_asset_tokens_do_not_trigger_io(monkeypatch):
    source = _deck()
    inspection = _inspection(source)
    photo = _named(inspection, "Photo")["id"]

    def forbidden(*args, **kwargs):
        raise AssertionError("No network fetch is permitted")

    monkeypatch.setattr(socket, "create_connection", forbidden)
    output, metadata = _transform(source, inspection, [
        _op("replace_image", photo, source="asset:no-resolver"),
        _op("replace_image", photo, source="local-path"),
    ])
    assert _files(output) == _files(source)
    assert len(metadata["transformation_report"]["skipped"]) == 2


def test_corrupt_shape_ids_and_resource_limits_fail_closed():
    source = _deck()
    files = _files(source)
    slide = _xml(files["ppt/slides/slide1.xml"])
    shapes = slide.findall(f"{_q('p', 'cSld')}/{_q('p', 'spTree')}/{_q('p', 'sp')}")
    duplicate = shapes[0].find(f"{_q('p', 'nvSpPr')}/{_q('p', 'cNvPr')}").get("id")
    shapes[1].find(f"{_q('p', 'nvSpPr')}/{_q('p', 'cNvPr')}").set("id", duplicate)
    files["ppt/slides/slide1.xml"] = _serialize(slide)
    with pytest.raises(DocumentTransformError) as caught:
        _inspection(_zip(files))
    assert caught.value.code == "invalid_package"
    with pytest.raises(DocumentTransformError) as caught:
        _inspection(source, limits=DocumentTransformLimits(max_targets=10))
    assert caught.value.code == "resource_limit"
    inspection = _inspection(source)
    count = len(inspection["targets"])
    with pytest.raises(DocumentTransformError) as caught:
        _transform(
            source, inspection, [_op("clone", "slide:256", clone_id="over-budget")],
            limits=DocumentTransformLimits(max_targets=count + 1),
        )
    assert caught.value.code == "resource_limit"
    with pytest.raises(DocumentTransformCancelled):
        _inspection(source, check=lambda: False)


def test_excessive_numeric_ids_have_stable_corrupt_package_errors():
    source = _deck()
    files = _files(source)
    main = _xml(files["ppt/presentation.xml"])
    main.find(f"{_q('p', 'sldIdLst')}/{_q('p', 'sldId')}").set("id", "9" * 6000)
    files["ppt/presentation.xml"] = _serialize(main)
    with pytest.raises(DocumentTransformError) as caught:
        _inspection(_zip(files))
    assert caught.value.code == "invalid_package"


def test_image_pixels_and_expanded_xml_enforce_bounded_output():
    source = _deck()
    inspection = _inspection(source)
    large = io.BytesIO()
    Image.new("RGB", (400, 400), "blue").save(large, "PNG")
    with pytest.raises(DocumentTransformError) as caught:
        _transform(
            source, inspection, [_op("replace_image", _named(inspection, "Photo")["id"], source="asset:large")],
            limits=DocumentTransformLimits(max_xml_nodes=1200), image_resolver=lambda token: large.getvalue(),
        )
    assert caught.value.code == "resource_limit"
    with pytest.raises(DocumentTransformError) as caught:
        _transform(
            source, inspection, [_op("replace_text", _named(inspection, "Mixed body")["id"], text="line\n" * 500)],
            limits=DocumentTransformLimits(max_xml_nodes=1200),
        )
    assert caught.value.code == "resource_limit"


def test_wrong_notes_owner_is_a_corrupt_package_not_a_soft_skip():
    source = _deck()
    files = _files(source)
    rels = _xml(files["ppt/notesSlides/_rels/notesSlide1.xml.rels"])
    for rel in rels:
        if rel.get("Type").endswith("/slide"):
            rel.set("Target", "../slides/slide2.xml")
    files["ppt/notesSlides/_rels/notesSlide1.xml.rels"] = _serialize(rels)
    with pytest.raises(DocumentTransformError) as caught:
        _inspection(_zip(files))
    assert caught.value.code == "invalid_package"


def test_created_notes_keep_master_placeholder_geometry_and_relationships():
    source = _deck()
    files = _files(source)
    master = _xml(files["ppt/notesMasters/notesMaster1.xml"])
    body = next(
        shape for shape in master.iter(_q("p", "sp"))
        if next(shape.iter(_q("p", "ph"))).get("type") == "body"
    )
    props = body.find(f"{_q('p', 'nvSpPr')}/{_q('p', 'cNvPr')}")
    etree.SubElement(props, _q("a", "hlinkClick"), {_q("r", "id"): "rIdMasterLink"})
    files["ppt/notesMasters/notesMaster1.xml"] = _serialize(master)
    rels = _xml(files["ppt/notesMasters/_rels/notesMaster1.xml.rels"])
    etree.SubElement(rels, _q("rel", "Relationship"), {
        "Id": "rIdMasterLink", "Type": REL_PREFIX + "hyperlink",
        "Target": "https://example.com/notes", "TargetMode": "External",
    })
    files["ppt/notesMasters/_rels/notesMaster1.xml.rels"] = _serialize(rels)
    source = _zip(files)
    inspection = _inspection(source)
    output, _metadata = _transform(source, inspection, [
        _op("replace_text", "slide:257:notes", text="Proper inherited note body"),
    ])
    package, presentation = _consistent(output)
    note_rels = package.relationships("ppt/notesSlides/notesSlide2.xml")
    assert any(rel["external"] and rel["target"] == "https://example.com/notes" for rel in note_rels)
    created = package.parse_xml("ppt/notesSlides/notesSlide2.xml")
    created_body = next(
        shape for shape in created.iter(_q("p", "sp"))
        if next(shape.iter(_q("p", "ph"))).get("type") == "body"
    )
    assert next(created_body.iter(_q("p", "ph"))).attrib == next(body.iter(_q("p", "ph"))).attrib
    assert _serialize(created_body.find(_q("p", "spPr"))) == _serialize(body.find(_q("p", "spPr")))
    assert package.files["ppt/notesMasters/notesMaster1.xml"] == files["ppt/notesMasters/notesMaster1.xml"]
    assert presentation.slides[1].notes_slide.notes_text_frame.text == "Proper inherited note body"


def test_outputs_are_deterministic_and_source_is_untouched():
    source = _deck(template=True)
    digest = hashlib.sha256(source).hexdigest()
    inspection = _inspection(source, ".potx")
    operations = [
        _op("clone", "slide:256", clone_id="copy"),
        _op("replace_text", "copy:shape:2", text="Deterministic output"),
        _op("replace_text", "copy:notes", text="Deterministic notes"),
        _op("replace_image", "copy" + _named(inspection, "Photo")["id"][len("slide:256"):], source="asset:new"),
        _op("remove", "slide:257"),
    ]
    first, first_metadata = _transform(source, inspection, operations, "create_from_reference", image_resolver=lambda token: _png("green"))
    second, second_metadata = _transform(source, inspection, operations, "create_from_reference", image_resolver=lambda token: _png("green"))
    _consistent(first)
    assert first == second and first_metadata == second_metadata
    assert hashlib.sha256(source).hexdigest() == digest


@pytest.mark.parametrize("optimized", [False, True])
def test_cold_adapter_import_and_transform_do_not_initialize_cloud_or_fetch(optimized):
    source = _deck(blank=True, template=True)
    inspection = _inspection(source, ".potx")
    layout = inspection["layouts"][0]["id"]
    payload = {
        "source": base64.b64encode(source).decode(),
        "prepared": _prepared(inspection, [
            _op("clone", layout, clone_id="cold-copy"),
            _op("replace_text", "cold-copy:notes", text="Cold notes"),
        ], "create_from_reference"),
    }
    probe = """
import base64, json, socket, sys
sys.path.insert(0, sys.argv[1])
def forbidden(*args, **kwargs):
    raise AssertionError('Network access is forbidden')
socket.create_connection = forbidden
import functions_document_transform_pptx
from functions_document_transform import transform_document, GuardedPackage
for name in ('config', 'functions_settings', 'functions_appinsights'):
    if name in sys.modules:
        raise AssertionError('Unexpected cloud/bootstrap import: ' + name)
value = json.load(sys.stdin)
with transform_document(base64.b64decode(value['source']), value['prepared']) as result:
    content = result.file_content.getvalue()
    package = GuardedPackage(content, '.pptx')
    if result.metadata['slide_count'] != 1:
        raise AssertionError('Slide clone did not execute')
    if len(result.metadata['transformation_report']['applied']) != 2:
        raise AssertionError('Required operations did not execute')
    print(json.dumps({'digest': result.content_sha256, 'metadata': result.metadata}, sort_keys=True))
"""
    args = [sys.executable, *(['-O'] if optimized else []), "-c", probe, APP_PATH]
    first = subprocess.run(
        args, input=json.dumps(payload), text=True, capture_output=True, timeout=30,
        env={**os.environ, "PYTHONHASHSEED": "7"},
    )
    second = subprocess.run(
        args, input=json.dumps(payload), text=True, capture_output=True, timeout=30,
        env={**os.environ, "PYTHONHASHSEED": "19"},
    )
    assert first.returncode == 0, first.stderr
    assert second.returncode == 0, second.stderr
    assert first.stdout == second.stdout


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-q"]))
