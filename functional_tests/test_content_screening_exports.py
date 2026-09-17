# test_content_screening_exports.py
"""
Functional tests for reviewed-content derivatives.
Version: 0.261.106
Implemented in: 0.261.106

Ensures clean text/CSV outputs use only the candidate's canonical units and
do not reintroduce removed data or spreadsheet formulas from the original.
"""

import csv
import io
import sys
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "application" / "single_app"))

# Import only the framework's pure export helpers.
from content_screening.contracts import ContentUnit, ScreeningValidationError
from content_screening.exports import build_clean_artifact, table_schema_text


class ScreeningExportTests(unittest.TestCase):
    def cell(self, row, column, text, value_type="text"):
        return ContentUnit(f"cell-{row}-{column}", text, {
            "kind": "table_cell", "sheet_index": 1, "row": row,
            "column": column, "value_type": value_type,
        })

    def test_text_derivative_does_not_copy_private_metadata(self):
        units = [
            ContentUnit("page", "Clean surviving text", {"page_number": 4}),
            ContentUnit("metadata", "private-old-summary", {"kind": "metadata"}),
        ]
        artifact = build_clean_artifact(units, "document", "original.pdf")
        self.assertTrue(artifact.file_name.endswith(".txt"))
        self.assertIn(b"Source page 4", artifact.content)
        self.assertNotIn(b"private-old-summary", artifact.content)

    def test_csv_preserves_surviving_cells_and_neutralizes_formulas(self):
        units = [
            self.cell(1, 1, "name"), self.cell(1, 2, "value"),
            self.cell(2, 1, "reviewed"), self.cell(2, 2, "=1+1"),
            self.cell(3, 1, "number"), self.cell(3, 2, "-12", "number"),
        ]
        artifact = build_clean_artifact(units, "document", "original.csv")
        rows = list(csv.reader(io.StringIO(artifact.content.decode("utf-8"))))
        self.assertEqual(rows[1], ["reviewed", "'=1+1"])
        self.assertEqual(rows[2], ["number", "-12"])
        self.assertIn("Rows: 2", table_schema_text(units))

    def test_removing_a_cell_does_not_pull_its_value_from_an_original(self):
        units = [self.cell(1, 1, "name"), self.cell(1, 2, "value"), self.cell(2, 1, "retained")]
        artifact = build_clean_artifact(units, "document", "original.csv")
        rows = list(csv.reader(io.StringIO(artifact.content.decode("utf-8"))))
        self.assertEqual(rows[1], ["retained", ""])

    def test_invalid_typed_edits_cannot_silently_change_data(self):
        with self.assertRaises(ScreeningValidationError):
            build_clean_artifact([self.cell(1, 1, "not a number", "number")], "document", "file.csv")

    def test_metadata_only_candidate_cannot_publish_original_as_fallback(self):
        with self.assertRaises(ScreeningValidationError):
            build_clean_artifact([
                ContentUnit("meta", "title", {"kind": "metadata"}),
            ], "document", "original.pdf")

    def test_workbook_derivative_is_repeatable_and_keeps_formula_text_inert(self):
        from openpyxl import load_workbook

        units = [self.cell(1, 1, "value"), self.cell(2, 1, "=1+1")]
        first = build_clean_artifact(units, "document", "original.xlsx")
        second = build_clean_artifact(units, "document", "original.xlsx")
        self.assertEqual(first.content, second.content)
        with zipfile.ZipFile(io.BytesIO(first.content)) as archive:
            self.assertTrue(all(item.date_time == (2000, 1, 1, 0, 0, 0) for item in archive.infolist()))
            self.assertIn(b"2000-01-01", archive.read("docProps/core.xml"))
        workbook = load_workbook(io.BytesIO(first.content), data_only=False)
        try:
            self.assertEqual(workbook.active["A2"].value, "=1+1")
            self.assertEqual(workbook.active["A2"].data_type, "s")
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
