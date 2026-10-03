#!/usr/bin/env python3
# test_workflow_merge_word.py
"""
Functional test for Word Merge tasks in workflows.
Version: 0.261.222
Implemented in: 0.261.222
Refs: microsoft/simplechat#1619

This test ensures that a workflow Merge task can append Word documents, in order, into one
Word document. The task contract accepts only Word options and the DOCX format, files of the
wrong type are refused when selected and skipped when found at run time, a damaged document
fails the task with its name, a repeated merge reuses its file, the runner attaches the
document, and chat can propose a Word merge whose option errors say which options each kind
takes. No model, Azure service or network is used.
"""

import io
import sys
from pathlib import Path

import pytest

sys.path.append(str(Path(__file__).resolve().parent))

from test_support.versioning import assert_app_version_at_least  # noqa: E402
from test_workflow_merge_task import (  # noqa: E402,F401
    MERGE_HANDLES,
    NORTH,
    MergeWorld,
    WorkflowInputError,
    _codes,
    app,
    drafts,
    dry_run,
    merge_blueprint,
    normalize,
    runner_merge,
)


def docx_bytes(title, lines):
    from docx import Document

    document = Document()
    document.add_heading(title, level=1)
    for line in lines:
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


FILES = {
    "march": ("March minutes.docx", docx_bytes("March", ["Budget approved."])),
    "april": ("April minutes.docx", docx_bytes("April", ["Hiring paused.", "Office move planned."])),
    "broken": ("broken.docx", b"PK\x03\x04garbage"),
    "north": NORTH,
    "notes": ("notes.pdf", b"%PDF-1.7"),
}


def world(**kwargs):
    return MergeWorld(dict(FILES), **kwargs)


def word_action(**changes):
    action = {"type": "merge", "merge_kind": "docx", "document_ids": ["march", "april"]}
    action.update(changes)
    return action


def texts(content):
    from docx import Document

    return [paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs if paragraph.text]


def test_version_includes_word_merge_tasks():
    assert_app_version_at_least("0.261.222")


def test_a_word_task_takes_only_word_options_and_creates_docx(app):
    options = {"formatting": "use_first", "page_breaks": False, "source_headings": True}
    action = normalize(app, word_action(merge_options=options))
    assert (action["merge_kind"], action["output_format"], action["merge_options"]) == ("docx", "docx", options)
    for payload, message in (
        (word_action(output_format="pdf"), "A Word merge creates DOCX files."),
        (word_action(merge_options={"bookmarks": True}), "Some merge options don't apply to Word merges. Remove them."),
        (word_action(merge_options={"formatting": "mixed"}), "Formatting must be keep_source or use_first."),
        (word_action(merge_options={"page_breaks": "yes"}), "The page_breaks merge option must be true or false."),
    ):
        with pytest.raises(app.actions.MergeActionError) as caught:
            normalize(app, payload)
        assert message in str(caught.value)


def test_a_word_merge_appends_the_documents_in_order(app):
    files = world()
    updates = []
    action = normalize(app, word_action(output_file_name="Minutes", merge_options={"source_headings": True}))
    result = files.run(app, action, progress=updates.append)

    [published] = files.published
    assert (published["file_name"], published["output_format"]) == ("Minutes.docx", "docx")
    assert published["summary"] == "Merged 2 file(s) into one Word file."
    assert texts(published["content"]) == [
        "March minutes.docx", "March", "Budget approved.",
        "April minutes.docx", "April", "Hiring paused.", "Office move planned.",
    ]
    assert files.reads == ["march", "april"]
    assert [update["index"] for update in updates] == [1, 2]
    assert result["merge_summary"] == {
        "status": "merged", "kind": "docx", "files": 2, "pages": 0, "slides": 0, "sheets": 0, "skipped": 0,
        "file_name": "Minutes.docx",
    }
    assert result["reply"].startswith("Merged **2 file(s)** into **Minutes.docx**.")
    assert "The merged document uses the first document's page setup, headers and footers." in result["reply"]


def test_files_found_at_run_time_skip_the_wrong_type(app):
    files = world(collected=["march", "north", "notes", "april"])
    result = files.run(app, normalize(app, word_action(target_mode="all", doc_scope="personal")))
    [(_action, extensions, _limit)] = files.collect_calls
    assert extensions == (".docx",)
    assert files.reads == ["march", "april"]
    assert result["merge_summary"]["files"] == 2 and result["merge_summary"]["skipped"] == 2
    assert "- Skipped 2 file(s) that aren't the right type: north.csv, notes.pdf." in result["reply"]


def test_failures_name_the_document_and_a_repeated_merge_reuses_its_file(app):
    with pytest.raises(app.merge.WorkflowMergeError) as wrong_type:
        world().run(app, normalize(app, word_action(document_ids=["march", "notes"])))
    assert str(wrong_type.value) == "notes.pdf isn't a Word document (.docx), so it can't be merged."

    files = world()
    with pytest.raises(app.merge.WorkflowMergeError) as broken:
        files.run(app, normalize(app, word_action(document_ids=["march", "broken"])))
    assert str(broken.value) == "broken.docx isn't a valid Word document."
    assert not files.published

    first, second = world(), world()
    first.run(app, normalize(app, word_action()))
    second.run(app, normalize(app, word_action()))
    assert first.published[0]["content"] == second.published[0]["content"]
    assert first.published[0]["idempotency_key"] == second.published[0]["idempotency_key"]


def test_the_runner_attaches_the_word_document(app, monkeypatch):
    run, uploads, thoughts = runner_merge(app, monkeypatch, world())
    result = run(normalize(app, word_action()))
    [upload] = uploads
    assert (upload["file_name"], upload["output_format"], upload["capability"]) == ("merged.docx", "docx", "file_merge")
    assert upload["content"].startswith(b"PK")
    [artifact] = result["generated_analysis_artifacts"]
    assert (artifact["file_name"], artifact["output_format"]) == ("merged.docx", "docx")
    assert [(entry["content"], entry["detail"]) for entry in thoughts] == [
        ("Merged file 1 of 2", None), ("Merged file 2 of 2", None), ("Merged files", "files=2"),
    ]

    run, uploads, _thoughts = runner_merge(app, monkeypatch, world())
    with pytest.raises(WorkflowInputError) as refused:
        run(normalize(app, word_action(document_ids=["march", "broken"])))
    assert refused.value.public_message == "broken.docx isn't a valid Word document."
    assert not uploads


def test_chat_can_propose_a_word_merge_task(drafts):
    result = dry_run(drafts, merge_blueprint(
        {"kind": "docx", "files": "inputs", "options": {"source_headings": True}}, inputs=["south", "north"],
    ), MERGE_HANDLES)
    assert result["ok"] is True, result["errors"]
    action = result["workflow"]["tasks"][0]["document_action"]
    assert (action["merge_kind"], action["output_format"], action["merge_options"]) == (
        "docx", "docx", {"source_headings": True},
    )
    assert action["document_ids"] == ["doc-south", "doc-north"]

    wrong_option = dry_run(drafts, merge_blueprint({"kind": "docx", "files": "all", "options": {"bookmarks": True}}))
    assert _codes(wrong_option) == [("merge_options_invalid", "/tasks/0/merge/options")]
    # The repair hint says which options each kind takes, so the planner can fix the right one.
    assert "docx takes formatting, page_breaks, source_headings" in wrong_option["errors"][0]["message"]
    wrong_format = merge_blueprint({"kind": "docx", "files": "all", "output_format": "pdf"})
    assert _codes(dry_run(drafts, wrong_format)) == [("merge_format_invalid", "/tasks/0/merge/output_format")]


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
