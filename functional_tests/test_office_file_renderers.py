# test_office_file_renderers.py
#!/usr/bin/env python3
"""Functional tests for the headless prepared-content Office rendering services.

Version: 0.261.126
Implemented in: 0.261.126
Refs: microsoft/simplechat#1509 (bounded M8 service preparation, not activation).

Reopen real XLSX, DOCX, PDF, and PPTX bytes; verify rich content and exact record
counts, access/cancellation fences, resource cleanup, and safe failure classes.
The 30,000-row workbook test uses a single-pass generator and measures renderer
allocation rather than constructing a second complete dataset. Fresh normal and
optimized Python processes forbid network, application bootstrap, model clients,
publication imports, and temporary-file creation while exercising all renderers.
No credentials, Flask app, cloud services, browsers, or external providers are used.
"""

import hashlib
import importlib
import io
import math
import os
import socket
import subprocess
import sys
import tempfile
import tracemalloc
import unittest
from contextlib import ExitStack
from dataclasses import replace
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree
from zipfile import ZipFile

import fitz
from docx import Document
from docx.document import Document as DocumentType
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation


REPO_ROOT = Path(__file__).resolve().parents[1]
APP_PATH = REPO_ROOT / "application" / "single_app"
NAMESPACES = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "rels": "http://schemas.openxmlformats.org/package/2006/relationships",
    "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
}

# Import the real standalone service without retaining a test-only search path.
with patch.object(sys, "path", [str(APP_PATH), *sys.path]):
    renderers = importlib.import_module("functions_office_file_renderers")
    export_visuals = importlib.import_module("functions_export_visuals")


def image_bytes(size=(32, 24), dpi=None):
    with Image.new("RGB", size, (34, 139, 99)) as image, io.BytesIO() as stream:
        image.save(stream, format="PNG", **({"dpi": dpi} if dpi is not None else {}))
        return stream.getvalue()


def rich_report():
    return renderers.PreparedReport(
        title="Prepared <report> & findings",
        content=(
            '# Café 中文 Ελληνικά\n\n'
            'A **bold** finding, *emphasis*, ~~old~~, `x < 4`, and '
            '[citation 1](https://example.test/report?a=1&b=2).\n\n'
            '3. Third result\n4. Fourth result\n\n'
            '- First bullet\n    - Nested bullet\n\n'
            '| Name | Count |\n| :--- | ---: |\n| Café & <literal> | 12 |\n\n'
            '```python\nvalue = "<&>"\nprint(value)\n```\n\n'
            '![Authorized image](asset:figure)\n\n'
            '> A quoted finding.\n\n'
            '---\n'
        ).replace("<literal>", "&lt;literal&gt;"),
        images={"asset:figure": image_bytes()},
    )


def prepared_deck():
    return renderers.PreparedDeck(
        title="Prepared findings",
        expected_slide_count=2,
        images={"figure": image_bytes()},
        slides=(
            renderers.PreparedSlide(
                title="Café 中文 & findings",
                shapes=(
                    renderers.SlideTextBox(
                        renderers.SlideBox(0.5, 1.5, 5, 3),
                        (
                            renderers.SlideParagraph("Prepared body <&>", bold=True),
                            renderers.SlideParagraph("First finding", list_kind="bullet"),
                            renderers.SlideParagraph("Nested evidence", list_kind="bullet", level=1),
                            renderers.SlideParagraph(
                                "Citation 1", italic=True, url="https://example.test/evidence?a=1&b=2",
                            ),
                        ),
                    ),
                    renderers.SlideTable(
                        renderers.SlideBox(6, 1.5, 6, 3),
                        (("Name", "Count"), ("Café & <item>", "12"), ("Last", "99")),
                    ),
                    renderers.SlideImage(renderers.SlideBox(0.5, 5, 2, 1.5), "figure", "Supplied visual"),
                ),
                notes="Speaker notes: preserve <&> and Ελληνικά.\nSecond line.",
            ),
            renderers.PreparedSlide(
                title="",
                layout="blank",
                shapes=(
                    renderers.SlideTextBox(
                        renderers.SlideBox(1, 1, 10, 4),
                        (
                            renderers.SlideParagraph("Last prepared slide"),
                            renderers.SlideParagraph("Numbered item", list_kind="number"),
                        ),
                    ),
                ),
                notes="Last notes",
            ),
        ),
    )


def small_workbook():
    return renderers.PreparedWorkbook((
        renderers.PreparedSheet("Records", ("id", "text"), ((1, "first"), (2, "last")), 2),
    ))


class OfficeRendererTests(unittest.TestCase):
    large_workbook_peak_bytes = 0

    def assert_render_error(self, code, renderer, source, **kwargs):
        with self.assertRaises(renderers.OfficeRenderError) as failure:
            result = renderer(source, **kwargs)
            result.close()
        self.assertEqual(failure.exception.code, code)
        self.assertEqual(failure.exception.retryable, code in {"source_unavailable", "render_io"})
        return failure.exception

    def test_xlsx_30000_complete_single_pass_rows_and_type_fidelity(self):
        row_count = 30_000
        observations = {"iterations": 0, "yielded": 0, "closed": False, "checks": 0}
        stamp = datetime(2026, 9, 21, 11, 4, 42, 123000)

        class Rows:
            def __iter__(self):
                observations["iterations"] += 1
                try:
                    for index in range(row_count):
                        observations["yielded"] += 1
                        yield (index, f"record-{index}-中文", index % 2 == 0, index + 0.5, stamp, None)
                finally:
                    observations["closed"] = True

        def recheck():
            observations["checks"] += 1

        columns = ("id", "label", "enabled", "value", "created", "empty")
        workbook = renderers.PreparedWorkbook((
            renderers.PreparedSheet("Ordered data", columns, Rows(), row_count),
        ))
        tracemalloc.start()
        try:
            with patch("openpyxl.worksheet._writer.create_temporary_file", side_effect=AssertionError("No disk spool")):
                output = renderers.render_prepared_xlsx(
                    workbook, checks=renderers.OfficeRenderChecks(source_check=recheck),
                )
            _, peak = tracemalloc.get_traced_memory()
        finally:
            tracemalloc.stop()
        self.__class__.large_workbook_peak_bytes = peak
        with output:
            book = load_workbook(output.file_content, read_only=True, data_only=False)
            try:
                rows = book["Ordered data"].iter_rows(values_only=True)
                header = next(rows)
                first = next(rows)
                last = first
                actual_count = 1
                for last in rows:
                    actual_count += 1
            finally:
                book.close()
            self.assertEqual(header, columns)
            self.assertEqual(actual_count, row_count)
            self.assertEqual(first, (0, "record-0-中文", True, 0.5, stamp, None))
            self.assertEqual(last, (29999, "record-29999-中文", False, 29999.5, stamp, None))
            self.assertIs(type(first[0]), int)
            self.assertIs(type(first[1]), str)
            self.assertIs(type(first[2]), bool)
            self.assertIs(type(first[3]), float)
            self.assertIs(type(first[4]), datetime)
            self.assertEqual(output.metadata["record_count"], row_count)
            self.assertEqual(output.metadata["cell_count"], (row_count + 1) * len(columns))
        self.assertEqual(observations["iterations"], 1)
        self.assertEqual(observations["yielded"], row_count)
        self.assertTrue(observations["closed"])
        self.assertGreater(observations["checks"], 10)
        self.assertLess(peak, 64 * 1024 * 1024, "Write-only rendering must not retain a cell grid")
        self.assertTrue(output.file_content.closed)

    def test_xlsx_order_safe_text_empty_strings_and_native_date_types(self):
        text_values = (
            "=HYPERLINK(\"https://example.test\", \"text\")", "+1+1", "-1+1", "@SUM(A1)",
            "\t=1+1", "#DIV/0!", "00123", "", "<tag>& \"quotes\"", "line 1\nline 2", "é 中文",
        )
        values = (
            *text_values, True, 42, 12.5, date(2026, 1, 2), datetime(2026, 1, 2, 3, 4),
            time(3, 4, 5, 123000), timedelta(days=2, seconds=3), None,
        )
        columns = tuple(f"col_{index}" for index in range(len(values)))
        row = dict(reversed(list(zip(columns, values))))
        workbook = renderers.PreparedWorkbook((
            renderers.PreparedSheet("Second", columns, [row], 1),
            renderers.PreparedSheet("First", ("zero",), (), 0),
        ))
        with renderers.render_prepared_xlsx(workbook) as output:
            book = load_workbook(output.file_content, read_only=True)
            try:
                names = book.sheetnames
                rendered_rows = list(book["Second"].iter_rows())
                actual = tuple(cell.value for cell in rendered_rows[1])
                text_types = [cell.data_type for cell in rendered_rows[1][:len(text_values)]]
                empty_sheet = list(book["First"].values)
            finally:
                book.close()
            output.file_content.seek(0)
            with ZipFile(output.file_content) as archive:
                worksheet = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
                formulas = worksheet.findall(".//s:f", NAMESPACES)
            self.assertEqual(names, ["Second", "First"])
            self.assertEqual(actual, values)
            self.assertEqual(text_types, ["s"] * len(text_values))
            self.assertEqual(empty_sheet, [("zero",)])
            self.assertEqual(formulas, [])
            self.assertEqual(output.metadata["record_count"], 1)

    def test_xlsx_rejects_preview_and_count_or_schema_mismatches(self):
        base = small_workbook()
        cases = (
            replace(base, complete=False),
            renderers.PreparedWorkbook((renderers.PreparedSheet("Data", ("id",), [(1,)], 2),)),
            renderers.PreparedWorkbook((renderers.PreparedSheet("Data", ("id",), [(1,), (2,)], 1),)),
            renderers.PreparedWorkbook((renderers.PreparedSheet("Data", ("id",), [{"id": 1, "extra": 2}], 1),)),
            renderers.PreparedWorkbook((renderers.PreparedSheet("Data", ("id",), [{"wrong": 1}], 1),)),
            renderers.PreparedWorkbook((renderers.PreparedSheet("Data", ("id",), [(1, 2)], 1),)),
        )
        for source in cases:
            with self.subTest(source=source):
                self.assert_render_error("incomplete_source", renderers.render_prepared_xlsx, source)

    def test_xlsx_supported_float_boundaries_reopen_as_equivalent_finite_numbers(self):
        values = (0.0, 1.0, -1.0, 0.1, 1.5, math.nextafter(1.0, 0.0), 1e308, -1e308, 5e-324)
        source = renderers.PreparedWorkbook((
            renderers.PreparedSheet("Numbers", ("value",), [(value,) for value in values], len(values)),
        ))
        with renderers.render_prepared_xlsx(source) as output:
            book = load_workbook(output.file_content, read_only=True)
            try:
                rows = list(book.active.iter_rows(min_row=2))
                actual = [row[0].value for row in rows]
                cell_types = [row[0].data_type for row in rows]
            finally:
                book.close()
        self.assertEqual(cell_types, ["n"] * len(values))
        for expected, restored in zip(values, actual):
            with self.subTest(value=expected):
                self.assertTrue(math.isfinite(restored))
                self.assertEqual(restored, expected)
                self.assertEqual(math.copysign(1, restored), math.copysign(1, expected))

    def test_xlsx_rejects_float_rounding_overflow_and_loss_of_negative_zero(self):
        for value in (1.0000000000000002, sys.float_info.max, -sys.float_info.max, sys.float_info.min, -0.0):
            with self.subTest(value=value):
                source = renderers.PreparedWorkbook((
                    renderers.PreparedSheet("Numbers", ("value",), [(value,)], 1),
                ))
                self.assert_render_error("unsupported_content", renderers.render_prepared_xlsx, source)

    def test_xlsx_enforces_excel_and_configured_limits_before_iteration(self):
        observed = []

        class UnreadableRows:
            def __iter__(self):
                observed.append(True)
                raise AssertionError("Invalid declared dimensions must fail before consuming rows")

        for sheet in (
            renderers.PreparedSheet("Rows", ("id",), UnreadableRows(), 1_048_576),
            renderers.PreparedSheet("Columns", tuple(f"c{i}" for i in range(16_385)), UnreadableRows(), 0),
        ):
            with self.subTest(name=sheet.name):
                self.assert_render_error("limit_exceeded", renderers.render_prepared_xlsx, renderers.PreparedWorkbook((sheet,)))
        self.assertEqual(observed, [])
        for limits in (
            renderers.OfficeRenderLimits(max_rows=1),
            renderers.OfficeRenderLimits(max_columns=1),
            renderers.OfficeRenderLimits(max_cells=5),
        ):
            self.assert_render_error("limit_exceeded", renderers.render_prepared_xlsx, small_workbook(), limits=limits)
        long_cell = renderers.PreparedWorkbook((renderers.PreparedSheet("Data", ("value",), [("x" * 32768,)], 1),))
        self.assert_render_error("limit_exceeded", renderers.render_prepared_xlsx, long_cell)

    def test_xlsx_rejects_invalid_names_and_lossy_or_unsupported_cells(self):
        for name in ("bad/name", "'quoted", "x" * 32, "   ", "bad\x00name"):
            source = renderers.PreparedWorkbook((renderers.PreparedSheet(name, ("value",), (), 0),))
            self.assert_render_error("invalid_input", renderers.render_prepared_xlsx, source)
        duplicates = renderers.PreparedWorkbook((
            renderers.PreparedSheet("Data", ("id",), (), 0),
            renderers.PreparedSheet("data", ("id",), (), 0),
        ))
        self.assert_render_error("invalid_input", renderers.render_prepared_xlsx, duplicates)
        for value in (
            {"nested": "value"}, ["array"], {"formula": "=1+1"}, float("nan"), float("inf"),
            10 ** 18, datetime(2026, 1, 1, tzinfo=timezone.utc), datetime(2026, 1, 1, microsecond=1),
        ):
            with self.subTest(value=value):
                source = renderers.PreparedWorkbook((renderers.PreparedSheet("Data", ("value",), [(value,)], 1),))
                self.assert_render_error("unsupported_content", renderers.render_prepared_xlsx, source)

    def test_docx_preserves_rich_blocks_images_links_unicode_and_escaping(self):
        source = rich_report()
        with renderers.render_prepared_docx(source) as output:
            doc = Document(output.file_content)
            paragraphs = [paragraph.text for paragraph in doc.paragraphs]
            styles = [paragraph.style.name for paragraph in doc.paragraphs]
            table_values = [[cell.text for cell in row.cells] for row in doc.tables[0].rows]
            image_count = len(doc.inline_shapes)
            output.file_content.seek(0)
            with ZipFile(output.file_content) as archive:
                xml = ElementTree.fromstring(archive.read("word/document.xml"))
                numbering = ElementTree.fromstring(archive.read("word/numbering.xml"))
                relationships = ElementTree.fromstring(archive.read("word/_rels/document.xml.rels"))
            targets = [node.attrib["Target"] for node in relationships]
            starts = [node.attrib[f"{{{NAMESPACES['w']}}}val"] for node in numbering.findall(".//w:start", NAMESPACES)]
            self.assertIn("Café 中文 Ελληνικά", paragraphs)
            self.assertIn("Heading 1", styles)
            self.assertIn('value = "<&>"\nprint(value)\n', paragraphs)
            self.assertIn("A quoted finding.", paragraphs)
            self.assertEqual(table_values, [["Name", "Count"], ["Café & <literal>", "12"]])
            self.assertEqual(image_count, 1)
            self.assertIn("https://example.test/report?a=1&b=2", targets)
            self.assertIn("3", starts)
            self.assertIsNotNone(xml.find(".//w:hyperlink", NAMESPACES))
            self.assertIsNotNone(xml.find(".//w:b", NAMESPACES))
            self.assertIsNotNone(xml.find(".//w:strike", NAMESPACES))
            self.assertIsNotNone(xml.find(".//w:numPr", NAMESPACES))
            self.assertEqual(output.metadata["table_count"], 1)
            self.assertEqual(output.metadata["image_count"], 1)
            self.assertEqual(output.metadata["content_char_count"], len(source.content))

    def test_pdf_preserves_rich_content_unicode_images_links_and_list_start(self):
        with renderers.render_prepared_pdf(rich_report()) as output:
            content = output.file_content.read()
            with fitz.open(stream=content, filetype="pdf") as doc:
                text = "\n".join(
                    page.get_text(flags=fitz.TEXTFLAGS_TEXT & ~fitz.TEXT_PRESERVE_LIGATURES)
                    for page in doc
                )
                images = sum(len(page.get_image_info()) for page in doc)
                links = [link for page in doc for link in page.get_links()]
                pages = len(doc)
            text = " ".join(text.split())
            for expected in (
                "Prepared <report> & findings", "Café", "中文", "Ελληνικά",
                "bold finding", "Nested bullet", "3. Third result", "4. Fourth result",
                'value = "<&>"', "print(value)", "Café & <literal>", "A quoted finding.",
            ):
                self.assertIn(expected, text)
            self.assertEqual(images, 1)
            self.assertTrue(any(link.get("uri") == "https://example.test/report?a=1&b=2" for link in links))
            self.assertEqual(output.metadata["page_count"], pages)

    def test_docx_anisotropic_dpi_images_fit_page_and_actual_containers(self):
        source = renderers.PreparedReport(
            "![Root](square)\n\n> ![Quote](large)\n\n- ![List](large)\n\n"
            "| Image | Description |\n| --- | --- |\n| ![Cell](large) | Prepared |\n\n"
            "![Portrait](portrait)",
            images={
                "square": image_bytes((600, 600), dpi=(300, 72)),
                "large": image_bytes((1200, 1200), dpi=(300, 72)),
                "portrait": image_bytes((600, 3600), dpi=(72, 300)),
            },
        )
        with renderers.render_prepared_docx(source) as output:
            document = Document(output.file_content)
            sizes = [(shape.width.inches, shape.height.inches) for shape in document.inline_shapes]
            section = document.sections[0]
            page_width = (section.page_width - section.left_margin - section.right_margin) / 914400
            page_height = (section.page_height - section.top_margin - section.bottom_margin) / 914400
            cell_width = document.tables[0].cell(1, 0).width.inches
        self.assertEqual(len(sizes), 5)
        self.assertAlmostEqual(sizes[0][0], 6.25, places=5)
        for width, height in sizes:
            self.assertGreater(width, 0)
            self.assertGreater(height, 0)
            self.assertLessEqual(width, page_width)
            self.assertLessEqual(height, page_height)
        for width, height in sizes[:4]:
            self.assertAlmostEqual(width, height, places=5)
        self.assertLessEqual(sizes[1][0], page_width - 0.3)
        self.assertLessEqual(sizes[2][0], page_width - 0.25)
        self.assertLessEqual(sizes[3][0], cell_width - 0.15)
        self.assertAlmostEqual(sizes[4][1], page_height, places=5)
        self.assertAlmostEqual(sizes[4][1] / sizes[4][0], 6, places=5)

    def test_docx_tight_lists_preserve_inline_separators_and_citations(self):
        source = renderers.PreparedReport(
            "- **Important** **finding**\n"
            "- **Evidence** [citation](https://example.test/citation)\n"
            "- `code` *emphasis* **strong** [1](https://example.test/one)\n"
            "- before **bold** after"
        )
        with renderers.render_prepared_docx(source) as output:
            document = Document(output.file_content)
            text = [paragraph.text.strip() for paragraph in document.paragraphs]
        self.assertEqual(text, [
            "Important finding", "Evidence citation", "code emphasis strong 1", "before bold after",
        ])

    def test_docx_block_and_nested_lists_ignore_only_structural_whitespace(self):
        source = renderers.PreparedReport(
            "- **First** *block*\n\n    `second` [source](https://example.test)\n\n"
            "    - **Nested** *item*"
        )
        with renderers.render_prepared_docx(source) as output:
            document = Document(output.file_content)
            text = [paragraph.text.strip() for paragraph in document.paragraphs]
        self.assertEqual(text, ["First block", "second source", "Nested item"])

    def test_plain_text_pdf_preserves_painted_indentation_and_column_gaps(self):
        source = renderers.PreparedReport("Origin\n    Name    Amount\n    A       42", content_format="text")
        with renderers.render_prepared_pdf(source) as output:
            payload = output.file_content.read()
            with fitz.open(stream=payload, filetype="pdf") as document:
                words = {word[4]: word for page in document for word in page.get_text("words")}
                text = "".join(page.get_text() for page in document)
        advance = words["A"][2] - words["A"][0]
        self.assertIn("    Name    Amount", text)
        self.assertIn("    A       42", text)
        self.assertAlmostEqual(words["Name"][0] - words["Origin"][0], 4 * advance, delta=0.05)
        self.assertAlmostEqual(words["Name"][0], words["A"][0], delta=0.05)
        self.assertAlmostEqual(words["Amount"][0], words["42"][0], delta=0.05)
        self.assertAlmostEqual(words["Amount"][0] - words["Name"][2], 4 * advance, delta=0.05)

    def test_plain_text_pdf_spacing_differs_from_markdown_whitespace_semantics(self):
        positions = {}
        for content_format in ("text", "markdown"):
            for content in ("A    B", "A B"):
                with renderers.render_prepared_pdf(renderers.PreparedReport(
                    content, content_format=content_format,
                )) as output:
                    payload = output.file_content.read()
                    with fitz.open(stream=payload, filetype="pdf") as document:
                        words = {word[4]: word for word in document[0].get_text("words")}
                positions[content_format, content] = words
        wide = positions["text", "A    B"]
        narrow = positions["text", "A B"]
        advance = narrow["A"][2] - narrow["A"][0]
        self.assertAlmostEqual(wide["B"][0] - narrow["B"][0], 3 * advance, delta=0.05)
        self.assertAlmostEqual(
            positions["markdown", "A    B"]["B"][0], positions["markdown", "A B"]["B"][0], delta=0.05,
        )

    def test_plain_text_pdf_preserves_blank_line_height(self):
        positions = []
        for text in ("First\nLast", "First\n\nLast"):
            with renderers.render_prepared_pdf(renderers.PreparedReport(text, content_format="text")) as output:
                payload = output.file_content.read()
                with fitz.open(stream=payload, filetype="pdf") as document:
                    words = {word[4]: word for word in document[0].get_text("words")}
            positions.append(words["Last"][1] - words["First"][1])
        self.assertAlmostEqual(positions[1], 2 * positions[0], delta=0.05)

    def test_reports_support_plain_text_empty_documents_and_internal_links(self):
        text_source = renderers.PreparedReport("<script>& literal\nSecond line", content_format="text")
        internal = renderers.PreparedReport(
            '<h2 id="finding">Prepared finding</h2>\n\n[See finding](#finding)',
        )
        for renderer in (renderers.render_prepared_docx, renderers.render_prepared_pdf):
            with self.subTest(renderer=renderer.__name__):
                with renderer(text_source) as output:
                    if output.output_format == "docx":
                        doc = Document(output.file_content)
                        text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
                    else:
                        data = output.file_content.read()
                        with fitz.open(stream=data, filetype="pdf") as doc:
                            text = "".join(page.get_text() for page in doc)
                    self.assertIn("<script>& literal", text)
                    self.assertIn("Second line", text)
                with renderer(renderers.PreparedReport("", expected_block_count=0)) as output:
                    self.assertEqual(output.metadata["block_count"], 0)
                    self.assertGreater(output.size_bytes, 0)
                with renderer(internal) as output:
                    self.assertEqual(output.metadata["link_count"], 1)

    def test_existing_authorized_export_visual_wrappers_preserve_image_and_caption(self):
        content = export_visuals.build_export_visual_html(
            "chart", export_visuals.build_export_visual_png_data_uri(image_bytes()),
            alt_text="Already prepared chart", caption_text="Existing authorized chart",
        )
        source = renderers.PreparedReport(content)
        with renderers.render_prepared_docx(source) as output:
            document = Document(output.file_content)
            caption = "\n".join(paragraph.text for paragraph in document.paragraphs)
            pictures = len(document.inline_shapes)
            self.assertEqual(pictures, 1)
            self.assertIn("Existing authorized chart", caption)
        with renderers.render_prepared_pdf(source) as output:
            payload = output.file_content.read()
            with fitz.open(stream=payload, filetype="pdf") as document:
                caption = "".join(page.get_text() for page in document)
                pictures = sum(len(page.get_image_info()) for page in document)
            self.assertEqual(pictures, 1)
            self.assertIn("Existing authorized chart", caption)

    def test_reports_reject_unsupported_html_and_missing_images_instead_of_dropping_them(self):
        invalid = (
            "<script>not inert</script>", "<iframe src='https://example.test'></iframe>",
            "<p style='font-size:300pt'>Unexpected style</p>",
            "<table><tr><td>A</td><td>B</td></tr><tr><td>C</td></tr></table>",
            "<table><tr><td colspan='2'>Merged cell</td></tr></table>",
            "<pre><img src='image'></pre>",
            "[Unsafe](javascript:invalid)", "[File](file:///private/document)",
            "![URL](https://example.test/unfetched.png)", "![Path](C:\\private\\image.png)",
            "[Unknown anchor](#missing)",
            "Before</div>\n\nContent after an unmatched wrapper.",
            '<a href="https://example.test/one" href="https://example.test/two">Ambiguous link</a>',
        )
        for renderer in (renderers.render_prepared_docx, renderers.render_prepared_pdf):
            for content in invalid:
                with self.subTest(renderer=renderer.__name__, content=content):
                    self.assert_render_error("unsupported_content", renderer, renderers.PreparedReport(content))

    def test_reports_enforce_completeness_node_and_image_limits(self):
        for renderer in (renderers.render_prepared_docx, renderers.render_prepared_pdf):
            self.assert_render_error("incomplete_source", renderer, renderers.PreparedReport("Text", complete=False))
            self.assert_render_error("incomplete_source", renderer, renderers.PreparedReport("Text", expected_block_count=2))
            self.assert_render_error(
                "limit_exceeded", renderer, rich_report(),
                limits=renderers.OfficeRenderLimits(max_report_nodes=3),
            )
            self.assert_render_error(
                "limit_exceeded", renderer, rich_report(),
                limits=renderers.OfficeRenderLimits(max_image_pixels=1),
            )
            self.assert_render_error(
                "limit_exceeded", renderer,
                renderers.PreparedReport("![One](asset)\n\n![Two](asset)", images={"asset": image_bytes()}),
                limits=renderers.OfficeRenderLimits(max_images=1),
            )
            self.assert_render_error(
                "invalid_input", renderer,
                renderers.PreparedReport("![Broken](asset)", images={"asset": b"not an image"}),
            )
            self.assert_render_error(
                "invalid_input", renderer, renderers.PreparedReport("<p>&#xD800;</p>"),
            )
            self.assert_render_error(
                "limit_exceeded", renderer,
                renderers.PreparedReport(
                    "![One](first)\n\n![Two](second)",
                    images={"first": image_bytes(), "second": image_bytes()},
                ),
                limits=renderers.OfficeRenderLimits(max_total_image_pixels=1000),
            )

    def test_pdf_paginates_complete_report_and_rejects_page_and_layout_overflow(self):
        source = renderers.PreparedReport("\n\n".join(f"Prepared record {index:04d} with its complete findings." for index in range(160)))
        with renderers.render_prepared_pdf(source) as output:
            payload = output.file_content.read()
            with fitz.open(stream=payload, filetype="pdf") as document:
                text = "\n".join(page.get_text() for page in document)
                count = len(document)
            self.assertGreater(count, 1)
            self.assertIn("Prepared record 0000", text)
            self.assertIn("Prepared record 0159", text)
            self.assertEqual(output.metadata["block_count"], 160)
        self.assert_render_error(
            "limit_exceeded", renderers.render_prepared_pdf, source,
            limits=renderers.OfficeRenderLimits(max_pages=1),
        )
        for content in (
            "```\n" + "wide" * 120 + "\n```",
            "wide" * 120,
            "| Long |\n| --- |\n| " + "wide" * 120 + " |",
        ):
            self.assert_render_error(
                "layout_overflow", renderers.render_prepared_pdf, renderers.PreparedReport(content),
            )

    def test_pptx_preserves_all_prepared_slides_shapes_tables_links_notes_and_images(self):
        source = prepared_deck()
        with renderers.render_prepared_pptx(source) as output:
            deck = Presentation(output.file_content)
            slides = list(deck.slides)
            titles = [slide.shapes.title.text if slide.shapes.title else "" for slide in slides]
            notes = [slide.notes_slide.notes_text_frame.text for slide in slides]
            body = "\n".join(
                shape.text for slide in slides for shape in slide.shapes if shape.has_text_frame
            )
            tables = [shape.table for slide in slides for shape in slide.shapes if shape.has_table]
            table_values = [[cell.text for cell in row.cells] for row in tables[0].rows]
            output.file_content.seek(0)
            with ZipFile(output.file_content) as archive:
                xml = ElementTree.fromstring(archive.read("ppt/slides/slide1.xml"))
                relationships = ElementTree.fromstring(archive.read("ppt/slides/_rels/slide1.xml.rels"))
                images = [name for name in archive.namelist() if name.startswith("ppt/media/")]
            self.assertEqual(titles, ["Café 中文 & findings", ""])
            self.assertEqual(notes, [slide.notes for slide in source.slides])
            self.assertIn("Prepared body <&>", body)
            self.assertIn("Nested evidence", body)
            self.assertIn("Last prepared slide", body)
            self.assertEqual(table_values, [["Name", "Count"], ["Café & <item>", "12"], ["Last", "99"]])
            self.assertEqual(len(images), 1)
            self.assertIsNotNone(xml.find(".//a:buChar", NAMESPACES))
            self.assertIsNotNone(xml.find(".//a:hlinkClick", NAMESPACES))
            self.assertTrue(any(node.attrib["Target"] == "https://example.test/evidence?a=1&b=2" for node in relationships))
            self.assertEqual(output.metadata["slide_count"], 2)
            self.assertEqual(output.metadata["shape_count"], 4)
            self.assertEqual(output.metadata["notes_count"], 2)

    def test_pptx_rejects_malformed_slides_and_unsupported_shapes_or_layouts(self):
        base = prepared_deck()
        bad_slides = (
            {"title": "Unexpected mapping", "shapes": []},
            replace(base.slides[0], layout="chart"),
            replace(base.slides[0], shapes=({"chart": [1, 2]},)),
            replace(base.slides[0], layout="blank"),
            replace(base.slides[0], title=""),
            replace(base.slides[0], shapes=(renderers.SlideTable(renderers.SlideBox(1, 2, 5, 3), (("a", "b"), ("c",))),)),
        )
        for slide in bad_slides:
            with self.subTest(slide=slide):
                code = "invalid_input" if isinstance(slide, renderers.PreparedSlide) and slide.title == "" else "unsupported_content"
                self.assert_render_error(code, renderers.render_prepared_pptx, replace(base, slides=(slide,), expected_slide_count=1))
        self.assert_render_error("incomplete_source", renderers.render_prepared_pptx, replace(base, expected_slide_count=3))
        self.assert_render_error("incomplete_source", renderers.render_prepared_pptx, replace(base, complete=False))
        self.assert_render_error(
            "limit_exceeded", renderers.render_prepared_pptx, base,
            limits=renderers.OfficeRenderLimits(max_slides=1),
        )

    def test_pptx_normalizes_crlf_cr_and_lf_in_all_prepared_text_surfaces(self):
        outputs = []
        for newline in ("\r\n", "\r", "\n"):
            source = renderers.PreparedDeck(
                (renderers.PreparedSlide(
                    "Prepared title",
                    shapes=(
                        renderers.SlideTextBox(
                            renderers.SlideBox(0.5, 1.5, 5, 3),
                            (renderers.SlideParagraph(f"First{newline}Second"),),
                        ),
                        renderers.SlideTable(
                            renderers.SlideBox(6, 1.5, 6, 3),
                            ((f"Head{newline}One", "Value"), (f"Cell{newline}Two", "42")),
                        ),
                    ),
                    notes=f"Speaker{newline}Notes",
                ),),
                1,
                title=f"Deck{newline}Title",
            )
            with renderers.render_prepared_pptx(source) as output:
                payload = output.file_content.read()
                with io.BytesIO(payload) as stream:
                    presentation = Presentation(stream)
                    slide = presentation.slides[0]
                    title_id = slide.shapes.title.shape_id
                    body = [shape.text for shape in slide.shapes if shape.has_text_frame and shape.shape_id != title_id]
                    table = next(shape.table for shape in slide.shapes if shape.has_table)
                    cells = [[cell.text for cell in row.cells] for row in table.rows]
                    notes = slide.notes_slide.notes_text_frame.text
                    title = presentation.core_properties.title
                with io.BytesIO(payload) as stream, ZipFile(stream) as archive:
                    xml = b"".join(archive.read(name) for name in (
                        "ppt/slides/slide1.xml", "ppt/notesSlides/notesSlide1.xml", "docProps/core.xml",
                    ))
            self.assertEqual(body, ["First\vSecond"])
            self.assertEqual(cells, [["Head\vOne", "Value"], ["Cell\vTwo", "42"]])
            self.assertEqual(notes, "Speaker\nNotes")
            self.assertEqual(title, "Deck\nTitle")
            self.assertNotIn(b"_x000D_", xml)
            outputs.append(payload)
        self.assertEqual(outputs[0], outputs[1])
        self.assertEqual(outputs[1], outputs[2])

    def test_pptx_normalized_slide_title_breaks_still_enforce_layout_limits(self):
        for newline in ("\r\n", "\r", "\n"):
            source = renderers.PreparedDeck((renderers.PreparedSlide(f"First{newline}Second"),), 1)
            self.assert_render_error("layout_overflow", renderers.render_prepared_pptx, source)

    def test_pptx_rejects_off_slide_overlap_and_text_clipping(self):
        shapes = (
            (renderers.SlideTextBox(renderers.SlideBox(13, 2, 3, 2), (renderers.SlideParagraph("Off slide"),)),),
            (renderers.SlideTextBox(renderers.SlideBox(1, 0.3, 4, 1), (renderers.SlideParagraph("Over title"),)),),
            (renderers.SlideTextBox(renderers.SlideBox(1, 2, 3, 0.3), (renderers.SlideParagraph("Too much text"),)),),
            (
                renderers.SlideTextBox(renderers.SlideBox(1, 2, 4, 2), (renderers.SlideParagraph("First"),)),
                renderers.SlideTextBox(renderers.SlideBox(2, 2, 4, 2), (renderers.SlideParagraph("Second"),)),
            ),
        )
        for prepared_shapes in shapes:
            deck = renderers.PreparedDeck((renderers.PreparedSlide("Title", prepared_shapes),), 1)
            self.assert_render_error("layout_overflow", renderers.render_prepared_pptx, deck)

    def test_pptx_enforces_total_shape_cell_and_image_counts(self):
        source = prepared_deck()
        for limits in (
            renderers.OfficeRenderLimits(max_shapes=2),
            renderers.OfficeRenderLimits(max_cells=3),
            renderers.OfficeRenderLimits(max_columns=1),
        ):
            self.assert_render_error("limit_exceeded", renderers.render_prepared_pptx, source, limits=limits)
        first = source.slides[0]
        duplicate_image = renderers.SlideImage(renderers.SlideBox(4, 5, 2, 1.5), "figure")
        deck = replace(source, slides=(replace(first, shapes=(*first.shapes, duplicate_image)),), expected_slide_count=1)
        self.assert_render_error(
            "limit_exceeded", renderers.render_prepared_pptx, deck,
            limits=renderers.OfficeRenderLimits(max_images=1),
        )

    def test_identical_prepared_content_has_deterministic_bytes_and_metadata(self):
        samples = (
            (renderers.render_prepared_xlsx, small_workbook()),
            (renderers.render_prepared_docx, rich_report()),
            (renderers.render_prepared_pdf, rich_report()),
            (renderers.render_prepared_pptx, prepared_deck()),
        )
        for renderer, source in samples:
            results = []
            for clock in ((2020, 1, 2, 3, 4, 6, 3, 2, -1), (2026, 8, 9, 10, 11, 12, 6, 221, -1)):
                with patch("zipfile.time.localtime", return_value=clock), renderer(source) as output:
                    payload = output.file_content.read()
                    results.append((payload, output.metadata, output.content_sha256))
            with self.subTest(renderer=renderer.__name__):
                self.assertEqual(results[0][1:], results[1][1:])
                self.assertEqual(results[0][0], results[1][0])

    def test_all_renderers_have_bounded_output_input_and_owned_stream_metadata(self):
        samples = (
            (renderers.render_prepared_xlsx, small_workbook()),
            (renderers.render_prepared_docx, renderers.PreparedReport("Prepared report")),
            (renderers.render_prepared_pdf, renderers.PreparedReport("Prepared report")),
            (renderers.render_prepared_pptx, prepared_deck()),
        )
        for renderer, source in samples:
            with self.subTest(renderer=renderer.__name__):
                with renderer(source) as output:
                    start = output.file_content.tell()
                    payload = output.file_content.read()
                    digest = hashlib.sha256(payload).hexdigest()
                    self.assertEqual(start, 0)
                    self.assertEqual(len(payload), output.size_bytes)
                    self.assertEqual(digest, output.content_sha256)
                    self.assertTrue(output.metadata["complete"])
                    self.assertTrue(output.media_type)
                self.assertTrue(output.file_content.closed)
                self.assert_render_error(
                    "limit_exceeded", renderer, source,
                    limits=renderers.OfficeRenderLimits(max_output_bytes=100),
                )
                self.assert_render_error(
                    "limit_exceeded", renderer, source,
                    limits=renderers.OfficeRenderLimits(max_input_bytes=1),
                )

    def test_cancellation_and_denied_access_prevent_source_consumption(self):
        observed = []

        class Rows:
            def __iter__(self):
                observed.append("consumed")
                yield ("secret",)

        source = renderers.PreparedWorkbook((renderers.PreparedSheet("Data", ("value",), Rows(), 1),))
        self.assert_render_error(
            "cancelled", renderers.render_prepared_xlsx, source,
            checks=renderers.OfficeRenderChecks(is_cancelled=lambda: True),
        )
        self.assert_render_error(
            "access_denied", renderers.render_prepared_xlsx, source,
            checks=renderers.OfficeRenderChecks(access_check=lambda: False),
        )
        self.assertEqual(observed, [])

    def test_mid_stream_cancel_and_revocation_close_iterators_and_output(self):
        for failure in ("cancelled", "access_denied", "source_changed"):
            state = {"revoked": False, "closed": False}
            streams = []
            original = renderers._RenderContext.buffer

            def tracked_buffer(context):
                stream = original(context)
                streams.append(stream)
                return stream

            def rows():
                try:
                    yield ("first",)
                    state["revoked"] = True
                    yield ("second",)
                finally:
                    state["closed"] = True

            checks = {
                "cancelled": renderers.OfficeRenderChecks(is_cancelled=lambda: state["revoked"]),
                "access_denied": renderers.OfficeRenderChecks(access_check=lambda: not state["revoked"]),
                "source_changed": renderers.OfficeRenderChecks(source_check=lambda: not state["revoked"]),
            }[failure]
            source = renderers.PreparedWorkbook((renderers.PreparedSheet("Data", ("value",), rows(), 2),))
            with patch.object(renderers._RenderContext, "buffer", tracked_buffer):
                self.assert_render_error(
                    failure, renderers.render_prepared_xlsx, source,
                    limits=renderers.OfficeRenderLimits(check_interval=1), checks=checks,
                )
            self.assertTrue(state["closed"])
            self.assertTrue(streams)
            self.assertTrue(all(stream.closed for stream in streams))

    def test_after_serialization_revocation_discards_output(self):
        state = {"saved": False}
        streams = []
        original_save = DocumentType.save
        original_buffer = renderers._RenderContext.buffer

        def save(document, stream):
            original_save(document, stream)
            state["saved"] = True

        def tracked_buffer(context):
            stream = original_buffer(context)
            streams.append(stream)
            return stream

        with patch.object(DocumentType, "save", save), patch.object(renderers._RenderContext, "buffer", tracked_buffer):
            self.assert_render_error(
                "access_denied", renderers.render_prepared_docx, renderers.PreparedReport("Prepared text"),
                checks=renderers.OfficeRenderChecks(access_check=lambda: not state["saved"]),
            )
        self.assertTrue(state["saved"])
        self.assertTrue(streams)
        self.assertTrue(all(stream.closed for stream in streams))

    def test_iterator_close_failure_preserves_primary_error_and_closes_resources(self):
        original_buffer = renderers._RenderContext.buffer
        for failure in ("cancelled", "access_denied", "source_changed", None):
            with self.subTest(failure=failure):
                state = {"started": False}
                streams = []

                class Rows:
                    def __init__(self):
                        self.reads = 0
                        self.close_calls = 0

                    def __iter__(self):
                        return self

                    def __next__(self):
                        if self.reads:
                            raise StopIteration
                        self.reads += 1
                        state["started"] = True
                        return ("row",)

                    def close(self):
                        self.close_calls += 1
                        raise OSError("private cleanup failure")

                def tracked_buffer(context):
                    stream = original_buffer(context)
                    streams.append(stream)
                    return stream

                rows = Rows()
                source = renderers.PreparedWorkbook((renderers.PreparedSheet("Data", ("value",), rows, 1),))
                checks = renderers.OfficeRenderChecks(
                    is_cancelled=lambda: state["started"] and failure == "cancelled",
                    access_check=lambda: not (state["started"] and failure == "access_denied"),
                    source_check=lambda: not (state["started"] and failure == "source_changed"),
                )
                with patch.object(renderers._RenderContext, "buffer", tracked_buffer):
                    error = self.assert_render_error(
                        failure or "render_io", renderers.render_prepared_xlsx, source,
                        checks=checks, limits=renderers.OfficeRenderLimits(check_interval=1),
                    )
                self.assertEqual(rows.close_calls, 1)
                self.assertTrue(streams)
                self.assertTrue(all(stream.closed for stream in streams))
                self.assertNotIn("private cleanup", str(error))
                notes = getattr(error, "__notes__", [])
                self.assertEqual(len(notes), 1 if failure else 0)
                self.assertNotIn("private cleanup", "\n".join(notes))

    def test_shared_iterator_cleanup_preserves_primary_exception_identity_and_cause(self):
        for code in ("cancelled", "access_denied", "source_changed", "source_unavailable"):
            with self.subTest(code=code):
                primary_error = renderers.OfficeRenderError(code)
                cause = OSError("private source failure")

                class Rows:
                    def __init__(self):
                        self.close_calls = 0

                    def __iter__(self):
                        return self

                    def __next__(self):
                        raise primary_error from cause

                    def close(self):
                        self.close_calls += 1
                        raise OSError("private cleanup failure")

                rows = Rows()
                source = renderers.PreparedWorkbook((renderers.PreparedSheet("Data", ("value",), rows, 1),))
                error = self.assert_render_error(code, renderers.render_prepared_xlsx, source)
                self.assertIs(error, primary_error)
                self.assertIs(error.__cause__, cause)
                self.assertIsNone(error.__context__)
                self.assertEqual(rows.close_calls, 1)
                notes = getattr(error, "__notes__", [])
                self.assertEqual(len(notes), 1)
                self.assertIn("cleanup", notes[0].lower())
                self.assertNotIn("private", "\n".join(notes))

    def test_iterator_cleanup_failure_is_not_hidden_by_an_outer_handled_exception(self):
        outer_error = RuntimeError("already handled failure")
        cleanup_error = OSError("private cleanup failure")

        class Rows:
            def __init__(self):
                self.values = iter((("row",),))
                self.close_calls = 0

            def __iter__(self):
                return self

            def __next__(self):
                return next(self.values)

            def close(self):
                self.close_calls += 1
                raise cleanup_error

        rows = Rows()
        source = renderers.PreparedWorkbook((renderers.PreparedSheet("Data", ("value",), rows, 1),))
        try:
            raise outer_error
        except RuntimeError:
            error = self.assert_render_error("render_io", renderers.render_prepared_xlsx, source)
        self.assertIs(error.__cause__, cleanup_error)
        self.assertEqual(rows.close_calls, 1)
        self.assertEqual(getattr(outer_error, "__notes__", []), [])

    def test_explicit_image_resolver_is_checked_and_never_silently_substituted(self):
        resolved = []
        image = image_bytes()

        def resolver(source):
            resolved.append(source)
            return image

        source = renderers.PreparedReport("![One](authorized-id)\n\n![Two](authorized-id)")
        with renderers.render_prepared_docx(source, image_resolver=resolver) as output:
            doc = Document(output.file_content)
            picture_count = len(doc.inline_shapes)
        self.assertEqual(resolved, ["authorized-id"])
        self.assertEqual(picture_count, 2)
        state = {"revoked": False}

        def revoked_resolver(source):
            state["revoked"] = True
            return image

        self.assert_render_error(
            "access_denied", renderers.render_prepared_pdf, source,
            image_resolver=revoked_resolver,
            checks=renderers.OfficeRenderChecks(access_check=lambda: not state["revoked"]),
        )

    def test_failure_classes_and_cleanup_are_safe_for_later_retries(self):
        def unavailable():
            raise TimeoutError("private provider address and credential")

        failure = self.assert_render_error(
            "source_unavailable", renderers.render_prepared_docx, renderers.PreparedReport("Prepared"),
            checks=renderers.OfficeRenderChecks(source_check=unavailable),
        )
        self.assertNotIn("private provider", str(failure))
        closed = []

        def unavailable_rows():
            try:
                yield (1,)
                raise TimeoutError("private source address")
            finally:
                closed.append(True)

        unavailable_workbook = renderers.PreparedWorkbook((
            renderers.PreparedSheet("Data", ("id",), unavailable_rows(), 2),
        ))
        self.assert_render_error("source_unavailable", renderers.render_prepared_xlsx, unavailable_workbook)
        self.assertEqual(closed, [True])
        streams = []
        original = renderers._RenderContext.buffer

        def tracked_buffer(context):
            stream = original(context)
            streams.append(stream)
            return stream

        def broken_save(document, stream):
            stream.write(b"partial")
            raise OSError("private disk detail")

        with patch.object(renderers._RenderContext, "buffer", tracked_buffer), patch.object(DocumentType, "save", broken_save):
            failure = self.assert_render_error(
                "render_io", renderers.render_prepared_docx, renderers.PreparedReport("Prepared"),
            )
        self.assertNotIn("private disk", str(failure))
        self.assertTrue(all(stream.closed for stream in streams))

    def test_output_limit_failures_leave_only_closed_bounded_memory_streams(self):
        streams = []
        sizes_at_close = []
        original_buffer = renderers._RenderContext.buffer
        original_close = renderers._BoundedBuffer.close

        def tracked_buffer(context):
            stream = original_buffer(context)
            streams.append(stream)
            return stream

        def tracked_close(stream):
            if not stream.closed:
                size = len(stream.getvalue())
                sizes_at_close.append(size)
            original_close(stream)

        with patch.object(renderers._RenderContext, "buffer", tracked_buffer), patch.object(
            renderers._BoundedBuffer, "close", tracked_close,
        ):
            for renderer, source in (
                (renderers.render_prepared_xlsx, small_workbook()),
                (renderers.render_prepared_docx, rich_report()),
                (renderers.render_prepared_pdf, rich_report()),
                (renderers.render_prepared_pptx, prepared_deck()),
            ):
                self.assert_render_error(
                    "limit_exceeded", renderer, source,
                    limits=renderers.OfficeRenderLimits(max_output_bytes=100),
                )
        self.assertGreaterEqual(len(streams), 4)
        self.assertTrue(all(stream.closed for stream in streams))
        self.assertTrue(all(size <= 100 for size in sizes_at_close))

    def test_renderers_never_use_temp_files_or_network(self):
        def forbidden(*args, **kwargs):
            raise AssertionError("The pure renderer attempted external I/O")

        with ExitStack() as stack:
            for target in ("TemporaryFile", "NamedTemporaryFile", "mkstemp"):
                stack.enter_context(patch.object(tempfile, target, forbidden))
            stack.enter_context(patch("openpyxl.worksheet._writer.create_temporary_file", forbidden))
            stack.enter_context(patch.object(socket.socket, "connect", forbidden))
            stack.enter_context(patch.object(socket, "create_connection", forbidden))
            for renderer, source in (
                (renderers.render_prepared_xlsx, small_workbook()),
                (renderers.render_prepared_docx, rich_report()),
                (renderers.render_prepared_pdf, rich_report()),
                (renderers.render_prepared_pptx, prepared_deck()),
            ):
                with renderer(source) as output:
                    self.assertGreater(output.size_bytes, 0)

    def test_real_cold_import_and_render_without_bootstrap_llm_or_uploads_normal_and_optimized(self):
        script = f"""
import importlib.abc
import io
import sys
import tempfile
sys.path.insert(0, {str(APP_PATH)!r})
forbidden = {{
    'config', 'flask', 'azure', 'openai', 'semantic_kernel', 'requests',
    'functions_settings', 'functions_simplechat_operations', 'functions_artifact_publication',
    'functions_generated_file_exports', 'route_backend_conversation_export',
}}
class DenyBootstrap(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname.split('.')[0] in forbidden:
            raise RuntimeError('Forbidden dependency: ' + fullname)
sys.meta_path.insert(0, DenyBootstrap())
def audit(event, args):
    if event in {{'socket.connect', 'socket.getaddrinfo', 'urllib.Request'}}:
        raise RuntimeError('Network access forbidden')
sys.addaudithook(audit)
def no_temp(*args, **kwargs):
    raise RuntimeError('Temporary file forbidden')
tempfile.TemporaryFile = tempfile.NamedTemporaryFile = tempfile.mkstemp = no_temp
import functions_office_file_renderers as r
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
import fitz
samples = (
    (r.render_prepared_xlsx, r.PreparedWorkbook((r.PreparedSheet('Data', ('id',), [(1,), (2,)], 2),))),
    (r.render_prepared_docx, r.PreparedReport('# Complete\\n\\nPrepared content.')),
    (r.render_prepared_pdf, r.PreparedReport('# Complete\\n\\nPrepared content.')),
    (r.render_prepared_pptx, r.PreparedDeck((r.PreparedSlide('Prepared title', notes='Prepared notes'),), 1)),
)
for renderer, source in samples:
    with renderer(source) as result:
        if not result.size_bytes or not result.metadata['complete']:
            raise RuntimeError('Empty or incomplete output')
        if result.output_format == 'xlsx':
            book = load_workbook(result.file_content, read_only=True)
            rows = list(book.active.values)
            book.close()
            if rows != [('id',), (1,), (2,)]:
                raise RuntimeError('Incomplete workbook')
        elif result.output_format == 'docx':
            doc = Document(result.file_content)
            text = '\\n'.join(p.text for p in doc.paragraphs)
            if 'Prepared content.' not in text:
                raise RuntimeError('Incomplete document')
        elif result.output_format == 'pdf':
            payload = result.file_content.read()
            with fitz.open(stream=payload, filetype='pdf') as doc:
                text = ''.join(p.get_text() for p in doc)
            if 'Prepared content.' not in text:
                raise RuntimeError('Incomplete PDF')
        else:
            deck = Presentation(result.file_content)
            if len(deck.slides) != 1 or deck.slides[0].notes_slide.notes_text_frame.text != 'Prepared notes':
                raise RuntimeError('Incomplete deck')
    if not result.file_content.closed:
        raise RuntimeError('Leaked stream')
if forbidden.intersection(sys.modules):
    raise RuntimeError('Bootstrap or provider imported')
"""
        for options in ([], ["-O"]):
            with self.subTest(options=options):
                outcome = subprocess.run(
                    [sys.executable, *options, "-c", script], cwd=REPO_ROOT,
                    capture_output=True, text=True, encoding="utf-8", timeout=60,
                    env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1", "PYTHONUTF8": "1"},
                )
                self.assertEqual(outcome.returncode, 0, outcome.stdout + outcome.stderr)


if __name__ == "__main__":
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(OfficeRendererTests)
    results = unittest.TextTestRunner(verbosity=2).run(suite)
    if OfficeRendererTests.large_workbook_peak_bytes:
        peak_mib = OfficeRendererTests.large_workbook_peak_bytes / (1024 * 1024)
        print(f"30,000-row renderer peak tracked allocation: {peak_mib:.2f} MiB")
    sys.exit(0 if results.wasSuccessful() else 1)
