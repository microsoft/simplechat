#!/usr/bin/env python3
# test_file_merge_scratch_dir.py
"""
Functional test for where V2 file merges spill their spooled data to disk.
Version: 0.261.244
Implemented in: 0.261.243

This test ensures that a merge too large to keep in memory spills its spooled rows or
assembled file into the container's dedicated scratch directory (/sc-temp-files) when the
process can write to it, and into the platform temp directory otherwise, never the working
directory, and that spilling doesn't change the merged result. It covers spreadsheet
merges and PDF, Word, PowerPoint and workbook merges.
"""

import io
import os
import sys
import tempfile

import pytest

sys.path.append(os.path.dirname(os.path.abspath(__file__)))
sys.path.append(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "application", "single_app"))

from test_support.versioning import assert_app_version_at_least

import functions_document_merge
import functions_document_merge_core
import functions_tabular_merge
import functions_temp_files

DocumentMergePart = functions_document_merge_core.DocumentMergePart
merge_documents = functions_document_merge.merge_documents
TabularMergeSource = functions_tabular_merge.TabularMergeSource
merge_tabular_sources = functions_tabular_merge.merge_tabular_sources


def pdf_bytes(pages):
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def docx_bytes(label):
    from docx import Document

    document = Document()
    document.add_heading(f"{label} title", level=1)
    document.add_paragraph(f"{label} body")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pptx_bytes(label):
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = f"{label} slide"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def xlsx_bytes(region, amount):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Region", "Amount"])
    sheet.append([region, amount])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


SPREADSHEETS = [("east.csv", b"Region,Amount\r\nEast,10\r\n"), ("west.xlsx", xlsx_bytes("West", 5))]
DOCUMENTS = {
    "pdf": [("a.pdf", pdf_bytes(2)), ("b.pdf", pdf_bytes(1))],
    "docx": [("a.docx", docx_bytes("A")), ("b.docx", docx_bytes("B"))],
    "pptx": [("a.pptx", pptx_bytes("A")), ("b.pptx", pptx_bytes("B"))],
    "workbook": SPREADSHEETS,
}
KINDS = ["spreadsheet", *DOCUMENTS]


def merge(kind):
    """Merge the kind's files and return the result as plain values."""
    if kind == "spreadsheet":
        sources = [
            TabularMergeSource(source_id=name, file_name=name, load_bytes=lambda content=content: content)
            for name, content in SPREADSHEETS
        ]
        with merge_tabular_sources(sources) as result:
            return list(result.columns), [list(row) for row in result.iter_rows()]
    parts = [DocumentMergePart(name, name, lambda content=content: content) for name, content in DOCUMENTS[kind]]
    with merge_documents(kind, parts) as result:
        return result.read_bytes()


def spill_every_spool(monkeypatch, scratch_directory):
    """Make every merge spool spill to disk on its first write and record each file's directory."""
    monkeypatch.setattr(functions_temp_files, "SC_TEMP_FILES_DIR", str(scratch_directory))
    monkeypatch.setattr(functions_document_merge_core, "_SPOOL_MEMORY_BYTES", 1)
    monkeypatch.setattr(functions_tabular_merge, "_SPOOL_MEMORY_BYTES", 1)
    real = tempfile.TemporaryFile
    directories = []

    def recording(*args, **kwargs):
        directories.append(kwargs.get("dir", args[6] if len(args) > 6 else None))
        return real(*args, **kwargs)

    # A spooled temporary file creates its file on disk through tempfile.TemporaryFile.
    monkeypatch.setattr(tempfile, "TemporaryFile", recording)
    return directories


def test_version_includes_merge_scratch_directory():
    assert_app_version_at_least("0.261.243")


@pytest.mark.parametrize("kind", KINDS)
def test_merges_spill_into_the_scratch_directory(tmp_path, monkeypatch, kind):
    print(f"🔍 Testing a {kind} merge spills into the scratch directory...")
    in_memory = merge(kind)
    scratch = tmp_path / "sc-temp-files"
    scratch.mkdir()
    directories = spill_every_spool(monkeypatch, scratch)
    spilled = merge(kind)
    assert directories, "The merge never spilled to disk."
    assert set(directories) == {str(scratch)}
    assert spilled == in_memory
    print(f"✅ {len(directories)} spool(s) spilled into the scratch directory with the same result.")


@pytest.mark.parametrize("kind", KINDS)
def test_merges_spill_into_the_platform_temp_directory_without_one(tmp_path, monkeypatch, kind):
    print(f"🔍 Testing a {kind} merge spills into the platform temp directory without a scratch directory...")
    in_memory = merge(kind)
    directories = spill_every_spool(monkeypatch, tmp_path / "missing")
    spilled = merge(kind)
    assert directories, "The merge never spilled to disk."
    assert set(directories) == {None}
    assert spilled == in_memory
    print(f"✅ {len(directories)} spool(s) used the platform temp directory with the same result.")


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
