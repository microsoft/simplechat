# test_m365_provider_core.py
#!/usr/bin/env python3
"""
Functional tests for source-bounded Microsoft 365 providers and file evidence.
Version: 0.261.038
Implemented in: 0.261.029

External Graph, authentication, logging, and storage I/O is mocked. The tests
execute the real provider, transport, extraction, metadata, and plugin modules.
"""

import importlib
import importlib.util
import json
import re
import sys
import types
import zipfile
from contextvars import copy_context
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import Mock
from urllib.parse import urlsplit

import fitz
import openpyxl
import pytest
import requests
import jsonschema
from flask import Flask, session
from opentelemetry.instrumentation.utils import is_http_instrumentation_enabled
from pptx import Presentation
from pptx.util import Inches


APP_DIR = Path(__file__).resolve().parents[1] / "application" / "single_app"
sys.path.insert(0, str(APP_DIR))

# Test-path setup must precede imports of the actual application modules.
import functions_m365_extraction as extraction  # noqa: E402
import functions_m365_operations as operations  # noqa: E402
import functions_m365_retrieval as retrieval  # noqa: E402
import functions_m365_transport as transport_module  # noqa: E402
from functions_m365_transport import M365CloudConfig, M365ProviderError, M365Transport  # noqa: E402
from conversation_memory_storage import BlobRecord, MemoryConflictError, MemoryNotFoundError  # noqa: E402
from functions_conversation_memory import ConversationMemoryStore, MemoryContext, PublicationGrant  # noqa: E402
from functions_m365_approvals import M365ApprovalRequired  # noqa: E402


REAL_AUTHORIZE_SOURCE = transport_module.authorize_m365_source
REAL_AUTHORIZE_CAPABILITY = transport_module.authorize_m365_capability


class FakeResponse:
    def __init__(self, payload=None, status=200, *, body=None, headers=None, chunks=None):
        self.status_code = status
        self.body = json.dumps(payload).encode() if body is None else body
        self.headers = {"Content-Type": "application/json", **(headers or {})}
        self.chunks = chunks
        self.closed = False

    def iter_content(self, chunk_size):
        if self.chunks is not None:
            yield from self.chunks
        else:
            for index in range(0, len(self.body), chunk_size):
                yield self.body[index:index + chunk_size]

    def close(self):
        self.closed = True


@pytest.fixture
def execution(monkeypatch):
    context = types.SimpleNamespace(
        actor_user_id="user-1", data_user_id="user-1", tenant_id="tenant-1",
        request_id="request-1", conversation_id="conversation-1",
        shared=False, workflow_id=None, run_id=None,
        group_id=None,
    )
    monkeypatch.setattr(transport_module, "get_m365_context", lambda **kwargs: context)
    monkeypatch.setattr(retrieval, "get_m365_context", lambda **kwargs: context)
    monkeypatch.setattr(
        transport_module, "authorize_m365_source",
        lambda source, action_id, policy, **kwargs: (context, {"source": source, "sharing_required": False}),
    )
    monkeypatch.setattr(retrieval, "authorize_m365_source", transport_module.authorize_m365_source)
    monkeypatch.setattr(retrieval, "authorize_m365_publication", transport_module.authorize_m365_source)
    monkeypatch.setattr(
        transport_module, "authorize_m365_capability",
        lambda action_id, operation_name, action_type, **kwargs: context,
    )
    monkeypatch.setattr(retrieval, "authorize_m365_capability", transport_module.authorize_m365_capability)
    monkeypatch.setattr(retrieval, "log_m365_failure", Mock())
    monkeypatch.setattr(transport_module, "log_m365_failure", Mock())
    monkeypatch.setattr(
        requests.sessions.Session, "request",
        Mock(side_effect=AssertionError("Unexpected real network access in provider tests.")),
    )
    return context


class GraphFixture:
    def __init__(self, source="onedrive", licensed=False):
        self.source = source
        self.licensed = licensed
        self.calls = []
        self.token_requests = []
        self.copilot_response = None
        self.extra_hits = []
        self.denied_items = set()
        self.more = False
        self.web_url = (
            "https://tenant-my.sharepoint.com/personal/user/Documents/report.txt"
            if source == "onedrive"
            else "https://tenant.sharepoint.com/sites/team/Documents/report.txt"
        )

    def item(self, item_id="item-1"):
        return {
            "id": item_id, "name": "report.txt", "size": 12, "file": {"mimeType": "text/plain"},
            "parentReference": {"driveId": "drive-1", "siteId": "site-1"},
            "webUrl": self.web_url, "eTag": '"version-1"', "cTag": '"content-1"',
            "lastModifiedDateTime": "2026-09-17T12:00:00Z",
            "@microsoft.graph.downloadUrl": "https://tenant.sharepoint.com/secret?token=do-not-return",
        }

    def hit(self, item_id="item-1"):
        return {
            "resource": self.item(item_id), "summary": "The <c0>report</c0> says &lt;safe&gt;.",
        }

    def request(self, method, url, **kwargs):
        self.calls.append((method, url, kwargs))
        path = urlsplit(url).path
        if path.endswith("/me"):
            return FakeResponse({
                "id": "user-1",
                "assignedLicenses": [{
                    "skuId": retrieval.M365_COPILOT_SKU_ID, "disabledPlans": [],
                }] if self.licensed else [],
                "assignedPlans": [{
                    "servicePlanId": retrieval.M365_COPILOT_SEARCH_PLAN_ID,
                    "capabilityStatus": "Enabled",
                }] if self.licensed else [],
            })
        if path.endswith("/copilot/retrieval"):
            if self.copilot_response is not None:
                return self.copilot_response
            return FakeResponse({
                "retrievalHits": [{
                    "webUrl": self.web_url, "resourceType": "listItem",
                    "extracts": [{"text": "The exact retained Copilot excerpt."}],
                    "sensitivityLabel": {"sensitivityLabelId": "label-1", "displayName": "Internal"},
                }],
            })
        if path.endswith("/search/query"):
            hits = [self.hit(), *self.extra_hits]
            return FakeResponse({
                "value": [{
                    "hitsContainers": [{
                        "total": len(hits), "moreResultsAvailable": self.more, "hits": hits,
                    }],
                }],
            })
        if "/items/" in path:
            item_id = path.rsplit("/", 1)[1]
            if item_id in self.denied_items:
                return FakeResponse({"error": {"code": "accessDenied", "message": "secret provider detail"}}, status=403)
            return FakeResponse(self.item(item_id))
        if path.endswith("/drives/drive-1"):
            return FakeResponse({
                "id": "drive-1",
                "driveType": "business" if self.source == "onedrive" else "documentLibrary",
                "webUrl": self.web_url.rsplit("/", 1)[0],
            })
        raise AssertionError(f"Unexpected mock Graph operation: {method} {path}")

    def transport(self, source=None, cloud=None):
        def token_provider(scopes, context):
            self.token_requests.append((scopes, context))
            return {"access_token": "unit-test-token"}

        return M365Transport(
            source or self.source, "action-1",
            cloud=cloud or M365CloudConfig("https://graph.microsoft.com/v1.0", "https://login.microsoftonline.com/tenant-1"),
            request=self.request, token_provider=token_provider,
        )


def test_metadata_is_authoritative_and_write_defaults_are_safe():
    expected = {
        "m365_calendar": "calendar", "m365_email": "email",
        "m365_onedrive": "onedrive", "m365_sharepoint": "spo",
    }
    assert set(operations.M365_ACTION_TYPES) == set(expected)
    for action_type, source in expected.items():
        definition = operations.get_m365_action_definition(action_type)
        assert definition["source"] == source
        defaults = operations.get_m365_default_capabilities(action_type)
        assert not any(defaults.get(name) for name in operations.M365_WRITE_FUNCTIONS)
        normalized = operations.normalize_m365_action_config(action_type, {
            "endpoint": "https://attacker.invalid",
            "scopes": {"search_files": ["https://attacker.invalid/.default"]},
            "m365_capabilities": {name: True for name in ("search_files", "get_my_messages", "get_my_events")},
            "enabled_functions": ["search_files", "get_my_messages", "get_my_events"],
        })
        allowed = {item["function_name"] for item in definition["capabilities"]}
        assert set(normalized["enabled_functions"]) <= allowed
        assert "endpoint" not in normalized and "scopes" not in normalized
        definition["capabilities"].clear()
        assert operations.get_m365_action_definition(action_type)["capabilities"]


def test_empty_or_cross_source_overrides_cannot_reenable_capabilities():
    empty = operations.get_m365_enabled_function_names("m365_email", {}, enabled_functions=[])
    narrowed = operations.get_m365_enabled_function_names(
        "m365_calendar", {"create_calendar_invite": False},
        agent_capabilities={"create_calendar_invite": True, "send_mail": True},
    )
    assert empty == []
    assert "create_calendar_invite" not in narrowed and "send_mail" not in narrowed
    with pytest.raises(ValueError):
        operations.normalize_m365_action_config("m365_onedrive", {"additionalFields": {"allowed_sites": ["anything"]}})


@pytest.mark.parametrize("base,host", [
    ("https://graph.microsoft.com/v1.0", "graph.microsoft.com"),
    ("https://graph.microsoft.us/v1.0", "graph.microsoft.us"),
    ("https://dod-graph.microsoft.us/v1.0", "dod-graph.microsoft.us"),
    ("https://graph.example.test/api/v1.0", "graph.example.test"),
])
def test_transport_binds_delegated_scopes_and_requests_to_cloud(execution, base, host):
    seen = []
    tokens = []
    cloud = M365CloudConfig(base, "https://identity.example.test/tenant")
    client = M365Transport(
        "onedrive", "action", cloud=cloud,
        request=lambda method, url, **kwargs: seen.append((url, kwargs)) or FakeResponse({"id": "file-1"}),
        token_provider=lambda scopes, context: tokens.append((scopes, context)) or {"access_token": "unit-test-token"},
    )
    result = client.request_json("GET", "/v1.0/me", ["Files.Read.All"])
    assert result["id"] == "file-1"
    assert urlsplit(seen[0][0]).hostname == host
    assert tokens[0][0] == [f"{cloud.resource_url}/Files.Read.All"]
    assert tokens[0][1] is execution
    assert seen[0][1]["allow_redirects"] is False
    with pytest.raises(M365ProviderError):
        client.qualify_scopes(["https://attacker.invalid/Files.Read.All"])


def test_foreign_next_link_is_rejected_before_following(execution):
    response = FakeResponse({"value": [], "@odata.nextLink": "https://attacker.invalid/v1.0/me"})
    request = Mock(return_value=response)
    client = M365Transport(
        "onedrive", cloud=M365CloudConfig("https://graph.microsoft.us/v1.0", "https://login.microsoftonline.us/tenant"),
        request=request, token_provider=lambda scopes, context: {"access_token": "unit-test-token"},
    )
    with pytest.raises(M365ProviderError) as caught:
        client.request_json("GET", "/me/drive", ["Files.Read.All"])
    assert caught.value.code == "untrusted_graph_url"
    assert request.call_count == 1 and response.closed


@pytest.mark.parametrize("status,code", [(401, "authentication_required"), (403, "access_denied"), (429, "throttled")])
def test_transport_errors_are_safe_and_do_not_retry(execution, status, code):
    request = Mock(return_value=FakeResponse(
        {"error": {"code": "providerDetail", "message": "private-token=secret"}},
        status=status, headers={"Retry-After": "41"},
    ))
    client = M365Transport(
        "spo", cloud=M365CloudConfig("https://graph.microsoft.com/v1.0", "https://login.microsoftonline.com/tenant"),
        request=request, token_provider=lambda scopes, context: {"access_token": "unit-test-token"},
    )
    with pytest.raises(M365ProviderError) as caught:
        client.request_json("POST", "/search/query", ["Files.Read.All"])
    assert caught.value.code == code
    assert caught.value.retry_after_seconds == 41
    assert "private-token" not in json.dumps(caught.value.as_dict())
    assert request.call_count == 1
    if status == 401:
        assert caught.value.details["scopes"] == ["https://graph.microsoft.com/Files.Read.All"]
        assert caught.value.details["sources"] == ["spo"]
    else:
        assert "scopes" not in caught.value.details


@pytest.mark.parametrize("status", [401, 403, 429, 503])
def test_only_graph_authentication_rejection_marks_the_current_chat_for_reconnect(execution, status):
    from functions_m365_connections import CHAT_RECONNECT_SESSION_KEY

    request = Mock(return_value=FakeResponse({"error": {"code": "rejected"}}, status=status))
    client = M365Transport(
        "spo", cloud=M365CloudConfig("https://graph.microsoft.com/v1.0", "https://login.microsoftonline.com/tenant"),
        request=request, token_provider=lambda scopes, context: {"access_token": "unexpired-but-rejected-token"},
    )
    app = Flask(__name__)
    app.secret_key = "test-only"
    with app.test_request_context():
        session["user"] = {"oid": execution.data_user_id, "tid": execution.tenant_id}
        session["token_cache"] = "existing-cache"
        with pytest.raises(M365ProviderError):
            client.request_json("POST", "/search/query", ["Files.Read.All", "Sites.Read.All"])
        marker = session.get(CHAT_RECONNECT_SESSION_KEY)
        retained_cache = session["token_cache"]
    assert bool(marker) is (status == 401)
    assert retained_cache == "existing-cache"
    assert request.call_count == 1


def test_expired_preauthenticated_download_link_does_not_mark_graph_login_invalid(execution):
    from functions_m365_connections import CHAT_RECONNECT_SESSION_KEY

    first = FakeResponse(status=302, headers={"Location": "https://tenant.sharepoint.com/download?temporary=opaque"})
    expired = FakeResponse(status=401, body=b"Expired download")
    client = M365Transport(
        "spo", cloud=M365CloudConfig("https://graph.microsoft.com/v1.0", "https://login.microsoftonline.com/tenant"),
        request=Mock(side_effect=[first, expired]), token_provider=lambda scopes, context: {"access_token": "valid-token"},
    )
    app = Flask(__name__)
    app.secret_key = "test-only"
    with app.test_request_context():
        session["user"] = {"oid": execution.data_user_id, "tid": execution.tenant_id}
        with pytest.raises(M365ProviderError) as raised:
            with client.download_file("drive", "item", suffix=".txt", allowed_mime_types=["text/plain"], max_bytes=100):
                pytest.fail("An expired download cannot yield a file.")
        marker = session.get(CHAT_RECONNECT_SESSION_KEY)
    assert raised.value.code == "download_link_expired"
    assert marker is None
    assert first.closed and expired.closed


def test_graph_file_results_have_canonical_source_and_never_download_url(execution):
    fixture = GraphFixture()
    result = retrieval.M365FileProvider(fixture.transport()).search("project report")
    file_result = result["results"][0]
    assert result["provider"] == "graph"
    assert result["fallback_reason"] == "copilot_license_not_assigned"
    assert file_result["canonical_id"]["drive_id"] == "drive-1"
    assert file_result["canonical_id"]["item_id"] == "item-1"
    assert file_result["web_url"] == fixture.web_url
    assert file_result["captured_version"]["observed_metadata"]["etag"] == '"version-1"'
    assert file_result["captured_version"]["source_version_verified"] is False
    assert file_result["excerpts"][0]["text"] == "The report says <safe>."
    assert file_result["coverage"]["complete"] is False
    assert "downloadUrl" not in json.dumps(result) and "do-not-return" not in json.dumps(result)
    assert not any(url.endswith("/copilot/retrieval") for _, url, _ in fixture.calls)


def test_source_classification_uses_drive_type_not_url_or_supplied_source(execution):
    fixture = GraphFixture(source="spo")
    fixture.web_url = "https://tenant-my.sharepoint.com/personal/user/Documents/report.txt"
    result = retrieval.M365FileProvider(fixture.transport(source="onedrive")).search("report")
    assert result["results"] == []
    assert result["coverage"]["excluded_other_source"] == 1


def test_unlicensed_unknown_disabled_and_wrong_plan_never_establish_eligibility():
    for profile in (
        {}, {"assignedLicenses": [], "assignedPlans": []},
        {"assignedLicenses": [{"skuId": retrieval.M365_COPILOT_SKU_ID, "disabledPlans": [retrieval.M365_COPILOT_SEARCH_PLAN_ID]}], "assignedPlans": []},
        {"assignedLicenses": [{"skuId": retrieval.M365_COPILOT_SKU_ID}], "assignedPlans": [{"servicePlanId": "other-plan", "capabilityStatus": "Enabled"}]},
    ):
        eligibility = retrieval.get_m365_license_eligibility(profile)
        assert eligibility["verified"] is False


@pytest.mark.parametrize("base", [
    "https://graph.microsoft.us/v1.0",
    "https://dod-graph.microsoft.us/v1.0",
    "https://graph.custom.example/v1.0",
])
def test_government_and_custom_do_not_probe_retrieval_or_commercial_license(execution, base):
    fixture = GraphFixture(licensed=True)
    cloud = M365CloudConfig(
        base, "https://identity.example.test/tenant", trusted_download_hosts=("sharepoint.com",),
    )
    result = retrieval.M365FileProvider(fixture.transport(cloud=cloud)).search("report")
    assert result["provider"] == "graph"
    assert result["fallback_reason"] == "copilot_retrieval_unsupported_in_cloud"
    assert all(url.startswith(base) for _, url, _ in fixture.calls)
    assert not any(url.endswith("/me") or url.endswith("/copilot/retrieval") for _, url, _ in fixture.calls)


def test_verified_retrieval_uses_v1_both_scopes_and_unordered_file_excerpts(execution):
    fixture = GraphFixture(source="spo", licensed=True)
    result = retrieval.M365FileProvider(fixture.transport()).search("report", top=7)
    request = next(call for call in fixture.calls if call[1].endswith("/copilot/retrieval"))
    assert request[1] == "https://graph.microsoft.com/v1.0/copilot/retrieval"
    assert request[2]["json"]["dataSource"] == "sharePoint"
    assert request[2]["json"]["maximumNumberOfResults"] == 7
    assert result["provider"] == "copilot_retrieval"
    assert result["ranking"] == "unordered_retrieval_hits"
    assert result["results"][0]["source_label"] == "SPO"
    assert result["results"][0]["excerpts"][0]["text"] == "The exact retained Copilot excerpt."
    assert any(
        scopes == ["https://graph.microsoft.com/Files.Read.All", "https://graph.microsoft.com/Sites.Read.All"]
        for scopes, _ in fixture.token_requests
    )
    assert all(context is execution for _, context in fixture.token_requests)


@pytest.mark.parametrize("status,code", [
    (401, "invalidAuthenticationToken"), (403, "accessDenied"),
    (403, "blockedByPolicy"), (429, "tooManyRequests"),
    (404, "itemNotFound"), (500, "internalError"),
])
def test_retrieval_errors_never_fall_through_to_graph(execution, status, code):
    fixture = GraphFixture(licensed=True)
    fixture.copilot_response = FakeResponse({"error": {"code": code}}, status=status)
    with pytest.raises(M365ProviderError):
        retrieval.M365FileProvider(fixture.transport()).search("report")
    assert not any(url.endswith("/search/query") for _, url, _ in fixture.calls)


def test_established_unsupported_api_can_fall_back_but_empty_results_cannot(execution):
    fixture = GraphFixture(licensed=True)
    fixture.copilot_response = FakeResponse({"error": {"code": "NotSupported"}}, status=501)
    fallback = retrieval.M365FileProvider(fixture.transport()).search("report")
    assert fallback["provider"] == "graph" and fallback["fallback_reason"] == "copilot_retrieval_api_unsupported"
    fixture.calls.clear()
    fixture.copilot_response = FakeResponse({"retrievalHits": []})
    empty = retrieval.M365FileProvider(fixture.transport()).search("report")
    assert empty["results"] == [] and empty["provider"] == "copilot_retrieval"
    assert not any(url.endswith("/search/query") for _, url, _ in fixture.calls)


def test_partial_file_access_errors_and_paging_are_visible(execution):
    fixture = GraphFixture()
    fixture.extra_hits = [fixture.hit("denied")]
    fixture.denied_items.add("denied")
    fixture.more = True
    result = retrieval.M365FileProvider(fixture.transport()).search("report", top=5)
    assert result["status"] == "partial" and len(result["results"]) == 1
    assert result["errors"][0]["code"] == "access_denied"
    assert result["next_offset"] == 2 and result["coverage"]["discovery_complete"] is False


def test_natural_query_cannot_inject_kql_or_broaden_folder_results(execution):
    fixture = GraphFixture()
    provider = retrieval.M365FileProvider(fixture.transport())
    result = provider.discover_page(
        'report") OR Path:"https://attacker.invalid" OR *',
        folder_url="https://tenant-my.sharepoint.com/personal/user/Restricted",
    )
    query = next(call[2]["json"]["requests"][0]["query"]["queryString"] for call in fixture.calls if call[1].endswith("/search/query"))
    assert 'AND Path:"https://tenant-my.sharepoint.com/personal/user/Restricted"' in query
    assert "* " not in query and '"OR"' in query
    assert result["results"] == []
    assert result["errors"][0]["code"] == "scope_mismatch"
    with pytest.raises(M365ProviderError):
        provider.resolve_folder('folder/../other')


def test_download_drops_bearer_validates_redirect_and_cleans_local_file(execution, tmp_path, monkeypatch):
    responses = [
        FakeResponse(status=302, headers={"Location": "https://tenant.sharepoint.com/_layouts/15/download.aspx?token=secret"}),
        FakeResponse(body=b"file content", headers={"Content-Type": "text/plain", "Content-Length": "12"}),
    ]
    calls = []
    request = lambda method, url, **kwargs: calls.append((url, kwargs)) or responses.pop(0)
    monkeypatch.setattr(transport_module.tempfile, "tempdir", str(tmp_path))
    client = M365Transport(
        "spo", cloud=M365CloudConfig("https://graph.microsoft.com/v1.0", "https://login.microsoftonline.com/tenant"),
        request=request, token_provider=lambda scopes, context: {"access_token": "unit-test-token"},
    )
    with client.download_file("drive-1", "item-1", suffix=".txt", allowed_mime_types=("text/plain",), max_bytes=100) as downloaded:
        local_path = Path(downloaded.path)
        contents = local_path.read_bytes()
        assert contents == b"file content" and downloaded.size_bytes == 12
    assert not local_path.exists()
    assert calls[0][1]["headers"]["Authorization"] == "Bearer unit-test-token"
    assert "Authorization" not in calls[1][1]["headers"]
    assert isinstance(calls[1][1]["auth"], requests.auth.AuthBase)


@pytest.mark.parametrize("url", [
    "https://tenant.sharepoint.com.attacker.invalid/file",
    "https://graph.microsoft.com@attacker.invalid/file",
    "http://tenant.sharepoint.com/file",
    "https://tenant.sharepoint.com:8443/file",
])
def test_download_rejects_untrusted_redirect_targets(execution, url):
    request = Mock(return_value=FakeResponse(status=302, headers={"Location": url}))
    client = M365Transport(
        "spo", cloud=M365CloudConfig("https://graph.microsoft.com/v1.0", "https://login.microsoftonline.com/tenant"),
        request=request, token_provider=lambda scopes, context: {"access_token": "unit-test-token"},
    )
    with pytest.raises(M365ProviderError):
        with client.download_file("drive-1", "item-1", suffix=".txt", allowed_mime_types=("text/plain",), max_bytes=100):
            pytest.fail("An untrusted download must not be yielded.")
    assert request.call_count == 1


@pytest.mark.parametrize("body,headers,limit,error", [
    (b"x" * 10, {"Content-Type": "text/plain"}, 5, "file_size_limit"),
    (b"x", {"Content-Type": "text/plain", "Content-Length": "10"}, 5, "file_size_limit"),
    (b"x", {"Content-Type": "text/plain", "Content-Length": "3"}, 10, "incomplete_download"),
    (b"x", {"Content-Type": "text/html"}, 10, "unsupported_content_type"),
])
def test_download_size_mime_and_incomplete_responses_never_leave_temporary_files(execution, tmp_path, monkeypatch, body, headers, limit, error):
    monkeypatch.setattr(transport_module.tempfile, "tempdir", str(tmp_path))
    client = M365Transport(
        "spo", cloud=M365CloudConfig("https://graph.microsoft.com/v1.0", "https://login.microsoftonline.com/tenant"),
        request=Mock(return_value=FakeResponse(body=body, headers=headers)),
        token_provider=lambda scopes, context: {"access_token": "unit-test-token"},
    )
    with pytest.raises(M365ProviderError) as caught:
        with client.download_file("drive-1", "item-1", suffix=".txt", allowed_mime_types=("text/plain",), max_bytes=limit):
            pytest.fail("An incomplete download must not be yielded.")
    assert caught.value.code == error
    remaining = list(tmp_path.iterdir())
    assert remaining == []


@pytest.fixture
def real_content_helpers(monkeypatch):
    config = types.ModuleType("config")
    config.re = re
    config.WORD_CHUNK_SIZE = 1000
    modules = {
        "config": config,
        "functions_debug": types.SimpleNamespace(debug_print=Mock()),
        "functions_settings": types.SimpleNamespace(get_settings=lambda: {}),
        "functions_logging": types.ModuleType("functions_logging"),
    }
    for name, module in modules.items():
        monkeypatch.setitem(sys.modules, name, module)
    spec = importlib.util.spec_from_file_location("functions_content", APP_DIR / "functions_content.py")
    module = importlib.util.module_from_spec(spec)
    monkeypatch.setitem(sys.modules, "functions_content", module)
    spec.loader.exec_module(module)
    return module


def test_text_pdf_word_powerpoint_and_all_spreadsheet_sheets_are_extracted(execution, tmp_path, real_content_helpers):
    text_file = tmp_path / "source.txt"
    text_file.write_text("Exact source text\nSecond line", encoding="utf-8")
    pdf_file = tmp_path / "source.pdf"
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_text((72, 72), "PDF evidence on page one.")
        pdf.save(pdf_file)
    word_file = tmp_path / "source.docx"
    with zipfile.ZipFile(word_file, "w") as word:
        word.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Word evidence text.</w:t></w:r></w:p></w:body></w:document>')
    ppt_file = tmp_path / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    textbox.text = "PowerPoint evidence on slide one."
    presentation.save(ppt_file)
    workbook_file = tmp_path / "source.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "First"
    workbook.active.append(["Value", "Formula"])
    workbook.active.append([7, "=2+5"])
    second = workbook.create_sheet("Second")
    second.append(["Other"])
    second.append([11])
    workbook.save(workbook_file)
    workbook.close()
    csv_file = tmp_path / "source.csv"
    csv_file.write_text('name,value\n"Multi\nline",42\n', encoding="utf-8", newline="")

    results = {}
    for path in (text_file, pdf_file, word_file, ppt_file, workbook_file, csv_file):
        results[path.suffix] = extraction.extract_m365_file(str(path), path.name)
    texts = {suffix: "".join(part.text for part in result.parts) for suffix, result in results.items()}
    assert "Exact source text" in texts[".txt"]
    assert "PDF evidence" in texts[".pdf"]
    assert "Word evidence" in texts[".docx"]
    assert "PowerPoint evidence" in texts[".pptx"]
    assert "11" in texts[".xlsx"] and "=2+5" in texts[".xlsx"]
    assert results[".xlsx"].coverage["sheet_names"] == ["First", "Second"]
    assert results[".pdf"].parts[0].location == {"pages": [1]}
    assert results[".pptx"].parts[0].location == {"slides": [1]}
    assert '"Multi\nline",42' in texts[".csv"]


def test_protected_pdf_and_oversized_office_entry_are_explicit_errors(execution, tmp_path, monkeypatch):
    path = tmp_path / "protected.pdf"
    with fitz.open() as document:
        document.new_page().insert_text((72, 72), "Protected content")
        document.save(path, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner-test", user_pw="user-test")
    with pytest.raises(M365ProviderError) as protected:
        extraction.extract_m365_file(str(path), path.name)
    assert protected.value.code == "unsupported_protected_file"
    package = tmp_path / "oversized.docx"
    with zipfile.ZipFile(package, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 100)
    monkeypatch.setattr(extraction, "OFFICE_DOCUMENT_PART_MAX_BYTES", 32)
    with pytest.raises(M365ProviderError) as oversized:
        extraction.extract_m365_file(str(package), package.name)
    assert oversized.value.code == "package_limit"


def test_extraction_hard_limit_reports_missing_coverage(execution, tmp_path, monkeypatch, real_content_helpers):
    path = tmp_path / "large.txt"
    path.write_text("0123456789" * 10, encoding="utf-8")
    monkeypatch.setattr(extraction, "M365_EXTRACTED_TEXT_MAX_CHARS", 15)
    result = extraction.extract_m365_file(str(path), path.name)
    assert "".join(part.text for part in result.parts) == "012345678901234"
    assert result.coverage["complete"] is False
    assert result.coverage["missing_ranges"][0]["char_start"] == 15


@pytest.fixture
def graph_plugins(monkeypatch, execution):
    pending = types.ModuleType("functions_msgraph_pending_actions")
    constants = {
        "MSGRAPH_PENDING_ACTION_DELAYED": "delayed", "MSGRAPH_PENDING_ACTION_MANUAL": "manual",
        "MSGRAPH_PENDING_OPERATION_CREATE_CALENDAR_INVITE": "create_calendar_invite",
        "MSGRAPH_PENDING_OPERATION_SEND_MAIL": "send_mail",
        "MSGRAPH_PENDING_RESOURCE_CALENDAR": "calendar", "MSGRAPH_PENDING_RESOURCE_MAIL": "mail",
        "MSGRAPH_PENDING_STATUS_PENDING": "pending", "MSGRAPH_PENDING_STATUS_SCHEDULED": "scheduled",
    }
    for name, value in constants.items():
        setattr(pending, name, value)
    pending.build_calendar_pending_action_summary = lambda payload: {"subject": payload["subject"]}
    pending.build_mail_pending_action_summary = lambda payload: {"subject": payload["subject"]}
    pending.create_msgraph_pending_action = Mock(side_effect=lambda user_id, **kwargs: {"id": "pending-1", "user_id": user_id, **kwargs})
    pending.sanitize_msgraph_pending_action_for_client = lambda action, **kwargs: action
    pending.schedule_msgraph_pending_action_auto_commit = Mock()
    dependencies = {
        "functions_authentication": types.SimpleNamespace(get_current_user_info=lambda: {"userId": execution.actor_user_id}),
        "functions_debug": types.SimpleNamespace(debug_print=Mock()),
        "functions_group": types.SimpleNamespace(assert_group_role=Mock(), find_group_by_id=Mock(), require_active_group=Mock()),
        "functions_msgraph_pending_actions": pending,
        "semantic_kernel_plugins.plugin_invocation_logger": types.SimpleNamespace(plugin_function_logger=lambda name: lambda function: function),
    }
    for name, module in dependencies.items():
        monkeypatch.setitem(sys.modules, name, module)
    modules = {}
    for module_name in ("msgraph_plugin", "m365_calendar_plugin", "m365_email_plugin"):
        full_name = f"semantic_kernel_plugins.{module_name}"
        spec = importlib.util.spec_from_file_location(full_name, APP_DIR / "semantic_kernel_plugins" / f"{module_name}.py")
        module = importlib.util.module_from_spec(spec)
        monkeypatch.setitem(sys.modules, full_name, module)
        spec.loader.exec_module(module)
        package = importlib.import_module("semantic_kernel_plugins")
        monkeypatch.setattr(package, module_name, module, raising=False)
        modules[module_name] = module
    base = modules["msgraph_plugin"]
    monkeypatch.setattr(base, "log_m365_failure", Mock())
    monkeypatch.setattr(base, "get_m365_context", lambda **kwargs: execution)
    return modules, pending


def test_direct_typed_invocation_cannot_access_inherited_other_source_methods(execution, graph_plugins):
    modules, _ = graph_plugins
    calendar = modules["m365_calendar_plugin"].M365CalendarPlugin({
        "id": "calendar", "enabled_functions": ["get_my_messages", "send_mail", "get_my_profile"],
        "m365_capabilities": {"get_my_messages": True, "send_mail": True, "get_my_profile": True},
    })
    request = Mock(side_effect=AssertionError("Cross-source call performed network I/O."))
    calendar._perform_graph_request = request
    mail = calendar.get_my_messages()
    profile = calendar.get_my_profile()
    send = calendar.send_mail("person@example.test", "Subject")
    assert mail["error"] == profile["error"] == send["error"] == "function_not_enabled"
    assert calendar.get_functions() == []
    assert request.call_count == 0


def test_legacy_sensitive_direct_methods_check_source_sharing_before_pending_or_remote_io(execution, graph_plugins, monkeypatch):
    modules, pending = graph_plugins
    checked = []

    def denied(source, action_id, policy, **kwargs):
        checked.append(source)
        raise M365ProviderError("source_declined", "Source declined.")

    monkeypatch.setattr(modules["msgraph_plugin"], "authorize_m365_source", denied)
    plugin = modules["msgraph_plugin"].MSGraphPlugin({"id": "legacy"})
    calls = [
        plugin.get_my_messages(),
        plugin.get_my_events(),
        plugin.list_drive_items(),
        plugin.create_calendar_invite("Subject", "2026-09-17T12:00:00", "2026-09-17T13:00:00", timezone="UTC"),
    ]
    assert checked == ["email", "calendar", "onedrive", "calendar"]
    assert all(result["error"] == "source_declined" for result in calls)
    assert pending.create_msgraph_pending_action.call_count == 0


def test_calendar_group_recipients_use_actor_authorization_and_workflow_group(execution, graph_plugins, monkeypatch):
    modules, _ = graph_plugins
    module = modules["msgraph_plugin"]
    execution.actor_user_id = "workflow-caller"
    execution.data_user_id = "run-as"
    execution.workflow_id = "workflow"
    execution.group_id = "workflow-group"
    checks = []
    monkeypatch.setattr(module, "get_current_user_info", lambda: {"userId": "stale-session-owner"})
    monkeypatch.setattr(module, "assert_group_role", lambda *args, **kwargs: checks.append(args))
    monkeypatch.setattr(module, "find_group_by_id", lambda group_id: {
        "owner": {"email": "author@example.test"},
        "users": [{"email": "run-as@example.test"}],
    })
    plugin = modules["m365_calendar_plugin"].M365CalendarPlugin({"id": "calendar"})
    attendees = {}
    group_id, count = plugin._resolve_group_attendees("", attendees, current_user_email="run-as@example.test")
    assert checks == [("workflow-caller", "workflow-group")]
    assert group_id == "workflow-group"
    assert count == 1
    assert set(attendees) == {"author@example.test"}


def test_calendar_manual_delivery_reuses_pending_helper_and_authoritative_principal(execution, graph_plugins):
    modules, pending = graph_plugins
    plugin = modules["m365_calendar_plugin"].M365CalendarPlugin({
        "id": "calendar", "m365_capabilities": {"create_calendar_invite": True},
        "additionalFields": {"msgraph_calendar_send_mode": "draft_manual"},
    })
    client = GraphFixture().transport(source="calendar")
    plugin._transports[None] = client
    result = plugin.create_calendar_invite(
        "Subject", "2026-09-17T12:00:00", "2026-09-17T13:00:00", timezone="UTC",
    )
    assert result["pending_user_action"] is True
    assert pending.create_msgraph_pending_action.call_args.args[0] == execution.data_user_id
    assert pending.create_msgraph_pending_action.call_args.kwargs["conversation_id"] == execution.conversation_id
    assert pending.schedule_msgraph_pending_action_auto_commit.call_count == 0


class MemoryBlobFixture:
    def __init__(self):
        self.records = {}
        self.generation = 0

    def read(self, container, name, *, max_bytes):
        record = self.records.get((container, name))
        if record is None:
            raise MemoryNotFoundError("No test blob.")
        if len(record.data) > max_bytes:
            raise AssertionError("The test attempted an oversized blob read.")
        return record

    def put(self, container, name, data, *, etag=None):
        current = self.records.get((container, name))
        if etag is None and current is not None or etag is not None and (current is None or current.etag != etag):
            raise MemoryConflictError("Test ETag conflict.")
        self.generation += 1
        record = BlobRecord(bytes(data), f"etag-{self.generation}")
        self.records[(container, name)] = record
        return record.etag

    def delete(self, container, name, *, etag):
        current = self.records.get((container, name))
        if current is None or current.etag != etag:
            raise MemoryConflictError("Test ETag conflict.")
        del self.records[(container, name)]


@pytest.fixture
def memory_runtime(execution, monkeypatch):
    blob = MemoryBlobFixture()
    persisted_request_runs = {}
    settings = {"model_room": 12000}

    def authorize(ctx, operation):
        return (
            ctx.tenant_id == "tenant-1" and ctx.conversation_id == "conversation-1"
            and ctx.principal_id in ("user-1", "user-2") and ctx.storage_owner == "user-1"
        )

    def publish(ctx, run, grant_context):
        assert grant_context["execution_context"] is execution
        assert grant_context["source"] in ("onedrive", "spo")
        return PublicationGrant(
            tenant_id=ctx.tenant_id, principal_id=ctx.principal_id,
            conversation_id=ctx.conversation_id, run_id=run["run_id"],
            request_id=run["request_id"], content_revision=run["content_revision"],
            approval_ids=("approval-1",), authorization_id="authorization-1",
            audience_fingerprint="audience-1",
            approved_at=datetime.now(timezone.utc).isoformat(),
        )

    store = ConversationMemoryStore(transport=blob, authorize_access=authorize, authorize_publish=publish, log_event=Mock())

    def binding(context):
        return store, MemoryContext(
            context.tenant_id, context.data_user_id, context.conversation_id, "user-1",
        )

    def request_run(context):
        key = (context.data_user_id, context.request_id, context.conversation_id)
        if key not in persisted_request_runs:
            persisted_request_runs[key] = retrieval.create_m365_request_budget(*binding(context), context)
        return persisted_request_runs[key]

    monkeypatch.setattr(retrieval, "_memory_resolver", binding)
    monkeypatch.setattr(retrieval, "_request_run_resolver", request_run)
    monkeypatch.setattr(retrieval, "_model_budget_resolver", lambda context: settings["model_room"])
    monkeypatch.setattr(retrieval, "_token_counter", lambda text, context: len(text))
    monkeypatch.setattr(retrieval, "_analysis_choice", Mock(return_value={"mode": "extended", "audit_id": "analysis-1"}))
    return types.SimpleNamespace(
        store=store, blob=blob, binding=binding, request_run=request_run,
        persisted_request_runs=persisted_request_runs, settings=settings,
    )


class ContentGraphFixture(GraphFixture):
    def __init__(self, text="file content", source="onedrive"):
        super().__init__(source=source)
        self.text = text
        self.downloads = 0
        self.hide_size = False

    def item(self, item_id="item-1"):
        item = super().item(item_id)
        item["size"] = None if self.hide_size else len(self.text.encode())
        return item

    def request(self, method, url, **kwargs):
        if urlsplit(url).path.endswith("/content"):
            self.calls.append((method, url, kwargs))
            self.downloads += 1
            return FakeResponse(
                body=self.text.encode(),
                headers={"Content-Type": "text/plain", "Content-Length": str(len(self.text.encode()))},
            )
        return super().request(method, url, **kwargs)

    def operations(self):
        action_type = "m365_onedrive" if self.source == "onedrive" else "m365_sharepoint"
        instance = retrieval.M365FileOperations(action_type, {"id": f"action-{self.source}"})
        instance.transport = self.transport()
        return instance


def test_prepare_and_range_reads_retain_all_evidence_across_context_windows(execution, memory_runtime, real_content_helpers):
    text = "Exact evidence " * 1800
    fixture = ContentGraphFixture(text=text)
    actions = fixture.operations()
    prepared = actions.prepare_file("drive-1", "item-1")
    assert prepared["total_chunks"] > 1 and prepared["coverage"]["complete"] is True
    assert fixture.downloads == 1
    captured = []
    for index in range(prepared["total_chunks"]):
        window = actions.read_file_chunk(prepared["memory_id"], index)
        captured.append(window["text"])
        assert window["provider"] == "conversation_memory"
    assert "".join(captured) == text
    repeat = ContentGraphFixture(text=text)
    restored_actions = repeat.operations()
    restored = restored_actions.prepare_file("drive-1", "item-1")
    assert restored["memory_id"] == prepared["memory_id"]
    assert repeat.downloads == 0
    assert len(memory_runtime.persisted_request_runs) == 1


def test_published_snapshot_is_readable_by_another_participant_without_remote_calls(execution, memory_runtime, real_content_helpers, monkeypatch):
    execution.shared = True
    fixture = ContentGraphFixture()
    actions = fixture.operations()
    prepared = actions.prepare_file("drive-1", "item-1")
    assert prepared["snapshot_state"] == "published_snapshot"
    execution.actor_user_id = execution.data_user_id = "user-2"
    execution.request_id = "request-2"
    monkeypatch.setattr(
        retrieval, "authorize_m365_source",
        Mock(side_effect=AssertionError("Published evidence must not require a new remote source approval.")),
    )
    fixture.calls.clear()
    window = actions.read_file_chunk(prepared["memory_id"])
    assert window["text"] == "file content"
    assert window["capture"]["principal_id"] == "user-1"
    assert fixture.calls == []
    spo = ContentGraphFixture(source="spo").operations()
    with pytest.raises(M365ProviderError) as wrong_source:
        spo.read_file_chunk(prepared["memory_id"])
    assert wrong_source.value.code == "source_not_allowed"


def test_private_staging_and_mismatched_request_budget_are_not_reusable_by_another_actor(execution, memory_runtime, real_content_helpers, monkeypatch):
    actions = ContentGraphFixture().operations()
    prepared = actions.prepare_file("drive-1", "item-1")
    original_budget = memory_runtime.request_run(execution)
    execution.actor_user_id = execution.data_user_id = "user-2"
    execution.request_id = "request-2"
    with pytest.raises(retrieval.MemoryAuthorizationError):
        actions.read_file_chunk(prepared["memory_id"])
    execution.actor_user_id = execution.data_user_id = "user-1"
    monkeypatch.setattr(retrieval, "_request_run_resolver", lambda context: original_budget)
    with pytest.raises(M365ProviderError) as wrong_request:
        actions.read_file_chunk(prepared["memory_id"])
    assert wrong_request.value.code == "request_memory_mismatch"


def _analysis_required():
    return M365ApprovalRequired({
        "id": "analysis-approval-1", "request_type": "m365_extended_analysis",
        "subject_user_id": "user-1", "resume_key": "resume-1",
        "execution_status": "awaiting_approval", "status": "pending",
    })


def test_three_downloads_are_per_logical_request_across_sources_and_reinstantiation(execution, memory_runtime, real_content_helpers, monkeypatch):
    onedrive, spo = ContentGraphFixture(), ContentGraphFixture(source="spo")
    onedrive.operations().prepare_file("drive-1", "item-1")
    spo.operations().prepare_file("drive-1", "item-2")
    onedrive.operations().prepare_file("drive-1", "item-3")
    ask = Mock(side_effect=_analysis_required())
    monkeypatch.setattr(retrieval, "_analysis_choice", ask)
    with pytest.raises(M365ApprovalRequired):
        spo.operations().prepare_file("drive-1", "item-4")
    assert onedrive.downloads + spo.downloads == 3
    assert ask.call_args.args[3]["download_count"] == 4
    monkeypatch.setattr(retrieval, "_analysis_choice", Mock(return_value={"mode": "extended", "approval_id": "analysis-approval-1"}))
    resumed = spo.operations().prepare_file("drive-1", "item-4")
    assert resumed["coverage"]["complete"] is True
    assert onedrive.downloads + spo.downloads == 4


def test_file_size_window_is_soft_and_unknown_length_can_resume_with_approval(execution, memory_runtime, real_content_helpers, monkeypatch):
    monkeypatch.setattr(retrieval, "M365_FAST_FILE_BYTES", 5)
    fixture = ContentGraphFixture(text="larger file content")
    fixture.hide_size = True
    actions = fixture.operations()
    monkeypatch.setattr(retrieval, "_analysis_choice", Mock(side_effect=_analysis_required()))
    with pytest.raises(M365ApprovalRequired):
        actions.prepare_file("drive-1", "item-1")
    assert fixture.downloads == 1
    monkeypatch.setattr(retrieval, "_analysis_choice", Mock(return_value={"mode": "extended", "approval_id": "analysis-approval-1"}))
    prepared = actions.prepare_file("drive-1", "item-1")
    assert prepared["coverage"]["complete"] is True
    assert prepared["file"]["size_bytes"] > 5
    assert fixture.downloads == 2


def test_context_soft_threshold_and_model_room_preserve_resumable_exact_ranges(execution, memory_runtime, real_content_helpers, monkeypatch):
    text = "0123456789" * 4
    actions = ContentGraphFixture(text=text).operations()
    prepared = actions.prepare_file("drive-1", "item-1")
    memory_runtime.settings["model_room"] = 13
    monkeypatch.setattr(retrieval, "M365_FAST_CONTEXT_TOKENS", 20)
    choice = Mock(return_value={"mode": "extended", "approval_id": "analysis-approval-1"})
    monkeypatch.setattr(retrieval, "_analysis_choice", choice)
    offset, read = 0, []
    while offset is not None:
        window = actions.read_file_chunk(prepared["memory_id"], 0, offset)
        read.append(window["text"])
        assert len(window["text"]) <= 13
        offset = window["next_char_offset"]
    assert "".join(read) == text
    assert choice.call_count >= 1
    assert window["coverage"]["logical_request_context_tokens"] == len(text)


def test_discovery_uses_checkpointed_pages_without_returning_file_context(execution, memory_runtime):
    fixture = GraphFixture()
    fixture.more = True
    actions = retrieval.M365FileOperations("m365_onedrive", {"id": "action-onedrive"})
    actions.transport = fixture.transport()
    first = actions.discover_files("report")
    assert first["next_offset"] == 1 and first["memory_id"]
    assert "excerpts" not in first["results"][0]
    fixture.more = False
    second = actions.discover_files("report", memory_id=first["memory_id"])
    assert second["memory_id"] == first["memory_id"] and second["next_offset"] is None
    assert second["coverage"]["logical_discovery_candidates_inspected"] == 2
    posts = [kwargs["json"]["requests"][0]["from"] for method, url, kwargs in fixture.calls if url.endswith("/search/query")]
    assert posts == [0, 1]


def test_search_retains_evidence_and_resume_does_not_repeat_remote_search(execution, memory_runtime):
    fixture = GraphFixture(licensed=True)
    actions = retrieval.M365FileOperations("m365_onedrive", {"id": "action-onedrive"})
    actions.transport = fixture.transport()
    first = actions.search_files("report")
    initial_calls = len(fixture.calls)
    second = actions.search_files("report")
    assert len(fixture.calls) == initial_calls
    assert first["results"][0]["memory_id"] == second["results"][0]["memory_id"]
    assert first["results"][0]["excerpts"][0]["text"] == "The exact retained Copilot excerpt."
    snapshot = actions.read_file_chunk(first["results"][0]["memory_id"])
    assert snapshot["coverage"]["complete"] is False


def test_real_file_plugins_filter_registration_and_direct_invocation(execution, memory_runtime):
    from semantic_kernel_plugins.m365_onedrive_plugin import M365OneDrivePlugin
    from semantic_kernel_plugins.m365_sharepoint_plugin import M365SharePointPlugin

    for plugin_type in (M365OneDrivePlugin, M365SharePointPlugin):
        plugin = plugin_type({
            "id": "action-1", "enabled_functions": ["send_mail", "get_my_events", "search_files"],
            "m365_capabilities": {"search_files": False},
        })
        functions = plugin.get_functions()
        result = plugin.search_files("report")
        assert functions == []
        assert result["error"]["code"] == "function_not_enabled"
        assert not hasattr(plugin, "send_mail") and not hasattr(plugin, "get_my_events")


def test_fast_windows_match_the_approved_limits():
    assert retrieval.M365_FAST_DOWNLOADS == 3
    assert retrieval.M365_FAST_FILE_BYTES == 25 * 1024 * 1024
    assert retrieval.M365_FAST_CONTEXT_TOKENS == 12000


def test_type_bounds_survive_legacy_loader_and_runtime_capability_mutations(execution, graph_plugins):
    modules, _ = graph_plugins
    legacy_loader = modules["msgraph_plugin"].MSGraphPlugin({
        "id": "calendar", "type": "m365_calendar",
        "enabled_functions": ["get_my_events", "get_my_messages", "get_my_profile"],
        "m365_capabilities": {"get_my_events": True, "get_my_messages": True},
    })
    legacy_loader._enabled_function_names.update({"get_my_messages", "get_my_profile"})
    legacy_loader._capabilities.update({"get_my_messages": True, "get_my_profile": True})
    registered = legacy_loader.get_functions()
    mail = legacy_loader.get_my_messages()
    profile = legacy_loader.get_my_profile()
    assert registered == ["get_my_events"]
    assert mail["error"] == profile["error"] == "source_not_allowed"
    email = modules["m365_email_plugin"].M365EmailPlugin({
        "id": "email", "m365_capabilities": {"search_users": True},
    })
    selected = email.search_users("Ada", select_fields="mailboxSettings")
    assert selected["error"] == "source_not_allowed"


def test_calendar_timezone_access_failure_cannot_silently_create_utc_invite(execution, graph_plugins):
    modules, pending = graph_plugins
    plugin = modules["m365_calendar_plugin"].M365CalendarPlugin({
        "id": "calendar", "m365_capabilities": {"create_calendar_invite": True},
    })
    plugin._perform_graph_request = Mock(return_value={"error": "access_denied"})
    result = plugin.create_calendar_invite("Meeting", "2026-09-17T12:00:00", "2026-09-17T13:00:00")
    assert result["error"] == "access_denied"
    assert pending.create_msgraph_pending_action.call_count == 0


def test_copilot_policy_denial_blocks_later_raw_reads_and_searches_for_that_source(execution, memory_runtime):
    fixture = GraphFixture(source="spo", licensed=True)
    fixture.copilot_response = FakeResponse({"error": {"code": "blockedByPolicy"}}, status=403)
    actions = retrieval.M365FileOperations("m365_sharepoint", {"id": "action-spo"})
    actions.transport = fixture.transport()
    with pytest.raises(M365ProviderError) as denied:
        actions.search_files("report")
    assert denied.value.as_result("spo")["provider"] == "copilot_retrieval"
    calls_after_denial = len(fixture.calls)
    for operation in (
        lambda: actions.prepare_file("drive-1", "item-1"),
        lambda: actions.discover_files("report"),
    ):
        with pytest.raises(M365ProviderError) as blocked:
            operation()
        assert blocked.value.code == "source_policy_blocked"
    assert len(fixture.calls) == calls_after_denial
    other_fixture = GraphFixture()
    other = retrieval.M365FileOperations("m365_onedrive", {"id": "action-onedrive"})
    other.transport = other_fixture.transport()
    allowed = other.search_files("report")
    assert allowed["results"]


def test_simultaneous_tools_cannot_use_separate_fast_budgets(execution, memory_runtime):
    run_id = memory_runtime.request_run(execution)
    store, context = memory_runtime.binding(execution)
    claim = store.claim(context, run_id)
    fixture = ContentGraphFixture()
    try:
        with pytest.raises(MemoryConflictError):
            fixture.operations().prepare_file("drive-1", "item-1")
        assert fixture.calls == []
    finally:
        store.release_claim(context, claim, status="queued")


def test_exact_12000_token_window_requires_approval_for_the_next_token(execution, memory_runtime, real_content_helpers, monkeypatch):
    actions = ContentGraphFixture(text="x" * 12001).operations()
    prepared = actions.prepare_file("drive-1", "item-1")
    first = actions.read_file_chunk(prepared["memory_id"], 0)
    assert len(first["text"]) == 8000
    ask = Mock(side_effect=_analysis_required())
    monkeypatch.setattr(retrieval, "_analysis_choice", ask)
    with pytest.raises(M365ApprovalRequired):
        actions.read_file_chunk(prepared["memory_id"], 1)
    assert ask.call_args.args[3]["context_tokens"] == 12001
    monkeypatch.setattr(retrieval, "_analysis_choice", Mock(return_value={"mode": "fast"}))
    fast = actions.read_file_chunk(prepared["memory_id"], 1)
    assert len(fast["text"]) == 4000
    assert fast["next_char_offset"] == 4000
    assert fast["coverage"]["logical_request_context_tokens"] == 12000
    assert fast["coverage"]["complete"] is False
    with pytest.raises(M365ProviderError) as limited:
        actions.read_file_chunk(prepared["memory_id"], 1, fast["next_char_offset"])
    assert limited.value.code == "fast_analysis_limit"


def test_metadata_size_mismatch_and_source_changes_cannot_publish_complete_capture(execution, memory_runtime, real_content_helpers, monkeypatch):
    execution.shared = True
    fixture = ContentGraphFixture()
    original = fixture.item

    def wrong_size(item_id="item-1"):
        item = original(item_id)
        item["size"] += 1
        return item

    monkeypatch.setattr(fixture, "item", wrong_size)
    with pytest.raises(M365ProviderError) as size_error:
        fixture.operations().prepare_file("drive-1", "item-1")
    assert size_error.value.code == "source_size_mismatch"
    store, ctx = memory_runtime.binding(execution)
    manifest = store.read_manifest(ctx, size_error.value.details["memory_id"])
    assert manifest["publication"] is None and manifest["evidence_count"] == 0

    def changed(item_id="item-1"):
        item = original(item_id)
        if fixture.downloads > 1:
            item["eTag"] = '"changed-version"'
        return item

    monkeypatch.setattr(fixture, "item", changed)
    with pytest.raises(M365ProviderError) as changed_error:
        fixture.operations().prepare_file("drive-1", "item-2")
    assert changed_error.value.code == "source_changed"


def test_invalid_folder_shapes_never_become_unscoped_search(execution):
    fixture = GraphFixture()
    provider = retrieval.M365FileProvider(fixture.transport())
    for folder in ([], {}, False, None):
        with pytest.raises(M365ProviderError):
            provider.search("report", folder=folder)
    assert fixture.calls == []


def test_unreadable_xml_and_scanned_pdf_return_explicit_extraction_errors(execution, tmp_path, real_content_helpers):
    word = tmp_path / "invalid.docx"
    with zipfile.ZipFile(word, "w") as archive:
        archive.writestr("word/document.xml", "<malformed")
    with pytest.raises(M365ProviderError) as malformed:
        extraction.extract_m365_file(str(word), word.name)
    assert malformed.value.code == "extraction_failed"
    pdf = tmp_path / "scanned.pdf"
    with fitz.open() as document:
        document.new_page()
        document.save(pdf)
    with pytest.raises(M365ProviderError) as scanned:
        extraction.extract_m365_file(str(pdf), pdf.name)
    assert scanned.value.code == "no_extractable_text"
    assert scanned.value.details["coverage"]["complete"] is False


def test_real_execution_contract_checks_saved_source_and_policy_before_token_acquisition(execution, monkeypatch):
    import functions_m365_execution as execution_module
    from functions_m365_approvals import M365PolicyError

    context = execution_module.M365ExecutionContext(
        actor_user_id="user-1", data_user_id="user-1", tenant_id="tenant-1",
        conversation_id="conversation-1", request_id="request-1",
        action_configs={
            "action-1": {"source": "onedrive", "maximum_sharing_acknowledgement": "today"},
        },
    )
    service = types.SimpleNamespace(authorize_sources=Mock(return_value={
        "onedrive": {"source": "onedrive", "sharing_required": False},
    }))
    monkeypatch.setattr(execution_module, "get_m365_approval_service", lambda: service)
    monkeypatch.setattr(transport_module, "get_m365_context", lambda **kwargs: execution_module.require_m365_execution_context())
    monkeypatch.setattr(transport_module, "authorize_m365_source", REAL_AUTHORIZE_SOURCE)
    token_provider = Mock(return_value={"access_token": "unit-test-token"})
    cloud = M365CloudConfig("https://graph.microsoft.us/v1.0", "https://login.microsoftonline.us/tenant-1")
    client = M365Transport("onedrive", "action-1", {"maximum_sharing_acknowledgement": "always"}, cloud=cloud, token_provider=token_provider)
    with execution_module.m365_execution_context(context):
        token, scopes = client.get_token(["Files.Read.All"])
        with pytest.raises(M365PolicyError) as other_source:
            M365Transport("spo", "action-1", cloud=cloud, token_provider=token_provider).get_token(["Files.Read.All"])
    assert token == "unit-test-token"
    assert scopes == ["https://graph.microsoft.us/Files.Read.All"]
    assert token_provider.call_count == 1
    assert service.authorize_sources.call_args.args[1] == {"onedrive": "today"}
    assert other_source.value.code == "m365_source_not_authorized"


def test_transport_strips_download_links_even_from_legacy_result_shapes(execution):
    fixture = GraphFixture()
    payload = fixture.transport().request_json("GET", "/drives/drive-1/items/item-1", ["Files.Read.All"])
    assert payload["webUrl"] == fixture.web_url
    assert "@microsoft.graph.downloadUrl" not in payload
    assert "do-not-return" not in json.dumps(payload)


def test_new_calendar_never_defaults_missing_or_invalid_delivery_to_auto_send(execution, graph_plugins):
    modules, _ = graph_plugins
    plugin = modules["m365_calendar_plugin"].M365CalendarPlugin({
        "id": "calendar", "msgraph_calendar_send_mode": None,
        "m365_capabilities": {"create_calendar_invite": True},
    })
    assert plugin._calendar_send_mode == "draft_manual"
    with pytest.raises(ValueError):
        modules["m365_calendar_plugin"].M365CalendarPlugin({
            "additionalFields": {"msgraph_calendar_send_mode": "invalid-mode"},
        })


def test_cancellation_prevents_a_later_remote_read_after_capture(execution, memory_runtime, real_content_helpers, monkeypatch):
    fixture = ContentGraphFixture()
    actions = fixture.operations()
    original_extract = retrieval.extract_m365_file

    def cancel_after_extract(*args):
        extracted = original_extract(*args)
        store, context = memory_runtime.binding(execution)
        store.cancel(context, memory_runtime.request_run(execution))
        return extracted

    monkeypatch.setattr(retrieval, "extract_m365_file", cancel_after_extract)
    with pytest.raises(MemoryConflictError):
        actions.prepare_file("drive-1", "item-1")
    item_reads = [url for method, url, _ in fixture.calls if "/items/" in url and not url.endswith("/content")]
    assert len(item_reads) == 1
    assert actions.transport.before_request is None and actions.transport.on_progress is None


def test_file_download_has_a_total_time_bound_as_well_as_byte_bounds(execution, tmp_path, monkeypatch):
    moments = iter((0.0, float(transport_module.M365_DOWNLOAD_MAX_SECONDS + 1)))
    monkeypatch.setattr(transport_module.time, "monotonic", lambda: next(moments))
    monkeypatch.setattr(transport_module.tempfile, "tempdir", str(tmp_path))
    client = M365Transport(
        "onedrive",
        cloud=M365CloudConfig("https://graph.microsoft.com/v1.0", "https://login.microsoftonline.com/tenant"),
        request=Mock(return_value=FakeResponse(body=b"content", headers={"Content-Type": "text/plain"})),
        token_provider=lambda scopes, context: {"access_token": "unit-test-token"},
    )
    with pytest.raises(M365ProviderError) as timed_out:
        with client.download_file("drive-1", "item-1", suffix=".txt", allowed_mime_types=("text/plain",), max_bytes=100):
            pytest.fail("An over-time download must not be returned.")
    assert timed_out.value.code == "timeout"
    remaining = list(tmp_path.iterdir())
    assert remaining == []


def test_real_source_decline_returns_a_legacy_tool_error_without_canceling_other_sources(execution, graph_plugins, monkeypatch):
    from functions_m365_approvals import M365SourceDenied

    modules, _ = graph_plugins

    def source_decision(source, action_id, policy, **kwargs):
        if source == "email":
            raise M365SourceDenied("email", "approval-1")
        return execution, {"source": source, "sharing_required": False}

    monkeypatch.setattr(modules["msgraph_plugin"], "authorize_m365_source", source_decision)
    plugin = modules["msgraph_plugin"].MSGraphPlugin({"id": "legacy"})
    denied = plugin.get_my_messages()
    calendar_denial = plugin._authorize_operation("get_my_events")
    assert denied["error"] == "m365_source_declined" and denied["source"] == "email"
    assert calendar_denial is None


def test_private_capture_requires_publication_approval_before_shared_chat_read(execution, memory_runtime, real_content_helpers, monkeypatch):
    from functions_m365_approvals import M365SourceDenied

    actions = ContentGraphFixture().operations()
    private = actions.prepare_file("drive-1", "item-1")
    execution.shared = True
    execution.request_id = "request-2"
    denied = Mock(side_effect=M365SourceDenied("onedrive", "denial-1"))
    monkeypatch.setattr(retrieval, "authorize_m365_publication", denied)
    with pytest.raises(M365SourceDenied):
        actions.read_file_chunk(private["memory_id"])
    store, ctx = memory_runtime.binding(execution)
    manifest = store.read_manifest(ctx, private["memory_id"])
    assert manifest["publication"] is None
    monkeypatch.setattr(
        retrieval, "authorize_m365_publication",
        lambda source, action_id, policy, **kwargs: (execution, {"source": source, "approval_id": "approval-1"}),
    )
    published = actions.read_file_chunk(private["memory_id"])
    assert published["text"] == "file content"
    assert published["snapshot_state"] == "published_snapshot"


def test_snapshot_only_action_publishes_a_prior_capture_without_remote_access(
    execution, memory_runtime, real_content_helpers, monkeypatch,
):
    from dataclasses import replace
    import functions_m365_approvals as approval_module
    import functions_m365_execution as execution_module
    from test_m365_runtime_adapters import load_module, module_stub
    from test_support.m365 import CosmosContainer, Notifications

    actions = ContentGraphFixture().operations()
    prepared = actions.prepare_file("drive-1", "item-1")
    manifest = {
        "id": actions.action_id, "type": "m365_onedrive", "source": "onedrive",
        "enabled_functions": ["read_file_chunk"],
    }
    actions.manifest["enabled_functions"] = ["read_file_chunk"]
    context = execution_module.M365ExecutionContext(
        actor_user_id="user-1", data_user_id="user-1", tenant_id="tenant-1",
        conversation_id="conversation-1", request_id="new-sharing-request",
        shared=True, audience_version="audience-1", action_configs={actions.action_id: manifest},
    )
    container = CosmosContainer()
    service = approval_module.M365ApprovalService(
        container_factory=lambda: container, notification_sender=Notifications(),
        decision_validator=lambda approval: True,
    )
    runtime = load_module("conversation_memory_runtime", {
        "config": module_stub(
            "config", CLIENTS={}, TENANT_ID="tenant-1", cosmos_conversations_container=None,
            cosmos_messages_container=None, build_enhanced_citations_blob_service_client=lambda settings: None,
        ),
        "functions_appinsights": module_stub("functions_appinsights", log_event=Mock()),
        "functions_collaboration": module_stub("functions_collaboration", build_conversation_participation_context=lambda *args: None),
        "functions_settings": module_stub("functions_settings", get_settings=lambda: {}),
    })
    monkeypatch.setattr(memory_runtime.store, "authorize_publish", runtime._authorize_memory_publication)
    monkeypatch.setattr(approval_module, "_service", service)
    monkeypatch.setattr(execution_module, "_action_config_resolver", lambda *args: manifest)
    monkeypatch.setattr(execution_module, "_action_selection_resolver", lambda context: [actions.action_id])
    for module in (retrieval, transport_module):
        monkeypatch.setattr(module, "get_m365_context", lambda **kwargs: execution_module.get_m365_execution_context())
    monkeypatch.setattr(retrieval, "authorize_m365_capability", REAL_AUTHORIZE_CAPABILITY)
    monkeypatch.setattr(retrieval, "authorize_m365_publication", transport_module.authorize_m365_publication)
    forbidden = Mock(side_effect=AssertionError("Publishing a retained snapshot must not fetch source data or tokens."))
    monkeypatch.setattr(retrieval, "authorize_m365_source", forbidden)
    monkeypatch.setattr(M365Transport, "get_token", forbidden)
    monkeypatch.setattr(M365Transport, "request_json", forbidden)
    with execution_module.m365_execution_context(context):
        with pytest.raises(M365ApprovalRequired) as pending:
            actions.read_file_chunk(prepared["memory_id"])
        service.decide(pending.value.approval_id, "user-1", {
            "decisions": {"onedrive": {"duration": "request", "timezone": "UTC"}},
        })
        published = actions.read_file_chunk(prepared["memory_id"])
    reader = replace(context, actor_user_id="user-2", data_user_id="user-2", request_id="reader-request")
    with execution_module.m365_execution_context(reader):
        reused = actions.read_file_chunk(prepared["memory_id"])
    assert published["snapshot_state"] == reused["snapshot_state"] == "published_snapshot"
    assert published["text"] == reused["text"] == "file content"
    assert forbidden.call_count == 0


def test_source_schemas_accept_defaults_and_reject_cross_source_capabilities():
    for action_type in operations.M365_ACTION_TYPES:
        schema = operations.get_m365_schema_for_type(action_type)
        config = operations.get_m365_default_config(action_type)
        jsonschema.Draft202012Validator.check_schema(schema)
        jsonschema.validate(config, schema)
        config["additionalFields"]["m365_capabilities"]["get_my_security_alerts"] = True
        with pytest.raises(jsonschema.ValidationError):
            jsonschema.validate(config, schema)


def test_transport_worker_callbacks_are_isolated_between_execution_contexts(execution):
    client = GraphFixture().transport()
    other_context = copy_context()
    first, second = Mock(), Mock()
    with client.callback_context(first, None):
        first_result = client.request_json("GET", "/drives/drive-1", ["Files.Read.All"])

        def other_call():
            assert client.before_request is None
            with client.callback_context(second, None):
                return client.request_json("GET", "/drives/drive-1", ["Files.Read.All"])

        second_result = other_context.run(other_call)
        assert client.before_request is first
    assert first_result["id"] == second_result["id"] == "drive-1"
    assert first.call_count == second.call_count == 1
    assert client.before_request is None


def test_calendar_and_email_default_kernel_names_are_source_specific(execution, graph_plugins):
    modules, _ = graph_plugins
    for module_name, class_name, action_type in (
        ("m365_calendar_plugin", "M365CalendarPlugin", "m365_calendar"),
        ("m365_email_plugin", "M365EmailPlugin", "m365_email"),
    ):
        plugin = getattr(modules[module_name], class_name)()
        kernel = plugin.get_kernel_plugin()
        assert plugin.metadata["name"] == action_type
        assert kernel.name == action_type


def test_all_typed_plugins_register_without_source_access_or_execution_context(graph_plugins, monkeypatch):
    from semantic_kernel_plugins.m365_onedrive_plugin import M365OneDrivePlugin
    from semantic_kernel_plugins.m365_sharepoint_plugin import M365SharePointPlugin

    modules, _ = graph_plugins
    forbidden = Mock(side_effect=AssertionError("Construction and registration must not access a source before preflight."))
    for module, names in (
        (modules["msgraph_plugin"], ("authorize_m365_source", "authorize_m365_capability", "get_m365_context", "get_current_user_info")),
        (transport_module, ("authorize_m365_source", "authorize_m365_capability", "get_m365_context", "get_m365_cloud_config", "_delegated_token")),
        (retrieval, ("authorize_m365_source", "authorize_m365_capability", "get_m365_context", "_memory_resolver", "_model_budget_resolver")),
    ):
        for name in names:
            monkeypatch.setattr(module, name, forbidden)
    plugin_types = (
        modules["m365_calendar_plugin"].M365CalendarPlugin,
        modules["m365_email_plugin"].M365EmailPlugin,
        M365OneDrivePlugin,
        M365SharePointPlugin,
    )
    for plugin_type in plugin_types:
        plugin = plugin_type()
        metadata = plugin.metadata
        functions = plugin.get_functions()
        kernel = plugin.get_kernel_plugin()
        assert functions
        assert set(kernel.functions) == set(functions)
        assert {method["name"] for method in metadata["methods"]} == set(functions)
    assert forbidden.call_count == 0


def test_all_typed_direct_guards_preserve_approval_required_exceptions(graph_plugins, monkeypatch):
    from semantic_kernel_plugins.m365_onedrive_plugin import M365OneDrivePlugin
    from semantic_kernel_plugins.m365_sharepoint_plugin import M365SharePointPlugin

    modules, _ = graph_plugins
    approval = _analysis_required()
    guard = Mock(side_effect=approval)
    monkeypatch.setattr(modules["msgraph_plugin"], "authorize_m365_source", guard)
    monkeypatch.setattr(retrieval, "authorize_m365_source", guard)
    remote = Mock(side_effect=AssertionError("No remote I/O is allowed before approval."))
    monkeypatch.setattr(M365Transport, "request_json", remote)
    monkeypatch.setattr(M365Transport, "get_token", remote)
    cases = (
        (modules["m365_calendar_plugin"].M365CalendarPlugin, "get_my_events", ()),
        (modules["m365_email_plugin"].M365EmailPlugin, "get_my_messages", ()),
        (M365OneDrivePlugin, "search_files", ("report",)),
        (M365SharePointPlugin, "search_files", ("report",)),
    )
    for plugin_type, method, arguments in cases:
        plugin = plugin_type({"id": "action-1"})
        with pytest.raises(M365ApprovalRequired) as caught:
            getattr(plugin, method)(*arguments)
        assert caught.value is approval
        assert caught.value.payload["approval_id"] == "analysis-approval-1"
    assert guard.call_count == 4
    assert remote.call_count == 0


def test_legacy_mail_pagination_preserves_result_shapes_and_delegated_scopes(execution, graph_plugins):
    modules, _ = graph_plugins
    responses = [
        FakeResponse({"value": [{"id": "1"}, {"id": "2"}], "@odata.nextLink": "https://graph.microsoft.com/v1.0/me/messages?page=2"}),
        FakeResponse({"value": [{"id": "3"}, {"id": "4"}]}),
    ]
    token_provider = Mock(return_value={"access_token": "unit-test-token"})
    request = Mock(side_effect=lambda *args, **kwargs: responses.pop(0))
    transport = M365Transport(
        "email", "legacy",
        cloud=M365CloudConfig("https://graph.microsoft.com/v1.0", "https://login.microsoftonline.com/tenant"),
        request=request, token_provider=token_provider,
    )
    plugin = modules["msgraph_plugin"].MSGraphPlugin({"id": "legacy"})
    plugin._transports["email"] = transport
    result = plugin.get_my_messages(top=3)
    assert result["count"] == 3 and result["truncated"] is True
    assert result["value"] == [{"id": "1"}, {"id": "2"}, {"id": "3"}]
    assert result["source"] == "email" and request.call_count == 2
    assert token_provider.call_args.args[0] == ["https://graph.microsoft.com/Mail.Read"]


@pytest.mark.parametrize("mode,expected_scope", [
    ("draft_manual", "Mail.ReadWrite"), ("auto_send", "Mail.Send"),
])
def test_typed_email_reuses_manual_and_auto_delivery_business_operations(execution, graph_plugins, mode, expected_scope):
    modules, pending = graph_plugins
    response = FakeResponse(
        {"id": "draft-1", "subject": "Subject", "changeKey": "draft-v1", "isDraft": True}, status=201,
    ) if mode == "draft_manual" else FakeResponse(body=b"", status=202)
    request = Mock(return_value=response)
    token_provider = Mock(return_value={"access_token": "unit-test-token"})
    client = M365Transport(
        "email", "email",
        cloud=M365CloudConfig("https://graph.microsoft.com/v1.0", "https://login.microsoftonline.com/tenant"),
        request=request, token_provider=token_provider,
    )
    plugin = modules["m365_email_plugin"].M365EmailPlugin({
        "id": "email", "m365_capabilities": {"send_mail": True},
        "additionalFields": {"msgraph_mail_send_mode": mode},
    })
    plugin._transports["email"] = plugin._transports[None] = client
    result = plugin.send_mail("person@example.test", "Subject", "Plain-text body")
    assert token_provider.call_args.args[0] == [f"https://graph.microsoft.com/{expected_scope}"]
    assert request.call_args.args[0] == "POST"
    if mode == "draft_manual":
        assert result["pending_user_action"] is True
        assert pending.create_msgraph_pending_action.call_args.args[0] == execution.data_user_id
        assert pending.create_msgraph_pending_action.call_args.kwargs["graph_message_id"] == "draft-1"
    else:
        assert result["mail_send_status"] == "sent"
        assert pending.create_msgraph_pending_action.call_count == 0


def test_office_viewer_links_resolve_to_canonical_file_paths_not_query_urls(execution):
    fixture = GraphFixture(source="spo")
    canonical = fixture.web_url
    original_item = fixture.item
    original_request = fixture.request

    def viewer_item(item_id="item-1"):
        item = original_item(item_id)
        item["webUrl"] = "https://tenant.sharepoint.com/sites/team/_layouts/15/Doc.aspx?sourcedoc=doc-id&file=report.txt"
        item["parentReference"]["path"] = "/drives/drive-1/root:"
        return item

    fixture.item = viewer_item

    def request(method, url, **kwargs):
        if urlsplit(url).path.endswith("/drives/drive-1"):
            fixture.calls.append((method, url, kwargs))
            return FakeResponse({
                "id": "drive-1", "driveType": "documentLibrary",
                "webUrl": "https://tenant.sharepoint.com/sites/team/Documents",
            })
        return original_request(method, url, **kwargs)

    fixture.request = request
    provider = retrieval.M365FileProvider(fixture.transport())
    resolved = provider.resolve_file(web_url=canonical)
    assert resolved["web_url"] == canonical
    assert "?" not in resolved["web_url"]
    assert "_layouts" not in resolved["web_url"]


def test_preauthenticated_download_urls_are_excluded_from_dependency_telemetry(execution):
    telemetry = []
    responses = [
        FakeResponse(status=302, headers={"Location": "https://tenant.sharepoint.com/_layouts/download.aspx?tempauth=secret"}),
        FakeResponse(body=b"data", headers={"Content-Type": "text/plain"}),
    ]

    def request(method, url, **kwargs):
        telemetry.append(is_http_instrumentation_enabled())
        return responses.pop(0)

    client = M365Transport(
        "onedrive",
        cloud=M365CloudConfig("https://graph.microsoft.com/v1.0", "https://login.microsoftonline.com/tenant"),
        request=request, token_provider=lambda scopes, context: {"access_token": "unit-test-token"},
    )
    with client.download_file("drive-1", "item-1", suffix=".txt", allowed_mime_types=("text/plain",), max_bytes=100) as downloaded:
        assert downloaded.size_bytes == 4
    assert telemetry == [True, False]
    assert is_http_instrumentation_enabled()


def test_capture_recovery_does_not_repeat_a_committed_content_download(execution, memory_runtime, real_content_helpers, monkeypatch):
    fixture = ContentGraphFixture()
    store = memory_runtime.store
    original = store.add_evidence

    def interrupted(*args, **kwargs):
        original(*args, **kwargs)
        raise RuntimeError("Simulated worker interruption after evidence commit.")

    monkeypatch.setattr(store, "add_evidence", interrupted)
    with pytest.raises(RuntimeError):
        fixture.operations().prepare_file("drive-1", "item-1")
    assert fixture.downloads == 1
    monkeypatch.setattr(store, "add_evidence", original)
    resumed = fixture.operations().prepare_file("drive-1", "item-1")
    assert resumed["coverage"]["complete"] is True
    assert fixture.downloads == 1
    window = fixture.operations().read_file_chunk(resumed["memory_id"])
    assert window["text"] == "file content"


def test_default_budget_binding_is_idempotent_and_survives_plugin_recreation(execution, memory_runtime, real_content_helpers, monkeypatch):
    monkeypatch.setattr(retrieval, "_request_run_resolver", None)
    fixture = ContentGraphFixture()
    first = fixture.operations().prepare_file("drive-1", "item-1")
    second = fixture.operations().prepare_file("drive-1", "item-1")
    store, ctx = memory_runtime.binding(execution)
    first_budget = retrieval.create_m365_request_budget(store, ctx, execution)
    second_budget = retrieval.create_m365_request_budget(store, ctx, execution)
    checkpoint = store.read_checkpoint(ctx, first_budget)
    assert first_budget == second_budget
    assert first["memory_id"] == second["memory_id"]
    assert fixture.downloads == 1
    assert checkpoint["checkpoint"]["download_count"] == 1


def test_copilot_eligibility_requires_a_matching_supported_cloud_profile(execution):
    client = GraphFixture(licensed=True).transport(cloud=M365CloudConfig(
        "https://graph.microsoft.com/v1.0", "https://login.microsoftonline.us/tenant",
    ))
    provider, reason = retrieval.select_m365_retrieval_provider(client)
    assert provider == "graph" and reason == "copilot_retrieval_unsupported_in_cloud"


@pytest.mark.parametrize("action_type,disabled,retained", [
    ("m365_calendar", "get_my_events", "create_calendar_invite"),
    ("m365_email", "get_my_messages", "send_mail"),
    ("m365_onedrive", "search_files", "read_file"),
    ("m365_sharepoint", "search_files", "read_file"),
])
def test_top_level_runtime_capabilities_cannot_broaden_saved_action(action_type, disabled, retained):
    config = {
        "additionalFields": {
            "m365_capabilities": {disabled: False, retained: True},
            "maximum_sharing_acknowledgement": "request",
        },
        "m365_capabilities": {disabled: True},
        "enabled_functions": [disabled, retained],
        "maximum_sharing_acknowledgement": "always",
    }
    enabled = operations.get_m365_enabled_function_names(action_type, config)
    normalized = operations.normalize_m365_action_config(action_type, config)
    normalized_again = operations.normalize_m365_action_config(action_type, normalized)
    assert enabled == [retained]
    assert normalized["enabled_functions"] == [retained]
    assert normalized["m365_capabilities"][disabled] is False
    assert normalized["m365_capabilities"][retained] is True
    assert normalized["additionalFields"]["m365_capabilities"][disabled] is False
    assert normalized["maximum_sharing_acknowledgement"] == "request"
    assert normalized["additionalFields"]["maximum_sharing_acknowledgement"] == "request"
    assert normalized_again == normalized
    restricted = operations.get_m365_enabled_function_names(
        action_type, config, agent_capabilities={disabled: True, retained: False},
    )
    assert restricted == []


def test_saved_capability_bounds_are_enforced_on_direct_typed_invocation(execution, graph_plugins):
    modules, _ = graph_plugins
    calendar = modules["m365_calendar_plugin"].M365CalendarPlugin({
        "id": "calendar",
        "additionalFields": {"m365_capabilities": {"get_my_events": False}},
        "m365_capabilities": {"get_my_events": True},
        "enabled_functions": ["get_my_events"],
    })
    remote = Mock(side_effect=AssertionError("A disabled saved capability must not perform remote I/O."))
    calendar._perform_graph_request = remote
    result = calendar.get_my_events()
    assert result["error"] == "function_not_enabled"
    assert calendar.get_functions() == []
    assert remote.call_count == 0

    for action_type in ("m365_onedrive", "m365_sharepoint"):
        file_actions = retrieval.M365FileOperations(action_type, {
            "id": "files",
            "additionalFields": {"m365_capabilities": {"search_files": False}},
            "m365_capabilities": {"search_files": True},
            "enabled_functions": ["search_files"],
        })
        with pytest.raises(M365ProviderError) as denied:
            file_actions.search_files("report")
        assert denied.value.code == "function_not_enabled"


def test_token_uses_the_context_refreshed_by_authoritative_action_policy(execution, monkeypatch):
    import functions_m365_execution as execution_module

    original_context = execution_module.M365ExecutionContext(
        actor_user_id="user-1", data_user_id="user-1", tenant_id="tenant-1",
        conversation_id="conversation-1", request_id="request-1",
        action_configs={"action-1": {
            "type": "m365_onedrive", "source": "onedrive",
            "maximum_sharing_acknowledgement": "always",
        }},
    )
    monkeypatch.setattr(execution_module, "_action_config_resolver", lambda context, action_id, source: {
        "type": "m365_onedrive", "source": "onedrive",
        "maximum_sharing_acknowledgement": "request",
    })
    service = types.SimpleNamespace(authorize_sources=Mock(return_value={
        "onedrive": {"source": "onedrive", "sharing_required": False},
    }))
    monkeypatch.setattr(execution_module, "get_m365_approval_service", lambda: service)
    monkeypatch.setattr(transport_module, "get_m365_context", lambda **kwargs: execution_module.require_m365_execution_context())
    monkeypatch.setattr(transport_module, "authorize_m365_source", REAL_AUTHORIZE_SOURCE)
    token_provider = Mock(return_value={"access_token": "unit-test-token"})
    client = M365Transport(
        "onedrive", "action-1",
        cloud=M365CloudConfig("https://graph.microsoft.com/v1.0", "https://login.microsoftonline.com/tenant-1"),
        token_provider=token_provider,
    )
    with execution_module.m365_execution_context(original_context):
        token, _ = client.get_token(["Files.Read.All"])
        current_context = execution_module.get_m365_execution_context()
    assert token == "unit-test-token"
    assert current_context is not original_context
    assert token_provider.call_args.args[1] is current_context
    assert current_context.action_configs["action-1"]["maximum_sharing_acknowledgement"] == "request"


def test_publication_uses_fresh_authoritative_policy_context(execution, monkeypatch):
    execution.shared = True
    refreshed = types.SimpleNamespace(**vars(execution))
    monkeypatch.setattr(
        retrieval, "authorize_m365_publication",
        lambda source, action_id, policy, **kwargs: (refreshed, {"source": source, "approval_id": "approval-1"}),
    )
    store = Mock()
    store.read_manifest.return_value = {"publication": None}
    store.publish.return_value = {"publication": {"approval_ids": ["approval-1"]}}
    actions = retrieval.M365FileOperations("m365_onedrive", {"id": "action-1"})
    result = actions._publish(store, Mock(), "run-1", execution, None)
    assert result["publication"]["approval_ids"] == ["approval-1"]
    assert store.publish.call_args.kwargs["grant_context"]["execution_context"] is refreshed


def test_workflow_permission_errors_preserve_safe_reconnect_information(execution):
    client = M365Transport(
        "email", "action-1",
        cloud=M365CloudConfig("https://graph.microsoft.com/v1.0", "https://login.microsoftonline.com/tenant-1"),
        token_provider=lambda scopes, context: {
            "error": "m365_consent_required",
            "scopes": ["Mail.Send"], "profile_url": "/profile",
        },
    )
    with pytest.raises(M365ProviderError) as consent:
        client.get_token(["Mail.Read", "Mail.Send"])
    assert consent.value.code == "m365_consent_required"
    assert consent.value.details["scopes"] == ["Mail.Send"]
    assert consent.value.details["profile_url"] == "/profile"
    assert consent.value.message.startswith("Reconnect Microsoft 365 in Profile")


def test_auth_error_metadata_cannot_introduce_foreign_scopes_or_redirects(execution):
    client = M365Transport(
        "email", "action-1",
        cloud=M365CloudConfig("https://graph.microsoft.us/v1.0", "https://login.microsoftonline.us/tenant-1"),
        token_provider=lambda scopes, context: {
            "error": "m365_consent_required",
            "scopes": ["https://foreign.example/Mail.Send"],
            "profile_url": "https://foreign.example/profile?token=secret",
        },
    )
    with pytest.raises(M365ProviderError) as consent:
        client.get_token(["Mail.Send"])
    assert consent.value.details["scopes"] == ["https://graph.microsoft.us/Mail.Send"]
    assert "profile_url" not in consent.value.details
    serialized = json.dumps(consent.value.as_dict())
    assert "foreign.example" not in serialized and "secret" not in serialized


def test_prepared_evidence_exposes_canonical_composite_and_compatible_run_id(execution, memory_runtime, real_content_helpers):
    actions = ContentGraphFixture().operations()
    prepared = actions.prepare_file("drive-1", "item-1")
    assert prepared["memory_run_id"] == prepared["memory_id"]
    assert prepared["evidence_reference"] == f"{prepared['memory_run_id']}:{prepared['evidence_id']}"
    composite = actions.read_file_chunk(prepared["evidence_reference"])
    compatible = actions.read_file_chunk(prepared["memory_id"])
    assert composite["text"] == compatible["text"] == "file content"
    assert composite["evidence_reference"] == compatible["evidence_reference"] == prepared["evidence_reference"]
    assert composite["memory_run_id"] == compatible["memory_run_id"] == prepared["memory_run_id"]
    read = actions.read_file("drive-1", "item-1")
    assert read["window"]["memory_id"] == prepared["evidence_reference"]


@pytest.mark.parametrize("use_default_resolver", [False, True])
def test_budget_commit_crash_recovers_reserved_downloads_without_refund(execution, memory_runtime, real_content_helpers, monkeypatch, use_default_resolver):
    if use_default_resolver:
        monkeypatch.setattr(retrieval, "_request_run_resolver", None)
    fixture = ContentGraphFixture()
    store = memory_runtime.store
    original_commit = store._commit_pending

    def interrupt_reserved_download(ctx, claim, record):
        state = record.get("checkpoint") or {}
        if state.get("kind") == "m365_request_budget" and state.get("download_count") == 1:
            raise MemoryConflictError("Simulated interruption after budget checkpoint upload.")
        return original_commit(ctx, claim, record)

    monkeypatch.setattr(store, "_commit_pending", interrupt_reserved_download)
    with pytest.raises(MemoryConflictError):
        fixture.operations().prepare_file("drive-1", "item-1")
    assert fixture.downloads == 0
    monkeypatch.setattr(store, "_commit_pending", original_commit)
    prepared = fixture.operations().prepare_file("drive-1", "item-1")
    _, ctx = memory_runtime.binding(execution)
    budget_id = retrieval.create_m365_request_budget(store, ctx, execution)
    checkpoint = store.read_checkpoint(ctx, budget_id)
    assert prepared["coverage"]["complete"] is True
    assert fixture.downloads == 1
    assert checkpoint["checkpoint"]["download_count"] == 2


def test_pending_policy_refusal_is_recovered_before_any_further_remote_read(execution, memory_runtime, monkeypatch):
    fixture = GraphFixture(source="spo", licensed=True)
    fixture.copilot_response = FakeResponse({"error": {"code": "blockedByPolicy"}}, status=403)
    actions = retrieval.M365FileOperations("m365_sharepoint", {"id": "action-spo"})
    actions.transport = fixture.transport()
    store = memory_runtime.store
    original_commit = store._commit_pending

    def interrupt_policy_record(ctx, claim, record):
        state = record.get("checkpoint") or {}
        if state.get("source_refusals", {}).get("spo"):
            raise MemoryConflictError("Simulated interruption after policy-refusal checkpoint upload.")
        return original_commit(ctx, claim, record)

    monkeypatch.setattr(store, "_commit_pending", interrupt_policy_record)
    with pytest.raises(MemoryConflictError):
        actions.search_files("report")
    calls_after_refusal = len(fixture.calls)
    monkeypatch.setattr(store, "_commit_pending", original_commit)
    with pytest.raises(M365ProviderError) as refused:
        actions.prepare_file("drive-1", "item-1")
    assert refused.value.code == "source_policy_blocked"
    assert len(fixture.calls) == calls_after_refusal


def test_incomplete_budget_checkpoint_does_not_reset_counters_or_resume_remote_calls(execution, memory_runtime, monkeypatch):
    from functions_conversation_memory import MemoryIncompleteCaptureError, MemoryUnavailableError

    fixture = ContentGraphFixture()
    blob = memory_runtime.blob
    original_put = blob.put

    def interrupt_budget_upload(container, name, data, *, etag=None):
        record = json.loads(data)
        state = record.get("checkpoint") or {}
        if record.get("kind") == "memory_checkpoint" and state.get("download_count") == 1:
            raise MemoryUnavailableError("Simulated interruption before budget checkpoint upload.")
        return original_put(container, name, data, etag=etag)

    monkeypatch.setattr(blob, "put", interrupt_budget_upload)
    with pytest.raises(MemoryUnavailableError):
        fixture.operations().prepare_file("drive-1", "item-1")
    calls_before_resume = len(fixture.calls)
    monkeypatch.setattr(blob, "put", original_put)
    with pytest.raises(MemoryIncompleteCaptureError):
        fixture.operations().prepare_file("drive-1", "item-1")
    assert fixture.downloads == 0
    assert len(fixture.calls) == calls_before_resume


def test_lost_file_run_creation_acknowledgement_reuses_original_keyed_capture(execution, memory_runtime, real_content_helpers, monkeypatch):
    from functions_conversation_memory import MemoryUnavailableError

    fixture = ContentGraphFixture()
    store = memory_runtime.store
    original_create = store.get_or_create_manifest
    created_file_ids = []

    def lose_first_file_ack(ctx, **kwargs):
        run = original_create(ctx, **kwargs)
        if kwargs.get("kind") == "m365_file_onedrive":
            created_file_ids.append(run["run_id"])
            if len(created_file_ids) == 1:
                raise MemoryUnavailableError("Simulated lost file-manifest creation acknowledgement.")
        return run

    monkeypatch.setattr(store, "get_or_create_manifest", lose_first_file_ack)
    with pytest.raises(MemoryUnavailableError):
        fixture.operations().prepare_file("drive-1", "item-1")
    assert fixture.downloads == 0
    prepared = fixture.operations().prepare_file("drive-1", "item-1")
    assert len(created_file_ids) == 2
    assert created_file_ids[0] == created_file_ids[1] == prepared["memory_run_id"]
    assert fixture.downloads == 1
    _, ctx = memory_runtime.binding(execution)
    runs = store.list_runs(ctx)
    captures = [run for run in runs["runs"] if run["purpose"] == "m365_file_onedrive"]
    assert len(captures) == 1 and captures[0]["status"] == "completed"


def test_request_budget_checkpoint_reads_always_hold_a_worker_claim(execution, memory_runtime, real_content_helpers, monkeypatch):
    monkeypatch.setattr(retrieval, "_request_run_resolver", None)
    store = memory_runtime.store
    original_read = store.read_checkpoint
    observed_statuses = []

    def claimed_read(ctx, run_id, **kwargs):
        manifest = store.read_manifest(ctx, run_id)
        if manifest["purpose"] == "m365_request_budget":
            observed_statuses.append(manifest["status"])
            if manifest["status"] != "running":
                raise AssertionError("Read the budget checkpoint only after claiming its run.")
        return original_read(ctx, run_id, **kwargs)

    monkeypatch.setattr(store, "read_checkpoint", claimed_read)
    fixture = ContentGraphFixture()
    first = fixture.operations().prepare_file("drive-1", "item-1")
    second = fixture.operations().prepare_file("drive-1", "item-1")
    assert first["memory_id"] == second["memory_id"]
    assert observed_statuses and set(observed_statuses) == {"running"}
    assert fixture.downloads == 1


@pytest.mark.parametrize("action_type", ["m365_onedrive", "m365_sharepoint"])
@pytest.mark.parametrize("snapshot_function", ["read_file_chunk", "analyze_file"])
def test_snapshot_only_capabilities_do_not_require_fresh_source_access(action_type, snapshot_function):
    config = {"enabled_functions": [snapshot_function]}
    enabled = operations.get_m365_enabled_function_names(action_type, config)
    remote = operations.get_m365_remote_function_names(action_type, config)
    definitions = operations.get_m365_function_definitions(action_type)
    selected = next(item for item in definitions if item["name"] == snapshot_function)
    assert enabled == [snapshot_function]
    assert remote == []
    assert selected["requires_remote_access"] is False
    disabled = operations.get_m365_remote_function_names(action_type, enabled_functions=[])
    assert disabled == []


def test_remote_capability_metadata_is_central_and_honors_all_restrictions():
    for action_type in operations.M365_ACTION_TYPES:
        definitions = operations.get_m365_action_definition(action_type)["capabilities"]
        remote = operations.get_m365_remote_function_names(action_type)
        defaults = operations.get_m365_default_capabilities(action_type)
        expected = [
            item["function_name"] for item in definitions
            if item["requires_remote_access"] and defaults[item["key"]]
        ]
        assert remote == expected
        assert all(type(item["requires_remote_access"]) is bool for item in definitions)
    config = {
        "additionalFields": {"m365_capabilities": {"search_files": False}},
        "m365_capabilities": {"search_files": True},
        "enabled_functions": ["search_files", "read_file", "read_file_chunk"],
        "metadata": {"methods": [{"name": "read_file", "requires_remote_access": False}]},
    }
    remote = operations.get_m365_remote_function_names("m365_onedrive", config)
    restricted = operations.get_m365_remote_function_names(
        "m365_onedrive", config, agent_capabilities={"read_file": False},
    )
    assert remote == ["read_file"]
    assert restricted == []


@pytest.mark.parametrize("action_type,function_name,arguments", [
    ("m365_calendar", "get_my_events", ()),
    ("m365_email", "get_my_messages", ()),
    ("m365_onedrive", "search_files", ("report",)),
    ("m365_sharepoint", "search_files", ("report",)),
    ("msgraph", "get_my_messages", ()),
])
def test_cached_plugins_check_current_saved_capabilities_before_remote_access(
    graph_plugins, monkeypatch, action_type, function_name, arguments,
):
    import functions_m365_execution as execution_module
    from semantic_kernel_plugins.m365_onedrive_plugin import M365OneDrivePlugin
    from semantic_kernel_plugins.m365_sharepoint_plugin import M365SharePointPlugin

    modules, _ = graph_plugins
    classes = {
        "m365_calendar": modules["m365_calendar_plugin"].M365CalendarPlugin,
        "m365_email": modules["m365_email_plugin"].M365EmailPlugin,
        "m365_onedrive": M365OneDrivePlugin,
        "m365_sharepoint": M365SharePointPlugin,
        "msgraph": modules["msgraph_plugin"].MSGraphPlugin,
    }
    capability_field = "msgraph_capabilities" if action_type == "msgraph" else "m365_capabilities"
    manifest = {
        "id": "action-1", "type": action_type,
        "additionalFields": {capability_field: {function_name: True}},
        capability_field: {function_name: True},
    }
    plugin = classes[action_type](deepcopy(manifest))
    context = execution_module.M365ExecutionContext(
        actor_user_id="user-1", data_user_id="user-1", tenant_id="tenant-1",
        conversation_id="conversation-1", request_id="request-1",
        action_configs={"action-1": manifest},
    )
    current = deepcopy(manifest)
    current["additionalFields"][capability_field][function_name] = False
    current[capability_field][function_name] = False
    monkeypatch.setattr(execution_module, "_action_config_resolver", lambda *args: current)
    source_approval = Mock(side_effect=AssertionError("A disabled function must not request source approval."))
    monkeypatch.setattr(
        execution_module, "get_m365_approval_service",
        lambda: types.SimpleNamespace(authorize_sources=source_approval),
    )
    monkeypatch.setattr(transport_module, "get_m365_context", lambda **kwargs: execution_module.require_m365_execution_context())
    for module in (transport_module, retrieval, modules["msgraph_plugin"]):
        monkeypatch.setattr(module, "authorize_m365_source", REAL_AUTHORIZE_SOURCE)
    remote = Mock(side_effect=AssertionError("A stale enabled function must not access the source."))
    monkeypatch.setattr(M365Transport, "request_json", remote)
    monkeypatch.setattr(M365Transport, "get_token", remote)
    with execution_module.m365_execution_context(context):
        result = getattr(plugin, function_name)(*arguments)
    error = result["error"]
    code = error["code"] if isinstance(error, dict) else error
    assert code in {"m365_function_not_authorized", "m365_source_not_authorized"}
    assert source_approval.call_count == 0 and remote.call_count == 0


def test_current_capabilities_are_rechecked_between_file_provider_requests(
    execution, memory_runtime, monkeypatch,
):
    import functions_m365_execution as execution_module

    manifest = {
        "id": "action-1", "type": "m365_onedrive",
        "additionalFields": {"m365_capabilities": {"search_files": True}},
    }
    context = execution_module.M365ExecutionContext(
        actor_user_id="user-1", data_user_id="user-1", tenant_id="tenant-1",
        conversation_id="conversation-1", request_id="request-1",
        action_configs={"action-1": manifest},
    )
    current = deepcopy(manifest)
    monkeypatch.setattr(execution_module, "_action_config_resolver", lambda *args: current)
    monkeypatch.setattr(execution_module, "get_m365_approval_service", lambda: types.SimpleNamespace(
        authorize_sources=lambda context, policies: {
            source: {"source": source, "sharing_required": False} for source in policies
        },
    ))
    monkeypatch.setattr(transport_module, "get_m365_context", lambda **kwargs: execution_module.require_m365_execution_context())
    monkeypatch.setattr(retrieval, "get_m365_context", lambda **kwargs: execution_module.require_m365_execution_context())
    monkeypatch.setattr(transport_module, "authorize_m365_source", REAL_AUTHORIZE_SOURCE)
    monkeypatch.setattr(retrieval, "authorize_m365_source", REAL_AUTHORIZE_SOURCE)
    requests_seen = []

    def request(method, url, **kwargs):
        requests_seen.append(url)
        current["additionalFields"]["m365_capabilities"]["search_files"] = False
        return FakeResponse({"id": "user-1", "assignedLicenses": [], "assignedPlans": []})

    actions = retrieval.M365FileOperations("m365_onedrive", manifest)
    actions.transport = M365Transport(
        "onedrive", "action-1", action_type="m365_onedrive",
        cloud=M365CloudConfig("https://graph.microsoft.com/v1.0", "https://login.microsoftonline.com/tenant-1"),
        request=request, token_provider=lambda scopes, context: {"access_token": "unit-test-token"},
    )
    with execution_module.m365_execution_context(context):
        with pytest.raises(retrieval.M365PolicyError) as denied:
            actions.search_files("report")
    assert denied.value.code == "m365_function_not_authorized"
    assert requests_seen == ["https://graph.microsoft.com/v1.0/me"]
    assert actions.transport.current_operation is None


def test_transport_operation_identity_is_scoped_between_calls(execution, monkeypatch):
    observed = []

    def authorize(source, action_id, policy, **kwargs):
        observed.append(kwargs.get("operation_name"))
        return execution, {"source": source}

    monkeypatch.setattr(transport_module, "authorize_m365_source", authorize)
    client = GraphFixture().transport()
    other_context = copy_context()
    with client.operation_context("search_files"):
        client.get_token(["Files.Read.All"])

        def other_operation():
            assert client.current_operation is None
            with client.operation_context("prepare_file"):
                client.get_token(["Files.Read.All"])

        other_context.run(other_operation)
        client.get_token(["Files.Read.All"])
    assert observed == ["search_files", "prepare_file", "search_files"]
    assert client.current_operation is None


@pytest.mark.parametrize("action_type,function_name,arguments", [
    ("msgraph", "get_my_profile", ()),
    ("msgraph", "search_users", ("Ada",)),
    ("msgraph", "get_my_security_alerts", ()),
    ("m365_onedrive", "read_file_chunk", ("0" * 32,)),
    ("m365_sharepoint", "read_file_chunk", ("0" * 32,)),
])
def test_cached_nonremote_guards_revalidate_capabilities_without_source_consent(
    graph_plugins, monkeypatch, action_type, function_name, arguments,
):
    import functions_m365_execution as execution_module
    from semantic_kernel_plugins.m365_onedrive_plugin import M365OneDrivePlugin
    from semantic_kernel_plugins.m365_sharepoint_plugin import M365SharePointPlugin

    modules, _ = graph_plugins
    classes = {
        "msgraph": modules["msgraph_plugin"].MSGraphPlugin,
        "m365_onedrive": M365OneDrivePlugin,
        "m365_sharepoint": M365SharePointPlugin,
    }
    field = "msgraph_capabilities" if action_type == "msgraph" else "m365_capabilities"
    manifest = {
        "id": "action-1", "type": action_type,
        "additionalFields": {field: {function_name: True}},
        field: {function_name: True},
    }
    plugin = classes[action_type](deepcopy(manifest))
    context = execution_module.M365ExecutionContext(
        actor_user_id="user-1", data_user_id="user-1", tenant_id="tenant-1",
        conversation_id="conversation-1", request_id="request-1",
        action_configs={"action-1": manifest},
    )
    current = deepcopy(manifest)
    current["additionalFields"][field][function_name] = False
    current[field][function_name] = False
    monkeypatch.setattr(execution_module, "_action_config_resolver", lambda *args: current)
    source_access = Mock(side_effect=AssertionError("Capability refresh must not acquire source consent or credentials."))
    monkeypatch.setattr(execution_module, "get_m365_approval_service", source_access)
    monkeypatch.setattr(transport_module, "get_m365_context", lambda **kwargs: execution_module.get_m365_execution_context())
    monkeypatch.setattr(retrieval, "get_m365_context", lambda **kwargs: execution_module.get_m365_execution_context())
    for module in (transport_module, retrieval, modules["msgraph_plugin"]):
        monkeypatch.setattr(module, "authorize_m365_capability", REAL_AUTHORIZE_CAPABILITY)
    monkeypatch.setattr(M365Transport, "request_json", source_access)
    monkeypatch.setattr(M365Transport, "get_token", source_access)
    with execution_module.m365_execution_context(context):
        result = getattr(plugin, function_name)(*arguments)
    error = result["error"]
    code = error["code"] if isinstance(error, dict) else error
    assert code == "m365_function_not_authorized"
    assert source_access.call_count == 0


def test_live_capability_revocation_stops_a_secret_download_redirect(execution, tmp_path, monkeypatch):
    import functions_m365_execution as execution_module

    manifest = {
        "id": "action-1", "type": "m365_onedrive",
        "additionalFields": {"m365_capabilities": {"prepare_file": True}},
    }
    context = execution_module.M365ExecutionContext(
        actor_user_id="user-1", data_user_id="user-1", tenant_id="tenant-1",
        conversation_id="conversation-1", request_id="request-1",
        action_configs={"action-1": manifest},
    )
    current = deepcopy(manifest)
    monkeypatch.setattr(execution_module, "_action_config_resolver", lambda *args: current)
    monkeypatch.setattr(execution_module, "get_m365_approval_service", lambda: types.SimpleNamespace(
        authorize_sources=lambda context, policies: {
            source: {"source": source, "sharing_required": False} for source in policies
        },
    ))
    monkeypatch.setattr(transport_module, "get_m365_context", lambda **kwargs: execution_module.require_m365_execution_context())
    monkeypatch.setattr(transport_module, "authorize_m365_source", REAL_AUTHORIZE_SOURCE)
    monkeypatch.setattr(transport_module.tempfile, "tempdir", str(tmp_path))
    requests_seen = []

    def request(method, url, **kwargs):
        requests_seen.append(url)
        current["additionalFields"]["m365_capabilities"]["prepare_file"] = False
        return FakeResponse(status=302, headers={
            "Location": "https://tenant.sharepoint.com/_layouts/download.aspx?tempauth=secret",
        })

    client = M365Transport(
        "onedrive", "action-1", action_type="m365_onedrive",
        cloud=M365CloudConfig("https://graph.microsoft.com/v1.0", "https://login.microsoftonline.com/tenant-1"),
        request=request, token_provider=lambda scopes, context: {"access_token": "unit-test-token"},
    )
    with execution_module.m365_execution_context(context), client.operation_context("prepare_file"):
        with pytest.raises(retrieval.M365PolicyError) as denied:
            with client.download_file(
                "drive-1", "item-1", suffix=".txt", allowed_mime_types=("text/plain",), max_bytes=100,
            ):
                pytest.fail("A revoked file capability must not follow a download URL.")
    assert denied.value.code == "m365_function_not_authorized"
    assert requests_seen == ["https://graph.microsoft.com/v1.0/drives/drive-1/items/item-1/content"]
    remaining = list(tmp_path.iterdir())
    assert remaining == []


def test_cached_mail_capability_is_rechecked_before_each_page(graph_plugins, monkeypatch):
    import functions_m365_execution as execution_module

    modules, _ = graph_plugins
    manifest = {
        "id": "action-1", "type": "m365_email",
        "additionalFields": {"m365_capabilities": {"get_my_messages": True}},
    }
    context = execution_module.M365ExecutionContext(
        actor_user_id="user-1", data_user_id="user-1", tenant_id="tenant-1",
        conversation_id="conversation-1", request_id="request-1",
        action_configs={"action-1": manifest},
    )
    current = deepcopy(manifest)
    monkeypatch.setattr(execution_module, "_action_config_resolver", lambda *args: current)
    monkeypatch.setattr(execution_module, "get_m365_approval_service", lambda: types.SimpleNamespace(
        authorize_sources=lambda context, policies: {
            source: {"source": source, "sharing_required": False} for source in policies
        },
    ))
    monkeypatch.setattr(transport_module, "get_m365_context", lambda **kwargs: execution_module.require_m365_execution_context())
    monkeypatch.setattr(transport_module, "authorize_m365_source", REAL_AUTHORIZE_SOURCE)
    monkeypatch.setattr(modules["msgraph_plugin"], "authorize_m365_source", REAL_AUTHORIZE_SOURCE)
    requests_seen = []

    def request(method, url, **kwargs):
        requests_seen.append(url)
        current["additionalFields"]["m365_capabilities"]["get_my_messages"] = False
        return FakeResponse({
            "value": [{"id": "message-1"}],
            "@odata.nextLink": "https://graph.microsoft.com/v1.0/me/messages?page=2",
        })

    plugin = modules["m365_email_plugin"].M365EmailPlugin(manifest)
    plugin._transports["email"] = M365Transport(
        "email", "action-1", action_type="m365_email",
        cloud=M365CloudConfig("https://graph.microsoft.com/v1.0", "https://login.microsoftonline.com/tenant-1"),
        request=request, token_provider=lambda scopes, context: {"access_token": "unit-test-token"},
    )
    with execution_module.m365_execution_context(context):
        result = plugin.get_my_messages(top=3)
    assert result["error"] in {"m365_function_not_authorized", "m365_source_not_authorized"}
    assert len(requests_seen) == 1


def test_live_snapshot_capability_gate_keeps_published_reads_graph_free(
    execution, memory_runtime, real_content_helpers, monkeypatch,
):
    import functions_m365_execution as execution_module

    execution.shared = True
    fixture = ContentGraphFixture()
    actions = fixture.operations()
    prepared = actions.prepare_file("drive-1", "item-1")
    manifest = {
        "id": "action-onedrive", "type": "m365_onedrive",
        "enabled_functions": ["read_file_chunk"],
        "additionalFields": {"m365_capabilities": {"read_file_chunk": True}},
    }
    viewer = execution_module.M365ExecutionContext(
        actor_user_id="user-2", data_user_id="user-2", tenant_id="tenant-1",
        conversation_id="conversation-1", request_id="viewer-request",
        shared=True, audience_version="audience-1",
        action_configs={"action-onedrive": manifest},
    )
    monkeypatch.setattr(execution_module, "_action_config_resolver", lambda *args: manifest)
    forbidden = Mock(side_effect=AssertionError("Published snapshots do not need source approval, tokens, or remote reads."))
    monkeypatch.setattr(execution_module, "get_m365_approval_service", forbidden)
    monkeypatch.setattr(transport_module, "get_m365_context", lambda **kwargs: execution_module.get_m365_execution_context())
    monkeypatch.setattr(retrieval, "get_m365_context", lambda **kwargs: execution_module.get_m365_execution_context())
    monkeypatch.setattr(retrieval, "authorize_m365_capability", REAL_AUTHORIZE_CAPABILITY)
    monkeypatch.setattr(retrieval, "authorize_m365_source", forbidden)
    monkeypatch.setattr(M365Transport, "get_token", forbidden)
    monkeypatch.setattr(M365Transport, "request_json", forbidden)
    before = len(fixture.calls)
    with execution_module.m365_execution_context(viewer):
        window = actions.read_file_chunk(prepared["evidence_reference"])
    assert window["text"] == "file content"
    assert window["snapshot_state"] == "published_snapshot"
    assert window["capture"]["principal_id"] == "user-1"
    assert len(fixture.calls) == before and forbidden.call_count == 0


def test_live_legacy_profile_capability_gate_does_not_require_source_sharing(graph_plugins, monkeypatch):
    import functions_m365_execution as execution_module

    modules, _ = graph_plugins
    manifest = {
        "id": "legacy", "type": "msgraph", "enabled_functions": ["get_my_profile"],
        "additionalFields": {"msgraph_capabilities": {"get_my_profile": True}},
    }
    context = execution_module.M365ExecutionContext(
        actor_user_id="user-1", data_user_id="user-1", tenant_id="tenant-1",
        conversation_id="conversation-1", request_id="request-1",
        action_configs={"legacy": manifest},
    )
    monkeypatch.setattr(execution_module, "_action_config_resolver", lambda *args: manifest)
    source_approval = Mock(side_effect=AssertionError("Legacy profile reads do not acquire a file/mail/calendar sharing grant."))
    monkeypatch.setattr(execution_module, "get_m365_approval_service", source_approval)
    monkeypatch.setattr(transport_module, "get_m365_context", lambda **kwargs: execution_module.require_m365_execution_context())
    monkeypatch.setattr(transport_module, "authorize_m365_capability", REAL_AUTHORIZE_CAPABILITY)
    monkeypatch.setattr(modules["msgraph_plugin"], "authorize_m365_capability", REAL_AUTHORIZE_CAPABILITY)
    token_provider = Mock(return_value={"access_token": "unit-test-token"})
    request = Mock(return_value=FakeResponse({"id": "user-1"}))
    plugin = modules["msgraph_plugin"].MSGraphPlugin(manifest)
    plugin._transports[None] = M365Transport(
        None, "legacy", action_type="msgraph",
        cloud=M365CloudConfig("https://graph.microsoft.com/v1.0", "https://login.microsoftonline.com/tenant-1"),
        request=request, token_provider=token_provider,
    )
    with execution_module.m365_execution_context(context):
        result = plugin.get_my_profile()
    assert result["id"] == "user-1"
    assert source_approval.call_count == 0
    assert request.call_count == 1
    assert token_provider.call_args.args[0] == ["https://graph.microsoft.com/User.Read"]


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-q"]))
