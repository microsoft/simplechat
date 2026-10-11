# test_document_transform_harness.py
"""Real chat publication and next-round native document transformation.

Version: 0.261.321
Implemented in: 0.261.321

Exercises bootstrap, retained results, native rendering, committed artifacts and
model-free delivery with external Azure and model I/O doubled.
"""

import importlib
import io
import json
from copy import deepcopy

import pytest
from docx import Document

from functions_document_transform import inspect_document
from functions_orchestration_registry import resolve_available_capability_ids
from functions_orchestration_result_contracts import ResultContractError
from test_orchestration_harness_execution import harness, initialized_application  # noqa: F401
from test_support.orchestration_harness_execution import decoded_frames, input_binding, render_step


PROFILE = "document_transform_v1"
HANDLE = {"kind": "document", "document_id": "reference"}


def document_bytes(text="Reference example"):
    document = Document()
    document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def pipeline(handle, mode):
    return [
        {
            "step_id": "inspect", "capability_id": "document_inspect",
            "arguments": {"reference": deepcopy(handle)},
        },
        {
            "step_id": "update",
            "capability_id": "document_refresh" if mode == "refresh" else "document_create_from_reference",
            "arguments": {"instruction": "Use the supplied new content and disclose missing facts."},
            "inputs": {"inspection": {"binding": input_binding("inspect", "inspection"), "allow_partial": False}},
        },
        render_step("draft", "docx", source="update", output="transformation", profile=PROFILE),
    ]


def preparation(content, handle, text, mode="refresh"):
    inspection = inspect_document(content, "reference.docx", handle)
    return {
        "schema_version": PROFILE, "output_format": "docx", "mode": mode,
        "reference": inspection["reference"],
        "operations": [
            {"op": "replace_text", "target": "word:body.p1", "text": text},
            {"op": "replace_text", "target": "unsupported:chart", "text": "An unsupported chart update"},
        ],
        "unresolved": [{"target": "missing:figure", "reason": "The latest figure was not provided."}],
    }


def bind_reference(harness, monkeypatch, content):
    original = harness.bootstrap.build_document_transform_sources
    source = {
        "id": "reference",
        "document_id": "reference", "scope": "personal", "scope_id": "owner",
        "authorization_status": "authorized", "file_name": "reference.docx", "source_kind": "narrative",
        "source_version": 1, "source_revision": "reference-revision",
    }
    reads = []
    document = {
        "id": "reference", "user_id": "owner", "file_name": "reference.docx",
        "version": 1, "_etag": "reference-revision",
    }

    def read_document(document_id, user_id, **kwargs):
        if document_id != "reference" or user_id != "owner" or kwargs.get("group_id") or kwargs.get("public_workspace_id"):
            raise PermissionError("The fixture reference belongs to the exact actor and personal scope.")
        return deepcopy(document)

    def document_contexts(**kwargs):
        if kwargs["user_id"] != "owner" or kwargs.get("doc_scope", "all") not in ("all", "personal"):
            raise PermissionError("The fixture reference belongs to the exact actor and personal scope.")
        return [
            {"document": deepcopy(document), "scope": "personal"}
            if identifier == "reference" else None
            for identifier in kwargs["document_ids"]
        ]

    def reader(handle):
        if handle != HANDLE:
            raise PermissionError("The fixture only admits the exact selected document.")
        reads.append(deepcopy(handle))
        return content

    def metadata(handle):
        if handle != HANDLE:
            raise PermissionError("The fixture only admits the exact selected document.")
        return {"file_name": "reference.docx", "source": deepcopy(source)}

    def sources(user_id, conversation_id):
        # The next-round artifact uses real conversation/file authorization and byte readers.
        bound = original(user_id, conversation_id)
        artifact_reader, artifact_metadata = bound.read_reference, bound.describe_reference
        bound.read_reference = lambda handle: reader(handle) if handle["kind"] == "document" else artifact_reader(handle)
        bound.describe_reference = lambda handle: metadata(handle) if handle["kind"] == "document" else artifact_metadata(handle)
        return bound

    monkeypatch.setattr(harness.bootstrap, "build_document_transform_sources", sources)
    monkeypatch.setattr(
        harness.bootstrap, "resolve_orchestration_source_manifest",
        lambda identifiers, **kwargs: [deepcopy(source) for identifier in identifiers if identifier == "reference"],
    )
    monkeypatch.setattr(
        harness.bootstrap, "read_orchestration_source_metadata",
        lambda identifier, user_id, **kwargs: deepcopy(source) if identifier == "reference" and user_id == "owner" else None,
    )
    # Source ownership modules load only after the offline application is initialized.
    mixed = importlib.import_module("functions_mixed_source_orchestration")
    screening = importlib.import_module("content_screening.access")
    monkeypatch.setattr(mixed, "_default_document_context_batch_resolver", document_contexts)
    monkeypatch.setattr(screening, "_read_authorized_document", read_document)
    return reads


@pytest.mark.parametrize("mode", ["refresh", "create_from_reference"])
@pytest.mark.parametrize("delivery_only", [False, True])
def test_chat_delivery_persists_actual_omissions_and_model_free_recovery(harness, monkeypatch, mode, delivery_only):
    content = document_bytes()
    reads = bind_reference(harness, monkeypatch, content)
    value = preparation(content, HANDLE, "Updated content", mode)
    harness.create(pipeline(HANDLE, mode), replies=[json.dumps({"transformation": value})])
    execution = harness.prepare()
    resolved_content = execution.context.document_transform_source_reader(HANDLE)
    assert resolved_content == content
    if delivery_only:
        try:
            result = harness.run_engine(execution)
            assert result["status"] == "completed", result
            services = harness.services()
            with harness.publication_only(services):
                frames = harness.execution.refresh_harness_delivery(
                    harness.read(), services=services, settings=harness.settings, lease=execution.lease,
                )
        finally:
            execution.close()
    else:
        frames = execution.execute()
    done = decoded_frames(frames)[-1]
    assert done["status"] == "completed", done
    services = harness.services()
    outputs = services.rendering.list_public_outputs("run-1")
    messages = harness.assistant_messages()
    report = outputs[0]["transformation_report"]
    assert "1 skipped" in done["full_content"] and "1 unresolved" in done["full_content"]
    assert "The latest figure was not provided" in done["full_content"]
    assert report["applied_count"] == 1 and report["skipped_count"] == report["unresolved_count"] == 1
    assert outputs == done["outputs"] == messages[0]["metadata"]["orchestration"]["outputs"]
    assert len(harness.model_calls) == harness.blobs.file_uploads == 1 and len(reads) >= 2
    assert "source_ref" not in json.dumps(done) and "content_sha256" not in json.dumps(outputs)


def test_follow_up_transforms_the_committed_file_in_the_same_conversation(harness, monkeypatch):
    original = document_bytes()
    bind_reference(harness, monkeypatch, original)
    first = preparation(original, HANDLE, "First revision")
    harness.create(pipeline(HANDLE, "refresh"), replies=[json.dumps({"transformation": first})])
    done = decoded_frames(harness.prepare().execute())[-1]
    assert done["status"] == "completed", done
    first_artifact = done["generated_artifacts"][0]
    handle = {
        "kind": "artifact", "conversation_id": "conversation-1",
        "message_id": first_artifact["artifact_message_id"],
    }
    services = harness.services()
    with services.rendering.open_download(first_artifact["output_id"]) as stream:
        first_bytes = stream.read()
    next_value = preparation(first_bytes, handle, "Second revision")
    harness.turn = {
        **harness.turn, "id": "user-turn-2", "content": "Update the file you just generated.",
    }
    harness.messages.create_item(harness.turn)
    harness.create(
        pipeline(handle, "refresh"), replies=[json.dumps({"transformation": next_value})],
        run_id="run-2", plan_id="plan-2", turn_id="turn-2",
        seeds={"document_transform_references": [{"handle": handle, "file_name": "draft.docx"}]},
    )
    second = decoded_frames(harness.prepare().execute())[-1]
    assert second["status"] == "completed", second
    second_artifact = second["generated_artifacts"][0]
    services = harness.services()
    with services.rendering.open_download(second_artifact["output_id"]) as stream:
        second_bytes = stream.read()
    with services.rendering.open_download(first_artifact["output_id"]) as stream:
        unchanged_bytes = stream.read()
    second_text = Document(io.BytesIO(second_bytes)).paragraphs[0].text
    first_text = Document(io.BytesIO(first_bytes)).paragraphs[0].text
    assert second_text == "Second revision"
    assert first_text == "First revision"
    assert unchanged_bytes == first_bytes
    assert second_artifact["artifact_message_id"] != first_artifact["artifact_message_id"]
    assert len(harness.model_calls) == harness.blobs.file_uploads == 2


@pytest.mark.parametrize("missing", ["document_transform_source_reader", "document_transform_source_resolver"])
def test_services_reject_incomplete_native_reference_bindings(harness, monkeypatch, missing):
    # These service helpers import only after the offline application owners initialize.
    from test_orchestration_services import services
    from test_support.orchestration_results import ResultFixture

    options = {
        "document_transform_source_reader": lambda handle: b"",
        "document_transform_source_resolver": lambda handle: {},
    }
    options[missing] = None
    with pytest.raises(ResultContractError) as raised:
        services(ResultFixture(), **options)
    assert raised.value.code == "result_native_reader_required"


def test_revision_refreshes_artifact_admission_without_mutating_saved_context(harness):
    editing = importlib.import_module("functions_orchestration_plan_editing")
    handle = {"kind": "artifact", "conversation_id": "conversation-1", "message_id": "prior-file"}
    candidates = [{"handle": handle, "file_name": "prior.docx"}]
    record = harness.create(pipeline(handle, "refresh"), seeds={"document_transform_references": candidates})
    context = editing._turn_context(record)
    original = deepcopy(context)
    services = harness.services()
    options = {
        **services.capability_request_bindings(),
        "export_catalog": services.export_catalog(),
        "composition_profiles": harness.service_bindings.composition_profiles(),
    }
    selected = []

    def fresh_candidates(document_ids):
        selected.append(document_ids)
        return deepcopy(candidates)

    checked = editing.validate_edited_plan(
        record["plan"], context, "owner", harness.settings, {},
        document_transform_candidate_reader=fresh_candidates, **options,
    )
    assert checked["steps"][0]["arguments"]["reference"] == handle
    with pytest.raises(editing.PlanRevisionError):
        editing.validate_edited_plan(
            record["plan"], context, "owner", harness.settings, {},
            document_transform_candidate_reader=lambda identifiers: [], **options,
        )
    assert context == original
    assert selected == [[]] and harness.model_calls == [] and harness.blobs.file_uploads == 0


def test_replanning_persists_fresh_reference_candidates_when_old_seeds_are_empty(harness, monkeypatch):
    editing = importlib.import_module("functions_orchestration_plan_editing")
    handle = {"kind": "artifact", "conversation_id": "conversation-1", "message_id": "prior-file"}
    candidates = [{"handle": handle, "file_name": "prior.docx"}]
    record = harness.create()
    before = deepcopy(record)
    services = harness.services()
    observed = []

    def replan(message, context, conversation_id, user_id, **options):
        observed.append(deepcopy(context["document_transform_references"]))
        available = resolve_available_capability_ids(
            harness.settings, request_context=options["request_context"], include_runtime_bindings=True,
            export_catalog=options["export_catalog"],
        )
        checked = harness.schema.normalize_plan(
            {"steps": pipeline(handle, "refresh")}, conversation_id, user_id,
            settings=harness.settings, contract_version=2,
            seeds=options["seeds"], composition_profiles=options["composition_profiles"],
            available_capability_ids=available,
        )
        checked["revised_request"] = message
        return "plan", checked

    monkeypatch.setattr(editing, "plan_request", replan)
    result = editing.build_plan_edit_outcome(
        record, {"action": "ask", "instruction": "Reuse the previous Word file."},
        "owner", harness.settings, identity={}, conversation=harness.conversation,
        conversation_context=record["conversation_context"],
        document_transform_candidate_reader=lambda identifiers: deepcopy(candidates),
        export_catalog=services.export_catalog(),
        composition_profiles=harness.service_bindings.composition_profiles(),
        **services.capability_request_bindings(),
    )
    assert result["kind"] == "plan"
    assert observed == [candidates]
    assert result["turn_context"]["seeds"]["document_transform_references"] == candidates
    assert record == before and harness.model_calls == [] and harness.blobs.file_uploads == 0


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-q", "--tb=short"]))
