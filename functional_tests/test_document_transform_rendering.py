# test_document_transform_rendering.py
"""
Functional tests for retained reference-document rendering and actual delivery reports.
Version: 0.261.321
Implemented in: 0.261.321

Real Office packages, transformation core, result lineage, output store and transport
run without model/provider calls. A scoped test adapter isolates renderer contracts
from the independently tested Word/PowerPoint editing implementations.
"""

import hashlib
import importlib
import io
import json
import sys
import zipfile
from contextlib import contextmanager
from copy import deepcopy
from dataclasses import replace
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

import pytest
from azure.core.exceptions import ResourceNotFoundError
from docx import Document
from pptx import Presentation

from test_generated_file_structured_renderers import TextSource, ValueSource
from test_support.orchestration_results import ResultFixture, complete, source
from test_support.orchestration_revisions import AtomicMemoryContainer


APP = Path(__file__).resolve().parents[1] / 'application' / 'single_app'
with patch.object(sys, 'path', [str(APP), *sys.path]):
    core = importlib.import_module('functions_document_transform')
    exports = importlib.import_module('functions_generated_file_exports')
    bridge = importlib.import_module('functions_generated_office_adapters')
    registry = importlib.import_module('functions_generated_export_registry')
    rendering = importlib.import_module('functions_orchestration_rendering')
    output_store = importlib.import_module('functions_orchestration_output_store')
    artifacts = importlib.import_module('functions_orchestration_artifacts')
    deliverables = importlib.import_module('functions_orchestration_deliverables')
    results_module = importlib.import_module('functions_orchestration_results')

PROFILE = 'document_transform_v1'
HANDLE = {'kind': 'document', 'document_id': 'document-1'}


def package_bytes(output_format='docx', *, template=False):
    stream = io.BytesIO()
    if output_format == 'docx':
        document = Document()
        document.add_paragraph('Original')
        document.save(stream)
    else:
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[0])
        slide.shapes.title.text = 'Original'
        document.save(stream)
    content = stream.getvalue()
    if template:
        content = rewrite_part(
            content, '[Content_Types].xml',
            lambda data: data.replace(
                core.MEDIA_TYPES[output_format].encode() + b'.main+xml',
                (
                    'application/vnd.openxmlformats-officedocument.'
                    + ('wordprocessingml' if output_format == 'docx' else 'presentationml')
                    + '.template.main+xml'
                ).encode(),
            ),
        )
    return content


def rewrite_part(content, name, mutate):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as original, zipfile.ZipFile(output, 'w') as updated:
        for entry in original.infolist():
            data = original.read(entry)
            updated.writestr(entry, mutate(data) if entry.filename == name else data)
    return output.getvalue()


def prepared_value(content, output_format='docx', *, handle=None, template=False, operations=None):
    extension = ('.dotx' if output_format == 'docx' else '.potx') if template else f'.{output_format}'
    return {
        'schema_version': PROFILE, 'output_format': output_format, 'mode': 'refresh',
        'reference': {
            'handle': deepcopy(HANDLE if handle is None else handle),
            'content_sha256': hashlib.sha256(content).hexdigest(), 'extension': extension,
        },
        'operations': deepcopy(operations if operations is not None else [
            {'op': 'replace_text', 'target': 'body', 'text': 'Updated'},
            {'op': 'replace_text', 'target': 'unsupported-chart', 'text': 'Proposed but not applied'},
        ]),
        'unresolved': [{'target': 'missing-data', 'reason': 'A revenue figure is still needed.'}],
    }


class TestAdapter:
    __test__ = False

    def __init__(self, output_format):
        self.output_format = output_format
        self.calls = []

    def inspect_package(self, package, **kwargs):
        return {'targets': [{'id': 'body', 'type': 'paragraph'}, {'id': 'untargeted', 'type': 'chart'}]}

    def transform_package(self, package, prepared, *, check, **kwargs):
        check()
        self.calls.append(deepcopy(prepared))
        report = {'applied': [], 'skipped': [], 'warnings': ['Existing chart data was not verified.']}
        for index, operation in enumerate(prepared['operations']):
            if operation['target'] != 'body':
                report['skipped'].append({
                    'op_index': index, 'target': operation['target'],
                    'reason': 'This chart cannot be updated by the adapter.',
                })
                continue
            part = package.main_part if self.output_format == 'docx' else 'ppt/slides/slide1.xml'
            root = package.parse_xml(part)
            namespace = (
                'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
                if self.output_format == 'docx'
                else 'http://schemas.openxmlformats.org/drawingml/2006/main'
            )
            root.find(f'.//{{{namespace}}}t').text = operation['text']
            package.write_xml(part, root)
            report['applied'].append({'op_index': index, 'target': 'body', 'reason': 'Text replaced.'})
        return report


@pytest.fixture
def adapters():
    instances = {name: TestAdapter(name) for name in ('docx', 'pptx')}
    previous = {
        name: core.register_document_transform_adapter(name, adapter)
        for name, adapter in instances.items()
    }
    try:
        yield instances
    finally:
        for name, adapter in previous.items():
            core.register_document_transform_adapter(name, adapter)


def export(content, prepared, *, source_value=None, reader=None, output_format=None, **kwargs):
    return exports.build_generated_file_export(
        source=ValueSource(prepared) if source_value is None else source_value,
        export_request=exports.GeneratedFileExportRequest(output_format or prepared['output_format'], PROFILE),
        max_output_bytes=kwargs.pop('max_output_bytes', 8 * 1024 * 1024),
        document_transform_source_reader=reader or (lambda handle: content), **kwargs,
    )


def test_catalog_advertises_only_explicit_office_transform_sources(adapters):
    catalog = {item['format_id']: item for item in exports.get_generated_file_export_catalog()}
    formats = [
        name for name, entry in catalog.items()
        if any(profile['profile'] == PROFILE for profile in entry['profiles'])
    ]
    assert formats == ['docx', 'pptx']
    for name in formats:
        profile = next(profile for profile in catalog[name]['profiles'] if profile['profile'] == PROFILE)
        assert profile['source_kinds'] == ['structured_value']
        assert profile['requires_complete'] is True and profile['supported_options'] == []
        assert profile['input_schema'] == core.get_document_transform_schema()
    with pytest.raises(exports.GeneratedFileExportError) as error:
        registry.resolve_generated_file_export_format(exports.GeneratedFileExportRequest('docx', PROFILE), 'text')
    assert error.value.code == 'unsupported_source'
    with pytest.raises(exports.GeneratedFileExportError) as error:
        registry.resolve_generated_file_export_format(exports.GeneratedFileExportRequest('pdf', PROFILE), 'structured_value')
    assert error.value.code == 'unsupported_profile'


@pytest.mark.parametrize('output_format', ['docx', 'pptx'])
@pytest.mark.parametrize('template', [False, True])
def test_transform_exports_real_packages_deterministically_with_owned_metadata(adapters, output_format, template):
    content = package_bytes(output_format, template=template)
    prepared = prepared_value(content, output_format, template=template)
    original = deepcopy(prepared)
    source_value = ValueSource(prepared)
    with export(content, prepared, source_value=source_value) as result:
        payload = result.file_content.read()
        assert result.file_content.seekable() and result.file_extension == output_format
        assert result.size_bytes == len(payload) and result.content_sha256 == hashlib.sha256(payload).hexdigest()
        assert result.record_count == 0 and result.character_count is None and result.profile == PROFILE
        report = result.metadata['transformation_report']
        assert report['applied'][0]['target'] == 'body' and report['skipped'][0]['target'] == 'unsupported-chart'
        assert report['unresolved'][0]['target'] == 'missing-data' and report['status'] == 'draft'
        assert report['output']['content_sha256'] == result.content_sha256
        document = Document(io.BytesIO(payload)) if output_format == 'docx' else Presentation(io.BytesIO(payload))
        text = document.paragraphs[0].text if output_format == 'docx' else document.slides[0].shapes.title.text
        assert text == 'Updated'
    assert result.file_content.closed and source_value.reads == 1 and prepared == original
    with export(content, prepared) as repeated:
        repeated_payload = repeated.file_content.read()
    assert repeated_payload == payload


@pytest.mark.parametrize('failure,code', [
    ('corrupt', 'invalid_data'), ('digest', 'source_changed'),
    ('active', 'unsupported_content'), ('output_limit', 'size_limit'), ('value_limit', 'value_limit'),
])
def test_reference_faults_map_to_user_safe_export_errors(adapters, failure, code):
    content = package_bytes()
    options = {}
    if failure == 'corrupt':
        content = b'not a package'
    elif failure == 'active':
        content = rewrite_part(content, 'word/document.xml', lambda data: data.replace(
            b'<w:body>', b'<w:body><w:p><w:fldSimple w:instr="DDEAUTO secret"/></w:p>',
        ))
    elif failure == 'output_limit':
        options['max_output_bytes'] = 8
    elif failure == 'value_limit':
        options['limits'] = exports.GeneratedFileExportLimits(max_value_bytes=8)
    prepared = prepared_value(content)
    if failure == 'digest':
        prepared['reference']['content_sha256'] = '0' * 64
    with pytest.raises(exports.GeneratedFileExportError) as error:
        export(content, prepared, **options)
    assert error.value.code == code and error.value.retryable is False
    assert 'secret' not in str(error.value) and 'not a package' not in str(error.value)


@pytest.mark.parametrize('error', [
    PermissionError('private authority'), TimeoutError('private storage'), InterruptedError('private stop'),
    OSError('private source'), ValueError('private callback'),
    core.DocumentTransformError('reference_changed', 'private callback error'),
])
def test_native_callback_fault_identity_is_preserved(adapters, error):
    content = package_bytes()

    def reader(handle):
        raise error

    with pytest.raises(type(error)) as failure:
        export(content, prepared_value(content), reader=reader)
    assert failure.value is error


def test_exact_output_format_is_checked_before_native_read(adapters):
    content = package_bytes()
    calls = []
    with pytest.raises(exports.GeneratedFileExportError) as error:
        export(content, prepared_value(content), output_format='pptx', reader=lambda handle: calls.append(handle))
    assert error.value.code == 'invalid_data' and calls == []


@pytest.mark.parametrize('invalid_reference', [
    {'kind': 'document', 'document_id': 'document-1', 'path': r'C:\private.docx'},
    {'kind': 'artifact', 'conversation_id': 'conversation-1', 'message_id': 'prior', 'url': 'https://private.test/file'},
])
def test_storage_locators_are_never_accepted_as_native_reference_handles(adapters, invalid_reference):
    content = package_bytes()
    prepared = prepared_value(content)
    prepared['reference']['handle'] = invalid_reference
    calls = []
    with pytest.raises(exports.GeneratedFileExportError) as error:
        export(content, prepared, reader=lambda handle: calls.append(handle))
    assert error.value.code == 'invalid_data' and calls == []


def test_transform_keeps_existing_complete_input_requirement(adapters):
    content = package_bytes()
    prepared = prepared_value(content)
    readiness = exports.GeneratedFileExportReadiness(state='ready', is_complete=False)
    with pytest.raises(exports.GeneratedFileExportError) as error:
        export(content, prepared, source_value=ValueSource(prepared, readiness))
    assert error.value.code == 'incomplete_source'


def test_stream_closes_when_post_transform_source_check_fails(adapters, monkeypatch):
    content = package_bytes()
    prepared = prepared_value(content)
    source_value = ValueSource(prepared)
    original = bridge.transform_document
    rendered = []
    failure = PermissionError('private revoked source')

    def revoke(*args, **kwargs):
        result = original(*args, **kwargs)
        rendered.append(result)

        def denied():
            raise failure

        source_value.on_recheck = denied
        return result

    monkeypatch.setattr(bridge, 'transform_document', revoke)
    with pytest.raises(PermissionError) as error:
        export(content, prepared, source_value=source_value)
    assert error.value is failure and rendered[0].file_content.closed


def test_new_reader_cannot_change_other_export_contracts(adapters):
    reader = lambda handle: b'not allowed'
    for kwargs in (
        {},
        {'source': ValueSource({}), 'export_request': exports.GeneratedFileExportRequest('json', 'structured_value_v1'),
         'max_output_bytes': 1024},
        {'source': TextSource(['report'], 6), 'export_request': exports.GeneratedFileExportRequest('docx', 'prepared_report_v1'),
         'max_output_bytes': 1024},
        {'source': ValueSource({}), 'export_request': exports.GeneratedFileExportRequest('docx', 'assembled_document_v1'),
         'max_output_bytes': 1024},
    ):
        with pytest.raises(exports.GeneratedFileExportError) as error:
            exports.build_generated_file_export(document_transform_source_reader=reader, **kwargs)
        assert error.value.code == 'invalid_options'


class RenderCase:
    """Real retained lineage and durable output state with only storage I/O doubled."""

    def __init__(
        self, content, prepared=None, *, include_inspection=True,
        inspector_capability='document_inspect', authorized_handle=None,
    ):
        self.content = content
        self.prepared = deepcopy(prepared or prepared_value(content))
        self.authorized_handle = deepcopy(HANDLE if authorized_handle is None else authorized_handle)
        self.now = datetime(2030, 1, 1, tzinfo=timezone.utc)
        self.fixture = ResultFixture()
        inspector = replace(
            self.fixture.producer, step_id='inspect', capability_id=inspector_capability,
            contract_version='document-inspect-v1',
        )
        self.fixture.add_producer(inspector)
        inspection = core.inspect_document(content, 'source.docx', self.authorized_handle)
        inspected = self.fixture.service.persist_task_result(
            producer=inspector, role='gather' if inspector_capability == 'document_inspect' else 'reason',
            status='complete', outputs=[results_module.NamedOutput('inspection', 'structured-v1', inspection, complete(1))],
            sources=[source()] if self.authorized_handle['kind'] == 'document' else [],
            origin='grounded' if self.authorized_handle['kind'] == 'document' else 'generated',
            guard_token='server-attempt-token',
        )
        transform_producer = replace(
            self.fixture.producer, step_id='prepare', capability_id='document_refresh',
            contract_version='document-transform-v1',
        )
        self.fixture.add_producer(transform_producer)
        self.task = self.fixture.service.persist_task_result(
            producer=transform_producer, role='reason', status='complete',
            outputs=[results_module.NamedOutput('transformation', 'structured-v1', self.prepared, complete(1))],
            sources=[], upstream=[inspected.output('inspection')] if include_inspection else [],
            origin='grounded' if include_inspection else 'generated', guard_token='server-attempt-token',
        )
        self.render_producer = replace(
            self.fixture.producer, step_id='save', capability_id='render_file', contract_version='render-file-v1',
        )
        self.fixture.add_producer(self.render_producer)
        self.runs = AtomicMemoryContainer('conversation_id')
        run = deepcopy(self.fixture.runs['run-1'])
        run.update(record_type='run', approval={'state': 'approved'},
                   execution_deadline_at=(self.now + timedelta(hours=1)).isoformat())
        self.runs.create_item(run)
        self.messages = {}
        self.bytes = {}
        self.reads = []
        self.failures = []
        self.streams = []
        self.uploads = 0
        self.restart()
        self.output = self.service.ensure_output(
            producer=self.render_producer, source_ref=self.task.output('transformation'),
            export_request=exports.GeneratedFileExportRequest('docx', PROFILE), file_name='draft.docx',
            approved_work_id='run-1', deadline_at=run['execution_deadline_at'],
        )

    def read_native(self, handle):
        self.reads.append(deepcopy(handle))
        if self.failures:
            raise self.failures.pop(0)
        if handle != self.authorized_handle:
            raise PermissionError('Unexpected identity')
        if handle['kind'] == 'document':
            self.fixture.metadata(handle['document_id'], 'owner')
        elif handle['conversation_id'] != 'conversation-1':
            raise PermissionError('Unexpected artifact conversation')
        return self.content

    def render(self, **kwargs):
        rendered = exports.build_generated_file_export(**kwargs)
        self.streams.append(rendered.file_content)
        return rendered

    def upload(self, user_id, conversation_id, file_name, stream, size, **kwargs):
        record = self.service.store.get(kwargs['generated_artifact_source']['output_id'])
        message = self.transport._synthetic_message(record, record['intent'])
        self.messages[message['id']] = message
        self.bytes[message['id']] = stream.read()
        self.uploads += 1
        return {'conversation_id': conversation_id, 'message': message}

    @contextmanager
    def open_stream(self, message, *, check):
        if message['id'] not in self.bytes:
            raise ResourceNotFoundError()
        check()
        with io.BytesIO(self.bytes[message['id']]) as stream:
            yield stream
        check()

    def restart(self):
        results = self.fixture.restart()
        results.access.read_run = lambda run_id: self.runs.read_item(run_id, 'conversation-1')
        store = output_store.OrchestrationOutputStore(
            self.runs, user_id='owner', conversation_id='conversation-1',
            read_conversation=results.access.read_conversation, clock=lambda: self.now,
        )
        self.transport = artifacts.OrchestrationArtifactTransport(
            upload=self.upload, read_message=lambda conversation_id, message_id: self.messages.get(message_id),
            open_stream=self.open_stream, delete=lambda *args, **kwargs: None, blob_container='chat',
        )
        self.service = rendering.OrchestrationRenderingService(
            store, results, self.transport, authorize_execution=lambda record, operation: None,
            max_output_bytes=8 * 1024 * 1024, renderer=self.render, jitter=lambda: 0.5,
            document_transform_source_reader=self.read_native,
        )
        return self.service

    def run(self):
        result = self.service.render_attempt(self.output['output_id'])
        assert all(stream.closed for stream in self.streams)
        return result


@pytest.mark.parametrize('tampering', ['missing_inspection', 'fake_inspector', 'handle', 'digest', 'extension'])
def test_model_supplied_reference_cannot_bypass_inspection_authority(adapters, tampering):
    content = package_bytes()
    prepared = prepared_value(content)
    options = {}
    if tampering == 'missing_inspection':
        options['include_inspection'] = False
    elif tampering == 'fake_inspector':
        options['inspector_capability'] = 'compose'
    elif tampering == 'handle':
        prepared['reference']['handle']['document_id'] = 'another-document'
    elif tampering == 'digest':
        prepared['reference']['content_sha256'] = '0' * 64
    else:
        prepared['reference']['extension'] = '.dotx'
    case = RenderCase(content, prepared, **options)
    result = case.run()
    assert result['state'] == 'failed' and result['error_code'] == 'output_source_changed'
    assert case.reads == [] and case.uploads == 0


def test_native_reader_rejects_unretained_handle_before_callback(adapters):
    case = RenderCase(package_bytes())
    record = case.service.store.get(case.output['output_id'])
    reader = case.service._document_transform_source_reader_for(record)
    with pytest.raises(output_store.OutputUnavailableError) as error:
        reader({'kind': 'artifact', 'conversation_id': 'conversation-1', 'message_id': 'unretained'})
    assert error.value.code == 'output_source_changed' and case.reads == []


def test_inspected_artifact_handle_uses_exact_same_native_read_boundary(adapters):
    content = package_bytes()
    handle = {'kind': 'artifact', 'conversation_id': 'conversation-1', 'message_id': 'prior-artifact'}
    case = RenderCase(content, prepared_value(content, handle=handle), authorized_handle=handle)
    result = case.run()
    assert result['state'] == 'completed' and case.reads == [handle]
    assert 'prior-artifact' not in json.dumps(result['transformation_report'])


def test_rendered_digest_is_verified_and_failed_stream_is_closed(adapters):
    case = RenderCase(package_bytes())

    def corrupt_metadata(**kwargs):
        result = case.render(**kwargs)
        result.content_sha256 = '0' * 64
        return result

    case.service.renderer = corrupt_metadata
    result = case.run()
    assert result['state'] == 'failed' and result['error_code'] == 'output_digest_mismatch'
    assert case.uploads == 0


def test_transform_does_not_use_an_unbound_generic_image_resolver(adapters):
    case = RenderCase(package_bytes())
    record = case.service.store.get(case.output['output_id'])
    request = exports.GeneratedFileExportRequest('docx', PROFILE)
    entry = registry.resolve_generated_file_export_format(request, 'structured_value')
    case.service.image_resolver = lambda token: pytest.fail('Unbound image resolver was used')
    resolver = case.service._image_resolver_for(record, entry)
    assert resolver is None
    case.service.image_asset_reader = lambda token: pytest.fail('An unretained image was read')
    resolver = case.service._image_resolver_for(record, entry)
    with pytest.raises(output_store.OutputUnavailableError) as error:
        resolver('asset:unretained')
    assert error.value.code == 'output_source_unavailable'


def test_retry_reuses_prepared_data_and_committed_report_survives_reload(adapters):
    case = RenderCase(package_bytes())
    writes = case.fixture.container.sequence
    case.failures = [TimeoutError('private storage outage')]
    pending = case.run()
    assert pending['state'] == 'retry_scheduled' and pending['error_code'] == 'output_storage_unavailable'
    record = case.service.store.get(pending['output_id'])
    case.now = output_store.parse_time(record['next_retry_at'])
    case.restart()
    completed = case.run()
    assert completed['state'] == 'completed' and completed['automatic_attempts'] == 2
    assert case.fixture.container.sequence == writes and len(adapters['docx'].calls) == 1
    assert case.uploads == 1 and case.reads == [HANDLE, HANDLE]
    report = completed['transformation_report']
    assert report['applied_count'] == report['skipped_count'] == report['unresolved_count'] == 1
    assert report['unchanged_count'] == 1
    private = case.service.store.get(completed['output_id'])
    assert private['renderer_metadata']['intent_id'] == private['committed_intent']['intent_id']
    encoded = json.dumps(report)
    assert 'document_id' not in encoded and 'content_sha256' not in encoded
    assert hashlib.sha256(case.content).hexdigest() not in encoded and 'document-1' not in encoded
    case.restart()
    reopened = case.service.read(completed['output_id'])
    assert reopened['transformation_report'] == report
    notes = deliverables.delivery_notes({}, {}, file_outputs=[reopened])
    assert 'Applied 1 change' in notes and 'unsupported-chart' in notes and 'missing-data' in notes
    assert 'Proposed but not applied' not in notes and 'document-1' not in notes
    repeat = case.run()
    assert repeat == reopened and case.uploads == 1 and len(adapters['docx'].calls) == 1
    hidden = output_store.public_output(private, withhold_details=True)
    unavailable = output_store.public_output(private, unavailable_code='output_access_denied')
    assert 'transformation_report' not in hidden and 'transformation_report' not in unavailable


@pytest.mark.parametrize('error,code', [
    (PermissionError('private denial'), 'output_access_denied'),
    (LookupError('private missing bytes'), 'output_source_unavailable'),
])
def test_rendering_preserves_source_access_failure_classification(adapters, error, code):
    case = RenderCase(package_bytes())
    case.failures = [error]
    result = case.run()
    assert result['state'] == 'failed' and result['error_code'] == code
    assert 'private' not in json.dumps(result) and case.uploads == 0


def test_changed_native_bytes_are_rejected_without_publication(adapters):
    case = RenderCase(package_bytes())
    case.content = package_bytes('pptx')
    result = case.run()
    assert result['state'] == 'failed' and result['error_code'] == 'output_source_changed'
    assert case.uploads == 0


def test_no_changes_are_explicitly_delivered_as_an_unchanged_reference_draft(adapters):
    content = package_bytes()
    prepared = prepared_value(content, operations=[])
    prepared['unresolved'] = []
    case = RenderCase(content, prepared)
    result = case.run()
    assert result['state'] == 'completed' and result['transformation_report']['applied_count'] == 0
    assert 'no changes applied' in result['message'].lower()
    notes = deliverables.delivery_notes({}, {}, file_outputs=[result])
    assert 'No changes applied; unchanged reference draft' in notes


def test_report_detail_bounds_preserve_exact_counts_and_remove_reference_secrets(adapters):
    content = package_bytes()
    prepared = prepared_value(content)
    with export(content, prepared) as result:
        report = deepcopy(result.metadata['transformation_report'])
    secret = prepared['reference']['content_sha256']
    report['unchanged'] = [
        {'target': f'paragraph-{index}', 'reason': f'Untouched source document-1 {secret}'}
        for index in range(100)
    ]
    retained = output_store._retained_transformation_report(report)
    projected = output_store.public_transformation_report({'transformation_report': retained})
    assert projected['unchanged_count'] == 100 and len(projected['unchanged']) == 32
    assert secret not in json.dumps(projected) and 'document-1' not in json.dumps(projected)
    assert 'reference' not in projected and 'output' not in projected


if __name__ == '__main__':
    raise SystemExit(pytest.main([__file__, '-q', '--tb=short']))
