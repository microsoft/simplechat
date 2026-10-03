#!/usr/bin/env python3
# test_orchestration_document_merge_capability.py
"""
Functional test for the document_merge chat orchestration capability.
Version: 0.261.224
Implemented in: 0.261.224

This test ensures that document_merge is a gated Reason capability beside tabular_merge:
plans must name the kind and may only use that kind's settings, sources are named
explicitly or bound from a search, the administrator's chat Merge limit applies, and a
merged file can only be rendered from a document merge's assembly in that merge's format.
The real executor assembles PDFs, Word documents, PowerPoint decks and workbooks once
without any model call and retains a description of the checked file plus a report,
never the file itself. Failures reach users only as application-owned messages, and
Render assembles the same files again and delivers the file only when it is
byte-for-byte the one that was checked.
"""

import hashlib
import io
import sys
from types import ModuleType

import pytest

from test_orchestration_dependency_runtime import binding, execute, runtime  # noqa: F401
from test_support.versioning import assert_app_version_at_least


AVAILABLE = ['document_merge', 'tabular_merge', 'document_search', 'compose', 'render_file']
ASSEMBLED = 'assembled_document_v1'


def pdf_bytes(pages, *, encrypt=False):
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=200, height=200)
    if encrypt:
        writer.encrypt(user_password='', owner_password='owner')
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def docx_bytes(label):
    from docx import Document

    document = Document()
    document.add_heading(f'{label} title', level=1)
    document.add_paragraph(f'{label} body')
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pptx_bytes(label, slides):
    from pptx import Presentation

    presentation = Presentation()
    for number in range(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = f'{label} slide {number + 1}'
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def xlsx_bytes(title):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(['Region', 'Amount'])
    sheet.append(['West', 5])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


FILES = {
    'report-a': ('Report A.pdf', pdf_bytes(2)),
    'report-b': ('Report B.pdf', pdf_bytes(1)),
    'locked': ('Locked.pdf', pdf_bytes(1, encrypt=True)),
    'letter-a': ('Letter A.docx', docx_bytes('Letter A')),
    'letter-b': ('Letter B.docx', docx_bytes('Letter B')),
    'deck-a': ('Deck A.pptx', pptx_bytes('Deck A', 2)),
    'deck-b': ('Deck B.pptx', pptx_bytes('Deck B', 1)),
    'east': ('east.csv', b'Region,Amount\r\nEast,10\r\n'),
    'west': ('west.xlsx', xlsx_bytes('West')),
}
KIND_FILES = {
    'pdf': ['report-a', 'report-b'],
    'docx': ['letter-a', 'letter-b'],
    'pptx': ['deck-a', 'deck-b'],
    'workbook': ['east', 'west'],
}
KIND_FORMATS = {'pdf': 'pdf', 'docx': 'docx', 'pptx': 'pptx', 'workbook': 'xlsx'}


def merge_step(document_ids=None, **arguments):
    step = {'step_id': 'merge', 'capability_id': 'document_merge', 'arguments': dict(arguments)}
    if document_ids is not None:
        step['arguments']['document_ids'] = list(document_ids)
    return step


def render_step(output_format, *, profile=ASSEMBLED, output='assembly', producer='merge'):
    return {
        'step_id': 'save', 'capability_id': 'render_file',
        'arguments': {
            'file_name': f'merged.{output_format}', 'output_format': output_format, 'profile': profile,
            'options': {},
        },
        'inputs': {'source': {'binding': binding(producer, output), 'allow_partial': False}},
        'outputs': [],
    }


def document_case(runtime, steps, *, settings=None, available=AVAILABLE):
    case = runtime.make(steps, settings=settings, available=available)
    for document_id, (file_name, _) in FILES.items():
        case.fixture.sources[document_id] = {
            'document_id': document_id, 'scope': 'personal', 'scope_id': 'owner',
            'source_version': 1, 'source_revision': f'{document_id}-revision',
            'file_name': file_name,
            'source_kind': 'tabular' if file_name.endswith(('.csv', '.xlsx')) else 'narrative',
        }
    reads = []

    def reader(source, user_id, group_id=None, public_workspace_id=None, *, purpose):
        assert user_id == 'owner' and purpose == 'native'
        assert group_id is None and public_workspace_id is None
        reads.append(source['document_id'])
        return {'id': source['document_id']}, FILES[source['document_id']][1]

    case.context.merge_source_reader = reader
    case.reads = reads
    return case


def merged_task(runtime, kind, **arguments):
    case = document_case(runtime, [merge_step(KIND_FILES[kind], kind=kind, **arguments)])
    result = execute(runtime, case)
    assert result['status'] == 'completed', result
    return case, result, runtime.contracts.TaskResult.from_dict(result['task_results']['merge'])


def render(case, task, output_format, reader, *, max_output_bytes=16 * 1024 * 1024):
    from functions_generated_file_exports import GeneratedFileExportRequest, build_generated_file_export
    from functions_orchestration_export_sources import open_orchestration_export_source

    source = open_orchestration_export_source(case.fixture.restart(), task.output('assembly'))
    with build_generated_file_export(
        source=source, export_request=GeneratedFileExportRequest(output_format, ASSEMBLED),
        max_output_bytes=max_output_bytes, document_reader=reader,
    ) as output:
        content = output.file_content.read()
        source.require_complete_consumption()
        return output, content


def original_reader(document_id):
    return FILES[document_id][1]


def test_version_includes_document_merge():
    assert_app_version_at_least('0.261.224')


def test_document_merge_is_a_gated_reason_capability(runtime):
    registry = runtime.registry
    capability = registry.get_capability('document_merge')
    assert capability['role'] == registry.ROLE_REASON
    assert capability['label'] == 'Merge documents'
    assert capability['document_action_type'] == 'merge'
    assert capability['result_outputs'] == {'assembly': 'structured-v1', 'report': 'structured-v1'}
    assert capability['result_input_kinds'] == {'sources': ('source-set-v1',)}
    assert capability['result_contract_version'] == 'document-merge-v1'
    assert capability['adapter'] == 'document_merge'
    inputs = capability['inputs']
    assert inputs['required'] == ['kind']
    assert inputs['properties']['kind']['enum'] == ['pdf', 'docx', 'pptx', 'workbook']
    assert inputs['properties']['document_ids']['minItems'] == 2
    # Settings a kind doesn't use stay absent, so a plan never restates a default.
    for name in ('bookmarks', 'formatting', 'page_breaks', 'source_headings', 'sections', 'sheet', 'sheets'):
        assert 'default' not in inputs['properties'][name]
    assert 'document_merge' in runtime.contracts.REASON_CAPABILITIES

    settings = {'enable_chat_orchestration': True, 'enable_user_workspace': True}
    available = {item['id'] for item in registry.resolve_available_capabilities(settings)}
    assert 'document_merge' in available
    disabled = {**settings, 'document_action_capabilities': {'merge': {'enabled': False}}}
    assert 'document_merge' not in {item['id'] for item in registry.resolve_available_capabilities(disabled)}
    no_workspaces = {'enable_chat_orchestration': True}
    assert 'document_merge' not in {item['id'] for item in registry.resolve_available_capabilities(no_workspaces)}

    assert registry.get_capability_document_limit(capability, settings=settings) == 10
    narrowed = {**settings, 'document_action_capabilities': {'merge': {'chat_max_documents': 4}}}
    assert registry.get_capability_document_limit(capability, settings=narrowed) == 4

    # Administrators narrow plans with the same capability list, so Merge documents is offered there.
    from admin_settings_fields import ADMIN_SETTINGS_FIELDS

    options = [
        option for fields in ADMIN_SETTINGS_FIELDS.values() for field in fields
        if field.get('key') == 'chat_orchestration_enabled_capabilities' for option in field['options']
    ]
    values = [option['value'] for option in options]
    assert {'value': 'document_merge', 'label': 'Merge documents'} in options
    assert values.index('tabular_merge') < values.index('document_merge')


def test_plans_need_a_kind_and_use_only_its_settings(runtime):
    schema = runtime.schema
    plan = runtime.make([merge_step(['report-a', 'report-b'], kind='pdf')], available=AVAILABLE).plan
    assert plan['steps'][0]['arguments'] == {
        'document_ids': ['report-a', 'report-b'], 'doc_scope': 'all', 'kind': 'pdf',
    }
    assert [output['name'] for output in plan['steps'][0]['outputs']] == ['assembly', 'report']
    runtime.make([merge_step(
        ['letter-a', 'letter-b'], kind='docx', formatting='use_first', page_breaks=False, source_headings=True,
    )], available=AVAILABLE)
    runtime.make([merge_step(['deck-a', 'deck-b'], kind='pptx', formatting='keep_source', sections=False)],
                 available=AVAILABLE)
    runtime.make([merge_step(['east', 'west'], kind='workbook', sheets='all')], available=AVAILABLE)

    with pytest.raises(schema.PlanValidationError):
        runtime.make([merge_step(['report-a', 'report-b'])], available=AVAILABLE)
    with pytest.raises(schema.PlanValidationError):
        runtime.make([merge_step(['report-a', 'report-b'], kind='zip')], available=AVAILABLE)
    for kind, wrong in (
        ('pdf', {'formatting': 'use_first'}), ('pdf', {'sheets': 'all'}), ('docx', {'sections': False}),
        ('pptx', {'page_breaks': False}), ('workbook', {'bookmarks': False}),
        ('workbook', {'sheet': 'Data', 'sheets': 'all'}),
    ):
        with pytest.raises(schema.PlanValidationError) as refused:
            runtime.make([merge_step(['report-a', 'report-b'], kind=kind, **wrong)], available=AVAILABLE)
        assert refused.value.code == 'merge_options_invalid', (kind, wrong)

    with pytest.raises(schema.PlanValidationError) as missing:
        runtime.make([merge_step(kind='pdf')], available=AVAILABLE)
    assert missing.value.code == 'source_binding_required'
    with pytest.raises(schema.PlanValidationError):
        runtime.make([merge_step(['report-a'], kind='pdf')], available=AVAILABLE)

    search = {'step_id': 'search', 'capability_id': 'document_search', 'arguments': {'query': 'board packs'}}
    bound = merge_step(kind='pdf')
    bound['inputs'] = {'sources': {'binding': binding('search', 'sources'), 'allow_partial': False}}
    plan = runtime.make([search, bound], available=AVAILABLE).plan
    assert plan['steps'][1]['depends_on'] == ['search']
    both = merge_step(['report-a', 'report-b'], kind='pdf')
    both['inputs'] = bound['inputs']
    with pytest.raises(schema.PlanValidationError):
        runtime.make([search, both], available=AVAILABLE)


def test_the_chat_limit_and_source_kinds_are_enforced_at_planning(runtime):
    schema = runtime.schema
    ids = [f'file-{index}' for index in range(4)]
    settings = {'document_action_capabilities': {'merge': {'chat_max_documents': 3}}}
    with pytest.raises(schema.PlanValidationError) as limited:
        runtime.make([merge_step(ids, kind='pdf')], settings=settings, available=AVAILABLE)
    assert 'document limit' in str(limited.value)

    plan = runtime.make([merge_step(['report-a', 'east'], kind='pdf')], available=AVAILABLE).plan
    with pytest.raises(schema.PlanValidationError) as spreadsheets:
        schema.validate_plan_document_source_kinds(plan, {'report-a': 'narrative', 'east': 'tabular'})
    assert spreadsheets.value.code == 'source_kind_invalid'
    schema.validate_plan_document_source_kinds(plan, {'report-a': 'narrative', 'east': 'narrative'})

    workbook = runtime.make([merge_step(['east', 'report-a'], kind='workbook')], available=AVAILABLE).plan
    with pytest.raises(schema.PlanValidationError) as documents:
        schema.validate_plan_document_source_kinds(workbook, {'east': 'tabular', 'report-a': 'narrative'})
    assert documents.value.code == 'source_kind_invalid'
    schema.validate_plan_document_source_kinds(workbook, {'east': 'tabular', 'report-a': 'tabular'})
    schema.validate_plan_document_source_kinds(workbook, {'east': 'tabular'})


@pytest.mark.parametrize('kind', ['pdf', 'docx', 'pptx', 'workbook'])
def test_a_merged_file_renders_only_from_an_assembly_in_its_own_format(runtime, kind):
    schema = runtime.schema
    merge = merge_step(KIND_FILES[kind], kind=kind)
    plan = runtime.make([merge, render_step(KIND_FORMATS[kind])], available=AVAILABLE).plan
    assert plan['steps'][1]['depends_on'] == ['merge']
    other = next(fmt for fmt in ('pdf', 'docx', 'pptx', 'xlsx') if fmt != KIND_FORMATS[kind])
    with pytest.raises(schema.PlanValidationError) as wrong_format:
        runtime.make([merge, render_step(other)], available=AVAILABLE)
    assert wrong_format.value.code == 'result_kind_incompatible'
    with pytest.raises(schema.PlanValidationError) as report:
        runtime.make([merge, render_step(KIND_FORMATS[kind], output='report')], available=AVAILABLE)
    assert report.value.code == 'result_kind_incompatible'


def test_the_assembled_profile_never_renders_other_results(runtime):
    schema = runtime.schema
    draft = {
        'step_id': 'draft', 'capability_id': 'compose',
        'arguments': {'instruction': 'Prepare a structured summary of the plan.'},
        'inputs': {}, 'outputs': [{'name': 'answer', 'kind': 'structured-v1'}], 'depends_on': [],
    }
    with pytest.raises(schema.PlanValidationError) as refused:
        runtime.make([draft, render_step('pdf', output='answer', producer='draft')], available=AVAILABLE)
    assert refused.value.code == 'result_kind_incompatible'


@pytest.mark.parametrize('kind', ['pdf', 'docx', 'pptx', 'workbook'])
def test_executor_retains_a_checked_description_and_report_without_a_model(runtime, kind):
    case, result, task = merged_task(runtime, kind)
    assert case.model.calls == []
    assert case.reads == KIND_FILES[kind]
    assert task.status == 'complete'
    service = case.fixture.restart()
    assembly = service.open_result(task.output('assembly')).read_value()
    assert assembly['profile'] == 'document_assembly_v1'
    assert assembly['kind'] == kind
    assert [part['document_id'] for part in assembly['parts']] == KIND_FILES[kind]
    assert [part['file_name'] for part in assembly['parts']] == [FILES[item][0] for item in KIND_FILES[kind]]
    assert assembly['output']['format'] == KIND_FORMATS[kind]
    assert assembly['output']['size_bytes'] > 0
    assert len(assembly['output']['content_sha256']) == 64
    # Only the description is retained; no file bytes are stored in the result.
    assert set(assembly) == {'profile', 'kind', 'options', 'parts', 'output'}
    report = service.open_result(task.output('report')).read_value()
    assert report['status'] == 'merged'
    assert report['kind'] == kind
    assert [entry['status'] for entry in report['parts']] == ['merged', 'merged']
    reference = task.output('assembly')
    assert reference.completeness.status == 'complete'
    assert reference.completeness.coverage.unit == 'sources'
    assert reference.completeness.coverage.expected == 2
    metadata = service.open_result(reference).metadata()
    assert metadata['origin'] == 'grounded'


def test_step_summaries_describe_the_merged_file(runtime):
    expectations = {
        'pdf': 'Merged 2 file(s) into one PDF with 3 page(s).',
        'docx': 'Merged 2 file(s) into one Word document.',
        'pptx': 'Merged 2 file(s) into one PowerPoint deck with 3 slide(s).',
        'workbook': 'Merged 2 file(s) into one Excel workbook with 2 sheet(s).',
    }
    for kind, expected in expectations.items():
        _case, result, _task = merged_task(runtime, kind)
        summaries = [step['summary'] for step in result['steps'] if step['step_id'] == 'merge']
        assert summaries == [expected], kind


@pytest.mark.parametrize('kind', ['pdf', 'docx', 'pptx', 'workbook'])
def test_render_assembles_the_same_files_and_delivers_the_checked_bytes(runtime, kind):
    case, _result, task = merged_task(runtime, kind)
    assembly = case.fixture.restart().open_result(task.output('assembly')).read_value()
    reads = []

    def reader(document_id):
        reads.append(document_id)
        return original_reader(document_id)

    output, content = render(case, task, KIND_FORMATS[kind], reader)
    assert reads == KIND_FILES[kind]
    assert hashlib.sha256(content).hexdigest() == assembly['output']['content_sha256']
    assert output.content_sha256 == assembly['output']['content_sha256']
    assert output.size_bytes == len(content) == assembly['output']['size_bytes']
    assert output.record_count == 0 and output.character_count is None
    assert output.output_format == KIND_FORMATS[kind]
    if kind == 'pdf':
        from pypdf import PdfReader

        merged = PdfReader(io.BytesIO(content))
        assert len(merged.pages) == 3
        assert [item.title for item in merged.outline] == ['Report A.pdf', 'Report B.pdf']
    elif kind == 'docx':
        from docx import Document

        texts = [paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs]
        assert texts.index('Letter A body') < texts.index('Letter B body')
    elif kind == 'pptx':
        from pptx import Presentation

        titles = [slide.shapes.title.text for slide in Presentation(io.BytesIO(content)).slides]
        assert titles == ['Deck A slide 1', 'Deck A slide 2', 'Deck B slide 1']
    else:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True)
        try:
            assert len(workbook.sheetnames) == 2
        finally:
            workbook.close()


def test_render_refuses_a_file_that_no_longer_matches_its_sources(runtime):
    from functions_generated_export_contracts import GeneratedFileExportError

    case, _result, task = merged_task(runtime, 'pdf')

    def changed(document_id):
        return pdf_bytes(4) if document_id == 'report-b' else original_reader(document_id)

    with pytest.raises(GeneratedFileExportError) as refused:
        render(case, task, 'pdf', changed)
    assert refused.value.code == 'source_changed'

    def missing(document_id):
        return b'not a pdf'

    with pytest.raises(GeneratedFileExportError) as unreadable:
        render(case, task, 'pdf', missing)
    assert unreadable.value.code == 'source_changed'


def test_render_passes_access_failures_through_unchanged(runtime):
    case, _result, task = merged_task(runtime, 'docx')

    class Revoked(PermissionError):
        pass

    def revoked(document_id):
        raise Revoked('access was removed')

    with pytest.raises(Revoked):
        render(case, task, 'docx', revoked)


def test_render_refuses_other_formats_missing_readers_and_small_limits(runtime):
    from functions_generated_export_contracts import GeneratedFileExportError, GeneratedFileExportRequest
    from functions_generated_file_exports import build_generated_file_export
    from functions_orchestration_export_sources import open_orchestration_export_source

    case, _result, task = merged_task(runtime, 'pdf')
    with pytest.raises(GeneratedFileExportError) as wrong_format:
        render(case, task, 'docx', original_reader)
    assert wrong_format.value.code == 'unsupported_source'
    with pytest.raises(GeneratedFileExportError) as no_reader:
        render(case, task, 'pdf', None)
    assert no_reader.value.code == 'unsupported_source'
    with pytest.raises(GeneratedFileExportError) as too_small:
        render(case, task, 'pdf', original_reader, max_output_bytes=16)
    assert too_small.value.code == 'size_limit'
    size = case.fixture.restart().open_result(task.output('assembly')).read_value()['output']['size_bytes']
    with pytest.raises(GeneratedFileExportError) as one_byte_short:
        render(case, task, 'pdf', original_reader, max_output_bytes=size - 1)
    assert one_byte_short.value.code == 'size_limit'
    render(case, task, 'pdf', original_reader, max_output_bytes=size)

    report = open_orchestration_export_source(case.fixture.restart(), task.output('report'))
    with pytest.raises(GeneratedFileExportError) as not_an_assembly:
        build_generated_file_export(
            source=report, export_request=GeneratedFileExportRequest('pdf', ASSEMBLED),
            max_output_bytes=16 * 1024 * 1024, document_reader=original_reader,
        )
    assert not_an_assembly.value.code == 'invalid_source'
    with pytest.raises(GeneratedFileExportError) as reader_elsewhere:
        build_generated_file_export(
            source=report, export_request=GeneratedFileExportRequest('json', 'structured_value_v1'),
            max_output_bytes=16 * 1024 * 1024, document_reader=original_reader,
        )
    assert reader_elsewhere.value.code == 'invalid_options'


def test_merge_failures_use_application_owned_messages(runtime):
    case = document_case(runtime, [merge_step(['report-a', 'locked'], kind='pdf')])
    result = execute(runtime, case)
    assert result['status'] == 'failed'
    assert result['failure']['code'] == 'document_merge_source_unreadable'
    assert 'Locked' not in result['message']
    assert "couldn't be read" in result['message']

    case = document_case(runtime, [merge_step(['report-a', 'letter-a'], kind='pdf')])
    result = execute(runtime, case)
    assert result['status'] == 'failed'
    assert result['failure']['code'] == 'document_merge_sources_invalid'
    assert 'Letter' not in result['message']
    assert case.reads == []


def test_merge_reads_are_refused_when_a_source_is_unavailable(runtime):
    from functions_orchestration_results import ResultUnavailableError

    case = document_case(runtime, [merge_step(['report-a', 'report-b'], kind='pdf')])
    case.fixture.denied.add('report-b')
    # The executor re-authorizes every planned source before any step runs.
    with pytest.raises(ResultUnavailableError):
        execute(runtime, case)
    assert case.reads == []


def test_merge_can_bind_the_files_a_search_found(runtime, monkeypatch):
    search = {'step_id': 'search', 'capability_id': 'document_search', 'arguments': {'query': 'board decks'}}
    bound = merge_step(kind='pptx')
    bound['inputs'] = {'sources': {'binding': binding('search', 'sources'), 'allow_partial': False}}
    case = document_case(runtime, [search, bound])
    module = ModuleType('functions_search')
    module.hybrid_search = lambda *args, **kwargs: [
        {'document_id': 'deck-a', 'id': 'hit-1', 'chunk_text': 'Deck A'},
        {'document_id': 'deck-b', 'id': 'hit-2', 'chunk_text': 'Deck B'},
    ]
    monkeypatch.setitem(sys.modules, 'functions_search', module)
    result = execute(runtime, case)
    assert result['status'] == 'completed', result
    assert case.reads == ['deck-a', 'deck-b']
    task = runtime.contracts.TaskResult.from_dict(result['task_results']['merge'])
    assembly = case.fixture.restart().open_result(task.output('assembly')).read_value()
    assert [part['document_id'] for part in assembly['parts']] == ['deck-a', 'deck-b']


def test_bound_sources_still_respect_the_chat_file_limit(runtime, monkeypatch):
    search = {'step_id': 'search', 'capability_id': 'document_search', 'arguments': {'query': 'reports'}}
    bound = merge_step(kind='pdf')
    bound['inputs'] = {'sources': {'binding': binding('search', 'sources'), 'allow_partial': False}}
    settings = {'document_action_capabilities': {'merge': {'chat_max_documents': 2}}}
    case = document_case(runtime, [search, bound], settings=settings)
    module = ModuleType('functions_search')
    module.hybrid_search = lambda *args, **kwargs: [
        {'document_id': document_id, 'id': f'hit-{document_id}', 'chunk_text': 'report'}
        for document_id in ('report-a', 'report-b', 'locked')
    ]
    monkeypatch.setitem(sys.modules, 'functions_search', module)
    result = execute(runtime, case)
    assert result['status'] == 'failed'
    assert result['failure']['code'] == 'document_merge_limit_exceeded'
    assert case.reads == []


if __name__ == '__main__':
    sys.exit(pytest.main([__file__, '-q']))
