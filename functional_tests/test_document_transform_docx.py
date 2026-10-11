# test_document_transform_docx.py
"""
Functional tests for package-preserving DOCX/DOTX transformation drafts.
Version: 0.261.321
Implemented in: 0.261.321

Generated in-memory python-docx references exercise both modes, mixed text runs,
headings, tables, headers/footers, shared embedded images, cloning/removal, empty
templates and preserved unsupported XML. Outputs reopen through python-docx and the
real guarded package. No Azure, configuration bootstrap, files or network are needed.
"""

from copy import deepcopy
from dataclasses import replace
import hashlib
import io
from pathlib import Path
import sys
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from lxml import etree
from PIL import Image
import pytest

# The application's modules use flat imports; config is deliberately not imported.
APP_PATH = Path(__file__).resolve().parents[1] / "application" / "single_app"
sys.path.insert(0, str(APP_PATH))

import functions_document_transform as core
import functions_document_transform_docx as adapter


WORD_NS = adapter.WORD_NS
W = f"{{{WORD_NS}}}"
R = f"{{{adapter.OFFICE_REL_NS}}}"
HANDLE = {"kind": "document", "document_id": "authorized-word-reference"}
UNKNOWN_XML = b'<?xml version="1.0" encoding="UTF-8"?>\n<!-- preserve bytes -->\n<opaque xmlns="urn:fixture">  unknown content  </opaque>'


def image_bytes(color="blue", format_name="PNG", size=(24, 12)):
    stream = io.BytesIO()
    with Image.new("RGB", size, color) as image:
        image.save(stream, format=format_name)
    return stream.getvalue()


def archive_bytes(files):
    result = io.BytesIO()
    with zipfile.ZipFile(result, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return result.getvalue()


def unpack(content):
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def xml_bytes(root):
    return etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)


def word_reference(extension=".docx", *, empty=False):
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.8)
    section.header.paragraphs[0].text = "Reference header"
    section.header.paragraphs[0].runs[0].bold = True
    section.header.add_paragraph().add_run().add_picture(io.BytesIO(image_bytes("green")), width=Inches(0.4))
    section.footer.paragraphs[0].text = "Reference footer"
    field = section.footer.add_paragraph().add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(W + "fldCharType", "begin")
    instruction = OxmlElement("w:instrText")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(W + "fldCharType", "separate")
    cached = OxmlElement("w:t")
    cached.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(W + "fldCharType", "end")
    for node in (begin, instruction, separate, cached, end):
        field._r.append(node)
    document.add_heading("Reference report", level=1)
    mixed = document.add_paragraph()
    mixed.add_run("Revenue: ")
    mixed.add_run("{{amount}}").bold = True
    mixed.add_run(" USD").italic = True
    document.add_paragraph("Repeated label")
    document.add_paragraph("Repeated label")
    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Old cell"
    table.cell(0, 0).paragraphs[0].runs[0].bold = True
    table.cell(0, 1).text = "Untouched cell"
    table.cell(1, 0).text = "First old paragraph"
    table.cell(1, 0).add_paragraph("Second old paragraph").runs[0].italic = True
    table.cell(1, 1).text = "Final cell"
    document.add_picture(io.BytesIO(image_bytes()), width=Inches(1.2))
    document.add_picture(io.BytesIO(image_bytes()), width=Inches(0.6))
    document.add_paragraph("Remove me")
    document.add_paragraph("")
    bookmarked = document.add_paragraph("Bookmarked reference")
    start = OxmlElement("w:bookmarkStart")
    start.set(W + "id", "1")
    start.set(W + "name", "reference_bookmark")
    bookmarked._p.insert(0, start)
    end = OxmlElement("w:bookmarkEnd")
    end.set(W + "id", "1")
    bookmarked._p.append(end)
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Controlled reference"
    run.append(text)
    paragraph.append(run)
    content.append(paragraph)
    control.append(content)
    document._element.body.insert(len(document._element.body) - 1, control)
    stream = io.BytesIO()
    document.save(stream)
    files = unpack(stream.getvalue())
    root = etree.fromstring(files["word/document.xml"])
    body = root.find(W + "body")
    if empty:
        for child in list(body):
            if child.tag != W + "sectPr":
                body.remove(child)
    else:
        etree.SubElement(body, "{urn:fixture}unknown", {"metadata": "retain"}).text = "Unsupported body object"
        for index, paragraph in enumerate(body.findall(W + "p"), start=1):
            paragraph.set(f"{{{adapter.WORD14_NS}}}paraId", f"{index:08X}")
    files["word/document.xml"] = xml_bytes(root)
    files["custom/opaque.xml"] = UNKNOWN_XML
    files["custom/opaque.bin"] = b"\x00\xfforiginal-package-part"
    types = etree.fromstring(files["[Content_Types].xml"])
    for name, content_type in (
        ("custom/opaque.xml", "application/xml"), ("custom/opaque.bin", "application/octet-stream"),
    ):
        etree.SubElement(types, f"{{{adapter.CT_NS}}}Override", PartName=f"/{name}", ContentType=content_type)
    if extension == ".dotx":
        for node in types:
            if node.get("PartName") == "/word/document.xml":
                node.set("ContentType", core._MAIN_TYPES[".dotx"])
    files["[Content_Types].xml"] = xml_bytes(types)
    return archive_bytes(files)


def prepared(source, extension=".docx", *, mode="refresh", operations=None, unresolved=None):
    return {
        "schema_version": core.DOCUMENT_TRANSFORM_SCHEMA_VERSION, "output_format": "docx", "mode": mode,
        "reference": {"handle": deepcopy(HANDLE), "content_sha256": hashlib.sha256(source).hexdigest(), "extension": extension},
        "operations": operations or [], "unresolved": unresolved or [],
    }


def inspect(source, extension=".docx", **kwargs):
    return core.inspect_document(source, f"reference{extension}", HANDLE, **kwargs)


def transform(source, value, **kwargs):
    with core.transform_document(source, value, **kwargs) as result:
        output = result.file_content.read()
        report = deepcopy(result.metadata["transformation_report"])
        metadata = deepcopy(result.metadata)
    reopened = Document(io.BytesIO(output))
    guarded = core.GuardedPackage(output, ".docx")
    guarded.validate()
    return output, report, metadata, reopened


@pytest.mark.parametrize("extension", [".docx", ".dotx"])
def test_inspection_is_structural_deterministic_and_useful(extension):
    source = word_reference(extension)
    source_digest = hashlib.sha256(source).hexdigest()
    manifest = inspect(source, extension)
    second = inspect(source, extension)
    assert manifest == second
    targets = {target["id"]: target for target in manifest["targets"]}
    assert targets["word:body.p1"]["type"] == "heading"
    assert targets["word:body.p1"]["heading_level"] == 1
    assert targets["word:body.p1"]["style"] == "Heading1"
    assert targets["word:body.p2.r2"]["text"] == "{{amount}}"
    assert targets["word:body.p3"]["text"] == targets["word:body.p4"]["text"] == "Repeated label"
    assert targets["word:body.tbl1.row1.cell1"]["type"] == "table_cell"
    assert targets["word:body.tbl1.row1.cell1.p1.r1"]["text"] == "Old cell"
    assert targets["word:body.p5.image1"]["operations"] == ["replace_image"]
    assert targets["word:header1.p1"]["text"] == "Reference header"
    assert targets["word:footer1.p1"]["text"] == "Reference footer"
    assert targets["word:footer1.p2"]["operations"] == []
    assert targets["word:body.object1"]["type"] == "content_control"
    assert targets["word:body.object1"]["operations"] == []
    assert targets["word:body"]["operations"] == ["append_text"]
    assert targets["word:body.tbl1"]["operations"] == ["clone", "remove"]
    assert manifest["layouts"][0]["page_size"]["w"] == "12240"
    assert manifest["layouts"][0]["margins"]["left"] == "1152"
    assert any("Fields" in limitation for limitation in manifest["limitations"])
    assert hashlib.sha256(source).hexdigest() == source_digest


@pytest.mark.parametrize("mode", ["refresh", "create_from_reference"])
@pytest.mark.parametrize("extension", [".docx", ".dotx"])
def test_mixed_run_replacement_preserves_formatting_and_untouched_parts(mode, extension):
    source = word_reference(extension)
    files = unpack(source)
    value = prepared(source, extension, mode=mode, operations=[
        {"op": "replace_text", "target": "word:body.p2", "text": "Revenue: 42 USD"},
    ])
    output, report, metadata, reopened = transform(source, value)
    assert reopened.paragraphs[1].text == "Revenue: 42 USD"
    assert [run.text for run in reopened.paragraphs[1].runs] == ["Revenue: ", "42", " USD"]
    assert reopened.paragraphs[1].runs[1].bold is True
    assert reopened.paragraphs[1].runs[2].italic is True
    assert reopened.paragraphs[0].style.style_id == "Heading1"
    assert report["applied"][0]["op_index"] == 0
    assert any("Mixed-style" in warning for warning in report["warnings"])
    assert metadata["adapter"] == "word-package-v1"
    result_files = unpack(output)
    for name, data in files.items():
        if name == "word/document.xml" or name == "[Content_Types].xml" and extension == ".dotx":
            continue
        assert result_files[name] == data, name
    old_root = etree.fromstring(files["word/document.xml"])
    new_root = etree.fromstring(result_files["word/document.xml"])
    assert etree.tostring(old_root.find(f"{W}body/{W}sectPr")) == etree.tostring(new_root.find(f"{W}body/{W}sectPr"))
    assert etree.tostring(old_root.find(f"{W}body/{W}sdt")) == etree.tostring(new_root.find(f"{W}body/{W}sdt"))
    assert result_files["custom/opaque.xml"] == UNKNOWN_XML
    assert files == unpack(source)


def test_repeated_text_run_targets_escaping_and_line_breaks_are_not_first_match():
    source = word_reference()
    output, report, _metadata, reopened = transform(source, prepared(source, operations=[
        {"op": "replace_text", "target": "word:body.p4.r1", "text": 'Updated <literal> & "text"'},
        {"op": "append_text", "target": "word:body.p3", "text": "\tmore\nnext line"},
        {"op": "replace_text", "target": "word:body.p1.r1", "text": "New report"},
    ]))
    assert reopened.paragraphs[2].text == "Repeated label\tmore\nnext line"
    assert reopened.paragraphs[3].text == 'Updated <literal> & "text"'
    assert reopened.paragraphs[0].text == "New report"
    assert reopened.paragraphs[0].style.style_id == "Heading1"
    assert len(report["applied"]) == 3
    root = etree.fromstring(unpack(output)["word/document.xml"])
    paragraph = root.find(W + "body").findall(W + "p")[2]
    assert paragraph.find(f".//{W}tab") is not None
    assert paragraph.find(f".//{W}br") is not None


@pytest.mark.parametrize("mode", ["refresh", "create_from_reference"])
def test_cells_multi_paragraphs_and_nested_run_targets(mode):
    source = word_reference()
    files = unpack(source)
    output, report, _metadata, reopened = transform(source, prepared(source, mode=mode, operations=[
        {"op": "replace_text", "target": "word:body.tbl1.row1.cell1", "text": "Updated cell"},
        {"op": "replace_text", "target": "word:body.tbl1.row2.cell1", "text": "First new\nSecond new\nThird new"},
        {"op": "append_text", "target": "word:body.tbl1.row2.cell2.p1.r1", "text": " appended"},
    ]))
    table = reopened.tables[0]
    assert table.cell(0, 0).text == "Updated cell"
    assert table.cell(0, 0).paragraphs[0].runs[0].bold is True
    assert table.cell(0, 1).text == "Untouched cell"
    assert table.cell(1, 0).text == "First new\nSecond new\nThird new"
    assert table.cell(1, 0).paragraphs[1].runs[0].italic is True
    assert table.cell(1, 0).paragraphs[2].runs[0].italic is True
    assert table.cell(1, 1).text == "Final cell appended"
    assert len(report["applied"]) == 3
    before = etree.fromstring(files["word/document.xml"]).find(f"{W}body/{W}tbl")
    after = etree.fromstring(unpack(output)["word/document.xml"]).find(f"{W}body/{W}tbl")
    for tag in ("tblPr", "tblGrid"):
        assert etree.tostring(before.find(W + tag)) == etree.tostring(after.find(W + tag))


def test_cell_replacement_can_reduce_paragraphs_and_append_new_ones():
    source = word_reference()
    _output, report, _metadata, reopened = transform(source, prepared(source, operations=[
        {"op": "replace_text", "target": "word:body.tbl1.row2.cell1", "text": "Only one"},
        {"op": "append_text", "target": "word:body.tbl1.row2.cell1", "text": "\nNew second"},
    ]))
    assert reopened.tables[0].cell(1, 0).text == "Only one\nNew second"
    assert len(reopened.tables[0].cell(1, 0).paragraphs) == 2
    assert len(report["applied"]) == 2


def test_headers_and_footer_change_only_when_explicitly_targeted():
    source = word_reference()
    files = unpack(source)
    output, report, _metadata, reopened = transform(source, prepared(source, operations=[
        {"op": "replace_text", "target": "word:header1.p1", "text": "New header"},
        {"op": "append_text", "target": "word:footer1.p1", "text": " updated"},
        {"op": "replace_text", "target": "word:footer1.p2", "text": "Not a page field"},
    ]))
    assert reopened.sections[0].header.paragraphs[0].text == "New header"
    assert reopened.sections[0].header.paragraphs[0].runs[0].bold is True
    assert reopened.sections[0].footer.paragraphs[0].text == "Reference footer updated"
    assert reopened.sections[0].footer.paragraphs[1].text == "1"
    assert [entry["op_index"] for entry in report["applied"]] == [0, 1]
    assert report["skipped"][0]["op_index"] == 2
    assert unpack(output)["word/document.xml"] == files["word/document.xml"]


@pytest.mark.parametrize("mode", ["refresh", "create_from_reference"])
@pytest.mark.parametrize("format_name", ["PNG", "JPEG"])
def test_image_replacement_is_slot_local_preserves_geometry_and_shared_source(mode, format_name):
    source = word_reference()
    before = core.GuardedPackage(source, ".docx")
    root = before.parse_xml(before.main_part)
    pictures = [node for node in root.iter() if adapter._qname(node)[0] in adapter.DRAWING_NAMESPACES
                and adapter._qname(node)[1] == "blip"]
    assert pictures[0].get(R + "embed") == pictures[1].get(R + "embed")
    old_rid = pictures[1].get(R + "embed")
    old_part = {rel["id"]: rel["target"] for rel in before.relationships(before.main_part)}[old_rid]
    calls = []
    new_image = image_bytes("red", format_name=format_name)

    def resolver(token):
        calls.append(token)
        return new_image

    output, report, _metadata, reopened = transform(source, prepared(source, mode=mode, operations=[
        {"op": "replace_image", "target": "word:body.p5.image1", "source": "asset:red"},
    ]), image_resolver=resolver)
    assert calls == ["asset:red"]
    assert len(reopened.inline_shapes) == 2
    assert reopened.inline_shapes[0].width == Inches(1.2)
    assert reopened.inline_shapes[1].width == Inches(0.6)
    guarded = core.GuardedPackage(output, ".docx")
    updated = [node for node in guarded.parse_xml(guarded.main_part).iter() if etree.QName(node).localname == "blip"]
    assert updated[0].get(R + "embed") != old_rid
    assert updated[1].get(R + "embed") == old_rid
    related = {rel["id"]: rel["target"] for rel in guarded.relationships(guarded.main_part)}
    assert guarded.files[related[updated[0].get(R + "embed")]] == new_image
    assert guarded.files[old_part] == before.files[old_part]
    assert guarded.files["word/header1.xml"] == before.files["word/header1.xml"]
    assert report["applied"][0]["op_index"] == 0


def test_clone_picture_ids_and_relationships_stay_valid_and_new_slot_is_editable():
    source = word_reference()
    output, report, metadata, reopened = transform(source, prepared(source, mode="create_from_reference", operations=[
        {"op": "clone", "target": "word:body.p5", "clone_id": "new-picture"},
        {"op": "replace_image", "target": "new-picture.image1", "source": "asset:red"},
        {"op": "remove", "target": "word:body.p6"},
    ]), image_resolver=lambda token: image_bytes("red"))
    assert len(report["applied"]) == 3
    assert len(reopened.inline_shapes) == 2
    assert metadata["clones"][0]["id"] == "new-picture"
    assert "new-picture.image1" in metadata["clones"][0]["targets"]
    root = etree.fromstring(unpack(output)["word/document.xml"])
    drawing_ids = [node.get("id") for node in root.iter()
                   if adapter._qname(node)[0] in adapter.WP_NAMESPACES and adapter._qname(node)[1] == "docPr"]
    assert len(drawing_ids) == len(set(drawing_ids))
    paragraph_ids = [node.get(f"{{{adapter.WORD14_NS}}}paraId") for node in root.iter(W + "p")
                     if node.get(f"{{{adapter.WORD14_NS}}}paraId") is not None]
    assert len(paragraph_ids) == len(set(paragraph_ids))


@pytest.mark.parametrize("mode", ["refresh", "create_from_reference"])
def test_reusable_heading_and_table_clones_removal_and_batch_target_stability(mode):
    source = word_reference()
    output, report, metadata, reopened = transform(source, prepared(source, mode=mode, operations=[
        {"op": "clone", "target": "word:body.p1", "clone_id": "new-heading"},
        {"op": "replace_text", "target": "new-heading", "text": "New findings"},
        {"op": "clone", "target": "word:body.tbl1", "clone_id": "new-table"},
        {"op": "replace_text", "target": "new-table.row1.cell1", "text": "New report cell"},
        {"op": "remove", "target": "word:body.p7"},
        {"op": "replace_text", "target": "word:body.p8", "text": "Filled previously empty paragraph"},
        {"op": "replace_text", "target": "word:body.p7.r1", "text": "Removed descendant"},
    ]))
    assert reopened.paragraphs[0].text == "Reference report"
    assert reopened.paragraphs[1].text == "New findings"
    assert reopened.paragraphs[1].style.style_id == "Heading1"
    assert reopened.tables[0].cell(0, 0).text == "Old cell"
    assert reopened.tables[1].cell(0, 0).text == "New report cell"
    assert reopened.tables[1].style.style_id == "TableGrid"
    assert "Remove me" not in [paragraph.text for paragraph in reopened.paragraphs]
    assert "Filled previously empty paragraph" in [paragraph.text for paragraph in reopened.paragraphs]
    assert len(report["applied"]) == 6
    assert report["skipped"][0]["op_index"] == 6
    assert [entry["id"] for entry in metadata["clones"]] == ["new-heading", "new-table"]
    assert unpack(output)["custom/opaque.xml"] == UNKNOWN_XML


@pytest.mark.parametrize("mode", ["refresh", "create_from_reference"])
@pytest.mark.parametrize("extension", [".docx", ".dotx"])
def test_empty_reference_has_insertion_anchor_and_normalizes_to_openable_docx(mode, extension):
    source = word_reference(extension, empty=True)
    manifest = inspect(source, extension)
    assert not any(target["id"].startswith("word:body.p") for target in manifest["targets"])
    output, report, _metadata, reopened = transform(source, prepared(source, extension, mode=mode, operations=[
        {"op": "append_text", "target": "word:body", "text": "A new report\nPrepared findings"},
    ]))
    assert [paragraph.text for paragraph in reopened.paragraphs] == ["A new report", "Prepared findings"]
    assert len(report["applied"]) == 1
    assert reopened.sections[0].header.paragraphs[0].text == "Reference header"
    guarded = core.GuardedPackage(output, ".docx")
    assert guarded.main_content_type == core._MAIN_TYPES[".docx"]
    assert unpack(source)["word/document.xml"] != unpack(output)["word/document.xml"]


def test_removing_last_body_block_retains_valid_blank_paragraph_and_section():
    source = word_reference(empty=True)
    files = unpack(source)
    root = etree.fromstring(files["word/document.xml"])
    body = root.find(W + "body")
    paragraph = etree.Element(W + "p")
    etree.SubElement(etree.SubElement(paragraph, W + "r"), W + "t").text = "Only block"
    body.insert(0, paragraph)
    files["word/document.xml"] = xml_bytes(root)
    source = archive_bytes(files)
    output, report, _metadata, reopened = transform(source, prepared(source, operations=[
        {"op": "remove", "target": "word:body.p1"},
    ]))
    assert len(report["applied"]) == 1
    assert [paragraph.text for paragraph in reopened.paragraphs] == [""]
    assert etree.fromstring(unpack(output)["word/document.xml"]).find(f"{W}body/{W}sectPr") is not None


@pytest.mark.parametrize("mode", ["refresh", "create_from_reference"])
def test_unsupported_and_unknown_targets_yield_unchanged_openable_draft(mode):
    source = word_reference()
    operations = [
        {"op": "replace_text", "target": "word:body.object1", "text": "Control value"},
        {"op": "replace_text", "target": "word:body.object2", "text": "Unknown XML"},
        {"op": "remove", "target": "word:body.p9"},
        {"op": "clone", "target": "word:body.p9", "clone_id": "bookmark-copy"},
        {"op": "remove", "target": "word:header1.p1"},
        {"op": "remove", "target": "word:body.tbl1.row1.cell1"},
        {"op": "replace_text", "target": "not-a-real-target", "text": "Unavailable"},
        {"op": "replace_text", "target": "word:body.p5", "text": "Don't lose the image"},
        {"op": "replace_text", "target": "word:footer1.p2.r1", "text": "Don't break field"},
    ]
    output, report, _metadata, reopened = transform(source, prepared(source, mode=mode, operations=operations))
    assert len(report["skipped"]) == len(operations)
    assert not report["applied"]
    assert report["status"] == "draft"
    assert all(entry["reason"] for entry in report["skipped"])
    assert any("No changes applied" in warning for warning in report["warnings"])
    assert unpack(output) == unpack(source)
    assert len(reopened.inline_shapes) == 2


@pytest.mark.parametrize("mode", ["refresh", "create_from_reference"])
def test_unresolved_target_is_preserved_even_with_conflicting_operation(mode):
    source = word_reference()
    value = prepared(source, mode=mode, operations=[
        {"op": "replace_text", "target": "word:body.p2", "text": "Invented revenue"},
    ], unresolved=[{"target": "word:body.p2", "reason": "No verified revenue figure is available."}])
    output, report, _metadata, reopened = transform(source, value)
    assert reopened.paragraphs[1].text == "Revenue: {{amount}} USD"
    assert not report["applied"]
    assert report["unresolved"] == value["unresolved"]
    assert report["unchanged"][0]["op_index"] == 0
    assert any("No changes applied" in warning for warning in report["warnings"])
    if mode == "create_from_reference":
        assert any("not newly verified" in warning for warning in report["warnings"])
    assert unpack(output) == unpack(source)


def test_supported_changes_survive_independent_unsupported_operation():
    source = word_reference()
    _output, report, _metadata, reopened = transform(source, prepared(source, operations=[
        {"op": "replace_text", "target": "word:body.object1", "text": "Unsupported"},
        {"op": "replace_text", "target": "word:body.p1", "text": "Useful draft"},
    ]))
    assert reopened.paragraphs[0].text == "Useful draft"
    assert report["status"] == "draft"
    assert report["applied"][0]["op_index"] == 1
    assert report["skipped"][0]["op_index"] == 0


def test_noops_and_clone_identifier_collision_are_not_false_successes():
    source = word_reference()
    output, report, _metadata, reopened = transform(source, prepared(source, operations=[
        {"op": "replace_text", "target": "word:body.p1", "text": "Reference report"},
        {"op": "append_text", "target": "word:body.p3", "text": ""},
        {"op": "replace_image", "target": "word:body.p5.image1", "source": "asset:original"},
        {"op": "clone", "target": "word:body.p1", "clone_id": "word:body.p2"},
    ]), image_resolver=lambda token: image_bytes())
    assert not report["applied"]
    assert [entry["op_index"] for entry in report["unchanged"] if "op_index" in entry] == [0, 1, 2]
    assert report["skipped"][0]["op_index"] == 3
    assert len(reopened.paragraphs) == 9
    assert unpack(output) == unpack(source)


@pytest.mark.parametrize("asset", [None, b"", b"not image bytes", "C:\\private\\image.png", {"url": "https://example.invalid/image"}])
def test_bad_or_missing_image_assets_are_skipped_without_losing_text_updates(asset):
    source = word_reference()
    _output, report, _metadata, reopened = transform(source, prepared(source, operations=[
        {"op": "replace_image", "target": "word:body.p5.image1", "source": "asset:bad"},
        {"op": "replace_text", "target": "word:body.p1", "text": "Useful draft"},
    ]), image_resolver=lambda token: asset)
    assert report["skipped"][0]["op_index"] == 0
    assert report["applied"][0]["op_index"] == 1
    assert reopened.paragraphs[0].text == "Useful draft"


def test_image_resolver_is_never_called_for_unknown_slots_or_non_asset_sources():
    source = word_reference()
    calls = []
    _output, report, _metadata, _reopened = transform(source, prepared(source, operations=[
        {"op": "replace_image", "target": "unknown", "source": "asset:image"},
        {"op": "replace_image", "target": "word:body.p5.image1", "source": "opaque-but-not-asset"},
        {"op": "replace_image", "target": "word:body.p1", "source": "asset:image"},
    ]), image_resolver=lambda token: calls.append(token))
    assert calls == []
    assert len(report["skipped"]) == 3


def test_authorization_and_callback_failures_propagate_without_downgrading_to_skips():
    source = word_reference()
    failure = PermissionError("private provider details")

    def denied(token):
        raise failure

    value = prepared(source, operations=[
        {"op": "replace_image", "target": "word:body.p5.image1", "source": "asset:image"},
    ])
    with pytest.raises(PermissionError) as error:
        core.transform_document(source, value, image_resolver=denied)
    assert error.value is failure
    with pytest.raises(core.DocumentTransformCancelled):
        adapter.inspect_package(core.GuardedPackage(source, ".docx"), limits=core.DocumentTransformLimits(), check=lambda: False)


def test_resource_limits_for_targets_and_replacement_images():
    source = word_reference()
    with pytest.raises(core.DocumentTransformError) as error:
        inspect(source, limits=replace(core.DocumentTransformLimits(), max_targets=2))
    assert error.value.code == "resource_limit"
    package = core.GuardedPackage(source, ".docx")
    limits = replace(core.DocumentTransformLimits(), max_part_bytes=10)
    value = prepared(source, operations=[
        {"op": "replace_image", "target": "word:body.p5.image1", "source": "asset:image"},
    ])
    with pytest.raises(core.DocumentTransformError) as error:
        adapter.transform_package(package, value, limits=limits, check=None, image_resolver=lambda token: image_bytes())
    assert error.value.code == "resource_limit"


def test_long_clone_ids_have_bounded_deterministic_descendant_aliases():
    source = word_reference()
    clone_id = "c" * 128
    value = prepared(source, operations=[{"op": "clone", "target": "word:body.p1", "clone_id": clone_id}])
    first, report, metadata, _reopened = transform(source, value)
    second, _report, second_metadata, _reopened = transform(source, value)
    assert first == second
    assert len(report["applied"]) == 1
    assert metadata["clones"] == second_metadata["clones"]
    assert clone_id in metadata["clones"][0]["targets"]
    assert all(len(identifier) <= 128 for identifier in metadata["clones"][0]["targets"])


def test_clone_descendant_collision_is_skipped_atomically():
    source = word_reference()
    _output, report, metadata, reopened = transform(source, prepared(source, operations=[
        {"op": "clone", "target": "word:body.p1", "clone_id": "copy.r1"},
        {"op": "clone", "target": "word:body.p1", "clone_id": "copy"},
        {"op": "replace_text", "target": "copy.r1", "text": "Preserved successful clone"},
    ]))
    assert [entry["op_index"] for entry in report["applied"]] == [0, 2]
    assert report["skipped"][0]["op_index"] == 1
    assert "descendant" in report["skipped"][0]["reason"]
    assert len(metadata["clones"]) == 1
    assert len(reopened.paragraphs) == 10
    assert reopened.paragraphs[1].text == "Preserved successful clone"


def test_unresolved_ancestors_and_descendants_cannot_be_overwritten_or_removed():
    source = word_reference()
    _output, report, _metadata, reopened = transform(source, prepared(source, operations=[
        {"op": "replace_text", "target": "word:body.p2.r2", "text": "Invented"},
        {"op": "remove", "target": "word:body.tbl1"},
        {"op": "replace_text", "target": "word:body.p3", "text": "Changed parent"},
        {"op": "append_text", "target": "word:body", "text": "Independent draft addition"},
    ], unresolved=[
        {"target": "word:body.p2", "reason": "Revenue is not verified."},
        {"target": "word:body.tbl1.row1.cell1", "reason": "No replacement cell value."},
        {"target": "word:body.p3.r1", "reason": "A value is missing."},
    ]))
    assert reopened.paragraphs[1].text == "Revenue: {{amount}} USD"
    assert reopened.paragraphs[2].text == "Repeated label"
    assert reopened.tables[0].cell(0, 0).text == "Old cell"
    assert reopened.paragraphs[-1].text == "Independent draft addition"
    assert [entry["op_index"] for entry in report["unchanged"] if "op_index" in entry] == [0, 1, 2]
    assert report["applied"][0]["op_index"] == 3


def test_field_results_spanning_paragraphs_are_preserved():
    source = word_reference()
    files = unpack(source)
    root = etree.fromstring(files["word/document.xml"])
    paragraphs = root.find(W + "body").findall(W + "p")
    first_run = paragraphs[2].find(W + "r")
    begin = etree.Element(W + "fldChar", {W + "fldCharType": "begin"})
    first_run.insert(0, begin)
    instruction = etree.Element(W + "instrText")
    instruction.text = " REF reference_bookmark "
    first_run.insert(1, instruction)
    first_run.insert(2, etree.Element(W + "fldChar", {W + "fldCharType": "separate"}))
    paragraphs[3].find(W + "r").append(etree.Element(W + "fldChar", {W + "fldCharType": "end"}))
    files["word/document.xml"] = xml_bytes(root)
    source = archive_bytes(files)
    manifest = inspect(source)
    targets = {target["id"]: target for target in manifest["targets"]}
    assert targets["word:body.p3"]["operations"] == []
    assert targets["word:body.p4.r1"]["operations"] == []
    output, report, _metadata, reopened = transform(source, prepared(source, operations=[
        {"op": "replace_text", "target": "word:body.p4.r1", "text": "Wrong cached result"},
        {"op": "clone", "target": "word:body.p3", "clone_id": "field-copy"},
        {"op": "replace_text", "target": "word:body.p1", "text": "Independent heading"},
    ]))
    assert [entry["op_index"] for entry in report["skipped"]] == [0, 1]
    assert report["applied"][0]["op_index"] == 2
    assert reopened.paragraphs[3].text == "Repeated label"
    changed = etree.fromstring(unpack(output)["word/document.xml"]).find(W + "body").findall(W + "p")
    assert etree.tostring(changed[2]) == etree.tostring(paragraphs[2])
    assert etree.tostring(changed[3]) == etree.tostring(paragraphs[3])


def test_section_break_blocks_and_extension_identities_are_not_cloned_or_removed():
    source = word_reference()
    files = unpack(source)
    root = etree.fromstring(files["word/document.xml"])
    body = root.find(W + "body")
    paragraphs = body.findall(W + "p")
    properties = etree.Element(W + "pPr")
    properties.append(deepcopy(body.find(W + "sectPr")))
    paragraphs[2].insert(0, properties)
    paragraphs[6].set("{urn:unknown-extension}uniqueId", "preserve-identity")
    files["word/document.xml"] = xml_bytes(root)
    source = archive_bytes(files)
    output, report, _metadata, reopened = transform(source, prepared(source, operations=[
        {"op": "clone", "target": "word:body.p3", "clone_id": "section-copy"},
        {"op": "remove", "target": "word:body.p3"},
        {"op": "remove", "target": "word:body.p7"},
        {"op": "replace_text", "target": "word:body.p3", "text": "Safe section heading"},
    ]))
    assert [entry["op_index"] for entry in report["skipped"]] == [0, 1, 2]
    assert report["applied"][0]["op_index"] == 3
    assert reopened.paragraphs[2].text == "Safe section heading"
    assert len(reopened.sections) == 2
    changed = etree.fromstring(unpack(output)["word/document.xml"])
    assert len(list(changed.iter(W + "sectPr"))) == 2
    assert changed.find(W + "body").findall(W + "p")[6].get("{urn:unknown-extension}uniqueId") == "preserve-identity"


def test_nested_table_whole_cell_edits_skip_but_nested_paragraphs_and_clones_work():
    source = word_reference()
    files = unpack(source)
    document = Document(io.BytesIO(source))
    nested = document.tables[0].cell(0, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested value"
    files["word/document.xml"] = xml_bytes(document._element)
    source = archive_bytes(files)
    output, report, _metadata, reopened = transform(source, prepared(source, operations=[
        {"op": "replace_text", "target": "word:body.tbl1.row1.cell2", "text": "Would discard nested table"},
        {"op": "replace_text", "target": "word:body.tbl1.row1.cell2.tbl1.row1.cell1.p1", "text": "New nested value"},
        {"op": "clone", "target": "word:body.tbl1", "clone_id": "nested-copy"},
        {"op": "replace_text", "target": "nested-copy.row1.cell2.tbl1.row1.cell1", "text": "New cloned value"},
    ]))
    assert report["skipped"][0]["op_index"] == 0
    assert len(report["applied"]) == 3
    assert reopened.tables[0].cell(0, 1).tables[0].cell(0, 0).text == "New nested value"
    assert reopened.tables[1].cell(0, 1).tables[0].cell(0, 0).text == "New cloned value"
    assert unpack(output)["custom/opaque.xml"] == UNKNOWN_XML


def test_hyperlinks_survive_text_editing_cloning_and_removal_without_new_fetches():
    source = word_reference()
    files = unpack(source)
    root = etree.fromstring(files["word/document.xml"])
    body = root.find(W + "body")
    paragraph = etree.Element(W + "p")
    hyperlink = etree.SubElement(paragraph, W + "hyperlink", {R + "id": "rExternalLink"})
    run = etree.SubElement(hyperlink, W + "r")
    etree.SubElement(run, W + "t").text = "Original linked label"
    body.insert(body.index(body.find(W + "sectPr")), paragraph)
    files["word/document.xml"] = xml_bytes(root)
    rels = etree.fromstring(files["word/_rels/document.xml.rels"])
    etree.SubElement(rels, f"{{{adapter.REL_NS}}}Relationship", {
        "Id": "rExternalLink", "Type": f"{adapter.OFFICE_REL_NS}/hyperlink",
        "Target": "https://example.invalid/inert-reference", "TargetMode": "External",
    })
    files["word/_rels/document.xml.rels"] = xml_bytes(rels)
    source = archive_bytes(files)
    output, report, _metadata, _reopened = transform(source, prepared(source, operations=[
        {"op": "replace_text", "target": "word:body.p10", "text": "New linked label"},
        {"op": "clone", "target": "word:body.p10", "clone_id": "linked-copy"},
        {"op": "remove", "target": "word:body.p10"},
    ]))
    assert len(report["applied"]) == 3
    guarded = core.GuardedPackage(output, ".docx")
    links = guarded.parse_xml(guarded.main_part).findall(f".//{W}hyperlink")
    assert len(links) == 1
    assert links[0].get(R + "id") == "rExternalLink"
    assert links[0].find(f"{W}r/{W}t").text == "New linked label"
    assert guarded.files["word/_rels/document.xml.rels"] == files["word/_rels/document.xml.rels"]


def test_header_image_is_replaceable_without_modifying_body_or_footer():
    source = word_reference()
    output, report, _metadata, reopened = transform(source, prepared(source, operations=[
        {"op": "replace_image", "target": "word:header1.p2.image1", "source": "asset:new-header"},
    ]), image_resolver=lambda token: image_bytes("yellow"))
    before, after = unpack(source), unpack(output)
    assert after["word/document.xml"] == before["word/document.xml"]
    assert after["word/footer1.xml"] == before["word/footer1.xml"]
    assert after["word/header1.xml"] != before["word/header1.xml"]
    assert len(report["applied"]) == 1
    assert reopened.sections[0].header.paragraphs[0].text == "Reference header"


def test_linked_images_are_preserved_and_never_resolved():
    source = word_reference()
    files = unpack(source)
    root = etree.fromstring(files["word/document.xml"])
    blip = next(node for node in root.iter() if adapter._qname(node)[1] == "blip")
    blip.set(R + "link", "rLinkedImage")
    files["word/document.xml"] = xml_bytes(root)
    rels = etree.fromstring(files["word/_rels/document.xml.rels"])
    etree.SubElement(rels, f"{{{adapter.REL_NS}}}Relationship", {
        "Id": "rLinkedImage", "Type": f"{adapter.OFFICE_REL_NS}/image",
        "Target": "https://example.invalid/inert-image.png", "TargetMode": "External",
    })
    files["word/_rels/document.xml.rels"] = xml_bytes(rels)
    source = archive_bytes(files)
    calls = []
    output, report, _metadata, _reopened = transform(source, prepared(source, operations=[
        {"op": "replace_image", "target": "word:body.p5.image1", "source": "asset:replacement"},
    ]), image_resolver=lambda token: calls.append(token))
    assert calls == []
    assert len(report["skipped"]) == 1
    assert "Linked" in report["skipped"][0]["reason"]
    assert unpack(source) == unpack(output)


def test_append_respects_combined_text_limit():
    source = word_reference()
    value = prepared(source, operations=[
        {"op": "append_text", "target": "word:body.p3", "text": "a" * 191},
    ])
    with pytest.raises(core.DocumentTransformError) as error:
        core.transform_document(source, value, limits=replace(core.DocumentTransformLimits(), max_text_chars=200))
    assert error.value.code == "resource_limit"


def test_image_pixel_limits_and_non_permission_callback_errors_propagate():
    source = word_reference()
    package = core.GuardedPackage(source, ".docx")
    value = prepared(source, operations=[
        {"op": "replace_image", "target": "word:body.p5.image1", "source": "asset:image"},
    ])
    with pytest.raises(core.DocumentTransformError) as error:
        adapter.transform_package(
            package, value, limits=replace(core.DocumentTransformLimits(), max_package_uncompressed_bytes=1000),
            check=None, image_resolver=lambda token: image_bytes(),
        )
    assert error.value.code == "resource_limit"
    failure = OSError("private callback failure")

    def failed_resolver(token):
        raise failure

    with pytest.raises(OSError) as error:
        core.transform_document(source, value, image_resolver=failed_resolver)
    assert error.value is failure


def test_unsupported_word_wrappers_are_preserved_not_structurally_duplicated():
    source = word_reference()
    files = unpack(source)
    root = etree.fromstring(files["word/document.xml"])
    body = root.find(W + "body")
    paragraph = etree.Element(W + "p")
    wrapper = etree.SubElement(paragraph, W + "smartTag", {W + "uri": "urn:fixture", W + "element": "fact"})
    etree.SubElement(etree.SubElement(wrapper, W + "r"), W + "t").text = "Smart-tagged reference"
    body.insert(body.index(body.find(W + "sectPr")), paragraph)
    files["word/document.xml"] = xml_bytes(root)
    source = archive_bytes(files)
    output, report, _metadata, _reopened = transform(source, prepared(source, operations=[
        {"op": "clone", "target": "word:body.p10", "clone_id": "unsupported-copy"},
        {"op": "remove", "target": "word:body.p10"},
        {"op": "replace_text", "target": "word:body.p1", "text": "Useful ordinary heading"},
    ]))
    assert [entry["op_index"] for entry in report["skipped"]] == [0, 1]
    assert report["applied"][0]["op_index"] == 2
    new_root = etree.fromstring(unpack(output)["word/document.xml"])
    assert etree.tostring(new_root.find(f".//{W}smartTag")) == etree.tostring(wrapper)


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
