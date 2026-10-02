#!/usr/bin/env python3
# test_orchestration_tabular_merge_capability.py
"""
Functional test for the tabular_merge chat orchestration capability.
Version: 0.261.218
Implemented in: 0.261.218

This test ensures that tabular_merge is a gated Reason capability alongside Analyze
and Compare: plans name its sources explicitly or bind a search's source set, only
tabular sources are admitted, the administrator's Merge limits apply, the real
executor retains the merged rows and a merge report without any model call, merge
failures reach users only as application-owned messages, and the retained rows
render to CSV and XLSX through the exact-schema export profiles.
"""

import csv
import io
import sys
from types import ModuleType

import pytest

from test_orchestration_dependency_runtime import binding, execute, runtime  # noqa: F401
from test_support.versioning import assert_app_version_at_least


AVAILABLE = ['tabular_merge', 'document_search', 'compose', 'render_file']
FILES = {
    'east-doc': ('east.csv', b'Region,Code,Amount\r\nEast,007,10.50\r\nEast,008,"1,200"\r\n'),
    'west-doc': ('west.csv', b'Amount;Region;Code\n5;West;0099\n'),
    'notes-doc': ('notes.docx', b'not a spreadsheet'),
    'other-doc': ('other.csv', b'Region,Total\nNorth,1\n'),
}


def merge_step(document_ids=None, **arguments):
    step = {'step_id': 'merge', 'capability_id': 'tabular_merge', 'arguments': dict(arguments)}
    if document_ids is not None:
        step['arguments']['document_ids'] = list(document_ids)
    return step


def tabular_case(runtime, steps, *, settings=None, available=AVAILABLE):
    case = runtime.make(steps, settings=settings, available=available)
    for document_id, (file_name, _) in FILES.items():
        case.fixture.sources[document_id] = {
            'document_id': document_id, 'scope': 'personal', 'scope_id': 'owner',
            'source_version': 1, 'source_revision': f'{document_id}-revision',
            'file_name': file_name,
            'source_kind': 'tabular' if file_name.endswith('.csv') else 'narrative',
        }
    reads = []

    def reader(source, user_id, group_id=None, public_workspace_id=None, *, purpose):
        assert user_id == 'owner' and purpose == 'native'
        reads.append(source['document_id'])
        return {'id': source['document_id']}, FILES[source['document_id']][1]

    case.context.merge_source_reader = reader
    case.reads = reads
    return case


def test_version_includes_tabular_merge():
    assert_app_version_at_least('0.261.218')


def test_tabular_merge_is_a_gated_reason_capability(runtime):
    registry = runtime.registry
    capability = registry.get_capability('tabular_merge')
    assert capability['role'] == registry.ROLE_REASON
    assert capability['label'] == 'Merge spreadsheets'
    assert capability['document_action_type'] == 'merge'
    assert capability['result_outputs'] == {'records': 'records-v1', 'report': 'structured-v1'}
    assert capability['result_input_kinds'] == {'sources': ('source-set-v1',)}
    assert capability['inputs']['properties']['document_ids']['minItems'] == 2
    assert capability['adapter'] == 'tabular_merge'
    assert 'tabular_merge' in runtime.contracts.REASON_CAPABILITIES

    settings = {'enable_chat_orchestration': True, 'enable_user_workspace': True}
    available = {item['id'] for item in registry.resolve_available_capabilities(settings)}
    assert 'tabular_merge' in available
    disabled = {**settings, 'document_action_capabilities': {'merge': {'enabled': False}}}
    assert 'tabular_merge' not in {item['id'] for item in registry.resolve_available_capabilities(disabled)}
    no_workspaces = {'enable_chat_orchestration': True}
    assert 'tabular_merge' not in {item['id'] for item in registry.resolve_available_capabilities(no_workspaces)}

    projection = registry.build_capability_client_projection([capability])
    assert projection == [{
        'id': 'tabular_merge', 'label': 'Merge spreadsheets', 'role': 'reason',
        'summary': capability['summary'], 'cost': capability['cost_class'],
    }]
    assert registry.get_capability_document_limit(capability, settings=settings) == 10
    narrowed = {**settings, 'document_action_capabilities': {'merge': {'chat_max_documents': 4}}}
    assert registry.get_capability_document_limit(capability, settings=narrowed) == 4


def test_plans_name_sources_explicitly_or_bind_a_source_set(runtime):
    schema = runtime.schema
    plan = runtime.make([merge_step(['east-doc', 'west-doc'])], available=AVAILABLE).plan
    arguments = plan['steps'][0]['arguments']
    assert arguments == {
        'document_ids': ['east-doc', 'west-doc'], 'doc_scope': 'all', 'schema_policy': 'by_name',
        'include_source_column': True, 'source_column_name': 'Source File',
    }
    assert [output['name'] for output in plan['steps'][0]['outputs']] == ['records', 'report']

    with pytest.raises(schema.PlanValidationError) as missing:
        runtime.make([merge_step()], available=AVAILABLE)
    assert missing.value.code == 'source_binding_required'

    search = {'step_id': 'search', 'capability_id': 'document_search', 'arguments': {'query': 'Q3 sales files'}}
    bound = merge_step()
    bound['inputs'] = {'sources': {'binding': binding('search', 'sources'), 'allow_partial': False}}
    plan = runtime.make([search, bound], available=AVAILABLE).plan
    assert plan['steps'][1]['depends_on'] == ['search']

    both = merge_step(['east-doc', 'west-doc'])
    both['inputs'] = bound['inputs']
    with pytest.raises(schema.PlanValidationError):
        runtime.make([search, both], available=AVAILABLE)
    with pytest.raises(schema.PlanValidationError):
        runtime.make([merge_step(['east-doc'])], available=AVAILABLE)
    with pytest.raises(schema.PlanValidationError):
        runtime.make([merge_step(['east-doc', 'west-doc'], schema_policy='union')], available=AVAILABLE)


def test_the_chat_document_limit_and_source_kinds_are_enforced_at_planning(runtime):
    schema = runtime.schema
    ids = [f'file-{index}' for index in range(4)]
    settings = {'document_action_capabilities': {'merge': {'chat_max_documents': 3}}}
    with pytest.raises(schema.PlanValidationError) as limited:
        runtime.make([merge_step(ids)], settings=settings, available=AVAILABLE)
    assert 'document limit' in str(limited.value)

    plan = runtime.make([merge_step(['east-doc', 'notes-doc'])], available=AVAILABLE).plan
    with pytest.raises(schema.PlanValidationError) as kinds:
        schema.validate_plan_document_source_kinds(plan, {'east-doc': 'tabular', 'notes-doc': 'narrative'})
    assert kinds.value.code == 'source_kind_invalid'
    schema.validate_plan_document_source_kinds(plan, {'east-doc': 'tabular'})
    schema.validate_plan_document_source_kinds(plan, {'east-doc': 'tabular', 'notes-doc': 'tabular'})


def test_merge_then_exact_render_plans_validate_without_naming_columns(runtime):
    render = {
        'step_id': 'save', 'capability_id': 'render_file',
        'arguments': {'file_name': 'merged.xlsx', 'output_format': 'xlsx', 'profile': 'exact_tabular_workbook_v1',
                      'options': {}},
        'inputs': {'source': {'binding': binding('merge', 'records'), 'allow_partial': False}},
        'outputs': [],
    }
    plan = runtime.make([merge_step(['east-doc', 'west-doc']), render], available=AVAILABLE).plan
    assert plan['steps'][1]['depends_on'] == ['merge']
    csv_render = {**render, 'arguments': {
        'file_name': 'merged.csv', 'output_format': 'csv', 'profile': 'exact_tabular_records_v1', 'options': {},
    }}
    runtime.make([merge_step(['east-doc', 'west-doc']), csv_render], available=AVAILABLE)
    with pytest.raises(runtime.schema.PlanValidationError):
        runtime.make([merge_step(['east-doc', 'west-doc']), {**render, 'arguments': {
            'file_name': 'merged.csv', 'output_format': 'csv', 'profile': 'tabular_records_v1', 'options': {},
        }}], available=AVAILABLE)


def test_executor_retains_exact_rows_and_a_report_without_a_model(runtime):
    case = tabular_case(runtime, [merge_step(['east-doc', 'west-doc'])])
    result = execute(runtime, case)
    assert result['status'] == 'completed', result
    assert case.model.calls == []
    assert case.reads == ['east-doc', 'west-doc']
    step = result['steps'][0] if isinstance(result.get('steps'), list) else None
    if step is not None:
        assert 'Merged 3 row(s) from 2 file(s)' in step.get('summary', '')

    task = runtime.contracts.TaskResult.from_dict(result['task_results']['merge'])
    assert task.status == 'complete'
    service = case.fixture.restart()
    rows = list(service.open_result(task.output('records')).iter_records())
    assert rows == [
        {'Source File': 'east.csv', 'Region': 'East', 'Code': '007', 'Amount': '10.50'},
        {'Source File': 'east.csv', 'Region': 'East', 'Code': '008', 'Amount': '1,200'},
        {'Source File': 'west.csv', 'Region': 'West', 'Code': '0099', 'Amount': '5'},
    ]
    records = task.output('records')
    assert records.item_count == 3
    assert [column.name for column in records.columns] == ['Source File', 'Region', 'Code', 'Amount']
    assert records.completeness.coverage.unit == 'sources'
    report = service.open_result(task.output('report')).read_value()
    assert report['status'] == 'merged'
    assert report['totals'] == {'sources': 2, 'rows': 3, 'columns': 4, 'blank_rows_skipped': 0}
    assert [entry['file_name'] for entry in report['sources']] == ['east.csv', 'west.csv']
    metadata = service.open_result(task.output('records')).metadata()
    assert metadata['origin'] == 'grounded'


def test_merge_failures_use_application_owned_messages(runtime):
    case = tabular_case(runtime, [merge_step(['east-doc', 'other-doc'])])
    result = execute(runtime, case)
    assert result['status'] == 'failed'
    assert result['failure']['code'] == 'merge_schema_mismatch'
    assert 'other.csv' not in result['message'] and 'Total' not in result['message']
    assert "don't all have the same columns" in result['message']

    case = tabular_case(runtime, [merge_step(['east-doc', 'notes-doc'])])
    result = execute(runtime, case)
    assert result['status'] == 'failed'
    assert result['failure']['code'] == 'merge_sources_invalid'
    assert case.reads == []


def test_merge_reads_are_refused_when_a_source_is_unavailable(runtime):
    from functions_orchestration_results import ResultUnavailableError

    case = tabular_case(runtime, [merge_step(['east-doc', 'west-doc'])])
    case.fixture.denied.add('west-doc')
    # The executor re-authorizes every planned source before any step runs.
    with pytest.raises(ResultUnavailableError):
        execute(runtime, case)
    assert case.reads == []


def test_merge_can_bind_the_sources_a_search_found(runtime, monkeypatch):
    search = {'step_id': 'search', 'capability_id': 'document_search', 'arguments': {'query': 'regional sales'}}
    bound = merge_step()
    bound['inputs'] = {'sources': {'binding': binding('search', 'sources'), 'allow_partial': False}}
    case = tabular_case(runtime, [search, bound])
    module = ModuleType('functions_search')
    module.hybrid_search = lambda *args, **kwargs: [
        {'document_id': 'east-doc', 'id': 'hit-1', 'chunk_text': 'East sales'},
        {'document_id': 'west-doc', 'id': 'hit-2', 'chunk_text': 'West sales'},
    ]
    monkeypatch.setitem(sys.modules, 'functions_search', module)
    result = execute(runtime, case)
    assert result['status'] == 'completed', result
    assert case.reads == ['east-doc', 'west-doc']
    task = runtime.contracts.TaskResult.from_dict(result['task_results']['merge'])
    assert task.output('records').item_count == 3


def test_bound_sources_still_respect_the_chat_file_limit(runtime, monkeypatch):
    search = {'step_id': 'search', 'capability_id': 'document_search', 'arguments': {'query': 'regional sales'}}
    bound = merge_step()
    bound['inputs'] = {'sources': {'binding': binding('search', 'sources'), 'allow_partial': False}}
    settings = {'document_action_capabilities': {'merge': {'chat_max_documents': 2}}}
    case = tabular_case(runtime, [search, bound], settings=settings)
    module = ModuleType('functions_search')
    module.hybrid_search = lambda *args, **kwargs: [
        {'document_id': document_id, 'id': f'hit-{document_id}', 'chunk_text': 'sales'}
        for document_id in ('east-doc', 'west-doc', 'other-doc')
    ]
    monkeypatch.setitem(sys.modules, 'functions_search', module)
    result = execute(runtime, case)
    assert result['status'] == 'failed'
    assert result['failure']['code'] == 'merge_limit_exceeded'
    assert case.reads == []


@pytest.mark.parametrize('output_format, profile', [
    ('csv', 'exact_tabular_records_v1'), ('xlsx', 'exact_tabular_workbook_v1'),
])
def test_retained_merge_renders_every_column_in_order(runtime, output_format, profile):
    from functions_generated_file_exports import GeneratedFileExportRequest, build_generated_file_export
    from functions_orchestration_export_sources import open_orchestration_export_source

    case = tabular_case(runtime, [merge_step(['east-doc', 'west-doc'])])
    result = execute(runtime, case)
    task = runtime.contracts.TaskResult.from_dict(result['task_results']['merge'])
    source = open_orchestration_export_source(case.fixture.restart(), task.output('records'))
    with build_generated_file_export(
        source=source, export_request=GeneratedFileExportRequest(output_format, profile),
        max_output_bytes=8 * 1024 * 1024,
    ) as output:
        content = output.file_content.read()
        source.require_complete_consumption()
        assert output.record_count == 3
    if output_format == 'csv':
        rows = list(csv.reader(io.StringIO(content.decode('utf-8'))))
        assert rows == [
            ['Source File', 'Region', 'Code', 'Amount'],
            ['east.csv', 'East', '007', '10.50'],
            ['east.csv', 'East', '008', '1,200'],
            ['west.csv', 'West', '0099', '5'],
        ]
    else:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True)
        sheet = workbook['Sheet1']
        values = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
        assert values[0] == ['Source File', 'Region', 'Code', 'Amount']
        assert values[1] == ['east.csv', 'East', '007', '10.50']
        assert values[3] == ['west.csv', 'West', '0099', '5']


if __name__ == '__main__':
    sys.exit(pytest.main([__file__, '-q']))
