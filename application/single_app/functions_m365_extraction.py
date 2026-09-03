# functions_m365_extraction.py
"""Non-ingesting, bounded file extraction for live Microsoft 365 evidence."""

import csv
import io
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List
from xml.etree.ElementTree import ParseError

from lxml.etree import XMLSyntaxError

from functions_m365_transport import M365ProviderError
from functions_office_media import (
    OFFICE_DOCUMENT_PART_MAX_BYTES,
    OFFICE_ZIP_ALLOWED_COMPRESSION,
    OFFICE_ZIP_MAX_ENTRIES,
    _read_zip_entry_bounded,
)


M365_EXTRACTED_TEXT_MAX_CHARS = 8 * 1024 * 1024
M365_PACKAGE_MAX_EXPANDED_BYTES = 256 * 1024 * 1024
M365_MAX_TABULAR_ROWS = 250000
M365_MAX_TABULAR_COLUMNS = 4096
M365_MAX_DOCUMENT_UNITS = 10000
M365_EVIDENCE_CHUNK_CHARS = 8000
M365_FILE_MIME_TYPES = {
    ".txt": ("text/plain",),
    ".md": ("text/markdown", "text/plain"),
    ".log": ("text/plain",),
    ".json": ("application/json", "text/plain"),
    ".xml": ("application/xml", "text/xml", "text/plain"),
    ".yaml": ("application/yaml", "application/x-yaml", "text/yaml", "text/plain"),
    ".yml": ("application/yaml", "application/x-yaml", "text/yaml", "text/plain"),
    ".html": ("text/html",),
    ".htm": ("text/html",),
    ".pdf": ("application/pdf",),
    ".doc": ("application/msword",),
    ".docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",),
    ".docm": ("application/vnd.ms-word.document.macroenabled.12",),
    ".ppt": ("application/vnd.ms-powerpoint",),
    ".pptx": ("application/vnd.openxmlformats-officedocument.presentationml.presentation",),
    ".pptm": ("application/vnd.ms-powerpoint.presentation.macroenabled.12",),
    ".xls": ("application/vnd.ms-excel",),
    ".xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",),
    ".xlsm": ("application/vnd.ms-excel.sheet.macroenabled.12",),
    ".csv": ("text/csv", "application/csv", "text/plain", "application/vnd.ms-excel"),
    ".tsv": ("text/tab-separated-values", "text/plain"),
}
_OFFICE_ZIP_EXTENSIONS = frozenset({".docx", ".docm", ".pptx", ".pptm", ".xlsx", ".xlsm"})
_TEXT_EXTENSIONS = frozenset({".txt", ".md", ".log", ".json", ".xml", ".yaml", ".yml", ".html", ".htm"})


@dataclass
class M365ExtractedPart:
    text: str
    location: Dict[str, Any] = field(default_factory=dict)


@dataclass
class M365ExtractionResult:
    parts: List[M365ExtractedPart] = field(default_factory=list)
    coverage: Dict[str, Any] = field(default_factory=lambda: {
        "complete": True,
        "text_only": True,
        "units_read": 0,
        "units_total": None,
        "characters_captured": 0,
        "limitations": [],
        "missing_ranges": [],
    })

    def add(self, text: str, location: Dict[str, Any]) -> bool:
        remaining = M365_EXTRACTED_TEXT_MAX_CHARS - self.coverage["characters_captured"]
        if len(text) > remaining:
            if remaining:
                self.parts.append(M365ExtractedPart(text[:remaining], {**location, "char_start": 0, "char_end": remaining}))
                self.coverage["characters_captured"] += remaining
            self.coverage.update({"complete": False, "hard_limit": "extracted_text_characters"})
            self.coverage["missing_ranges"].append({**location, "char_start": remaining, "char_end": len(text)})
            return False
        self.parts.append(M365ExtractedPart(text, location))
        self.coverage["characters_captured"] += len(text)
        self.coverage["units_read"] += 1
        return True


def m365_file_format(name: str, mime_type: str = ""):
    suffix = Path(str(name or "")).suffix.lower()
    allowed = M365_FILE_MIME_TYPES.get(suffix)
    if allowed is None:
        raise M365ProviderError("unsupported_format", "This file format is not supported for Microsoft 365 text extraction.")
    mime = str(mime_type or "").split(";", 1)[0].strip().lower()
    if mime and mime not in {*allowed, "application/octet-stream", "binary/octet-stream"}:
        raise M365ProviderError("unsupported_content_type", "The file's type does not match its supported extension.")
    return suffix, allowed


def _validate_office_package(path: str) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > OFFICE_ZIP_MAX_ENTRIES or len({entry.filename for entry in entries}) != len(entries):
                raise M365ProviderError("package_limit", "The Office package exceeds safe package limits.")
            expanded_bytes = 0
            for entry in entries:
                if (
                    entry.flag_bits & 1
                    or entry.compress_type not in OFFICE_ZIP_ALLOWED_COMPRESSION
                    or entry.filename.startswith(("/", "\\"))
                    or ".." in entry.filename.replace("\\", "/").split("/")
                ):
                    raise M365ProviderError("unsupported_protected_file", "The Office file is encrypted or has an unsupported package structure.")
                if entry.is_dir():
                    continue
                content = _read_zip_entry_bounded(archive, entry.filename, OFFICE_DOCUMENT_PART_MAX_BYTES)
                if content is None:
                    raise M365ProviderError("package_limit", "An Office package entry is unreadable or exceeds its safe size limit.")
                expanded_bytes += len(content)
                if expanded_bytes > M365_PACKAGE_MAX_EXPANDED_BYTES:
                    raise M365ProviderError("package_limit", "The expanded Office package exceeds its safe size limit.")
    except (zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise M365ProviderError("unsupported_protected_file", "The Office file is encrypted or is not a readable Office package.") from exc


def _legacy_office_text(path: str, suffix: str):
    # Legacy extractors depend on application config; load only when a legacy file is requested.
    from functions_content import extract_legacy_ppt_pages, extract_word_text
    import olefile

    if olefile.isOleFile(path):
        with olefile.OleFileIO(path) as document:
            if document.exists("EncryptedPackage") or document.exists("EncryptionInfo"):
                raise M365ProviderError("unsupported_protected_file", "Protected Office content cannot be extracted.")
    # The established legacy parsers raise plain Exception for malformed OLE streams.
    try:
        if suffix == ".ppt":
            return extract_legacy_ppt_pages(path)
        return extract_word_text(path, suffix)
    except Exception as exc:
        raise M365ProviderError("extraction_failed", "The existing Office extractor could not read this file.") from exc


def _extract_text(path: str, result: M365ExtractionResult) -> None:
    # Shared text decoding is deliberately deferred with the config-dependent extractor owner.
    from functions_content import extract_text_file

    text = extract_text_file(path)
    result.coverage["unit_kind"] = "text"
    result.coverage["units_total"] = 1
    result.add(text, {"char_start": 0, "char_end": len(text)})


def _extract_word(path: str, suffix: str, result: M365ExtractionResult) -> None:
    # Reuse the established Word parser without any document or search-index ingestion.
    from functions_content import extract_word_text

    text = _legacy_office_text(path, suffix) if suffix == ".doc" else extract_word_text(path, suffix)
    result.coverage.update({"unit_kind": "text", "units_total": 1})
    result.coverage["limitations"].append("Word text is retained without a layout-derived page map or image OCR.")
    result.add(text, {"char_start": 0, "char_end": len(text)})


def _extract_pdf(path: str, result: M365ExtractionResult) -> None:
    # Native parsing is loaded only for this format; it never invokes a remote OCR service.
    import fitz

    with open(path, "rb") as source:
        if not source.read(1024).lstrip().startswith(b"%PDF-"):
            raise M365ProviderError("invalid_file", "The file is not a readable PDF.")
    try:
        with fitz.open(path) as document:
            if document.needs_pass:
                raise M365ProviderError("unsupported_protected_file", "Password-protected PDFs cannot be extracted.")
            if not document.permissions & fitz.PDF_PERM_COPY:
                raise M365ProviderError("unsupported_protected_file", "This PDF does not permit text copying.")
            result.coverage.update({"unit_kind": "page", "units_total": len(document)})
            result.coverage["limitations"].append("Text extraction does not OCR scanned pages or describe figures.")
            for index, page in enumerate(document):
                if index >= M365_MAX_DOCUMENT_UNITS:
                    result.coverage.update({"complete": False, "hard_limit": "document_units"})
                    result.coverage["missing_ranges"].append({"page_start": index + 1, "page_end": len(document)})
                    break
                text = page.get_text("text", sort=True)
                if not text.strip():
                    result.coverage["complete"] = False
                    result.coverage["missing_ranges"].append({"pages": [index + 1], "reason": "no_extractable_text"})
                if not result.add(text, {"pages": [index + 1]}):
                    if index + 1 < len(document):
                        result.coverage["missing_ranges"].append({"page_start": index + 2, "page_end": len(document)})
                    break
    except (fitz.FileDataError, fitz.EmptyFileError, RuntimeError) as exc:
        raise M365ProviderError("extraction_failed", "The PDF text could not be extracted.") from exc


def _slide_text(shapes) -> Iterable[str]:
    for shape in shapes:
        if shape.shape_type == 6:
            yield from _slide_text(shape.shapes)
        elif shape.has_text_frame:
            yield shape.text_frame.text
        elif shape.has_table:
            for row in shape.table.rows:
                yield "\t".join(cell.text for cell in row.cells)


def _extract_powerpoint(path: str, suffix: str, result: M365ExtractionResult) -> None:
    if suffix == ".ppt":
        slides = _legacy_office_text(path, suffix)
        result.coverage.update({"unit_kind": "slide", "units_total": len(slides)})
        for slide in slides:
            if not result.add(slide["content"], {"slides": [slide["page_number"]]}):
                break
        result.coverage["limitations"].append("Legacy PowerPoint extraction retains slide text, not embedded media.")
        return
    # Office parsers are format-specific; no macros, external links, or slide code are executed.
    from pptx import Presentation
    from pptx.exc import InvalidXmlError

    try:
        presentation = Presentation(path)
    except InvalidXmlError as exc:
        raise M365ProviderError("extraction_failed", "The PowerPoint package could not be parsed.") from exc
    result.coverage.update({"unit_kind": "slide", "units_total": len(presentation.slides)})
    result.coverage["limitations"].append("Slide and speaker-note text is retained; charts, pictures, and embedded media are not interpreted.")
    for index, slide in enumerate(presentation.slides):
        if index >= M365_MAX_DOCUMENT_UNITS:
            result.coverage.update({"complete": False, "hard_limit": "document_units"})
            result.coverage["missing_ranges"].append({"slide_start": index + 1, "slide_end": len(presentation.slides)})
            break
        paragraphs = list(_slide_text(slide.shapes))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            paragraphs.append(slide.notes_slide.notes_text_frame.text)
        if not result.add("\n".join(paragraphs), {"slides": [index + 1]}):
            if index + 1 < len(presentation.slides):
                result.coverage["missing_ranges"].append({"slide_start": index + 2, "slide_end": len(presentation.slides)})
            break


def _row_text(values) -> str:
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(values)
    return output.getvalue()


def _add_rows(rows, sheet_name: str, result: M365ExtractionResult) -> bool:
    count = 0
    for row_index, values in enumerate(rows, start=1):
        if result.coverage["units_read"] >= M365_MAX_TABULAR_ROWS:
            result.coverage.update({"complete": False, "hard_limit": "tabular_rows"})
            result.coverage["missing_ranges"].append({"sheet": sheet_name, "row_start": row_index, "reason": "remaining_rows_not_read"})
            return False
        values = tuple(values)
        if len(values) > M365_MAX_TABULAR_COLUMNS:
            result.coverage.update({"complete": False, "hard_limit": "tabular_columns"})
            result.coverage["missing_ranges"].append({"sheet": sheet_name, "row_start": row_index, "reason": "row_exceeds_column_limit"})
            return False
        if not result.add(_row_text(values), {"sheet": sheet_name, "row_start": row_index, "row_end": row_index}):
            return False
        count += 1
    result.coverage.setdefault("sheets", []).append({"name": sheet_name, "rows_captured": count, "complete": True})
    return True


def _extract_spreadsheet(path: str, suffix: str, result: M365ExtractionResult) -> None:
    result.coverage["unit_kind"] = "row"
    result.coverage["limitations"].append("Cells are retained as source rows; macros are inert and formulas are not recalculated.")
    if suffix in (".csv", ".tsv"):
        with open(path, encoding="utf-8-sig", newline="") as source:
            reader = csv.reader(source, delimiter="\t" if suffix == ".tsv" else ",", strict=True)
            _add_rows(reader, "CSV" if suffix == ".csv" else "TSV", result)
    elif suffix == ".xls":
        import xlrd

        try:
            workbook = xlrd.open_workbook(path, on_demand=True)
        except xlrd.XLRDError as exc:
            raise M365ProviderError("extraction_failed", "The legacy spreadsheet is protected or could not be parsed.") from exc
        try:
            result.coverage["sheet_names"] = workbook.sheet_names()
            for index in range(workbook.nsheets):
                sheet = workbook.sheet_by_index(index)
                if not _add_rows((sheet.row_values(row) for row in range(sheet.nrows)), sheet.name, result):
                    result.coverage["unread_sheets"] = workbook.sheet_names()[index + 1:]
                    break
        finally:
            workbook.release_resources()
    else:
        import openpyxl

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            result.coverage["sheet_names"] = list(workbook.sheetnames)
            for index, sheet in enumerate(workbook.worksheets):
                sheet.reset_dimensions()
                if not _add_rows(sheet.iter_rows(values_only=True), sheet.title, result):
                    result.coverage["unread_sheets"] = workbook.sheetnames[index + 1:]
                    break
        finally:
            workbook.close()
    if result.coverage["complete"]:
        result.coverage["units_total"] = result.coverage["units_read"]


def extract_m365_file(path: str, name: str, mime_type: str = "") -> M365ExtractionResult:
    suffix, _ = m365_file_format(name, mime_type)
    if suffix in _OFFICE_ZIP_EXTENSIONS:
        _validate_office_package(path)
    result = M365ExtractionResult()
    try:
        if suffix in _TEXT_EXTENSIONS:
            _extract_text(path, result)
        elif suffix in (".doc", ".docx", ".docm"):
            _extract_word(path, suffix, result)
        elif suffix == ".pdf":
            _extract_pdf(path, result)
        elif suffix in (".ppt", ".pptx", ".pptm"):
            _extract_powerpoint(path, suffix, result)
        else:
            _extract_spreadsheet(path, suffix, result)
    except (UnicodeError, csv.Error, OSError, ValueError, KeyError, ParseError, XMLSyntaxError, zipfile.BadZipFile) as exc:
        raise M365ProviderError("extraction_failed", "The file could not be decoded or parsed in its supported format.") from exc
    if not any(part.text.strip() for part in result.parts):
        raise M365ProviderError(
            "no_extractable_text", "The file contains no extractable text; OCR or a different supported source is required.",
            details={"coverage": result.coverage},
        )
    return result


def iter_m365_evidence_chunks(extraction: M365ExtractionResult):
    """Coalesce short rows without losing exact row/sheet boundaries; split long units losslessly."""
    from functions_conversation_memory import EvidenceChunk, EvidenceLocation

    text_parts = []
    location = None
    length = 0
    for part in extraction.parts:
        current = part.location
        same_sheet = (
            location is not None and current.get("sheet") is not None
            and location.get("sheet") == current.get("sheet")
            and location.get("row_end", -1) + 1 == current.get("row_start")
        )
        if text_parts and not (same_sheet and length + len(part.text) <= M365_EVIDENCE_CHUNK_CHARS):
            yield EvidenceChunk("".join(text_parts), EvidenceLocation(**_memory_location(location)))
            text_parts, location, length = [], None, 0
        if len(part.text) > M365_EVIDENCE_CHUNK_CHARS:
            for start in range(0, len(part.text), M365_EVIDENCE_CHUNK_CHARS):
                end = min(len(part.text), start + M365_EVIDENCE_CHUNK_CHARS)
                base_offset = current.get("char_start", 0)
                locator = {**current, "char_start": base_offset + start, "char_end": base_offset + end}
                yield EvidenceChunk(part.text[start:end], EvidenceLocation(**_memory_location(locator)))
            continue
        if location is None:
            location = dict(current)
        elif same_sheet:
            location["row_end"] = current["row_end"]
        text_parts.append(part.text)
        length += len(part.text)
    if text_parts:
        yield EvidenceChunk("".join(text_parts), EvidenceLocation(**_memory_location(location)))


def _memory_location(location: Dict[str, Any]) -> Dict[str, Any]:
    result = dict(location)
    for name in ("pages", "slides"):
        if name in result:
            result[name] = tuple(result[name])
    return result
