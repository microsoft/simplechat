# functions_document_merge_pptx.py
"""PowerPoint assembly for V2 file merge: slides appended in order into one deck.

Version: 0.261.223

The merge works on the Office Open XML package itself, so every slide is copied with
everything it relates to — pictures, media, charts and their embedded workbooks,
diagrams and speaker notes — exactly as it was. ``keep_source`` formatting brings each
deck's slide layouts, masters and themes along (identical ones are reused), so slides
keep their original look. ``use_first`` places each slide on the first deck's matching
layout, by name and then by type, so every slide takes the first deck's theme. Each
source deck can become its own section. The first deck decides the slide size. Entries
and section IDs are fixed, so the same decks always give the same bytes. Hyperlinks are
copied as they are; content linked from other files or the web stays linked and is reported.
"""

import hashlib
import io
import posixpath
import re
import uuid
import zipfile

from functions_document_merge import (
    DocumentMergeError,
    FORMATTING_KEEP_SOURCE,
    finish_output,
    fixed_zip_entry,
    guard_ooxml_package,
    new_output_spool,
    selected_indices,
)


_NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
}
_RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
RT_OFFICE_DOCUMENT = _RT + "officeDocument"
RT_SLIDE = _RT + "slide"
RT_SLIDE_LAYOUT = _RT + "slideLayout"
RT_SLIDE_MASTER = _RT + "slideMaster"
RT_NOTES_SLIDE = _RT + "notesSlide"
RT_NOTES_MASTER = _RT + "notesMaster"
RT_HYPERLINK = _RT + "hyperlink"
_COMMENT_RELATIONSHIP_SUFFIXES = ("/comments", "/commentAuthors")
_SECTION_EXTENSION_URI = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"
_SECTION_ID_NAMESPACE = uuid.uuid5(uuid.NAMESPACE_URL, "https://github.com/microsoft/simplechat/file-merge/sections")
_FIRST_MASTER_OR_LAYOUT_ID = 2147483648
_FIRST_SLIDE_ID = 256
_MAX_SLIDE_ID = 2147483647
_PRESENTATION_CHILD_ORDER = (
    "sldMasterIdLst", "notesMasterIdLst", "handoutMasterIdLst", "sldIdLst", "sldSz", "notesSz",
)
_NUMBERED_NAME = re.compile(r"^(?P<stem>.*?)(?P<number>\d*)(?P<extension>\.[^./]+)$")


def _parser():
    # lxml ships with python-pptx; entity resolution and network access stay off.
    from lxml import etree

    return etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False, remove_blank_text=False)


def _parse(data):
    from lxml import etree

    return etree.fromstring(data, _parser())


def _serialize(element):
    from lxml import etree

    return etree.tostring(element, xml_declaration=True, encoding="UTF-8", standalone=True)


def _q(prefix, name):
    return f"{{{_NS[prefix]}}}{name}"


class _Relationship:
    __slots__ = ("rid", "reltype", "target", "external")

    def __init__(self, rid, reltype, target, external):
        self.rid = rid
        self.reltype = reltype
        self.target = target
        self.external = external


class _Package:
    """An Office Open XML package as part names, bytes, relationships and content types."""

    def __init__(self, content, part, limits):
        guard_ooxml_package(content, part, limits, "PowerPoint presentation")
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                self.files = {info.filename: archive.read(info) for info in archive.infolist() if not info.is_dir()}
        except (zipfile.BadZipFile, KeyError, RuntimeError) as exc:
            raise DocumentMergeError("unreadable_document", f"{part.display_name()} isn't a valid PowerPoint presentation.") from exc
        self.part = part
        if "[Content_Types].xml" not in self.files:
            raise DocumentMergeError("unreadable_document", f"{part.display_name()} isn't a valid PowerPoint presentation.")
        types = _parse(self.files["[Content_Types].xml"])
        self.defaults = {
            element.get("Extension", "").lower(): element.get("ContentType")
            for element in types.iter(_q("ct", "Default"))
        }
        self.overrides = {
            element.get("PartName", "").lstrip("/"): element.get("ContentType")
            for element in types.iter(_q("ct", "Override"))
        }
        root_rels = self.relationships("")
        documents = [rel for rel in root_rels if rel.reltype == RT_OFFICE_DOCUMENT and not rel.external]
        if not documents:
            raise DocumentMergeError("unreadable_document", f"{part.display_name()} isn't a valid PowerPoint presentation.")
        self.presentation_name = documents[0].target
        content_type = self.content_type(self.presentation_name) or ""
        if "presentationml" not in content_type or "macroEnabled" in content_type:
            raise DocumentMergeError(
                "unreadable_document", f"{part.display_name()} isn't a PowerPoint (.pptx) presentation.",
            )

    @staticmethod
    def rels_name(part_name):
        directory, file_name = posixpath.split(part_name)
        return posixpath.join(directory, "_rels", f"{file_name}.rels")

    def relationships(self, part_name):
        data = self.files.get(self.rels_name(part_name) if part_name else "_rels/.rels")
        if data is None:
            return []
        directory = posixpath.dirname(part_name)
        rels = []
        for element in _parse(data).iter(_q("rel", "Relationship")):
            external = element.get("TargetMode") == "External"
            target = element.get("Target", "")
            if not external:
                target = target.lstrip("/") if target.startswith("/") else posixpath.normpath(
                    posixpath.join(directory, target),
                )
            rels.append(_Relationship(element.get("Id"), element.get("Type"), target, external))
        return rels

    def content_type(self, part_name):
        if part_name in self.overrides:
            return self.overrides[part_name]
        extension = posixpath.splitext(part_name)[1].lstrip(".").lower()
        return self.defaults.get(extension)

    def slide_names(self):
        presentation = _parse(self.files[self.presentation_name])
        by_id = {rel.rid: rel.target for rel in self.relationships(self.presentation_name)}
        names = []
        for element in presentation.iter(_q("p", "sldId")):
            target = by_id.get(element.get(_q("r", "id")))
            if target in self.files:
                names.append(target)
        return names

    def slide_size(self):
        size = _parse(self.files[self.presentation_name]).find(_q("p", "sldSz"))
        return (size.get("cx"), size.get("cy")) if size is not None else None


def _write_relationships(rels, owner=None):
    """Serialize relationships; with ``owner``, internal targets are absolute names made relative."""
    from lxml import etree

    root = etree.Element(_q("rel", "Relationships"), nsmap={None: _NS["rel"]})
    for rel in rels:
        target = rel.target if rel.external or owner is None else _relative(owner, rel.target)
        attributes = {"Id": rel.rid, "Type": rel.reltype, "Target": target}
        if rel.external:
            attributes["TargetMode"] = "External"
        etree.SubElement(root, _q("rel", "Relationship"), attributes)
    return _serialize(root)


def _relative(from_part, to_part):
    return posixpath.relpath(to_part, posixpath.dirname(from_part) or ".")


class _Destination:
    """The first deck, growing as other decks' slides and parts are copied in."""

    def __init__(self, package, context):
        self.package = package
        self.context = context
        self.files = dict(package.files)
        self.defaults = dict(package.defaults)
        self.overrides = dict(package.overrides)
        self.used = {name.lower() for name in self.files}
        self.presentation_name = package.presentation_name
        self.presentation = _parse(self.files[self.presentation_name])
        self.presentation_rels = package.relationships(self.presentation_name)
        self.layout_fingerprints = {}
        self.layouts = []
        self.notes_master = next(
            (rel.target for rel in self.presentation_rels if rel.reltype == RT_NOTES_MASTER and not rel.external), None,
        )
        self._index_layouts(package)

    def _index_layouts(self, package):
        for rel in self.presentation_rels:
            if rel.reltype != RT_SLIDE_MASTER or rel.external:
                continue
            master_name = rel.target
            master_print = _master_fingerprint(package, master_name)
            for layout_rel in package.relationships(master_name):
                if layout_rel.reltype != RT_SLIDE_LAYOUT or layout_rel.external:
                    continue
                layout_name = layout_rel.target
                self.layout_fingerprints.setdefault(
                    _layout_fingerprint(package, layout_name, master_print), layout_name,
                )
                self.layouts.append((layout_name, _layout_identity(package.files.get(layout_name, b""))))

    def unique_name(self, name):
        if name.lower() not in self.used:
            self.used.add(name.lower())
            return name
        directory, file_name = posixpath.split(name)
        match = _NUMBERED_NAME.match(file_name)
        stem, extension = (match.group("stem"), match.group("extension")) if match else (file_name, "")
        number = 1
        while True:
            candidate = posixpath.join(directory, f"{stem}{number}{extension}")
            if candidate.lower() not in self.used:
                self.used.add(candidate.lower())
                return candidate
            number += 1

    def next_presentation_rid(self):
        taken = {rel.rid for rel in self.presentation_rels}
        number = len(taken) + 1
        while f"rId{number}" in taken:
            number += 1
        return f"rId{number}"

    def add_presentation_rel(self, reltype, part_name):
        rid = self.next_presentation_rid()
        self.presentation_rels.append(_Relationship(rid, reltype, part_name, False))
        return rid

    def add_content_type(self, source, source_name, new_name):
        if source_name in source.overrides:
            self.overrides[new_name] = source.overrides[source_name]
            return
        extension = posixpath.splitext(new_name)[1].lstrip(".").lower()
        if extension not in self.defaults:
            content_type = source.content_type(source_name)
            if content_type:
                self.defaults[extension] = content_type

    def child(self, name):
        element = self.presentation.find(_q("p", name))
        if element is not None:
            return element
        from lxml import etree

        element = etree.Element(_q("p", name))
        position = _PRESENTATION_CHILD_ORDER.index(name)
        anchor = 0
        for index, existing in enumerate(self.presentation):
            local = etree.QName(existing).localname
            if local in _PRESENTATION_CHILD_ORDER and _PRESENTATION_CHILD_ORDER.index(local) < position:
                anchor = index + 1
        self.presentation.insert(anchor, element)
        return element

    def next_master_or_layout_id(self):
        values = [_FIRST_MASTER_OR_LAYOUT_ID - 1]
        values.extend(int(element.get("id")) for element in self.presentation.iter(_q("p", "sldMasterId")))
        for name, data in self.files.items():
            if name.endswith(".xml") and b"sldLayoutIdLst" in data:
                values.extend(int(element.get("id")) for element in _parse(data).iter(_q("p", "sldLayoutId")))
        return max(values) + 1

    def next_slide_id(self):
        ids = [int(element.get("id")) for element in self.presentation.iter(_q("p", "sldId"))]
        value = max(ids + [_FIRST_SLIDE_ID - 1]) + 1
        if value > _MAX_SLIDE_ID:
            raise DocumentMergeError("slide_limit_exceeded", "The merged deck has too many slides.")
        return value

    def write(self, spool):
        from lxml import etree

        self.files[self.presentation_name] = _serialize(self.presentation)
        self.files[_Package.rels_name(self.presentation_name)] = _write_relationships(
            self.presentation_rels, owner=self.presentation_name,
        )
        types = etree.Element(_q("ct", "Types"), nsmap={None: _NS["ct"]})
        for extension, content_type in sorted(self.defaults.items()):
            etree.SubElement(types, _q("ct", "Default"), {"Extension": extension, "ContentType": content_type})
        for part_name, content_type in sorted(self.overrides.items()):
            if part_name in self.files:
                etree.SubElement(types, _q("ct", "Override"), {"PartName": f"/{part_name}", "ContentType": content_type})
        with zipfile.ZipFile(spool, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(fixed_zip_entry("[Content_Types].xml"), _serialize(types))
            for name in sorted(self.files):
                if name != "[Content_Types].xml":
                    archive.writestr(fixed_zip_entry(name), self.files[name])


def _master_fingerprint(package, master_name):
    digest = hashlib.sha256(package.files.get(master_name, b""))
    for rel in package.relationships(master_name):
        if not rel.external and rel.reltype.endswith("/theme"):
            digest.update(package.files.get(rel.target, b""))
    return digest.hexdigest()


def _layout_fingerprint(package, layout_name, master_print):
    return hashlib.sha256(package.files.get(layout_name, b"") + master_print.encode("ascii")).hexdigest()


def _layout_identity(data):
    """A layout's name and type, used to match layouts across decks."""
    if not data:
        return ("", "")
    layout = _parse(data)
    common = layout.find(_q("p", "cSld"))
    return (
        ((common.get("name") if common is not None else "") or "").strip().casefold(),
        (layout.get("type") or "").strip(),
    )


def _master_of_layout(package, layout_name):
    return next(
        (rel.target for rel in package.relationships(layout_name) if rel.reltype == RT_SLIDE_MASTER and not rel.external),
        None,
    )


class _DeckCopier:
    """Copies one source deck's selected slides into the destination."""

    def __init__(self, destination, source, index, slide_names):
        self.destination = destination
        self.source = source
        self.index = index
        self.mapping = {}
        self.copied_masters = []
        self.keep_source = destination.context.options.formatting == FORMATTING_KEEP_SOURCE
        self.fallback_layouts = 0
        self.dropped_comments = 0
        self.dropped_links = 0
        self.linked_content = 0
        self.slide_names = set(slide_names)
        self.source_masters = {
            rel.target for rel in source.relationships(source.presentation_name)
            if rel.reltype == RT_SLIDE_MASTER and not rel.external
        }
        for name in slide_names:
            self.mapping[name] = destination.unique_name(name)

    def copy_slide(self, slide_name):
        return self._copy(slide_name, force=True)

    def _copy(self, source_name, force=False):
        if source_name in self.mapping and not force:
            return self.mapping[source_name]
        new_name = self.mapping.get(source_name) or self.destination.unique_name(source_name)
        self.mapping[source_name] = new_name
        data = self.source.files[source_name]
        new_rels = []
        dropped = set()
        for rel in self.source.relationships(source_name):
            if rel.external:
                if rel.reltype != RT_HYPERLINK:
                    self.linked_content += 1
                new_rels.append(rel)
                continue
            target = self._target_for(source_name, rel)
            if target is None:
                dropped.add(rel.rid)
                continue
            new_rels.append(_Relationship(rel.rid, rel.reltype, _relative(new_name, target), False))
        if dropped:
            data = _remove_references(data, dropped)
        if source_name in self.source_masters:
            self.copied_masters.append(new_name)
        self.destination.files[new_name] = data
        if new_rels:
            self.destination.files[_Package.rels_name(new_name)] = _write_relationships(new_rels)
        self.destination.add_content_type(self.source, source_name, new_name)
        return new_name

    def _target_for(self, source_name, rel):
        """The destination part for a relationship, or None when it is not carried over."""
        target = rel.target
        if target not in self.source.files:
            return None
        if rel.reltype.endswith(_COMMENT_RELATIONSHIP_SUFFIXES) or "/comments" in rel.reltype:
            self.dropped_comments += 1
            return None
        if rel.reltype == RT_SLIDE:
            if target in self.slide_names or target in self.mapping:
                return self.mapping.get(target) or self._copy(target)
            self.dropped_links += 1
            return None
        if rel.reltype == RT_SLIDE_LAYOUT and source_name in self.slide_names:
            return self._layout_for(target)
        if rel.reltype == RT_NOTES_MASTER:
            if self.destination.notes_master is not None:
                return self.destination.notes_master
            copied = self._copy(target)
            rid = self.destination.add_presentation_rel(RT_NOTES_MASTER, copied)
            notes = self.destination.child("notesMasterIdLst")
            from lxml import etree

            etree.SubElement(notes, _q("p", "notesMasterId"), {_q("r", "id"): rid})
            self.destination.notes_master = copied
            return copied
        return self._copy(target)

    def _layout_for(self, layout_name):
        master_name = _master_of_layout(self.source, layout_name)
        master_print = _master_fingerprint(self.source, master_name) if master_name else ""
        fingerprint = _layout_fingerprint(self.source, layout_name, master_print)
        existing = self.destination.layout_fingerprints.get(fingerprint)
        if existing is not None:
            return existing
        if self.keep_source:
            copied = self._copy(layout_name)
            # Every layout of a copied master is now in the destination; reuse them by fingerprint.
            for rel in self.source.relationships(master_name or ""):
                if rel.reltype == RT_SLIDE_LAYOUT and rel.target in self.mapping:
                    self.destination.layout_fingerprints.setdefault(
                        _layout_fingerprint(self.source, rel.target, master_print), self.mapping[rel.target],
                    )
            return copied
        name, kind = _layout_identity(self.source.files.get(layout_name, b""))
        for candidate, (candidate_name, candidate_kind) in self.destination.layouts:
            if name and candidate_name == name:
                return candidate
        for candidate, (_, candidate_kind) in self.destination.layouts:
            if kind and candidate_kind == kind:
                return candidate
        self.fallback_layouts += 1
        for candidate, (_, candidate_kind) in self.destination.layouts:
            if candidate_kind == "blank":
                return candidate
        return self.destination.layouts[0][0]


def _remove_references(data, rids):
    """Remove elements that point at relationships that were not carried over."""
    root = _parse(data)
    attributes = (_q("r", "id"), _q("r", "embed"), _q("r", "link"), _q("r", "pict"))
    for element in list(root.iter()):
        if any(element.get(attribute) in rids for attribute in attributes):
            parent = element.getparent()
            if parent is not None:
                parent.remove(element)
    return _serialize(root)


class _IdAllocator:
    """Hands out master and layout IDs that stay unique across the merged presentation."""

    def __init__(self, destination):
        self.next_value = destination.next_master_or_layout_id()

    def take(self):
        value = self.next_value
        self.next_value += 1
        return value


class _MergeState:
    """What the decks merged so far have built: the destination package, IDs, sections and size."""

    __slots__ = ("destination", "allocator", "sections", "first_size", "total_slides")

    def __init__(self):
        self.destination = None
        self.allocator = None
        self.sections = []
        self.first_size = None
        self.total_slides = 0


def assemble_pptx(context):
    """Append every part's selected slides, in order, to the first deck."""
    state = _MergeState()
    for index, part in enumerate(context.parts):
        content = context.load(index)
        try:
            _add_deck(context, state, index, part, content)
        except DocumentMergeError:
            raise
        except Exception as exc:
            if not context.blames_file(exc):
                raise
            # Library errors can quote slide text, so only the file's name is reported.
            raise DocumentMergeError(
                "unreadable_document", f"{part.display_name()} couldn't be combined with the other decks.",
            ) from exc
        context.progress(index)

    context.report["totals"]["slides"] = state.total_slides
    if context.options.formatting == FORMATTING_KEEP_SOURCE:
        context.limit("Slides keep their original layouts, masters and themes. The first deck decides the slide size.")
    else:
        context.limit("Slides use the first deck's theme and its matching layouts. The first deck decides the slide size.")
    spool = new_output_spool()
    try:
        _collect_garbage(state.destination)
        if context.options.sections:
            _write_sections(state.destination, state.sections)
        state.destination.write(spool)
        spool.seek(0)
        _validate(spool, state.total_slides)
    except DocumentMergeError:
        spool.close()
        raise
    except Exception as exc:
        spool.close()
        if not context.blames_file(exc):
            raise
        raise DocumentMergeError("merge_failed", "The merged PowerPoint deck couldn't be written.") from exc
    except BaseException:
        spool.close()
        raise
    return finish_output(context, spool)


def _add_deck(context, state, index, part, content):
    """Append one deck's selected slides; the first deck becomes the destination."""
    from lxml import etree

    package = _Package(content, part, context.limits)
    slide_names = package.slide_names()
    if not slide_names:
        raise DocumentMergeError("empty_source", f"{part.display_name()} has no slides.")
    chosen = [slide_names[position] for position in selected_indices(part, len(slide_names), "slide")]
    if state.total_slides + len(chosen) > context.limits.max_slides:
        raise DocumentMergeError(
            "slide_limit_exceeded", f"The merged deck would have more than {context.limits.max_slides:,} slides.",
        )
    context.check_cancel()
    entry = context.part_entry(index)
    entry["slides"] = len(chosen)
    entry["source_slides"] = len(slide_names)
    if index == 0:
        state.destination = _Destination(package, context)
        state.first_size = package.slide_size()
        state.allocator = _IdAllocator(state.destination)
        state.sections.append((part.display_name(), _keep_first_deck_slides(state.destination, chosen)))
        state.total_slides += len(chosen)
        return
    destination = state.destination
    if package.slide_size() != state.first_size:
        context.warn(
            "slide_size_differs",
            "This deck's slide size differs from the first deck's, so its slides may look stretched or cropped.",
            index,
        )
    if _parse(package.files[package.presentation_name]).find(_q("p", "embeddedFontLst")) is not None:
        context.warn("embedded_fonts", "Fonts embedded in this deck were not carried over.", index)
    copier = _DeckCopier(destination, package, index, chosen)
    slide_list = destination.child("sldIdLst")
    section_ids = []
    for slide_name in chosen:
        context.check_cancel()
        new_name = copier.copy_slide(slide_name)
        rid = destination.add_presentation_rel(RT_SLIDE, new_name)
        slide_id = str(destination.next_slide_id())
        etree.SubElement(slide_list, _q("p", "sldId"), {"id": slide_id, _q("r", "id"): rid})
        section_ids.append(slide_id)
    for master_name in copier.copied_masters:
        master = _parse(destination.files[master_name])
        for layout_id in master.iter(_q("p", "sldLayoutId")):
            layout_id.set("id", str(state.allocator.take()))
        destination.files[master_name] = _serialize(master)
        rid = destination.add_presentation_rel(RT_SLIDE_MASTER, master_name)
        etree.SubElement(destination.child("sldMasterIdLst"), _q("p", "sldMasterId"), {
            "id": str(state.allocator.take()), _q("r", "id"): rid,
        })
    if copier.fallback_layouts:
        context.warn(
            "layout_fallback",
            f"{copier.fallback_layouts} slide(s) had no matching layout in the first deck and use a blank layout.",
            index,
        )
    if copier.dropped_comments:
        context.warn("comments_removed", "Slide comments were not carried over.", index)
    if copier.dropped_links:
        context.warn("links_removed", "Links to slides that were not merged were removed.", index)
    if copier.linked_content:
        context.warn(
            "linked_content", "This deck links to content in other files or on the web; the links were kept.",
            index,
        )
    state.sections.append((part.display_name(), section_ids))
    state.total_slides += len(chosen)


def _keep_first_deck_slides(destination, chosen):
    """Keep the first deck's selected slides, in the selected order; return their slide IDs."""
    from lxml import etree

    if len(set(chosen)) != len(chosen):
        raise DocumentMergeError("range_invalid", "A slide can appear only once in the merged deck.")
    slide_list = destination.child("sldIdLst")
    targets = {rel.rid: rel.target for rel in destination.presentation_rels}
    by_name = {}
    for element in list(slide_list):
        by_name[targets.get(element.get(_q("r", "id")))] = element
        slide_list.remove(element)
    kept_ids = []
    for name in chosen:
        element = by_name.pop(name)
        slide_list.append(element)
        kept_ids.append(element.get("id"))
    removed = {element.get(_q("r", "id")) for element in by_name.values()}
    if removed:
        destination.presentation_rels = [rel for rel in destination.presentation_rels if rel.rid not in removed]
    return kept_ids


def _reachable_parts(destination):
    """Every part reachable from the package root, using the presentation's current relationships."""
    seen = set()
    stack = [rel.target for rel in destination.package.relationships("") if not rel.external]
    while stack:
        name = stack.pop()
        if name in seen:
            continue
        seen.add(name)
        if name == destination.presentation_name:
            stack.extend(rel.target for rel in destination.presentation_rels if not rel.external)
            continue
        rels_data = destination.files.get(_Package.rels_name(name))
        if rels_data is None:
            continue
        directory = posixpath.dirname(name)
        for element in _parse(rels_data).iter(_q("rel", "Relationship")):
            if element.get("TargetMode") == "External":
                continue
            target = element.get("Target", "")
            stack.append(
                target.lstrip("/") if target.startswith("/")
                else posixpath.normpath(posixpath.join(directory, target))
            )
    return seen


def _collect_garbage(destination):
    """Remove parts nothing refers to any more, such as the first deck's unselected slides."""
    reachable = _reachable_parts(destination)
    for name in list(destination.files):
        if name == "[Content_Types].xml" or name.endswith(".rels"):
            continue
        if name not in reachable:
            del destination.files[name]
    for name in list(destination.files):
        if not name.endswith(".rels") or name == "_rels/.rels":
            continue
        directory, file_name = posixpath.split(name)
        owner = posixpath.join(posixpath.dirname(directory), file_name[: -len(".rels")])
        if owner not in destination.files and owner != destination.presentation_name:
            del destination.files[name]


def _write_sections(destination, sections):
    from lxml import etree

    presentation = destination.presentation
    extension_list = presentation.find(_q("p", "extLst"))
    if extension_list is None:
        extension_list = etree.SubElement(presentation, _q("p", "extLst"))
    for extension in list(extension_list):
        if extension.get("uri") == _SECTION_EXTENSION_URI:
            extension_list.remove(extension)
    extension = etree.Element(_q("p", "ext"), {"uri": _SECTION_EXTENSION_URI})
    section_list = etree.SubElement(extension, _q("p14", "sectionLst"), nsmap={"p14": _NS["p14"]})
    for position, (name, slide_ids) in enumerate(sections):
        if not slide_ids:
            continue
        # Derived, not random, so the same decks always give the same bytes.
        section_id = uuid.uuid5(_SECTION_ID_NAMESPACE, f"{position}:{name}")
        section = etree.SubElement(section_list, _q("p14", "section"), {
            "name": _section_name(name), "id": "{" + str(section_id).upper() + "}",
        })
        ids = etree.SubElement(section, _q("p14", "sldIdLst"))
        for slide_id in slide_ids:
            etree.SubElement(ids, _q("p14", "sldId"), {"id": str(slide_id)})
    extension_list.insert(0, extension)


def _section_name(file_name):
    stem = posixpath.splitext(str(file_name or ""))[0].strip() or "Deck"
    return stem[:255]


def _validate(spool, expected_slides):
    # python-pptx re-opens the package, so a broken relationship fails here, not for the user.
    from pptx import Presentation

    presentation = Presentation(spool)
    if len(presentation.slides) != expected_slides:
        raise DocumentMergeError("merge_failed", "The merged deck did not contain every selected slide.")
    for slide in presentation.slides:
        _ = slide.slide_layout
    spool.seek(0)


__all__ = ["assemble_pptx"]
