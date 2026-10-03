# functions_document_merge_docx.py
"""Word assembly for V2 file merge: documents appended in order into one .docx.

Version: 0.261.224

Composition uses docxcompose, which carries over each document's styles, numbering,
images, footnotes, diagrams and shapes. With ``keep_source`` formatting, a style that
shares a name with the first document's but looks different is copied under a new name,
so every document keeps its own look; with ``use_first``, the first document's styles
win. The merged file uses the first document's page setup, headers and footers, and a
fixed modified date, and the IDs of copied lists are derived rather than random, so the
same files always give the same bytes. Comments are not carried over. A document with
fields that start other programs (DDE) is refused, and the first document's link to its
template is removed so Word never fetches it; other fields and links are copied as they
are, and linked content is reported.
"""

from datetime import datetime
import hashlib
import io
import re

from functions_document_merge import (
    DocumentMergeError,
    FORMATTING_KEEP_SOURCE,
    finish_output,
    guard_ooxml_package,
    new_output_spool,
    normalized_package,
)


_W_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_R_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_W = f"{{{_W_NAMESPACE}}}"
_COMMENT_TAGS = ("commentRangeStart", "commentRangeEnd", "commentReference")
_PACKAGE_DATE = datetime(2000, 1, 1)
# DDE and DDEAUTO fields ask Word to start another program and read data from it.
_PROGRAM_FIELD = re.compile(r"\s*DDE(?:AUTO)?(?!\w)", re.IGNORECASE)
_FIELD_TAGS = (f"{_W}fldSimple", f"{_W}fldChar", f"{_W}instrText", f"{_W}delInstrText")
_RT_SETTINGS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/settings"
_RT_HYPERLINK = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
_RT_NUMBERING = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering"
_LIST_ID_SEED = b"simplechat/file-merge/word-list-id:"


def assemble_docx(context):
    """Append every part's body, in order, to the first document."""
    # python-docx and docxcompose load only when a Word merge runs.
    from docx import Document
    from docxcompose.composer import Composer

    composer = None
    master = None
    first_lists = set()
    for index, part in enumerate(context.parts):
        content = context.load(index)
        guard_ooxml_package(content, part, context.limits, "Word document", check=context.check_cancel)
        try:
            document = Document(io.BytesIO(content))
            starts_programs = _has_program_fields(document)
        except Exception as exc:
            if not context.blames_file(exc):
                raise
            raise DocumentMergeError(
                "unreadable_document", f"{part.display_name()} couldn't be opened as a Word document.",
            ) from exc
        if starts_programs:
            raise DocumentMergeError(
                "active_content",
                f"{part.display_name()} has fields that start other programs (DDE), so it can't be merged.",
            )
        entry = context.part_entry(index)
        entry["paragraphs"] = len(document.paragraphs)
        entry["tables"] = len(document.tables)
        context.check_cancel()
        try:
            if index == 0:
                master = document
                first_lists = _list_definition_ids(master)
                if _remove_linked_template(master):
                    context.warn(
                        "linked_template_removed",
                        "This document's link to its template was removed, so Word won't fetch the template.",
                        index,
                    )
                composer = Composer(master, preserve_styles=context.options.formatting == FORMATTING_KEEP_SOURCE)
                if context.options.source_headings:
                    _insert_heading(master, part.display_name(), at_start=True)
            else:
                if _remove_comment_marks(document):
                    context.warn("comments_removed", "Comments in this document were not carried over.", index)
                if context.options.page_breaks:
                    master.add_page_break()
                if context.options.source_headings:
                    _insert_heading(master, part.display_name(), at_start=False)
                composer.append(document)
        except DocumentMergeError:
            raise
        except Exception as exc:
            if not context.blames_file(exc):
                raise
            # Library errors can quote document text, so only the file's name is reported.
            raise DocumentMergeError(
                "unreadable_document", f"{part.display_name()} couldn't be combined with the other documents.",
            ) from exc
        context.progress(index)

    try:
        _derive_list_ids(master, first_lists)
        linked = _linked_content_count(master)
    except Exception as exc:
        if not context.blames_file(exc):
            raise
        raise DocumentMergeError("merge_failed", "The merged Word document couldn't be prepared.") from exc
    if linked:
        context.warn(
            "linked_content",
            "The merged document links to content in other files or on the web; the links were kept.",
        )
    context.limit("The merged document uses the first document's page setup, headers and footers.")
    if context.options.formatting == FORMATTING_KEEP_SOURCE:
        context.limit(
            "Styles that share a name but look different are kept as separate copies, so each document keeps its look."
        )
    else:
        context.limit("Styles that share a name use the first document's definition.")
    return finish_output(context, _write_package(composer, master, context))


def _write_package(composer, master, context):
    """Save the composed document with a fixed modified date and ZIP metadata, then re-open it."""
    from docx import Document

    package = new_output_spool()
    spool = None
    try:
        # A document without core properties gets defaults stamped with the current time.
        master.core_properties.modified = _PACKAGE_DATE
        composer.save(package)
        spool = normalized_package(package, context)
        # Re-open the result so a malformed package fails here, not in the user's Word.
        Document(spool)
        return spool
    except Exception as exc:
        if spool is not None:
            spool.close()
        if isinstance(exc, DocumentMergeError) or not context.blames_file(exc):
            raise
        raise DocumentMergeError("merge_failed", "The merged Word document couldn't be written.") from exc
    except BaseException:
        if spool is not None:
            spool.close()
        raise
    finally:
        package.close()


def _insert_heading(document, text, *, at_start):
    """Add a heading naming a source document, using Heading 1 when the document has it."""
    try:
        paragraph = document.add_heading(text, level=1)
    except KeyError:
        paragraph = document.add_paragraph()
        paragraph.add_run(text).bold = True
    if at_start:
        element = paragraph._p
        body = element.getparent()
        body.remove(element)
        body.insert(0, element)
    return paragraph


def _remove_comment_marks(document):
    """Comments live in a part that is not carried over, so remove their anchors."""
    removed = False
    for tag in _COMMENT_TAGS:
        for element in document.element.body.iter(f"{{{_W_NAMESPACE}}}{tag}"):
            removed = True
    if not removed:
        return False
    for tag in _COMMENT_TAGS:
        for element in list(document.element.body.iter(f"{{{_W_NAMESPACE}}}{tag}")):
            parent = element.getparent()
            if tag == "commentReference" and parent is not None and parent.tag == f"{{{_W_NAMESPACE}}}r":
                run_parent = parent.getparent()
                if run_parent is not None:
                    run_parent.remove(parent)
                continue
            if parent is not None:
                parent.remove(element)
    return True


def _part_root(part):
    """A Word XML part's tree; a part python-docx keeps as bytes is parsed with entities off."""
    element = getattr(part, "_element", None)
    if element is not None:
        return element
    content_type = str(getattr(part, "content_type", "") or "")
    if "wordprocessingml" not in content_type or not content_type.endswith("+xml"):
        return None
    from lxml import etree

    parser = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False)
    return etree.fromstring(part.blob, parser)


def _has_program_fields(document):
    """Whether a field anywhere in the document, its headers, footers or notes starts a program."""
    for part in document.part.package.iter_parts():
        root = _part_root(part)
        if root is not None and _starts_programs(root):
            return True
    return False


def _starts_programs(root):
    # A complex field's code sits between its "begin" and "separate" (or "end") marks; it can be
    # split across runs and contain other fields, so the code of every open field is collected.
    codes = []
    for element in root.iter(*_FIELD_TAGS):
        if element.tag == f"{_W}fldSimple":
            if _PROGRAM_FIELD.match(element.get(f"{_W}instr") or ""):
                return True
        elif element.tag == f"{_W}fldChar":
            kind = element.get(f"{_W}fldCharType")
            if kind == "begin":
                codes.append([])
            elif kind in ("separate", "end") and codes:
                if codes[-1] is not None and _PROGRAM_FIELD.match("".join(codes[-1])):
                    return True
                codes[-1] = None
                if kind == "end":
                    codes.pop()
        elif codes and codes[-1] is not None:
            codes[-1].append(element.text or "")
    # A field still open at the end of a part counts too.
    return any(code is not None and _PROGRAM_FIELD.match("".join(code)) for code in codes)


def _remove_linked_template(document):
    """Remove the link to the document's template, so Word never fetches it; True when removed."""
    settings = next(
        (rel.target_part for rel in document.part.rels.values() if rel.reltype == _RT_SETTINGS and not rel.is_external),
        None,
    )
    root = getattr(settings, "_element", None)
    if root is None:
        return False
    templates = list(root.iter(f"{_W}attachedTemplate"))
    for template in templates:
        rid = template.get(f"{{{_R_NAMESPACE}}}id")
        template.getparent().remove(template)
        if rid and rid in settings.rels:
            settings.drop_rel(rid)
    return bool(templates)


def _linked_content_count(document):
    """How many relationships in the merged package point at other files or the web, links aside."""
    return sum(
        1 for part in document.part.package.iter_parts() for rel in part.rels.values()
        if rel.is_external and rel.reltype != _RT_HYPERLINK
    )


def _numbering_root(document):
    """The document's numbering XML, without creating a numbering part; None when there is none."""
    for rel in document.part.rels.values():
        if rel.reltype == _RT_NUMBERING and not rel.is_external:
            return getattr(rel.target_part, "_element", None)
    return None


def _list_definition_ids(document):
    root = _numbering_root(document)
    if root is None:
        return set()
    return {element.get(f"{_W}abstractNumId") for element in root.iter(f"{_W}abstractNum")}


def _derive_list_ids(document, first_lists):
    """Give each list definition docxcompose copied a derived ID instead of its random one.

    docxcompose gives every copied list a new random ``w:nsid``, so each later document's
    lists restart rather than continue the first's. A value derived from the definition's
    number, unique within the document, does the same and keeps the bytes reproducible.
    """
    root = _numbering_root(document)
    if root is None:
        return
    definitions = list(root.iter(f"{_W}abstractNum"))
    copied = [element for element in definitions if element.get(f"{_W}abstractNumId") not in first_lists]
    taken = {
        (nsid.get(f"{_W}val") or "").upper()
        for element in definitions if element.get(f"{_W}abstractNumId") in first_lists
        for nsid in element.iter(f"{_W}nsid")
    }
    for element in copied:
        nsid = next(element.iter(f"{_W}nsid"), None)
        if nsid is None:
            continue
        seed = hashlib.sha256(_LIST_ID_SEED + (element.get(f"{_W}abstractNumId") or "").encode("utf-8")).digest()
        value = int.from_bytes(seed[:4], "big")
        while f"{value:08X}" in taken:
            value = (value + 1) & 0xFFFFFFFF
        nsid.set(f"{_W}val", f"{value:08X}")
        taken.add(f"{value:08X}")


__all__ = ["assemble_docx"]
