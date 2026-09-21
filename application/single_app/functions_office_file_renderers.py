# functions_office_file_renderers.py
"""Headless binary renderers for complete, already prepared Office content.

This is a preparation-layer service, not an orchestration capability or a publisher.
It does not import the export registry, routes, settings, app bootstrap, storage, or
model clients. Callers own preparation, source authorization, and publication.

Public APIs are ``render_prepared_xlsx(PreparedWorkbook(...))``,
``render_prepared_docx(PreparedReport(...))``,
``render_prepared_pdf(PreparedReport(...))``, and
``render_prepared_pptx(PreparedDeck(...))``. Each accepts keyword-only ``limits``
and ``checks``; document/deck renderers also accept an ``image_resolver``.

Reports contain complete Markdown or plain text, not instructions for a writer.
Markdown supports headings, paragraphs, nested lists, rectangular tables, code,
quotes, rules, links/citations, and explicitly supplied PNG/JPEG images. Unsupported
HTML, attributes, and structures fail rather than being flattened or discarded.
Slides contain positioned text boxes, rectangular text tables, and supplied images;
there is deliberately no Markdown-to-deck composition, chart generation, or AI.

Images are resolved only from the input's bytes mapping, a validated inline PNG,
or an explicit caller-owned authorized resolver. No URL or file is ever fetched.
Checks must be supplied by adapters consuming retained or access-controlled data.
They run before consumption, periodically during work, and after serialization.
Publication must independently reauthorize, since a returned stream is not a grant.

Successful results own a seekable, bounded in-memory binary stream positioned at
zero. The caller MUST close it (prefer ``with render_prepared_...(...) as result``).
Failures close all owned streams and acquired row iterators. No temporary files
are used, including openpyxl's normally disk-backed write-only worksheet staging.
Counts describe the complete input, never a preview. Package timestamps and PDF
identifiers are deterministic; library upgrades can still change output encoding
or layout. SHA-256 always describes the actual returned bytes.
``OfficeRenderError.code`` and ``retryable`` are stable, user-safe classifications;
original exceptions are retained only as server-side exception causes.
"""

from __future__ import annotations

import hashlib
import io
import math
import re
from collections.abc import Callable, Iterable, Mapping, Sequence
from contextlib import contextmanager, suppress
from dataclasses import dataclass, field, fields
from datetime import date, datetime, time, timedelta
from html import escape
from html.parser import HTMLParser
from typing import Any, BinaryIO, Literal, TypeAlias
from unicodedata import normalize
from urllib.parse import urlsplit
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import fitz
import markdown2
from bs4 import BeautifulSoup, Comment, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, Twips
from docx.table import _Cell
from docx.text.run import Run
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.rich_text import CellRichText
from openpyxl.compat.strings import safe_string
from openpyxl.styles import Font
from openpyxl.worksheet._reader import _cast_number
from openpyxl.worksheet._writer import WorksheetWriter
from openpyxl.writer.excel import ExcelWriter
from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement as PptxElement
from pptx.util import Inches as PptxInches, Pt as PptxPt

from functions_export_cleanup import ClosingExportResource
from functions_export_visuals import (
    EXPORT_VISUAL_ASSET_MAX_BYTES,
    EXPORT_VISUAL_ASSET_MAX_PIXELS,
    EXPORT_VISUAL_CAPTION_CLASSES,
    EXPORT_VISUAL_WRAPPER_CLASSES,
    decode_export_visual_png,
)


_XML_INVALID = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_SHEET_INVALID = re.compile(r"[\[\]:*?/\\]")
_XLSX_FONT = Font(name="Arial", size=11)
_XLSX_HEADER_FONT = Font(name="Arial", size=11, bold=True)
_MEDIA_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
_ERROR_MESSAGES = {
    "invalid_input": "The prepared content does not match the renderer's input contract.",
    "incomplete_source": "The complete prepared content does not match its declared count.",
    "unsupported_content": "The prepared content requests an unsupported structure or value.",
    "limit_exceeded": "The complete output exceeds a configured renderer limit.",
    "layout_overflow": "The prepared content does not fit the requested page or slide layout.",
    "cancelled": "Rendering was cancelled.",
    "access_denied": "Access to the prepared content is no longer authorized.",
    "source_changed": "The prepared content is no longer the authorized source revision.",
    "source_unavailable": "The prepared source could not be rechecked or read.",
    "check_failed": "The prepared content could not be verified.",
    "render_io": "A renderer I/O operation could not be completed.",
    "render_failed": "The prepared content could not be rendered.",
}


class OfficeRenderError(ValueError):
    """Stable errors; only source-unavailable and renderer I/O failures are retryable."""

    def __init__(self, code: str):
        self.code = code
        self.retryable = code in {"source_unavailable", "render_io"}
        super().__init__(_ERROR_MESSAGES[code])


@dataclass(frozen=True)
class OfficeRenderLimits:
    """Server-configured ceilings, not permission to truncate or split content."""

    max_output_bytes: int = 32 * 1024 * 1024
    max_input_bytes: int = 64 * 1024 * 1024
    max_rows: int = 1_048_575
    max_columns: int = 16_384
    max_cells: int = 5_000_000
    max_sheets: int = 32
    max_report_nodes: int = 50_000
    max_pages: int = 500
    max_slides: int = 200
    max_shapes: int = 2_000
    max_images: int = 60
    max_image_bytes: int = EXPORT_VISUAL_ASSET_MAX_BYTES
    max_image_pixels: int = EXPORT_VISUAL_ASSET_MAX_PIXELS
    max_total_image_pixels: int = EXPORT_VISUAL_ASSET_MAX_PIXELS
    check_interval: int = 100


@dataclass(frozen=True)
class OfficeRenderChecks:
    """Caller-owned source/access checks raise on failure; cancellation returns bool.

    Rechecks may return ``None`` or ``True`` for success, ``False`` for failure.
    PermissionError becomes access_denied; OSError/TimeoutError become retryable
    source_unavailable. Raise OfficeRenderError for a more specific classification.
    No callback, identity, or client is copied into returned metadata.
    """

    source_check: Callable[[], None | bool] | None = None
    access_check: Callable[[], None | bool] | None = None
    is_cancelled: Callable[[], bool] | None = None


@dataclass
class RenderedOfficeFile:
    """Caller-owned, complete binary output; close explicitly or use a context manager."""

    file_content: BinaryIO
    output_format: str
    media_type: str
    size_bytes: int
    content_sha256: str
    metadata: dict[str, Any]
    profile: str = "prepared_office_v1"

    def close(self) -> None:
        self.file_content.close()

    def __enter__(self) -> RenderedOfficeFile:
        return self

    def __exit__(self, exc_type, exc_value, traceback) -> None:
        self.close()


CellValue: TypeAlias = str | bool | int | float | date | datetime | time | timedelta | None


@dataclass(frozen=True)
class PreparedSheet:
    """Ordered columns and complete, single-pass rows with a required data-row count.

    A row is a sequence of exactly len(columns) typed cells, or a mapping with
    exactly those keys. String cells are always text, including formula/error-like
    text. Formulas, nested values, nonfinite numbers, timezone-aware dates, numeric
    integers above Excel's 15-digit precision, and sub-millisecond times fail.
    Floats must retain their finite numeric value through openpyxl's actual numeric
    encoding/decoding. Negative zero is rejected because the reader loses its sign;
    positive zero and integral floats may reopen as equivalent integer numbers.
    Iterators acquired from rows are closed on both success and failure.
    """

    name: str
    columns: Sequence[str]
    rows: Iterable[Sequence[CellValue] | Mapping[str, CellValue]]
    row_count: int


@dataclass(frozen=True)
class PreparedWorkbook:
    sheets: Sequence[PreparedSheet]
    complete: bool = True


@dataclass(frozen=True)
class PreparedReport:
    """Complete Markdown/plain text shared by Word and PDF.

    images maps exact Markdown image references to authorized PNG/JPEG bytes.
    expected_block_count, when supplied, counts top-level rendered blocks (not
    the separately supplied title). Empty reports are valid, complete documents.
    """

    content: str
    content_format: Literal["markdown", "text"] = "markdown"
    title: str = ""
    images: Mapping[str, bytes] = field(default_factory=dict)
    complete: bool = True
    expected_block_count: int | None = None


@dataclass(frozen=True)
class SlideBox:
    """Explicit position/size in inches; boxes must fit and must not overlap."""

    left: float
    top: float
    width: float
    height: float


@dataclass(frozen=True)
class SlideParagraph:
    """Prepared literal text with optional paragraph-wide styling and hyperlink."""

    text: str
    list_kind: Literal["none", "bullet", "number"] = "none"
    level: int = 0
    bold: bool = False
    italic: bool = False
    url: str | None = None


@dataclass(frozen=True)
class SlideTextBox:
    box: SlideBox
    paragraphs: Sequence[SlideParagraph]
    font_size: float = 18


@dataclass(frozen=True)
class SlideTable:
    box: SlideBox
    rows: Sequence[Sequence[str]]
    header: bool = True
    font_size: float = 12


@dataclass(frozen=True)
class SlideImage:
    box: SlideBox
    source: str
    alt: str = ""


SlideShape: TypeAlias = SlideTextBox | SlideTable | SlideImage


@dataclass(frozen=True)
class PreparedSlide:
    """Title, positioned shapes, and literal notes; no generated/implicit content.

    title_and_content reserves (0.5, 0.25, slide_width - 1, 1) for the title.
    blank requires an empty title. All supplied shapes are rendered in order.
    """

    title: str
    shapes: Sequence[SlideShape] = ()
    notes: str = ""
    layout: Literal["title_and_content", "blank"] = "title_and_content"


@dataclass(frozen=True)
class PreparedDeck:
    slides: Sequence[PreparedSlide]
    expected_slide_count: int
    title: str = ""
    size: Literal["wide", "standard"] = "wide"
    images: Mapping[str, bytes] = field(default_factory=dict)
    complete: bool = True


def _require(condition: bool, code: str = "invalid_input") -> None:
    if not condition:
        raise OfficeRenderError(code)


def _sequence(value: Any) -> bool:
    return isinstance(value, Sequence) and not isinstance(value, (str, bytes, bytearray))


def _text(value: Any) -> str:
    _require(type(value) is str and not _XML_INVALID.search(value))
    return value


class _RenderContext:
    def __init__(self, limits: OfficeRenderLimits | None, checks: OfficeRenderChecks | None):
        self.limits = limits if limits is not None else OfficeRenderLimits()
        self.checks = checks if checks is not None else OfficeRenderChecks()
        _require(isinstance(self.limits, OfficeRenderLimits))
        _require(isinstance(self.checks, OfficeRenderChecks))
        for definition in fields(self.limits):
            value = getattr(self.limits, definition.name)
            _require(type(value) is int and value > 0)
        for definition in fields(self.checks):
            value = getattr(self.checks, definition.name)
            _require(value is None or callable(value))
        self.input_bytes = 0
        self.units = 0
        self.buffers: list[_BoundedBuffer] = []
        self.check()

    def check(self) -> None:
        for stream in self.buffers:
            stream.raise_failure()
        try:
            if self.checks.is_cancelled is not None:
                cancelled = self.checks.is_cancelled()
                _require(type(cancelled) is bool)
                _require(not cancelled, "cancelled")
            for callback, failure in (
                (self.checks.access_check, "access_denied"),
                (self.checks.source_check, "source_changed"),
            ):
                if callback is not None:
                    outcome = callback()
                    _require(outcome is None or type(outcome) is bool)
                    _require(outcome is not False, failure)
        except OfficeRenderError:
            raise
        except PermissionError as exc:
            raise OfficeRenderError("access_denied") from exc
        except OSError as exc:
            raise OfficeRenderError("source_unavailable") from exc
        except Exception as exc:
            raise OfficeRenderError("check_failed") from exc

    def tick(self) -> None:
        self.units += 1
        if self.units % self.limits.check_interval == 0:
            self.check()

    def consume(self, size: int) -> None:
        self.input_bytes += size
        _require(self.input_bytes <= self.limits.max_input_bytes, "limit_exceeded")

    def text(self, value: Any) -> str:
        _text(value)
        _require(len(value) <= self.limits.max_input_bytes - self.input_bytes, "limit_exceeded")
        self.consume(len(value.encode("utf-8")))
        return value

    def buffer(self) -> _BoundedBuffer:
        stream = _BoundedBuffer(self)
        self.buffers.append(stream)
        return stream

    def result(self, stream: _BoundedBuffer, output_format: str, metadata: dict[str, Any]) -> RenderedOfficeFile:
        stream.raise_failure()
        self.check()
        if output_format in {"xlsx", "docx", "pptx"}:
            stream = _normalize_office_package(stream, self)
        size = stream.seek(0, io.SEEK_END)
        _require(0 < size <= self.limits.max_output_bytes, "limit_exceeded")
        digest = hashlib.sha256()
        stream.seek(0)
        while chunk := stream.read(65536):
            digest.update(chunk)
            self.check()
        self.check()
        stream.seek(0)
        stream.context = None
        self.buffers.remove(stream)
        return RenderedOfficeFile(
            file_content=stream,
            output_format=output_format,
            media_type=_MEDIA_TYPES[output_format],
            size_bytes=size,
            content_sha256=digest.hexdigest(),
            metadata={"complete": True, "input_bytes": self.input_bytes, **metadata},
        )

    def close(self) -> None:
        for stream in self.buffers:
            stream.discarding = True
            stream.close()


class _BoundedBuffer(io.BytesIO):
    """Latch write failures until a safe library boundary, discarding later writes.

    ZIP package writers and MuPDF's C++ output callbacks cannot reliably unwind a
    Python write exception. A failed write never grows the buffer; virtual seeks
    let their close/finalization finish before the original failure is re-raised.
    """

    def __init__(self, context: _RenderContext):
        super().__init__()
        self.context = context
        self.failure: Exception | None = None
        self.discarding = False
        self.checked_bytes = 0

    def write(self, data) -> int:
        if self.discarding or self.failure is not None:
            super().seek(len(data), io.SEEK_CUR)
            return len(data)
        try:
            if self.context is not None:
                _require(self.tell() + len(data) <= self.context.limits.max_output_bytes, "limit_exceeded")
                self.checked_bytes += len(data)
                if self.checked_bytes >= 65536:
                    self.context.check()
                    self.checked_bytes = 0
            return super().write(data)
        except Exception as exc:
            self.failure = exc
            super().seek(len(data), io.SEEK_CUR)
            return len(data)

    def raise_failure(self) -> None:
        if self.failure is not None:
            raise self.failure


@contextmanager
def _rendering(limits, checks):
    context = _RenderContext(limits, checks)
    try:
        yield context
    except OfficeRenderError:
        raise
    except PermissionError as exc:
        raise OfficeRenderError("access_denied") from exc
    except OSError as exc:
        raise OfficeRenderError("render_io") from exc
    except Exception as exc:
        for stream in context.buffers:
            stream.raise_failure()
        raise OfficeRenderError("render_failed") from exc
    finally:
        context.close()


def _normalize_office_package(stream, context):
    """Canonical ZIP metadata without materializing uncompressed package parts."""
    stream.seek(0)
    normalized = context.buffer()
    with ZipFile(stream, "r") as source, ZipFile(normalized, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            context.check()
            info = ZipInfo(entry.filename, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.create_system = 0
            with source.open(entry) as reader, target.open(info, "w", force_zip64=True) as writer:
                while chunk := reader.read(65536):
                    writer.write(chunk)
                    context.check()
    normalized.raise_failure()
    return normalized


def _xlsx_cell(worksheet, value: CellValue, context: _RenderContext):
    if isinstance(value, str):
        context.text(value)
        _require(len(value.encode("utf-16-le")) // 2 <= 32767, "limit_exceeded")
    elif type(value) is bool or value is None:
        context.consume(1)
    elif type(value) is int:
        _require(abs(value) <= 999_999_999_999_999, "unsupported_content")
        context.consume(8)
    elif type(value) is float:
        _require(math.isfinite(value), "unsupported_content")
        restored = _cast_number(safe_string(value))
        _require(
            math.isfinite(restored) and restored == value
            and math.copysign(1, restored) == math.copysign(1, value),
            "unsupported_content",
        )
        context.consume(8)
    elif type(value) in {date, datetime, time, timedelta}:
        if isinstance(value, (datetime, time)):
            _require(value.tzinfo is None and value.microsecond % 1000 == 0, "unsupported_content")
        if isinstance(value, (date, datetime)):
            _require(value.year >= 1900, "unsupported_content")
        if isinstance(value, timedelta):
            _require(value.microseconds % 1000 == 0, "unsupported_content")
        context.consume(16)
    else:
        raise OfficeRenderError("unsupported_content")
    cell = WriteOnlyCell(worksheet, value=CellRichText([""]) if value == "" else value)
    cell.font = _XLSX_FONT
    if isinstance(value, str):
        cell.data_type = "s"
    context.tick()
    return cell


def _next_source_row(iterator):
    try:
        return next(iterator)
    except (StopIteration, OfficeRenderError):
        raise
    except PermissionError as exc:
        raise OfficeRenderError("access_denied") from exc
    except OSError as exc:
        raise OfficeRenderError("source_unavailable") from exc
    except Exception as exc:
        raise OfficeRenderError("incomplete_source") from exc


class _StreamingWorkbookWriter(ExcelWriter):
    """Use openpyxl 3.1's writer against ZIP entries instead of worksheet temp files.

    This seam and the numeric fidelity check use openpyxl 3.1 internals. Worksheet
    XML is generated a row at a time directly into the compressed output; no cell
    grid or XML spool is retained. Tests reopen workbooks and forbid temp files.
    """

    def __init__(self, workbook, archive, sheets, context):
        super().__init__(workbook, archive)
        self.sheets = sheets
        self.context = context

    def write_worksheet(self, worksheet):
        spec = self.sheets[worksheet.title]
        context = self.context
        context.check()
        with self._archive.open(worksheet.path[1:], "w", force_zip64=True) as entry:
            writer = WorksheetWriter(worksheet, out=entry)
            worksheet._writer = writer
            writer.write_top()
            try:
                header = [_xlsx_cell(worksheet, column, context) for column in spec.columns]
                for cell in header:
                    cell.font = _XLSX_HEADER_FONT
                worksheet.append(header)
                with ClosingExportResource(iter(spec.rows)) as iterator:
                    count = 0
                    while True:
                        try:
                            row = _next_source_row(iterator)
                        except StopIteration:
                            break
                        context.tick()
                        _require(count < spec.row_count, "incomplete_source")
                        if isinstance(row, Mapping):
                            _require(set(row) == set(spec.columns), "incomplete_source")
                            values = [row[column] for column in spec.columns]
                        else:
                            _require(_sequence(row) and len(row) == len(spec.columns), "incomplete_source")
                            values = row
                        worksheet.append([_xlsx_cell(worksheet, value, context) for value in values])
                        count += 1
                    _require(count == spec.row_count, "incomplete_source")
                    worksheet.close()
                    worksheet._rels = writer._rels
                    self.manifest.append(worksheet)
                    context.check()
            finally:
                if not worksheet.closed:
                    with suppress(Exception):
                        if worksheet._rows is not None:
                            worksheet._rows.close()
                    with suppress(Exception):
                        writer.close()
        context.check()


def render_prepared_xlsx(
    workbook: PreparedWorkbook, *, limits: OfficeRenderLimits | None = None,
    checks: OfficeRenderChecks | None = None,
) -> RenderedOfficeFile:
    """Render complete typed rows once, without formulas, flattening, or hidden queries.

    row_count excludes the header. Excel's hard row/column/cell-text limits apply
    even when configured limits are higher. Sheet names are validated, not renamed;
    case-insensitive duplicates fail. Zero data rows still produce the header.
    """
    with _rendering(limits, checks) as context:
        _require(isinstance(workbook, PreparedWorkbook))
        _require(workbook.complete is True, "incomplete_source")
        _require(_sequence(workbook.sheets) and len(workbook.sheets) > 0)
        _require(len(workbook.sheets) <= context.limits.max_sheets, "limit_exceeded")
        names = set()
        total_rows = 0
        total_cells = 0
        metadata_sheets = []
        book = Workbook(write_only=True, iso_dates=True)
        book.properties.creator = "SimpleChat"
        book.properties.created = book.properties.modified = datetime(2000, 1, 1)
        try:
            for spec in workbook.sheets:
                _require(isinstance(spec, PreparedSheet))
                name = context.text(spec.name)
                _require(
                    bool(name.strip()) and len(name.encode("utf-16-le")) // 2 <= 31 and not _SHEET_INVALID.search(name)
                    and not name.startswith("'") and not name.endswith("'")
                    and name.casefold() not in names
                )
                names.add(name.casefold())
                _require(_sequence(spec.columns) and len(spec.columns) > 0)
                _require(len(spec.columns) <= min(16384, context.limits.max_columns), "limit_exceeded")
                _require(all(type(column) is str and column for column in spec.columns))
                _require(len(set(spec.columns)) == len(spec.columns))
                _require(type(spec.row_count) is int and spec.row_count >= 0)
                _require(spec.row_count <= min(1_048_575, context.limits.max_rows), "limit_exceeded")
                _require(isinstance(spec.rows, Iterable) and not isinstance(spec.rows, (str, bytes, Mapping)))
                total_rows += spec.row_count
                total_cells += (spec.row_count + 1) * len(spec.columns)
                _require(total_cells <= context.limits.max_cells, "limit_exceeded")
                sheet = book.create_sheet(name)
                sheet.freeze_panes = "A2"
                metadata_sheets.append({
                    "name": name, "row_count": spec.row_count, "column_count": len(spec.columns),
                })
                context.check()
            stream = context.buffer()
            with ZipFile(stream, "w", ZIP_DEFLATED, allowZip64=True) as archive:
                writer = _StreamingWorkbookWriter(
                    book, archive, {spec.name: spec for spec in workbook.sheets}, context,
                )
                writer.write_data()
            return context.result(stream, "xlsx", {
                "record_count": total_rows, "sheet_count": len(metadata_sheets),
                "cell_count": total_cells, "sheets": metadata_sheets,
            })
        finally:
            book.close()


@dataclass(frozen=True)
class _ImageData:
    content: bytes
    width: int
    height: int
    extension: str


class _ImageAssets:
    def __init__(self, images, resolver, context):
        _require(isinstance(images, Mapping))
        _require(resolver is None or callable(resolver))
        _require(len(images) <= context.limits.max_images, "limit_exceeded")
        self.images = images
        self.resolver = resolver
        self.context = context
        self.resolved: dict[str, _ImageData] = {}
        self.total_pixels = 0
        for source, content in images.items():
            context.text(source)
            _require(type(content) is bytes and bool(content))
            _require(len(content) <= context.limits.max_image_bytes, "limit_exceeded")
            context.consume(len(content))
            context.tick()

    def resolve(self, source: str) -> _ImageData:
        context = self.context
        context.check()
        _text(source)
        _require(bool(source))
        if source in self.resolved:
            return self.resolved[source]
        _require(len(self.resolved) < context.limits.max_images, "limit_exceeded")
        if source in self.images:
            content = self.images[source]
        elif source.startswith("data:image/png;base64,"):
            _require(len(source) <= ((context.limits.max_image_bytes + 2) // 3) * 4 + 32, "limit_exceeded")
            content = decode_export_visual_png(source)
            _require(content is not None)
            context.consume(len(content))
        elif self.resolver is not None:
            try:
                content = self.resolver(source)
            except OfficeRenderError:
                raise
            except PermissionError as exc:
                raise OfficeRenderError("access_denied") from exc
            except OSError as exc:
                raise OfficeRenderError("source_unavailable") from exc
            except Exception as exc:
                raise OfficeRenderError("invalid_input") from exc
            _require(type(content) is bytes and bool(content))
            context.consume(len(content))
        else:
            raise OfficeRenderError("unsupported_content")
        context.check()
        _require(type(content) is bytes and bool(content))
        _require(len(content) <= context.limits.max_image_bytes, "limit_exceeded")
        try:
            with io.BytesIO(content) as stream, Image.open(stream) as image:
                _require(image.format in {"PNG", "JPEG"}, "unsupported_content")
                _require(getattr(image, "n_frames", 1) == 1, "unsupported_content")
                width, height = image.size
                _require(0 < width * height <= context.limits.max_image_pixels, "limit_exceeded")
                self.total_pixels += width * height
                _require(self.total_pixels <= context.limits.max_total_image_pixels, "limit_exceeded")
                extension = "png" if image.format == "PNG" else "jpeg"
                image.verify()
            with io.BytesIO(content) as stream, Image.open(stream) as image:
                image.load()
        except OfficeRenderError:
            raise
        except Image.DecompressionBombError as exc:
            raise OfficeRenderError("limit_exceeded") from exc
        except (UnidentifiedImageError, OSError, SyntaxError, ValueError) as exc:
            raise OfficeRenderError("invalid_input") from exc
        context.check()
        result = _ImageData(content, width, height, extension)
        self.resolved[source] = result
        return result


_INLINE_TAGS = {"a", "b", "strong", "em", "i", "s", "del", "strike", "code", "span", "br", "img"}
_BLOCK_TAGS = {"p", "h1", "h2", "h3", "h4", "h5", "h6", "ul", "ol", "pre", "blockquote", "table", "hr", "div"}
_REPORT_TAGS = _INLINE_TAGS | _BLOCK_TAGS | {"li", "thead", "tbody", "tr", "th", "td"}
_MARKDOWN_EXTRAS = ["fenced-code-blocks", "tables", "break-on-newline", "cuddled-lists", "strike"]


class _ReportMarkupGuard(HTMLParser):
    """Bound allocation and reject markup repairs that could hide requested content."""

    def __init__(self, context):
        super().__init__(convert_charrefs=True)
        self.context = context
        self.stack = []
        self.count = 0
        self.previous_data = False

    def node(self):
        self.count += 1
        _require(self.count <= self.context.limits.max_report_nodes, "limit_exceeded")
        self.context.tick()

    def handle_starttag(self, tag, attrs):
        self.previous_data = False
        self.node()
        _require(tag in _REPORT_TAGS, "unsupported_content")
        _require(len({name for name, _ in attrs}) == len(attrs), "unsupported_content")
        if tag not in {"br", "img", "hr"}:
            self.stack.append(tag)
            _require(len(self.stack) <= 32, "limit_exceeded")

    def handle_endtag(self, tag):
        self.previous_data = False
        _require(bool(self.stack) and self.stack[-1] == tag, "unsupported_content")
        self.stack.pop()

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag not in {"br", "img", "hr"}:
            self.handle_endtag(tag)

    def handle_data(self, data):
        if not self.previous_data and (self.stack or data.strip()):
            self.node()
        self.previous_data = True
        _text(data)

    def handle_comment(self, data):
        raise OfficeRenderError("unsupported_content")

    def handle_decl(self, decl):
        raise OfficeRenderError("unsupported_content")

    def unknown_decl(self, data):
        raise OfficeRenderError("unsupported_content")

    def handle_pi(self, data):
        raise OfficeRenderError("unsupported_content")


def _safe_link(value: Any, allow_fragment: bool = True) -> str:
    text = _text(value)
    _require(bool(text) and not re.search(r"[\x00-\x20\\]", text), "unsupported_content")
    try:
        parsed = urlsplit(text)
        hostname = parsed.hostname
    except ValueError as exc:
        raise OfficeRenderError("unsupported_content") from exc
    if text.startswith("#") and allow_fragment:
        _require(bool(re.fullmatch(r"#[A-Za-z][\w.-]*", text)), "unsupported_content")
    elif parsed.scheme in {"https", "http"}:
        _require(bool(hostname) and parsed.username is None and parsed.password is None, "unsupported_content")
    elif parsed.scheme == "mailto":
        _require(bool(parsed.path) and not parsed.netloc, "unsupported_content")
    else:
        raise OfficeRenderError("unsupported_content")
    return text


@dataclass
class _ReportData:
    root: Tag
    images: dict[int, _ImageData]
    metadata: dict[str, Any]


def _validate_report_node(node: Tag) -> None:
    name = node.name
    _require(name in _REPORT_TAGS, "unsupported_content")
    allowed = {"id"}
    if name == "a":
        allowed |= {"href", "title"}
    elif name == "img":
        allowed |= {"src", "alt", "title"}
    elif name == "ol":
        allowed.add("start")
    elif name in {"td", "th"}:
        allowed.add("style")
    elif name in {"span", "div", "p"}:
        allowed.add("class")
    _require(set(node.attrs) <= allowed, "unsupported_content")
    for value in node.attrs.values():
        for text in value if isinstance(value, list) else [value]:
            _text(text)
    if name == "div":
        classes = node.get("class") or []
        _require(len(classes) == 1, "unsupported_content")
        if classes == ["codehilite"]:
            children = {"pre"}
        else:
            _require(classes[0] in EXPORT_VISUAL_WRAPPER_CLASSES, "unsupported_content")
            children = {"p"}
        _require(all(child.name in children for child in node.children if isinstance(child, Tag)), "unsupported_content")
        _require(not any(isinstance(child, NavigableString) and child.strip() for child in node.children), "unsupported_content")
    if name == "p" and node.get("class"):
        _require(
            len(node["class"]) == 1 and node["class"][0] in EXPORT_VISUAL_CAPTION_CLASSES,
            "unsupported_content",
        )
    if name == "span" and node.get("class"):
        _require(node.find_parent("pre") is not None, "unsupported_content")
    if node.get("style"):
        _require(bool(re.fullmatch(r"text-align:\s*(left|right|center);?", node["style"])), "unsupported_content")
    if name == "a" and node.has_attr("href"):
        _safe_link(node["href"])
        _require(node.find_parent("a") is None, "unsupported_content")
    if name == "img":
        _require(type(node.get("src")) is str and bool(node["src"]))
    if name == "ol" and node.has_attr("start"):
        _require(bool(re.fullmatch(r"[1-9][0-9]{0,8}", str(node["start"]))), "unsupported_content")
    if node.has_attr("id"):
        _require(bool(re.fullmatch(r"[A-Za-z][\w.-]*", node["id"])), "unsupported_content")
        _require(name not in {"thead", "tbody", "tr", "div"}, "unsupported_content")
    if node.find_parent("pre") is not None:
        _require(name in {"span", "code"} and not node.has_attr("id"), "unsupported_content")
    children = [child.name for child in node.children if isinstance(child, Tag)]
    if name in _INLINE_TAGS | {"p", "h1", "h2", "h3", "h4", "h5", "h6", "td", "th", "pre"}:
        _require(set(children) <= _INLINE_TAGS, "unsupported_content")
    elif name in {"ul", "ol"}:
        _require(set(children) <= {"li"}, "unsupported_content")
    elif name == "li":
        _require(set(children) <= _INLINE_TAGS | {"p", "ul", "ol"}, "unsupported_content")
    elif name == "table":
        _require(set(children) <= {"thead", "tbody", "tr"}, "unsupported_content")
    elif name in {"thead", "tbody"}:
        _require(set(children) <= {"tr"}, "unsupported_content")
    elif name == "tr":
        _require(set(children) <= {"td", "th"}, "unsupported_content")
    if name in {"ul", "ol", "table", "thead", "tbody", "tr"}:
        _require(not any(isinstance(child, NavigableString) and child.strip() for child in node.children), "unsupported_content")
    if name in {"thead", "tbody", "tr", "td", "th", "li"}:
        parents = {
            "thead": {"table"}, "tbody": {"table"}, "tr": {"table", "thead", "tbody"},
            "td": {"tr"}, "th": {"tr"}, "li": {"ul", "ol"},
        }
        _require(node.parent.name in parents[name], "unsupported_content")


def _prepare_report(report, context, image_resolver) -> _ReportData:
    _require(isinstance(report, PreparedReport))
    _require(report.complete is True, "incomplete_source")
    content = context.text(report.content)
    context.text(report.title)
    _require(type(report.content_format) is str and report.content_format in {"markdown", "text"}, "unsupported_content")
    assets = _ImageAssets(report.images, image_resolver, context)
    context.check()
    if not content:
        rendered = ""
    elif report.content_format == "markdown":
        rendered = markdown2.markdown(content, extras=_MARKDOWN_EXTRAS)
    else:
        rendered = "".join(f"<p>{escape(line)}</p>" for line in content.split("\n")) if content else ""
    _require(len(rendered.encode("utf-8")) <= context.limits.max_input_bytes, "limit_exceeded")
    context.check()
    guard = _ReportMarkupGuard(context)
    for offset in range(0, len(rendered), 65536):
        guard.feed(rendered[offset:offset + 65536])
        context.check()
    guard.close()
    _require(not guard.stack, "unsupported_content")
    soup = BeautifulSoup(rendered, "html.parser")
    root = soup
    images = {}
    ids = set()
    anchors = set()
    node_count = 0
    table_cells = 0
    for child in list(root.children):
        if isinstance(child, NavigableString) and not isinstance(child, Comment):
            if child.strip():
                paragraph = soup.new_tag("p")
                child.replace_with(paragraph)
                paragraph.append(child)
            else:
                child.extract()
    stack = [(node, 1) for node in reversed(list(root.children))]
    while stack:
        node, depth = stack.pop()
        node_count += 1
        context.tick()
        _require(node_count <= context.limits.max_report_nodes and depth <= 32, "limit_exceeded")
        _require(not isinstance(node, Comment), "unsupported_content")
        if isinstance(node, NavigableString):
            _text(str(node))
            continue
        _require(isinstance(node, Tag), "unsupported_content")
        _validate_report_node(node)
        if node.has_attr("id"):
            _require(node["id"] not in ids, "unsupported_content")
            ids.add(node["id"])
        if node.name == "a" and node.get("href", "").startswith("#"):
            anchors.add(node["href"][1:])
        if node.name == "img":
            _require(len(images) < context.limits.max_images, "limit_exceeded")
            images[id(node)] = assets.resolve(node["src"])
        if node.name == "table":
            rows = node.find_all("tr")
            _require(bool(rows), "unsupported_content")
            widths = [len(row.find_all(["td", "th"], recursive=False)) for row in rows]
            _require(widths[0] > 0 and len(set(widths)) == 1, "unsupported_content")
            table_cells += sum(widths)
            _require(table_cells <= context.limits.max_cells, "limit_exceeded")
            _require(widths[0] <= context.limits.max_columns, "limit_exceeded")
        stack.extend((child, depth + 1) for child in reversed(list(node.children)))
    _require(anchors <= ids, "unsupported_content")
    block_count = len(list(root.children))
    if report.expected_block_count is not None:
        _require(type(report.expected_block_count) is int and report.expected_block_count >= 0)
        _require(block_count == report.expected_block_count, "incomplete_source")
    context.check()
    return _ReportData(root, images, {
        "block_count": block_count, "node_count": node_count,
        "content_char_count": len(content), "image_count": len(images),
        "table_count": len(root.find_all("table")), "link_count": len(root.find_all("a", href=True)),
        "table_cell_count": table_cells,
    })


class _WordReportWriter:
    """Request-independent adaptation of the existing Markdown/HTML-to-Word path."""

    def __init__(self, document, report, context):
        self.document = document
        self.report = report
        self.context = context
        self.bookmarks = {
            node["id"]: f"report_{index}"
            for index, node in enumerate(report.root.find_all(id=True), start=1)
        }
        self.bookmark_count = 0

    def bookmark(self, paragraph, node):
        name = node.get("id") if isinstance(node, Tag) else None
        if not name:
            return
        self.bookmark_count += 1
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(self.bookmark_count))
        start.set(qn("w:name"), self.bookmarks[name])
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(self.bookmark_count))
        paragraph._p.extend([start, end])

    def picture_size(self, paragraph, asset):
        section = self.document.sections[0]
        width = section.page_width - section.left_margin - section.right_margin
        height = section.page_height - section.top_margin - section.bottom_margin
        cell = paragraph._parent
        if isinstance(cell, _Cell):
            _require(cell.width is not None, "layout_overflow")
            width = min(width, cell.width)
            table = cell._parent
            margin_groups = [
                cell._tc.tcPr.find(qn("w:tcMar")),
                table._tbl.tblPr.find(qn("w:tblCellMar")),
            ]
            style = table.style
            while style is not None:
                properties = style.element.find(qn("w:tblPr"))
                if properties is not None:
                    margin_groups.append(properties.find(qn("w:tblCellMar")))
                style = style.base_style
            margins = {}
            for group in margin_groups:
                if group is None:
                    continue
                for side in ("left", "right", "top", "bottom"):
                    edge = group.find(qn(f"w:{side}"))
                    if edge is not None and side not in margins:
                        margins[side] = Twips(int(edge.get(qn("w:w"), "0")))
            width -= margins.get("left", 0) + margins.get("right", 0)
            height -= margins.get("top", 0) + margins.get("bottom", 0)
        width -= max(0, paragraph.paragraph_format.left_indent or 0)
        width -= max(0, paragraph.paragraph_format.right_indent or 0)
        _require(width > 0 and height > 0, "layout_overflow")
        fitted_width = int(min(width, Inches(asset.width / 96), height * asset.width / asset.height))
        fitted_height = int(fitted_width * asset.height / asset.width)
        _require(0 < fitted_width <= width and 0 < fitted_height <= height, "layout_overflow")
        return fitted_width, fitted_height

    def inline(self, paragraph, node, formatting=None, parent=None):
        formatting = dict(formatting or {})
        self.context.tick()
        if isinstance(node, NavigableString):
            if parent is None:
                run = paragraph.add_run(str(node))
            else:
                element = OxmlElement("w:r")
                parent.append(element)
                run = Run(element, paragraph)
                run.text = str(node)
            run.bold = formatting.get("bold", False)
            run.italic = formatting.get("italic", False)
            run.font.strike = formatting.get("strike", False)
            if formatting.get("link"):
                run.font.underline = True
            if formatting.get("code"):
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            return
        self.bookmark(paragraph, node)
        name = node.name
        if name == "br":
            if parent is None:
                run = paragraph.add_run()
            else:
                element = OxmlElement("w:r")
                parent.append(element)
                run = Run(element, paragraph)
            run.add_break()
            return
        if name == "img":
            asset = self.report.images[id(node)]
            width, height = self.picture_size(paragraph, asset)
            if parent is None:
                run = paragraph.add_run()
            else:
                element = OxmlElement("w:r")
                parent.append(element)
                run = Run(element, paragraph)
            with io.BytesIO(asset.content) as stream:
                picture = run.add_picture(stream, width=width, height=height)
            picture._inline.docPr.set("descr", node.get("alt", ""))
            if node.get("title"):
                picture._inline.docPr.set("title", node["title"])
            return
        if name in {"b", "strong"}:
            formatting["bold"] = True
        elif name in {"i", "em"}:
            formatting["italic"] = True
        elif name in {"s", "del", "strike"}:
            formatting["strike"] = True
        elif name == "code":
            formatting["code"] = True
        elif name == "a" and node.has_attr("href"):
            link = OxmlElement("w:hyperlink")
            href = node["href"]
            if href.startswith("#"):
                link.set(qn("w:anchor"), self.bookmarks[href[1:]])
            else:
                relationship = paragraph.part.relate_to(href, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
                link.set(qn("r:id"), relationship)
            if node.get("title"):
                link.set(qn("w:tooltip"), node["title"])
            paragraph._p.append(link)
            parent = link
            formatting["link"] = True
        for child in node.children:
            self.inline(paragraph, child, formatting, parent)

    def numbering(self, ordered, start, depth):
        numbering = self.document.part.numbering_part.element
        abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
        abstract_id = max(abstract_ids, default=-1) + 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        for tag, value in (("start", str(start)), ("numFmt", "decimal" if ordered else "bullet"),
                           ("lvlText", "%1." if ordered else "\u2022"), ("lvlJc", "left")):
            element = OxmlElement(f"w:{tag}")
            element.set(qn("w:val"), value)
            level.append(element)
        properties = OxmlElement("w:pPr")
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), str(360 * (depth + 1)))
        indent.set(qn("w:hanging"), "180")
        properties.append(indent)
        level.append(properties)
        abstract.append(level)
        numbering.append(abstract)
        return numbering.add_num(abstract_id).numId

    def list_items(self, node, depth=0):
        ordered = node.name == "ol"
        number_id = self.numbering(ordered, int(node.get("start", 1)), depth)
        for item in node.find_all("li", recursive=False):
            self.context.tick()
            paragraph = self.document.add_paragraph(style="List Number" if ordered else "List Bullet")
            numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
            numbering.get_or_add_ilvl().val = 0
            numbering.get_or_add_numId().val = number_id
            paragraph.paragraph_format.left_indent = Inches(0.25 * (depth + 1))
            self.bookmark(paragraph, item)
            for child in item.children:
                if isinstance(child, Tag) and child.name in {"ul", "ol"}:
                    self.list_items(child, depth + 1)
                    paragraph = None
                elif isinstance(child, Tag) and child.name == "p":
                    if paragraph is None or paragraph.text:
                        paragraph = self.document.add_paragraph()
                        paragraph.paragraph_format.left_indent = Inches(0.25 * (depth + 1))
                    self.inline(paragraph, child)
                elif isinstance(child, Tag) or str(child).strip() or self.inline_separator(child):
                    if paragraph is None:
                        paragraph = self.document.add_paragraph()
                        paragraph.paragraph_format.left_indent = Inches(0.25 * (depth + 1))
                    self.inline(paragraph, child)

    @staticmethod
    def inline_separator(node):
        if not isinstance(node, NavigableString) or node.strip():
            return False
        for neighbor in (node.previous_sibling, node.next_sibling):
            if isinstance(neighbor, Tag):
                if neighbor.name not in _INLINE_TAGS or neighbor.name == "br":
                    return False
            elif not isinstance(neighbor, NavigableString) or not neighbor.strip():
                return False
        return True

    def block(self, node, quote_depth=0):
        self.context.tick()
        name = node.name
        if name == "div":
            for child in node.children:
                if isinstance(child, Tag):
                    self.block(child, quote_depth)
            return
        if name in {"ul", "ol"}:
            if node.get("id"):
                self.bookmark(self.document.add_paragraph(), node)
            self.list_items(node)
            return
        if name == "blockquote":
            if node.get("id"):
                self.bookmark(self.document.add_paragraph(), node)
            for child in node.children:
                if isinstance(child, Tag) and child.name in _BLOCK_TAGS:
                    self.block(child, quote_depth + 1)
                elif isinstance(child, Tag) or str(child).strip():
                    paragraph = self.document.add_paragraph()
                    paragraph.paragraph_format.left_indent = Inches(0.3 * (quote_depth + 1))
                    self.inline(paragraph, child)
            return
        if name == "table":
            rows = node.find_all("tr")
            column_count = len(rows[0].find_all(["th", "td"], recursive=False))
            _require(column_count <= 63, "limit_exceeded")
            table = self.document.add_table(rows=len(rows), cols=column_count)
            table.style = "Table Grid"
            table.autofit = False
            for column in table.columns:
                column.width = Inches(6.5 / column_count)
            self.bookmark(table.cell(0, 0).paragraphs[0], node)
            for row_index, row in enumerate(rows):
                cells = table.rows[row_index].cells
                for column_index, cell_node in enumerate(row.find_all(["th", "td"], recursive=False)):
                    cell = cells[column_index]
                    cell.width = Inches(6.5 / column_count)
                    paragraph = cell.paragraphs[0]
                    self.inline(paragraph, cell_node, {"bold": cell_node.name == "th"})
                    alignment = cell_node.get("style", "")
                    if "right" in alignment:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    elif "center" in alignment:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            return
        if name.startswith("h") and len(name) == 2 and name[1].isdigit():
            paragraph = self.document.add_heading("", level=int(name[1]))
        else:
            paragraph = self.document.add_paragraph()
        if quote_depth:
            paragraph.paragraph_format.left_indent = Inches(0.3 * quote_depth)
        if name == "pre":
            self.bookmark(paragraph, node)
            run = paragraph.add_run(node.get_text())
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(6)
        elif name == "hr":
            self.bookmark(paragraph, node)
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "4")
            border.append(bottom)
            paragraph._p.get_or_add_pPr().append(border)
        else:
            self.inline(paragraph, node)


def render_prepared_docx(
    report: PreparedReport, *, limits: OfficeRenderLimits | None = None,
    checks: OfficeRenderChecks | None = None,
    image_resolver: Callable[[str], bytes] | None = None,
) -> RenderedOfficeFile:
    """Render prepared report content as Word, preserving supported rich semantics."""
    with _rendering(limits, checks) as context:
        data = _prepare_report(report, context, image_resolver)
        document = Document()
        section = document.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(1)
        document.styles["Normal"].font.name = "Arial"
        document.styles["Normal"].font.size = Pt(11)
        document.core_properties.title = report.title
        document.core_properties.author = "SimpleChat"
        document.core_properties.created = document.core_properties.modified = datetime(2000, 1, 1)
        if report.title:
            document.add_heading(report.title, level=0)
        writer = _WordReportWriter(document, data, context)
        for node in data.root.children:
            writer.block(node)
        context.check()
        stream = context.buffer()
        document.save(stream)
        return context.result(stream, "docx", data.metadata)


_REPORT_PDF_CSS = """
body { font-family: sans-serif; font-size: 11pt; line-height: 1.35; color: #222; }
h1 { font-size: 20pt; } h2 { font-size: 16pt; } h3 { font-size: 14pt; }
h4, h5, h6 { font-size: 12pt; }
p { margin-top: 3pt; margin-bottom: 6pt; }
p.office-plain-text { font-family: monospace; white-space: pre-wrap; margin: 0; }
table { border-collapse: collapse; width: 100%; margin: 8pt 0; font-size: 9pt; }
th, td { border: 0.6pt solid #777; padding: 4pt; }
th { background-color: #eee; font-weight: bold; }
pre, code { font-family: monospace; font-size: 9pt; }
pre { white-space: pre-wrap; background-color: #f3f3f3; padding: 6pt; }
blockquote { margin-left: 18pt; font-style: italic; }
a { color: #174ea6; text-decoration: underline; }
del, s, strike { text-decoration: line-through; }
"""


def render_prepared_pdf(
    report: PreparedReport, *, limits: OfficeRenderLimits | None = None,
    checks: OfficeRenderChecks | None = None,
    image_resolver: Callable[[str], bytes] | None = None,
) -> RenderedOfficeFile:
    """Paginate the shared report with PyMuPDF Story and bundled local fallback fonts.

    Only generated archive resource names reach Story. No CSS, webfont, directory,
    URL loader, or arbitrary HTML is accepted. Page/byte limits and out-of-page
    text/image geometry fail closed instead of returning a clipped preview.
    Extracted non-whitespace text must cover the prepared text in order (Unicode
    compatibility normalization accommodates typography such as ligatures).
    This also detects Story silently clipping long, unbreakable words or code.
    Plain text uses monospaced, whitespace-preserving paragraphs, including blank
    lines; Markdown keeps its normal HTML whitespace semantics.
    """
    with _rendering(limits, checks) as context:
        data = _prepare_report(report, context, image_resolver)
        archive = fitz.Archive()
        if report.content_format == "text":
            for paragraph in data.root.find_all("p", recursive=False):
                paragraph["class"] = ["office-plain-text"]
                if not paragraph.contents:
                    paragraph.append(data.root.new_tag("br"))
        for ordered_list in data.root.find_all("ol"):
            ordered_list["style"] = "list-style-type:none"
            start = int(ordered_list.get("start", 1))
            for index, item in enumerate(ordered_list.find_all("li", recursive=False), start=start):
                first = next((child for child in item.children if isinstance(child, Tag) or str(child).strip()), None)
                target = first if isinstance(first, Tag) and first.name == "p" else item
                target.insert(0, NavigableString(f"{index}. "))
        for index, node in enumerate(data.root.find_all("img")):
            image = data.images[id(node)]
            name = f"office-image-{index}.{image.extension}"
            archive.add((image.content, name))
            node["src"] = name
            width = min(480, image.width * 0.75, 620 * image.width / image.height)
            node["style"] = f"width:{width:.3f}pt;height:{width * image.height / image.width:.3f}pt"
        html = data.root.decode_contents()
        if report.title:
            html = f"<h1>{escape(report.title)}</h1>{html}"
        identity = hashlib.sha256((_REPORT_PDF_CSS + html).encode("utf-8"))
        for image in data.images.values():
            identity.update(image.content)
        story = fitz.Story(html=html, user_css=_REPORT_PDF_CSS, archive=archive)
        page_rect = fitz.paper_rect("letter")
        bounds = page_rect + (36, 36, -36, -36)
        initial = context.buffer()
        writer = fitz.DocumentWriter(initial)
        page_count = 0
        positions = []
        closed = False
        try:
            more = True
            while more:
                context.check()
                _require(page_count < context.limits.max_pages, "limit_exceeded")
                page_count += 1
                # Story's continuation rectangle can include next-block spacing;
                # painted text/image geometry is checked after all pages are drawn.
                more, _ = story.place(bounds)

                def collect_position(position):
                    position.page_num = page_count
                    positions.append(position)

                story.element_positions(collect_position)
                device = writer.begin_page(page_rect)
                story.draw(device)
                writer.end_page()
                initial.raise_failure()
                context.check()
            writer.close()
            closed = True
            initial.raise_failure()
        finally:
            if not closed:
                initial.discarding = True
                with suppress(Exception):
                    writer.close()
        context.check()
        with fitz.open(stream=initial.getvalue(), filetype="pdf") as document:
            fitz.Story.add_pdf_links(document, positions)
            document.set_metadata({"title": report.title, "author": "SimpleChat"})
            pdf_id = identity.hexdigest()[:32]
            document.xref_set_key(-1, "ID", f"[<{pdf_id}><{pdf_id}>]")
            expected = iter(character for character in normalize(
                "NFKC", report.title + data.root.get_text(),
            ) if not character.isspace())
            expected_character = next(expected, None)
            painted_images = 0
            for page in document:
                context.check()
                page_bounds = page.rect
                text = page.get_text(
                    flags=fitz.TEXTFLAGS_TEXT & ~fitz.TEXT_PRESERVE_LIGATURES,
                    clip=fitz.INFINITE_RECT(),
                )
                for index, character in enumerate(normalize("NFKC", text)):
                    if index % 4096 == 0:
                        context.check()
                    if character == expected_character:
                        expected_character = next(expected, None)
                words = page.get_text("words", clip=fitz.INFINITE_RECT())
                rectangles = [fitz.Rect(word[:4]) for word in words]
                page_images = page.get_image_info()
                painted_images += len(page_images)
                rectangles.extend(fitz.Rect(image["bbox"]) for image in page_images)
                for rectangle in rectangles:
                    _require(
                        rectangle.x0 >= page_bounds.x0 and rectangle.x1 <= page_bounds.x1
                        and rectangle.y0 >= page_bounds.y0 and rectangle.y1 <= page_bounds.y1,
                        "layout_overflow",
                    )
                    context.tick()
            _require(expected_character is None and painted_images == len(data.images), "layout_overflow")
            final = context.buffer()
            document.save(final, garbage=3, deflate=True, no_new_id=True)
            final.raise_failure()
        return context.result(final, "pdf", {
            **data.metadata, "page_count": page_count, "text_coverage_verified": True,
        })


def _slide_box(box, width, height):
    _require(type(box) is SlideBox, "unsupported_content")
    for value in (box.left, box.top, box.width, box.height):
        _require(type(value) in {int, float} and math.isfinite(value))
    _require(box.width > 0 and box.height > 0)
    _require(
        box.left >= 0 and box.top >= 0
        and box.left + box.width <= width + 0.001
        and box.top + box.height <= height + 0.001,
        "layout_overflow",
    )


def _boxes_overlap(left, right):
    return (
        min(left.left + left.width, right.left + right.width) > max(left.left, right.left) + 0.001
        and min(left.top + left.height, right.top + right.height) > max(left.top, right.top) + 0.001
    )


def _slide_text_height(text, width, font_size, level=0, list_kind="none"):
    """Conservative fit budget: no invisible overflow or viewer-dependent auto-shrink."""
    inset = 0 if list_kind == "none" else font_size * (level + 1)
    usable = width * 72 - 9 - inset
    _require(usable >= font_size, "layout_overflow")
    font = fitz.Font("hebo")
    lines = 0
    for line in _normalize_slide_text(text).expandtabs(4).split("\n"):
        used = 0
        lines += 1
        for token in re.findall(r"\S+|[^\S\n]+", line):
            measured = sum(
                font.glyph_advance(ord(character)) if font.has_glyph(ord(character)) else 1.1
                for character in token
            ) * font_size * 1.15
            _require(measured <= usable or not token.strip(), "layout_overflow")
            if used and used + measured > usable:
                lines += 1
                used = 0
            used += measured
    return lines * font_size * 1.4 + 6


def _normalize_slide_text(text):
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _write_slide_paragraphs(frame, paragraphs, font_size):
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = frame.margin_right = PptxInches(0.06)
    frame.margin_top = frame.margin_bottom = PptxInches(0.06)
    for index, spec in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.level = spec.level
        paragraph.space_after = PptxPt(6)
        paragraph.line_spacing = 1.15
        properties = paragraph._p.get_or_add_pPr()
        if spec.list_kind == "none":
            properties.append(PptxElement("a:buNone"))
        else:
            properties.set("marL", str(int(PptxPt(font_size * (spec.level + 1)))))
            properties.set("indent", str(-int(PptxPt(font_size * 0.7))))
            if spec.list_kind == "number":
                bullet = PptxElement("a:buAutoNum")
                bullet.set("type", "arabicPeriod")
            else:
                bullet = PptxElement("a:buChar")
                bullet.set("char", "\u2022")
            properties.append(bullet)
        paragraph.text = _normalize_slide_text(spec.text)
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = PptxPt(font_size)
            run.font.bold = spec.bold
            run.font.italic = spec.italic
            run.font.color.rgb = RGBColor(34, 34, 34)
            if spec.url:
                run.hyperlink.address = spec.url


def _render_slide_shape(slide, shape, context, assets):
    box = shape.box
    position = tuple(PptxInches(value) for value in (box.left, box.top, box.width, box.height))
    if type(shape) is SlideTextBox:
        _require(_sequence(shape.paragraphs) and bool(shape.paragraphs))
        _require(type(shape.font_size) in {float, int} and 8 <= shape.font_size <= 40)
        needed_height = 9
        for paragraph in shape.paragraphs:
            _require(type(paragraph) is SlideParagraph, "unsupported_content")
            context.text(paragraph.text)
            _require(
                type(paragraph.list_kind) is str and paragraph.list_kind in {"none", "bullet", "number"},
                "unsupported_content",
            )
            _require(type(paragraph.level) is int and 0 <= paragraph.level <= 8)
            _require(paragraph.list_kind != "none" or paragraph.level == 0, "unsupported_content")
            _require(type(paragraph.bold) is bool and type(paragraph.italic) is bool)
            if paragraph.url is not None:
                context.text(paragraph.url)
                _safe_link(paragraph.url, allow_fragment=False)
            needed_height += _slide_text_height(
                paragraph.text, box.width, shape.font_size, paragraph.level, paragraph.list_kind,
            )
            context.tick()
        _require(needed_height <= box.height * 72, "layout_overflow")
        frame = slide.shapes.add_textbox(*position).text_frame
        _write_slide_paragraphs(frame, shape.paragraphs, shape.font_size)
    elif type(shape) is SlideTable:
        _require(type(shape.header) is bool)
        _require(type(shape.font_size) in {float, int} and 8 <= shape.font_size <= 32)
        _require(_sequence(shape.rows) and bool(shape.rows))
        _require(_sequence(shape.rows[0]) and bool(shape.rows[0]))
        columns = len(shape.rows[0])
        _require(len(shape.rows) * columns <= context.limits.max_cells, "limit_exceeded")
        for row in shape.rows:
            _require(_sequence(row) and len(row) == columns, "unsupported_content")
            for cell in row:
                context.text(cell)
                needed_height = _slide_text_height(cell, box.width / columns, shape.font_size) + 9
                _require(needed_height <= box.height * 72 / len(shape.rows), "layout_overflow")
                context.tick()
        table = slide.shapes.add_table(len(shape.rows), columns, *position).table
        table.first_row = shape.header
        for row_index, row in enumerate(shape.rows):
            for column_index, text in enumerate(row):
                cell = table.cell(row_index, column_index)
                cell.margin_left = cell.margin_right = PptxInches(0.06)
                cell.margin_top = cell.margin_bottom = PptxInches(0.06)
                _write_slide_paragraphs(
                    cell.text_frame, [SlideParagraph(text, bold=shape.header and row_index == 0)], shape.font_size,
                )
    else:
        context.text(shape.source)
        context.text(shape.alt)
        asset = assets.resolve(shape.source)
        width = min(box.width, box.height * asset.width / asset.height)
        height = width * asset.height / asset.width
        left = box.left + (box.width - width) / 2
        top = box.top + (box.height - height) / 2
        with io.BytesIO(asset.content) as stream:
            picture = slide.shapes.add_picture(
                stream, PptxInches(left), PptxInches(top), PptxInches(width), PptxInches(height),
            )
        picture._element.nvPicPr.cNvPr.set("descr", _normalize_slide_text(shape.alt))
    context.check()


def render_prepared_pptx(
    deck: PreparedDeck, *, limits: OfficeRenderLimits | None = None,
    checks: OfficeRenderChecks | None = None,
    image_resolver: Callable[[str], bytes] | None = None,
) -> RenderedOfficeFile:
    """Render every prepared slide/shape, retaining titles, literal body, tables, notes.

    Supported sizes are wide (13 1/3 x 7.5 inches) and standard (10 x 7.5).
    Positioned shapes cannot overlap or cross the reserved title region.
    Text uses Arial with conservative fit checks, not truncation, slide invention,
    automatic reflow, or font shrinking. Exact typography still depends on the
    viewer's local fonts. Unsupported layouts/shapes must be changed in preparation.
    CRLF and CR are normalized to LF before measuring or writing all text surfaces.
    Paragraph/cell/title LF becomes a PowerPoint soft break; notes LF separates
    paragraphs. Normalization does not relax the existing layout limits.
    """
    with _rendering(limits, checks) as context:
        _require(isinstance(deck, PreparedDeck))
        _require(deck.complete is True, "incomplete_source")
        _require(_sequence(deck.slides) and bool(deck.slides))
        _require(type(deck.expected_slide_count) is int and deck.expected_slide_count >= 0)
        _require(len(deck.slides) == deck.expected_slide_count, "incomplete_source")
        _require(len(deck.slides) <= context.limits.max_slides, "limit_exceeded")
        _require(type(deck.size) is str and deck.size in {"wide", "standard"}, "unsupported_content")
        context.text(deck.title)
        assets = _ImageAssets(deck.images, image_resolver, context)
        width, height = (40 / 3, 7.5) if deck.size == "wide" else (10, 7.5)
        presentation = Presentation()
        presentation.slide_width = PptxInches(width)
        presentation.slide_height = PptxInches(height)
        presentation.core_properties.title = _normalize_slide_text(deck.title)
        presentation.core_properties.author = "SimpleChat"
        presentation.core_properties.created = presentation.core_properties.modified = datetime(2000, 1, 1)
        shape_count = 0
        table_count = 0
        table_cell_count = 0
        image_count = 0
        paragraph_count = 0
        notes_count = 0
        for spec in deck.slides:
            context.check()
            _require(type(spec) is PreparedSlide, "unsupported_content")
            _require(
                type(spec.layout) is str and spec.layout in {"title_and_content", "blank"},
                "unsupported_content",
            )
            context.text(spec.title)
            context.text(spec.notes)
            _require(_sequence(spec.shapes))
            shape_count += len(spec.shapes)
            _require(shape_count <= context.limits.max_shapes, "limit_exceeded")
            occupied = []
            if spec.layout == "title_and_content":
                _require(bool(spec.title))
                title_box = SlideBox(0.5, 0.25, width - 1, 1)
                title_height = _slide_text_height(spec.title, title_box.width, 30) + 9
                _require(title_height <= 72, "layout_overflow")
                occupied.append(title_box)
                slide = presentation.slides.add_slide(presentation.slide_layouts[5])
                title = slide.shapes.title
                title.left, title.top = PptxInches(0.5), PptxInches(0.25)
                title.width, title.height = PptxInches(width - 1), PptxInches(1)
                _write_slide_paragraphs(title.text_frame, [SlideParagraph(spec.title, bold=True)], 30)
            else:
                _require(spec.title == "", "unsupported_content")
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            for shape in spec.shapes:
                _require(type(shape) in {SlideTextBox, SlideTable, SlideImage}, "unsupported_content")
                if type(shape) is SlideImage:
                    _require(image_count < context.limits.max_images, "limit_exceeded")
                if type(shape) is SlideTable:
                    _require(_sequence(shape.rows) and bool(shape.rows))
                    _require(_sequence(shape.rows[0]) and bool(shape.rows[0]))
                    columns = len(shape.rows[0])
                    table_cell_count += len(shape.rows) * columns
                    _require(table_cell_count <= context.limits.max_cells, "limit_exceeded")
                    _require(columns <= context.limits.max_columns, "limit_exceeded")
                _slide_box(shape.box, width, height)
                _require(not any(_boxes_overlap(shape.box, previous) for previous in occupied), "layout_overflow")
                occupied.append(shape.box)
                _render_slide_shape(slide, shape, context, assets)
                if type(shape) is SlideTextBox:
                    paragraph_count += len(shape.paragraphs)
                elif type(shape) is SlideTable:
                    table_count += 1
                else:
                    image_count += 1
            if spec.notes:
                slide.notes_slide.notes_text_frame.text = _normalize_slide_text(spec.notes)
                notes_count += 1
            context.check()
        stream = context.buffer()
        presentation.save(stream)
        return context.result(stream, "pptx", {
            "slide_count": len(deck.slides), "shape_count": shape_count,
            "table_count": table_count, "image_count": image_count,
            "table_cell_count": table_cell_count,
            "text_paragraph_count": paragraph_count, "notes_count": notes_count,
        })
