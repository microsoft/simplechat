# exports.py
"""Build clean derivatives exclusively from reviewed canonical units."""

import csv
import io
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path

from content_screening.contracts import ScreeningValidationError, hash_payload, normalize_units
from functions_assistant_table_exports import neutralize_csv_spreadsheet_formula


@dataclass(frozen=True)
class CleanArtifact:
    file_name: str
    content_type: str
    content: bytes


def _cell_value(unit):
    text = unit.text
    kind = unit.locator.get("value_type", "text")
    if not text:
        return None
    if kind == "boolean":
        if text not in {"true", "false"}:
            raise ScreeningValidationError("A reviewed boolean cell is invalid.")
        return text == "true"
    if kind == "number":
        try:
            value = Decimal(text)
        except InvalidOperation as error:
            raise ScreeningValidationError("A reviewed numeric cell is invalid.") from error
        if not value.is_finite():
            raise ScreeningValidationError("A reviewed numeric cell is invalid.")
        return int(value) if value == value.to_integral_value() else float(value)
    if kind == "date":
        for parser in (datetime.fromisoformat, date.fromisoformat, time.fromisoformat):
            try:
                return parser(text)
            except ValueError:
                continue
        raise ScreeningValidationError("A reviewed date or time cell is invalid.")
    return text


def structured_sheets(units):
    sheets = {}
    for unit in normalize_units(units):
        kind = unit.locator.get("kind")
        if kind not in {"table_sheet", "table_cell", "table_formula"}:
            continue
        index = unit.locator.get("sheet_index")
        if type(index) is not int or index <= 0:
            raise ScreeningValidationError("The source sheet locator is invalid.")
        sheet = sheets.setdefault(index, {"name": f"Sheet{index}", "cells": {}})
        if kind == "table_sheet":
            sheet["name"] = unit.text.strip() or f"Sheet{index}"
        elif kind == "table_cell":
            row, column = unit.locator.get("row"), unit.locator.get("column")
            if (
                type(row) is not int or type(column) is not int
                or not 1 <= row <= 1048576 or not 1 <= column <= 16384
                or (row, column) in sheet["cells"]
            ):
                raise ScreeningValidationError("The source cell locator is invalid.")
            sheet["cells"][(row, column)] = _cell_value(unit)
    return sheets


def _csv_artifact(sheets, base_name):
    sheet = next(iter(sheets.values()))
    cells = sheet["cells"]
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    max_row = max((row for row, _column in cells), default=0)
    max_column = max((column for _row, column in cells), default=0)
    for row in range(1, max_row + 1):
        writer.writerow([
            neutralize_csv_spreadsheet_formula(cells.get((row, column)))
            for column in range(1, max_column + 1)
        ])
    return CleanArtifact(f"{base_name}.csv", "text/csv; charset=utf-8", output.getvalue().encode("utf-8"))


def _workbook_artifact(sheets, base_name):
    # Load the existing workbook dependency only when a workbook is requested.
    from openpyxl import Workbook
    from openpyxl.writer.excel import ExcelWriter

    workbook = Workbook()
    workbook.properties.created = datetime(2000, 1, 1)
    workbook.properties.modified = datetime(2000, 1, 1)
    workbook.remove(workbook.active)
    names = set()
    for index, sheet in sorted(sheets.items()):
        name = re.sub(r"[:\\/*?\[\]]", "_", sheet["name"]).strip("'")[:31] or f"Sheet{index}"
        root_name = name
        suffix = 2
        while name.casefold() in names:
            tail = f" {suffix}"
            name = f"{root_name[:31 - len(tail)]}{tail}"
            suffix += 1
        names.add(name.casefold())
        worksheet = workbook.create_sheet(name)
        for (row, column), value in sorted(sheet["cells"].items()):
            cell = worksheet.cell(row=row, column=column, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        ExcelWriter(workbook, archive).save()
    workbook.close()
    # A reviewed revision must produce identical bytes when publication is retried.
    deterministic = io.BytesIO()
    with zipfile.ZipFile(output, "r") as source:
        with zipfile.ZipFile(deterministic, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for name in sorted(source.namelist()):
                entry = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
                entry.compress_type = zipfile.ZIP_DEFLATED
                target.writestr(entry, source.read(name))
    return CleanArtifact(
        f"{base_name}.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        deterministic.getvalue(),
    )


def build_clean_artifact(units, document_id, original_file_name):
    units = normalize_units(units)
    base_name = f"reviewed-{hash_payload(document_id)[:12]}"
    sheets = structured_sheets(units)
    if sheets:
        if Path(original_file_name).suffix.lower() == ".csv" and len(sheets) == 1:
            return _csv_artifact(sheets, base_name)
        return _workbook_artifact(sheets, base_name)
    body = []
    for unit in units:
        if unit.locator.get("kind") == "metadata" or not unit.text.strip():
            continue
        page = unit.locator.get("page_number")
        if type(page) is int and page > 0:
            body.append(f"[Source page {page}]")
        body.append(unit.text)
    if not body:
        raise ScreeningValidationError("No content remains. Delete the document instead of publishing an empty replacement.")
    return CleanArtifact(f"{base_name}.txt", "text/plain; charset=utf-8", "\n\n".join(body).encode("utf-8"))


def table_schema_text(units):
    sections = []
    for _index, sheet in sorted(structured_sheets(units).items()):
        cells = sheet["cells"]
        row_count = max((row for row, _column in cells), default=0)
        column_count = max((column for _row, column in cells), default=0)
        headers = [str(cells.get((1, column)) or f"Column {column}") for column in range(1, column_count + 1)]
        sections.append(
            f"Sheet: {sheet['name']}\nRows: {max(0, row_count - 1)}\n"
            f"Columns: {', '.join(headers)}"
        )
    return "\n\n".join(sections)
