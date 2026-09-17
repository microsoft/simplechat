# extraction.py
"""Private canonical-content capture for the existing extraction processors."""

import csv
import json
import math
from contextlib import contextmanager
from contextvars import ContextVar
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path

from content_screening.contracts import (
    CONTENT_METADATA_FIELDS,
    ContentUnit,
    ScreeningError,
    ScreeningValidationError,
    Subject,
    normalize_units,
)


MAX_CANONICAL_CHARACTERS = 64 * 1024 * 1024
MAX_TABLE_CELLS = 1000000
TABLE_EXTENSIONS = frozenset({".csv", ".xlsx", ".xls", ".xlsm"})
_EXTRACTION = ContextVar("content_screening_extraction", default=None)
_PUBLICATION = ContextVar("content_screening_publication", default=None)


@dataclass
class ExtractionCapture:
    subject: Subject
    scan_id: str
    file_name: str
    max_characters: int = MAX_CANONICAL_CHARACTERS
    units: list[ContentUnit] = field(default_factory=list)
    fallback_chunks: list[ContentUnit] = field(default_factory=list)
    supplemental_units: list[ContentUnit] = field(default_factory=list)
    mode: str = "chunks"
    source_page_offset: int = 0
    characters: int = 0
    staged_count: int = 0
    completed: bool = False
    preloaded_text: bool = False
    failure_code: str | None = None

    def _append(self, target, unit):
        self.characters += len(unit.text)
        if self.characters > self.max_characters:
            self.failure_code = "screening_content_budget_exceeded"
            raise ScreeningError(code="screening_content_budget_exceeded")
        target.append(unit)

    def add_text(self, text, *, kind="text"):
        if self.mode in {"pages", "table"} or self.preloaded_text:
            return
        if not isinstance(text, str):
            self.failure_code = "screening_invalid_extraction"
            raise ScreeningValidationError("Extracted content must be text.")
        self.mode = "text"
        self._append(self.units, ContentUnit(
            f"source-text-{len(self.units) + 1}", text, {"kind": kind},
        ))

    def add_pages(self, pages, *, kind="page"):
        if not isinstance(pages, list):
            raise ScreeningError(code="screening_invalid_extraction")
        self.mode = "pages"
        last_page = 0
        for page in pages:
            number = page.get("page_number")
            if type(number) is not int or number <= 0:
                raise ScreeningError(code="screening_invalid_page_mapping")
            absolute_page = self.source_page_offset + number
            self._append(self.units, ContentUnit(
                f"source-page-{absolute_page}",
                page.get("content", ""),
                {"kind": kind, "page_number": absolute_page},
            ))
            last_page = max(last_page, number)
        self.source_page_offset += last_page

    def add_supplement(self, text, *, kind="figure", page_number=None):
        if not isinstance(text, str) or not text.strip():
            return
        locator = {"kind": kind}
        if type(page_number) is int and page_number > 0:
            locator["page_number"] = page_number
        self._append(self.supplemental_units, ContentUnit(
            f"supplement-{len(self.supplemental_units) + 1}", text, locator,
        ))

    def add_chunk(self, text, page_number, *, start_time=None):
        if not isinstance(text, str):
            self.failure_code = "screening_invalid_extraction"
            raise ScreeningValidationError("Extracted content must be text.")
        self.staged_count += 1
        if self.mode != "chunks":
            return
        locator = {"kind": "segment", "segment_number": self.staged_count}
        if start_time is not None:
            locator.update({"kind": "transcript", "start_time": str(start_time)})
        self._append(self.fallback_chunks, ContentUnit(
            f"segment-{self.staged_count}", text, locator,
        ))

    def finish(self, document):
        if not self.completed or self.failure_code:
            raise ScreeningError(code=self.failure_code or "screening_extraction_incomplete")
        body = list(self.units if self.mode != "chunks" else self.fallback_chunks)
        body.extend(self.supplemental_units)
        if document.get("vision_analysis"):
            body.append(ContentUnit(
                "source-vision", json.dumps(document["vision_analysis"], ensure_ascii=False),
                {"kind": "vision"},
            ))
        normalize_units(body)
        for name in CONTENT_METADATA_FIELDS:
            value = document.get(name)
            if value in (None, "", [], {}):
                continue
            text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False)
            body.append(ContentUnit(f"metadata-{name}", text, {"kind": "metadata", "field": name}))
        return normalize_units(body)


def current_extraction(document_id=None):
    capture = _EXTRACTION.get()
    if capture is not None and (document_id is None or capture.subject.document_id == document_id):
        return capture
    return None


@contextmanager
def capture_extraction(capture):
    token = _EXTRACTION.set(capture)
    try:
        yield capture
    finally:
        _EXTRACTION.reset(token)


@contextmanager
def publication_context(subject, scan_id, heartbeat=None):
    token = _PUBLICATION.set((subject.key, scan_id, heartbeat))
    try:
        yield
    finally:
        _PUBLICATION.reset(token)


def is_publication(subject, scan_id):
    current = _PUBLICATION.get()
    return current is not None and current[:2] == (subject.key, scan_id)


def heartbeat_publication():
    current = _PUBLICATION.get()
    if current is not None and current[2] is not None:
        current[2]()


def capture_text_before_chunking(text):
    capture = current_extraction()
    if capture is not None:
        capture.add_text(text)


def _cell_text(value):
    if value is None:
        return "", "empty"
    if isinstance(value, bool):
        return "true" if value else "false", "boolean"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(), "date"
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            raise ScreeningError(code="screening_invalid_table_value")
        return str(value), "number"
    return str(value), "text"


def _capture_sheet(capture, sheet_index, sheet_name, rows, *, formula_rows=None):
    capture._append(capture.units, ContentUnit(
        f"sheet-{sheet_index}-name", str(sheet_name),
        {"kind": "table_sheet", "sheet_index": sheet_index},
    ))
    cells_seen = 0
    formulas = iter(formula_rows) if formula_rows is not None else None
    for row_number, row in enumerate(rows, start=1):
        formula_row = next(formulas, ()) if formulas is not None else ()
        for column_number, value in enumerate(row, start=1):
            cells_seen += 1
            if cells_seen > MAX_TABLE_CELLS:
                raise ScreeningError(code="screening_table_budget_exceeded")
            text, value_type = _cell_text(value)
            locator = {
                "kind": "table_cell", "sheet_index": sheet_index, "sheet": str(sheet_name),
                "row": row_number, "column": column_number, "value_type": value_type,
            }
            capture._append(capture.units, ContentUnit(
                f"sheet-{sheet_index}-row-{row_number}-column-{column_number}", text, locator,
            ))
            if column_number <= len(formula_row):
                formula = formula_row[column_number - 1]
                if isinstance(formula, str) and formula.startswith("="):
                    capture._append(capture.units, ContentUnit(
                        f"sheet-{sheet_index}-formula-{row_number}-{column_number}",
                        formula,
                        {**locator, "kind": "table_formula"},
                    ))


def capture_table_source(capture, source_path):
    """Inspect full native table data, not the bounded Search schema summary."""
    extension = Path(source_path).suffix.lower()
    capture.mode = "table"
    if extension == ".csv":
        with open(source_path, "r", encoding="utf-8-sig", newline="") as source:
            _capture_sheet(capture, 1, "Sheet1", csv.reader(source))
        return
    if extension in {".xlsx", ".xlsm"}:
        # Spreadsheet dependencies are loaded only for native table extraction.
        from openpyxl import load_workbook

        values_book = load_workbook(source_path, read_only=True, data_only=True, keep_links=False)
        try:
            formula_book = load_workbook(source_path, read_only=True, data_only=False, keep_links=False)
            try:
                for sheet_index, sheet in enumerate(values_book.worksheets, start=1):
                    if (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_TABLE_CELLS:
                        raise ScreeningError(code="screening_table_budget_exceeded")
                    _capture_sheet(
                        capture, sheet_index, sheet.title, sheet.iter_rows(values_only=True),
                        formula_rows=formula_book[sheet.title].iter_rows(values_only=True),
                    )
            finally:
                formula_book.close()
        finally:
            values_book.close()
        return
    if extension == ".xls":
        # xlrd is the existing legacy-workbook reader; avoid loading it for text files.
        import xlrd

        workbook = xlrd.open_workbook(source_path, on_demand=True)
        try:
            for sheet_index, sheet in enumerate(workbook.sheets(), start=1):
                if sheet.nrows * sheet.ncols > MAX_TABLE_CELLS:
                    raise ScreeningError(code="screening_table_budget_exceeded")

                def iter_rows(current_sheet=sheet):
                    for row_number in range(current_sheet.nrows):
                        values = []
                        for cell in current_sheet.row(row_number):
                            if cell.ctype == xlrd.XL_CELL_DATE:
                                values.append(xlrd.xldate_as_datetime(cell.value, workbook.datemode))
                            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                                values.append(bool(cell.value))
                            else:
                                values.append(cell.value)
                        yield values

                _capture_sheet(capture, sheet_index, sheet.name, iter_rows())
        finally:
            workbook.release_resources()
        return
    raise ScreeningValidationError("This table format is not supported for screening.")


def publication_chunks(units, max_characters, *, overlap=0):
    """Split complete canonical units without losing source page identity."""
    if type(max_characters) is not int or max_characters <= 0:
        raise ScreeningValidationError("A valid publication chunk size is required.")
    if type(overlap) is not int or not 0 <= overlap < max_characters:
        raise ScreeningValidationError("Publication overlap is invalid.")
    chunks = []
    for unit in normalize_units(units):
        if unit.locator.get("kind") in {"metadata", "table_cell", "table_formula", "table_sheet"}:
            continue
        if not unit.text.strip():
            continue
        start = 0
        while start < len(unit.text):
            end = min(start + max_characters, len(unit.text))
            chunks.append({
                "chunk_text": unit.text[start:end],
                "chunk_sequence": len(chunks) + 1,
                "page_number": unit.locator.get("page_number"),
                "unit_id": unit.unit_id,
                "start": start,
                "end": end,
                "start_time": unit.locator.get("start_time"),
            })
            if end == len(unit.text):
                break
            start = end - overlap
    return chunks
