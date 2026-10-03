# functions_tabular_merge.py
"""Deterministic merging and inspection of CSV and Excel files.

Version: 0.261.224
Implemented in: 0.261.218
Reconciliation policies, sheet modes, duplicate removal, sorting and inspection added in: 0.261.219
Single-file merges for workflow files found at run time (min_sources=1) added in: 0.261.220
Workbooks with unsafe or unreadable XML refused in: 0.261.224

The engine is pure: it receives already-authorized byte loaders, never resolves
documents, settings, storage, routes, or models, and performs no model work.
Values are kept as text so leading zeros, codes, and identifiers survive exactly.
Rows are spooled to a bounded temporary file during the single parse, so callers
learn the exact row count before they persist or render the merged result.

Schema policies decide which columns the merged table has:

- ``by_name``: every file has the same columns, in any order (the first file's spelling is kept).
- ``exact_order``: every file has the same columns in the same order.
- ``union``: every column from every file; a file without a column leaves it blank (null).
- ``mapped``: exactly the requested ``columns``; other columns are left out and reported.

``column_aliases`` rename alternative headers to one target name before matching, under
every policy. Files whose columns do not fit fail the merge, or are left out and reported
when ``on_incompatible`` is ``exclude``.
"""

import bisect
import codecs
import csv
from array import array
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta, timezone
import hashlib
import io
import json
import math
import os
import re
import tempfile
from typing import Callable, List, Mapping, Optional, Sequence, Tuple
import unicodedata
import zipfile

from functions_ooxml_package_guard import UnreadablePackageError, first_unsafe_xml_part

TABULAR_MERGE_REPORT_VERSION = "tabular-merge-report-v1"
TABULAR_INSPECTION_VERSION = "tabular-inspection-v1"

SCHEMA_POLICY_BY_NAME = "by_name"
SCHEMA_POLICY_EXACT_ORDER = "exact_order"
SCHEMA_POLICY_UNION = "union"
SCHEMA_POLICY_MAPPED = "mapped"
TABULAR_MERGE_SCHEMA_POLICIES = (
    SCHEMA_POLICY_BY_NAME, SCHEMA_POLICY_EXACT_ORDER, SCHEMA_POLICY_UNION, SCHEMA_POLICY_MAPPED,
)

SHEETS_FIRST = "first"
SHEETS_ALL = "all"
TABULAR_MERGE_SHEET_MODES = (SHEETS_FIRST, SHEETS_ALL)

ON_INCOMPATIBLE_FAIL = "fail"
ON_INCOMPATIBLE_EXCLUDE = "exclude"
TABULAR_MERGE_INCOMPATIBLE_MODES = (ON_INCOMPATIBLE_FAIL, ON_INCOMPATIBLE_EXCLUDE)

DEDUPE_NONE = "none"
DEDUPE_EXACT_ROWS = "exact_rows"
DEDUPE_KEY_COLUMNS = "key_columns"
TABULAR_MERGE_DEDUPE_MODES = (DEDUPE_NONE, DEDUPE_EXACT_ROWS, DEDUPE_KEY_COLUMNS)
DEDUPE_KEEP_FIRST = "first"
DEDUPE_KEEP_LAST = "last"
TABULAR_MERGE_DEDUPE_KEEP = (DEDUPE_KEEP_FIRST, DEDUPE_KEEP_LAST)

SORT_TYPE_TEXT = "text"
SORT_TYPE_NUMBER = "number"
SORT_TYPE_DATE = "date"
TABULAR_MERGE_SORT_TYPES = (SORT_TYPE_TEXT, SORT_TYPE_NUMBER, SORT_TYPE_DATE)

TABULAR_MERGE_FORMAT_CSV = "csv"
TABULAR_MERGE_FORMAT_XLSX = "xlsx"
TABULAR_MERGE_FORMAT_XLS = "xls"
TABULAR_MERGE_EXTENSION_FORMATS = {
    ".csv": TABULAR_MERGE_FORMAT_CSV,
    ".xlsx": TABULAR_MERGE_FORMAT_XLSX,
    ".xlsm": TABULAR_MERGE_FORMAT_XLSX,
    ".xls": TABULAR_MERGE_FORMAT_XLS,
}
TABULAR_MERGE_SUPPORTED_EXTENSIONS = tuple(TABULAR_MERGE_EXTENSION_FORMATS)

DEFAULT_SOURCE_COLUMN_NAME = "Source File"
DEFAULT_SHEET_COLUMN_NAME = "Source Sheet"
MAX_SOURCE_COLUMN_NAME_CHARS = 128
MAX_SHEET_NAME_CHARS = 31
MAX_COLUMN_NAME_BYTES = 256
MAX_HEADER_ROW = 1000
MAX_MAPPED_COLUMNS = 255
MAX_ALIAS_TARGETS = 256
MAX_ALIASES_PER_TARGET = 32
MAX_DEDUPE_COLUMNS = 16
MAX_SORT_COLUMNS = 3
DEFAULT_INSPECT_SAMPLE_ROWS = 3
MAX_INSPECT_SAMPLE_ROWS = 10

SOURCE_STATUS_MERGED = "merged"
SOURCE_STATUS_COMPATIBLE = "compatible"
SOURCE_STATUS_SCHEMA_MISMATCH = "schema_mismatch"
SOURCE_STATUS_EXCLUDED = "excluded"
SOURCE_STATUS_SKIPPED = "skipped"
SOURCE_STATUS_NOT_PROCESSED = "not_processed"

_CSV_DELIMITERS = (",", ";", "\t", "|")
_CSV_SNIFF_BYTES = 64 * 1024
_SPOOL_MEMORY_BYTES = 8 * 1024 * 1024
_CANCEL_CHECK_INTERVAL = 1000
_ZIP_MAGIC = b"PK\x03\x04"
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_MAX_EXACT_INTEGER_FLOAT = float(2 ** 53)
_INSPECT_CELL_CHARS = 80
_INSPECT_MAX_BYTES = 96 * 1024
_NUMBER_TEXT = re.compile(
    r"[+-]?(?:(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?|\.\d+)(?:[eE][+-]?\d+)?"
)
_CURRENCY_PREFIXES = "$\u20ac\u00a3\u00a5"
_NOT_WORD = re.compile(r"[\W_]+")

_EXCEL_LIMITATIONS = (
    "Excel formulas are merged as their last calculated values; a workbook saved without "
    "calculated values merges those cells as blanks.",
    "Excel numbers, booleans and dates are written as plain text, with dates in ISO 8601 form.",
)
_EXCLUDED_LIMITATION = (
    "{count} of {total} file(s) or sheet(s) were left out because their columns don't fit the "
    "merge; the merge report lists each one and why."
)
_BLANK_COLUMNS_LIMITATION = (
    "Columns that some files don't have are left blank for those files' rows; the merge report "
    "lists the missing columns."
)
_IGNORED_COLUMNS_LIMITATION = (
    "Columns that aren't in the requested column list were left out; the merge report lists them."
)


class TabularMergeError(ValueError):
    """A safe, user-presentable merge failure with a stable code."""

    def __init__(self, code, message, *, report=None, details=None):
        super().__init__(message)
        self.code = code
        self.message = message
        self.report = report
        self.details = dict(details or {})


class TabularMergeCancelled(TabularMergeError):
    """The caller asked the merge to stop."""

    def __init__(self):
        super().__init__("cancelled", "The merge was cancelled.")


@dataclass(frozen=True)
class TabularMergeLimits:
    """Bounds for one merge; every value must be a positive integer."""

    max_sources: int = 10
    max_total_rows: int = 250_000
    max_source_bytes: int = 100 * 1024 * 1024
    max_columns: int = 256
    max_header_bytes: int = 256
    max_cell_chars: int = 32_767
    max_workbook_uncompressed_bytes: int = 512 * 1024 * 1024
    max_workbook_entries: int = 10_000
    max_sort_rows: int = 250_000
    max_sheets_per_source: int = 100

    def __post_init__(self):
        for name in (
            "max_sources", "max_total_rows", "max_source_bytes", "max_columns",
            "max_header_bytes", "max_cell_chars", "max_workbook_uncompressed_bytes",
            "max_workbook_entries", "max_sort_rows", "max_sheets_per_source",
        ):
            value = getattr(self, name)
            if type(value) is not int or value < 1:
                raise ValueError(f"{name} must be a positive integer.")
        if self.max_sources < 2:
            raise ValueError("max_sources must allow at least two sources.")


@dataclass(frozen=True)
class TabularMergeSort:
    """One sort key: a merged column, its direction and how its values compare."""

    column: str
    descending: bool = False
    value_type: str = SORT_TYPE_TEXT

    def __post_init__(self):
        name = _display_header(self.column) if isinstance(self.column, str) else ""
        if not name or len(name.encode("utf-8")) > MAX_COLUMN_NAME_BYTES:
            raise TabularMergeError("invalid_options", "Each sort key needs a column name.")
        object.__setattr__(self, "column", name)
        if type(self.descending) is not bool:
            raise TabularMergeError("invalid_options", "A sort direction must be true or false.")
        if self.value_type not in TABULAR_MERGE_SORT_TYPES:
            raise TabularMergeError(
                "invalid_options", f"A sort type must be one of: {', '.join(TABULAR_MERGE_SORT_TYPES)}.",
            )

    def to_dict(self):
        return {"column": self.column, "descending": self.descending, "value_type": self.value_type}


@dataclass(frozen=True)
class TabularMergeOptions:
    """How sources are matched, combined, de-duplicated, ordered and attributed."""

    schema_policy: str = SCHEMA_POLICY_BY_NAME
    sheet: Optional[str] = None
    include_source_column: bool = True
    source_column_name: str = DEFAULT_SOURCE_COLUMN_NAME
    sheets: str = SHEETS_FIRST
    header_row: Optional[int] = None
    columns: Tuple[str, ...] = ()
    column_aliases: Tuple[Tuple[str, Tuple[str, ...]], ...] = ()
    on_incompatible: str = ON_INCOMPATIBLE_FAIL
    dedupe: str = DEDUPE_NONE
    dedupe_columns: Tuple[str, ...] = ()
    dedupe_keep: str = DEDUPE_KEEP_FIRST
    sort_by: Tuple[TabularMergeSort, ...] = ()

    def __post_init__(self):
        if self.schema_policy not in TABULAR_MERGE_SCHEMA_POLICIES:
            raise TabularMergeError(
                "invalid_options",
                f"Schema policy must be one of: {', '.join(TABULAR_MERGE_SCHEMA_POLICIES)}.",
            )
        _validate_table_selection(self.sheet, self.sheets, self.header_row)
        if type(self.include_source_column) is not bool:
            raise TabularMergeError("invalid_options", "include_source_column must be true or false.")
        name = _display_header(self.source_column_name) if isinstance(self.source_column_name, str) else ""
        if not name or len(name) > MAX_SOURCE_COLUMN_NAME_CHARS or len(name.encode("utf-8")) > 240:
            raise TabularMergeError(
                "invalid_options",
                f"The source column name must be 1 to {MAX_SOURCE_COLUMN_NAME_CHARS} characters.",
            )

        columns = _column_names(self.columns, "The requested columns", maximum=MAX_MAPPED_COLUMNS)
        object.__setattr__(self, "columns", columns)
        if self.schema_policy == SCHEMA_POLICY_MAPPED and not columns:
            raise TabularMergeError("invalid_options", "The mapped policy needs the list of columns to keep.")
        if self.schema_policy != SCHEMA_POLICY_MAPPED and columns:
            raise TabularMergeError("invalid_options", "A list of columns is used only by the mapped policy.")

        aliases = _alias_pairs(self.column_aliases)
        object.__setattr__(self, "column_aliases", aliases)
        if self.schema_policy == SCHEMA_POLICY_MAPPED:
            targets = {normalize_header_name(column) for column in columns}
            if any(normalize_header_name(target) not in targets for target, _ in aliases):
                raise TabularMergeError(
                    "invalid_options", "Every alias must rename a header to one of the requested columns.",
                )

        if self.on_incompatible not in TABULAR_MERGE_INCOMPATIBLE_MODES:
            raise TabularMergeError(
                "invalid_options",
                f"on_incompatible must be one of: {', '.join(TABULAR_MERGE_INCOMPATIBLE_MODES)}.",
            )
        if self.dedupe not in TABULAR_MERGE_DEDUPE_MODES:
            raise TabularMergeError(
                "invalid_options", f"dedupe must be one of: {', '.join(TABULAR_MERGE_DEDUPE_MODES)}.",
            )
        if self.dedupe_keep not in TABULAR_MERGE_DEDUPE_KEEP:
            raise TabularMergeError(
                "invalid_options", f"dedupe_keep must be one of: {', '.join(TABULAR_MERGE_DEDUPE_KEEP)}.",
            )
        dedupe_columns = _column_names(
            self.dedupe_columns, "The duplicate key columns", maximum=MAX_DEDUPE_COLUMNS,
        )
        object.__setattr__(self, "dedupe_columns", dedupe_columns)
        if self.dedupe == DEDUPE_KEY_COLUMNS and not dedupe_columns:
            raise TabularMergeError("invalid_options", "Removing duplicates by key needs the key columns.")
        if self.dedupe != DEDUPE_KEY_COLUMNS and dedupe_columns:
            raise TabularMergeError("invalid_options", "Key columns are used only when removing duplicates by key.")

        sort_by = tuple(
            item if isinstance(item, TabularMergeSort) else _sort_from_mapping(item)
            for item in (self.sort_by or ())
        )
        if len(sort_by) > MAX_SORT_COLUMNS:
            raise TabularMergeError("invalid_options", f"Sort by at most {MAX_SORT_COLUMNS} columns.")
        if len({normalize_header_name(item.column) for item in sort_by}) != len(sort_by):
            raise TabularMergeError("invalid_options", "Each column can be sorted on only once.")
        object.__setattr__(self, "sort_by", sort_by)

    def policy_report(self):
        """The options as plain data for the merge report."""
        return {
            "schema_policy": self.schema_policy,
            "sheet": self.sheet,
            "sheets": self.sheets,
            "header_row": self.header_row,
            "include_source_column": self.include_source_column,
            "source_column_name": _display_header(self.source_column_name) if self.include_source_column else None,
            "columns": list(self.columns),
            "column_aliases": {target: list(aliases) for target, aliases in self.column_aliases},
            "on_incompatible": self.on_incompatible,
            "dedupe": self.dedupe,
            "dedupe_columns": list(self.dedupe_columns),
            "dedupe_keep": self.dedupe_keep if self.dedupe != DEDUPE_NONE else None,
            "sort_by": [item.to_dict() for item in self.sort_by],
        }


@dataclass(frozen=True)
class TabularInspectOptions:
    """Which sheets and header row to inspect, and how many sample rows to keep."""

    sheet: Optional[str] = None
    sheets: str = SHEETS_FIRST
    header_row: Optional[int] = None
    sample_rows: int = DEFAULT_INSPECT_SAMPLE_ROWS

    def __post_init__(self):
        _validate_table_selection(self.sheet, self.sheets, self.header_row)
        if type(self.sample_rows) is not int or not 0 <= self.sample_rows <= MAX_INSPECT_SAMPLE_ROWS:
            raise TabularMergeError(
                "invalid_options", f"Sample rows must be a whole number from 0 to {MAX_INSPECT_SAMPLE_ROWS}.",
            )


@dataclass(frozen=True)
class TabularMergeSource:
    """One authorized source. ``load_bytes`` returns the original file bytes."""

    source_id: str
    file_name: str
    load_bytes: Callable[[], bytes] = field(repr=False)
    extension: str = ""

    def resolved_extension(self):
        extension = (self.extension or os.path.splitext(self.file_name or "")[1]).strip().lower()
        if extension and not extension.startswith("."):
            extension = f".{extension}"
        return extension


def _validate_table_selection(sheet, sheets, header_row):
    if sheet is not None:
        if not isinstance(sheet, str) or not sheet.strip():
            raise TabularMergeError("invalid_options", "A sheet name must be nonempty text.")
        if len(sheet) > MAX_SHEET_NAME_CHARS:
            raise TabularMergeError(
                "invalid_options", f"A sheet name has at most {MAX_SHEET_NAME_CHARS} characters.",
            )
    if sheets not in TABULAR_MERGE_SHEET_MODES:
        raise TabularMergeError(
            "invalid_options", f"sheets must be one of: {', '.join(TABULAR_MERGE_SHEET_MODES)}.",
        )
    if sheet is not None and sheets == SHEETS_ALL:
        raise TabularMergeError("invalid_options", "Name one sheet, or read all sheets, but not both.")
    if header_row is not None and (type(header_row) is not int or not 1 <= header_row <= MAX_HEADER_ROW):
        raise TabularMergeError(
            "invalid_options", f"The header row must be a whole number from 1 to {MAX_HEADER_ROW}.",
        )


def _column_name(value, label):
    name = _display_header(value) if isinstance(value, str) else ""
    if not name or len(name.encode("utf-8")) > MAX_COLUMN_NAME_BYTES:
        raise TabularMergeError(
            "invalid_options", f"{label} must be nonempty names of at most {MAX_COLUMN_NAME_BYTES} bytes.",
        )
    return name


def _column_names(values, label, *, maximum):
    if values is None:
        return ()
    if isinstance(values, (str, bytes)) or not isinstance(values, (list, tuple)):
        raise TabularMergeError("invalid_options", f"{label} must be a list of column names.")
    names = tuple(_column_name(value, label) for value in values)
    if len(names) > maximum:
        raise TabularMergeError("invalid_options", f"{label} can name at most {maximum} columns.")
    if len({name.casefold() for name in names}) != len(names):
        raise TabularMergeError("invalid_options", f"{label} name a column more than once.")
    return names


def _alias_pairs(value):
    """Validated ``((target, (alias, ...)), ...)``; identity aliases are dropped."""
    if not value:
        return ()
    if isinstance(value, Mapping):
        items = list(value.items())
    elif isinstance(value, (list, tuple)):
        items = list(value)
    else:
        raise TabularMergeError("invalid_options", "Column aliases must map each column to its other names.")
    if len(items) > MAX_ALIAS_TARGETS:
        raise TabularMergeError("invalid_options", f"Column aliases can rename at most {MAX_ALIAS_TARGETS} columns.")
    pairs = []
    targets = {}
    for item in items:
        if not isinstance(item, (list, tuple)) or len(item) != 2:
            raise TabularMergeError("invalid_options", "Column aliases must map each column to its other names.")
        target, aliases = item
        target = _column_name(target, "Alias targets")
        if target.casefold() in targets:
            raise TabularMergeError("invalid_options", "Column aliases name a target column more than once.")
        if isinstance(aliases, (str, bytes)) or not isinstance(aliases, (list, tuple)) or not aliases:
            raise TabularMergeError("invalid_options", "Each aliased column needs a list of its other names.")
        if len(aliases) > MAX_ALIASES_PER_TARGET:
            raise TabularMergeError(
                "invalid_options", f"A column can have at most {MAX_ALIASES_PER_TARGET} other names.",
            )
        names = []
        for alias in aliases:
            name = _column_name(alias, "Column aliases")
            if name.casefold() != target.casefold() and name.casefold() not in {item.casefold() for item in names}:
                names.append(name)
        targets[target.casefold()] = names
        pairs.append((target, tuple(names)))
    seen = {}
    for target, names in pairs:
        for name in names:
            key = name.casefold()
            if key in targets:
                raise TabularMergeError(
                    "invalid_options", f'"{name}" is both a column name and another column\'s alias.',
                )
            if key in seen and seen[key] != target.casefold():
                raise TabularMergeError("invalid_options", f'"{name}" is an alias of more than one column.')
            seen[key] = target.casefold()
    return tuple((target, names) for target, names in pairs if names)


def _sort_from_mapping(value):
    if not isinstance(value, Mapping):
        raise TabularMergeError("invalid_options", "Each sort key must name a column.")
    unknown = set(value) - {"column", "descending", "value_type"}
    if unknown:
        raise TabularMergeError("invalid_options", "A sort key has an unsupported setting.")
    return TabularMergeSort(
        column=value.get("column"),
        descending=value.get("descending", False),
        value_type=value.get("value_type", SORT_TYPE_TEXT),
    )


def tabular_merge_format_for_file_name(file_name):
    """The merge reader format for a file name, or None when unsupported."""
    extension = os.path.splitext(str(file_name or ""))[1].strip().lower()
    return TABULAR_MERGE_EXTENSION_FORMATS.get(extension)


def tabular_merge_options_from_arguments(arguments, *, mapping=None, mapping_bound=False):
    """Engine options from plan arguments; a bound column mapping makes the policy ``mapped``.

    At planning time the mapping is not prepared yet, so ``mapping_bound`` checks that the
    other arguments can be combined with one. Invalid combinations raise ``invalid_options``.
    """
    arguments = arguments if isinstance(arguments, Mapping) else {}
    policy = arguments.get("schema_policy") or SCHEMA_POLICY_BY_NAME
    columns = arguments.get("columns") or ()
    aliases = arguments.get("column_aliases") or ()
    if mapping is not None or mapping_bound:
        if columns or aliases:
            raise TabularMergeError(
                "invalid_options", "Use either a bound column mapping or explicit columns and aliases, not both.",
            )
        if policy not in (SCHEMA_POLICY_BY_NAME, SCHEMA_POLICY_MAPPED):
            raise TabularMergeError("invalid_options", "A bound column mapping uses the mapped schema policy.")
        policy = SCHEMA_POLICY_MAPPED
        if mapping is None:
            # The prepared mapping supplies the real columns when the merge runs.
            columns = ("Mapped column",)
        else:
            columns = mapping["columns"]
            aliases = mapping["column_aliases"]
    return TabularMergeOptions(
        schema_policy=policy,
        sheet=arguments.get("sheet") or None,
        include_source_column=arguments.get("include_source_column", True) is not False,
        source_column_name=arguments.get("source_column_name") or DEFAULT_SOURCE_COLUMN_NAME,
        sheets=arguments.get("sheets") or SHEETS_FIRST,
        header_row=arguments.get("header_row"),
        columns=columns,
        column_aliases=aliases,
        on_incompatible=arguments.get("on_incompatible") or ON_INCOMPATIBLE_FAIL,
        dedupe=arguments.get("dedupe") or DEDUPE_NONE,
        dedupe_columns=arguments.get("dedupe_columns") or (),
        dedupe_keep=arguments.get("dedupe_keep") or DEDUPE_KEEP_FIRST,
        sort_by=arguments.get("sort_by") or (),
    )


def tabular_inspect_options_from_arguments(arguments):
    """Inspection options from plan arguments."""
    arguments = arguments if isinstance(arguments, Mapping) else {}
    sample_rows = arguments.get("sample_rows")
    return TabularInspectOptions(
        sheet=arguments.get("sheet") or None,
        sheets=arguments.get("sheets") or SHEETS_FIRST,
        header_row=arguments.get("header_row"),
        sample_rows=DEFAULT_INSPECT_SAMPLE_ROWS if sample_rows is None else sample_rows,
    )


# ---------------------------------------------------------------------------
# Prepared column mappings
# ---------------------------------------------------------------------------

TABULAR_COLUMN_MAPPING_PROFILE = "tabular_column_mapping_v1"
MAPPING_CONFIDENCE_LEVELS = ("high", "medium", "low")
MAX_MAPPING_ENTRIES = 1024
MAX_MAPPING_NOTES_CHARS = 2000


def tabular_column_mapping_schema():
    """The prepared-content profile a compose step fills to line up differently named columns."""
    name = {"type": "string", "minLength": 1, "maxLength": MAX_COLUMN_NAME_BYTES}
    return {
        "type": "object",
        "description": (
            "How to line up the columns of several CSV or Excel files before merging them. columns "
            "is the merged table's columns in order; use the clearest existing header for each. "
            "mappings lists every header whose name differs from the column it fills, and every "
            "header to leave out (target_column null). A header that already matches a column, "
            "ignoring case and spacing, needs no entry. Map a header only when it holds the same "
            "kind of value; when unsure, keep it as its own column or mark confidence low."
        ),
        "properties": {
            "columns": {
                "type": "array", "items": dict(name), "minItems": 1, "maxItems": MAX_MAPPED_COLUMNS,
                "uniqueItems": True,
            },
            "mappings": {
                "type": "array",
                "maxItems": MAX_MAPPING_ENTRIES,
                "items": {
                    "type": "object",
                    "properties": {
                        "source_column": dict(name, description="A header exactly as it appears in a file."),
                        "target_column": {
                            "type": ["string", "null"], "maxLength": MAX_COLUMN_NAME_BYTES,
                            "description": "The column in columns this header fills, or null to leave it out.",
                        },
                        "confidence": {"type": "string", "enum": list(MAPPING_CONFIDENCE_LEVELS)},
                    },
                    "required": ["source_column", "target_column", "confidence"],
                    "additionalProperties": False,
                },
            },
            "notes": {
                "type": "string", "maxLength": MAX_MAPPING_NOTES_CHARS,
                "description": "A short explanation of the mapping for the user.",
            },
        },
        "required": ["columns", "mappings"],
        "additionalProperties": False,
    }


def tabular_mapping_from_profile(value):
    """Validate a prepared ``tabular_column_mapping_v1`` value; return merge columns and aliases.

    Raises ``invalid_mapping`` when the value breaks the profile: unknown fields, repeated
    columns or headers, a target that isn't a listed column, or a header that is itself a
    different listed column.
    """
    def invalid(message):
        return TabularMergeError("invalid_mapping", message)

    if not isinstance(value, Mapping) or set(value) - {"columns", "mappings", "notes"}:
        raise invalid("A column mapping has only columns, mappings and notes.")
    raw_columns = value.get("columns")
    raw_mappings = value.get("mappings")
    notes = value.get("notes", "")
    if not isinstance(raw_columns, list) or not 1 <= len(raw_columns) <= MAX_MAPPED_COLUMNS:
        raise invalid(f"A column mapping lists 1 to {MAX_MAPPED_COLUMNS} columns.")
    if not isinstance(raw_mappings, list) or len(raw_mappings) > MAX_MAPPING_ENTRIES:
        raise invalid(f"A column mapping has at most {MAX_MAPPING_ENTRIES} mappings.")
    if not isinstance(notes, str) or len(notes) > MAX_MAPPING_NOTES_CHARS:
        raise invalid(f"Mapping notes have at most {MAX_MAPPING_NOTES_CHARS} characters.")
    try:
        columns = _column_names(raw_columns, "The mapped columns", maximum=MAX_MAPPED_COLUMNS)
    except TabularMergeError as exc:
        raise invalid(exc.message) from exc
    targets = {column.casefold(): column for column in columns}
    aliases = {}
    ignored = []
    low_confidence = []
    sources = set()
    for item in raw_mappings:
        if not isinstance(item, Mapping) or set(item) != {"source_column", "target_column", "confidence"}:
            raise invalid("Each mapping has exactly source_column, target_column and confidence.")
        if item["confidence"] not in MAPPING_CONFIDENCE_LEVELS:
            raise invalid(f"Mapping confidence is one of: {', '.join(MAPPING_CONFIDENCE_LEVELS)}.")
        try:
            source = _column_name(item["source_column"], "Mapped headers")
        except TabularMergeError as exc:
            raise invalid(exc.message) from exc
        key = source.casefold()
        if key in sources:
            raise invalid(f'"{source}" is mapped more than once.')
        sources.add(key)
        target = item["target_column"]
        if target is None:
            if key in targets:
                raise invalid(f'"{source}" is a listed column, so it can\'t be left out.')
            ignored.append(source)
            continue
        target_key = _display_header(target).casefold() if isinstance(target, str) else ""
        if target_key not in targets:
            raise invalid(f'"{source}" is mapped to a column that isn\'t listed.')
        if key == target_key:
            continue
        if key in targets:
            raise invalid(f'"{source}" is a listed column, so it can\'t also fill "{targets[target_key]}".')
        aliases.setdefault(targets[target_key], []).append(source)
        if item["confidence"] == "low":
            low_confidence.append({"source_column": source, "target_column": targets[target_key]})
    for target, names in aliases.items():
        if len(names) > MAX_ALIASES_PER_TARGET:
            raise invalid(f'At most {MAX_ALIASES_PER_TARGET} headers can fill "{target}".')
    return {
        "columns": columns,
        "column_aliases": aliases,
        "ignored": ignored,
        "low_confidence": low_confidence,
        "notes": notes.strip(),
    }


def iter_csv_rows(content, file_name="A selected file"):
    """Decode CSV bytes; return ``(encoding, delimiter, rows)`` where rows yields lists of text."""
    source = TabularMergeSource(file_name, file_name, lambda: content)
    text, encoding = _decode_csv(content, source)
    delimiter = _detect_delimiter(text)
    reader = csv.reader(
        io.StringIO(text, newline=""), delimiter=delimiter, quotechar='"', doublequote=True, strict=False,
    )

    def rows():
        try:
            for values in reader:
                yield values
        except csv.Error as exc:
            raise TabularMergeError(
                "unreadable_csv", f"{_safe_name(source)} couldn't be read as CSV near line {reader.line_num}.",
            ) from exc

    return encoding, delimiter, rows()


def normalize_header_name(value):
    """The comparison key for a header: NFC, collapsed whitespace, casefolded."""
    return _display_header(value).casefold()


def squash_header_name(value):
    """A looser key that also ignores spaces, punctuation and underscores, for near matches."""
    return _NOT_WORD.sub("", normalize_header_name(value))


def _display_header(value):
    text = unicodedata.normalize("NFC", str(value if value is not None else ""))
    return " ".join(text.split())


# ---------------------------------------------------------------------------
# Spooled rows and the merge result
# ---------------------------------------------------------------------------


class _RowSpool:
    """Validated rows as compact JSON lines; small merges stay in memory."""

    def __init__(self):
        self._file = tempfile.SpooledTemporaryFile(max_size=_SPOOL_MEMORY_BYTES, mode="w+b")
        self._offsets = array("q")
        self._end = 0
        self.count = 0
        self.closed = False

    def append(self, values):
        data = json.dumps(values, ensure_ascii=False, separators=(",", ":")).encode("utf-8") + b"\n"
        self._file.seek(self._end)
        self._file.write(data)
        self._offsets.append(self._end)
        self._end += len(data)
        self.count += 1

    def _read(self, index):
        self._file.seek(self._offsets[index])
        line = self._file.readline()
        if not line:
            raise TabularMergeError("count_mismatch", "The merged rows could not be read back completely.")
        return json.loads(line)

    def iter_indexed(self, order=None):
        """``(index, values)`` in spool order, or in ``order`` when given."""
        if self.closed:
            raise TabularMergeError("result_closed", "The merged rows are no longer available.")
        self._file.flush()
        indices = range(self.count) if order is None else order
        for index in indices:
            if self.closed:
                raise TabularMergeError("result_closed", "The merged rows are no longer available.")
            yield index, self._read(index)

    def close(self):
        if not self.closed:
            self.closed = True
            self._file.close()


class TabularMergeResult:
    """The merged table: ordered columns, an exact row count, rows and a report.

    Rows can be iterated more than once until ``close()``; each pass verifies the
    count. Use it as a context manager so the spooled rows are always released.
    """

    def __init__(self, columns, report, spool, *, nullable=(), row_count=None, order=None, keep=None):
        self._columns = tuple(columns)
        self._nullable = frozenset(nullable)
        self._report = report
        self._spool = spool
        self._row_count = spool.count if row_count is None else row_count
        self._order = order
        self._keep = keep

    @property
    def columns(self):
        return self._columns

    @property
    def nullable_columns(self):
        """Columns that hold null for rows from files that don't have them."""
        return self._nullable

    @property
    def row_count(self):
        return self._row_count

    @property
    def report(self):
        return json.loads(json.dumps(self._report))

    def iter_rows(self):
        width = len(self._columns)
        yielded = 0
        for index, values in self._spool.iter_indexed(self._order):
            if len(values) < width:
                values.extend([None] * (width - len(values)))
            if self._keep is not None and not self._keep(index, values):
                continue
            yielded += 1
            yield values
        if yielded != self._row_count:
            raise TabularMergeError("count_mismatch", "The merged rows could not be read back completely.")

    def iter_records(self):
        columns = self._columns
        for values in self.iter_rows():
            yield dict(zip(columns, values))

    def close(self):
        self._spool.close()

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, traceback):
        self.close()
        return False


# ---------------------------------------------------------------------------
# Merging
# ---------------------------------------------------------------------------


def merge_tabular_sources(
    sources: Sequence[TabularMergeSource],
    *,
    options: Optional[TabularMergeOptions] = None,
    limits: Optional[TabularMergeLimits] = None,
    cancel_requested: Optional[Callable[[], bool]] = None,
    on_progress: Optional[Callable[[dict], None]] = None,
    min_sources: int = 2,
) -> TabularMergeResult:
    """Append the rows of the sources, in the given order, into one table.

    The first file (or sheet) defines the column order for ``by_name`` and ``exact_order``;
    ``union`` adds each new column the first time a file has it; ``mapped`` uses exactly
    the requested columns. Under ``on_incompatible="fail"`` a mismatch reads every remaining
    header so the report lists all mismatches, then nothing is merged. ``min_sources`` is 1
    only for a scheduled merge of whatever files arrived, which may be a single file.
    """
    if min_sources not in (1, 2):
        raise ValueError("min_sources must be 1 or 2.")
    options = options or TabularMergeOptions()
    limits = limits or TabularMergeLimits()
    sources = list(sources or ())
    _validate_sources(sources, limits, minimum=min_sources)
    merger = _Merger(options, limits, len(sources), cancel_requested)
    try:
        for index, source in enumerate(sources):
            _check_cancel(cancel_requested)
            rows = merger.add_source(source)
            if on_progress is not None:
                on_progress({
                    "index": index + 1, "total": len(sources), "file_name": source.file_name, "rows": rows,
                })
        return merger.finish()
    except BaseException:
        merger.discard()
        raise


class _Merger:
    """The single-pass state of one merge."""

    def __init__(self, options, limits, source_count, cancel_requested):
        self.options = options
        self.limits = limits
        self.cancel_requested = cancel_requested
        self.report = _new_report(options, source_count)
        self.alias_index = _alias_index(options)
        self.provenance = None
        self.data_display = []
        self.data_keys = []
        self.data_index = {}
        self.seen_counts = []
        self.reference = None
        self.mismatch = False
        self.merged_tables = 0
        self.excluded_tables = 0
        self.rows_read = 0
        self.blank_rows = 0
        self.duplicates_removed = 0
        self.sheet_null = False
        self.dedupe_keys = [normalize_header_name(name) for name in options.dedupe_columns]
        self.sort_keys = [normalize_header_name(item.column) for item in options.sort_by]
        self.seen_digests = set()
        self.last_digest = {}
        self.row_starts = []
        self.row_entries = []
        self.sort_values = []
        if options.schema_policy == SCHEMA_POLICY_MAPPED:
            for name in options.columns:
                self._add_data_column(name, normalize_header_name(name))
            self._init_provenance(set(self.data_keys))
            self._require_named_columns(self.dedupe_keys, "dedupe_column_not_found", "remove duplicates")
            self._require_named_columns(self.sort_keys, "sort_column_not_found", "sort")
        # Created last, so a refused configuration leaves nothing to release.
        self.spool = _RowSpool()

    def discard(self):
        self.spool.close()

    def add_source(self, source):
        content = _load_source_bytes(source, self.limits)
        reader = _open_reader(source, content, self.limits)
        rows = 0
        try:
            try:
                plan = reader.plan(self.options, self.limits)
            except TabularMergeError as exc:
                if exc.code != "sheet_not_found" or self.options.on_incompatible != ON_INCOMPATIBLE_EXCLUDE:
                    raise
                entry = _new_source_entry(source, reader, self.options.sheet)
                self.report["sources"].append(entry)
                self._incompatible(entry, "sheet_not_found")
                return 0
            for sheet, skip_reason in plan:
                _check_cancel(self.cancel_requested)
                entry = _new_source_entry(source, reader, sheet)
                self.report["sources"].append(entry)
                if skip_reason is not None:
                    entry["status"] = SOURCE_STATUS_SKIPPED
                    entry["reason"] = skip_reason
                    continue
                try:
                    table = reader.open_table(sheet, self.options.header_row)
                except TabularMergeError as exc:
                    if exc.code != "empty_source" or isinstance(exc, TabularMergeCancelled):
                        raise
                    if self.options.sheets == SHEETS_ALL:
                        entry["status"] = SOURCE_STATUS_SKIPPED
                        entry["reason"] = "no_header_row"
                        continue
                    if self.options.on_incompatible != ON_INCOMPATIBLE_EXCLUDE:
                        raise
                    self._incompatible(entry, "no_header_row")
                    continue
                try:
                    entry["header_row"] = table.header_row
                    self._add_table(source, table, entry)
                finally:
                    table.close()
                rows += entry["rows"]
        finally:
            reader.close()
        return rows

    # -- columns -----------------------------------------------------------

    def _add_data_column(self, display, key):
        self.data_index[key] = len(self.data_keys)
        self.data_keys.append(key)
        self.data_display.append(display)
        self.seen_counts.append(0)

    def _init_provenance(self, taken_keys):
        names = []
        if self.options.include_source_column:
            taken = set(taken_keys)
            names.append(self._claim_provenance_name(
                _display_header(self.options.source_column_name), taken, "source_column_renamed",
                "file name column",
            ))
            if self.options.sheets == SHEETS_ALL:
                taken.add(names[0].casefold())
                names.append(self._claim_provenance_name(
                    DEFAULT_SHEET_COLUMN_NAME, taken, "sheet_column_renamed", "sheet name column",
                ))
        self.provenance = names
        self.report["source_column"] = names[0] if names else None
        self.report["sheet_column"] = names[1] if len(names) > 1 else None
        self._check_width()

    def _claim_provenance_name(self, base, taken, code, label):
        candidate = base
        suffix = 2
        while candidate.casefold() in taken:
            candidate = _suffixed_name(base, suffix)
            suffix += 1
        if candidate != base:
            self.report["warnings"].append({
                "code": code,
                "message": f'The files already have a "{base}" column, so the {label} is "{candidate}".',
            })
        return candidate

    def _unique_data_name(self, display):
        taken = {name.casefold() for name in (self.provenance or []) + self.data_display}
        if display.casefold() not in taken:
            return display
        suffix = 2
        candidate = _suffixed_name(display, suffix)
        while candidate.casefold() in taken:
            suffix += 1
            candidate = _suffixed_name(display, suffix)
        self.report["warnings"].append({
            "code": "column_renamed",
            "message": f'A column named "{display}" is shown as "{candidate}" because that name is already used.',
        })
        return candidate

    def _check_width(self):
        width = len(self.provenance or ()) + len(self.data_keys)
        if width > self.limits.max_columns:
            raise TabularMergeError(
                "too_many_columns",
                f"The merged table would have {width} columns; at most {self.limits.max_columns} are supported.",
            )

    def _require_named_columns(self, keys, code, purpose):
        missing = [key for key in keys if key not in self.data_index]
        if missing:
            raise TabularMergeError(
                code, f"The column(s) to {purpose} by aren't in the merged table: {', '.join(_quoted(missing))}.",
            )

    # -- tables ------------------------------------------------------------

    def _add_table(self, source, table, entry):
        header = _header_names(table.header, source, self.limits, entry)
        header = _apply_aliases(header, self.alias_index, source, entry)
        policy = self.options.schema_policy
        if policy in (SCHEMA_POLICY_BY_NAME, SCHEMA_POLICY_EXACT_ORDER):
            if self.reference is None:
                self.reference = header
                for display, key in header:
                    self._add_data_column(display, key)
                self._init_provenance(set(self.data_keys))
                self._require_named_columns(self.dedupe_keys, "dedupe_column_not_found", "remove duplicates")
                self._require_named_columns(self.sort_keys, "sort_column_not_found", "sort")
                mapping = list(range(len(header)))
            else:
                mapping = _match_header(self.reference, header, policy, entry)
                if mapping is None:
                    self._incompatible(entry, "columns_differ")
                    return
        elif policy == SCHEMA_POLICY_UNION:
            mapping = self._union_mapping(header, entry)
        else:
            mapping = self._mapped_mapping(header, entry)
            if mapping is None:
                self._incompatible(entry, "no_matching_columns")
                return

        if self.mismatch:
            entry["status"] = SOURCE_STATUS_NOT_PROCESSED
            return
        self._append_rows(source, table, entry, mapping, len(header))
        entry["status"] = SOURCE_STATUS_MERGED
        self.merged_tables += 1
        for position, column in enumerate(mapping):
            if column is not None:
                self.seen_counts[position] += 1

    def _union_mapping(self, header, entry):
        first = self.provenance is None
        if first:
            self._init_provenance({key for _, key in header})
        added = []
        for display, key in header:
            if key not in self.data_index:
                name = self._unique_data_name(display)
                self._add_data_column(name, key)
                self._check_width()
                added.append(name)
        positions = {key: index for index, (_, key) in enumerate(header)}
        mapping = [positions.get(key) for key in self.data_keys]
        if not first:
            entry["missing_columns"] = [
                name for name, column in zip(self.data_display, mapping) if column is None
            ]
            entry["extra_columns"] = added
        return mapping

    def _mapped_mapping(self, header, entry):
        positions = {key: index for index, (_, key) in enumerate(header)}
        mapping = [positions.get(key) for key in self.data_keys]
        entry["missing_columns"] = [name for name, column in zip(self.data_display, mapping) if column is None]
        entry["ignored_columns"] = [display for display, key in header if key not in self.data_index]
        if all(column is None for column in mapping):
            return None
        return mapping

    def _incompatible(self, entry, reason):
        entry["reason"] = reason
        if self.options.on_incompatible == ON_INCOMPATIBLE_EXCLUDE:
            entry["status"] = SOURCE_STATUS_EXCLUDED
            self.excluded_tables += 1
        else:
            entry["status"] = SOURCE_STATUS_SCHEMA_MISMATCH
            self.mismatch = True

    # -- rows --------------------------------------------------------------

    def _append_rows(self, source, table, entry, mapping, width):
        options = self.options
        limits = self.limits
        prefix = []
        if self.provenance:
            prefix.append(source.file_name)
            if len(self.provenance) > 1:
                prefix.append(table.sheet)
                if table.sheet is None:
                    self.sheet_null = True
        self.row_starts.append(self.spool.count)
        self.row_entries.append(entry)
        dedupe_positions = (
            [self.data_index.get(key) for key in self.dedupe_keys]
            if options.dedupe == DEDUPE_KEY_COLUMNS else None
        )
        sort_positions = [self.data_index.get(key) for key in self.sort_keys]
        keep_last = options.dedupe != DEDUPE_NONE and options.dedupe_keep == DEDUPE_KEEP_LAST
        checked = 0
        for row_number, values in table.rows():
            if checked % _CANCEL_CHECK_INTERVAL == 0:
                _check_cancel(self.cancel_requested)
            checked += 1
            cells = list(values)
            if len(cells) > width:
                if any(_cell_has_content(value) for value in cells[width:]):
                    raise TabularMergeError(
                        "row_has_extra_values",
                        f"Row {row_number} in {_safe_name(source)} has more values than there are columns.",
                    )
                cells = cells[:width]
            if not any(_cell_has_content(value) for value in cells):
                entry["blank_rows_skipped"] += 1
                self.blank_rows += 1
                continue
            if len(cells) < width:
                entry["short_rows_padded"] += 1
                cells.extend([""] * (width - len(cells)))
            for value in cells:
                if len(value) > limits.max_cell_chars:
                    raise TabularMergeError(
                        "cell_too_large",
                        f"Row {row_number} in {_safe_name(source)} has a value longer than "
                        f"{limits.max_cell_chars} characters.",
                    )
            if self.rows_read >= limits.max_total_rows:
                raise TabularMergeError(
                    "row_limit_exceeded",
                    f"The selected files hold more than {limits.max_total_rows:,} rows. "
                    "Merge fewer files here, or use a workflow for larger merges.",
                )
            self.rows_read += 1
            data = [cells[column] if column is not None else None for column in mapping]
            if options.dedupe != DEDUPE_NONE:
                digest = _row_digest(
                    data if dedupe_positions is None else _values_at(data, dedupe_positions)
                )
                if digest is not None:
                    if keep_last:
                        previous = self.last_digest.get(digest)
                        if previous is not None:
                            owner = self._entry_for_row(previous)
                            owner["rows"] -= 1
                            owner["duplicates_removed"] += 1
                            self.duplicates_removed += 1
                        self.last_digest[digest] = self.spool.count
                    elif digest in self.seen_digests:
                        entry["duplicates_removed"] += 1
                        self.duplicates_removed += 1
                        continue
                    else:
                        self.seen_digests.add(digest)
            if self.sort_keys:
                if self.spool.count >= limits.max_sort_rows:
                    raise TabularMergeError(
                        "sort_limit_exceeded",
                        f"Sorting is available for merges of up to {limits.max_sort_rows:,} rows.",
                    )
                self.sort_values.append(tuple(
                    _sort_value(data[column] if column is not None else None, item)
                    for column, item in zip(sort_positions, options.sort_by)
                ))
            self.spool.append(prefix + data)
            entry["rows"] += 1

    def _entry_for_row(self, index):
        return self.row_entries[bisect.bisect_right(self.row_starts, index) - 1]

    # -- finish ------------------------------------------------------------

    def finish(self):
        report = self.report
        if self.mismatch:
            report["status"] = "failed"
            report["columns"] = list(self.provenance or ()) + list(self.data_display)
            for entry in report["sources"]:
                if entry["status"] in (SOURCE_STATUS_MERGED, SOURCE_STATUS_NOT_PROCESSED):
                    entry["status"] = SOURCE_STATUS_COMPATIBLE
                    entry["rows"] = 0
                    entry["blank_rows_skipped"] = 0
                    entry["short_rows_padded"] = 0
                    entry["duplicates_removed"] = 0
            raise TabularMergeError("schema_mismatch", _mismatch_message(report), report=report)
        if self.merged_tables == 0:
            report["status"] = "failed"
            raise TabularMergeError(
                "nothing_to_merge",
                "None of the selected files has a sheet with column headers and matching columns to merge.",
                report=report,
            )
        self._require_named_columns(self.dedupe_keys, "dedupe_column_not_found", "remove duplicates")
        self._require_named_columns(self.sort_keys, "sort_column_not_found", "sort")

        columns = tuple(self.provenance) + tuple(self.data_display)
        nullable = set()
        if len(self.provenance) > 1 and self.sheet_null:
            nullable.add(self.provenance[1])
        for name, seen in zip(self.data_display, self.seen_counts):
            if seen < self.merged_tables:
                nullable.add(name)
        keep_last = self.options.dedupe != DEDUPE_NONE and self.options.dedupe_keep == DEDUPE_KEEP_LAST
        row_count = self.spool.count - (self.duplicates_removed if keep_last else 0)
        order = self._sort_order() if self.sort_keys else None
        keep = self._keep_last_predicate() if keep_last else None

        report["status"] = "merged"
        report["columns"] = list(columns)
        report["nullable_columns"] = [name for name in columns if name in nullable]
        report["totals"] = {
            "sources": report["totals"]["sources"],
            "tables_merged": self.merged_tables,
            "rows": row_count,
            "columns": len(columns),
            "blank_rows_skipped": self.blank_rows,
            "duplicates_removed": self.duplicates_removed,
            "excluded": self.excluded_tables,
        }
        limitations = []
        if any(entry.get("format") != TABULAR_MERGE_FORMAT_CSV for entry in report["sources"]):
            limitations.extend(_EXCEL_LIMITATIONS)
        if self.excluded_tables:
            considered = sum(
                entry["status"] in (SOURCE_STATUS_MERGED, SOURCE_STATUS_EXCLUDED) for entry in report["sources"]
            )
            limitations.append(_EXCLUDED_LIMITATION.format(count=self.excluded_tables, total=considered))
        if any(name in nullable for name in self.data_display):
            limitations.append(_BLANK_COLUMNS_LIMITATION)
        if any(entry.get("ignored_columns") for entry in report["sources"] if entry["status"] == SOURCE_STATUS_MERGED):
            limitations.append(_IGNORED_COLUMNS_LIMITATION)
        report["limitations"] = limitations
        result = TabularMergeResult(
            columns, report, self.spool, nullable=nullable, row_count=row_count, order=order, keep=keep,
        )
        self.sort_values = []
        return result

    def _sort_order(self):
        values = self.sort_values
        order = list(range(self.spool.count))
        for position in reversed(range(len(self.options.sort_by))):
            descending = self.options.sort_by[position].descending
            order.sort(key=lambda index, position=position: values[index][position], reverse=descending)
        return order

    def _keep_last_predicate(self):
        offset = len(self.provenance)
        positions = (
            [self.data_index[key] for key in self.dedupe_keys]
            if self.options.dedupe == DEDUPE_KEY_COLUMNS else None
        )
        last = self.last_digest

        def keep(index, values):
            data = values[offset:]
            digest = _row_digest(data if positions is None else _values_at(data, positions))
            return digest is None or last.get(digest) == index

        return keep


def _validate_sources(sources, limits, *, minimum=2):
    if len(sources) < minimum:
        raise TabularMergeError(
            "too_few_sources", "Merging needs at least two files." if minimum > 1 else "Select a file to inspect.",
        )
    if len(sources) > limits.max_sources:
        raise TabularMergeError(
            "too_many_sources",
            f"{len(sources)} files were selected; at most {limits.max_sources} can be used here.",
        )
    seen = set()
    for source in sources:
        if not isinstance(source, TabularMergeSource):
            raise TabularMergeError("invalid_source", "Each merge source must be a TabularMergeSource.")
        if not isinstance(source.source_id, str) or not source.source_id.strip():
            raise TabularMergeError("invalid_source", "Each merge source needs an identifier.")
        if source.source_id in seen:
            raise TabularMergeError("duplicate_source", "The same file was selected more than once.")
        seen.add(source.source_id)
        if not callable(source.load_bytes):
            raise TabularMergeError("invalid_source", "Each merge source needs a byte loader.")
        if TABULAR_MERGE_EXTENSION_FORMATS.get(source.resolved_extension()) is None:
            raise TabularMergeError(
                "unsupported_format",
                f"{_safe_name(source)} isn't a CSV or Excel file (.csv, .xlsx, .xlsm, .xls).",
            )


def _new_report(options, source_count):
    return {
        "version": TABULAR_MERGE_REPORT_VERSION,
        "status": "running",
        "policy": options.policy_report(),
        "columns": [],
        "nullable_columns": [],
        "source_column": None,
        "sheet_column": None,
        "totals": {
            "sources": source_count, "tables_merged": 0, "rows": 0, "columns": 0,
            "blank_rows_skipped": 0, "duplicates_removed": 0, "excluded": 0,
        },
        "sources": [],
        "warnings": [],
        "limitations": [],
    }


def _new_source_entry(source, reader, sheet):
    entry = {
        "source_id": source.source_id,
        "file_name": _safe_name(source),
        "format": TABULAR_MERGE_EXTENSION_FORMATS.get(source.resolved_extension()),
        "sheet": sheet,
        "encoding": None,
        "delimiter": None,
        "header_row": None,
        "rows": 0,
        "blank_rows_skipped": 0,
        "short_rows_padded": 0,
        "duplicates_removed": 0,
        "status": SOURCE_STATUS_NOT_PROCESSED,
        "reason": None,
        "missing_columns": [],
        "extra_columns": [],
        "ignored_columns": [],
        "renamed_columns": [],
        "order_differs": False,
        "warnings": [],
    }
    entry.update(reader.describe())
    return entry


def _safe_name(source):
    name = str(getattr(source, "file_name", "") or "").strip()
    return name or "A selected file"


def _suffixed_name(base, suffix):
    ending = f" ({suffix})"
    name = base
    while len((name + ending).encode("utf-8")) > MAX_COLUMN_NAME_BYTES and name:
        name = name[:-1].rstrip()
    return name + ending


def _check_cancel(cancel_requested):
    if cancel_requested is not None and cancel_requested():
        raise TabularMergeCancelled()


def _load_source_bytes(source, limits):
    content = source.load_bytes()
    if isinstance(content, bytearray):
        content = bytes(content)
    if not isinstance(content, bytes):
        raise TabularMergeError("source_unavailable", f"{_safe_name(source)} could not be read.")
    if len(content) > limits.max_source_bytes:
        raise TabularMergeError(
            "source_too_large",
            f"{_safe_name(source)} is larger than the {limits.max_source_bytes // (1024 * 1024)} MB merge limit.",
        )
    return content


def _alias_index(options):
    index = {}
    for target, aliases in options.column_aliases:
        target_key = normalize_header_name(target)
        for alias in aliases:
            index[normalize_header_name(alias)] = (target, target_key)
    return index


def _header_names(raw_header, source, limits, entry):
    """Display and comparison names for one header row; trailing blanks are dropped."""
    cells = list(raw_header)
    while cells and not _display_header(cells[-1]):
        cells.pop()
    if not cells:
        raise TabularMergeError("empty_source", f"{_safe_name(source)} has no header row.")
    names = []
    seen = {}
    for position, value in enumerate(cells, start=1):
        display = _display_header(value)
        if not display:
            display = f"Column {position}"
            entry["warnings"].append({
                "code": "blank_header",
                "message": f'Column {position} has no header, so it is named "{display}".',
            })
        if len(display.encode("utf-8")) > limits.max_header_bytes:
            raise TabularMergeError(
                "header_too_long",
                f"A column header in {_safe_name(source)} is longer than {limits.max_header_bytes} bytes.",
            )
        normalized = display.casefold()
        if normalized in seen:
            raise TabularMergeError(
                "duplicate_columns",
                f'{_safe_name(source)} has more than one "{display}" column.',
            )
        seen[normalized] = position
        names.append((display, normalized))
    return names


def _apply_aliases(header, alias_index, source, entry):
    """Rename headers that are aliases of a target column; renaming must not duplicate one."""
    if not alias_index:
        return header
    renamed = []
    result = []
    seen = set()
    for display, key in header:
        target = alias_index.get(key)
        if target is not None:
            renamed.append({"from": display, "to": target[0]})
            display, key = target
        if key in seen:
            raise TabularMergeError(
                "duplicate_columns",
                f'{_safe_name(source)} has more than one column that becomes "{display}".',
            )
        seen.add(key)
        result.append((display, key))
    entry["renamed_columns"] = renamed
    return result


def _match_header(reference, header, schema_policy, entry):
    """Positions in ``header`` for each reference column, or None on a mismatch."""
    reference_keys = [normalized for _, normalized in reference]
    header_keys = [normalized for _, normalized in header]
    positions = {key: index for index, key in enumerate(header_keys)}
    missing = [display for display, key in reference if key not in positions]
    reference_set = set(reference_keys)
    extra = [display for display, key in header if key not in reference_set]
    entry["missing_columns"] = missing
    entry["extra_columns"] = extra
    if missing or extra:
        return None
    if reference_keys != header_keys:
        entry["order_differs"] = True
        if schema_policy == SCHEMA_POLICY_EXACT_ORDER:
            return None
    return [positions[key] for key in reference_keys]


def _mismatch_message(report):
    reference = report["sources"][0]["file_name"] if report["sources"] else "the first file"
    parts = []
    for entry in report["sources"]:
        if entry["status"] != SOURCE_STATUS_SCHEMA_MISMATCH:
            continue
        details = []
        if entry["reason"] == "no_matching_columns":
            details.append("none of the requested columns")
        if entry["missing_columns"] and entry["reason"] != "no_matching_columns":
            details.append(f"missing {', '.join(_quoted(entry['missing_columns']))}")
        if entry["extra_columns"]:
            details.append(f"extra {', '.join(_quoted(entry['extra_columns']))}")
        if not details and entry["order_differs"]:
            details.append("columns in a different order")
        label = entry["file_name"] if not entry.get("sheet") else f'{entry["file_name"]} sheet "{entry["sheet"]}"'
        parts.append(f"{label} ({'; '.join(details)})")
    return (
        f"These files don't have the same columns as {reference}: {'; '.join(parts)}. "
        "Nothing was merged."
    )


def _quoted(values, limit=8):
    values = list(values)
    shown = [f'"{value}"' for value in values[:limit]]
    if len(values) > limit:
        shown.append(f"and {len(values) - limit} more")
    return shown


def _cell_has_content(value):
    return bool(value) and bool(str(value).strip())


def _values_at(data, positions):
    return [data[position] if position is not None and position < len(data) else None for position in positions]


def _row_digest(values):
    """A collision-resistant key for duplicate detection, or None when every value is blank.

    Missing (null) and empty values compare equal, and trailing blanks are ignored, so a row
    read before a later file added a column matches the same row read after it.
    """
    normalized = ["" if value is None else value for value in values]
    while normalized and normalized[-1] == "":
        normalized.pop()
    if not any(_cell_has_content(value) for value in normalized):
        return None
    encoded = json.dumps(normalized, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    return hashlib.blake2b(encoded, digest_size=16).digest()


def _sort_value(value, sort):
    """A comparable key; typed values sort before text that isn't that type, blanks last."""
    ascending = not sort.descending
    typed_rank = 0 if ascending else 2
    blank_rank = 2 if ascending else 0
    if value is None or not str(value).strip():
        return (blank_rank, "")
    if sort.value_type == SORT_TYPE_NUMBER:
        number = _parse_number(value)
        if number is not None:
            return (typed_rank, number)
        return (1, value.casefold())
    if sort.value_type == SORT_TYPE_DATE:
        moment = _parse_date(value)
        if moment is not None:
            return (typed_rank, moment)
        return (1, value.casefold())
    return (typed_rank, value.casefold())


def _parse_number(value):
    text = value.strip()
    if text[:1] in _CURRENCY_PREFIXES:
        text = text[1:].strip()
    if not _NUMBER_TEXT.fullmatch(text):
        return None
    try:
        number = float(text.replace(",", ""))
    except ValueError:
        return None
    return number if math.isfinite(number) else None


def _parse_date(value):
    text = value.strip()
    if len(text) < 10 or not text[:4].isdigit():
        return None
    try:
        if len(text) == 10:
            day = date.fromisoformat(text)
            return datetime(day.year, day.month, day.day)
        moment = datetime.fromisoformat(text)
    except ValueError:
        return None
    if moment.tzinfo is not None:
        moment = moment.astimezone(timezone.utc).replace(tzinfo=None)
    return moment


# ---------------------------------------------------------------------------
# Inspection
# ---------------------------------------------------------------------------


def inspect_tabular_sources(
    sources: Sequence[TabularMergeSource],
    *,
    options: Optional[TabularInspectOptions] = None,
    limits: Optional[TabularMergeLimits] = None,
    cancel_requested: Optional[Callable[[], bool]] = None,
) -> dict:
    """Describe each file's sheets, headers, row counts and samples, and how their columns line up.

    A file that can't be read is reported with its problem rather than failing the inspection.
    """
    options = options or TabularInspectOptions()
    limits = limits or TabularMergeLimits()
    sources = list(sources or ())
    _validate_sources(sources, limits, minimum=1)
    report = {
        "version": TABULAR_INSPECTION_VERSION,
        "options": {
            "sheet": options.sheet, "sheets": options.sheets,
            "header_row": options.header_row, "sample_rows": options.sample_rows,
        },
        "sources": [],
        "compatibility": None,
        "samples_omitted": False,
        "limitations": [],
    }
    inspected = []
    for source in sources:
        _check_cancel(cancel_requested)
        entry = {
            "source_id": source.source_id,
            "file_name": _safe_name(source),
            "format": TABULAR_MERGE_EXTENSION_FORMATS.get(source.resolved_extension()),
            "status": "inspected",
            "problem": None,
            "encoding": None,
            "delimiter": None,
            "sheets": [],
            "tables": [],
        }
        report["sources"].append(entry)
        reader = None
        try:
            content = _load_source_bytes(source, limits)
            reader = _open_reader(source, content, limits)
            entry.update(reader.describe())
            entry["sheets"] = [{"name": name, "visible": state == "visible"} for name, state in reader.sheet_states()]
            for sheet, skip_reason in reader.plan(options, limits):
                _check_cancel(cancel_requested)
                table_entry = {"sheet": sheet, "status": "inspected", "reason": None, "problem": None}
                entry["tables"].append(table_entry)
                if skip_reason is not None:
                    table_entry.update(status=SOURCE_STATUS_SKIPPED, reason=skip_reason)
                    continue
                try:
                    table = reader.open_table(sheet, options.header_row)
                except TabularMergeCancelled:
                    raise
                except TabularMergeError as exc:
                    table_entry.update(status="problem", problem={"code": exc.code, "message": exc.message})
                    continue
                try:
                    header = _inspect_table(source, table, table_entry, options, limits, cancel_requested)
                except TabularMergeCancelled:
                    raise
                except TabularMergeError as exc:
                    table_entry.update(status="problem", problem={"code": exc.code, "message": exc.message})
                    continue
                finally:
                    table.close()
                inspected.append((entry, table_entry, header))
        except TabularMergeCancelled:
            raise
        except TabularMergeError as exc:
            entry["status"] = "problem"
            entry["problem"] = {"code": exc.code, "message": exc.message}
        finally:
            if reader is not None:
                reader.close()

    report["compatibility"] = _compatibility(inspected)
    limitations = ["Row counts leave out blank rows."]
    if options.sample_rows:
        limitations.append(
            f"Samples show up to {options.sample_rows} row(s) per sheet, and sample values longer than "
            f"{_INSPECT_CELL_CHARS} characters are shortened."
        )
    if any(entry["format"] != TABULAR_MERGE_FORMAT_CSV for entry in report["sources"]):
        limitations.append("Excel formulas are shown as their last calculated values.")
    report["limitations"] = limitations
    _fit_inspection_budget(report)
    return report


def _inspect_table(source, table, table_entry, options, limits, cancel_requested):
    probe = {"warnings": []}
    header = _header_names(table.header, source, limits, probe)
    width = len(header)
    samples = []
    count = 0
    capped = False
    extra_rows = 0
    first_data_row = None
    for row_number, values in table.rows():
        if count % _CANCEL_CHECK_INTERVAL == 0:
            _check_cancel(cancel_requested)
        cells = list(values)
        if not any(_cell_has_content(value) for value in cells):
            continue
        if first_data_row is None:
            first_data_row = (row_number, cells)
        if len(cells) > width and any(_cell_has_content(value) for value in cells[width:]):
            extra_rows += 1
        count += 1
        if len(samples) < options.sample_rows:
            row = [_clip_sample(value) for value in cells[:width]]
            samples.append(row + [""] * (width - len(row)))
        if count >= limits.max_total_rows:
            capped = True
            break
    warnings = list(probe["warnings"])
    header_values = sum(_cell_has_content(value) for value in table.header)
    if (
        options.header_row is None and header_values == 1 and first_data_row is not None
        and sum(_cell_has_content(value) for value in first_data_row[1]) > 1
    ):
        warnings.append({
            "code": "possible_title_row",
            "message": (
                f"Row {table.header_row} has only one value, so it may be a title rather than column "
                f"headers; row {first_data_row[0]} may hold the headers."
            ),
            "suggested_header_row": first_data_row[0],
        })
    if extra_rows:
        warnings.append({
            "code": "rows_with_extra_values",
            "message": f"{extra_rows} row(s) have more values than there are column headers.",
        })
    table_entry.update({
        "header_row": table.header_row,
        "columns": [display for display, _ in header],
        "row_count": count,
        "row_count_capped": capped,
        "sample_rows": samples,
        "warnings": warnings,
    })
    return header


def _clip_sample(value):
    text = "" if value is None else str(value)
    if len(text) > _INSPECT_CELL_CHARS:
        return text[:_INSPECT_CELL_CHARS - 1] + "\u2026"
    return text


def _compatibility(inspected):
    groups = {}
    order = None
    same_order = True
    presence = {}
    for entry, table_entry, header in inspected:
        keys = [key for _, key in header]
        group = groups.setdefault(
            tuple(sorted(keys)), {"columns": [display for display, _ in header], "tables": []},
        )
        group["tables"].append({
            "source_id": entry["source_id"], "file_name": entry["file_name"], "sheet": table_entry["sheet"],
        })
        if order is None:
            order = keys
        elif keys != order:
            same_order = False
        for display, key in header:
            item = presence.setdefault(key, {"name": display, "present_in": 0})
            item["present_in"] += 1
    near = {}
    for key, item in presence.items():
        squashed = _NOT_WORD.sub("", key)
        if squashed:
            near.setdefault(squashed, []).append(item["name"])
    similar = [{"names": names} for names in near.values() if len(names) > 1]
    identical = len(groups) == 1
    if not inspected:
        suggested = None
    elif identical:
        suggested = SCHEMA_POLICY_BY_NAME
    else:
        suggested = SCHEMA_POLICY_MAPPED if similar else SCHEMA_POLICY_UNION
    return {
        "tables": len(inspected),
        "identical_columns": bool(inspected) and identical,
        "same_order": bool(inspected) and identical and same_order,
        "groups": list(groups.values()),
        "columns": [{"name": item["name"], "present_in": item["present_in"]} for item in presence.values()],
        "similar_columns": similar,
        "suggested_policy": suggested,
    }


def _json_size(value):
    return len(json.dumps(value, ensure_ascii=False, separators=(",", ":")).encode("utf-8"))


def _fit_inspection_budget(report):
    if _json_size(report) <= _INSPECT_MAX_BYTES:
        return
    for entry in report["sources"]:
        for table in entry["tables"]:
            if table.get("sample_rows"):
                table["sample_rows"] = []
    report["samples_omitted"] = True
    report["limitations"].append("Sample rows were left out to keep the inspection small.")
    if _json_size(report) > _INSPECT_MAX_BYTES:
        raise TabularMergeError(
            "inspection_too_large", "The selected files have too many sheets or columns to inspect here.",
        )


# ---------------------------------------------------------------------------
# Readers
# ---------------------------------------------------------------------------


class _Table:
    """A header plus a row iterator over one sheet or CSV body."""

    format = ""
    sheet = None
    header_row = None

    @property
    def header(self):
        raise NotImplementedError

    def rows(self):
        raise NotImplementedError

    def close(self):
        return None


class _Reader:
    """One opened source file; workbooks can hold several tables."""

    format = ""

    def __init__(self, source):
        self.source = source

    def describe(self):
        return {}

    def sheet_states(self):
        return []

    def plan(self, options, limits):
        """``(sheet, skip_reason)`` for each table to read, in order."""
        states = self.sheet_states()
        if options.sheets == SHEETS_ALL:
            if not states:
                raise TabularMergeError("empty_source", f"{_safe_name(self.source)} has no worksheets.")
            if len(states) > limits.max_sheets_per_source:
                raise TabularMergeError(
                    "too_many_sheets",
                    f"{_safe_name(self.source)} has more than {limits.max_sheets_per_source} sheets.",
                )
            return [(name, None if state == "visible" else "hidden_sheet") for name, state in states]
        return [(_select_sheet(states, options.sheet, self.source), None)]

    def open_table(self, sheet, header_row):
        raise NotImplementedError

    def close(self):
        return None


def _open_reader(source, content, limits):
    table_format = TABULAR_MERGE_EXTENSION_FORMATS[source.resolved_extension()]
    if table_format == TABULAR_MERGE_FORMAT_CSV:
        return _CsvReader(source, content)
    if table_format == TABULAR_MERGE_FORMAT_XLSX:
        return _XlsxReader(source, content, limits)
    return _XlsReader(source, content)


class _CsvReader(_Reader):
    format = TABULAR_MERGE_FORMAT_CSV

    def __init__(self, source, content):
        super().__init__(source)
        self._text, self._encoding = _decode_csv(content, source)
        self._delimiter = _detect_delimiter(self._text)

    def describe(self):
        return {"encoding": self._encoding, "delimiter": self._delimiter}

    def plan(self, options, limits):
        return [(None, None)]

    def open_table(self, sheet, header_row):
        return _CsvTable(self.source, self._text, self._delimiter, header_row)


class _CsvTable(_Table):
    format = TABULAR_MERGE_FORMAT_CSV

    def __init__(self, source, text, delimiter, header_row):
        self._source = source
        self._reader = csv.reader(
            io.StringIO(text, newline=""), delimiter=delimiter, quotechar='"', doublequote=True, strict=False,
        )
        self._header = None
        # Rows are counted as records, as a spreadsheet shows them, not as physical lines.
        record = 0
        try:
            for values in self._reader:
                record += 1
                if header_row is not None and record < header_row:
                    continue
                if any(_cell_has_content(value) for value in values):
                    self._header = values
                    self.header_row = record
                    self._record = record
                    break
                if header_row is not None:
                    raise TabularMergeError(
                        "empty_source", f"Row {header_row} in {_safe_name(source)} is blank, so it has no column headers.",
                    )
        except csv.Error as exc:
            raise TabularMergeError(
                "unreadable_csv", f"{_safe_name(source)} couldn't be read as CSV near line {self._reader.line_num}.",
            ) from exc
        if self._header is None:
            raise TabularMergeError(
                "empty_source",
                f"{_safe_name(source)} has no header row." if header_row is None
                else f"{_safe_name(source)} has fewer than {header_row} rows.",
            )

    @property
    def header(self):
        return self._header

    def rows(self):
        reader = self._reader
        try:
            for values in reader:
                self._record += 1
                yield self._record, values
        except csv.Error as exc:
            raise TabularMergeError(
                "unreadable_csv", f"{_safe_name(self._source)} couldn't be read as CSV near line {reader.line_num}.",
            ) from exc


def _decode_csv(content, source):
    for bom, encoding in (
        (codecs.BOM_UTF32_LE, "utf-32"), (codecs.BOM_UTF32_BE, "utf-32"),
        (codecs.BOM_UTF8, "utf-8-sig"),
        (codecs.BOM_UTF16_LE, "utf-16"), (codecs.BOM_UTF16_BE, "utf-16"),
    ):
        if content.startswith(bom):
            try:
                return content.decode(encoding), encoding
            except UnicodeDecodeError as exc:
                raise TabularMergeError(
                    "unreadable_csv", f"{_safe_name(source)} has text that doesn't match its encoding.",
                ) from exc
    if b"\x00" in content[:_CSV_SNIFF_BYTES]:
        raise TabularMergeError("unreadable_csv", f"{_safe_name(source)} doesn't look like a text CSV file.")
    for encoding in ("utf-8", "cp1252"):
        try:
            return content.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1"), "latin-1"


def _detect_delimiter(text):
    """The most frequent supported delimiter outside quotes on the first nonblank line."""
    counts = {delimiter: 0 for delimiter in _CSV_DELIMITERS}
    in_quotes = False
    seen_content = False
    for character in text[:_CSV_SNIFF_BYTES]:
        if character == '"':
            in_quotes = not in_quotes
            seen_content = True
            continue
        if in_quotes:
            continue
        if character in "\r\n":
            if seen_content:
                break
            continue
        if character in counts:
            counts[character] += 1
        if not character.isspace():
            seen_content = True
    best = max(_CSV_DELIMITERS, key=lambda delimiter: (counts[delimiter], -_CSV_DELIMITERS.index(delimiter)))
    return best if counts[best] else ","


class _XlsxReader(_Reader):
    format = TABULAR_MERGE_FORMAT_XLSX

    def __init__(self, source, content, limits):
        super().__init__(source)
        _guard_workbook_archive(content, source, limits)
        # openpyxl is imported only when a workbook is merged; it is a large binary dependency.
        from openpyxl import load_workbook

        try:
            self._workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=True, keep_links=False,
            )
        except Exception as exc:
            raise TabularMergeError(
                "unreadable_workbook", f"{_safe_name(source)} couldn't be opened as an Excel workbook.",
            ) from exc

    def sheet_states(self):
        try:
            return [
                (worksheet.title, getattr(worksheet, "sheet_state", "visible"))
                for worksheet in self._workbook.worksheets
            ]
        except Exception as exc:
            raise TabularMergeError(
                "unreadable_workbook", f"{_safe_name(self.source)} couldn't be read as an Excel workbook.",
            ) from exc

    def open_table(self, sheet, header_row):
        try:
            return _XlsxTable(self.source, self._workbook[sheet], sheet, header_row)
        except TabularMergeError:
            raise
        except Exception as exc:
            raise TabularMergeError(
                "unreadable_workbook", f"{_safe_name(self.source)} couldn't be read as an Excel workbook.",
            ) from exc

    def close(self):
        workbook = getattr(self, "_workbook", None)
        if workbook is not None:
            self._workbook = None
            try:
                workbook.close()
            except Exception:
                pass


class _XlsxTable(_Table):
    format = TABULAR_MERGE_FORMAT_XLSX

    def __init__(self, source, worksheet, sheet, header_row):
        self._source = source
        self.sheet = sheet
        worksheet.reset_dimensions()
        self._iterator = worksheet.iter_rows(values_only=True)
        self._row_number = 0
        self._header = None
        for values in self._iterator:
            self._row_number += 1
            if header_row is not None and self._row_number < header_row:
                continue
            cells = [_excel_text(value) for value in values]
            if any(_cell_has_content(value) for value in cells):
                self._header = cells
                self.header_row = self._row_number
                break
            if header_row is not None:
                raise TabularMergeError(
                    "empty_source",
                    f'Row {header_row} of sheet "{sheet}" in {_safe_name(source)} is blank, so it has no column headers.',
                )
        if self._header is None:
            raise TabularMergeError(
                "empty_source", f'Sheet "{sheet}" in {_safe_name(source)} has no header row.',
            )

    @property
    def header(self):
        return self._header

    def rows(self):
        try:
            for values in self._iterator:
                self._row_number += 1
                yield self._row_number, [_excel_text(value) for value in values]
        except TabularMergeError:
            raise
        except Exception as exc:
            raise TabularMergeError(
                "unreadable_workbook", f"{_safe_name(self._source)} couldn't be read near row {self._row_number}.",
            ) from exc


def _guard_workbook_archive(content, source, limits):
    if content[:8] == _OLE_MAGIC:
        raise TabularMergeError(
            "encrypted_workbook",
            f"{_safe_name(source)} is password-protected or encrypted, so it can't be merged.",
        )
    if content[:4] != _ZIP_MAGIC:
        raise TabularMergeError("unreadable_workbook", f"{_safe_name(source)} isn't a valid Excel workbook.")
    try:
        archive = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile as exc:
        raise TabularMergeError(
            "unreadable_workbook", f"{_safe_name(source)} isn't a valid Excel workbook.",
        ) from exc
    with archive:
        entries = archive.infolist()
        if len(entries) > limits.max_workbook_entries:
            raise TabularMergeError("workbook_too_large", f"{_safe_name(source)} has too many internal parts to merge.")
        if sum(entry.file_size for entry in entries) > limits.max_workbook_uncompressed_bytes:
            raise TabularMergeError(
                "workbook_too_large", f"{_safe_name(source)} is too large to merge once uncompressed.",
            )
        try:
            unsafe = first_unsafe_xml_part(archive, entries)
        except UnreadablePackageError as exc:
            raise TabularMergeError(
                "unreadable_workbook", f"{_safe_name(source)} isn't a valid Excel workbook.",
            ) from exc
    if unsafe is not None:
        raise TabularMergeError(
            "unreadable_workbook",
            f"{_safe_name(source)} isn't a valid Excel workbook; it has XML that Excel files can't contain.",
        )


class _XlsReader(_Reader):
    format = TABULAR_MERGE_FORMAT_XLS

    def __init__(self, source, content):
        super().__init__(source)
        if content[:8] != _OLE_MAGIC:
            raise TabularMergeError(
                "unreadable_workbook", f"{_safe_name(source)} isn't a valid Excel 97-2003 (.xls) workbook.",
            )
        # xlrd is imported only when a legacy workbook is merged.
        import xlrd

        self._xlrd = xlrd
        try:
            self._book = xlrd.open_workbook(file_contents=content, on_demand=True)
        except Exception as exc:
            code = "encrypted_workbook" if "encrypt" in str(exc).lower() else "unreadable_workbook"
            message = (
                f"{_safe_name(source)} is password-protected or encrypted, so it can't be merged."
                if code == "encrypted_workbook"
                else f"{_safe_name(source)} couldn't be opened as an Excel workbook."
            )
            raise TabularMergeError(code, message) from exc

    def sheet_states(self):
        try:
            names = self._book.sheet_names()
            visibility = getattr(self._book, "_sheet_visibility", None)
            if not isinstance(visibility, list) or len(visibility) != len(names):
                visibility = [
                    getattr(self._book.sheet_by_index(index), "visibility", 0) for index in range(len(names))
                ]
            return [(name, "visible" if state == 0 else "hidden") for name, state in zip(names, visibility)]
        except Exception as exc:
            raise TabularMergeError(
                "unreadable_workbook", f"{_safe_name(self.source)} couldn't be read as an Excel workbook.",
            ) from exc

    def open_table(self, sheet, header_row):
        try:
            worksheet = self._book.sheet_by_name(sheet)
            return _XlsTable(self.source, self._xlrd, worksheet, self._book.datemode, sheet, header_row)
        except TabularMergeError:
            raise
        except Exception as exc:
            raise TabularMergeError(
                "unreadable_workbook", f"{_safe_name(self.source)} couldn't be read as an Excel workbook.",
            ) from exc

    def close(self):
        book = getattr(self, "_book", None)
        if book is not None:
            self._book = None
            try:
                book.release_resources()
            except Exception:
                pass


class _XlsTable(_Table):
    format = TABULAR_MERGE_FORMAT_XLS

    def __init__(self, source, xlrd, worksheet, datemode, sheet, header_row):
        self._source = source
        self._xlrd = xlrd
        self._worksheet = worksheet
        self._datemode = datemode
        self.sheet = sheet
        self._next_row = 0
        self._header = None
        while self._next_row < worksheet.nrows:
            row_number = self._next_row + 1
            cells = self._row_values(self._next_row)
            self._next_row += 1
            if header_row is not None and row_number < header_row:
                continue
            if any(_cell_has_content(value) for value in cells):
                self._header = cells
                self.header_row = row_number
                break
            if header_row is not None:
                raise TabularMergeError(
                    "empty_source",
                    f'Row {header_row} of sheet "{sheet}" in {_safe_name(source)} is blank, so it has no column headers.',
                )
        if self._header is None:
            raise TabularMergeError(
                "empty_source", f'Sheet "{sheet}" in {_safe_name(source)} has no header row.',
            )

    @property
    def header(self):
        return self._header

    def rows(self):
        while self._next_row < self._worksheet.nrows:
            row_number = self._next_row + 1
            cells = self._row_values(self._next_row)
            self._next_row += 1
            yield row_number, cells

    def _row_values(self, row_index):
        xlrd = self._xlrd
        values = []
        for cell in self._worksheet.row(row_index):
            if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                values.append("")
            elif cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    values.append(_excel_text(xlrd.xldate_as_datetime(cell.value, self._datemode)))
                except Exception:
                    values.append(_excel_text(cell.value))
            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                values.append("TRUE" if cell.value else "FALSE")
            elif cell.ctype == xlrd.XL_CELL_ERROR:
                values.append(xlrd.error_text_from_code.get(cell.value, "#ERROR"))
            else:
                values.append(_excel_text(cell.value))
        return values


def _select_sheet(sheets: List[Tuple[str, str]], requested, source):
    names = [name for name, _ in sheets]
    if not names:
        raise TabularMergeError("empty_source", f"{_safe_name(source)} has no worksheets.")
    if requested is not None:
        if requested in names:
            return requested
        folded = requested.casefold()
        for name in names:
            if name.casefold() == folded:
                return name
        raise TabularMergeError(
            "sheet_not_found",
            f'{_safe_name(source)} has no sheet named "{requested}". Its sheets are: '
            f"{', '.join(_quoted(names))}.",
            details={"available_sheets": names},
        )
    for name, state in sheets:
        if state == "visible":
            return name
    return names[0]


def _excel_text(value):
    """Canonical text for one Excel cell value."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer() and abs(value) < _MAX_EXACT_INTEGER_FLOAT:
            return str(int(value))
        return repr(value)
    if isinstance(value, datetime):
        if value.tzinfo is None and value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(timespec="seconds" if value.microsecond == 0 else "microseconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat(timespec="seconds" if value.microsecond == 0 else "microseconds")
    if isinstance(value, timedelta):
        total_seconds = int(value.total_seconds())
        sign = "-" if total_seconds < 0 else ""
        hours, remainder = divmod(abs(total_seconds), 3600)
        minutes, seconds = divmod(remainder, 60)
        return f"{sign}{hours}:{minutes:02d}:{seconds:02d}"
    return str(value)


__all__ = [
    "DEDUPE_EXACT_ROWS",
    "DEDUPE_KEEP_FIRST",
    "DEDUPE_KEEP_LAST",
    "DEDUPE_KEY_COLUMNS",
    "DEDUPE_NONE",
    "DEFAULT_INSPECT_SAMPLE_ROWS",
    "DEFAULT_SHEET_COLUMN_NAME",
    "DEFAULT_SOURCE_COLUMN_NAME",
    "MAX_INSPECT_SAMPLE_ROWS",
    "ON_INCOMPATIBLE_EXCLUDE",
    "ON_INCOMPATIBLE_FAIL",
    "SCHEMA_POLICY_BY_NAME",
    "SCHEMA_POLICY_EXACT_ORDER",
    "SCHEMA_POLICY_MAPPED",
    "SCHEMA_POLICY_UNION",
    "SHEETS_ALL",
    "SHEETS_FIRST",
    "SORT_TYPE_DATE",
    "SORT_TYPE_NUMBER",
    "SORT_TYPE_TEXT",
    "TABULAR_COLUMN_MAPPING_PROFILE",
    "TABULAR_INSPECTION_VERSION",
    "TABULAR_MERGE_DEDUPE_KEEP",
    "TABULAR_MERGE_DEDUPE_MODES",
    "TABULAR_MERGE_INCOMPATIBLE_MODES",
    "TABULAR_MERGE_REPORT_VERSION",
    "TABULAR_MERGE_SCHEMA_POLICIES",
    "TABULAR_MERGE_SHEET_MODES",
    "TABULAR_MERGE_SORT_TYPES",
    "TABULAR_MERGE_SUPPORTED_EXTENSIONS",
    "TabularInspectOptions",
    "TabularMergeCancelled",
    "TabularMergeError",
    "TabularMergeLimits",
    "TabularMergeOptions",
    "TabularMergeResult",
    "TabularMergeSort",
    "TabularMergeSource",
    "inspect_tabular_sources",
    "iter_csv_rows",
    "merge_tabular_sources",
    "normalize_header_name",
    "squash_header_name",
    "tabular_column_mapping_schema",
    "tabular_inspect_options_from_arguments",
    "tabular_mapping_from_profile",
    "tabular_merge_format_for_file_name",
    "tabular_merge_options_from_arguments",
]
